"""Classify the three user-indicated ETI ED2 EANs using official ETI cards."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "ETI_modulowka_IMPORT_BaseLinker_atrybuty_nazwy_mala_litera.xlsx"
OUTPUT = HERE / "ETI_ED2_trzy_EAN_korekta_kategorii_2026-09-25.xlsx"
PRODUCTS = {
    "3838895411048": (
        "004671271/ETI", "ED2 125/3", 125,
        "https://www.etigroup.eu/products-and-solutions/moulded-case-circuit-breakers-and-disconnectors-etibreak/switch-disconnectors-125-1600a/004671271-ed2-125-3",
    ),
    "3838895411055": (
        "004671272/ETI", "ED2 160/3", 160,
        "https://www.etipolam.com.pl/katalog-elektroniczny/wylaczniki-i-rozlaczniki-kompaktowe/rozlaczniki-ed2-125-1600a/004671272-ed2-160-3",
    ),
    "3838895411062": (
        "004671273/ETI", "ED2 250/3", 250,
        "https://www.etigroup.eu/products-and-solutions/moulded-case-circuit-breakers-and-disconnectors-etibreak/switch-disconnectors-125-1600a/004671273-ed2-250-3",
    ),
}
RECOMMENDED = "Wyłączniki i rozłączniki kompaktowe > Rozłączniki kompaktowe"


def format_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 34
    for number, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + number)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row[0].row].height = 46


def main() -> None:
    source = load_workbook(SOURCE, read_only=True, data_only=True)["Import"]
    rows = {str(row[3]): row for row in list(source.values)[1:]}
    assert len(rows) == 127
    book = Workbook()
    summary = book.active
    summary.title = "Wniosek"
    summary.append(["Temat", "Ustalenie", "Źródło"])
    summary.append(["Rodzina", "Tak: trzy wskazane produkty są aparatami z kompaktowej rodziny ETIBREAK ED2.", "https://www.etipolam.com.pl/katalog-elektroniczny/wylaczniki-i-rozlaczniki-kompaktowe/rozlaczniki-ed2-125-1600a/004671272-ed2-160-3"])
    summary.append(["Funkcja", "Rozłączniki izolacyjne, nie automatyczne wyłączniki kompaktowe MCCB; ETI podaje IEC 60947-3 i funkcję switch disconnector.", "https://www.etigroup.eu/products-and-solutions/moulded-case-circuit-breakers-and-disconnectors-etibreak/switch-disconnectors-125-1600a/004671273-ed2-250-3"])
    summary.append(["Rekomendowana kategoria", RECOMMENDED, "Wniosek z klasyfikacji producenta; nazwa kategorii sklepowej do utworzenia, jeżeli nie istnieje."])
    summary.append(["Kategoria awaryjna w obecnym drzewie", "Aparatura modułowa > Rozłączniki izolacyjne — technicznie poprawna, ale mniej precyzyjna.", SOURCE.name])
    summary.append(["Nie przypisywać", "Wyłączniki kompaktowe (MCCB) — sugerowałoby zabezpieczenie samoczynne przed zwarciem/przeciążeniem.", "https://www.etigroup.eu/products-and-solutions/moulded-case-circuit-breakers-and-disconnectors-etibreak/circuit-breakers-eb2-series-20-1600a/004671352-eb2-250-3le-125a-3p"])
    summary.append(["BaseLinker", "Nie zmieniano kategorii ani żadnych danych w BaseLinkerze.", ""])
    format_sheet(summary, [38, 115, 110])
    for number in (2, 3, 6):
        cell = summary.cell(number, 3)
        cell.hyperlink = cell.value
        cell.font = Font(name="Arial", color="0563C1", underline="single")

    details = book.create_sheet("Trzy EAN")
    details.append(["EAN z listy", "SKU", "Model", "Prąd znamionowy [A]", "Funkcja wg ETI", "Norma wg ETI", "Obecna kategoria", "Rekomendowana kategoria", "Karta ETI"])
    for ean, (sku, model, current, url) in PRODUCTS.items():
        row = rows[ean]
        attributes = json.loads(row[7])
        assert row[2] == sku and attributes["Symbol"] == model
        assert attributes["Typ produktu"] == "Rozłącznik izolacyjny"
        assert attributes["Prąd znamionowy [A]"] == str(current)
        details.append([
            ean, sku, model, current, "Rozłącznik izolacyjny (switch disconnector)",
            "IEC 60947-3", row[5], RECOMMENDED, url,
        ])
    format_sheet(details, [22, 23, 19, 19, 42, 22, 52, 64, 110])
    for row in range(2, details.max_row + 1):
        cell = details.cell(row, 9)
        cell.hyperlink = cell.value
        cell.font = Font(name="Arial", color="0563C1", underline="single")

    book.save(OUTPUT)
    verify = load_workbook(OUTPUT, read_only=True, data_only=True)
    assert [(s.title, s.max_row) for s in verify.worksheets] == [("Wniosek", 7), ("Trzy EAN", 4)]
    print(f"products=3; output={OUTPUT}")


if __name__ == "__main__":
    main()
