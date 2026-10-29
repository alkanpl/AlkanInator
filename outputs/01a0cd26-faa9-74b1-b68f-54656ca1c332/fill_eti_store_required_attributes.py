"""Fill manufacturer-confirmed missing ETI attributes in a copy of the store export."""

from __future__ import annotations

from collections import Counter
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
SOURCE = Path(r"C:\Users\Handlowiec\Downloads\eti-modulowka-25.09.xlsx")
MISSING_REPORT = HERE / "ETI_eksport_sklepu_audyt_atrybutow_i_tagi_wariantow_2026-09-25.xlsx"
OUTPUT = HERE / "eti-modulowka-25.09_UZUPELNIENIA_ZIELONE.xlsx"
REPORT = HERE / "ETI_raport_uzupelnien_atrybutow_2026-09-25.xlsx"
PREFIX = "Atrybut Produktu: "
GREEN = PatternFill(fill_type="solid", fgColor="C6EFCE")

ETIMAT_1P = "https://www.etigroup.eu/products-and-solutions/modular-current-and-residual-protective-devices-asti/miniature-circuit-breakers-mcb-and-accessories-etimat/miniature-circuit-breakers-mcb/002111516-etimat-6-1p-b16"
ETIMAT_3P = "https://www.etigroup.eu/products-and-solutions/modular-current-and-residual-protective-devices-asti/miniature-circuit-breakers-mcb-and-accessories-etimat/miniature-circuit-breakers-mcb/002115516-etimat-6-3p-b16"
ETIMAT_3P_D = "https://www.etigroup.eu/products-and-solutions/modular-current-and-residual-protective-devices-asti/miniature-circuit-breakers-mcb-and-accessories-etimat/miniature-circuit-breakers-mcb/002164521-etimat-6-3p-d50"
ETIMAT_10_B6 = "https://www.eti.ua/produktsiia/modulni-pristroi-zahistu-etimat-efi-kzs-ms-mpe-msp/modulni-avtomaticni-vimikaci-ta-aksesuari/modulni-avtomaticni-vimikaci/002125712-etimat-10-3p-b6"
ETIMAT_SM = "https://www.etipolam.com.pl/katalog-elektroniczny/aparaty-modulowe-zabezpieczajace/wylaczniki-nadpradowe-selektywne-i-akcesoria/wylaczniki-nadpradowe-selektywne/002129226-etimat-sm-3p-e63"
KZS_4M = "https://www.etigroup.eu/products-and-solutions/modular-current-and-residual-protective-devices-asti/residual-and-overcurrent-protection-rcbo-and-accessories-kzs/residual-and-overcurrent-protection-rcbo/002174002-kzs-4m-3p-n-ac-b10-0-03"

FACTS: dict[tuple[str, str], tuple[object, str, str]] = {}


def add(sku: str, attr: str, value: object, source: str, note: str) -> None:
    key = (sku, attr)
    if key in FACTS:
        raise ValueError(f"Duplicate fact: {key}")
    FACTS[key] = (value, source, note)


def make_facts() -> None:
    one_p = [
        "002111516", "002111514", "002111517", "002111512", "002111518",
        "002141514", "002141516", "002141517", "002141518",
    ]
    three_p = [
        "002115516", "002115517", "002115514", "002115518", "002125712",
        "002145516", "002145517", "002145518", "002145519",
    ]
    for code in one_p:
        add(f"{code}/ETI", "Liczba Modułów", 1, ETIMAT_1P,
            "Reguła rodziny ETIMAT 6 1P; kod i liczba biegunów zgodne z eksportem. Dla 002111516 karta dokładnego SKU.")
    for code in three_p:
        source = ETIMAT_10_B6 if code == "002125712" else ETIMAT_3P
        note = ("Karta dokładnego SKU ETIMAT 10 3P; w eksporcie błędnie nazwany ETIMAT 6."
                if code == "002125712" else
                "Reguła rodziny ETIMAT 6 3P; kod i liczba biegunów zgodne z eksportem.")
        add(f"{code}/ETI", "Liczba Modułów", 3, source, note)
    add("002111516/ETI", "Układ Biegunów", "1P", ETIMAT_1P, "Karta dokładnego SKU: 1 biegun.")

    contactors = {
        "002461210": (1, "002461210-r-20-20-230v"),
        "002461211": (1, "002461211-r-20-20-24v"),
        "002461220": (1, "002461220-r-20-11-230v"),
        "002461230": (1, "002461230-r-20-02-230v"),
        "002462310": (2, "002462310-r-25-40-230v"),
        "002463410": (3, "002463410-r-40-40-230v"),
        "002463450": (3, "002463450-r-63-40-230v"),
    }
    contactor_base = "https://www.etigroup.eu/products-and-solutions/contactors-eticon/modular-contactors-for-installation-into-distribution-boards-and-accesories/modular-contactors-for-instalation-into-distribution-boards/"
    for code, (modules, slug) in contactors.items():
        add(f"{code}/ETI", "Liczba Modułów", modules, contactor_base + slug,
            "Karta dokładnego SKU ETI: szerokość w liczbie modułów.")

    for attr, fact in {
        "Liczba Modułów": 4,
        "Znamionowy Prąd Różnicowy IΔn [mA]": 30,
        "Typ Wyłącznika": "AC",
        "Znamionowa Zwarciowa Zdolność Wyłączania Icn [kA]": 6,
    }.items():
        add("002174002/ETI", attr, fact, KZS_4M,
            "Karta dokładnego SKU ETI; dla Icn użyto zdolności wyłączania wg EN 61009, nie Inc.")

    rails = {
        "002921026": (54, 1000, "002921026-iz12-1f-54"),
        "002921024": (54, 1000, "002921024-iz12-3f-54"),
        "002921020": (12, 210, "002921020-iz12-3f-12"),
        "002921022": (18, 310, "002921022-iz12-3f-18"),
        "002921018": (12, 210, "002921018-iz12-1f-12"),
    }
    rail_base = "https://www.etigroup.eu/products-and-solutions/enclosures-distribution-boards-equipment-etibox/section-rails/"
    for code, (modules, length_mm, slug) in rails.items():
        source = rail_base + slug
        add(f"{code}/ETI", "Liczba Modułów", modules, source,
            "Karta dokładnego SKU ETI: długość szyny w modułach (1×N).")
        add(f"{code}/ETI", "Długość", f"{length_mm} mm", source,
            "Karta dokładnego SKU ETI: wymiar Length z klasyfikacji ETIM, nie szerokość netto opakowania.")


def style_report(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 35
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    make_facts()
    original = load_workbook(SOURCE)
    if original.sheetnames != ["Worksheet"]:
        raise ValueError("Unexpected input workbook layout")
    sheet = original.active
    headers = [cell.value for cell in sheet[1]]
    columns = {name: headers.index(name) + 1 for name in set(headers) if name is not None}
    sku_col = columns["SKU"]
    sku_rows = {str(sheet.cell(row, sku_col).value): row for row in range(2, sheet.max_row + 1)}
    if len(sku_rows) != 136:
        raise ValueError("Unexpected product count or duplicate SKU")

    missing_book = load_workbook(MISSING_REPORT, read_only=True, data_only=True)
    missing = {(str(row[1]), str(row[5])) for row in list(missing_book["Braki atrybutów"].values)[1:]}
    if len(missing) != 41 or len(FACTS) != 40 or not set(FACTS).issubset(missing):
        raise ValueError("Facts do not match the previously audited missing cells")

    log = []
    for (sku, attr), (fact, source, note) in FACTS.items():
        row = sku_rows[sku]
        cell = sheet.cell(row, columns[PREFIX + attr])
        if cell.value not in (None, ""):
            raise ValueError(f"Refusing to overwrite existing value: {sku}, {attr}")
        if attr == "Liczba Modułów" and sku.startswith(("002111", "002141", "002115", "002145")):
            poles = sheet.cell(row, columns[PREFIX + "Liczba Biegunów"]).value
            expected = 1 if sku.startswith(("002111", "002141")) else 3
            if poles != expected or fact != expected:
                raise ValueError(f"MCB family/pole mismatch: {sku}")
        if sku == "002125712/ETI" and attr == "Liczba Modułów" and fact != 3:
            raise ValueError("ETIMAT 10 width mismatch")
        cell.value = fact
        cell.fill = copy(GREEN)
        cell.comment = Comment(f"Źródło: {source}\nPewność: wysoka\nStatus: OK\n{note}", "ETI / Codex")
        log.append([sku, sheet.cell(row, columns["Title"]).value, attr, fact, cell.coordinate,
                    source, "wysoka — karta SKU lub potwierdzona rodzina", "OK", note])

    if len(log) != 40:
        raise ValueError("Unexpected fill count")
    original.save(OUTPUT)

    report = Workbook()
    summary = report.active
    summary.title = "Podsumowanie"
    summary.append(["Miara", "Wynik"])
    for line in [
        ("Produkty w pliku", 136), ("Braki przed", len(missing)), ("Uzupełniono zielonym", len(log)),
        ("Pozostałe braki", len(missing - set(FACTS))), ("Plik źródłowy", str(SOURCE)),
        ("Plik wynikowy", str(OUTPUT)), ("Zmiany online", "Nie wykonano"),
    ]:
        summary.append(line)
    for attr, count in sorted(Counter(row[2] for row in log).items()):
        summary.append((f"Uzupełniono: {attr}", count))
    style_report(summary, [57, 110])

    filled = report.create_sheet("Uzupełnione")
    filled.append(["SKU", "Produkt", "Atrybut", "Wartość", "Komórka", "Źródło", "Pewność", "Status", "Uzasadnienie"])
    for row in sorted(log):
        filled.append(row)
    style_report(filled, [20, 75, 65, 18, 14, 110, 50, 16, 105])

    remaining = report.create_sheet("Pozostałe braki")
    remaining.append(["SKU", "Atrybut", "Powód niewypełnienia", "Dane potwierdzone", "Źródło", "Status"])
    remaining.append(["002129226/ETI", "Liczba Modułów",
                      "Karta ETI podaje szerokość 81 mm, lecz nie podaje liczby szerokości modułowych. Nie przeliczano wymiaru na moduły.",
                      "Szerokość fizyczna 81 mm", ETIMAT_SM, "DO_SPRAWDZENIA"])
    style_report(remaining, [20, 35, 105, 36, 110, 23])

    alerts = report.create_sheet("Istniejące dane do korekty")
    alerts.append(["SKU", "Pole w źródle", "Wartość w sklepie", "Wartość ETI", "Źródło", "Uwaga"])
    alerts.append(["002125712/ETI", "Tytuł / seria", "ETIMAT 6", "ETIMAT 10", ETIMAT_10_B6,
                   "Istniejącej treści nie zmieniano; prośba dotyczyła brakujących komórek."])
    alerts.append(["002125712/ETI", "Icn [kA]", "6", "10", ETIMAT_10_B6,
                   "Istniejącej wartości nie nadpisano."])
    alerts.append(["002129226/ETI", "Zdolność zwarciowa [kA]", "10", "25", ETIMAT_SM,
                   "ETI podaje zdolność łączeniową 25 kA; w eksporcie pole Icn ma 10 kA. Wymaga rozstrzygnięcia właściwej definicji pola."])
    style_report(alerts, [20, 43, 28, 26, 110, 108])

    report.save(REPORT)
    print(f"filled={len(log)} remaining={len(missing - set(FACTS))} source_products={len(sku_rows)}")
    print(OUTPUT)
    print(REPORT)


if __name__ == "__main__":
    main()
