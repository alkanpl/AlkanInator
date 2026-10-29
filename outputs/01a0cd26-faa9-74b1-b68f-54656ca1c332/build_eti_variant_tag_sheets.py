"""One worksheet per ETI variant tag, based on the enriched store export."""

from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "eti-modulowka-25.09_UZUPELNIENIA_ZIELONE.xlsx"
PREVIOUS = HERE / "ETI_eksport_sklepu_audyt_atrybutow_i_tagi_wariantow_2026-09-25.xlsx"
OUTPUT = HERE / "ETI_tagi_wariantow_osobne_arkusze_2026-09-25.xlsx"
P = "Atrybut Produktu: "

GROUPS = [
    ("01_ETIMAT6_MCB", "Wariant - ETI ETIMAT 6 wyłącznik nadprądowy",
     ["Charakterystyka Wyzwalania", "Prąd Znamionowy [A]", "Układ Biegunów"],
     ["Znamionowa Zwarciowa Zdolność Wyłączania Icn [kA]", "Liczba Modułów"],
     "Wyłącznik nadprądowy ETIMAT 6, producent ETI; nie łączono z ETIMAT 10 ani ETIMAT SM."),
    ("02_KZS2M_RCBO", "Wariant - ETI KZS-2M wyłącznik różnicowo-nadprądowy",
     ["Charakterystyka Wyzwalania", "Prąd Znamionowy [A]", "Typ Wyłącznika"],
     ["Układ Biegunów", "Znamionowy Prąd Różnicowy IΔn [mA]", "Znamionowa Zwarciowa Zdolność Wyłączania Icn [kA]", "Liczba Modułów"],
     "Wyłącznik RCBO KZS-2M, producent ETI; KZS-4M pozostaje osobno."),
    ("03_EFI6_RCCB", "Wariant - ETI EFI6 wyłącznik różnicowoprądowy",
     ["Układ Biegunów", "Prąd Znamionowy [A]"],
     ["Typ Wyłącznika", "Znamionowy Prąd Różnicowy IΔn [mA]", "Znamionowy Warunkowy Prąd Zwarciowy [kA]"],
     "Wyłącznik RCCB EFI6, producent ETI."),
    ("04_SV_rozlaczniki", "Wariant - ETI SV rozłącznik izolacyjny",
     ["Układ Biegunów", "Prąd Znamionowy [A]"], [],
     "Rozłącznik izolacyjny rodziny SV, producent ETI; dwa zapisy typu produktu w eksporcie połączono po oznaczeniu SV."),
    ("05_SSQ_przelaczniki", "Wariant - ETI SSQ przełącznik I-0-II",
     ["Układ Biegunów", "Prąd Znamionowy [A]"], [],
     "Przełącznik trójpołożeniowy SSQ, producent ETI."),
    ("06_ETIMAT_T", "Wariant - ETI ETIMAT T ogranicznik mocy",
     ["Prąd Znamionowy [A]"], ["Układ Biegunów"],
     "Ogranicznik mocy ETIMAT T 3P, producent ETI."),
    ("07_ED2_rozlaczniki", "Wariant - ETI ED2 rozłącznik kompaktowy",
     ["Prąd Znamionowy [A]"], ["Układ Biegunów"],
     "Rozłączniki kompaktowe ED2, producent ETI. Przed użyciem tagu skoryguj kategorię MCCB w sklepie."),
    ("08_STV_D02", "Wariant - ETI STV D02 rozłącznik bezpiecznikowy",
     ["Układ Biegunów"], ["Prąd Znamionowy [A]", "Rozmiar Wkładki Bezpiecznikowej"],
     "Rozłącznik bezpiecznikowy STV D02, producent ETI."),
    ("09_SON_H", "Wariant - ETI SON H sygnalizator obecności napięcia",
     ["Barwa Lampek"], ["Napięcie [V]"],
     "Sygnalizator SON H, producent ETI."),
    ("10_IZ12_szyny", "Wariant - ETI IZ12 szyna grzebieniowa",
     ["Liczba Faz", "Liczba Modułów"], ["Długość", "Prąd Znamionowy [A]", "Przekrój Poprzeczny [mm²]"],
     "Szyny IZ12, producent ETI; długość jest powiązana z liczbą modułów i pokazana jako kontrola, nie osobna oś."),
]


def v(item: dict, attribute: str) -> str:
    raw = item.get(P + attribute)
    return "" if raw is None else str(raw).strip()


def header(sheet, row: int, widths: list[int]) -> None:
    sheet.freeze_panes = f"A{row+1}"
    sheet.auto_filter.ref = f"A{row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    for cell in sheet[row]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row].height = 42
    for idx, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    for cells in sheet.iter_rows():
        for cell in cells:
            if cell.row != row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    source_wb = load_workbook(SOURCE, read_only=True, data_only=True)
    rows = list(source_wb.active.values)
    names = list(rows[0])
    items = {str(row[names.index("SKU")]): dict(zip(names, row)) for row in rows[1:] if row[names.index("SKU")]}
    if len(items) != 136:
        raise ValueError("Unexpected source product count")

    old_wb = load_workbook(PREVIOUS, read_only=True, data_only=True)
    grouped: dict[str, list[str]] = defaultdict(list)
    for row in list(old_wb["Tagi do importu"].values)[1:]:
        grouped[str(row[2])].append(str(row[1]))
    mcb_tag = GROUPS[0][1]
    grouped[mcb_tag].append("002111516/ETI")
    rails_tag = GROUPS[-1][1]
    grouped[rails_tag] = ["002921026/ETI", "002921024/ETI", "002921020/ETI",
                          "002921022/ETI", "002921018/ETI"]
    if set(grouped) != {g[1] for g in GROUPS}:
        raise ValueError("Unexpected tag vocabulary")
    all_tagged = [sku for members in grouped.values() for sku in members]
    if len(all_tagged) != 118 or len(set(all_tagged)) != 118 or not set(all_tagged).issubset(items):
        raise ValueError("Unexpected variant membership")

    book = Workbook()
    summary = book.active
    summary.title = "Podsumowanie"
    summary.append(["Arkusz", "Tag", "Produkty", "Osie wariantu", "Wartości osi w grupie", "Wspólna podstawa / uwagi"])
    import_sheet = book.create_sheet("Tagi do importu")
    import_sheet.append(["id", "SKU", "Product Tags"])

    for sheet_name, tag, axes, extra, basis in GROUPS:
        members = [items[sku] for sku in grouped[tag]]
        for axis in axes:
            if any(not v(item, axis) for item in members):
                raise ValueError(f"Missing axis {axis} in {tag}")
        combos = [tuple(v(item, a).casefold() for a in axes) for item in members]
        if len(combos) != len(set(combos)):
            raise ValueError(f"Duplicate variant combination in {tag}")
        axis_summary = "; ".join(f"{axis}: {', '.join(sorted({v(i, axis) for i in members}, key=str.casefold))}" for axis in axes)
        summary.append([sheet_name, tag, len(members), "; ".join(axes), axis_summary, basis])

        sheet = book.create_sheet(sheet_name)
        for line in [
            ("TAG", tag), ("Liczba produktów", len(members)),
            ("Połączone na podstawie", basis), ("Osie wariantu", "; ".join(axes)),
            ("Wartości osi w grupie", axis_summary),
            ("Źródło", "Eksport ETI z uzupełnionymi atrybutami; wartości atrybutów zapisane dla każdego SKU."),
        ]:
            sheet.append(line)
        sheet.append([])
        sheet.append(["id", "SKU", "Nazwa produktu", "Kategoria w sklepie", "Typ produktu",
                      *axes, *extra, "Wartości rozróżniające ten SKU", "Tag"])
        for item in sorted(members, key=lambda x: str(x["SKU"])):
            row = [item["id"], item["SKU"], item["Title"], item["Kategorie produktów"],
                   v(item, "Typ Produktu"), *[v(item, a) for a in axes], *[v(item, a) for a in extra],
                   "; ".join(f"{a}: {v(item, a)}" for a in axes), tag]
            sheet.append(row)
            import_sheet.append([item["id"], item["SKU"], tag])
        header(sheet, 9, [12, 20, 75, 50, 46, *([29] * len(axes)), *([29] * len(extra)), 115, 70])
        sheet.column_dimensions["B"].width = 65
        sheet.column_dimensions["A"].width = 26
        for rr in range(1, 7):
            sheet.cell(rr, 1).font = Font(name="Arial", size=10, bold=True, color="17365D")
            sheet.cell(rr, 2).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[3].height = 40
        sheet.row_dimensions[5].height = 55

    header(summary, 1, [29, 71, 15, 85, 115, 105])
    header(import_sheet, 1, [12, 20, 72])
    import_sheet.auto_filter.ref = f"A1:C{import_sheet.max_row}"

    without = book.create_sheet("Bez tagu")
    without.append(["id", "SKU", "Nazwa produktu", "Kategoria", "Powód pozostawienia bez tagu"])
    old_audit = {str(row[1]): row for row in list(old_wb["Kontrola produktów"].values)[1:]}
    for sku in sorted(set(items) - set(all_tagged)):
        item = items[sku]
        reason = old_audit[sku][15] or "Brak bezpiecznej grupy wariantowej."
        without.append([item["id"], sku, item["Title"], item["Kategorie produktów"], reason])
    if without.max_row != 19:
        raise ValueError("Unexpected untagged count")
    header(without, 1, [12, 20, 75, 55, 115])

    notes = book.create_sheet("Metodyka")
    notes.append(["Temat", "Informacja"])
    for entry in [
        ("Źródło danych", str(SOURCE)),
        ("Znaczenie tagu", "Jeden identyczny tag łączy produkty proste należące do tej samej technicznej rodziny. Nie tworzy produktu variable."),
        ("Osie", "Atrybuty podane w każdym arkuszu odróżniają SKU. Kombinacja wartości osi jest unikalna w danej grupie."),
        ("Uzupełnienia względem poprzedniego pliku", "Do ETIMAT 6 dołączono 002111516 po potwierdzeniu 1P. Utworzono nową grupę pięciu szyn IZ12 po uzupełnieniu liczby modułów i długości."),
        ("Szyny", "Długość jest skorelowana z liczbą modułów: 12→210 mm, 18→310 mm, 54→1000 mm. Nie powielano tej samej osi jako niezależnego wariantu."),
        ("ED2", "Tag przygotowany, ale trzy ED2 w eksporcie mają niewłaściwą kategorię MCCB. Zalecana korekta kategorii przed wykorzystaniem tagu."),
        ("Bez tagu", "18 produktów nie otrzymało tagu z powodu braku kompletnych osi, odrębnej rodziny lub zbyt małej grupy. Szczegóły w arkuszu Bez tagu."),
        ("Do importu", "Arkusz Tagi do importu zawiera tylko 118 produktów z tagiem; nie ma pustych wartości Product Tags."),
        ("Zmiany online", "Nie wykonano żadnego zapisu do sklepu ani BaseLinkera."),
    ]:
        notes.append(entry)
    header(notes, 1, [37, 125])

    book.save(OUTPUT)
    print(f"tags={len(GROUPS)} tagged={len(all_tagged)} untagged={without.max_row-1}")
    print(OUTPUT)


if __name__ == "__main__":
    main()
