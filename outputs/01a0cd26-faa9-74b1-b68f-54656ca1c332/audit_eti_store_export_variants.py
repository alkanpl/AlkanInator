"""Audit the ETI store export and prepare conservative variant tags.

Input workbooks are read-only. No live WooCommerce/BaseLinker changes are made.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
SOURCE = Path(r"C:\Users\Handlowiec\Downloads\eti-modulowka-25.09.xlsx")
RULES = Path(r"C:\Users\Handlowiec\Desktop\ETI MODUŁÓWKA\Aparatura_modulowa_atrybuty_filtry_i_warianty_2026-09-24.xlsx")
OUTPUT = HERE / "ETI_eksport_sklepu_audyt_atrybutow_i_tagi_wariantow_2026-09-25.xlsx"

P = "Atrybut Produktu: "
AXES = {
    "ETIMAT 6": ["Charakterystyka Wyzwalania", "Prąd Znamionowy [A]", "Układ Biegunów"],
    "KZS-2M": ["Charakterystyka Wyzwalania", "Prąd Znamionowy [A]", "Typ Wyłącznika"],
    "EFI6": ["Układ Biegunów", "Prąd Znamionowy [A]"],
    "SV": ["Układ Biegunów", "Prąd Znamionowy [A]"],
    "SSQ": ["Układ Biegunów", "Prąd Znamionowy [A]"],
    "ETIMAT T": ["Prąd Znamionowy [A]"],
    "ED2": ["Prąd Znamionowy [A]"],
    "STV D02": ["Układ Biegunów"],
    "SON H": ["Barwa Lampek"],
}
TAGS = {
    "ETIMAT 6": "Wariant - ETI ETIMAT 6 wyłącznik nadprądowy",
    "KZS-2M": "Wariant - ETI KZS-2M wyłącznik różnicowo-nadprądowy",
    "EFI6": "Wariant - ETI EFI6 wyłącznik różnicowoprądowy",
    "SV": "Wariant - ETI SV rozłącznik izolacyjny",
    "SSQ": "Wariant - ETI SSQ przełącznik I-0-II",
    "ETIMAT T": "Wariant - ETI ETIMAT T ogranicznik mocy",
    "ED2": "Wariant - ETI ED2 rozłącznik kompaktowy",
    "STV D02": "Wariant - ETI STV D02 rozłącznik bezpiecznikowy",
    "SON H": "Wariant - ETI SON H sygnalizator obecności napięcia",
}


def value(item: dict, attribute: str) -> str:
    raw = item.get(P + attribute)
    return "" if raw is None else str(raw).strip()


def base_category(item: dict) -> str:
    return str(item["Kategorie produktów"] or "").split(">")[-1].strip()


def expected_category(item: dict) -> tuple[str, str]:
    typ = value(item, "Typ Produktu")
    sku = str(item["SKU"])
    if sku in {"004671271/ETI", "004671272/ETI", "004671273/ETI"}:
        return "Rozłączniki i aparatura bezpiecznikowa", "ED2 jest rozłącznikiem kompaktowym, nie wyłącznikiem MCCB. Docelowo: Wyłączniki i rozłączniki kompaktowe > Rozłączniki kompaktowe. Wzorzec rozłącznika użyty tymczasowo do audytu."
    if typ == "Wstawka redukcyjna":
        return "Akcesoria do aparatury modułowej", "Wstawka redukcyjna jest akcesorium; przypisanie do kategorii rozłączników wymaga przeglądu."
    if typ == "Wyłącznik różnicowo-nadprądowy RCBO":
        return "Wyłączniki różnicowo-nadprądowe (RCBO) — KATEGORIA DO DODANIA", ""
    return base_category(item), ""


def variant_family(item: dict) -> str:
    title = str(item["Title"] or "").upper().replace(" ", "")
    sku = str(item["SKU"])
    typ = value(item, "Typ Produktu")
    if typ == "Wyłącznik nadprądowy MCB":
        if sku == "002125712/ETI":
            return "ETIMAT 10 — konflikt danych"
        return "ETIMAT 10" if "ETIMAT10" in title else "ETIMAT 6" if "ETIMAT6" in title else "nieustalona rodzina MCB"
    if typ == "Wyłącznik różnicowo-nadprądowy RCBO":
        return "KZS-2M" if "KZS-2M" in title else "KZS-4M" if "KZS-4M" in title else "nieustalona rodzina RCBO"
    if typ == "Wyłącznik różnicowoprądowy RCCB":
        return "EFI6" if "EFI6" in title else "nieustalona rodzina RCCB"
    if typ in {"Rozłącznik izolacyjny", "Rozłącznik izolacyjny modułowy"}:
        if sku.startswith("00467127"):
            return "ED2"
        return "SV" if "SV" in title else "nieustalona rodzina rozłączników"
    if typ == "Przełącznik trójpołożeniowy":
        return "SSQ"
    if typ == "Ogranicznik mocy":
        return "ETIMAT T"
    if typ == "Rozłącznik bezpiecznikowy":
        return "STV D02" if "STVD02" in title else "EFD" if "EFD" in title else "nieustalona rodzina bezpiecznikowa"
    if typ == "Sygnalizator obecności napięcia":
        return "SON H"
    if typ == "Stycznik modułowy":
        return "R styczniki"
    if typ == "Szyna grzebieniowa":
        return "IZ12 szyny"
    return typ or "nieustalony typ"


def family_decision(family: str, items: list[dict]) -> tuple[bool, str]:
    axes = AXES.get(family)
    if not axes:
        if family == "R styczniki":
            return False, "Brak napięcia cewki w atrybucie Napięcie [V] dla części SKU; powtarzają się kombinacje prądu i styków."
        if family == "IZ12 szyny":
            return False, "Brak wymaganych atrybutów Liczba Modułów i Długość; szyn o tej samej liczbie faz nie da się odróżnić atrybutami."
        if family == "EFD":
            return False, "Dwa produkty różnią się jednocześnie rozmiarem wkładki i prądem; nie tworzono pary wieloosiowej."
        return False, "Brak drugiego produktu tej samej rodziny lub bezpiecznych osi wariantu."
    if len(items) < 2:
        return False, "Tylko jeden produkt w rodzinie."
    for axis in axes:
        if any(not value(item, axis) for item in items):
            return False, f"Niepełna oś wariantu: {axis}."
    combos = [tuple(value(item, axis).casefold() for axis in axes) for item in items]
    if len(set(combos)) != len(combos):
        return False, "Powtórzone kombinacje wartości atrybutów wariantowych."
    varying = [axis for axis in axes if len({value(item, axis) for item in items}) > 1]
    if not varying:
        return False, "Brak różnicy w przypisanych atrybutach wariantowych."
    if len(items) == 2 and len(varying) > 1:
        return False, "Dwa produkty różnią się na więcej niż jednej osi."
    if len(varying) > 1:
        isolated = [item for i, item in enumerate(items) if not any(
            sum(combos[i][k] != combos[j][k] for k in range(len(axes))) == 1
            for j in range(len(items)) if j != i
        )]
        if isolated:
            return False, "Co najmniej jeden produkt nie ma sąsiada różniącego się jedną osią."
    return True, "Wspólna funkcja i rodzina; kompletne, unikalne wartości na bezpiecznych osiach."


def style(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 36
    for idx, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.values)
    headers = list(rows[0])
    needed = {"id", "SKU", "Title", "Kategorie produktów"}
    needed |= {P + axis for axes in AXES.values() for axis in axes}
    needed |= {P + "Typ Produktu", P + "Producent", P + "Konfiguracja Styków (NO/NC)", P + "Napięcie [V]"}
    if any(headers.count(name) != 1 for name in needed):
        raise ValueError("Missing or ambiguous audit columns")
    items = [dict(zip(headers, row)) for row in rows[1:] if row[headers.index("SKU")]]
    if len(items) != 136 or len({str(x["SKU"]) for x in items}) != len(items):
        raise ValueError("Unexpected row count or duplicate SKU")

    rule_wb = load_workbook(RULES, read_only=True, data_only=True)
    category_rules: dict[str, list[str]] = defaultdict(list)
    for row in list(rule_wb["Filtry per kategoria"].values)[3:]:
        if row[0] and row[4]:
            category_rules[str(row[0])].append(str(row[4]))

    missing_rows = []
    product_rows = []
    category_rows = []
    family_items: dict[str, list[dict]] = defaultdict(list)
    applicable_count = Counter()
    filled_count = Counter()
    for item in items:
        family_items[variant_family(item)].append(item)
        expected, category_note = expected_category(item)
        if expected not in category_rules:
            raise ValueError(f"No requirement template for {expected}")
        typ = value(item, "Typ Produktu")
        attrs = [a for a in category_rules[expected]
                 if a != "Rozmiar Wkładki Bezpiecznikowej"
                 or typ in {"Rozłącznik bezpiecznikowy", "Podstawa bezpiecznikowa DC"}]
        # ED2 uses disconnector obligations, not the inappropriate MCCB contacts filter.
        absent = [a for a in attrs if not value(item, a)]
        for attr in attrs:
            applicable_count[expected] += 1
            filled_count[expected] += bool(value(item, attr))
        for attr in absent:
            missing_rows.append([item["id"], item["SKU"], item["Title"], base_category(item), expected,
                                 attr, "Brak wartości w eksporcie; nie uzupełniano na podstawie tytułu ani domysłów.",
                                 "Eksport sklepu + arkusz Filtry per kategoria", "brak", "DO_SPRAWDZENIA"])
        if category_note:
            category_rows.append([item["id"], item["SKU"], item["Title"], base_category(item),
                                  expected, category_note, "DO_SPRAWDZENIA"])
        product_rows.append({"item": item, "expected": expected, "absent": absent, "category_note": category_note})

    qualified = {family: [item for item in group if all(value(item, axis) for axis in AXES.get(family, []))]
                 for family, group in family_items.items()}
    decisions = {family: family_decision(family, qualified[family]) for family in family_items}
    result = Workbook()
    import_tab = result.active
    import_tab.title = "Tagi do importu"
    import_tab.append(["id", "SKU", "Product Tags", "Status tagu", "Uwagi"])
    audit_tab = result.create_sheet("Kontrola produktów")
    audit_tab.append(["id", "SKU", "Nazwa", "Kategoria w sklepie", "Kategoria wzorcowa", "Typ produktu",
                      "Obowiązkowe atrybuty", "Uzupełnione", "Brakujące atrybuty", "Status atrybutów",
                      "Rodzina wariantowa", "Tag wariantu", "Osie wariantu", "Wartości osi tego SKU",
                      "Status wariantu", "Powód / uwagi", "Źródło", "Pewność"])
    group_tab = result.create_sheet("Grupy wariantów")
    group_tab.append(["Rodzina", "Liczba produktów", "SKU z tagiem", "Wspólny tag", "Osie wyznaczające warianty",
                      "Wartości osi w grupie", "Decyzja", "Uzasadnienie", "SKU"])
    missing_tab = result.create_sheet("Braki atrybutów")
    missing_tab.append(["id", "SKU", "Nazwa", "Kategoria w sklepie", "Kategoria wzorcowa", "Brakujący atrybut",
                        "Uwagi", "Źródło", "Pewność", "Status"])
    for row in missing_rows:
        missing_tab.append(row)
    category_tab = result.create_sheet("Kategorie do weryfikacji")
    category_tab.append(["id", "SKU", "Nazwa", "Kategoria w sklepie", "Wzorzec do audytu", "Uwaga", "Status"])
    for row in category_rows:
        category_tab.append(row)
    data_tab = result.create_sheet("Dane do weryfikacji")
    data_tab.append(["SKU", "Pole", "Wartość w eksporcie", "Wartość wg ETI", "Źródło", "Status"])
    eti_002125712 = "https://www.etipolam.com.pl/katalog-elektroniczny/aparaty-modulowe-zabezpieczajace/wylaczniki-nadpradowe-i-akcesoria/wylaczniki-nadpradowe/002125712-etimat-10-3p-b6"
    data_tab.append(["002125712/ETI", "Rodzina / nazwa", "ETIMAT 6", "ETIMAT 10", eti_002125712, "DO_SPRAWDZENIA"])
    data_tab.append(["002125712/ETI", "Znamionowa zwarciowa zdolność wyłączania Icn [kA]", "6", "10", eti_002125712, "DO_SPRAWDZENIA"])

    tagged_count = 0
    for record in product_rows:
        item = record["item"]
        family = variant_family(item)
        allowed, reason = decisions[family]
        is_qualified = item in qualified[family]
        axes = AXES.get(family, []) if allowed and is_qualified else []
        tag = TAGS[family] if allowed and is_qualified else ""
        if not is_qualified and family in AXES:
            reason = "Brak wartości osi: " + "; ".join(a for a in AXES[family] if not value(item, a))
        if tag:
            tagged_count += 1
            import_tab.append([item["id"], item["SKU"], tag, "GOTOWY",
                               "Najpierw skoryguj kategorię ED2." if family == "ED2" else ""])
        expected = record["expected"]
        typ = value(item, "Typ Produktu")
        attrs = [a for a in category_rules[expected]
                 if a != "Rozmiar Wkładki Bezpiecznikowej"
                 or typ in {"Rozłącznik bezpiecznikowy", "Podstawa bezpiecznikowa DC"}]
        audit_tab.append([item["id"], item["SKU"], item["Title"], base_category(item), expected, typ,
                          "; ".join(attrs), len(attrs) - len(record["absent"]), "; ".join(record["absent"]),
                          "OK" if not record["absent"] else "DO_SPRAWDZENIA", family, tag,
                          "; ".join(axes), "; ".join(f"{a}: {value(item, a)}" for a in axes),
                          "GOTOWY" if tag else "BEZ TAGU", reason + (" " + record["category_note"] if record["category_note"] else ""),
                          "Eksport sklepu; arkusz Filtry per kategoria; arkusz Osie wariantów",
                          "wysoka dla wartości w eksporcie; grupowanie konserwatywne"])

    for family, group in sorted(family_items.items()):
        allowed, reason = decisions[family]
        axes = AXES.get(family, []) if allowed else []
        summary = "; ".join(f"{axis}: {', '.join(sorted({value(item, axis) for item in qualified[family]}))}" for axis in axes)
        skipped = len(group) - len(qualified[family])
        if allowed and skipped:
            reason += f" {skipped} SKU bez kompletnych osi pozostaje bez tagu."
        group_tab.append([family, len(group), len(qualified[family]) if allowed else 0, TAGS.get(family, "") if allowed else "",
                          "; ".join(axes), summary, "TAG" if allowed else "WSTRZYMANE", reason,
                          ", ".join(str(item["SKU"]) for item in group)])

    overview = result.create_sheet("Podsumowanie", 0)
    overview.append(["Miara", "Wynik"])
    for line in [
        ("Produkty w eksporcie", len(items)),
        ("Kontrole atrybutów obowiązkowych", sum(applicable_count.values())),
        ("Uzupełnione wartości obowiązkowe", sum(filled_count.values())),
        ("Brakujące wartości obowiązkowe", len(missing_rows)),
        ("Produkty z kompletem obowiązkowych atrybutów", sum(not r["absent"] for r in product_rows)),
        ("Produkty z tagiem wariantowym", tagged_count),
        ("Produkty bez tagu", len(items) - tagged_count),
        ("Grupy z tagiem", sum(allowed for allowed, _ in decisions.values())),
        ("Kategorie do weryfikacji", len(category_rows)),
    ]:
        overview.append(line)
    overview.append(["Kategoria wzorcowa", "Kontrole", "Wypełnione", "Braki"])
    for cat in sorted(applicable_count):
        overview.append([cat, applicable_count[cat], filled_count[cat], applicable_count[cat] - filled_count[cat]])

    notes = result.create_sheet("Metodyka")
    notes.append(["Temat", "Zasada / źródło"])
    for row in [
        ("Plik źródłowy", str(SOURCE)),
        ("Plik wzorcowy", str(RULES)),
        ("Zakres", "136 SKU ETI z eksportu sklepu; oryginał bez zmian. Sprawdzano tylko niepuste wartości atrybutów, bez ponownego potwierdzania każdej wartości u producenta."),
        ("Obowiązkowość", "Każdy filtr z arkusza Filtry per kategoria traktowany jako wymagany dla właściwego rodzaju produktu. Rozmiar wkładki dotyczy wyłącznie bezpiecznikowych aparatów."),
        ("Warianty", "Produkty pozostają simple. Wspólny Product Tags oznacza grupę; wartości osi z przypisanych atrybutów odróżniają SKU. Nie tworzono rodzica variable."),
        ("Bezpieczne tagi", "Tylko grupa co najmniej dwóch SKU, kompletne osie, unikalne kombinacje. Para różniąca się na dwóch osiach nie otrzymuje tagu."),
        ("Braki", "Nie wypełniano braków z tytułów, opisów ani domysłów. Arkusz Tagi do importu zawiera wyłącznie SKU z bezpiecznym tagiem."),
        ("ED2", "Trzy rozłączniki ED2 są błędnie w kategorii MCCB. Audytowane według wymagań rozłączników, a tag grupy ED2 jest przygotowany, lecz przed użyciem rekomendowana korekta kategorii."),
        ("Źródło producenta ED2", "https://www.etigroup.eu/products-and-solutions/moulded-case-circuit-breakers-and-disconnectors-etibreak/switch-disconnectors-125-1600a/004671271-ed2-125-3"),
        ("Konflikt 002125712", "ETI identyfikuje 002125712 jako ETIMAT 10 3p B6, 10 kA; eksport sklepu podaje ETIMAT 6 i 6 kA. Dane pozostawiono bez zmian, a SKU nie ma tagu."),
        ("Źródło ETI 002125712", eti_002125712),
        ("Seria", "Atrybut Seria nie jest osią wariantu ani wymogiem audytu; rodziny techniczne rozpoznano z nazw/oznaczeń produktów."),
        ("Brak zmian online", "Nie modyfikowano WooCommerce ani BaseLinkera."),
    ]:
        notes.append(row)

    for tab, widths in [
        (overview, [58, 24, 20, 20]), (import_tab, [12, 20, 62, 31, 67]),
        (audit_tab, [12, 20, 75, 48, 58, 44, 90, 18, 90, 20, 30, 62, 78, 100, 20, 115, 75, 63]),
        (group_tab, [28, 18, 17, 67, 80, 115, 21, 105, 105]),
        (missing_tab, [12, 20, 75, 48, 58, 65, 85, 70, 15, 20]),
        (category_tab, [12, 20, 75, 48, 58, 115, 20]),
        (data_tab, [20, 60, 33, 35, 110, 22]),
        (notes, [31, 130]),
    ]:
        style(tab, widths)
    result.save(OUTPUT)
    print(f"products={len(items)} checks={sum(applicable_count.values())} missing={len(missing_rows)} tagged={tagged_count} groups={sum(a for a,_ in decisions.values())} category_issues={len(category_rows)}")
    print(OUTPUT)


if __name__ == "__main__":
    main()
