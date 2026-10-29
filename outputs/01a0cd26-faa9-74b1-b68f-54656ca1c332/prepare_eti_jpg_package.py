"""Download exact-SKU ETI product images and make a local-photo updater package.

Input workbooks are never overwritten. Run from a network-enabled environment.
"""

from __future__ import annotations

import hashlib
import io
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from PIL import Image, ImageOps
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "ETI_modulowka_IMPORT_BaseLinker_ze_zdjeciami_2026-09-25.xlsx"
IMAGE_DIR = HERE / "ETI_JPG_1000x1000_2026-09-25"
UPDATER = HERE / "ETI_modulowka_Alkan_BL_Updater_JPG_1000x1000_2026-09-25.xlsx"
REPORT = HERE / "ETI_modulowka_RAPORT_JPG_1000x1000_2026-09-25.xlsx"
PACKAGE = HERE / "ETI_modulowka_Alkan_BL_Updater_JPG_1000x1000_2026-09-25.zip"

MISSING = {
    "002181094", "002181060", "002181098", "002181061", "002181062",
    "002181063", "002181064", "002181065", "002181088", "002181089",
}
EXPECTED_HEADERS = [
    "product_id", "name", "sku", "ean", "manufacturer_name", "category",
    "description", "features", "images_urls",
]


def fetch_and_convert(code: str, middle_url: str) -> dict:
    big_url = middle_url.replace("_Photo_T_MIDDLE.webp", "_Photo_T_BIG.webp")
    if big_url == middle_url:
        raise ValueError(f"Unexpected source URL: {middle_url}")
    referer = (
        f"https://www.etipolam.com.pl/katalog-elektroniczny/{code}"
        if "etipolam.com.pl" in middle_url
        else f"https://www.etigroup.eu/products-services/{code}"
    )
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/131.0.0.0 Safari/537.36",
        "Referer": referer,
        "Accept": "image/avif,image/webp,image/apng,image/*,*/*;q=0.8",
    }
    failures = []
    for source_url in (big_url, middle_url):
        for attempt in range(3):
            try:
                with urlopen(Request(source_url, headers=headers), timeout=25) as response:
                    content_type = response.headers.get("Content-Type", "").lower()
                    data = response.read(20 * 1024 * 1024 + 1)
                if len(data) > 20 * 1024 * 1024 or not content_type.startswith("image/"):
                    raise ValueError(f"Invalid image response: {content_type}, {len(data)} bytes")
                with Image.open(io.BytesIO(data)) as opened:
                    opened.load()
                    original_size = opened.size
                    image = ImageOps.exif_transpose(opened)
                    if image.mode in {"RGBA", "LA"}:
                        background = Image.new("RGB", image.size, "white")
                        background.paste(image, mask=image.getchannel("A"))
                        image = background
                    else:
                        image = image.convert("RGB")
                    image = ImageOps.pad(
                        image,
                        (1000, 1000),
                        method=Image.Resampling.LANCZOS,
                        color=(255, 255, 255),
                    )
                    output = IMAGE_DIR / f"{code}.jpg"
                    image.save(output, "JPEG", quality=90, optimize=True, progressive=True)
                with Image.open(output) as check:
                    check.verify()
                with Image.open(output) as check:
                    assert check.format == "JPEG" and check.size == (1000, 1000)
                return {
                    "code": code,
                    "source_url": source_url,
                    "original_size": original_size,
                    "output": output,
                    "bytes": output.stat().st_size,
                    "sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
                    "status": "OK",
                    "detail": "",
                }
            except (HTTPError, URLError, OSError, ValueError, AssertionError) as exc:
                failures.append(f"{source_url}: {type(exc).__name__}: {exc}")
                time.sleep(0.5 * (attempt + 1))
    return {
        "code": code,
        "source_url": "",
        "original_size": None,
        "output": None,
        "bytes": 0,
        "sha256": "",
        "status": "POBRANIE NIEUDANE",
        "detail": " | ".join(failures),
    }


def style(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", color="000000")
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> None:
    source_book = load_workbook(SOURCE)
    assert source_book.sheetnames == ["Import"]
    sheet = source_book["Import"]
    assert [cell.value for cell in sheet[1]] == EXPECTED_HEADERS
    assert sheet.max_row == 128
    IMAGE_DIR.mkdir(exist_ok=True)

    products = {}
    for row in range(2, sheet.max_row + 1):
        sku = sheet.cell(row, 3).value
        assert isinstance(sku, str) and sku.endswith("/ETI")
        code = sku[:-4]
        assert code not in products
        url = sheet.cell(row, 9).value or ""
        assert (code in MISSING) == (not url), (code, url)
        if url:
            assert f"/{code}/" in url and f"/{code}_Photo_T_MIDDLE.webp" in url
        products[code] = {"row": row, "sku": sku, "name": sheet.cell(row, 2).value, "middle_url": url}

    result = {}
    with ThreadPoolExecutor(max_workers=4) as executor:
        futures = {
            executor.submit(fetch_and_convert, code, product["middle_url"]): code
            for code, product in products.items() if code not in MISSING
        }
        for future in as_completed(futures):
            code = futures[future]
            result[code] = future.result()
            print(f"{len(result)}/{len(futures)} {code}: {result[code]['status']}", flush=True)

    failed = [code for code, item in result.items() if item["status"] != "OK"]
    if failed:
        print(f"Failed downloads/conversions: {failed}")

    sheet.cell(1, 10, "images_paths")
    for code, product in products.items():
        row = product["row"]
        sheet.cell(row, 9).value = None
        if code in result and result[code]["status"] == "OK":
            sheet.cell(row, 10, f"{IMAGE_DIR.name}/{code}.jpg")
        else:
            sheet.cell(row, 10).value = None
    source_book.save(UPDATER)

    report = Workbook()
    overview = report.active
    overview.title = "Podsumowanie"
    overview.append(["Kontrola zdjęć", "Wartość"])
    for row in [
        ("Produkty", 127),
        ("JPG 1000 × 1000 px", sum(item["status"] == "OK" for item in result.values())),
        ("Brak zdjęcia u ETI", len(MISSING)),
        ("Nieudane pobrania", len(failed)),
        ("Źródło", "Oficjalne zdjęcia ETI przypisane do dokładnego kodu produktu; preferowana wersja BIG."),
        ("Format wyjściowy", "JPEG (JPG), RGB, 1000 × 1000 px, białe marginesy bez kadrowania produktu."),
        ("Uwaga o jakości", "Jeśli źródło miało mniej niż 1000 px, wymiar wynika z powiększenia; nie oznacza to większej ilości detali."),
        ("Wykorzystanie", "Arkusz używa images_paths i jest przeznaczony do Alkan BL Updatera; rozpakuj ZIP przed użyciem."),
        ("Nie wykonywano", "Nie przesłano produktów ani zdjęć do BaseLinkera. Potwierdzenie warunków wykorzystania zdjęć pozostaje po stronie użytkownika/ETI."),
    ]:
        overview.append(row)
    style(overview, [36, 115])

    details = report.create_sheet("Kontrola JPG")
    details.append(["SKU", "Nazwa", "Status", "Źródło", "Wymiary źródła", "Plik JPG", "Bajty", "SHA-256", "Uwagi"])
    for code, product in products.items():
        item = result.get(code)
        if item is None:
            status, source_url, dimensions, output, size, digest, detail = (
                "BRAK U ETI", "", "", "", "", "", "Na karcie ETI jest tylko placeholder.",
            )
        else:
            status = item["status"]
            source_url = item["source_url"]
            dimensions = f"{item['original_size'][0]} × {item['original_size'][1]}" if item["original_size"] else ""
            output = str(item["output"].relative_to(HERE)) if item["output"] else ""
            size = item["bytes"]
            digest = item["sha256"]
            detail = item["detail"]
        details.append([product["sku"], product["name"], status, source_url, dimensions, output, size, digest, detail])
        cell = details.cell(details.max_row, 4)
        if cell.value:
            cell.hyperlink = cell.value
            cell.font = Font(name="Arial", color="0563C1", underline="single")
    style(details, [20, 75, 24, 115, 24, 60, 14, 70, 95])
    report.save(REPORT)

    with zipfile.ZipFile(PACKAGE, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.write(UPDATER, UPDATER.name)
        archive.write(REPORT, REPORT.name)
        for code, item in result.items():
            if item["status"] == "OK":
                archive.write(item["output"], f"{IMAGE_DIR.name}/{code}.jpg")

    # Independent read-back of the deliverables and unchanged non-photo fields.
    verify_book = load_workbook(UPDATER, read_only=True)
    verify_sheet = verify_book["Import"]
    source_check = load_workbook(SOURCE, read_only=True)["Import"]
    assert verify_sheet.max_row == source_check.max_row == 128
    assert verify_sheet.max_column == 10
    assert sum(bool(verify_sheet.cell(r, 10).value) for r in range(2, 129)) == len(result) - len(failed)
    assert all(not verify_sheet.cell(r, 9).value for r in range(2, 129))
    for source_row, target_row in zip(source_check.iter_rows(), verify_sheet.iter_rows()):
        assert [c.value for c in source_row[:8]] == [c.value for c in target_row[:8]]
    with zipfile.ZipFile(PACKAGE) as archive:
        assert archive.testzip() is None
        assert len(archive.namelist()) == len(result) - len(failed) + 2
    print(f"JPG: {len(result) - len(failed)}; missing at ETI: {len(MISSING)}; failed: {len(failed)}")
    print(f"Workbook: {UPDATER}")
    print(f"Package: {PACKAGE}")


if __name__ == "__main__":
    main()
