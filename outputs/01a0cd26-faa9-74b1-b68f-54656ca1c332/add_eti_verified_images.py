"""Add SKU-matched official ETI photo URLs without changing the input workbook.

The code sets below are the result of checking each ETI product card in a browser
on 2026-09-25.  The ten ETIMAT T cards expose an empty-image placeholder rather
than a product photograph, so they intentionally remain blank in the import.
"""

from __future__ import annotations

from collections import Counter
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "ETI_modulowka_IMPORT_BaseLinker_atrybuty_nazwy_mala_litera.xlsx"
TARGET = HERE / "ETI_modulowka_IMPORT_BaseLinker_ze_zdjeciami_2026-09-25.xlsx"
REPORT = HERE / "ETI_modulowka_RAPORT_zdjec_2026-09-25.xlsx"

POLISH_PHOTO_CODES = {
    "002061261", "002061262", "002061661", "002061662", "002061663",
}
MISSING_PHOTO_CODES = {
    "002181094", "002181060", "002181098", "002181061", "002181062",
    "002181063", "002181064", "002181065", "002181088", "002181089",
}

HEADERS = [
    "product_id", "name", "sku", "ean", "manufacturer_name", "category",
    "description", "features", "images_urls",
]


def official_photo_url(code: str) -> str:
    host = "www.etipolam.com.pl" if code in POLISH_PHOTO_CODES else "www.etigroup.eu"
    return f"https://{host}/product_db/idents/{code}/en-GB/photo/{code}_Photo_T_MIDDLE.webp"


def product_page_url(code: str) -> str:
    if code in POLISH_PHOTO_CODES or code in MISSING_PHOTO_CODES:
        return f"https://www.etipolam.com.pl/katalog-elektroniczny/{code}"
    return f"https://www.etigroup.eu/products-services/{code}"


def style_table(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="F2F6FA")


def main() -> None:
    original = load_workbook(SOURCE)
    assert original.sheetnames == ["Import"], original.sheetnames
    sheet = original["Import"]
    assert [c.value for c in sheet[1]] == HEADERS
    assert sheet.max_row == 128, sheet.max_row

    records: list[dict[str, str | int]] = []
    seen: set[str] = set()
    for row_number in range(2, sheet.max_row + 1):
        row = {header: sheet.cell(row_number, col).value for col, header in enumerate(HEADERS, 1)}
        sku = row["sku"]
        assert isinstance(sku, str) and sku.endswith("/ETI"), (row_number, sku)
        code = sku.removesuffix("/ETI")
        assert len(code) == 9 and code.isdigit(), (row_number, sku)
        assert code not in seen, code
        assert not row["images_urls"], (row_number, row["images_urls"])
        seen.add(code)

        image_url = "" if code in MISSING_PHOTO_CODES else official_photo_url(code)
        sheet.cell(row_number, 9).value = image_url or None
        records.append({
            "row": row_number,
            "code": code,
            "sku": sku,
            "name": row["name"],
            "category": row["category"],
            "page": product_page_url(code),
            "photo": image_url,
            "status": "OK" if image_url else "BRAK ZDJECIA",
            "source": "ETI – zdjęcie przypisane do SKU" if image_url else "ETI – placeholder bez zdjęcia produktu",
        })

    assert len(records) == 127
    assert seen & MISSING_PHOTO_CODES == MISSING_PHOTO_CODES
    assert seen & POLISH_PHOTO_CODES == POLISH_PHOTO_CODES
    counts = Counter(record["status"] for record in records)
    assert counts == {"OK": 117, "BRAK ZDJECIA": 10}, counts
    original.save(TARGET)

    audit = Workbook()
    summary = audit.active
    summary.title = "Podsumowanie"
    summary.append(["Kontrola zdjęć ETI", "Wartość"])
    summary_rows = [
        ("Data weryfikacji", date(2026, 9, 25).isoformat()),
        ("Plik wejściowy", SOURCE.name),
        ("Plik importowy", TARGET.name),
        ("Liczba produktów", 127),
        ("Zdjęcie produktu z oficjalnej karty ETI", 117),
        ("Brak zdjęcia produktu na karcie ETI", 10),
        ("Format importu", "images_urls: jeden bezpośredni URL do zdjęcia dla każdego SKU; puste dla braków"),
        ("Weryfikacja", "Każda karta produktu i każde z 117 zdjęć zostały otwarte w przeglądarce; kod URL odpowiada SKU."),
        ("Braki", "Na kartach 10 produktów ETIMAT T producent pokazuje eti-product-empty.png."),
        ("Uwaga", "Nie użyto zdjęć podobnych produktów ani placeholderów. Importu do BaseLinkera nie wykonywano."),
        ("Prawa do użycia", "Publiczne adresy zdjęć ETI; potwierdzenie warunków wykorzystania w sklepie pozostaje po stronie użytkownika/ETI."),
    ]
    for line in summary_rows:
        summary.append(line)
    style_table(summary, [36, 115])

    check = audit.create_sheet("Kontrola zdjęć")
    check.append(["Wiersz importu", "SKU", "Nazwa", "Kategoria", "Status", "Adres zdjęcia", "Karta ETI", "Źródło / uwaga"])
    for record in records:
        check.append([
            record["row"], record["sku"], record["name"], record["category"],
            record["status"], record["photo"], record["page"], record["source"],
        ])
        row = check.max_row
        for column in (6, 7):
            cell = check.cell(row, column)
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
        if record["status"] != "OK":
            check.cell(row, 5).fill = PatternFill("solid", fgColor="FFF2CC")
    style_table(check, [16, 20, 74, 54, 18, 110, 80, 45])

    missing = audit.create_sheet("Brak zdjęcia")
    missing.append(["SKU", "Nazwa", "Kategoria", "Karta ETI", "Powód", "Dalsze działanie"])
    for record in records:
        if record["status"] == "OK":
            continue
        missing.append([
            record["sku"], record["name"], record["category"], record["page"],
            "Na karcie ETI jest placeholder zamiast zdjęcia produktu.",
            "Poprosić ETI o zdjęcie przypisane do tego kodu; nie podmieniać zdjęciem podobnego modelu.",
        ])
        cell = missing.cell(missing.max_row, 4)
        cell.hyperlink = cell.value
        cell.font = Font(color="0563C1", underline="single")
    style_table(missing, [20, 75, 54, 80, 60, 75])
    audit.save(REPORT)

    result = load_workbook(TARGET, read_only=True)
    result_sheet = result["Import"]
    source_check = load_workbook(SOURCE, read_only=True)["Import"]
    assert result_sheet.max_row == source_check.max_row
    assert result_sheet.max_column == source_check.max_column
    for source_row, target_row in zip(source_check.iter_rows(), result_sheet.iter_rows()):
        assert all(a.value == b.value for a, b in zip(source_row[:-1], target_row[:-1]))
    assert sum(bool(result_sheet.cell(row, 9).value) for row in range(2, 129)) == 117
    print(f"Import: {TARGET}")
    print(f"Report: {REPORT}")
    print(f"Matched photos: {counts['OK']}; missing photos: {counts['BRAK ZDJECIA']}")


if __name__ == "__main__":
    main()
