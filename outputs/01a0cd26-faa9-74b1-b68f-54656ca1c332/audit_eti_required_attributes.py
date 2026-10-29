"""Audit ETI BaseLinker features against the supplied category-filter workbook.

Never writes to either input. Facts copied from an existing feature are tracked
separately from the three manufacturer-confirmed modular-width family rules.
"""

from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "ETI_modulowka_IMPORT_BaseLinker_opisy_unikalne_FAQ.xlsx"
RULES = Path(r"C:\Users\Handlowiec\Desktop\ETI MODUŁÓWKA\Aparatura_modulowa_atrybuty_filtry_i_warianty_2026-09-24.xlsx")
IMPORT_OUT = HERE / "ETI_modulowka_IMPORT_BaseLinker_atrybuty_wymagane.xlsx"
REPORT_OUT = HERE / "ETI_modulowka_RAPORT_atrybuty_wymagane.xlsx"

ALIASES = {
    "Typ produktu": "Typ Produktu",
    "Prąd znamionowy [A]": "Prąd Znamionowy [A]",
    "Układ biegunów": "Układ Biegunów",
    "Charakterystyka wyzwalania": "Charakterystyka Wyzwalania",
    "Znamionowa zdolność zwarciowa Icn [kA]": "Znamionowa Zwarciowa Zdolność Wyłączania Icn [kA]",
    "Prąd różnicowy IΔn [mA]": "Znamionowy Prąd Różnicowy IΔn [mA]",
    "Typ prądu różnicowego": "Typ Wyłącznika",
    "Warunkowy prąd zwarciowy Inc [kA]": "Znamionowy Warunkowy Prąd Zwarciowy [kA]",
    "Konfiguracja styków": "Konfiguracja Styków (NO/NC)",
    "Liczba modułów": "Liczba Modułów",
    "Barwa lampek": "Barwa Lampek",
    "Rozmiar": "Rozmiar Wkładki Bezpiecznikowej",
}

FAMILY_CATEGORY = {
    "Wyłącznik nadprądowy MCB": "Wyłączniki nadprądowe",
    "Ogranicznik mocy": "Ograniczniki mocy",
    "Wyłącznik różnicowoprądowy RCCB": "Wyłączniki różnicowoprądowe",
    "Wyłącznik różnicowo-nadprądowy RCBO": "Wyłączniki różnicowo-nadprądowe (RCBO) — KATEGORIA DO DODANIA",
    "Rozłącznik izolacyjny": "Rozłączniki i aparatura bezpiecznikowa",
    "Przełącznik trójpołożeniowy": "Przełączniki Sieć-Agregat",
    "Sygnalizator obecności napięcia": "Sterowanie i sygnalizacja",
    "Stycznik modułowy": "Styczniki",
    "Rozłącznik bezpiecznikowy": "Rozłączniki i aparatura bezpiecznikowa",
    "Wstawka redukcyjna": "Akcesoria do aparatury modułowej",
    "Podstawa bezpiecznikowa DC": "Rozłączniki i aparatura bezpiecznikowa",
}

MODULE_SOURCES = {
    "ETIMAT 6 1p": "https://www.etigroup.eu/products-and-solutions/modular-current-and-residual-protective-devices-asti/miniature-circuit-breakers-mcb-and-accessories-etimat/miniature-circuit-breakers-mcb/002111512-etimat-6-1p-b6",
    "ETIMAT 6 3p": "https://www.etigroup.eu/products-and-solutions/modular-current-and-residual-protective-devices-asti/miniature-circuit-breakers-mcb-and-accessories-etimat/miniature-circuit-breakers-mcb/002115512-etimat-6-3p-b6",
    "KZS-2M": "https://www.etigroup.eu/products-and-solutions/modular-current-and-residual-protective-devices-asti/residual-and-overcurrent-protection-rcbo-and-accessories-kzs/residual-and-overcurrent-protection-rcbo/002173201-kzs-2m-a-b6-0-03",
}


def filter_rules() -> dict[str, list[str]]:
    book = load_workbook(RULES, read_only=True, data_only=True)
    sheet = book["Filtry per kategoria"]
    result: dict[str, list[str]] = defaultdict(list)
    for row in list(sheet.values)[3:]:
        category, attribute = row[0], row[4]
        if category and attribute:
            result[str(category)].append(str(attribute))
    return dict(result)


def normalise(features: dict[str, str], family: str) -> tuple[dict[str, str], list[tuple[str, str, str]]]:
    result: dict[str, str] = {}
    changes = []
    for old_key, raw in features.items():
        if old_key == "Seria":
            raise ValueError("Seria must remain excluded")
        new_key = ALIASES.get(old_key, old_key)
        value = str(raw).strip()
        if old_key == "Napięcie znamionowe AC [V]" and family == "Sygnalizator obecności napięcia":
            new_key = "Napięcie [V]"
            value = f"{value} V AC"
        if new_key in result and result[new_key] != value:
            raise ValueError(f"Conflicting feature {new_key}: {result[new_key]} vs {value}")
        result[new_key] = value
        if new_key != old_key or value != raw:
            changes.append((old_key, new_key, value))
    return result, changes


def confirmed_modules(symbol: str, family: str, poles: str) -> tuple[str, str, str] | None:
    if family == "Wyłącznik nadprądowy MCB":
        match = re.fullmatch(r"ETIMAT 6 ([13])p [BCD]\d+", symbol)
        if match and poles == f"{match.group(1)}P":
            key = f"ETIMAT 6 {match.group(1)}p"
            return match.group(1), MODULE_SOURCES[key], "Reguła rodziny ETIMAT 6: symbol i układ biegunów zgodne; karta ETI dla wykonania 1P/3P."
    if family == "Wyłącznik różnicowo-nadprądowy RCBO":
        if re.fullmatch(r"KZS-2M (?:A|AC) [BC]\d+/0\.03", symbol) and poles == "1P+N":
            return "2", MODULE_SOURCES["KZS-2M"], "Reguła rodziny KZS-2M: 2 moduły; symbol i układ 1P+N zgodne z kartą ETI."
    return None


def applicable_attributes(category: str, family: str, requirements: dict[str, list[str]]) -> list[str]:
    attrs = requirements[category]
    # The reference denominator is 5/7 for fuse size, not 5/56 for every
    # product in the combined disconnector/fuse category.
    if family not in {"Rozłącznik bezpiecznikowy", "Podstawa bezpiecznikowa DC"}:
        attrs = [attr for attr in attrs if attr != "Rozmiar Wkładki Bezpiecznikowej"]
    return attrs


def style_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    for column, width in zip(sheet.columns, widths):
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = width
        for cell in list(column)[1:]:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    requirements = filter_rules()
    book = load_workbook(SOURCE)
    sheet = book["Import"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    required = {"sku", "name", "category", "features", "description"}
    if not required.issubset(headers):
        raise ValueError(f"Missing import columns: {required - set(headers)}")

    audit = []
    changes = []
    missing = []
    category_issues = []
    counts_before = Counter()
    counts_after = Counter()
    counts_applicable = Counter()
    by_category = Counter()
    for index in range(2, sheet.max_row + 1):
        sku = sheet.cell(index, headers["sku"]).value
        name = sheet.cell(index, headers["name"]).value
        current_category = sheet.cell(index, headers["category"]).value
        original = json.loads(sheet.cell(index, headers["features"]).value)
        family = original["Typ produktu"]
        expected_category = FAMILY_CATEGORY[family]
        expected_path = f"Aparatura modułowa > {expected_category}"
        if current_category != expected_path:
            category_issues.append([sku, name, current_category, expected_category,
                "Kategorie nie zostały zmienione w pliku importowym; mapowanie służy tylko audytowi atrybutów."])
        features, renamed = normalise(original, family)
        for old_key, new_key, value in renamed:
            changes.append([sku, old_key, new_key, value, "PRZENIESIONO / UJEDNOLICONO",
                "Wartość w poprzednim imporcie ETI; zmieniono etykietę atrybutu do arkusza filtrów.",
                "wysoka", "OK"])
        module_fact = None
        if "Liczba Modułów" not in features:
            module_fact = confirmed_modules(features["Symbol"], family, features.get("Układ Biegunów", ""))
            if module_fact:
                value, url, note = module_fact
                features["Liczba Modułów"] = value
                changes.append([sku, "", "Liczba Modułów", value, "DODANO",
                    f"{note} {url}", "wysoka — potwierdzona rodzina/model", "OK"])

        for attr in applicable_attributes(expected_category, family, requirements):
            counts_applicable[(expected_category, attr)] += 1
            before = next((str(v).strip() for k, v in original.items() if ALIASES.get(k, k) == attr), "")
            if attr == "Napięcie [V]" and family == "Sygnalizator obecności napięcia":
                before = original.get("Napięcie znamionowe AC [V]", "")
            after = features.get(attr, "")
            counts_before[(expected_category, attr)] += bool(before)
            counts_after[(expected_category, attr)] += bool(after)
            status = "OK" if after else "DO_SPRAWDZENIA"
            source = "Poprzedni import ETI — atrybut istniejący"
            confidence = "wysoka"
            if attr == "Liczba Modułów" and module_fact:
                source, confidence = module_fact[1], "wysoka — potwierdzona rodzina/model"
            if not after:
                source, confidence = "Brak potwierdzonej wartości", "brak"
                missing.append([sku, name, current_category, expected_category, attr,
                    features.get("Symbol", ""), "Nie wpisano wartości bez potwierdzenia."])
            audit.append([sku, name, current_category, expected_category, attr, before, after,
                source, confidence, status])
        by_category[expected_category] += 1
        sheet.cell(index, headers["features"], json.dumps(features, ensure_ascii=False, separators=(",", ":")))

    if sheet.max_row != 128:
        raise ValueError(f"Unexpected ETI product count: {sheet.max_row - 1}")
    book.save(IMPORT_OUT)

    report = Workbook()
    summary = report.active
    summary.title = "Podsumowanie"
    summary.append(["Kategoria wzorcowa", "Produkty ETI", "Wymagane atrybuty", "Wypełnione przed", "Wypełnione po", "Pozostałe braki"])
    for category, product_count in sorted(by_category.items()):
        attrs = requirements[category]
        before = sum(counts_before[(category, attr)] for attr in attrs)
        after = sum(counts_after[(category, attr)] for attr in attrs)
        applicable = sum(counts_applicable[(category, attr)] for attr in attrs)
        summary.append([category, product_count, applicable, before, after, applicable - after])
    summary.append(["RAZEM", sum(by_category.values()), len(audit), sum(counts_before.values()),
        sum(counts_after.values()), len(missing)])
    style_sheet(summary, [65, 17, 21, 20, 20, 20])

    for title, header, rows, widths in [
        ("Kontrola SKU", ["SKU", "Produkt", "Kategoria w imporcie", "Kategoria wg arkusza", "Atrybut wymagany", "Wartość przed", "Wartość po", "Źródło", "Pewność", "Status"], audit,
         [19, 58, 53, 64, 57, 25, 25, 90, 34, 18]),
        ("Zmiany", ["SKU", "Stara etykieta", "Nowa etykieta", "Wartość", "Zmiana", "Źródło / uzasadnienie", "Pewność", "Status"], changes,
         [19, 48, 58, 28, 35, 110, 36, 18]),
        ("Niepotwierdzone braki", ["SKU", "Produkt", "Kategoria w imporcie", "Kategoria wg arkusza", "Brakujący atrybut", "Symbol", "Uwagi"], missing,
         [19, 58, 53, 64, 57, 36, 65]),
        ("Kategorie do weryfikacji", ["SKU", "Produkt", "Kategoria w imporcie", "Kategoria wg arkusza", "Uwaga"], category_issues,
         [19, 58, 53, 64, 95]),
    ]:
        tab = report.create_sheet(title)
        tab.append(header)
        for row in rows:
            tab.append(row)
        style_sheet(tab, widths)

    notes = report.create_sheet("Metodyka i źródła")
    notes.append(["Temat", "Informacja"])
    for row in [
        ("Zakres", "127 pozycji ETI z poprzedniego importu z opisami i FAQ. Import źródłowy oraz plik wymagań pozostały bez zmian."),
        ("Wymagania", "Arkusz „Filtry per kategoria” z dostarczonego pliku; wymagany jest każdy wymieniony tam atrybut dla produktu właściwego typu."),
        ("Mapowanie", "Kategorie w imporcie ETI są mniej szczegółowe niż w arkuszu wymagań. W audycie zastosowano kategorię wzorcową według typu produktu; nie zmieniono kategorii importu."),
        ("Nazwy atrybutów", "Istniejące wartości zachowano, ale etykiety potrzebne filtrom ujednolicono do dokładnego zapisu z arkusza wymagań."),
        ("Liczba Modułów — ETIMAT 6", "Dla symboli ETIMAT 6 1p / 3p ETI podaje odpowiednio 1 / 3 szerokości modułowe na kartach produktów. W obrębie tej rodziny zastosowano regułę wyłącznie przy zgodności symbolu i układu biegunów."),
        ("Liczba Modułów — KZS-2M", "ETI podaje 2 szerokości modułowe dla rodziny KZS-2M. Regułę zastosowano wyłącznie do symboli KZS-2M A/AC z układem 1P+N."),
        ("Źródło 1P", MODULE_SOURCES["ETIMAT 6 1p"]),
        ("Źródło 3P", MODULE_SOURCES["ETIMAT 6 3p"]),
        ("Źródło KZS-2M", MODULE_SOURCES["KZS-2M"]),
        ("Ważne", "Źródła modułów są kartami reprezentatywnymi rodzin/modeli, a nie osobną kartą każdego SKU. W razie wymogu potwierdzenia 1:1 przed importem sprawdź każdy kod producenta."),
        ("Warianty", "Nie tworzono tagów wariantowych; plik wymagań opisuje też osie wariantów, ale to osobna konfiguracja sklepu poza uzupełnianiem atrybutów."),
    ]:
        notes.append(row)
    style_sheet(notes, [35, 125])
    report.save(REPORT_OUT)
    print(f"products={sum(by_category.values())}; required_checks={len(audit)}; before={sum(counts_before.values())}; after={sum(counts_after.values())}; missing={len(missing)}; changes={len(changes)}; category_issues={len(category_issues)}")
    print(IMPORT_OUT)
    print(REPORT_OUT)


if __name__ == "__main__":
    main()
