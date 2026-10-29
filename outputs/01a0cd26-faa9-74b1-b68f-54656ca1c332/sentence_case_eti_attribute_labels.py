"""Change only feature labels to Polish sentence case in the ETI import."""

import json
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "ETI_modulowka_IMPORT_BaseLinker_atrybuty_wymagane.xlsx"
TARGET = HERE / "ETI_modulowka_IMPORT_BaseLinker_atrybuty_nazwy_mala_litera.xlsx"

LABELS = {
    "Barwa Lampek": "Barwa lampek",
    "Charakterystyka Wyzwalania": "Charakterystyka wyzwalania",
    "Konfiguracja Styków (NO/NC)": "Konfiguracja styków (NO/NC)",
    "Liczba Modułów": "Liczba modułów",
    "Prąd Znamionowy [A]": "Prąd znamionowy [A]",
    "Rozmiar Wkładki Bezpiecznikowej": "Rozmiar wkładki bezpiecznikowej",
    "Typ Produktu": "Typ produktu",
    "Typ Wyłącznika": "Typ wyłącznika",
    "Układ Biegunów": "Układ biegunów",
    "Znamionowa Zwarciowa Zdolność Wyłączania Icn [kA]": "Znamionowa zwarciowa zdolność wyłączania Icn [kA]",
    "Znamionowy Prąd Różnicowy IΔn [mA]": "Znamionowy prąd różnicowy IΔn [mA]",
    "Znamionowy Warunkowy Prąd Zwarciowy [kA]": "Znamionowy warunkowy prąd zwarciowy [kA]",
}


def main() -> None:
    book = load_workbook(SOURCE)
    sheet = book["Import"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    feature_column = columns["features"]
    changed = 0
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row, feature_column)
        original = json.loads(cell.value)
        renamed = {}
        for key, value in original.items():
            new_key = LABELS.get(key, key)
            if new_key in renamed:
                raise ValueError(f"Duplicate key after renaming: {new_key}")
            renamed[new_key] = value
            changed += new_key != key
        cell.value = json.dumps(renamed, ensure_ascii=False, separators=(",", ":"))
    book.save(TARGET)
    print(f"products={sheet.max_row - 1}; labels_changed={changed}; output={TARGET}")


if __name__ == "__main__":
    main()
