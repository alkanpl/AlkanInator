"""Build a BaseLinker import from the reviewed ETI attributes workbook."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "ETI_modulowka_tytuly_atrybuty.xlsx"
TARGET = HERE / "ETI_modulowka_IMPORT_BaseLinker.xlsx"
MANUFACTURER_DATA = "ETI, Jana Pawła II 18 06-100 Pułtusk, etipolam@etipolam.com.pl"

CATEGORY_BY_TYPE = {
    "Wyłącznik nadprądowy MCB": "Aparatura modułowa > Wyłączniki nadprądowe",
    "Ogranicznik mocy": "Aparatura modułowa > Ograniczniki mocy",
    "Wyłącznik różnicowoprądowy RCCB": "Aparatura modułowa > Wyłączniki różnicowoprądowe",
    # No dedicated RCBO/fuse-holder path is present in the saved Woo taxonomy.
    "Wyłącznik różnicowo-nadprądowy RCBO": "Aparatura modułowa",
    "Rozłącznik izolacyjny": "Aparatura modułowa > Rozłączniki izolacyjne",
    "Przełącznik trójpołożeniowy": "Aparatura modułowa > Przełączniki Sieć-Agregat",
    "Sygnalizator obecności napięcia": "Aparatura modułowa > Lampki sygnalizacyjne",
    "Stycznik modułowy": "Aparatura modułowa > Styczniki",
    "Rozłącznik bezpiecznikowy": "Aparatura modułowa",
    "Wstawka redukcyjna": "Aparatura modułowa",
    "Podstawa bezpiecznikowa DC": "Aparatura modułowa",
}

FEATURE_NAMES = {
    "Typ Produktu": "Typ produktu",
    "Charakterystyka Wyzwalania": "Charakterystyka wyzwalania",
    "Prąd Znamionowy [A]": "Prąd znamionowy [A]",
    "Układ Biegunów": "Układ biegunów",
    "Znamionowa Zdolność Zwarciowa Icn [kA]": "Znamionowa zdolność zwarciowa Icn [kA]",
    "Prąd Różnicowy IΔn [mA]": "Prąd różnicowy IΔn [mA]",
    "Typ Prądu Różnicowego": "Typ prądu różnicowego",
    "Warunkowy Prąd Zwarciowy Inc [kA]": "Warunkowy prąd zwarciowy Inc [kA]",
    "Napięcie Znamionowe AC [V]": "Napięcie znamionowe AC [V]",
    "Napięcie Znamionowe DC [V]": "Napięcie znamionowe DC [V]",
    "Napięcie Cewki [V]": "Napięcie cewki [V]",
    "Konfiguracja Styków": "Konfiguracja styków",
    "Rodzaj Wkładki": "Rodzaj wkładki",
    "Rozmiar": "Rozmiar",
    "Barwa Lampek": "Barwa lampek",
    "Ilość Lampek": "Ilość lampek",
    "Pozycje Przełącznika": "Pozycje przełącznika",
    "Kompatybilność": "Kompatybilność",
    "Rodzaj Akcesoria": "Rodzaj akcesoria",
    "Liczba Modułów": "Liczba modułów",
    "Stopień Ochrony [IP]": "Stopień ochrony [IP]",
    "Kategoria Użytkowania": "Kategoria użytkowania",
}

IMPORT_HEADERS = [
    "product_id", "name", "sku", "ean", "manufacturer_name",
    "category", "description", "features", "images_urls",
]


def as_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def main() -> None:
    source_wb = load_workbook(SOURCE, read_only=True, data_only=True)
    source_ws = source_wb["Do wystawienia"]
    source_rows = source_ws.values
    headers = list(next(source_rows))
    counts = Counter()
    imported = []

    for values in source_rows:
        item = dict(zip(headers, values))
        code = as_text(item["Nr kodowy"])
        ean = as_text(item["EAN dla 1szt"])
        family = as_text(item["Typ Produktu"])
        if not code or not ean or not as_text(item["Nowy tytuł"]):
            raise ValueError(f"Missing identity data for {code or '[no code]'}")
        if item["Status P0"] != "OK" or item["Brakujące P0"]:
            raise ValueError(f"Mandatory attributes incomplete for {code}")
        if family not in CATEGORY_BY_TYPE:
            raise ValueError(f"Unmapped family: {family}")
        if as_text(item["Producent"]) != "ETI":
            raise ValueError(f"Unexpected producer for {code}")

        features = {
            "Producent": "ETI",
            "Dane producenta": MANUFACTURER_DATA,
            "Symbol": as_text(item["Symbol"]),
        }
        for source_name, target_name in FEATURE_NAMES.items():
            value = as_text(item[source_name])
            if value:
                features[target_name] = value

        imported.append([
            None,
            as_text(item["Nowy tytuł"]),
            f"{code}/ETI",
            ean,
            "ETI",
            CATEGORY_BY_TYPE[family],
            None,  # No confirmed product description in the input workbook.
            json.dumps(features, ensure_ascii=False, separators=(",", ":")),
            None,  # No confirmed product image URLs in the input workbook.
        ])
        counts[CATEGORY_BY_TYPE[family]] += 1

    source_wb.close()
    if len(imported) != 127:
        raise ValueError(f"Expected 127 products, found {len(imported)}")
    if len({row[2] for row in imported}) != len(imported):
        raise ValueError("Duplicate SKUs")
    if len({row[3] for row in imported}) != len(imported):
        raise ValueError("Duplicate EANs")

    output_wb = Workbook()
    output_ws = output_wb.active
    output_ws.title = "Import"
    output_ws.append(IMPORT_HEADERS)
    for row in imported:
        output_ws.append(row)

    output_ws.freeze_panes = "A2"
    output_ws.auto_filter.ref = f"A1:I{len(imported) + 1}"
    widths = {"A": 13, "B": 76, "C": 19, "D": 18, "E": 21, "F": 57, "G": 18, "H": 90, "I": 18}
    for column, width in widths.items():
        output_ws.column_dimensions[column].width = width
    for cell in output_ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(vertical="center")
    output_ws.row_dimensions[1].height = 25
    for row in output_ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top")
        for index in (3, 4):  # SKU and GTIN must remain text, including leading zeros.
            row[index - 1].number_format = "@"

    output_wb.save(TARGET)
    print(f"Saved {TARGET}; products={len(imported)}; categories={dict(counts)}")


if __name__ == "__main__":
    main()
