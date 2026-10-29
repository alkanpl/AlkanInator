"""Create a minimal image-only workbook, excluding products with existing photos.

Uses the read-only BaseLinker dry-run audit from 2026-09-25. No original file is
overwritten and no BaseLinker changes are made by this script.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "ETI_modulowka_Alkan_BL_Updater_JPG_1000x1000_2026-09-25.xlsx"
AUDIT = HERE / "baselinker_dry_run_ETI_2026-09-25" / "audit.csv"
TARGET = HERE / "ETI_modulowka_TYLKO_BRAKUJACE_ZDJECIA_2026-09-25.xlsx"
REPORT = HERE / "ETI_modulowka_RAPORT_przed_wysylka_2026-09-25.xlsx"


def style(sheet, widths: tuple[int, ...]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> None:
    source = load_workbook(SOURCE, read_only=True)["Import"]
    fields = {cell.value: cell.column for cell in source[1]}
    assert {"sku", "ean", "images_urls", "images_paths"} <= fields.keys()
    assert source.max_row == 128
    with AUDIT.open(encoding="utf-8-sig", newline="") as handle:
        audit = list(csv.DictReader(handle, delimiter=";"))
    assert len(audit) == 127
    by_sku = {row["sku"]: row for row in audit}
    assert len(by_sku) == 127

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Import"
    sheet.append(["sku", "ean", "images_paths"])
    exclusions = []
    eligible = []
    for row in range(2, source.max_row + 1):
        sku = source.cell(row, fields["sku"]).value
        ean = source.cell(row, fields["ean"]).value
        image_path = source.cell(row, fields["images_paths"]).value
        image_url = source.cell(row, fields["images_urls"]).value
        assert not image_url
        check = by_sku[sku]
        assert int(check["row_number"]) == row
        before = json.loads(check["before_json"] or "{}")
        existing_photo_count = before.get("images_count", 0)
        if check["status"] == "READY" and check["changed_fields"] == "images" and existing_photo_count == 0:
            assert image_path and (HERE / image_path).is_file()
            sheet.append([sku, ean, image_path])
            eligible.append(sku)
        else:
            if existing_photo_count > 0:
                reason = "Ma już zdjęcie w BaseLinkerze — bez aktualizacji"
            elif check["status"] == "INVALID":
                reason = check["error"]
            else:
                reason = "Brak zdjęcia produktu u ETI"
            exclusions.append((sku, check["status"], existing_photo_count, reason))
    assert len(eligible) == 89, len(eligible)
    assert len(exclusions) == 38, len(exclusions)
    assert sum("Ma już zdjęcie" in row[3] for row in exclusions) == 27
    assert sum("EAN" in row[3] for row in exclusions) == 1
    assert sum("Brak zdjęcia produktu" in row[3] for row in exclusions) == 10
    style(sheet, (22, 22, 64))
    workbook.save(TARGET)

    report = Workbook()
    overview = report.active
    overview.title = "Podsumowanie"
    overview.append(["Kontrola", "Liczba / informacja"])
    overview.append(["Zdjęcia dopuszczone do wysyłki", 89])
    overview.append(["Już mają zdjęcie w BaseLinkerze — nie aktualizować", 27])
    overview.append(["Brak zdjęcia u ETI", 10])
    overview.append(["Konflikt EAN — nie aktualizować", 1])
    overview.append(["Zakres pliku", "Tylko SKU, EAN i lokalna ścieżka JPG; bez nazw, opisów, cech i kategorii."])
    overview.append(["Źródło statusu", "Odczyt BaseLinkera w trybie DRY_RUN 2026-09-25, katalog Ecommerce Alkan (ID 47159)."])
    style(overview, (52, 100))
    excluded_sheet = report.create_sheet("Wykluczone")
    excluded_sheet.append(["SKU", "Status próby", "Zdjęcia w BaseLinkerze", "Powód wykluczenia"])
    for record in exclusions:
        excluded_sheet.append(record)
    style(excluded_sheet, (22, 22, 28, 100))
    report.save(REPORT)

    verified = load_workbook(TARGET, read_only=True)["Import"]
    assert verified.max_row == 90 and verified.max_column == 3
    assert [cell.value for cell in verified[1]] == ["sku", "ean", "images_paths"]
    print(f"Eligible for upload: {len(eligible)}")
    print(f"Excluded with existing photo: 27; no ETI photo: 10; EAN conflict: 1")
    print(f"Workbook: {TARGET}")
    print(f"Report: {REPORT}")


if __name__ == "__main__":
    main()
