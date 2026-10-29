"""Read-only classification report for ETI SSQ, ED2 and the broader MCCB list."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
ETI_IMPORT = HERE / "ETI_modulowka_IMPORT_BaseLinker_atrybuty_nazwy_mala_litera.xlsx"
TAXONOMY = Path(r"C:\Users\Handlowiec\Desktop\ETI MODUŁÓWKA\Aparatura_modulowa_atrybuty_filtry_i_warianty_2026-09-24.xlsx")
MASTER_EXPORT = ROOT / "archived_input_files" / "wszystko.csv"
OUTPUT = HERE / "ETI_klasyfikacja_przelacznikow_i_MCCB_2026-09-25.xlsx"

SSQ_CATEGORY = "Aparatura modułowa > Przełączniki Sieć-Agregat"
ED2_CATEGORY = "Aparatura modułowa > Rozłączniki izolacyjne"
ED2_OPTION = "Wyłączniki i rozłączniki kompaktowe > Rozłączniki kompaktowe (jeśli dodamy osobną kategorię)"


def style(sheet, widths: tuple[int, ...]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", color="000000")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row[0].row].height = 34


def link(cell) -> None:
    if cell.value:
        cell.hyperlink = cell.value
        cell.font = Font(name="Arial", color="0563C1", underline="single")


def main() -> None:
    rows = list(load_workbook(ETI_IMPORT, read_only=True, data_only=True)["Import"].values)
    headers = rows[0]
    assert headers[1:3] == ("name", "sku")
    assert len(rows) == 128
    eti = []
    for row in rows[1:]:
        details = json.loads(row[7])
        eti.append({
            "name": row[1], "sku": row[2], "category": row[5],
            "symbol": details.get("Symbol", ""),
            "type": details.get("Typ produktu", ""),
            "poles": details.get("Układ biegunów", ""),
            "current": details.get("Prąd znamionowy [A]", ""),
        })

    ssq = [p for p in eti if p["type"] == "Przełącznik trójpołożeniowy"]
    ed2 = [p for p in eti if p["symbol"].startswith("ED2 ")]
    assert len(ssq) == 9 and all(p["category"] == SSQ_CATEGORY for p in ssq)
    assert len(ed2) == 3 and all(p["category"] == ED2_CATEGORY for p in ed2)
    assert not any("MCCB" in p["type"].upper() or "EB2" in p["symbol"].upper() for p in eti)

    taxonomy = load_workbook(TAXONOMY, read_only=True, data_only=True)
    master_mccb = set()
    for row in taxonomy["Osie wariantów"].values:
        if row[0] and "Wyłączniki kompaktowe (MCCB)" in str(row[0]):
            master_mccb.update(sku.strip() for sku in str(row[7] or "").split(",") if sku.strip())
    assert master_mccb == {"36290/KAN", "36291/KAN", "36292/KAN", "37282/KAN", "116437/EAT"}
    master_names = {}
    with MASTER_EXPORT.open(encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            sku = row.get("SKU", "")
            if sku in master_mccb:
                master_names[sku] = row.get("Nazwa", "")
    assert set(master_names) == master_mccb

    book = Workbook()
    summary = book.active
    summary.title = "Wnioski"
    summary.append(["Temat", "Wniosek / zalecenie", "Źródło"])
    summary.append(["Bieżąca lista ETI", "127 produktów; 9 przełączników SSQ; 3 rozłączniki ED2; 0 wyłączników MCCB.", ETI_IMPORT.name])
    summary.append(["Przełączniki SSQ", f"Zostawić w istniejącej kategorii: {SSQ_CATEGORY}. Precyzyjniejsza nazwa ewentualna: Przełączniki modułowe I-0-II.", "https://www.etipolam.com.pl/katalog-elektroniczny/aparaty-automatyki-sygnalizacji-i-sterowania/laczniki-i-lampki-modulowe/002421414-ssq-125"])
    summary.append(["Wyłączniki kompaktowe MCCB", "Żaden z 127 produktów ETI nie jest MCCB. Kategoria oznacza wyłączniki automatyczne w obudowie kompaktowej, chroniące m.in. przed przeciążeniem i zwarciem.", "https://www.etipolam.com.pl/katalogi-pdf/catalog/wylaczniki-eb2-20-1600a"])
    summary.append(["Rozłączniki ED2", f"Nie przypisywać do Wyłączników kompaktowych MCCB. Pozostawić w {ED2_CATEGORY}; przy rozbudowie drzewa można dodać {ED2_OPTION}.", "https://www.etipolam.com.pl/katalog-elektroniczny/wylaczniki-i-rozlaczniki-kompaktowe/rozlaczniki-ed2-125-1600a/004671272-ed2-160-3"])
    summary.append(["Szerszy arkusz 556 produktów", "Zawiera 5 kandydatów do Wyłączników kompaktowych MCCB: 4 Kanlux i 1 Eaton. Nie należą do bieżącej listy ETI.", TAXONOMY.name])
    summary.append(["Stan BaseLinkera", "Nie zmieniano żadnej kategorii w BaseLinkerze; raport ma charakter rekomendacji.", ""])
    style(summary, (34, 116, 112))
    for row in (3, 4, 5):
        link(summary.cell(row, 3))

    ssq_sheet = book.create_sheet("ETI SSQ I-0-II")
    ssq_sheet.append(["SKU", "Nazwa", "Symbol", "Bieguny", "Prąd [A]", "Obecna i zalecana kategoria", "Karta ETI"])
    for product in ssq:
        code = product["sku"].split("/")[0]
        ssq_sheet.append([
            product["sku"], product["name"], product["symbol"], product["poles"],
            product["current"], SSQ_CATEGORY,
            f"https://www.etipolam.com.pl/katalog-elektroniczny/{code}",
        ])
    style(ssq_sheet, (22, 80, 20, 14, 14, 55, 80))
    for row in range(2, ssq_sheet.max_row + 1):
        link(ssq_sheet.cell(row, 7))

    ed2_sheet = book.create_sheet("ETI ED2 nie MCCB")
    ed2_sheet.append(["SKU", "Nazwa", "Symbol", "Prąd [A]", "Obecna kategoria", "Zalecenie", "Karta ETI"])
    for product in ed2:
        code = product["sku"].split("/")[0]
        ed2_sheet.append([
            product["sku"], product["name"], product["symbol"], product["current"],
            product["category"], "Rozłącznik izolacyjny; nie wyłącznik MCCB",
            f"https://www.etipolam.com.pl/katalog-elektroniczny/{code}",
        ])
    style(ed2_sheet, (22, 75, 20, 14, 55, 55, 80))
    for row in range(2, ed2_sheet.max_row + 1):
        link(ed2_sheet.cell(row, 7))

    wider = book.create_sheet("MCCB poza lista ETI")
    wider.append(["SKU", "Nazwa", "Producent", "Rekomendowana kategoria", "Podstawa"])
    for sku in sorted(master_mccb):
        brand = "Kanlux" if sku.endswith("/KAN") else "Eaton"
        wider.append([sku, master_names[sku], brand, "Wyłączniki kompaktowe (MCCB)", "Szerszy arkusz 556 produktów; poza bieżącą listą ETI"])
    style(wider, (22, 80, 20, 45, 65))

    book.save(OUTPUT)
    verify = load_workbook(OUTPUT, read_only=True, data_only=True)
    assert [(s.title, s.max_row) for s in verify.worksheets] == [
        ("Wnioski", 7), ("ETI SSQ I-0-II", 10), ("ETI ED2 nie MCCB", 4), ("MCCB poza lista ETI", 6),
    ]
    print(f"ETI products={len(eti)}; SSQ={len(ssq)}; ED2={len(ed2)}; ETI MCCB=0; wider MCCB={len(master_mccb)}")
    print(OUTPUT)


if __name__ == "__main__":
    main()
