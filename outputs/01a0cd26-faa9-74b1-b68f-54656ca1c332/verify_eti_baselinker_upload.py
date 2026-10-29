"""Read-back verification of ETI image-only BaseLinker upload."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
BEFORE_DIR = HERE / "baselinker_dry_run_ETI_2026-09-25"
APPLY_DIR = HERE / "baselinker_UPLOAD_ETI_missing_only_2026-09-25"
REPORT = HERE / "ETI_modulowka_RAPORT_po_wgraniu_zdjec_2026-09-25.xlsx"
CONFIG = ROOT / "src" / "baselinker_api" / "config.json"
ENDPOINT = "https://api.baselinker.com/connector.php"
INVENTORY_ID = 47159


def read_audit(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=";"))


def api_call(token: str, method: str, parameters: dict) -> dict:
    encoded = urlencode({"method": method, "parameters": json.dumps(parameters, separators=(",", ":"))}).encode()
    request = Request(ENDPOINT, data=encoded, headers={"X-BLToken": token})
    with urlopen(request, timeout=45) as response:
        payload = json.load(response)
    if payload.get("status") != "SUCCESS":
        raise RuntimeError(f"{method}: {payload.get('error_code')} {payload.get('error_message')}")
    return payload


def make_style(sheet, widths: tuple[int, ...]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    token = json.loads(CONFIG.read_text(encoding="utf-8-sig"))["token"]
    original_audit = read_audit(BEFORE_DIR / "audit.csv")
    apply_audit = read_audit(APPLY_DIR / "audit.csv")
    before = json.loads((BEFORE_DIR / "backup_before.json").read_text(encoding="utf-8"))
    assert len(original_audit) == 127
    assert len(apply_audit) == 89
    assert all(row["status"] == "SUCCESS" and row["changed_fields"] == "images" for row in apply_audit)
    applied = {row["sku"] for row in apply_audit}
    assert len(applied) == 89

    ids = [int(row["product_id"]) for row in original_audit]
    current = {}
    for start in range(0, len(ids), 50):
        response = api_call(token, "getInventoryProductsData", {
            "inventory_id": INVENTORY_ID,
            "products": ids[start:start + 50],
        })
        current.update({str(key): value for key, value in response.get("products", {}).items()})
    assert len(current) == 127, len(current)

    details = []
    for row in original_audit:
        sku = row["sku"]
        product_id = row["product_id"]
        previous_images = before[product_id].get("images") or {}
        new_images = current[product_id].get("images") or {}
        old_count = len(previous_images)
        new_count = len(new_images)
        if sku in applied:
            group = "Wysłane – wcześniej bez zdjęcia"
            status = "POTWIERDZONE" if old_count == 0 and new_count >= 1 else "DO SPRAWDZENIA"
        elif old_count >= 1:
            group = "Istniejące zdjęcie – nie aktualizowano"
            status = "BEZ ZMIAN" if previous_images == new_images else "DO SPRAWDZENIA"
        elif row["status"] == "INVALID":
            group = "Konflikt EAN – nie wysłano"
            status = "BEZ ZMIAN" if previous_images == new_images else "DO SPRAWDZENIA"
        else:
            group = "Brak zdjęcia u ETI – nie wysłano"
            status = "BEZ ZMIAN" if previous_images == new_images else "DO SPRAWDZENIA"
        details.append((sku, product_id, group, old_count, new_count, status))

    report = Workbook()
    overview = report.active
    overview.title = "Podsumowanie"
    overview.append(["Wynik kontroli", "Liczba / uwaga"])
    overview.append(["Wysłano z sukcesem według API", 89])
    overview.append(["Po odczycie mają zdjęcie", sum(x[2].startswith("Wysłane") and x[5] == "POTWIERDZONE" for x in details)])
    overview.append(["Wcześniejsze zdjęcia bez zmian", sum(x[2].startswith("Istniejące") and x[5] == "BEZ ZMIAN" for x in details)])
    overview.append(["Brak zdjęcia u ETI – nie wysłano", sum(x[2].startswith("Brak zdjęcia") for x in details)])
    overview.append(["Konflikt EAN – nie wysłano", sum(x[2].startswith("Konflikt") for x in details)])
    overview.append(["Do sprawdzenia", sum(x[5] == "DO SPRAWDZENIA" for x in details)])
    overview.append(["Zakres", "Katalog Ecommerce Alkan (ID 47159); wysłano wyłącznie pole images."])
    overview.append(["Podstawa", "Porównanie zdjęć przed wysyłką i po ponownym odczycie z BaseLinkera."])
    make_style(overview, (48, 105))
    check = report.create_sheet("Kontrola SKU")
    check.append(["SKU", "ID BaseLinker", "Grupa", "Zdjęcia przed", "Zdjęcia po", "Status"])
    for detail in details:
        check.append(detail)
    make_style(check, (22, 20, 48, 22, 22, 22))
    report.save(REPORT)

    counts = {"uploaded_confirmed": 0, "untouched_existing": 0, "review": 0}
    for row in details:
        if row[2].startswith("Wysłane") and row[5] == "POTWIERDZONE":
            counts["uploaded_confirmed"] += 1
        if row[2].startswith("Istniejące") and row[5] == "BEZ ZMIAN":
            counts["untouched_existing"] += 1
        if row[5] == "DO SPRAWDZENIA":
            counts["review"] += 1
    print(counts)
    print(REPORT)
    assert counts == {"uploaded_confirmed": 89, "untouched_existing": 27, "review": 0}, counts


if __name__ == "__main__":
    main()
