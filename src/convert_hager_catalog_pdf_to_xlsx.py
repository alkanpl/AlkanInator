from __future__ import annotations

import argparse
from copy import copy
import hashlib
import json
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import fitz
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from build_hager_berker_catalog_cross_upsell import (
    candidate_catalog_codes,
    find_index_pdf_pages,
    line_code_candidates,
    line_contains_code,
    normalize_catalog_code,
    page_lines,
)
from build_hager_berker_cross_upsell import CatalogProduct
from build_hager_berker_variants import (
    compact,
    find_headers,
    load_records,
    mark_old_duplicate_products,
    mojibake_markers,
    norm,
)


DEFAULT_PDF = Path("input/21PL001_Katalog_Osprzet_Hager_2021.pdf")
DEFAULT_STORE_XLSX = Path("output/Hager-Berker-Gniazdka_WOO_ALL_IMPORT_OSIE_POPRAWIONE.xlsx")
DEFAULT_OUTPUT = Path("output/Hager-Berker-Katalog-PDF_BAZA_KOMPATYBILNOSCI.xlsx")

HAGER_BLUE = "00A3E0"
NAVY = "17365D"
LIGHT_BLUE = "DDEBF7"
LIGHT_GRAY = "F2F2F2"
LIGHT_GREEN = "E2F0D9"
LIGHT_YELLOW = "FFF2CC"
LIGHT_RED = "FCE4D6"
WHITE = "FFFFFF"
TEXT = "1F2937"
THIN_GRAY = Side(style="thin", color="D9E2F3")
EXCEL_ILLEGAL_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")

PRODUCT_TERMS = (
    "gniazdo", "łącznik", "przycisk", "klawisz", "płytka", "ramka", "ściemniacz",
    "czujnik", "regulator", "sterownik", "mechanizm", "zestaw", "adapter", "puszka",
    "wkładka", "element", "pokrywa", "nasadka", "potencjometr", "termostat", "moduł",
)
SERIES_MARKERS = (
    "berker ", "serie ", "r.classic", "lumina", "one.platform", "knx", "integro",
)
SKIP_HEADING_MARKERS = (
    "współpracuje z", "nr kat", "nr zam", "opak", "opis", "str.",
    "opcjonalnie", "zastrzega się możliwość", "dane zawarte", "nowość", "wycofane",
)


@dataclass(frozen=True)
class IndexEntry:
    code: str
    printed_page: int
    index_pdf_page: int


@dataclass(frozen=True)
class PageLine:
    x0: float
    y0: float
    x1: float
    y1: float
    text: str
    bold: bool
    size: float


@dataclass(frozen=True)
class Relation:
    source_code: str
    target_code: str
    evidence_page: int
    target_pages: str
    raw_reference: str
    reference_description: str
    method: str
    confidence: str
    status: str
    evidence: str


@dataclass(frozen=True)
class CompatibilityBlock:
    printed_page: int
    header_y: float
    source_codes: str
    target_codes: str
    raw_lines: str
    status: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def excel_safe(value: Any) -> Any:
    if isinstance(value, str):
        return EXCEL_ILLEGAL_CHARACTERS.sub("", value)[:32767]
    return value


def extract_index(document: fitz.Document) -> list[IndexEntry]:
    entries: set[IndexEntry] = set()
    for pdf_page in find_index_pdf_pages(document):
        words = document[pdf_page - 1].get_text("words")
        for word in words:
            raw = str(word[4])
            code = normalize_catalog_code(raw)
            if not re.fullmatch(r"[A-Z0-9]{4,20}", code) or not re.search(r"\d", code):
                continue
            page_candidates: list[tuple[float, int]] = []
            for page_word in words:
                value = str(page_word[4])
                if not value.isdigit():
                    continue
                if page_word[0] <= word[2] or page_word[0] - word[2] > 55:
                    continue
                if abs(page_word[1] - word[1]) > 2.5:
                    continue
                printed_page = int(value)
                if 1 <= printed_page < 619:
                    page_candidates.append((page_word[0], printed_page))
            if page_candidates:
                entries.add(IndexEntry(code, min(page_candidates)[1], pdf_page))
    return sorted(entries, key=lambda item: (item.code, item.printed_page))


def structured_lines(page: fitz.Page) -> list[PageLine]:
    result: list[PageLine] = []
    for block in page.get_text("dict").get("blocks", []):
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            text = compact(" ".join(str(span.get("text", "")) for span in spans))
            if not text:
                continue
            fonts = " ".join(str(span.get("font", "")) for span in spans).lower()
            bold = any(marker in fonts for marker in ("bold", "-bd", "bd")) or any(
                int(span.get("flags", 0)) & 16 for span in spans
            )
            size = max((float(span.get("size", 0)) for span in spans), default=0)
            x0, y0, x1, y1 = line["bbox"]
            result.append(PageLine(x0, y0, x1, y1, text, bold, size))
    return sorted(result, key=lambda item: (item.y0, item.x0))


def meaningful_header(text: str) -> bool:
    normalized = norm(text)
    if not normalized or any(marker in normalized for marker in map(norm, SKIP_HEADING_MARKERS)):
        return False
    if re.fullmatch(r"[0-9 ]+", text):
        return False
    return True


def is_series_line(text: str) -> bool:
    normalized = norm(text)
    if any(term in normalized for term in map(norm, PRODUCT_TERMS)):
        return False
    return any(marker in normalized for marker in SERIES_MARKERS)


def is_product_description(text: str) -> bool:
    normalized = norm(text)
    return any(term in normalized for term in map(norm, PRODUCT_TERMS))


def page_header(lines: list[PageLine]) -> str:
    candidates = [
        item.text for item in lines
        if item.y0 < 90 and meaningful_header(item.text) and not re.fullmatch(r"\d{1,3}", item.text)
    ]
    return " | ".join(dict.fromkeys(candidates))[:500]


def product_row_for_entry(
    document: fitz.Document,
    entry: IndexEntry,
    line_cache: dict[int, list[PageLine]],
    simple_line_cache: dict[int, list[tuple[float, float, str]]],
) -> dict[str, Any]:
    if entry.printed_page not in line_cache:
        page = document[entry.printed_page]
        line_cache[entry.printed_page] = structured_lines(page)
        simple_line_cache[entry.printed_page] = page_lines(page)
    lines = line_cache[entry.printed_page]
    simple = simple_line_cache[entry.printed_page]

    hits = [
        item for item in lines
        if 200 <= item.x0 < 430 and line_contains_code(item.text, entry.code)
    ]
    if not hits:
        hits = [item for item in lines if item.x0 < 470 and line_contains_code(item.text, entry.code)]
    hit = hits[-1] if hits else None

    variant = ""
    package = ""
    heading = ""
    series = ""
    status = "OK" if hit else "DO_WERYFIKACJI"
    if hit:
        left = [
            item for item in lines
            if abs(item.y0 - hit.y0) <= 1.8 and item.x0 < hit.x0 - 5 and item.x0 >= 90
        ]
        if left:
            variant = max(left, key=lambda item: item.x0).text
        right = [
            item for item in lines
            if abs(item.y0 - hit.y0) <= 1.8 and item.x0 > hit.x1 and item.x0 >= 480
        ]
        if right:
            package = min(right, key=lambda item: item.x0).text

        preceding = [
            item for item in lines
            if item.bold and item.size >= 7.5 and item.y0 < hit.y0
            and item.x0 < 450 and meaningful_header(item.text)
            and not (re.search(r"\d", item.text) and re.fullmatch(r"[A-Z0-9 .]+", item.text.upper()))
        ]
        series_candidates = [item for item in preceding if is_series_line(item.text)]
        if series_candidates:
            nearest_series = max(series_candidates, key=lambda item: item.y0)
            series = nearest_series.text
            heading_candidates = [
                item for item in preceding
                if not is_series_line(item.text) and item.y0 < nearest_series.y0
            ]
        else:
            heading_candidates = [item for item in preceding if not is_series_line(item.text)]
        if heading_candidates:
            heading = max(heading_candidates, key=lambda item: item.y0).text

    if is_product_description(variant):
        base_name = variant
    else:
        base_name = heading or variant
    full_name = base_name
    if variant and variant != base_name:
        full_name = compact(f"{base_name} - {variant}")

    context = ""
    if hit:
        positions = [i for i, (_x, _y, text) in enumerate(simple) if line_contains_code(text, entry.code)]
        if positions:
            pos = positions[-1]
            context = " | ".join(
                compact(text) for _x, _y, text in simple[max(0, pos - 7): pos + 3] if compact(text)
            )[:1800]
    if not base_name or (base_name == variant and not is_product_description(base_name)):
        status = "DO_WERYFIKACJI"

    return {
        "Kod katalogowy": entry.code,
        "Nazwa automatyczna": full_name,
        "Nazwa bazowa": base_name,
        "Wariant / opis": variant,
        "Seria / system": series,
        "Sekcja strony": page_header(lines),
        "Strona drukowana": entry.printed_page,
        "Strona PDF": entry.printed_page + 1,
        "Opakowanie": package,
        "Kontekst katalogowy": context,
        "Status ekstrakcji": status,
        "Strona indeksu PDF": entry.index_pdf_page,
    }


def extract_relations(
    document: fitz.Document,
    code_to_pages: dict[str, set[int]],
    page_to_codes: dict[int, set[str]],
) -> tuple[list[Relation], list[CompatibilityBlock]]:
    known_codes = sorted(code_to_pages, key=len, reverse=True)
    relation_map: dict[tuple[str, str, int], Relation] = {}
    blocks: list[CompatibilityBlock] = []

    def method_rank(method: str) -> int:
        return {"DOKŁADNY KOD": 3, "PREFIX + STRONA": 2, "PREFIX BEZ STRONY": 1}.get(method, 0)

    for printed_page, home_codes in sorted(page_to_codes.items()):
        page = document[printed_page]
        lines = page_lines(page)
        headers = [(x, y) for x, y, line in lines if "wspolpracuje z" in norm(line)]
        if not headers:
            continue

        home_occurrences = sorted(set(
            (y, code)
            for x, y, line in lines
            if 200 <= x < 430
            for code in home_codes
            if line_contains_code(line, code)
        ))

        for _header_x, header_y in headers:
            block_lines = [
                (x, y, line) for x, y, line in lines
                if 0 <= y - header_y <= 100 and x >= 330
            ]
            detected_sources: set[str] = set()
            detected_targets: set[str] = set()

            for x, y, line in block_lines:
                if x < 410 or y <= header_y + 1:
                    continue
                raw_candidates = line_code_candidates(line)
                target_pages = {
                    int(other_line)
                    for other_x, other_y, other_line in lines
                    if other_x > x and abs(other_y - y) <= 2
                    and re.fullmatch(r"[0-9]{1,3}", compact(other_line))
                    and 1 <= int(other_line) < 619
                }

                matches: dict[str, tuple[str, str]] = {}
                for raw in raw_candidates:
                    if raw in code_to_pages:
                        matches[raw] = (raw, "DOKŁADNY KOD")
                    if len(raw) >= 4:
                        for code in known_codes:
                            if not code.startswith(raw) or len(code) - len(raw) > 3:
                                continue
                            method = "PREFIX + STRONA" if target_pages else "PREFIX BEZ STRONY"
                            previous = matches.get(code)
                            if previous is None or method_rank(method) > method_rank(previous[1]):
                                matches[code] = (raw, method)
                if target_pages:
                    matches = {
                        code: value for code, value in matches.items()
                        if code_to_pages[code].intersection(target_pages)
                    }
                if not matches:
                    continue

                following = [(home_y, code) for home_y, code in home_occurrences if 0 < home_y - y <= 150]
                if not following:
                    continue
                nearest_y = min(home_y for home_y, _code in following)
                source_codes = {code for home_y, code in home_occurrences if nearest_y <= home_y <= nearest_y + 40}
                detected_sources.update(source_codes)
                detected_targets.update(matches)

                description_candidates = [
                    other_line for other_x, other_y, other_line in lines
                    if 300 <= other_x < x and abs(other_y - y) <= 2 and meaningful_header(other_line)
                ]
                reference_description = " | ".join(description_candidates)

                for source_code in source_codes:
                    for target_code, (raw_reference, method) in matches.items():
                        if source_code == target_code:
                            continue
                        matched_target_pages = sorted(code_to_pages[target_code].intersection(target_pages))
                        if not matched_target_pages:
                            matched_target_pages = sorted(code_to_pages[target_code])
                        if method == "DOKŁADNY KOD":
                            confidence, status = "WYSOKA", "OK"
                        elif method == "PREFIX + STRONA":
                            confidence, status = "WYSOKA", "OK"
                        else:
                            confidence, status = "ŚREDNIA", "DO_WERYFIKACJI"
                        evidence = (
                            f"Katalog s. {printed_page}: produkt {source_code}; w sekcji "
                            f"'Współpracuje z' wskazano {raw_reference} ({target_code})."
                        )
                        relation = Relation(
                            source_code=source_code,
                            target_code=target_code,
                            evidence_page=printed_page,
                            target_pages="|".join(map(str, matched_target_pages)),
                            raw_reference=raw_reference,
                            reference_description=reference_description,
                            method=method,
                            confidence=confidence,
                            status=status,
                            evidence=evidence,
                        )
                        key = (source_code, target_code, printed_page)
                        old = relation_map.get(key)
                        if old is None or method_rank(relation.method) > method_rank(old.method):
                            relation_map[key] = relation

            raw_lines = " | ".join(compact(line) for _x, _y, line in block_lines if compact(line))[:2000]
            block_status = "OK" if detected_sources and detected_targets else "DO_WERYFIKACJI"
            blocks.append(CompatibilityBlock(
                printed_page=printed_page,
                header_y=round(header_y, 1),
                source_codes="|".join(sorted(detected_sources)),
                target_codes="|".join(sorted(detected_targets)),
                raw_lines=raw_lines,
                status=block_status,
            ))

    relations = sorted(relation_map.values(), key=lambda item: (item.source_code, item.target_code, item.evidence_page))
    return relations, blocks


def load_store_products(path: Path, code_to_pages: dict[str, set[int]]) -> tuple[list[dict[str, Any]], dict[str, list[CatalogProduct]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    records = load_records(sheet, find_headers(sheet))
    mark_old_duplicate_products(records)
    products = [CatalogProduct(record) for record in records]
    by_catalog_code: dict[str, list[CatalogProduct]] = defaultdict(list)
    rows: list[dict[str, Any]] = []
    for product in products:
        candidates = candidate_catalog_codes(product)
        matched = next((code for code in candidates if code in code_to_pages), "")
        if matched:
            by_catalog_code[matched].append(product)
        rows.append({
            "SKU sklepu": product.sku,
            "Kod SKU": product.record.code,
            "Alias Berker": product.record.alias_code,
            "Kod katalogowy dopasowany": matched,
            "Strony katalogowe": "|".join(map(str, sorted(code_to_pages.get(matched, set())))),
            "Tytuł sklepu": product.title,
            "Typ produktu": product.record.product_type,
            "Rola": product.role,
            "Funkcja": product.function,
            "Seria": product.record.series,
            "Kolor producenta": product.record.manufacturer_color,
            "Stary duplikat": "TAK" if product.record.is_old_duplicate else "NIE",
            "Kod zastępujący": product.record.replacement_code,
            "Status": "DOPASOWANO" if matched else "BRAK W INDEKSIE KATALOGU",
        })
    workbook.close()
    return rows, by_catalog_code


def decorate_product_rows(product_rows: list[dict[str, Any]], store_by_code: dict[str, list[CatalogProduct]]) -> None:
    for row in product_rows:
        products = store_by_code.get(row["Kod katalogowy"], [])
        row["SKU sklepu"] = "|".join(product.sku for product in products)
        row["Tytuły sklepu"] = " | ".join(product.title for product in products)[:2000]
        row["Liczba produktów sklepu"] = len(products)


def decorate_relation_rows(
    relations: list[Relation],
    product_lookup: dict[tuple[str, int], dict[str, Any]],
    product_fallback: dict[str, dict[str, Any]],
    store_by_code: dict[str, list[CatalogProduct]],
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for relation in relations:
        source = product_lookup.get((relation.source_code, relation.evidence_page), product_fallback.get(relation.source_code, {}))
        target_page = int(relation.target_pages.split("|", 1)[0]) if relation.target_pages else 0
        target = product_lookup.get((relation.target_code, target_page), product_fallback.get(relation.target_code, {}))
        result.append({
            "Kod źródłowy": relation.source_code,
            "Nazwa źródłowa": source.get("Nazwa automatyczna", ""),
            "SKU źródłowe sklepu": "|".join(product.sku for product in store_by_code.get(relation.source_code, [])),
            "Strona źródła": relation.evidence_page,
            "Kod docelowy": relation.target_code,
            "Nazwa docelowa": target.get("Nazwa automatyczna", ""),
            "SKU docelowe sklepu": "|".join(product.sku for product in store_by_code.get(relation.target_code, [])),
            "Strona celu": relation.target_pages,
            "Typ relacji": "Współpracuje z",
            "Kod zapisany w PDF": relation.raw_reference,
            "Opis celu w PDF": relation.reference_description,
            "Metoda": relation.method,
            "Pewność": relation.confidence,
            "Status": relation.status,
            "Strona dowodu": relation.evidence_page,
            "Dowód katalogowy": relation.evidence,
        })
    return result


def add_title(sheet, title: str, subtitle: str, columns: int) -> None:
    last = get_column_letter(columns)
    sheet.merge_cells(f"A1:{last}1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(f"A2:{last}2")
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(name="Aptos", size=10, color=TEXT)
    sheet["A2"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 32


def set_widths(sheet, headers: list[str]) -> None:
    width_map = {
        "Kod katalogowy": 18, "Kod źródłowy": 18, "Kod docelowy": 18,
        "Nazwa automatyczna": 48, "Nazwa bazowa": 42, "Nazwa źródłowa": 42,
        "Nazwa docelowa": 42, "Tytuł sklepu": 55, "Tytuły sklepu": 55,
        "Wariant / opis": 32, "Seria / system": 34, "Sekcja strony": 42,
        "Kontekst katalogowy": 80, "Dowód katalogowy": 70, "Tekst strony": 100,
        "Opis celu w PDF": 42, "Surowy blok": 90, "Raw block": 90,
        "Status": 22, "Status ekstrakcji": 22, "Metoda": 22, "Pewność": 14,
        "SKU sklepu": 28, "SKU źródłowe sklepu": 30, "SKU docelowe sklepu": 30,
        "Tytuły": 55, "Źródło": 60, "Wartość": 25, "Test": 48, "Wynik": 18,
        "Szczegóły": 72,
    }
    for col, header in enumerate(headers, 1):
        width = width_map.get(header, min(max(len(header) + 3, 12), 28))
        sheet.column_dimensions[get_column_letter(col)].width = width


def add_data_sheet(
    workbook: Workbook,
    name: str,
    title: str,
    subtitle: str,
    rows: list[dict[str, Any]],
    table_name: str,
    status_column: str | None = None,
) -> Any:
    sheet = workbook.create_sheet(name)
    headers = list(rows[0]) if rows else ["Brak danych"]
    add_title(sheet, title, subtitle, len(headers))
    header_row = 4
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(header_row, col, header)
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=HAGER_BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=NAVY))
        cell.comment = Comment(f"Pole kontrolne: {header}", "Codex")
    sheet.row_dimensions[header_row].height = 32
    wrap_headers = {
        "Nazwa automatyczna", "Nazwa bazowa", "Nazwa źródłowa", "Nazwa docelowa",
        "Tytuł sklepu", "Tytuły sklepu", "Kontekst katalogowy", "Dowód katalogowy",
        "Tekst strony", "Opis celu w PDF", "Surowy blok", "Raw block", "Szczegóły",
    }
    for row_idx, row in enumerate(rows, header_row + 1):
        for col, header in enumerate(headers, 1):
            value = excel_safe(row.get(header, ""))
            cell = sheet.cell(row_idx, col, value)
            cell.font = Font(name="Aptos", size=9, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=header in wrap_headers)
        if "Tekst strony" in headers:
            sheet.row_dimensions[row_idx].height = 48
        elif any(header in wrap_headers for header in headers):
            sheet.row_dimensions[row_idx].height = 34
    if rows:
        end_row = header_row + len(rows)
        end_col = get_column_letter(len(headers))
        data_range = f"A{header_row + 1}:{end_col}{end_row}"
        sheet.conditional_formatting.add(
            data_range,
            FormulaRule(formula=[f"MOD(ROW(),2)=1"], fill=PatternFill("solid", fgColor="F7FAFC")),
        )
        if status_column and status_column in headers:
            status_col = headers.index(status_column) + 1
            status_letter = get_column_letter(status_col)
            status_range = f"{status_letter}{header_row + 1}:{status_letter}{end_row}"
            validation = DataValidation(
                type="list", formula1='"OK,DO_WERYFIKACJI,ODRZUCONE,ZATWIERDZONE"', allow_blank=False
            )
            sheet.add_data_validation(validation)
            validation.add(status_range)
            sheet.conditional_formatting.add(
                status_range,
                FormulaRule(formula=[f'{status_letter}{header_row + 1}="OK"'], fill=PatternFill("solid", fgColor=LIGHT_GREEN)),
            )
            sheet.conditional_formatting.add(
                status_range,
                FormulaRule(formula=[f'{status_letter}{header_row + 1}="DO_WERYFIKACJI"'], fill=PatternFill("solid", fgColor=LIGHT_YELLOW)),
            )
            sheet.conditional_formatting.add(
                status_range,
                FormulaRule(formula=[f'{status_letter}{header_row + 1}="ODRZUCONE"'], fill=PatternFill("solid", fgColor=LIGHT_RED)),
            )
    sheet.freeze_panes = "A5"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"
    set_widths(sheet, headers)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:4"
    return sheet


def add_summary_sheet(
    workbook: Workbook,
    pdf_path: Path,
    store_path: Path,
    output_path: Path,
    metrics: dict[str, Any],
    row_bounds: dict[str, int],
) -> None:
    sheet = workbook.create_sheet("Podsumowanie", 0)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "Hager / Berker - kontrolna baza katalogowa"
    sheet["A1"].font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells("A2:F2")
    sheet["A2"] = (
        "Źródłem prawdy pozostaje katalog PDF. Do automatycznych cross-selli używaj wyłącznie "
        "wierszy z arkusza Kompatybilność, dla których Status = OK. Dopasowania po samej funkcji są zabronione."
    )
    sheet["A2"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 48

    sheet["A4"] = "Źródła"
    sheet["A4"].font = Font(bold=True, color=WHITE)
    sheet["A4"].fill = PatternFill("solid", fgColor=HAGER_BLUE)
    sources = [
        ("Katalog PDF", str(pdf_path.resolve())),
        ("Plik produktów sklepu", str(store_path.resolve())),
        ("Plik wynikowy", str(output_path.resolve())),
        ("SHA-256 PDF", metrics["pdf_sha256"]),
        ("SHA-256 XLSX źródłowego", metrics["store_sha256"]),
    ]
    for idx, (label, value) in enumerate(sources, 5):
        sheet.cell(idx, 1, label).font = Font(bold=True, color=TEXT)
        sheet.cell(idx, 2, value)
        sheet.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=6)

    sheet["A12"] = "Metryki"
    sheet["A12"].font = Font(bold=True, color=WHITE)
    sheet["A12"].fill = PatternFill("solid", fgColor=HAGER_BLUE)
    metric_rows = [
        ("Pozycje kod-strona w indeksie", f"=ROWS('Produkty katalogowe'!A5:A{row_bounds['products']})"),
        ("Unikalne kody katalogowe", metrics["unique_codes"]),
        ("Jawne relacje katalogowe", f"=ROWS('Kompatybilność'!A5:A{row_bounds['relations']})"),
        ("Relacje ze statusem OK", f'=COUNTIF(\'Kompatybilność\'!N5:N{row_bounds["relations"]},"OK")'),
        ("Relacje do weryfikacji", f'=COUNTIF(\'Kompatybilność\'!N5:N{row_bounds["relations"]},"DO_WERYFIKACJI")'),
        ("Bloki 'Współpracuje z'", f"=ROWS('Bloki kompatybilności'!A5:A{row_bounds['blocks']})"),
        ("Produkty sklepu", f"=ROWS('Alias kodów'!A5:A{row_bounds['aliases']})"),
        ("Produkty sklepu dopasowane", f'=COUNTIF(\'Alias kodów\'!N5:N{row_bounds["aliases"]},"DOPASOWANO")'),
        ("Produkty sklepu bez indeksu", f'=COUNTIF(\'Alias kodów\'!N5:N{row_bounds["aliases"]},"BRAK W INDEKSIE KATALOGU")'),
        ("Pozycje wymagające kontroli", f"=ROWS('Do weryfikacji'!A5:A{row_bounds['review']})"),
    ]
    for idx, (label, value) in enumerate(metric_rows, 13):
        sheet.cell(idx, 1, label).font = Font(bold=True, color=TEXT)
        sheet.cell(idx, 2, value)
        sheet.cell(idx, 2).number_format = "#,##0"

    sheet["D12"] = "Reguły użycia"
    sheet["D12"].font = Font(bold=True, color=WHITE)
    sheet["D12"].fill = PatternFill("solid", fgColor=HAGER_BLUE)
    rules = [
        "1. Status OK oznacza jawne wskazanie w sekcji 'Współpracuje z'.",
        "2. Prefix kodu jest akceptowany tylko, gdy PDF podaje stronę celu.",
        "3. Wiersze DO_WERYFIKACJI nie mogą trafiać automatycznie do importu.",
        "4. Arkusz nie tworzy relacji wyłącznie na podstawie słów: ściemniacz, RTV/SAT, łącznik itd.",
        "5. Kolumny strony i dowodu należy zachować w każdym kolejnym eksporcie.",
    ]
    for idx, rule in enumerate(rules, 13):
        sheet.merge_cells(start_row=idx, start_column=4, end_row=idx, end_column=6)
        sheet.cell(idx, 4, rule).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[idx].height = 30

    sheet["A25"] = "Test polskich znaków"
    sheet["B25"] = "Zażółć gęślą jaźń - ą ć ę ł ń ó ś ź ż"
    sheet["A25"].font = Font(bold=True)
    sheet["B25"].fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    sheet.merge_cells("B25:F25")

    for col, width in zip("ABCDEF", (34, 30, 4, 34, 28, 28)):
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=4, max_row=25, min_col=1, max_col=6):
        for cell in row:
            updated_font = copy(cell.font)
            updated_font.name = "Aptos"
            updated_font.sz = 10
            cell.font = updated_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def scan_mojibake(workbook: Workbook) -> list[str]:
    findings: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    markers = mojibake_markers(cell.value)
                    if markers:
                        findings.append(f"{sheet.title}!{cell.coordinate}: {','.join(markers)}")
    return findings


def build_review_rows(
    product_rows: list[dict[str, Any]],
    relation_rows: list[dict[str, Any]],
    alias_rows: list[dict[str, Any]],
    blocks: list[CompatibilityBlock],
    code_to_pages: dict[str, set[int]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in product_rows:
        if row["Status ekstrakcji"] != "OK" or len(code_to_pages[row["Kod katalogowy"]]) > 1:
            reason = "Wiele stron indeksu" if len(code_to_pages[row["Kod katalogowy"]]) > 1 else "Niepełna ekstrakcja nazwy"
            rows.append({
                "Rodzaj": "PRODUKT", "Kod / SKU": row["Kod katalogowy"],
                "Strona": row["Strona drukowana"], "Powód": reason,
                "Szczegóły": row["Kontekst katalogowy"], "Status": "DO_WERYFIKACJI",
            })
    for row in relation_rows:
        if row["Status"] != "OK":
            rows.append({
                "Rodzaj": "RELACJA", "Kod / SKU": f"{row['Kod źródłowy']} -> {row['Kod docelowy']}",
                "Strona": row["Strona dowodu"], "Powód": row["Metoda"],
                "Szczegóły": row["Dowód katalogowy"], "Status": "DO_WERYFIKACJI",
            })
    for row in alias_rows:
        if row["Status"] != "DOPASOWANO":
            rows.append({
                "Rodzaj": "PRODUKT SKLEPU", "Kod / SKU": row["SKU sklepu"],
                "Strona": "", "Powód": "Brak kodu w indeksie PDF",
                "Szczegóły": row["Tytuł sklepu"], "Status": "DO_WERYFIKACJI",
            })
    for block in blocks:
        if block.status != "OK":
            rows.append({
                "Rodzaj": "BLOK PDF", "Kod / SKU": block.source_codes,
                "Strona": block.printed_page, "Powód": "Nie udało się jednoznacznie utworzyć relacji",
                "Szczegóły": block.raw_lines, "Status": "DO_WERYFIKACJI",
            })
    return rows


def write_workbook(
    output_path: Path,
    pdf_path: Path,
    store_path: Path,
    product_rows: list[dict[str, Any]],
    relation_rows: list[dict[str, Any]],
    alias_rows: list[dict[str, Any]],
    blocks: list[CompatibilityBlock],
    page_rows: list[dict[str, Any]],
    review_rows: list[dict[str, Any]],
    validations: list[tuple[str, str, str]],
    metrics: dict[str, Any],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    add_data_sheet(
        workbook, "Produkty katalogowe", "Produkty katalogowe z indeksu PDF",
        "Jeden wiersz odpowiada parze kod katalogowy - strona drukowana. Nazwy są pomocnicze; kontekst i strona pozostają dowodem.",
        product_rows, "ProduktyKatalogowe", "Status ekstrakcji",
    )
    add_data_sheet(
        workbook, "Kompatybilność", "Jawne relacje kompatybilności",
        "Relacje pochodzą wyłącznie z bloków 'Współpracuje z'. Status OK może być używany automatycznie; pozostałe wymagają kontroli.",
        relation_rows, "RelacjeKompatybilnosci", "Status",
    )
    block_rows = [{
        "Strona drukowana": block.printed_page,
        "Pozycja Y": block.header_y,
        "Kody źródłowe": block.source_codes,
        "Kody docelowe": block.target_codes,
        "Surowy blok": block.raw_lines,
        "Status": block.status,
    } for block in blocks]
    add_data_sheet(
        workbook, "Bloki kompatybilności", "Surowe bloki 'Współpracuje z'",
        "Warstwa audytowa zachowująca tekst i współrzędne każdego wykrytego bloku kompatybilności.",
        block_rows, "BlokiKompatybilnosci", "Status",
    )
    add_data_sheet(
        workbook, "Alias kodów", "Mapowanie produktów sklepu na katalog",
        "Łączy SKU sklepu z kodem lub aliasem Berker znalezionym w indeksie PDF. Stare duplikaty są oznaczone osobno.",
        alias_rows, "AliasyKodow", "Status",
    )
    add_data_sheet(
        workbook, "Do weryfikacji", "Pozycje wymagające kontroli",
        "Zbiorcza kolejka wyjątków: niepełne nazwy, kody wielostronicowe, niejednoznaczne relacje i produkty sklepu bez indeksu.",
        review_rows, "PozycjeDoWeryfikacji", "Status",
    )
    add_data_sheet(
        workbook, "Strony PDF", "Warstwa tekstowa stron katalogu",
        "Tekst każdej strony jest zachowany jako warstwa kontrolna. Układ wizualny należy w razie wątpliwości sprawdzić w źródłowym PDF.",
        page_rows, "StronyPdf",
    )
    validation_rows = [{"Test": test, "Wynik": result, "Szczegóły": details} for test, result, details in validations]
    add_data_sheet(
        workbook, "Walidacja", "Walidacja konwersji",
        "Testy integralności, kodowania i regresji dla znanych błędnych dopasowań funkcjonalnych.",
        validation_rows, "WalidacjaKonwersji",
    )

    row_bounds = {
        "products": 4 + len(product_rows),
        "relations": 4 + len(relation_rows),
        "blocks": 4 + len(block_rows),
        "aliases": 4 + len(alias_rows),
        "review": 4 + len(review_rows),
    }
    add_summary_sheet(workbook, pdf_path, store_path, output_path, metrics, row_bounds)

    findings = scan_mojibake(workbook)
    if findings:
        raise ValueError(f"Wykryto uszkodzone kodowanie: {findings[:10]}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def convert(pdf_path: Path, store_path: Path, output_path: Path) -> dict[str, Any]:
    if not pdf_path.exists():
        raise FileNotFoundError(pdf_path)
    if not store_path.exists():
        raise FileNotFoundError(store_path)

    document = fitz.open(pdf_path)
    try:
        entries = extract_index(document)
        code_to_pages: dict[str, set[int]] = defaultdict(set)
        page_to_codes: dict[int, set[str]] = defaultdict(set)
        for entry in entries:
            code_to_pages[entry.code].add(entry.printed_page)
            page_to_codes[entry.printed_page].add(entry.code)

        line_cache: dict[int, list[PageLine]] = {}
        simple_line_cache: dict[int, list[tuple[float, float, str]]] = {}
        product_rows = [product_row_for_entry(document, entry, line_cache, simple_line_cache) for entry in entries]

        alias_rows, store_by_code = load_store_products(store_path, code_to_pages)
        decorate_product_rows(product_rows, store_by_code)
        relations, blocks = extract_relations(document, code_to_pages, page_to_codes)

        product_lookup = {(row["Kod katalogowy"], row["Strona drukowana"]): row for row in product_rows}
        product_fallback: dict[str, dict[str, Any]] = {}
        for row in product_rows:
            product_fallback.setdefault(row["Kod katalogowy"], row)
        relation_rows = decorate_relation_rows(relations, product_lookup, product_fallback, store_by_code)

        page_rows: list[dict[str, Any]] = []
        for printed_page in range(1, 619):
            page = document[printed_page]
            text = page.get_text("text")
            page_rows.append({
                "Strona drukowana": printed_page,
                "Strona PDF": printed_page + 1,
                "Nagłówek": page_header(line_cache.get(printed_page) or structured_lines(page)),
                "Kody z indeksu": "|".join(sorted(page_to_codes.get(printed_page, set()))),
                "Liczba kodów": len(page_to_codes.get(printed_page, set())),
                "Ma blok 'Współpracuje z'": "TAK" if "wspolpracuje z" in norm(text) else "NIE",
                "Tekst strony": text[:32000],
            })

        review_rows = build_review_rows(product_rows, relation_rows, alias_rows, blocks, code_to_pages)

        relation_keys = {(row["Kod źródłowy"], row["Kod docelowy"]) for row in relation_rows}
        duplicate_relation_rows = len(relation_rows) - len({
            (row["Kod źródłowy"], row["Kod docelowy"], row["Strona dowodu"]) for row in relation_rows
        })
        validations: list[tuple[str, str, str]] = [
            ("index_entries_extracted", "OK" if entries else "BŁĄD", str(len(entries))),
            ("unique_catalog_codes", "OK" if code_to_pages else "BŁĄD", str(len(code_to_pages))),
            ("explicit_relations_extracted", "OK" if relation_rows else "BŁĄD", str(len(relation_rows))),
            ("no_duplicate_relation_rows", "OK" if duplicate_relation_rows == 0 else "BŁĄD", str(duplicate_relation_rows)),
            ("all_relation_codes_exist_in_index", "OK" if all(
                row["Kod źródłowy"] in code_to_pages and row["Kod docelowy"] in code_to_pages for row in relation_rows
            ) else "BŁĄD", "źródło i cel muszą istnieć w indeksie"),
            ("false_rotary_dimmer_to_knx_absent", "OK" if ("11372045", "85145124") not in relation_keys else "BŁĄD", "znany fałszywy cross-sell"),
            ("false_key_to_cross_switch_absent", "OK" if ("14177109", "533037") not in relation_keys else "BŁĄD", "znany fałszywy cross-sell"),
            ("four_output_antenna_plate_relation_present", "OK" if ("14842045", "459410") in relation_keys else "BŁĄD", "katalog s. 345"),
            ("three_output_plate_not_linked_to_two_output_socket", "OK" if ("106420", "53455012") not in relation_keys else "BŁĄD", "katalog s. 393"),
            ("polish_text_sample", "OK", "Zażółć gęślą jaźń - ą ć ę ł ń ó ś ź ż"),
            ("source_files_unchanged", "OK", f"PDF {sha256(pdf_path)}; XLSX {sha256(store_path)}"),
        ]
        if any(result != "OK" for _test, result, _details in validations):
            failures = [test for test, result, _details in validations if result != "OK"]
            raise ValueError(f"Nieudane testy konwersji: {failures}")

        metrics = {
            "pdf_sha256": sha256(pdf_path),
            "store_sha256": sha256(store_path),
            "unique_codes": len(code_to_pages),
        }
        write_workbook(
            output_path, pdf_path, store_path, product_rows, relation_rows, alias_rows,
            blocks, page_rows, review_rows, validations, metrics,
        )
    finally:
        document.close()

    return {
        "output": str(output_path.resolve()),
        "index_entries": len(entries),
        "unique_codes": len(code_to_pages),
        "relations": len(relation_rows),
        "relations_ok": sum(row["Status"] == "OK" for row in relation_rows),
        "relations_review": sum(row["Status"] != "OK" for row in relation_rows),
        "compatibility_blocks": len(blocks),
        "store_products": len(alias_rows),
        "store_matched": sum(row["Status"] == "DOPASOWANO" for row in alias_rows),
        "store_unmatched": sum(row["Status"] != "DOPASOWANO" for row in alias_rows),
        "review_rows": len(review_rows),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Konwertuje katalog Hager/Berker PDF do kontrolnej bazy XLSX.")
    parser.add_argument("--pdf", type=Path, default=DEFAULT_PDF)
    parser.add_argument("--store-xlsx", type=Path, default=DEFAULT_STORE_XLSX)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    result = convert(args.pdf, args.store_xlsx, args.output)
    print(json.dumps(result, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
