from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XML = ROOT / "input" / "BMEcat_PL 17.04.2026.xml"
DEFAULT_INPUT = ROOT / "output" / "BMEcat_Hager_PL_produkty_ETIM_PL_ZDJECIA.xlsx"
DEFAULT_OUTPUT = (
    ROOT / "output" / "BMEcat_Hager_PL_produkty_ETIM_PL_ZDJECIA_RELACJE.xlsx"
)
DEFAULT_REPORT = (
    ROOT / "reports" / "BMEcat_Hager_PL_produkty_ETIM_PL_ZDJECIA_RELACJE_report.json"
)

SHEET_NAME = "Relacje produktowe"
RELATION_TYPES = (
    ("accessories", "Akcesoria"),
    ("mandatory", "Elementy obowiązkowe"),
    ("sparepart", "Części zamienne"),
    ("followup", "Następcy"),
    ("similar", "Produkty podobne"),
    ("others", "Inne relacje"),
)
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
WARNING = "FFF2CC"


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def logical_workbook_hash(path: Path, excluded_sheet: str | None = None) -> str:
    digest = hashlib.sha256()
    workbook = load_workbook(path, read_only=True, data_only=False)
    for sheet in workbook.worksheets:
        if sheet.title == excluded_sheet:
            continue
        digest.update(sheet.title.encode("utf-8"))
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                digest.update(repr(value).encode("utf-8"))
                digest.update(b"\x1f")
            digest.update(b"\x1e")
    workbook.close()
    return digest.hexdigest()


def extract_relations(
    xml_path: Path,
) -> tuple[dict[str, dict[str, list[str]]], Counter, int]:
    relations: dict[str, dict[str, list[str]]] = {}
    totals: Counter = Counter()
    duplicate_entries = 0

    for _event, product in ET.iterparse(xml_path, events=("end",)):
        if local_name(product.tag) != "PRODUCT":
            continue

        product_code = ""
        for child in product:
            if local_name(child.tag) == "SUPPLIER_PID":
                product_code = (child.text or "").strip()
                break

        grouped: dict[str, list[str]] = defaultdict(list)
        for reference in product.iter():
            if local_name(reference.tag) != "PRODUCT_REFERENCE":
                continue
            relation_type = reference.attrib.get("type", "")
            target = ""
            for child in reference:
                if local_name(child.tag) == "PROD_ID_TO":
                    target = (child.text or "").strip()
                    break
            if relation_type and target:
                grouped[relation_type].append(target)
                totals[relation_type] += 1

        duplicate_entries += sum(
            len(targets) - len(set(targets)) for targets in grouped.values()
        )
        if product_code in relations:
            raise ValueError(f"Duplicate SUPPLIER_PID in XML: {product_code}")
        relations[product_code] = dict(grouped)
        product.clear()

    return relations, totals, duplicate_entries


def formula_count(workbook_path: Path) -> int:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    count = 0
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            count += sum(
                1 for value in row if isinstance(value, str) and value.startswith("=")
            )
    workbook.close()
    return count


def build_workbook(
    input_path: Path, xml_path: Path, output_path: Path, report_path: Path
) -> dict:
    input_hash = sha256(input_path)
    xml_hash = sha256(xml_path)
    input_logical_hash = logical_workbook_hash(input_path)
    formulas_before = formula_count(input_path)
    relations, relation_totals, duplicate_entries = extract_relations(xml_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(input_path, output_path)
    workbook = load_workbook(output_path, data_only=False)
    products_sheet = workbook["Produkty"]
    if SHEET_NAME in workbook.sheetnames:
        del workbook[SHEET_NAME]
    sheet = workbook.create_sheet(SHEET_NAME)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C2"

    headers = ["Kod produktu", "Nazwa"]
    for _relation_type, label in RELATION_TYPES:
        headers.extend((f"Liczba: {label}", label))
    headers.extend(("Łączna liczba relacji", "Status źródła"))
    sheet.append(headers)

    product_codes: list[str] = []
    product_names: dict[str, str] = {}
    for row in products_sheet.iter_rows(
        min_row=2, min_col=1, max_col=3, values_only=True
    ):
        code = str(row[0] or "").strip()
        product_codes.append(code)
        product_names[code] = str(row[2] or "").strip()

    product_code_set = set(product_codes)
    xml_code_set = set(relations)
    missing_xml_codes = [code for code in product_codes if code not in xml_code_set]
    extra_xml_codes = sorted(xml_code_set - product_code_set)
    external_targets = Counter()
    max_cell_length = 0
    products_with_relations = 0
    rows_with_duplicates = 0

    for code in product_codes:
        grouped = relations.get(code, {})
        row_values: list[object] = [code, product_names[code]]
        total_for_product = 0
        has_duplicates = False
        for relation_type, _label in RELATION_TYPES:
            targets = grouped.get(relation_type, [])
            total_for_product += len(targets)
            if len(targets) != len(set(targets)):
                has_duplicates = True
            for target in targets:
                if target not in product_code_set:
                    external_targets[relation_type] += 1
            joined = " | ".join(targets)
            max_cell_length = max(max_cell_length, len(joined))
            row_values.extend((len(targets), joined or None))
        if total_for_product:
            products_with_relations += 1
        if has_duplicates:
            rows_with_duplicates += 1
        row_values.extend(
            (
                total_for_product,
                "DUPLIKATY W XML" if has_duplicates else "OK",
            )
        )
        sheet.append(row_values)

    if max_cell_length > 32767:
        raise ValueError(f"Relation cell exceeds Excel limit: {max_cell_length}")

    header_fill = PatternFill("solid", fgColor=BLUE)
    for cell in sheet[1]:
        cell.fill = copy(header_fill)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34

    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 46
    for column in ("C", "E", "G", "I", "K", "M", "O"):
        sheet.column_dimensions[column].width = 15
    for column in ("D", "F", "H", "J", "L", "N"):
        sheet.column_dimensions[column].width = 58
    sheet.column_dimensions["P"].width = 18

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        row[0].number_format = "@"
        row[0].alignment = Alignment(vertical="top")
        row[1].alignment = Alignment(vertical="top")
        for index in (2, 4, 6, 8, 10, 12, 14):
            row[index].number_format = "#,##0"
            row[index].alignment = Alignment(horizontal="right", vertical="top")
        for index in (3, 5, 7, 9, 11, 13):
            row[index].number_format = "@"
            row[index].alignment = Alignment(vertical="top", wrap_text=False)
        row[15].alignment = Alignment(vertical="top")
        if row[15].value != "OK":
            row[15].fill = PatternFill("solid", fgColor=WARNING)

    sheet.auto_filter.ref = f"A1:P{sheet.max_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"

    workbook.save(output_path)
    workbook.close()

    output_logical_hash = logical_workbook_hash(output_path, excluded_sheet=SHEET_NAME)
    formulas_after = formula_count(output_path)
    validation_book = load_workbook(output_path, read_only=True, data_only=False)
    relation_sheet = validation_book[SHEET_NAME]
    relation_sums = Counter()
    target_codes_seen = set()
    for row in relation_sheet.iter_rows(min_row=2, values_only=True):
        for index, (relation_type, _label) in enumerate(RELATION_TYPES):
            count_value = int(row[2 + index * 2] or 0)
            relation_sums[relation_type] += count_value
            list_value = str(row[3 + index * 2] or "")
            if list_value:
                split_targets = list_value.split(" | ")
                if len(split_targets) != count_value:
                    raise AssertionError(
                        f"Count/list mismatch for {row[0]} {relation_type}"
                    )
                target_codes_seen.update(split_targets)
    validation = {
        "sheet_rows": relation_sheet.max_row,
        "sheet_columns": relation_sheet.max_column,
        "product_order_matches": [
            row[0]
            for row in relation_sheet.iter_rows(
                min_row=2, min_col=1, max_col=1, values_only=True
            )
        ]
        == product_codes,
        "relation_totals_match": relation_sums == relation_totals,
        "all_targets_exist_in_products": target_codes_seen <= product_code_set,
        "original_sheets_unchanged": input_logical_hash == output_logical_hash,
        "formulas_unchanged": formulas_before == formulas_after,
        "max_relation_cell_length": max_cell_length,
    }
    validation["passed"] = all(
        (
            validation["sheet_rows"] == 14543,
            validation["sheet_columns"] == 16,
            validation["product_order_matches"],
            validation["relation_totals_match"],
            validation["all_targets_exist_in_products"],
            validation["original_sheets_unchanged"],
            validation["formulas_unchanged"],
            validation["max_relation_cell_length"] <= 32767,
            not missing_xml_codes,
            not extra_xml_codes,
            not external_targets,
        )
    )
    validation_book.close()
    if not validation["passed"]:
        raise AssertionError(json.dumps(validation, ensure_ascii=False, indent=2))

    metrics = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "input_file": str(input_path),
        "input_sha256": input_hash,
        "xml_file": str(xml_path),
        "xml_sha256": xml_hash,
        "output_file": str(output_path),
        "products": len(product_codes),
        "products_with_relations": products_with_relations,
        "products_without_relations": len(product_codes) - products_with_relations,
        "relation_totals": dict(relation_totals),
        "total_relations": sum(relation_totals.values()),
        "duplicate_source_entries": duplicate_entries,
        "products_with_duplicate_source_entries": rows_with_duplicates,
        "missing_xml_codes": missing_xml_codes,
        "extra_xml_codes": extra_xml_codes,
        "external_targets": dict(external_targets),
        "formula_count_before": formulas_before,
        "formula_count_after": formulas_after,
        "validation": validation,
        "output_sha256": sha256(output_path),
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return metrics


def main() -> None:
    parser = argparse.ArgumentParser(description="Add Hager BMEcat product relations to XLSX")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--xml", type=Path, default=DEFAULT_XML)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    metrics = build_workbook(args.input, args.xml, args.output, args.report)
    print(json.dumps(metrics, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
