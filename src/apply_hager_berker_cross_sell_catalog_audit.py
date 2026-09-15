from __future__ import annotations

import argparse
import hashlib
import json
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_INPUT = Path(
    "output/Hager-Berker-Gniazdka_WOO_ALL_IMPORT_CROSS_UP_SELL_KATALOG_UPSELL_ROZSZERZONY.xlsx"
)
DEFAULT_CATALOG = Path("output/Hager-Berker-Katalog-PDF_BAZA_KOMPATYBILNOSCI.xlsx")
DEFAULT_OUTPUT = Path(
    "output/Hager-Berker-Gniazdka_WOO_ALL_IMPORT_CROSS_UP_SELL_PO_AUDYCIE_PDF.xlsx"
)

MAX_CROSS_SELLS = 4

# Pełny kod jest prawidłowy, ale te dwa wpisy nie powinny trafić do bieżącego importu.
# WDS1500 i WDS1700 są alternatywnymi mechanizmami, a WDE3032 ma już cztery trafne pozycje.
EXCLUDED_EXACT_PAIRS = {
    ("WDS1700/HAG", "WDS1500/HAG"),
    ("WDE3032/HAG", "WDE160020/HAG"),
}

# Błędne albo alternatywne produkty, które nie są cross-sellem.
INVALID_CURRENT_PAIRS = {
    ("11372045/HAG", "85145124/HAG"),
    ("14177109/HAG", "WDEM3037/HAG"),
    ("WTF4593PW/HAG", "WDF459212/HAG"),
    ("WTF4593BG/HAG", "WDF459212/HAG"),
    ("WDF459212/HAG", "WDF459314/HAG"),
    ("WDF459314/HAG", "WDF459212/HAG"),
    ("WDF459212/HAG", "WDF459316/HAG"),
    ("WDF459316/HAG", "WDF459212/HAG"),
    ("2909/HAG", "296110/HAG"),
    ("296110/HAG", "2909/HAG"),
    ("2973/HAG", "296110/HAG"),
    ("296110/HAG", "2973/HAG"),
    ("385103/HAG", "WDE385203/HAG"),
    ("WDE385203/HAG", "385103/HAG"),
    ("10152204/HAG", "WDA1101/HAG"),
    ("10152204/HAG", "WDA1102/HAG"),
}

# Zwolnienie pojedynczego miejsca na dokładniej potwierdzony element katalogowy.
SLOT_REPLACEMENTS = {
    ("12036084/HAG", "WDF459315/HAG"),
    ("12036086/HAG", "WDF459315/HAG"),
    ("12036089/HAG", "WDF459315/HAG"),
    ("WDE382610/HAG", "15078989/HAG"),
    ("WDE503501/HAG", "WDE1687W/HAG"),
    ("WDEM3035/HAG", "WDE1687W/HAG"),
    ("WDEM5035/HAG", "WDE1687W/HAG"),
    ("WDEM503503/HAG", "WDE1687W/HAG"),
    ("WDN9021/HAG", "WTN4011BG/HAG"),
}

MOJIBAKE_MARKERS = (
    "\ufffd",
    "Ă…â€š",
    "Ă„â€¦",
    "Ă…â€ş",
    "ĂÂł",
    "Ă˘â‚¬â€ť",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def split_skus(value: Any) -> list[str]:
    return [part.strip() for part in str(value or "").split("|") if part.strip()]


def find_header_map(sheet, row: int) -> dict[str, int]:
    return {
        str(cell.value): cell.column
        for cell in sheet[row]
        if cell.value not in (None, "")
    }


def load_old_duplicate_flags(catalog_path: Path) -> dict[str, bool]:
    workbook = load_workbook(catalog_path, read_only=True, data_only=True)
    sheet = workbook["Alias kodów"]
    headers = find_header_map(sheet, 4)
    result: dict[str, bool] = {}
    for row in sheet.iter_rows(min_row=5, values_only=True):
        sku = str(row[headers["SKU sklepu"] - 1] or "").strip()
        if sku:
            result[sku] = str(row[headers["Stary duplikat"] - 1] or "").upper() == "TAK"
    workbook.close()
    return result


def collect_exact_additions(
    catalog_path: Path,
    store_skus: set[str],
    current_pairs: set[tuple[str, str]],
    old_duplicate: dict[str, bool],
) -> tuple[set[tuple[str, str]], dict[tuple[str, str], tuple[int, str, str]]]:
    workbook = load_workbook(catalog_path, read_only=True, data_only=True)
    sheet = workbook["Kompatybilność"]
    headers = find_header_map(sheet, 4)
    candidates: set[tuple[str, str]] = set()
    evidence: dict[tuple[str, str], tuple[int, str, str]] = {}
    for row in sheet.iter_rows(min_row=5, values_only=True):
        get = lambda name: row[headers[name] - 1]
        if get("Status") != "OK" or get("Metoda") != "DOKŁADNY KOD":
            continue
        for source_sku in split_skus(get("SKU źródłowe sklepu")):
            for target_sku in split_skus(get("SKU docelowe sklepu")):
                pair = (source_sku, target_sku)
                if (
                    source_sku not in store_skus
                    or target_sku not in store_skus
                    or source_sku == target_sku
                    or old_duplicate.get(source_sku, False)
                    or old_duplicate.get(target_sku, False)
                    or pair in current_pairs
                    or pair in EXCLUDED_EXACT_PAIRS
                ):
                    continue
                candidates.add(pair)
                evidence[pair] = (
                    int(get("Strona dowodu")),
                    str(get("Kod źródłowy")),
                    str(get("Kod docelowy")),
                )
    workbook.close()
    if len(candidates) != 65:
        raise ValueError(f"Oczekiwano 65 dokładnych dodatków, otrzymano {len(candidates)}")
    return candidates, evidence


def scan_mojibake(workbook) -> list[tuple[str, str, str]]:
    bad: list[tuple[str, str, str]] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(marker in cell.value for marker in MOJIBAKE_MARKERS):
                    bad.append((sheet.title, cell.coordinate, cell.value[:120]))
    return bad


def build_output(input_path: Path, catalog_path: Path, output_path: Path) -> dict[str, Any]:
    input_hash_before = sha256(input_path)
    catalog_hash_before = sha256(catalog_path)
    workbook = load_workbook(input_path)
    if workbook.sheetnames != ["Worksheet"]:
        raise ValueError(f"Nieoczekiwane arkusze: {workbook.sheetnames}")
    sheet = workbook["Worksheet"]
    headers = find_header_map(sheet, 1)
    required = {"SKU", "Cross-Sells"}
    if not required.issubset(headers):
        raise ValueError(f"Brak kolumn: {sorted(required - set(headers))}")

    sku_col = headers["SKU"]
    cross_col = headers["Cross-Sells"]
    row_by_sku: dict[str, int] = {}
    current_by_sku: dict[str, list[str]] = {}
    for row_idx in range(2, sheet.max_row + 1):
        sku = str(sheet.cell(row_idx, sku_col).value or "").strip()
        if not sku:
            continue
        if sku in row_by_sku:
            raise ValueError(f"Powtórzone SKU w imporcie: {sku}")
        row_by_sku[sku] = row_idx
        values = split_skus(sheet.cell(row_idx, cross_col).value)
        if len(values) != len(set(values)):
            raise ValueError(f"Powtórzony cross-sell przed zmianą dla {sku}")
        current_by_sku[sku] = values

    store_skus = set(row_by_sku)
    current_pairs = {(source, target) for source, targets in current_by_sku.items() for target in targets}
    old_duplicate = load_old_duplicate_flags(catalog_path)
    additions, evidence = collect_exact_additions(
        catalog_path, store_skus, current_pairs, old_duplicate
    )

    removals = INVALID_CURRENT_PAIRS | SLOT_REPLACEMENTS
    missing_removals = sorted(pair for pair in removals if pair not in current_pairs)
    if missing_removals:
        raise ValueError(f"Brakuje oczekiwanych relacji do usunięcia: {missing_removals}")

    final_by_sku = {sku: list(targets) for sku, targets in current_by_sku.items()}
    for source, target in sorted(removals):
        final_by_sku[source].remove(target)

    additions_by_source: dict[str, list[str]] = defaultdict(list)
    for source, target in additions:
        additions_by_source[source].append(target)
    for source, targets in additions_by_source.items():
        targets.sort(key=lambda target: (evidence[(source, target)][2], target))
        for target in targets:
            if target not in final_by_sku[source]:
                final_by_sku[source].append(target)

    errors: list[str] = []
    final_pairs: set[tuple[str, str]] = set()
    for source, targets in final_by_sku.items():
        if len(targets) > MAX_CROSS_SELLS:
            errors.append(f"{source}: {len(targets)} cross-selli")
        if len(targets) != len(set(targets)):
            errors.append(f"{source}: powtórzone cele")
        for target in targets:
            if target == source:
                errors.append(f"{source}: relacja do samego siebie")
            if target not in store_skus:
                errors.append(f"{source}: nieistniejący cel {target}")
            final_pairs.add((source, target))
    if errors:
        raise ValueError("; ".join(errors[:20]))
    if not additions.issubset(final_pairs):
        raise ValueError("Nie wszystkie dodatki znalazły się w wyniku")
    if removals.intersection(final_pairs):
        raise ValueError("Nie wszystkie usunięcia zostały zastosowane")

    changed_rows = 0
    for sku, row_idx in row_by_sku.items():
        before = current_by_sku[sku]
        after = final_by_sku[sku]
        if before == after:
            continue
        changed_rows += 1
        sheet.cell(row_idx, cross_col).value = "|".join(after) if after else None

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

    if sha256(input_path) != input_hash_before or sha256(catalog_path) != catalog_hash_before:
        raise RuntimeError("Plik źródłowy zmienił się podczas operacji")

    check = load_workbook(output_path, read_only=True, data_only=False)
    bad_encoding = scan_mojibake(check)
    check.close()
    if bad_encoding:
        raise ValueError(f"Wykryto uszkodzone kodowanie: {bad_encoding[:5]}")

    return {
        "input": str(input_path.resolve()),
        "catalog": str(catalog_path.resolve()),
        "output": str(output_path.resolve()),
        "rows": sheet.max_row,
        "columns": sheet.max_column,
        "changed_products": changed_rows,
        "cross_sell_pairs_before": len(current_pairs),
        "invalid_pairs_removed": len(INVALID_CURRENT_PAIRS),
        "slot_replacements_removed": len(SLOT_REPLACEMENTS),
        "exact_pairs_added": len(additions),
        "cross_sell_pairs_after": len(final_pairs),
        "products_with_cross_sell_after": sum(bool(targets) for targets in final_by_sku.values()),
        "max_cross_sells_after": max(map(len, final_by_sku.values())),
        "input_sha256_unchanged": sha256(input_path) == input_hash_before,
        "catalog_sha256_unchanged": sha256(catalog_path) == catalog_hash_before,
        "encoding_check": "OK",
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Stosuje audyt cross-selli Hager/Berker oparty na kontrolnej bazie katalogowej."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--catalog", type=Path, default=DEFAULT_CATALOG)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    metrics = build_output(args.input, args.catalog, args.output)
    print(json.dumps(metrics, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
