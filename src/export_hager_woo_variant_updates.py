from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "output" / "Hager-Berker-Gniazdka_WARIANTY_ATRYBUTY_I_TAGI.xlsx"
DEFAULT_WOO_EXPORT = ROOT / "archived_input_files" / "wszystko.csv"
DEFAULT_OUTPUT = ROOT / "output" / "Hager-Berker-WooCommerce_AKTUALIZACJA_ATRYBUTOW_I_TAGOW.csv"
DEFAULT_REPORT = ROOT / "reports" / "Hager-Berker-WooCommerce_AKTUALIZACJA_ATRYBUTOW_I_TAGOW_report.json"

TARGET_COLUMNS = {
    "Kolor": "Kolor",
    "Kolor Producenta": "Kolor producenta",
    "Materiał": "Materiał",
    "Krotność": "Krotność",
    "Orientacja": "Orientacja",
}


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalize(value: Any) -> str:
    source = clean(value).translate(str.maketrans({"ł": "l", "Ł": "L"}))
    text = "".join(
        char
        for char in unicodedata.normalize("NFKD", source)
        if not unicodedata.combining(char)
    ).lower()
    return " ".join(re.findall(r"[a-z0-9]+", text))


def slugify(value: Any) -> str:
    return "-".join(normalize(value).split())


def woo_values(value: Any) -> str:
    parts = [clean(part) for part in str(value or "").split("|") if clean(part)]
    return ", ".join(dict.fromkeys(parts))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class PhpUnserializer:
    """Minimal PHP unserializer for WooCommerce's Product Attributes payload."""

    def __init__(self, value: str):
        self.data = value.encode("utf-8")
        self.pos = 0

    def read_until(self, marker: bytes) -> bytes:
        end = self.data.index(marker, self.pos)
        value = self.data[self.pos:end]
        self.pos = end + len(marker)
        return value

    def expect(self, value: bytes) -> None:
        if self.data[self.pos : self.pos + len(value)] != value:
            raise ValueError(f"Nieprawidłowy payload PHP w pozycji {self.pos}")
        self.pos += len(value)

    def parse(self) -> Any:
        kind = chr(self.data[self.pos])
        self.pos += 1
        if kind == "N":
            self.expect(b";")
            return None
        self.expect(b":")
        if kind in {"i", "b", "d"}:
            raw = self.read_until(b";").decode("ascii")
            if kind == "i":
                return int(raw)
            if kind == "b":
                return raw == "1"
            return float(raw)
        if kind == "s":
            length = int(self.read_until(b":"))
            self.expect(b'"')
            raw = self.data[self.pos : self.pos + length]
            self.pos += length
            self.expect(b'";')
            return raw.decode("utf-8")
        if kind == "a":
            length = int(self.read_until(b":"))
            self.expect(b"{")
            result: dict[Any, Any] = {}
            for _ in range(length):
                key = self.parse()
                value = self.parse()
                result[key] = value
            self.expect(b"}")
            return result
        raise ValueError(f"Nieobsługiwany typ PHP: {kind}")


@dataclass
class Attribute:
    name: str
    value: str
    visible: str = "1"
    global_attribute: str = "1"


def read_woo_catalog(path: Path) -> tuple[dict[str, dict[str, str]], dict[str, str]]:
    rows: dict[str, dict[str, str]] = {}
    slug_labels: dict[str, str] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        for row in reader:
            sku = clean(row.get("SKU"))
            if sku:
                rows[sku] = row
            for index in range(1, 31):
                name = clean(row.get(f"Nazwa atrybutu {index}"))
                if name and clean(row.get(f"Atrybut {index}: globalny")) == "1":
                    slug_labels.setdefault(slugify(name), name)
    return rows, slug_labels


def attributes_from_woo(row: dict[str, str]) -> list[Attribute]:
    result: list[Attribute] = []
    for index in range(1, 31):
        name = clean(row.get(f"Nazwa atrybutu {index}"))
        value = clean(row.get(f"Wartości atrybutu {index}"))
        if not name:
            continue
        if not value:
            raise ValueError(f"Pusty istniejący atrybut WooCommerce: {name}")
        result.append(
            Attribute(
                name=name,
                value=value,
                visible=clean(row.get(f"Atrybut {index} widoczny")) or "1",
                global_attribute=clean(row.get(f"Atrybut {index}: globalny")) or "0",
            )
        )
    return result


def declared_attributes(
    serialized: str,
    product_values: dict[str, str],
    hager_slug_labels: dict[str, str],
    woo_slug_labels: dict[str, str],
) -> list[Attribute]:
    payload = PhpUnserializer(serialized).parse() if serialized else {}
    entries: list[tuple[int, Attribute]] = []
    for item in payload.values():
        if not isinstance(item, dict):
            continue
        raw_name = clean(item.get("name"))
        taxonomy = bool(item.get("is_taxonomy"))
        visible = "1" if item.get("is_visible", 1) else "0"
        position = int(item.get("position", len(entries)))
        if taxonomy:
            raw_slug = raw_name[3:] if raw_name.startswith("pa_") else raw_name
            slug = raw_slug.replace("_", "-")
            label = woo_slug_labels.get(slug) or hager_slug_labels.get(slug) or raw_name
            value = product_values.get(slug, "")
            if not value:
                continue
            entries.append((position, Attribute(label, value, visible, "1")))
        else:
            value = woo_values(item.get("value"))
            if raw_name and value:
                entries.append((position, Attribute(raw_name, value, visible, "0")))
    return [attribute for _position, attribute in sorted(entries, key=lambda item: item[0])]


def set_attribute(attributes: list[Attribute], name: str, value: str) -> None:
    value = woo_values(value)
    if not value:
        return
    target = normalize(name)
    for attribute in attributes:
        if normalize(attribute.name) == target:
            attribute.name = name
            attribute.value = value
            attribute.visible = "1"
            attribute.global_attribute = "1"
            return
    attributes.append(Attribute(name=name, value=value))


def build_rows(input_path: Path, woo_export: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    woo_rows, woo_slug_labels = read_woo_catalog(woo_export)
    workbook = load_workbook(input_path, read_only=True, data_only=False)
    try:
        sheet = workbook["Produkty warianty"]
        headers = [cell.value for cell in sheet[1]]
        columns = {value: index for index, value in enumerate(headers)}
        attribute_columns = {
            header.removeprefix("Atrybut Produktu: "): index
            for index, header in enumerate(headers)
            if isinstance(header, str) and header.startswith("Atrybut Produktu: ")
        }
        hager_slug_labels = {slugify(name): name for name in attribute_columns}
        before_tags_header = next(
            header
            for header in headers
            if isinstance(header, str) and header.startswith("Product Tags") and header != "Product Tags"
        )

        output_rows: list[dict[str, Any]] = []
        matched_woo = 0
        reconstructed = 0
        old_duplicates_included = 0
        parse_errors: list[dict[str, str]] = []

        for values in sheet.iter_rows(min_row=2, values_only=True):
            variant_tag = clean(values[columns["Wariant: tag"]])
            if not variant_tag:
                continue
            sku = clean(values[columns["SKU"]])
            if clean(values[columns["Wariant: decyzja"]]).startswith("POMINIĘTY"):
                old_duplicates_included += 1
                continue

            product_values = {
                slugify(name): woo_values(values[index])
                for name, index in attribute_columns.items()
                if clean(values[index])
            }
            woo_row = woo_rows.get(sku)
            if woo_row:
                attributes = attributes_from_woo(woo_row)
                matched_woo += 1
            else:
                try:
                    attributes = declared_attributes(
                        clean(values[columns["Product Attributes"]]),
                        product_values,
                        hager_slug_labels,
                        woo_slug_labels,
                    )
                except Exception as error:  # noqa: BLE001 - captured in audit report.
                    parse_errors.append({"sku": sku, "error": str(error)})
                    attributes = []
                reconstructed += 1

            desired: dict[str, str] = {}
            for source_name, woo_name in TARGET_COLUMNS.items():
                source_index = attribute_columns[source_name]
                value = clean(values[source_index])
                if value:
                    desired[woo_name] = value
                    set_attribute(attributes, woo_name, value)

            duplicates = [name for name, count in Counter(normalize(item.name) for item in attributes).items() if count > 1]
            if duplicates:
                raise ValueError(f"Powtórzone atrybuty dla SKU {sku}: {duplicates}")

            output_rows.append(
                {
                    "sku": sku,
                    "tags": clean(values[columns["Product Tags"]]),
                    "variant_tag": variant_tag,
                    "attributes": attributes,
                    "desired": desired,
                    "source_tags": clean(values[columns[before_tags_header]]),
                    "source": "Woo export" if woo_row else "Product Attributes XLSX",
                }
            )
    finally:
        workbook.close()

    stats = {
        "matched_to_woo_export": matched_woo,
        "reconstructed_from_xlsx": reconstructed,
        "old_duplicates_included": old_duplicates_included,
        "parse_errors": parse_errors,
    }
    return output_rows, stats


def write_csv(rows: list[dict[str, Any]], output_path: Path) -> int:
    max_attributes = max(len(row["attributes"]) for row in rows)
    headers = ["SKU", "Tagi"]
    for index in range(1, max_attributes + 1):
        headers.extend(
            [
                f"Nazwa atrybutu {index}",
                f"Wartości atrybutu {index}",
                f"Atrybut {index} widoczny",
                f"Atrybut {index}: globalny",
            ]
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for item in rows:
            row: dict[str, str] = {"SKU": item["sku"], "Tagi": item["tags"]}
            for index, attribute in enumerate(item["attributes"], 1):
                row[f"Nazwa atrybutu {index}"] = attribute.name
                row[f"Wartości atrybutu {index}"] = attribute.value
                row[f"Atrybut {index} widoczny"] = attribute.visible
                row[f"Atrybut {index}: globalny"] = attribute.global_attribute
            writer.writerow(row)
    return max_attributes


def validate_csv(path: Path, expected_rows: list[dict[str, Any]], max_attributes: int) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        rows = list(reader)
        headers = reader.fieldnames or []
    expected = {item["sku"]: item for item in expected_rows}
    actual_skus = [clean(row.get("SKU")) for row in rows]
    tag_checks = []
    attribute_checks = []
    empty_attribute_values = []
    for row in rows:
        sku = clean(row.get("SKU"))
        source = expected[sku]
        tag_checks.append(source["variant_tag"] in clean(row.get("Tagi")))
        actual_attributes = {}
        for index in range(1, max_attributes + 1):
            name = clean(row.get(f"Nazwa atrybutu {index}"))
            value = clean(row.get(f"Wartości atrybutu {index}"))
            if not name:
                continue
            actual_attributes[normalize(name)] = value
            if not value:
                empty_attribute_values.append({"sku": sku, "attribute": name})
        for name, value in source["desired"].items():
            attribute_checks.append(actual_attributes.get(normalize(name)) == woo_values(value))
    validation = {
        "row_count_matches": len(rows) == len(expected_rows),
        "all_skus_nonempty": all(actual_skus),
        "skus_unique": len(actual_skus) == len(set(actual_skus)),
        "all_variant_tags_present": all(tag_checks),
        "all_target_attributes_match": all(attribute_checks),
        "no_empty_named_attributes": not empty_attribute_values,
        "only_update_columns_present": headers[:2] == ["SKU", "Tagi"] and all(
            header.startswith(("Nazwa atrybutu ", "Wartości atrybutu ", "Atrybut "))
            for header in headers[2:]
        ),
    }
    validation["passed"] = all(validation.values())
    if not validation["passed"]:
        raise AssertionError(json.dumps(validation, ensure_ascii=False, indent=2))
    return validation


def main() -> None:
    parser = argparse.ArgumentParser(description="Eksport aktualizacji atrybutów i tagów Hager/Berker do WooCommerce")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--woo-export", type=Path, default=DEFAULT_WOO_EXPORT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()

    rows, stats = build_rows(args.input, args.woo_export)
    max_attributes = write_csv(rows, args.output)
    validation = validate_csv(args.output, rows, max_attributes)
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "input_file": str(args.input.resolve()),
        "input_sha256": sha256(args.input),
        "woo_export_file": str(args.woo_export.resolve()),
        "woo_export_sha256": sha256(args.woo_export),
        "output_file": str(args.output.resolve()),
        "output_sha256": sha256(args.output),
        "rows": len(rows),
        "max_attribute_slots": max_attributes,
        **stats,
        "validation": validation,
    }
    if report["parse_errors"] or report["old_duplicates_included"]:
        raise AssertionError(json.dumps(report, ensure_ascii=False, indent=2))
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
