"""Dodaje PODSERIE Kontakt Simon (Line, Duo, Nature, GO, Premium...) do eksportu Woo.

Zrodlo prawdy: oficjalny e-katalog Kontakt-Simon (``src/simon_api.py``) -
parametr "Seria" artykulu oraz drzewo kategorii katalogu. Reguly tlumaczenia
na podserie sklepu: ``dictionaries/simon_subseries_rules.yaml``.

Model (decyzja uzytkownika 2026-09-28):
- atrybut "Seria" = pelna nazwa ("Simon 55 Line"), nowy atrybut "Podseria" = "Line";
- tytul: "<Seria> <Podseria>" tuz przed kodem producenta ("... Simon 55 Line TR1/111");
- podserie dostaja TYLKO produkty, ktore producent przypisuje do podserii
  (ramki, sterowniki GO, puszki natynkowe do ramek Premium); mechanizmy nie;
- produkt pasujacy do kilku podserii (puszki natynkowe do ramek Line i Duo,
  uszczelki IP44) dostaje wszystkie po "|" w atrybutach, a tytul zostaje;
- seria w Woo rozna od serii producenta jest poprawiana (status WARNING,
  arkusz "Korekty serii"); ``--keep-woo-series`` wylacza te korekte.

Wynik: kopia eksportu z podmienionymi kolumnami ``Title`` i
``Atrybut Produktu: Seria`` oraz nowa kolumna ``Atrybut Produktu: Podseria``
(za kolumna Seria). Plik wejsciowy nie jest nadpisywany.

Uzycie:
    py src/build_simon_subseries_woo_update.py --input input/Sajmon-caly.xlsx
    py src/build_simon_subseries_woo_update.py --input input/Sajmon-caly.xlsx --offline
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))

from simon_api import (  # noqa: E402
    SimonApiClient,
    flatten_category_tree,
    is_missing,
    load_api_key,
    normalize_symbol,
    product_params,
)

ROOT = Path(__file__).resolve().parents[1]
TODAY = date.today().isoformat()
DEFAULT_RULES = ROOT / "dictionaries" / "simon_subseries_rules.yaml"
DEFAULT_OUTPUT = ROOT / "output" / f"Simon_podseria_Woo_{TODAY}.xlsx"
DEFAULT_CONTROL = ROOT / "reports" / f"Simon_podseria_{TODAY}_kontrola.xlsx"
DEFAULT_REPORT = ROOT / "reports" / f"Simon_podseria_{TODAY}.json"

COL_TITLE = "Title"
COL_SKU = "SKU"
COL_ID = "id"
COL_SERIES = "Atrybut Produktu: Seria"
COL_SUBSERIES = "Atrybut Produktu: Podseria"

STATUS_ORDER = {"OK": 0, "WARNING": 1, "DO_SPRAWDZENIA": 2}
MULTI_SEPARATOR = "|"  # separator wielu wartosci atrybutu w eksporcie/imporcie Woo
SUBSERIES_WORDS = "premium|nature|go|line|duo|flash"


@dataclass
class Resolution:
    series: str | None = None
    subseries: str | None = None
    status: str = "OK"
    source: str = ""
    confidence: float = 0.0
    note: str = ""
    suggestion: str = ""
    series_full: bool = True
    touch_title: bool = True


@dataclass
class RowResult:
    row_id: str
    sku: str
    symbol: str
    title_before: str
    title_after: str
    series_before: str
    series_after: str
    subseries: str
    api_series: str
    api_categories: str
    source: str
    confidence: float
    status: str
    note: str
    suggestion: str = ""
    changes: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- reguly
def load_rules(path: Path = DEFAULT_RULES) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as stream:
        return yaml.safe_load(stream)


def clean_api_series(value: str | None) -> str:
    return re.sub(r"\s+", " ", (value or "").replace("\n", " ")).strip()


def _matches(rule: dict[str, Any], paths: list[list[str]]) -> bool:
    pattern = re.compile(rule["match"], flags=re.I)
    if rule.get("require_single_path") and len(paths) > 1:
        return False
    return any(pattern.search(name) for path in paths for name in path)


def resolve_subseries(api_series: str, paths: list[list[str]], rules: dict[str, Any]) -> Resolution:
    """Ustala serie i podserie producenta dla artykulu z API."""
    api_series = clean_api_series(api_series)
    base_series = set(rules.get("base_series") or [])
    aliases = rules.get("series_aliases") or {}

    entry = (rules.get("api_series") or {}).get(api_series)
    if entry:
        return Resolution(
            series=entry["series"],
            subseries=entry["subseries"],
            status=entry.get("status", "OK"),
            source="api:Seria",
            confidence=0.98 if entry.get("status", "OK") == "OK" else 0.85,
            note=entry.get("note", ""),
            series_full=bool(entry.get("series_full", True)),
            touch_title=bool(entry.get("title", True)),
        )

    for rule in rules.get("multi_subseries_category_rules") or []:
        if _matches(rule, paths):
            # Kilka podserii naraz: wszystkie po "|" w atrybutach, tytul nietkniety.
            return Resolution(
                series=rule["series"],
                subseries=MULTI_SEPARATOR.join(rule["subseries"]),
                status=rule.get("status", "WARNING"),
                source="api:kategoria",
                confidence=0.8,
                note=rule.get("note", ""),
                touch_title=False,
            )

    for rule in rules.get("category_rules") or []:
        if _matches(rule, paths):
            return Resolution(
                series=rule["series"],
                subseries=rule["subseries"],
                status=rule.get("status", "OK"),
                source="api:kategoria",
                confidence=0.95 if rule.get("status", "OK") == "OK" else 0.8,
                note=rule.get("note", ""),
                series_full=bool(rule.get("series_full", True)),
                touch_title=bool(rule.get("title", True)),
            )

    if api_series in base_series:
        return Resolution(series=api_series, source="api:Seria", confidence=0.95)
    if api_series in aliases:
        return Resolution(series=aliases[api_series], source="api:Seria", confidence=0.9,
                          note=f"Seria API '{api_series}' -> {aliases[api_series]}")
    return Resolution(series=None, source="api:Seria", confidence=0.6,
                      note=f"Seria API '{api_series}' nie jest seria sklepu - Seria bez zmian")


# --------------------------------------------------------------------------- tytul
def _squash(text: str) -> str:
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s+([:,.;])", r"\1", text)
    return text.strip()


def _series_regex(series: str, rules: dict[str, Any]) -> re.Pattern | None:
    pattern = (rules.get("title", {}).get("series_patterns") or {}).get(series)
    if not pattern:
        return None
    return re.compile(rf"(?<![\w.])(?:{pattern})(?:\s+(?:{SUBSERIES_WORDS}))?(?![\w])", flags=re.I)


def build_title(title: str, series: str | None, subseries: str | None, symbol: str,
                rules: dict[str, Any]) -> tuple[str, str]:
    """Zwraca (nowy tytul, uwaga). Pusta uwaga = zmiana bezproblemowa lub brak zmiany."""
    title = title or ""
    if not series:
        return title, ""
    title_rules = rules.get("title", {})
    target = f"{series} {subseries}" if subseries else series
    if subseries and subseries in (title_rules.get("skip_if_token_present") or []):
        if re.search(rf"(?<![\w-]){re.escape(subseries)}(?![\w-])", title, flags=re.I):
            return title, ""
    target_pattern = re.escape(target).replace("\\ ", "\\s*")
    if subseries:
        target_regex = re.compile(rf"(?<![\w.]){target_pattern}(?![\w])", flags=re.I)
    else:
        # Sama seria: "Simon 54" liczy sie tylko, gdy NIE stoi za nia stara podseria.
        target_regex = re.compile(rf"(?<![\w.]){target_pattern}(?![\w])(?!\s+(?:{SUBSERIES_WORDS})(?![\w]))", flags=re.I)
    if target_regex.search(title):
        return title, ""

    work = title
    if subseries:
        strip = (title_rules.get("strip_tokens") or {}).get(subseries)
        if strip:
            work = re.sub(rf"(?<!\S)(?:{strip})(?=[\s:,.;]|$)", "", work)
            work = _squash(work)

    series_re = _series_regex(series, rules)
    if series_re and series_re.search(work):
        new_title = series_re.sub(target, work, count=1)
        return _squash(new_title), ""

    normalized_symbol = normalize_symbol(symbol)
    if normalized_symbol:
        tail = re.compile(rf"\s+{re.escape(normalized_symbol)}\s*$", flags=re.I)
        match = tail.search(work)
        if match:
            new_title = work[: match.start()] + f" {target} {normalized_symbol}"
            return _squash(new_title), ""
    return _squash(f"{work} {target}"), "kod producenta nie stoi na koncu tytulu - seria dopisana na koncu"


# --------------------------------------------------------------------------- wiersz
def process_row(row: dict[str, Any], product: dict | None, paths: list[list[str]],
                rules: dict[str, Any], keep_woo_series: bool = False) -> RowResult:
    sku = str(row.get(COL_SKU) or "").strip()
    symbol = normalize_symbol(sku)
    title_before = str(row.get(COL_TITLE) or "").strip()
    series_before = str(row.get(COL_SERIES) or "").strip()
    row_id = str(row.get(COL_ID) or "").strip()

    if is_missing(product):
        return RowResult(row_id, sku, symbol, title_before, title_before, series_before, series_before,
                         "", "", "", "api", 0.0, "DO_SPRAWDZENIA",
                         "brak artykulu w API Simon - bez zmian")

    params = product_params(product)
    api_series = clean_api_series(params.get("Seria"))
    categories = " || ".join(" > ".join(path[2:]) for path in paths) if paths else ""
    resolution = resolve_subseries(api_series, paths, rules)

    changes: list[str] = []
    notes: list[str] = []
    if resolution.note:
        notes.append(resolution.note)
    status = resolution.status

    # --- Seria
    series_after = series_before
    if resolution.series:
        if resolution.subseries and resolution.series_full:
            wanted = MULTI_SEPARATOR.join(
                f"{resolution.series} {part}" for part in resolution.subseries.split(MULTI_SEPARATOR)
            )
        else:
            wanted = resolution.series
        if wanted != series_before:
            if not series_before:
                series_after = wanted
                changes.append("Seria")
                notes.append("uzupelniono pusta serie wg API")
                status = max(status, "WARNING", key=STATUS_ORDER.get)
            elif keep_woo_series and not resolution.subseries:
                notes.append(f"seria Woo '{series_before}' rozna od API '{wanted}' - zostawiona (--keep-woo-series)")
                status = max(status, "WARNING", key=STATUS_ORDER.get)
            else:
                series_after = wanted
                changes.append("Seria")
                if not resolution.subseries:
                    notes.append(f"korekta serii wg API: '{series_before}' -> '{wanted}'")
                    status = max(status, "WARNING", key=STATUS_ORDER.get)

    # --- Podseria
    subseries = resolution.subseries or ""
    if subseries:
        changes.append("Podseria")

    # --- Tytul
    title_after = title_before
    if resolution.series and resolution.touch_title and (subseries or "Seria" in changes):
        title_after, title_note = build_title(title_before, resolution.series, resolution.subseries, symbol, rules)
        if title_after != title_before:
            changes.append("Title")
        if title_note:
            notes.append(title_note)
            status = max(status, "WARNING", key=STATUS_ORDER.get)

    return RowResult(
        row_id=row_id, sku=sku, symbol=symbol,
        title_before=title_before, title_after=title_after,
        series_before=series_before, series_after=series_after,
        subseries=subseries, api_series=api_series, api_categories=categories,
        source=resolution.source, confidence=resolution.confidence,
        status=status, note="; ".join(notes), suggestion=resolution.suggestion, changes=changes,
    )


# --------------------------------------------------------------------------- I/O
def read_export(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell) if cell is not None else "" for cell in next(rows)]
    records = []
    for values in rows:
        record = {}
        for index, name in enumerate(header):
            if name and name not in record:
                record[name] = values[index] if index < len(values) else None
        records.append(record)
    workbook.close()
    return header, records


def write_output(source: Path, target: Path, results: list[RowResult]) -> None:
    workbook = load_workbook(source)
    sheet = workbook.worksheets[0]
    header = [cell.value for cell in sheet[1]]
    col_title = header.index(COL_TITLE) + 1
    col_sku = header.index(COL_SKU) + 1
    col_series = header.index(COL_SERIES) + 1
    if COL_SUBSERIES in header:
        col_sub = header.index(COL_SUBSERIES) + 1
    else:
        col_sub = col_series + 1
        sheet.insert_cols(col_sub)
        sheet.cell(row=1, column=col_sub, value=COL_SUBSERIES)
        if col_title >= col_sub:
            col_title += 1
        if col_sku >= col_sub:
            col_sku += 1
    by_row = {}
    for result in results:
        by_row.setdefault(result.sku, []).append(result)
    for row_index in range(2, sheet.max_row + 1):
        sku = str(sheet.cell(row=row_index, column=col_sku).value or "").strip()
        queue = by_row.get(sku)
        if not queue:
            continue
        result = queue.pop(0) if len(queue) > 1 else queue[0]
        if result.title_after != result.title_before:
            sheet.cell(row=row_index, column=col_title, value=result.title_after)
        if result.series_after != result.series_before:
            sheet.cell(row=row_index, column=col_series, value=result.series_after or None)
        sheet.cell(row=row_index, column=col_sub, value=result.subseries or None)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)


CONTROL_COLUMNS = [
    ("id", "row_id"), ("SKU", "sku"), ("Symbol", "symbol"),
    ("Tytul przed", "title_before"), ("Tytul po", "title_after"),
    ("Seria przed", "series_before"), ("Seria po", "series_after"), ("Podseria", "subseries"),
    ("Seria wg API", "api_series"), ("Kategoria e-katalogu", "api_categories"),
    ("Zrodlo", "source"), ("Pewnosc", "confidence"), ("Status", "status"), ("Uwagi", "note"),
    ("Zmienione pola", "changes"),
]


def _write_sheet(workbook: Workbook, title: str, rows: list[RowResult]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append([label for label, _ in CONTROL_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    fills = {"WARNING": PatternFill("solid", fgColor="FFF2CC"), "DO_SPRAWDZENIA": PatternFill("solid", fgColor="F8CBAD")}
    for result in rows:
        values = []
        for _, attr in CONTROL_COLUMNS:
            value = getattr(result, attr)
            values.append(", ".join(value) if isinstance(value, list) else value)
        sheet.append(values)
        fill = fills.get(result.status)
        if fill:
            sheet.cell(row=sheet.max_row, column=13).fill = fill
    for index, (label, _) in enumerate(CONTROL_COLUMNS, start=1):
        width = 60 if "Tytul" in label or "Kategoria" in label or label == "Uwagi" else 18
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def write_control(path: Path, results: list[RowResult], summary: dict[str, Any]) -> None:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Podsumowanie"
    for key, value in summary.items():
        if isinstance(value, dict):
            overview.append([key])
            for sub_key, sub_value in value.items():
                overview.append(["", sub_key, sub_value])
        else:
            overview.append([key, ", ".join(map(str, value)) if isinstance(value, list) else value])
    overview.column_dimensions["A"].width = 34
    overview.column_dimensions["B"].width = 40
    changed = [r for r in results if r.changes]
    _write_sheet(workbook, "Zmiany", changed)
    _write_sheet(workbook, "Podserie", [r for r in results if r.subseries])
    _write_sheet(workbook, "Korekty serii", [r for r in results if "Seria" in r.changes and not r.subseries])
    _write_sheet(workbook, "Do sprawdzenia", [r for r in results if r.status == "DO_SPRAWDZENIA"])
    _write_sheet(workbook, "Wszystkie", results)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def build_summary(results: list[RowResult], extra: dict[str, Any]) -> dict[str, Any]:
    def count(predicate) -> int:
        return sum(1 for r in results if predicate(r))

    subseries_counts: dict[str, int] = {}
    for result in results:
        if result.subseries:
            key = f"{result.series_after} / {result.subseries}"
            subseries_counts[key] = subseries_counts.get(key, 0) + 1
    return {
        "data": TODAY,
        "wiersze": len(results),
        "z_podseria": count(lambda r: bool(r.subseries)),
        "z_wieloma_podseriami": count(lambda r: MULTI_SEPARATOR in r.subseries),
        "zmienione_tytuly": count(lambda r: "Title" in r.changes),
        "zmieniona_seria": count(lambda r: "Seria" in r.changes),
        "korekty_serii_bez_podserii": count(lambda r: "Seria" in r.changes and not r.subseries),
        "status": {s: count(lambda r, s=s: r.status == s) for s in STATUS_ORDER},
        "brak_w_api": count(lambda r: "brak artykulu w API" in r.note),
        "podserie": dict(sorted(subseries_counts.items())),
        **extra,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Dodaje podserie Kontakt Simon do eksportu Woo.")
    parser.add_argument("--input", required=True, type=Path, help="Eksport WP All Export (XLSX)")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--control", type=Path, default=DEFAULT_CONTROL)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--rules", type=Path, default=DEFAULT_RULES)
    parser.add_argument("--keep-woo-series", action="store_true",
                        help="Nie poprawiaj serii Woo, gdy rozni sie od serii producenta (tylko raport)")
    parser.add_argument("--offline", action="store_true", help="Tylko cache API (cache/simon_api), bez sieci")
    parser.add_argument("--refresh", action="store_true", help="Pobierz artykuly z API na nowo (ignoruj cache)")
    parser.add_argument("--api-key", default=None)
    args = parser.parse_args()

    if args.input.resolve() == args.output.resolve():
        raise SystemExit("Plik wyjsciowy nie moze byc plikiem wejsciowym")

    rules = load_rules(args.rules)
    header, records = read_export(args.input)
    for required in (COL_TITLE, COL_SKU, COL_SERIES):
        if required not in header:
            raise SystemExit(f"Brak kolumny '{required}' w pliku wejsciowym")

    client = SimonApiClient(load_api_key(args.api_key) if not args.offline else "offline", verbose=True)
    skus = [str(r.get(COL_SKU) or "") for r in records]
    if args.offline:
        products = {normalize_symbol(s): client.cached_product(s) for s in skus if s}
        tree = client.fetch_category_tree() if (client.cache_dir / "category_tree.json").exists() else {}
    else:
        products = client.fetch_products(skus, refresh=args.refresh)
        tree = client.fetch_category_tree(refresh=args.refresh)
    paths_by_id = flatten_category_tree(tree) if tree else {}

    results: list[RowResult] = []
    for record in records:
        symbol = normalize_symbol(str(record.get(COL_SKU) or ""))
        product = products.get(symbol)
        paths = paths_by_id.get(int(product["id"]), []) if product and not is_missing(product) else []
        results.append(process_row(record, product, paths, rules, keep_woo_series=args.keep_woo_series))

    seen: dict[str, int] = {}
    for sku in skus:
        seen[sku] = seen.get(sku, 0) + 1
    duplicates = sorted(s for s, n in seen.items() if n > 1)

    write_output(args.input, args.output, results)
    summary = build_summary(results, {
        "plik_wejsciowy": str(args.input), "plik_wyjsciowy": str(args.output),
        "kontrola": str(args.control), "duplikaty_sku": duplicates,
        "zapytania_api": client.requests_made, "keep_woo_series": args.keep_woo_series,
    })
    write_control(args.control, results, summary)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in summary.items() if k not in ("podserie",)}, ensure_ascii=False, indent=1))
    print("Podserie:", json.dumps(summary["podserie"], ensure_ascii=False))


if __name__ == "__main__":
    main()
