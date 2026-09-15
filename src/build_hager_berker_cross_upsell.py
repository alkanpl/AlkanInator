from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from build_hager_berker_variants import (
    Record,
    compact,
    find_headers,
    load_records,
    mark_old_duplicate_products,
    mojibake_markers,
    norm,
    remove_series_phrases,
    split_tags,
    title_without_codes,
)


DEFAULT_INPUT = Path("output/Hager-Berker-Gniazdka_WOO_ALL_IMPORT_OSIE_POPRAWIONE.xlsx")
DEFAULT_OUTPUT = Path("output/Hager-Berker-Gniazdka_WOO_ALL_IMPORT_CROSS_UP_SELL.xlsx")
DEFAULT_REPORT = Path("reports/hager_berker_cross_up_sell/cross_up_sell_report.xlsx")

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
WHITE = "FFFFFF"

KNOWN_TEXT_REPAIRS = {
    "Ĺšrubowo-windowe": "Śrubowo-windowe",
    "Ĺšrubowe": "Śrubowe",
}

UPSELL_NOISE = {
    "podswietlany",
    "podswietlana",
    "podswietlane",
    "podswietleniem",
    "podswietlenie",
    "przeslona",
    "przeslonami",
    "przeslony",
    "komfort",
    "komfortowy",
    "komfortowa",
    "ekranowany",
    "nieekranowany",
}

COLOR_NOISE = {
    "bialy", "biala", "biale", "czarny", "czarna", "czarne", "kremowy", "kremowa",
    "antracyt", "antracytowy", "aluminium", "alu", "srebrny", "srebrna", "zloty",
    "zloty", "mosiadz", "stal", "szlachetna", "nierdzewna", "czerwony", "czerwona",
    "zielony", "zielona", "niebieski", "niebieska", "zolty", "zolta", "pomaranczowy",
    "mat", "matowy", "matowa", "polysk", "aksamit", "lakierowany", "lakierowana",
    "szklo", "szklany", "szklana", "transparentny", "przezroczysty",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def split_series(value: str) -> frozenset[str]:
    return frozenset(norm(item) for item in str(value or "").split("|") if compact(item))


def variant_tags(value: str) -> frozenset[str]:
    return frozenset(norm(tag) for tag in split_tags(value) if norm(tag).startswith("wariant "))


def classify_function(record: Record) -> str:
    text = norm(f"{record.title} {record.product_type}")
    if "usb" in text:
        return "usb"
    if any(token in text for token in ("komputer", "rj45", "uae", "sieciow", "teleinformat")):
        return "data"
    if any(token in text for token in ("anten", "rtv", "sat", "koncentrycz", "coax")):
        return "tv_sat"
    if "hdmi" in text:
        return "hdmi"
    if any(token in text for token in ("glosnik", "cinch", "audio", "video")):
        return "audio_video"
    if "sciemniacz" in text:
        return "dimmer"
    if "zaluzj" in text:
        return "shutter_control"
    if "termostat" in text or "regulator temperatur" in text:
        return "thermostat"
    if "czujnik ruch" in text:
        return "motion_sensor"
    if "ciegl" in text:
        return "pull_switch"
    if "na klucz" in text or "kluczyk" in text:
        return "key_switch"
    if "hotel" in text or "na karte" in text:
        return "hotel_switch"
    if "czasow" in text:
        return "timer_switch"
    if "obrotow" in text:
        return "rotary_switch"
    if "sygnalizator" in text:
        return "signaling"
    if "krzyzow" in text:
        return "cross_switch"
    if "schodow" in text:
        return "stair_switch"
    if any(token in text for token in ("swiecznik", "seryjny", "klawisz podwojny", "lacznik podwojny")):
        return "double_switch"
    multiplicity = norm(record.multiplicity)
    if "klawisz" in text and re.search(r"\b2\s*x?\b", multiplicity):
        return "double_switch"
    if "klawisz" in text and re.search(r"\b[3-5]\s*x?\b", multiplicity):
        return "multi_key"
    if "przycisk" in text or "dzwonk" in text:
        return "button"
    if "lacznik" in text or "klawisz" in text:
        return "single_switch"
    if "gniazdo" in text or "schuko" in text:
        return "power_socket"
    if "zaslep" in text:
        return "blanking"
    return "other"


def classify_role(record: Record) -> str:
    text = norm(f"{record.title} {record.product_type}")
    product_type = norm(record.product_type)
    if record.frame:
        return "frame"
    if "zestaw uszczeln" in text:
        return "seal"
    if "puszka natynk" in text or "adapter natynk" in text:
        return "mounting"
    if "klawisz" in product_type or text.startswith("klawisz "):
        return "key"
    if any(token in product_type for token in ("plytka", "element centralny", "nasadka", "pokryw")):
        return "cover"
    if text.startswith(("plytka ", "element centralny ", "nasadka ", "pokrywa ")):
        return "cover"
    if "mechanizm" in product_type or "mechanizm" in text:
        return "mechanism"
    if any(token in text for token in ("gniazdo", "lacznik", "przycisk", "sciemniacz", "regulator", "czujnik")):
        return "complete"
    if "zaslep" in text:
        return "complete"
    return "other"


def color_key(record: Record) -> str:
    return norm(record.manufacturer_color or record.color)


def colors_compatible(left: "CatalogProduct", right: "CatalogProduct") -> bool:
    return bool(left.color and right.color and left.color == right.color)


def series_compatible(left: "CatalogProduct", right: "CatalogProduct") -> bool:
    return bool(left.series and right.series and left.series.intersection(right.series))


def same_design_family(left: "CatalogProduct", right: "CatalogProduct") -> bool:
    left_primary = norm(left.record.primary_series)
    right_primary = norm(right.record.primary_series)
    return bool(
        left_primary
        and right_primary
        and left_primary in right.series
        and right_primary in left.series
    )


def same_variant_group(left: "CatalogProduct", right: "CatalogProduct") -> bool:
    return bool(left.variant_tags and right.variant_tags and left.variant_tags.intersection(right.variant_tags))


def parse_ip(text: str) -> int:
    values = [int(value) for value in re.findall(r"\bip\s*([0-9]{2})\b", norm(text))]
    return max(values, default=0)


def parse_category(record: Record) -> int:
    text = norm(f"{record.network_category} {record.title}")
    match = re.search(r"\b(?:kat|kategoria)\s*([56])\s*([ae]?)\b", text)
    if not match:
        return 0
    return {("5", ""): 50, ("5", "e"): 51, ("6", ""): 60, ("6", "a"): 61}.get(
        (match.group(1), match.group(2)), 0,
    )


def parse_power(text: str) -> int:
    values = [int(value) for value in re.findall(r"\b([0-9]{1,4})\s*w\b", norm(text))]
    return max(values, default=0)


def configuration_count(record: Record, role: str, function: str) -> int:
    multiplicity = norm(record.multiplicity)
    match = re.search(r"\b([1-5])\s*x?\b", multiplicity)
    if match:
        return int(match.group(1))
    text = norm(record.title)
    words = (("pieci", 5), ("poczwor", 4), ("potroj", 3), ("podwoj", 2), ("pojedyncz", 1))
    for prefix, value in words:
        if prefix in text:
            return value
    if role == "key" or function == "power_socket":
        return 1
    return 0


def upgrade_signature(record: Record, function: str) -> str:
    text = norm(record.core or title_without_codes(record.title, record.sku))
    text = remove_series_phrases(text)
    text = re.sub(r"\bip\s*[0-9]{2}\b", " ", text)
    text = re.sub(r"\b(?:kat|kategoria)\s*[56]\s*[ae]?\b", " ", text)
    tokens = [token for token in text.split() if token not in UPSELL_NOISE and token not in COLOR_NOISE]
    if function == "dimmer":
        tokens = [token for token in tokens if not re.fullmatch(r"[0-9]+", token)]
        tokens = [token for token in tokens if token not in {"w", "v", "led", "obciazenie", "obc"}]
    return " ".join(tokens)


@dataclass
class CatalogProduct:
    record: Record
    role: str = field(init=False)
    function: str = field(init=False)
    series: frozenset[str] = field(init=False)
    color: str = field(init=False)
    variant_tags: frozenset[str] = field(init=False)
    upgrade_key: str = field(init=False)
    improvements: tuple[int, int, int, int, int, int] = field(init=False)
    configuration_count: int = field(init=False)

    def __post_init__(self) -> None:
        self.role = classify_role(self.record)
        self.function = classify_function(self.record)
        self.series = split_series(self.record.series)
        self.color = color_key(self.record)
        self.variant_tags = variant_tags(self.record.tags_before)
        self.upgrade_key = upgrade_signature(self.record, self.function)
        self.configuration_count = configuration_count(self.record, self.role, self.function)
        text = norm(self.record.title)
        shutters = 1 if norm(self.record.shutters) in {"tak", "yes", "1"} or "przeslon" in text else 0
        screened = 1 if norm(self.record.screened) in {"tak", "yes", "1"} or "ekranowan" in text else 0
        self.improvements = (
            1 if "podswietl" in text else 0,
            parse_ip(text),
            shutters,
            parse_category(self.record),
            screened,
            (1 if "komfort" in text else 0) * 10000 + parse_power(text),
        )

    @property
    def sku(self) -> str:
        return self.record.sku

    @property
    def title(self) -> str:
        return self.record.title


@dataclass(frozen=True)
class Recommendation:
    source_sku: str
    target_sku: str
    kind: str
    score: int
    confidence: str
    reason: str


def mechanisms_compatible(source_function: str, target_function: str) -> bool:
    if source_function == target_function:
        return source_function != "other"
    if source_function == "single_switch" and target_function in {"stair_switch", "cross_switch", "button"}:
        return True
    if source_function == "stair_switch" and target_function == "single_switch":
        return True
    return False


def cross_candidate(source: CatalogProduct, target: CatalogProduct) -> Recommendation | None:
    if source.sku == target.sku or target.record.is_old_duplicate or same_variant_group(source, target):
        return None

    if source.role == "frame" and target.role in {"key", "cover", "complete"}:
        if series_compatible(source, target) and colors_compatible(source, target):
            function_priority = {
                "power_socket": 12,
                "single_switch": 11,
                "stair_switch": 10,
                "double_switch": 10,
                "button": 9,
                "shutter_control": 9,
                "dimmer": 8,
                "usb": 8,
                "data": 7,
                "tv_sat": 5,
                "hdmi": 4,
                "audio_video": 2,
                "pull_switch": 1,
                "key_switch": 1,
                "hotel_switch": 1,
                "timer_switch": 1,
                "rotary_switch": 1,
                "signaling": 1,
                "other": 1,
            }.get(target.function, 0)
            return Recommendation(
                source.sku, target.sku, "Cross-sell", 90 + function_priority, "WYSOKA",
                "element pasujący do ramki: zgodna seria i kolor producenta",
            )

    if source.role in {"key", "cover", "complete"} and target.role == "frame":
        if target.record.multiplicity == "1x" and series_compatible(source, target) and colors_compatible(source, target):
            return Recommendation(
                source.sku, target.sku, "Cross-sell", 100, "WYSOKA",
                "ramka 1-krotna: zgodna seria i kolor producenta",
            )

    if source.role in {"key", "cover"} and target.role == "mechanism":
        if mechanisms_compatible(source.function, target.function):
            return Recommendation(
                source.sku, target.sku, "Cross-sell", 100 if source.function == target.function else 97, "WYSOKA",
                "mechanizm dopasowany do funkcji klawisza lub płytki czołowej",
            )

    source_text = norm(source.title)
    target_text = norm(target.title)
    if source.role == "seal" and target.role == "complete":
        if "ip44" in target_text and series_compatible(source, target) and colors_compatible(source, target):
            return Recommendation(
                source.sku, target.sku, "Cross-sell", 97, "WYSOKA",
                "produkt IP44 zgodny z zestawem uszczelniającym, serią i kolorem",
            )
    if target.role == "seal" and source.role == "complete":
        if "ip44" in source_text and series_compatible(source, target) and colors_compatible(source, target):
            return Recommendation(
                source.sku, target.sku, "Cross-sell", 99, "WYSOKA",
                "zestaw uszczelniający zgodny z produktem IP44, serią i kolorem",
            )

    return None


def up_sell_candidate(source: CatalogProduct, target: CatalogProduct) -> Recommendation | None:
    if source.sku == target.sku or target.record.is_old_duplicate or same_variant_group(source, target):
        return None
    if source.role in {"frame", "seal", "mounting", "other"} or source.role != target.role:
        return None
    if source.function == "other" or source.function != target.function:
        return None
    if not same_design_family(source, target) or not colors_compatible(source, target):
        return None
    if source.configuration_count != target.configuration_count:
        return None
    if not source.upgrade_key or source.upgrade_key != target.upgrade_key:
        return None

    source_values = source.improvements
    target_values = target.improvements
    if not all(target_value >= source_value for source_value, target_value in zip(source_values, target_values)):
        return None
    improved = [index for index, (source_value, target_value) in enumerate(zip(source_values, target_values)) if target_value > source_value]
    if not improved:
        return None

    labels = ["podświetlenie", "wyższy stopień IP", "przesłony torów prądowych", "wyższa kategoria sieciowa", "ekranowanie", "większe możliwości/moc"]
    reason = "ulepszenie tej samej funkcji: " + ", ".join(labels[index] for index in improved)
    return Recommendation(source.sku, target.sku, "Up-sell", 100 + len(improved), "WYSOKA", reason)


def diversity_key(product: CatalogProduct) -> tuple[str, str]:
    if product.role == "frame":
        return ("frame", "frame")
    return (product.role, product.function)


def select_recommendations(
    products: list[CatalogProduct],
    max_cross_sells: int = 4,
    max_up_sells: int = 3,
) -> tuple[dict[str, list[Recommendation]], dict[str, list[Recommendation]]]:
    cross: dict[str, list[Recommendation]] = {}
    up: dict[str, list[Recommendation]] = {}
    by_sku = {product.sku: product for product in products}
    eligible = [product for product in products if product.sku and not product.record.is_old_duplicate]

    for source in eligible:
        candidates = [candidate for target in eligible if (candidate := cross_candidate(source, target))]
        candidates.sort(key=lambda item: (-item.score, by_sku[item.target_sku].title, item.target_sku))
        chosen: list[Recommendation] = []
        seen_diversity: set[tuple[str, str]] = set()
        for candidate in candidates:
            key = diversity_key(by_sku[candidate.target_sku])
            if key in seen_diversity:
                continue
            chosen.append(candidate)
            seen_diversity.add(key)
            if len(chosen) >= max_cross_sells:
                break
        if chosen:
            cross[source.sku] = chosen

        cross_targets = {item.target_sku for item in chosen}
        up_candidates = [candidate for target in eligible if (candidate := up_sell_candidate(source, target))]
        up_candidates = [item for item in up_candidates if item.target_sku not in cross_targets]
        up_candidates.sort(key=lambda item: (-item.score, by_sku[item.target_sku].title, item.target_sku))
        if up_candidates:
            up[source.sku] = up_candidates[:max_up_sells]

    return cross, up


def validate_recommendations(
    products: list[CatalogProduct],
    cross: dict[str, list[Recommendation]],
    up: dict[str, list[Recommendation]],
    max_cross_sells: int,
    max_up_sells: int,
) -> dict[str, bool]:
    skus = {product.sku for product in products}
    all_relations = [item for mapping in (cross, up) for items in mapping.values() for item in items]
    no_unknown = all(item.source_sku in skus and item.target_sku in skus for item in all_relations)
    no_self = all(item.source_sku != item.target_sku for item in all_relations)
    unique_lists = all(
        len(items) == len({item.target_sku for item in items})
        for mapping in (cross, up) for items in mapping.values()
    )
    disjoint = all(
        {item.target_sku for item in cross.get(sku, [])}.isdisjoint(
            item.target_sku for item in up.get(sku, [])
        )
        for sku in skus
    )
    encoding_ok = all(
        not mojibake_markers(value)
        for item in all_relations
        for value in (item.reason, item.source_sku, item.target_sku)
    )
    return {
        "all_referenced_skus_exist": no_unknown,
        "no_self_references": no_self,
        "no_duplicate_targets_per_product": unique_lists,
        "cross_sell_and_up_sell_are_disjoint": disjoint,
        "cross_sell_limit_respected": all(len(items) <= max_cross_sells for items in cross.values()),
        "up_sell_limit_respected": all(len(items) <= max_up_sells for items in up.values()),
        "polish_text_encoding_intact": encoding_ok,
    }


def write_report(
    path: Path,
    products: list[CatalogProduct],
    cross: dict[str, list[Recommendation]],
    up: dict[str, list[Recommendation]],
    validation: dict[str, bool],
    encoding_repairs: list[dict[str, Any]],
    source_path: Path,
    output_path: Path,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Podsumowanie"
    by_sku = {product.sku: product for product in products}
    cross_count = sum(len(items) for items in cross.values())
    up_count = sum(len(items) for items in up.values())
    rows = [
        ("Plik źródłowy", str(source_path.resolve())),
        ("Plik do importu", str(output_path.resolve())),
        ("Produkty", len(products)),
        ("Produkty z cross-sell", len(cross)),
        ("Powiązania cross-sell", cross_count),
        ("Produkty z up-sell", len(up)),
        ("Powiązania up-sell", up_count),
        ("Stare duplikaty pominięte", sum(product.record.is_old_duplicate for product in products)),
        ("Naprawione odziedziczone wartości kodowania", len(encoding_repairs)),
        ("Zasada", "Tylko relacje wysokiej pewności; stare pola wygenerowano od nowa."),
    ]
    summary.append(["Metryka", "Wartość"])
    for row in rows:
        summary.append(row)

    relations = workbook.create_sheet("Powiązania")
    relations.append([
        "Rodzaj", "SKU źródłowe", "Tytuł źródłowy", "SKU docelowe", "Tytuł docelowy",
        "Ocena", "Pewność", "Uzasadnienie",
    ])
    for kind, mapping in (("Cross-sell", cross), ("Up-sell", up)):
        for source_sku in sorted(mapping):
            for item in mapping[source_sku]:
                relations.append([
                    kind,
                    source_sku,
                    by_sku[source_sku].title,
                    item.target_sku,
                    by_sku[item.target_sku].title,
                    item.score,
                    item.confidence,
                    item.reason,
                ])

    validation_sheet = workbook.create_sheet("Walidacja")
    validation_sheet.append(["Test", "Wynik"])
    for name, passed in validation.items():
        validation_sheet.append([name, "OK" if passed else "BŁĄD"])

    repairs_sheet = workbook.create_sheet("Naprawy kodowania")
    repairs_sheet.append(["Wiersz", "Kolumna", "Wartość przed", "Wartość po"])
    for repair in encoding_repairs:
        repairs_sheet.append([
            repair["row"], repair["column"], repair["before"], repair["after"],
        ])

    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=BLUE)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            width = min(max(len(compact(cell.value)) for cell in column) + 2, 70)
            sheet.column_dimensions[column[0].column_letter].width = max(width, 12)
    workbook.save(path)
    workbook.close()


def compare_preserved_cells(input_path: Path, output_path: Path, sheet_name: str) -> bool:
    source = load_workbook(input_path, read_only=True, data_only=False)
    result = load_workbook(output_path, read_only=True, data_only=False)
    try:
        if source.sheetnames != result.sheetnames:
            return False
        source_sheet = source[sheet_name]
        result_sheet = result[sheet_name]
        source_headers = [cell.value for cell in source_sheet[1]]
        result_headers = [cell.value for cell in result_sheet[1]]
        if source_headers != result_headers:
            return False
        editable = {source_headers.index("Up-Sells"), source_headers.index("Cross-Sells")}
        for source_row, result_row in zip(
            source_sheet.iter_rows(values_only=True), result_sheet.iter_rows(values_only=True), strict=True,
        ):
            for index, (left, right) in enumerate(zip(source_row, result_row, strict=True)):
                if index in editable or left == right:
                    continue
                if KNOWN_TEXT_REPAIRS.get(str(left)) == right:
                    continue
                return False
        return True
    finally:
        source.close()
        result.close()


def build_cross_up_sell(
    input_path: Path,
    output_path: Path,
    report_path: Path,
    max_cross_sells: int = 4,
    max_up_sells: int = 3,
) -> dict[str, Any]:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("Plik wynikowy musi być inny niż źródłowy.")
    source_hash = sha256(input_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(input_path, output_path)

    workbook = load_workbook(output_path, data_only=False)
    try:
        sheet = workbook["Worksheet"]
        headers = [cell.value for cell in sheet[1]]
        up_column = headers.index("Up-Sells") + 1
        cross_column = headers.index("Cross-Sells") + 1
        records = load_records(sheet, find_headers(sheet))
        mark_old_duplicate_products(records)
        products = [CatalogProduct(record) for record in records]
        cross, up = select_recommendations(products, max_cross_sells, max_up_sells)
        encoding_repairs: list[dict[str, Any]] = []

        for product in products:
            sheet.cell(product.record.excel_row, up_column).value = "|".join(
                item.target_sku for item in up.get(product.sku, [])
            )
            sheet.cell(product.record.excel_row, cross_column).value = "|".join(
                item.target_sku for item in cross.get(product.sku, [])
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                before = str(cell.value or "")
                if before not in KNOWN_TEXT_REPAIRS:
                    continue
                cell.value = KNOWN_TEXT_REPAIRS[before]
                encoding_repairs.append({
                    "row": cell.row,
                    "column": headers[cell.column - 1],
                    "before": before,
                    "after": cell.value,
                })
        workbook.save(output_path)
    finally:
        workbook.close()

    validation = validate_recommendations(products, cross, up, max_cross_sells, max_up_sells)
    validation["source_file_unchanged"] = sha256(input_path) == source_hash
    validation["all_non_target_cells_preserved"] = compare_preserved_cells(input_path, output_path, "Worksheet")

    reopened = load_workbook(output_path, read_only=True, data_only=False)
    try:
        encoding_violations = []
        for sheet in reopened.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    markers = mojibake_markers(value)
                    if markers:
                        encoding_violations.append((sheet.title, compact(value), markers))
                        if len(encoding_violations) >= 20:
                            break
                if len(encoding_violations) >= 20:
                    break
        validation["workbook_polish_text_encoding_intact"] = not encoding_violations
    finally:
        reopened.close()

    if not all(validation.values()):
        raise AssertionError(json.dumps({"validation": validation, "encoding": encoding_violations}, ensure_ascii=False, indent=2))

    write_report(report_path, products, cross, up, validation, encoding_repairs, input_path, output_path)
    metrics = {
        "products": len(products),
        "products_with_cross_sell": len(cross),
        "cross_sell_relations": sum(len(items) for items in cross.values()),
        "products_with_up_sell": len(up),
        "up_sell_relations": sum(len(items) for items in up.values()),
        "old_duplicates_skipped": sum(product.record.is_old_duplicate for product in products),
        "encoding_repairs": len(encoding_repairs),
        "output": str(output_path.resolve()),
        "report": str(report_path.resolve()),
        "validation": validation,
    }
    return metrics


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Buduje cross-sell i up-sell dla katalogu Hager/Berker.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--max-cross-sells", type=int, default=4)
    parser.add_argument("--max-up-sells", type=int, default=3)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    metrics = build_cross_up_sell(
        args.input,
        args.output,
        args.report,
        max_cross_sells=args.max_cross_sells,
        max_up_sells=args.max_up_sells,
    )
    print(json.dumps(metrics, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
