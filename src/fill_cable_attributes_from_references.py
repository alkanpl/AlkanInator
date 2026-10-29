"""Uzupelnia puste atrybuty kabli z referencji: rodzenstwa, tytulu i typu kabla.

Kolejnosc zrodel dla pustego atrybutu (``dictionaries/cable_attribute_rules.yaml``):
1. rodzenstwo - ten sam SKU bazowy w innych konfekcjach (``/1m`` <-> ``/100m``),
2. tytul - napiecie (``450/750V`` -> 750, ``0,6/1kV`` -> 1000) i ``żo`` (zyla ochronna),
3. wartosc typowa dla typu kabla (YDY: PVC, nieekranowany, 750 V ...).

Istniejace wartosci zostaja; gdy tytul albo typ mowia co innego, produkt trafia do
arkusza "Konflikty" (status WARNING) bez zmiany. Wyjatki: ``Producent`` kabli to
zawsze ``Polski Producent`` (decyzja 2026-10-07), a ``Długość`` ``1m`` -> ``1 m``.

Wejscie to eksport WP All Export (albo jego kopia z ``output/``), wynik to pelna
kopia pliku - zmieniaja sie tylko kolumny ``Atrybut produktu: ...``.

Uzycie:
    py src/fill_cable_attributes_from_references.py --input <eksport.xlsx> --output <kopia.xlsx>
"""

from __future__ import annotations

import argparse
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

import yaml

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RULES = ROOT / "dictionaries" / "cable_attribute_rules.yaml"
ATTR_PREFIX = "Atrybut produktu: "
LENGTH_SUFFIX = re.compile(r"/\d+\s*m$", re.IGNORECASE)
ZO = re.compile(r"żo", re.IGNORECASE)
KNOWN_VOLTAGES = {300, 500, 750, 1000}


@dataclass
class Change:
    row_id: str
    sku: str
    title: str
    attribute: str
    old: str
    new: str
    source: str
    status: str
    note: str = ""


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def load_rules(path: Path = DEFAULT_RULES) -> dict[str, Any]:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def family_key(sku: str) -> str:
    return LENGTH_SUFFIX.sub("", sku.strip()).lower()


def voltage_from_title(title: str) -> str:
    """U z pary U0/U albo pojedynczej wartosci z jednostka; '0,6/1kV' -> 1000."""
    if re.search(r"0[,.]6\s*/\s*1\s*kV|\b1\s*kV", title, re.IGNORECASE):
        return "1000"
    for first, second in re.findall(r"\b(\d{3})\s*/\s*(\d{3})", title):
        value = max(int(first), int(second))
        if value in KNOWN_VOLTAGES:
            return str(value)
    for value in re.findall(r"\b(\d{3,4})\s*V\b", title):
        if int(value) in KNOWN_VOLTAGES:
            return value
    return ""


def type_from_title(title: str, rules: dict[str, Any]) -> str:
    for item in rules.get("type_patterns", []):
        if re.search(item["pattern"], title):
            return item["type"]
    return ""


def color_from_title(title: str, rules: dict[str, Any]) -> str:
    for item in rules.get("color_patterns", []):
        if re.search(item["pattern"], title, re.IGNORECASE):
            return item["color"]
    return ""


def lgy_voltage(cross_section: str, rules: dict[str, Any]) -> str:
    try:
        mm2 = float(cross_section.replace(",", "."))
    except ValueError:
        return ""
    cfg = rules["lgy_voltage"]
    return cfg["h05"] if mm2 <= cfg["max_h05_mm2"] else cfg["h07"]


def derived_values(title: str, cable_type: str, attrs: dict[str, str], rules: dict[str, Any]) -> dict[str, tuple[str, str]]:
    """Wartosci z tytulu i typu: atrybut -> (wartosc, zrodlo)."""
    values: dict[str, tuple[str, str]] = {}
    for name, value in (rules.get("types", {}).get(cable_type) or {}).items():
        values[name] = (str(value), f"typ {cable_type}")
    if cable_type == "LGY" and attrs.get("Przekrój żyły [mm]"):
        voltage = lgy_voltage(attrs["Przekrój żyły [mm]"], rules)
        if voltage:
            values["Napięcie [V]"] = (voltage, "typ LGY wg przekroju")
    # Napiecie z tytulu tylko dla typow, ktore w ogole maja napiecie w regulach (nie YnTKSY, PV).
    voltage = voltage_from_title(title)
    if voltage and "Napięcie [V]" in values:
        values["Napięcie [V]"] = (voltage, "tytul")
    if cable_type and ZO.search(title):
        values["Żyła ochronna"] = ("Tak", "tytul (żo)")
    return values


def length_from(sku: str, title: str) -> str:
    """Dlugosc konfekcji: koncowka SKU ('/100m'), a dla odcinka - koniec tytulu ('... bialy 1m')."""
    match = re.search(r"/(\d+)\s*m$", sku, re.IGNORECASE) or re.search(r"\s(\d+)\s?m\b(?:\s+\S+){0,2}$", title)
    return match.group(1) if match else ""


def majority(values: list[str]) -> tuple[str, bool]:
    filled = [v for v in values if v]
    if not filled:
        return "", True
    counts = Counter(filled)
    return counts.most_common(1)[0][0], len(counts) == 1


def decide(rows: list[dict[str, Any]], rules: dict[str, Any]) -> tuple[list[Change], list[Change]]:
    families: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        families[family_key(text(row.get("Sku")))].append(row)
    sibling_attrs = rules.get("sibling_attributes", [])
    polish_types = set(rules.get("polish_producer_types", []))
    changes: list[Change] = []
    conflicts: list[Change] = []

    for row in rows:
        row_id, sku, title = text(row.get("id")), text(row.get("Sku")), text(row.get("Title"))
        attrs = {k[len(ATTR_PREFIX):]: text(v) for k, v in row.items() if k.startswith(ATTR_PREFIX)}
        new: dict[str, tuple[str, str, str, str]] = {}

        def put(name: str, value: str, source: str, status: str = "OK", note: str = "") -> None:
            if name in attrs and not attrs[name] and name not in new and value:
                new[name] = (value, source, status, note)

        siblings = [s for s in families[family_key(sku)] if s is not row]
        for name in sibling_attrs:
            value, unanimous = majority([text(s.get(ATTR_PREFIX + name)) for s in siblings])
            put(name, value, "rodzenstwo", "OK" if unanimous else "WARNING",
                "" if unanimous else "rodzenstwo niezgodne - wartosc wiekszosci")

        if not attrs.get("Typ") and "Typ" not in new:
            put("Typ", type_from_title(title, rules), "tytul")
        cable_type = attrs.get("Typ") or new.get("Typ", ("",))[0]
        if cable_type:
            put("Kolor", color_from_title(title, rules), "tytul")
        current = {**attrs, **{k: v[0] for k, v in new.items()}}
        for name, (value, source) in derived_values(title, cable_type, current, rules).items():
            if current.get(name) and current[name] != value:
                conflicts.append(Change(row_id, sku, title, name, current[name], value, source, "WARNING",
                                        "wartosc produktu rozna od tytulu/typu - bez zmiany"))
            else:
                put(name, value, source)

        if cable_type in polish_types and "Producent" in attrs and attrs["Producent"] != rules["polish_producer"]:
            new["Producent"] = (rules["polish_producer"], "decyzja uzytkownika 2026-10-07", "OK", "")
        length = attrs.get("Długość", "")
        if re.fullmatch(r"\d+m", length):
            new["Długość"] = (length[:-1] + " m", "regula", "OK", "")
        elif not length and cable_type:
            meters = length_from(sku, title)
            put("Długość", f"{meters} m" if meters else "", "SKU/tytul")

        for name, (value, source, status, note) in new.items():
            changes.append(Change(row_id, sku, title, name, attrs.get(name, ""), value, source, status, note))
    return changes, conflicts


def read_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    values = list(load_workbook(path, read_only=True).worksheets[0].iter_rows(values_only=True))
    header = [text(h) for h in values[0]]
    return header, [dict(zip(header, r)) for r in values[1:]]


def write_output(input_path: Path, output_path: Path, header: list[str], changes: list[Change]) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(input_path)
    sheet = workbook.worksheets[0]
    col = {name: idx + 1 for idx, name in enumerate(header)}
    row_by_id = {text(sheet.cell(r, col["id"]).value): r for r in range(2, sheet.max_row + 1)}
    for change in changes:
        sheet.cell(row_by_id[change.row_id], col[ATTR_PREFIX + change.attribute]).value = change.new
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_control(path: Path, changes: list[Change], conflicts: list[Change], missing: list[list[str]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    for title, items in (("Zmiany", changes), ("Konflikty", conflicts)):
        sheet = workbook.active if title == "Zmiany" else workbook.create_sheet(title)
        sheet.title = title
        sheet.append(["id", "SKU", "Tytul", "Atrybut", "Bylo", "Jest" if title == "Zmiany" else "Wg tytulu/typu",
                      "source", "status", "Uwagi"])
        for c in items:
            sheet.append([c.row_id, c.sku, c.title, c.attribute, c.old, c.new, c.source, c.status, c.note])
            if c.status != "OK":
                for cell in sheet[sheet.max_row]:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
    gaps = workbook.create_sheet("Nadal puste")
    gaps.append(["id", "SKU", "Tytul", "Typ", "Puste atrybuty"])
    for item in missing:
        gaps.append(item)
    for ws in workbook.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(max(len(text(c.value)) for c in column) + 2, 70)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def still_missing(rows: list[dict[str, Any]], changes: list[Change], rules: dict[str, Any]) -> list[list[str]]:
    filled = {(c.row_id, c.attribute) for c in changes}
    report = []
    for row in rows:
        cable_type = text(row.get(ATTR_PREFIX + "Typ")) or next(
            (c.new for c in changes if c.row_id == text(row.get("id")) and c.attribute == "Typ"), "")
        expected = set((rules.get("types", {}).get(cable_type) or {})) | {"Kolor", "Liczba żył", "Przekrój żyły [mm]"}
        if not cable_type:
            continue
        empty = sorted(name for name in expected
                       if not text(row.get(ATTR_PREFIX + name)) and (text(row.get("id")), name) not in filled)
        if empty:
            report.append([text(row.get("id")), text(row.get("Sku")), text(row.get("Title")), cable_type, ", ".join(empty)])
    return report


def main(argv: list[str] | None = None) -> int:
    today = date.today().isoformat()
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", type=Path, default=ROOT / "output" / f"Kable-atrybuty_UZUPELNIONE_{today}.xlsx")
    parser.add_argument("--control", type=Path, default=ROOT / "reports" / f"Kable-atrybuty_kontrola_{today}.xlsx")
    parser.add_argument("--rules", type=Path, default=DEFAULT_RULES)
    args = parser.parse_args(argv)
    if args.output.resolve() == args.input.resolve():
        raise SystemExit("Plik wynikowy nie moze nadpisac pliku wejsciowego.")

    rules = load_rules(args.rules)
    header, rows = read_rows(args.input)
    changes, conflicts = decide(rows, rules)
    write_output(args.input, args.output, header, changes)
    missing = still_missing(rows, changes, rules)
    write_control(args.control, changes, conflicts, missing)
    by_attr = Counter(c.attribute for c in changes)
    print(f"Zmian: {len(changes)} {dict(by_attr)}; konfliktow: {len(conflicts)}; produktow z pustymi atrybutami: {len(missing)}")
    print(f"Wynik: {args.output}")
    print(f"Kontrola: {args.control}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
