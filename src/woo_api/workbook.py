"""Pliki XLSX w układzie WP All Export: arkusz "Worksheet", nagłówki w 1. wierszu."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils.escape import unescape


SHEET_NAME = "Worksheet"
INTEGER_COLUMNS = {"id", "Parent Product ID"}
EXCEL_CELL_LIMIT = 32767
ROW_NUMBER_KEY = "_row"


def merge_columns(column_lists: Iterable[list[str]]) -> list[str]:
    """Kolumny bazowe w kolejności sklepu, potem wszystkie atrybuty alfabetycznie."""
    base: list[str] = []
    attributes: set[str] = set()
    for columns in column_lists:
        for column in columns:
            if column.startswith("Atrybut Produktu: "):
                attributes.add(column)
            elif column not in base:
                base.append(column)
    return base + sorted(attributes, key=str.casefold)


def _cell_value(column: str, value: Any) -> Any:
    if value is None or value == "":
        return None
    if column in INTEGER_COLUMNS:
        try:
            return int(value)
        except (TypeError, ValueError):
            return str(value)
    return str(value)


def write_products_xlsx(path: str | Path, columns: list[str], rows: list[dict[str, Any]]) -> list[str]:
    """Zapisuje produkty; zwraca ostrzeżenia o komórkach dłuższych niż limit Excela."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(columns)
    warnings: list[str] = []
    for row in rows:
        values = []
        for column in columns:
            value = _cell_value(column, row.get(column))
            if isinstance(value, str) and len(value) > EXCEL_CELL_LIMIT:
                warnings.append(f"id {row.get('id')}: kolumna {column} ma {len(value)} znaków - Excel ucina powyżej {EXCEL_CELL_LIMIT}")
            values.append(value)
        sheet.append(values)
    sheet.freeze_panes = "B2"
    workbook.save(path)
    return warnings


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value)


def read_products_xlsx(path: str | Path, sheet: str = "") -> tuple[list[str], list[dict[str, Any]]]:
    """Czyta arkusz; każdy wiersz dostaje numer wiersza Excela pod kluczem "_row"."""
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        iterator = worksheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if not header_row:
            return [], []
        headers = [_text(value).strip() for value in header_row]
        rows: list[dict[str, Any]] = []
        for number, values in enumerate(iterator, start=2):
            if values is None or all(value is None or _text(value).strip() == "" for value in values):
                continue
            row: dict[str, Any] = {ROW_NUMBER_KEY: number}
            for header, value in zip(headers, values):
                if header:
                    # Excel stores "\r" and other control characters as "_x000D_"-style escapes.
                    row[header] = unescape(value) if isinstance(value, str) else _text(value)
            rows.append(row)
        return [header for header in headers if header], rows
    finally:
        workbook.close()
