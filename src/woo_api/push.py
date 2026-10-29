"""Wysyłka zmian z pliku XLSX do sklepu: zawsze najpierw dry-run całego pliku, raport, dopiero potem zapis."""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Callable

import requests
from openpyxl import Workbook
from PIL import Image, UnidentifiedImageError

from woo_api.client import WooBridgeClient
from woo_api.workbook import EXCEL_CELL_LIMIT, ROW_NUMBER_KEY, read_products_xlsx


IDENTITY_COLUMNS = ("id", "SKU", "Post Modified Date")
# Kolumny tylko do odczytu, których nie ma sensu wysyłać (sklep i tak je pomija).
READ_ONLY_COLUMNS = {
    "Product Type",
    "Parent Product ID",
    "Regular Price",
    "Sale Price",
    "Stock",
    "Stock Status",
    "hwp_product_gtin",
    "Permalink",
    "Product Attributes",
}
MAX_BATCH_SIZE = 200
# Zdjecia produktow zawsze 1000x1000 JPG na bialym tle (CLAUDE.md). Sklep pobiera obce
# URL-e z Image URL bez obrobki, wiec format sprawdza dry-run, zanim cokolwiek trafi do mediow.
IMAGE_COLUMN = "Image URL"
IMAGE_SIZE = (1000, 1000)
WHITE_CORNER_MIN = 240


@dataclass
class PushOptions:
    apply: bool = False
    columns: list[str] = field(default_factory=list)
    allow_new_terms: bool = False
    skip_errors: bool = False
    batch_size: int = 100
    changeset: str = ""
    sheet: str = ""
    check_images: bool = True


@dataclass
class PushResult:
    dry_run: dict[str, Any]
    applied: dict[str, Any] | None
    changeset: str
    report_paths: list[Path]
    blocked: bool


def make_changeset_id(source: str | Path, now: datetime | None = None) -> str:
    stem = Path(source).stem.translate(str.maketrans("łŁ", "lL"))  # NFKD does not decompose ł
    stem = unicodedata.normalize("NFKD", stem).encode("ascii", "ignore").decode("ascii")
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", stem).strip("-")[:30] or "plik"
    return f"ai-{(now or datetime.now()).strftime('%Y%m%d-%H%M%S')}-{stem}"


def prepare_rows(headers: list[str], rows: list[dict[str, Any]], columns: list[str]) -> list[dict[str, Any]]:
    if "id" not in headers and "SKU" not in headers:
        raise ValueError("Plik musi mieć kolumnę id albo SKU.")
    missing = [column for column in columns if column not in headers]
    if missing:
        raise ValueError("W pliku nie ma kolumn: " + ", ".join(missing))

    keep = set(columns) if columns else {header for header in headers if header not in READ_ONLY_COLUMNS}
    keep.update(column for column in IDENTITY_COLUMNS if column in headers)

    payload: list[dict[str, Any]] = []
    seen: dict[str, int] = {}
    problems: list[str] = []
    for row in rows:
        key = f"id:{row['id']}" if row.get("id") else f"sku:{row.get('SKU', '')}"
        if key in {"id:", "sku:"}:
            problems.append(f"wiersz {row[ROW_NUMBER_KEY]}: brak id i SKU")
            continue
        if key in seen:
            problems.append(f"wiersz {row[ROW_NUMBER_KEY]}: ten sam produkt co w wierszu {seen[key]} ({key})")
            continue
        seen[key] = row[ROW_NUMBER_KEY]
        payload.append({column: value for column, value in row.items() if column in keep or column == ROW_NUMBER_KEY})
    if problems:
        raise ValueError("Popraw plik przed wysyłką:\n" + "\n".join(problems[:50]))
    return payload


def merge_results(parts: list[dict[str, Any]]) -> dict[str, Any]:
    merged: dict[str, Any] = {
        "changeset": "",
        "summary": {},
        "products": [],
        "new_terms": [],
        "written_columns": [],
        "read_only_columns": [],
        "ignored_columns": [],
    }
    seen_terms: set[tuple[str, str]] = set()
    for part in parts:
        merged["changeset"] = part.get("changeset") or merged["changeset"]
        merged["products"].extend(part.get("products", []))
        for key, value in part.get("summary", {}).items():
            merged["summary"][key] = merged["summary"].get(key, 0) + value
        for term in part.get("new_terms", []):
            marker = (term["taxonomy"], term["name"].casefold())
            if marker not in seen_terms:
                seen_terms.add(marker)
                merged["new_terms"].append(term)
        for key in ("written_columns", "read_only_columns", "ignored_columns"):
            for column in part.get(key, []):
                if column not in merged[key]:
                    merged[key].append(column)
    return merged


def image_problems(body: bytes) -> tuple[list[str], list[str]]:
    """(bledy, ostrzezenia) dla pliku zdjecia: ma byc JPG 1000x1000 RGB, rogi biale."""
    try:
        with Image.open(BytesIO(body)) as image:
            image.load()
            errors = []
            if image.format != "JPEG":
                errors.append(f"format {image.format}, a ma być JPG")
            if image.size != IMAGE_SIZE:
                errors.append(f"{image.width}x{image.height}, a ma być {IMAGE_SIZE[0]}x{IMAGE_SIZE[1]}")
            if image.mode not in ("RGB", "L"):
                errors.append(f"tryb kolorów {image.mode}, a ma być RGB")
            rgb = image.convert("RGB")
            corners = [(0, 0), (rgb.width - 1, 0), (0, rgb.height - 1), (rgb.width - 1, rgb.height - 1)]
            dark = [xy for xy in corners if min(rgb.getpixel(xy)) < WHITE_CORNER_MIN]
    except (UnidentifiedImageError, OSError):
        return ["to nie jest plik graficzny"], []
    warnings = ["rogi zdjęcia nie są białe - sprawdź tło"] if dark else []
    return errors, warnings


def fetch_image(url: str, *, verify: bool = True, timeout: float = 60.0) -> bytes:
    # Osobne zapytanie, nie sesja klienta - token sklepu nie moze isc do obcych hostow.
    response = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=timeout, verify=verify)
    response.raise_for_status()
    return response.content


def check_images(result: dict[str, Any], fetch: Callable[[str], bytes]) -> int:
    """Sprawdza zdjecia, ktore zapis doda produktom; produkt ze zlym zdjeciem dostaje blad. Zwraca liczbe plikow."""
    checked: dict[str, tuple[list[str], list[str]]] = {}
    for product in result.get("products", []):
        for change in product.get("changes", []):
            if change.get("column") != IMAGE_COLUMN:
                continue
            old = set(_split_urls(change.get("old")))
            for url in _split_urls(change.get("new")):
                if url in old:
                    continue
                if url not in checked:
                    try:
                        checked[url] = image_problems(fetch(url))
                    except requests.RequestException as exc:
                        checked[url] = ([f"nie udało się pobrać ({exc})"], [])
                errors, warnings = checked[url]
                if errors:
                    product.setdefault("errors", []).append(f"{IMAGE_COLUMN}: {url}: {', '.join(errors)}")
                    product["status"] = "error"
                for warning in warnings:
                    product.setdefault("warnings", []).append(f"{IMAGE_COLUMN}: {url}: {warning}")
    products = result.get("products", [])
    summary = result.setdefault("summary", {})
    summary["errors"] = sum(bool(p.get("errors")) for p in products)
    summary["warnings"] = sum(bool(p.get("warnings")) for p in products)
    summary["changed"] = sum(p.get("status") in ("changed", "applied") for p in products)
    summary["images_checked"] = len(checked)
    return len(checked)


def _split_urls(value: Any) -> list[str]:
    return [part.strip() for part in str(value or "").split("|") if part.strip()]


def _clip(value: Any) -> Any:
    if isinstance(value, str) and len(value) > EXCEL_CELL_LIMIT - 100:
        return value[: EXCEL_CELL_LIMIT - 100] + " …[ucięte w raporcie]"
    return value


def write_report(path: Path, result: dict[str, Any], *, mode: str, source: Path, changeset: str = "") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Podsumowanie"
    summary.append(["Tryb", mode])
    summary.append(["Plik", str(source)])
    if changeset:
        summary.append(["Zapis (changeset)", changeset])
        summary.append(["Cofnięcie", f"py src/woo_api_cli.py revert {changeset} --apply"])
    for key, value in result.get("summary", {}).items():
        summary.append([key, value])
    summary.append(["Zapisywane kolumny", ", ".join(result.get("written_columns", []))])
    summary.append(["Kolumny tylko do odczytu", ", ".join(result.get("read_only_columns", []))])
    summary.append(["Kolumny nieznane (pominięte)", ", ".join(result.get("ignored_columns", []))])

    changes = workbook.create_sheet("Zmiany")
    changes.append(["wiersz", "id", "SKU", "tytuł", "status", "kolumna", "było", "będzie"])
    problems = workbook.create_sheet("Problemy")
    problems.append(["wiersz", "id", "SKU", "tytuł", "rodzaj", "opis"])
    for product in result.get("products", []):
        base = [product.get("row"), product.get("id"), product.get("sku"), product.get("title")]
        for change in product.get("changes", []):
            changes.append(base + [product.get("status"), change["column"], _clip(change["old"]), _clip(change["new"])])
        for error in product.get("errors", []):
            problems.append(base + ["błąd", error])
        for warning in product.get("warnings", []):
            problems.append(base + ["ostrzeżenie", warning])

    terms = workbook.create_sheet("Nowe terminy")
    terms.append(["taksonomia", "nazwa"])
    for term in result.get("new_terms", []):
        terms.append([term["taxonomy"], term["name"]])

    for sheet in workbook.worksheets[1:]:
        sheet.freeze_panes = "A2"
    workbook.save(path)
    return path


def _chunks(items: list[dict[str, Any]], size: int) -> list[list[dict[str, Any]]]:
    return [items[index : index + size] for index in range(0, len(items), size)]


def run_push(
    client: WooBridgeClient,
    input_path: str | Path,
    options: PushOptions,
    *,
    report_dir: str | Path,
    log: Callable[[str], None] = print,
    now: datetime | None = None,
    image_fetcher: Callable[[str], bytes] | None = None,
) -> PushResult:
    input_path = Path(input_path)
    batch_size = max(1, min(MAX_BATCH_SIZE, options.batch_size))
    headers, rows = read_products_xlsx(input_path, options.sheet)
    payload = prepare_rows(headers, rows, options.columns)
    stamp = (now or datetime.now()).strftime("%Y%m%d-%H%M%S")
    report_dir = Path(report_dir)

    parts = []
    for number, chunk in enumerate(_chunks(payload, batch_size), start=1):
        log(f"Dry-run {number}/{-(-len(payload) // batch_size)} ({len(chunk)} wierszy)")
        parts.append(client.sync(chunk, dry_run=True, allow_new_terms=options.allow_new_terms, columns=options.columns))
    dry = merge_results(parts)
    if options.check_images:
        verify = getattr(getattr(client, "session", None), "verify", True)
        fetch = image_fetcher or (lambda url: fetch_image(url, verify=verify))
        checked = check_images(dry, fetch)
        if checked:
            log(f"Sprawdzono zdjęcia: {checked} (wymóg: JPG {IMAGE_SIZE[0]}x{IMAGE_SIZE[1]} na białym tle)")
    reports = [write_report(report_dir / f"{stamp}_{input_path.stem}_dry-run.xlsx", dry, mode="DRY_RUN", source=input_path)]

    if not options.apply:
        return PushResult(dry, None, "", reports, blocked=False)

    if dry["summary"].get("errors") and not options.skip_errors:
        log("Zapis wstrzymany: dry-run ma błędy (arkusz Problemy). Popraw plik albo użyj --skip-errors, żeby zapisać resztę.")
        return PushResult(dry, None, "", reports, blocked=True)

    changed_rows = {product["row"] for product in dry["products"] if product["status"] == "changed"}
    to_apply = [row for row in payload if row[ROW_NUMBER_KEY] in changed_rows]
    if not to_apply:
        log("Brak zmian do zapisania.")
        return PushResult(dry, None, "", reports, blocked=False)

    changeset = options.changeset or make_changeset_id(input_path, now)
    parts = []
    for number, chunk in enumerate(_chunks(to_apply, batch_size), start=1):
        log(f"Zapis {number}/{-(-len(to_apply) // batch_size)} ({len(chunk)} produktów)")
        parts.append(
            client.sync(
                chunk,
                dry_run=False,
                allow_new_terms=options.allow_new_terms,
                columns=options.columns,
                changeset=changeset,
            )
        )
    applied = merge_results(parts)
    reports.append(write_report(report_dir / f"{stamp}_{input_path.stem}_apply.xlsx", applied, mode="APPLY", source=input_path, changeset=changeset))
    return PushResult(dry, applied, changeset, reports, blocked=False)


def run_revert(
    client: WooBridgeClient,
    changeset: str,
    *,
    apply: bool,
    report_dir: str | Path,
    now: datetime | None = None,
) -> tuple[dict[str, Any], Path]:
    result = client.revert(changeset, dry_run=not apply)
    stamp = (now or datetime.now()).strftime("%Y%m%d-%H%M%S")
    mode = "REVERT_APPLY" if apply else "REVERT_DRY_RUN"
    report = write_report(
        Path(report_dir) / f"{stamp}_revert_{changeset}_{'apply' if apply else 'dry-run'}.xlsx",
        result,
        mode=mode,
        source=Path(changeset),
        changeset=result.get("changeset", ""),
    )
    return result, report
