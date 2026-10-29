"""Uzupelnia wymiary rur karbowanych (AROT) w eksporcie Woo.

Zrodla: dictionaries/corrugated_pipe_dimensions_verified.yaml (karty producentow),
awaryjnie nazwa ("AROT 32/25", "AROT 50 N"). Zapisuje atrybuty sklepu
"Srednica zewnetrzna" i "Srednica wewnetrzna" ("50 mm"), normalizuje "Dlugosc"
("25 m" zostaje w metrach - to dlugosc krazka). Woo Length/Width/Height (cm)
= wymiary krazka, a Weight (kg) = masa krazka - wg producenta (decyzja
uzytkownika 2026-10-05: dane producenta nadpisuja stare wartosci sklepu,
poprzednie wartosci trafiaja do raportu). Plik wejsciowy nie jest nadpisywany.
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import re
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import PatternFill

PREFIX = "Atrybut produktu: "
COL_OUT = PREFIX + "Średnica zewnętrzna"
COL_IN = PREFIX + "Średnica wewnętrzna"
COL_L = PREFIX + "Długość"
WOO = ("Length", "Width", "Height")
WEIGHT = "Weight"


def fmt(n: float) -> str:
    return (f"{n:.2f}".rstrip("0").rstrip(".")).replace(".", ",")


def num(v):
    try:
        return float(str(v).replace(",", ".")) if v not in (None, "") else None
    except ValueError:
        return None


def parse_title(title: str) -> dict:
    m = re.search(r"AROT\s+(\d{2,3})\s*/\s*(\d{2,3})", title, re.I)
    if m:
        return {"outer": int(m.group(1)), "inner": int(m.group(2))}
    m = re.search(r"AROT\s+(\d{2,3})\s*N\b", title, re.I)
    return {"outer": int(m.group(1))} if m else {}


def main() -> None:
    today = dt.date.today().isoformat()
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--input", default="input/Rury-karbowane.xlsx")
    ap.add_argument("--verified", default="dictionaries/corrugated_pipe_dimensions_verified.yaml")
    ap.add_argument("--output", default=f"output/Rury-karbowane_WYMIARY_UZUPELNIONE_{today}.xlsx")
    ap.add_argument("--report", default=f"reports/Rury-karbowane_WYMIARY_KONTROLA_{today}.xlsx")
    args = ap.parse_args()

    verified = yaml.safe_load(Path(args.verified).read_text(encoding="utf-8"))["products"]
    wb = openpyxl.load_workbook(args.input)
    ws = wb.active
    header = [c.value for c in ws[1]]
    for name in (COL_OUT, COL_IN):
        if name not in header:
            header.append(name)
            ws.cell(row=1, column=len(header), value=name)
    col = {n: i + 1 for i, n in enumerate(header)}
    fill = PatternFill("solid", fgColor="FFF2CC")

    report = []
    for r in range(2, ws.max_row + 1):
        get = lambda n: ws.cell(r, col[n]).value
        sku, title = str(get("Sku") or ""), str(get("Title") or "")
        old = {k: get(k) for k in (COL_OUT, COL_IN, COL_L, WEIGHT) + WOO}
        new = dict(old)
        notes, status = [], "OK"
        v = verified.get(sku)
        t = parse_title(title)
        data = v or t
        source = f"karta producenta: {v['source']}" if v else ("nazwa" if t else "-")
        if v and v.get("note"):
            notes.append(v["note"])
        if t.get("outer") and v and t["outer"] != v["outer"]:
            notes.append(f"nazwa podaje Ø{t['outer']}, producent Ø{v['outer']}")
            status = "WARNING"
        if data.get("outer"):
            new[COL_OUT] = f"{fmt(data['outer'])} mm"
        if data.get("inner"):
            new[COL_IN] = f"{fmt(data['inner'])} mm"
            if v and v.get("inner_min"):
                notes.append(f"min. średnica wewnętrzna wg karty: {fmt(v['inner_min'])} mm")
        else:
            notes.append("brak średnicy wewnętrznej")
            status = "DO_SPRAWDZENIA"
        m = re.fullmatch(r"\s*(\d+(?:[.,]\d+)?)\s*(m|mb)\s*", str(old[COL_L] or ""), re.I)
        if m:
            new[COL_L] = f"{fmt(num(m.group(1)))} m"
        coil = v.get("coil") if v else None
        existing = [num(old[k]) or 0 for k in WOO]
        if coil:
            want = [coil[0], coil[0], coil[1]]
            if any(existing) and existing != want:
                notes.append(f"Woo L/W/H {'/'.join(fmt(x) for x in existing)} cm zastąpione krążkiem wg producenta")
            new.update(dict(zip(WOO, want)))
            if v.get("coil_kg"):
                kg = v["coil_kg"]
                new[WEIGHT] = int(kg) if float(kg).is_integer() else kg
        elif not any(existing):
            notes.append("brak wymiarów krążka")
            if status == "OK":
                status = "DO_SPRAWDZENIA"

        for k, val in new.items():
            if val != old[k]:
                ws.cell(r, col[k]).value = val
                ws.cell(r, col[k]).fill = fill
        report.append({
            "id": get("id"), "Sku": sku, "Title": title,
            "Średnica zewnętrzna": new[COL_OUT], "Średnica wewnętrzna": new[COL_IN],
            "Długość przed": old[COL_L], "Długość po": new[COL_L],
            "Woo L/W/H przed": "/".join(str(old[k]) for k in WOO),
            "Woo L/W/H po": "/".join(str(new[k]) for k in WOO),
            "Weight przed": old[WEIGHT], "Weight po": new[WEIGHT],
            "Krążek wg producenta [cm]": f"Ø{coil[0]} x {coil[1]}, {v['coil_kg']} kg" if coil else None,
            "source": source,
            "confidence": {"OK": 0.95, "WARNING": 0.8, "DO_SPRAWDZENIA": 0.5}[status],
            "status": status, "uwagi": "; ".join(notes),
        })

    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    rb = openpyxl.Workbook()
    rs = rb.active
    rs.title = "Wymiary"
    rs.append(list(report[0]))
    for row in report:
        rs.append(list(row.values()))
    rs.freeze_panes = "D2"
    rs.auto_filter.ref = rs.dimensions
    rb.save(args.report)
    print(f"rows={len(report)} {dict(collections.Counter(x['status'] for x in report))} -> {args.output} | {args.report}")


if __name__ == "__main__":
    main()
