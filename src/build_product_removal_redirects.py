"""Usuwanie produktow ze sklepu: przekierowania 301, status draft i raport.

Wejscie to eksport WP All Export z dopisanymi recznie kolumnami:

- ``Do usunięcia`` - ``tak`` / ``x`` / ``1`` oznacza produkt do usuniecia,
- ``Zamiennik`` - SKU, ID albo URL produktu, na ktory ma prowadzic redirect
  (puste = redirect na kategorie produktu).

Eksport musi miec tez ``Permalink`` (albo ``Slug``) i ``Product categories``
(nazwy kategorii, hierarchia ``Rodzic > Dziecko``, kilka kategorii po ``|``).

Cel przekierowania:
1. zamiennik - opublikowany produkt z eksportu (SKU/ID) albo podany URL;
   zamiennik, ktory sam jest do usuniecia, prowadzi dalej na swoj cel;
2. bez zamiennika - najglebsza kategoria produktu (remis: nizsze term_id,
   tak samo jak wtyczka alkan-product-category-redirect); kategoria, ktora po
   usunieciu zostalaby pusta, ustepuje rodzicowi (WARNING).
Kazdy cel przechodzi przez istniejace przekierowania sklepu (eksport JSON
wtyczki Redirection), zeby nie tworzyc lancuchow 301 -> 301.

Produkty bez zamiennika dostaja w raporcie propozycje zamiennikow (ta sama
kategoria, zgodne atrybuty i tytul) ze statusem DO_SPRAWDZENIA - nic nie
trafia do importu bez decyzji uzytkownika.

Produkt, dla ktorego nie da sie ustalic bezpiecznego celu (niejednoznaczny
SKU zamiennika, brak URL, petla), ma status DO_SPRAWDZENIA: nie dostaje
przekierowania ani statusu draft, dopoki dane nie zostana poprawione.

Wyniki:
- kopia eksportu z ``Status`` = ``draft`` dla usuwanych produktow (WP All Import),
- CSV do importu w Redirection (Narzedzia -> Import, wybierz grupe),
- JSON w formacie eksportu Redirection (import tworzy grupe "Alkan - usunięte produkty"),
- skoroszyt kontrolny w ``reports/``.

Uzycie:
    py src/build_product_removal_redirects.py --input input/usuwanie.xlsx \
        --redirects input/redirection-export_2026-10-12.json --categories input/kategorie-produkcja.xlsx
"""

from __future__ import annotations

import argparse
import csv
import html
import json
import re
import urllib.error
import urllib.request
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import unquote, urlparse

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CATEGORY_SNAPSHOT = ROOT / "dictionaries" / "woo_product_categories_snapshot.tsv"
SITE_URL = "https://alkan.pl"
PRODUCT_BASE = "sklep"
CATEGORY_BASE = "kategoria-produktu"
GROUP_NAME = "Alkan - usunięte produkty"

OK, WARNING, CHECK = "OK", "WARNING", "DO_SPRAWDZENIA"
SEVERITY = {OK: 0, WARNING: 1, CHECK: 2}

DELETE_COLUMNS = ("Do usunięcia", "Do usuniecia", "Usunac", "Usunąć")
REPLACEMENT_COLUMNS = ("Zamiennik", "Zamiennik SKU", "Zamiennik (SKU/ID/URL)")
PERMALINK_COLUMNS = ("Permalink", "URL", "Link", "Product URL")
SLUG_COLUMNS = ("Slug", "post_name", "Product Slug")
CATEGORY_COLUMNS = ("Product categories", "Product Categories", "Kategorie produktu", "Kategorie", "Categories")
SKU_COLUMNS = ("Sku", "SKU")
TRUE_VALUES = {"tak", "t", "x", "1", "yes", "y", "true", "usun", "usuń", "do usunięcia", "do usuniecia"}
ATTR_PREFIX = "atrybut produktu:"
# Identyfikatory i dane GPSR nie mowia nic o podobienstwie produktow.
SKIP_SUGGESTION_ATTRS = {"symbol", "dane producenta"}
SUGGESTION_LIMIT = 3
SUGGESTION_MIN_SCORE = 0.4


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def find_column(header: Iterable[str], candidates: Iterable[str]) -> str:
    lookup = {text(h).lower(): text(h) for h in header}
    for candidate in candidates:
        if candidate.lower() in lookup:
            return lookup[candidate.lower()]
    return ""


def worst(*statuses: str) -> str:
    return max(statuses, key=lambda s: SEVERITY[s])


# --- adresy -----------------------------------------------------------------

def site_path(value: str, site_hosts: set[str]) -> str | None:
    """Sciezka bez koncowego ukosnika, malymi literami; None dla obcej domeny."""
    value = text(value)
    if not value:
        return None
    parsed = urlparse(value)
    if parsed.netloc and parsed.netloc.lower() not in site_hosts:
        return None
    path = unquote(parsed.path or "/").strip().lower()
    if not path.startswith("/"):
        path = "/" + path
    if len(path) > 1:
        path = path.rstrip("/")
    return path


def as_url(path: str, site_url: str) -> str:
    return site_url.rstrip("/") + (path if path == "/" else path + "/")


def hosts_for(site_url: str) -> set[str]:
    host = urlparse(site_url).netloc.lower()
    bare = host[4:] if host.startswith("www.") else host
    return {bare, "www." + bare}


# --- kategorie --------------------------------------------------------------

@dataclass
class Category:
    term_id: int
    parent_id: int
    name: str
    slug: str
    count: int | None = None


def _int(value: Any) -> int | None:
    value = text(value)
    return int(value) if re.fullmatch(r"\d+", value) else None


def categories_from_records(records: list[dict[str, Any]]) -> dict[int, Category]:
    """Tabela kategorii z dowolnego zrodla: zrzut TSV, eksport WP All Export, CSV z phpMyAdmin."""
    if not records:
        return {}
    header = list(records[0].keys())
    col_id = find_column(header, ("term_id", "Term ID", "ID", "id"))
    col_parent = find_column(header, ("parent_id", "Parent ID", "Term Parent ID", "Parent", "parent", "Term Parent"))
    col_parent_slug = find_column(header, ("Parent Slug", "Term Parent Slug", "parent_slug"))
    col_name = find_column(header, ("name", "Term Name", "Name", "Title"))
    col_slug = find_column(header, ("slug", "Term Slug", "Slug"))
    col_count = find_column(header, ("count", "Count", "Term Count"))
    if not (col_id and col_name and col_slug):
        raise SystemExit(f"Plik kategorii musi miec kolumny ID, nazwy i sluga (sa: {header})")
    cats: dict[int, Category] = {}
    raw_parent: dict[int, str] = {}
    for record in records:
        term_id = _int(record.get(col_id))
        if term_id is None:
            continue
        parent_value = text(record.get(col_parent)) if col_parent else ""
        if col_parent_slug and not parent_value:
            parent_value = text(record.get(col_parent_slug))
        raw_parent[term_id] = parent_value
        count = _int(record.get(col_count)) if col_count else None
        cats[term_id] = Category(term_id, 0, html.unescape(text(record.get(col_name))), text(record.get(col_slug)).lower(), count)
    by_slug = {c.slug: c.term_id for c in cats.values()}
    by_name = defaultdict(list)
    for c in cats.values():
        by_name[c.name.lower()].append(c.term_id)
    for term_id, parent_value in raw_parent.items():
        parent_id = _int(parent_value)
        if parent_id is None and parent_value:
            parent_id = by_slug.get(parent_value.lower())
            if parent_id is None and len(by_name.get(parent_value.lower(), [])) == 1:
                parent_id = by_name[parent_value.lower()][0]
        cats[term_id].parent_id = parent_id if parent_id in cats else 0
    return cats


def load_table(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        values = list(load_workbook(path, read_only=True).worksheets[0].iter_rows(values_only=True))
        header = [text(h) for h in values[0]]
        return [dict(zip(header, row)) for row in values[1:]]
    raw = path.read_text(encoding="utf-8-sig").splitlines()
    lines = [line for line in raw if line.strip() and not line.startswith("#")]
    delimiter = "\t" if "\t" in lines[0] else (";" if lines[0].count(";") > lines[0].count(",") else ",")
    return list(csv.DictReader(lines, delimiter=delimiter))


def category_chain(term_id: int, cats: dict[int, Category]) -> list[Category]:
    chain: list[Category] = []
    seen: set[int] = set()
    while term_id in cats and term_id not in seen:
        seen.add(term_id)
        chain.append(cats[term_id])
        term_id = cats[term_id].parent_id
    return list(reversed(chain))


def category_path(term_id: int, cats: dict[int, Category], base: str = CATEGORY_BASE) -> str:
    return "/" + "/".join([base] + [c.slug for c in category_chain(term_id, cats)])


def category_label(term_id: int, cats: dict[int, Category]) -> str:
    return " > ".join(c.name for c in category_chain(term_id, cats))


def descendants(term_id: int, cats: dict[int, Category]) -> set[int]:
    children = defaultdict(list)
    for c in cats.values():
        children[c.parent_id].append(c.term_id)
    found, stack = {term_id}, [term_id]
    while stack:
        for child in children[stack.pop()]:
            if child not in found:
                found.add(child)
                stack.append(child)
    return found


def parse_product_categories(value: Any) -> list[list[str]]:
    """``A > B|C > D`` z WP All Export -> [["A", "B"], ["C", "D"]]."""
    paths = []
    for part in html.unescape(text(value)).split("|"):
        names = [n.strip() for n in re.split(r"\s*>\s*", part) if n.strip()]
        if names:
            paths.append(names)
    return paths


def match_category_paths(paths: list[list[str]], cats: dict[int, Category]) -> tuple[list[int], list[str]]:
    """Zamienia sciezki nazw na term_id. Zwraca (znalezione, nierozpoznane)."""
    by_path = {tuple(c.name.lower() for c in category_chain(t, cats)): t for t in cats}
    by_name = defaultdict(list)
    for c in cats.values():
        by_name[c.name.lower()].append(c.term_id)
    found, missing = [], []
    for names in paths:
        key = tuple(n.lower() for n in names)
        term_id = by_path.get(key)
        if term_id is None and len(by_name.get(key[-1], [])) == 1:
            term_id = by_name[key[-1]][0]
        if term_id is None:
            missing.append(" > ".join(names))
        elif term_id not in found:
            found.append(term_id)
    return found, missing


# --- istniejace przekierowania ---------------------------------------------

@dataclass
class Redirect:
    id: int
    source: str
    source_path: str
    target: str
    target_path: str | None
    group_id: int


def load_redirects(path: Path, site_hosts: set[str]) -> list[Redirect]:
    data = json.loads(path.read_text(encoding="utf-8"))
    items = []
    for raw in data.get("redirects", []):
        if not raw.get("enabled", True) or raw.get("regex") or raw.get("match_type") != "url" or raw.get("action_type") != "url":
            continue
        source = text(raw.get("url"))
        # Reguly z parametrami zapytania (?show_products=...) dotycza innego adresu niz goly produkt.
        if "?" in source:
            continue
        action = raw.get("action_data") or {}
        target = text(action.get("url") if isinstance(action, dict) else action)
        source_path = site_path(source, site_hosts)
        if not source_path:
            continue
        items.append(Redirect(int(raw.get("id") or 0), source, source_path, target, site_path(target, site_hosts), int(raw.get("group_id") or 0)))
    return items


class RedirectIndex:
    def __init__(self, redirects: list[Redirect]):
        self.redirects = redirects
        self.by_source: dict[str, Redirect] = {}
        for item in redirects:
            self.by_source.setdefault(item.source_path, item)

    def follow(self, path: str, max_hops: int = 10) -> tuple[str, list[Redirect], bool]:
        """Zwraca (koncowa sciezka, przebyte reguly, petla)."""
        hops: list[Redirect] = []
        seen = {path}
        while path in self.by_source and len(hops) < max_hops:
            rule = self.by_source[path]
            if rule.target_path is None:
                hops.append(rule)
                return path, hops, False
            hops.append(rule)
            path = rule.target_path
            if path in seen:
                return path, hops, True
            seen.add(path)
        return path, hops, False

    def pointing_to(self, path: str) -> list[Redirect]:
        return [r for r in self.redirects if r.target_path == path]

    def chains(self) -> list[tuple[Redirect, str, int]]:
        """Istniejace reguly, ktorych cel sam jest przekierowany (lancuch 301 -> 301)."""
        result = []
        for rule in self.redirects:
            if rule.target_path and rule.target_path in self.by_source and rule.target_path != rule.source_path:
                final, hops, _loop = self.follow(rule.target_path)
                result.append((rule, final, len(hops) + 1))
        return result


# --- decyzje ----------------------------------------------------------------

@dataclass
class Decision:
    row_id: str
    sku: str
    title: str
    status_before: str
    source_path: str = ""
    target_path: str = ""
    target_kind: str = ""
    target_label: str = ""
    status: str = OK
    notes: list[str] = field(default_factory=list)
    add_redirect: bool = False
    set_draft: bool = False

    def flag(self, status: str, note: str) -> None:
        self.status = worst(self.status, status)
        self.notes.append(note)


@dataclass
class Suggestion:
    deleted_id: str
    deleted_sku: str
    rank: int
    row_id: str
    sku: str
    title: str
    url: str
    score: float
    differences: str


@dataclass
class Plan:
    decisions: list[Decision]
    suggestions: list[Suggestion]
    updates: list[tuple[Redirect, Decision]]
    chains: list[tuple[Redirect, str, int]]


@dataclass
class Context:
    header: list[str]
    rows: list[dict[str, Any]]
    cats: dict[int, Category]
    index: RedirectIndex
    site_url: str
    product_base: str
    category_base: str

    def __post_init__(self) -> None:
        h = self.header
        self.col_delete = find_column(h, DELETE_COLUMNS)
        self.col_repl = find_column(h, REPLACEMENT_COLUMNS)
        self.col_link = find_column(h, PERMALINK_COLUMNS)
        self.col_slug = find_column(h, SLUG_COLUMNS)
        self.col_cats = find_column(h, CATEGORY_COLUMNS)
        self.col_sku = find_column(h, SKU_COLUMNS)
        self.col_status = find_column(h, ("Status", "post_status"))
        self.col_type = find_column(h, ("Product Type", "Type"))
        self.col_title = find_column(h, ("Title", "Nazwa"))
        self.attr_cols = [c for c in h if c.lower().startswith(ATTR_PREFIX)
                          and c[len(ATTR_PREFIX):].strip().lower() not in SKIP_SUGGESTION_ATTRS]
        self.hosts = hosts_for(self.site_url)
        self.by_id = {text(r.get("id")): r for r in self.rows}
        self.by_sku = defaultdict(list)
        for r in self.rows:
            if self.col_sku and text(r.get(self.col_sku)):
                self.by_sku[text(r.get(self.col_sku)).lower()].append(r)
        self.marked = {text(r.get("id")) for r in self.rows if self.is_marked(r)}
        self.row_cats = {text(r.get("id")): match_category_paths(parse_product_categories(r.get(self.col_cats)), self.cats)
                         for r in self.rows} if self.col_cats else {}

    def is_marked(self, row: dict[str, Any]) -> bool:
        return bool(self.col_delete) and text(row.get(self.col_delete)).lower() in TRUE_VALUES

    def status_of(self, row: dict[str, Any]) -> str:
        return text(row.get(self.col_status)).lower() if self.col_status else ""

    def is_variation(self, row: dict[str, Any]) -> bool:
        return "variation" in text(row.get(self.col_type)).lower() and "variable" not in text(row.get(self.col_type)).lower()

    def product_path(self, row: dict[str, Any]) -> str:
        link = text(row.get(self.col_link)) if self.col_link else ""
        if link and "?" not in link:
            path = site_path(link, self.hosts)
            if path:
                return path
        slug = text(row.get(self.col_slug)) if self.col_slug else ""
        if slug:
            return f"/{self.product_base}/{slug.lower()}"
        return ""

    def remaining_in_category(self, term_id: int) -> int:
        """Opublikowane produkty z eksportu w kategorii (z podkategoriami), ktore zostaja."""
        branch = descendants(term_id, self.cats)
        count = 0
        for r in self.rows:
            rid = text(r.get("id"))
            if rid in self.marked or self.status_of(r) not in ("publish", "") or self.is_variation(r):
                continue
            if branch & set(self.row_cats.get(rid, ([], []))[0]):
                count += 1
        return count

    def marked_in_category(self, term_id: int) -> int:
        branch = descendants(term_id, self.cats)
        return sum(1 for rid in self.marked if branch & set(self.row_cats.get(rid, ([], []))[0]))


def resolve_replacement(value: str, ctx: Context) -> tuple[dict[str, Any] | None, str | None, str]:
    """Zwraca (wiersz zamiennika, sciezka URL, blad)."""
    value = text(value)
    if re.match(r"^(https?://|/)", value, re.I):
        path = site_path(value, ctx.hosts)
        if path is None:
            return None, None, f"zamiennik {value} jest spoza sklepu"
        for r in ctx.rows:
            if ctx.product_path(r) == path:
                return r, path, ""
        return None, path, ""
    if value in ctx.by_id:
        return ctx.by_id[value], None, ""
    matches = ctx.by_sku.get(value.lower(), [])
    if len(matches) == 1:
        return matches[0], None, ""
    if len(matches) > 1:
        ids = ", ".join(text(m.get("id")) for m in matches)
        return None, None, f"SKU zamiennika {value} jest niejednoznaczne (id: {ids}) - podaj ID"
    return None, None, f"nie znaleziono zamiennika {value} w eksporcie - podaj ID albo URL"


def category_target(row: dict[str, Any], ctx: Context, decision: Decision) -> str:
    rid = text(row.get("id"))
    if not ctx.col_cats:
        decision.flag(CHECK, "eksport nie ma kolumny Product categories - dodaj ja w WP All Export")
        return ""
    found, missing = ctx.row_cats.get(rid, ([], []))
    for name in missing:
        decision.flag(WARNING, f"kategorii '{name}' nie ma w drzewie kategorii")
    if not found:
        decision.flag(CHECK, "nie rozpoznano zadnej kategorii produktu")
        return ""
    depth = {t: len(category_chain(t, ctx.cats)) for t in found}
    term_id = sorted(found, key=lambda t: (-depth[t], t))[0]
    if len(found) > 1:
        decision.notes.append("kilka kategorii - wybrano najglebsza: " + category_label(term_id, ctx.cats))
    # Kategoria, w ktorej nic nie zostanie, to slaby cel - idziemy do rodzica.
    while ctx.cats[term_id].parent_id in ctx.cats:
        cat = ctx.cats[term_id]
        remaining = ctx.remaining_in_category(term_id)
        known_empty = cat.count is not None and cat.count - ctx.marked_in_category(term_id) <= 0
        if remaining > 0 or (cat.count is not None and not known_empty):
            break
        if known_empty:
            decision.flag(WARNING, f"kategoria {category_label(term_id, ctx.cats)} zostanie pusta - redirect na rodzica")
            term_id = cat.parent_id
            continue
        decision.flag(WARNING, f"w eksporcie nie ma innych produktow z kategorii {category_label(term_id, ctx.cats)} - sprawdz, czy nie bedzie pusta")
        break
    decision.target_label = category_label(term_id, ctx.cats)
    return category_path(term_id, ctx.cats, ctx.category_base)


def target_for(row: dict[str, Any], ctx: Context, decision: Decision, stack: tuple[str, ...] = ()) -> str:
    rid = text(row.get("id"))
    if rid in stack:
        decision.flag(CHECK, "petla zamiennikow: " + " -> ".join(stack + (rid,)))
        return ""
    value = text(row.get(ctx.col_repl)) if ctx.col_repl else ""
    if not value:
        decision.target_kind = decision.target_kind or "kategoria"
        return category_target(row, ctx, decision)
    other, path, error = resolve_replacement(value, ctx)
    if error:
        decision.flag(CHECK, error)
        return ""
    decision.target_kind = "zamiennik"
    if other is not None:
        other_id = text(other.get("id"))
        if other_id == rid:
            decision.flag(CHECK, "zamiennik wskazuje ten sam produkt")
            return ""
        if other_id in ctx.marked:
            decision.notes.append(f"zamiennik {value} tez jest do usuniecia - redirect idzie na jego cel")
            return target_for(other, ctx, decision, stack + (rid,))
        if ctx.status_of(other) not in ("publish", ""):
            decision.flag(CHECK, f"zamiennik {value} nie jest opublikowany ({ctx.status_of(other)})")
            return ""
        other_path = ctx.product_path(other)
        if not other_path:
            decision.flag(CHECK, f"zamiennik {value} nie ma adresu (Permalink/Slug) w eksporcie")
            return ""
        decision.target_label = f"{text(other.get(ctx.col_sku))} {text(other.get(ctx.col_title))}".strip()
        return other_path
    decision.target_label = value
    decision.notes.append("zamiennik podany jako URL spoza eksportu - sprawdz, czy strona dziala")
    decision.status = worst(decision.status, WARNING)
    return path or ""


def decide(row: dict[str, Any], ctx: Context, seen_sources: dict[str, str]) -> Decision:
    rid = text(row.get("id"))
    decision = Decision(rid, text(row.get(ctx.col_sku)), text(row.get(ctx.col_title)), ctx.status_of(row))
    if ctx.is_variation(row):
        decision.flag(CHECK, "wariant nie ma wlasnego adresu - wylacz go w produkcie nadrzednym")
        return decision
    decision.source_path = ctx.product_path(row)
    if not decision.source_path:
        if decision.status_before not in ("publish", ""):
            decision.notes.append("produkt nie byl publiczny i nie ma adresu - bez przekierowania")
            decision.set_draft = decision.status_before != "draft"
            return decision
        decision.flag(CHECK, "brak adresu produktu - dodaj Permalink albo Slug do eksportu")
        return decision
    if decision.source_path in seen_sources:
        decision.flag(CHECK, f"ten sam adres co produkt id {seen_sources[decision.source_path]}")
        return decision
    seen_sources[decision.source_path] = rid

    path = target_for(row, ctx, decision)
    if not path:
        decision.status = CHECK
        return decision
    final, hops, loop = ctx.index.follow(path)
    if loop:
        decision.flag(CHECK, "cel wpada w petle istniejacych przekierowan: " + " -> ".join(str(h.id) for h in hops))
        return decision
    decision.target_path = final
    if hops:
        decision.notes.append(f"cel {as_url(path, '')} jest juz przekierowany (id {', '.join(str(h.id) for h in hops)}) - uzyto celu koncowego")
        if hops[-1].target_path is None:
            decision.flag(WARNING, f"cel koncowy jest poza sklepem: {hops[-1].target}")
            decision.target_path = decision.target_label = hops[-1].target
    if decision.target_path == decision.source_path:
        decision.flag(CHECK, "przekierowanie wskazywaloby na siebie")
        return decision

    existing = ctx.index.by_source.get(decision.source_path)
    if existing:
        if existing.target_path == decision.target_path:
            decision.notes.append(f"przekierowanie juz istnieje (id {existing.id}) - nie dodano drugiego")
        else:
            decision.flag(WARNING, f"z tego adresu jest juz przekierowanie id {existing.id} -> {existing.target}; nowego nie dodano, zmien recznie, jesli trzeba")
    else:
        decision.add_redirect = True
    decision.set_draft = decision.status_before != "draft"
    return decision


def tokens(value: str) -> set[str]:
    return set(re.findall(r"[0-9a-ząćęłńóśźż]+(?:[.,x×/][0-9]+)*", value.lower()))


def suggest(row: dict[str, Any], ctx: Context) -> list[tuple[float, dict[str, Any], list[str]]]:
    rid = text(row.get("id"))
    own_cats = set(ctx.row_cats.get(rid, ([], []))[0])
    values = {c: text(row.get(c)).lower() for c in ctx.attr_cols if text(row.get(c))}
    title = tokens(text(row.get(ctx.col_title)))
    scored = []
    for other in ctx.rows:
        oid = text(other.get("id"))
        if oid == rid or oid in ctx.marked or ctx.status_of(other) not in ("publish", "") or ctx.is_variation(other):
            continue
        if own_cats and not own_cats & set(ctx.row_cats.get(oid, ([], []))[0]):
            continue
        diffs = []
        same = 0
        for col, value in values.items():
            other_value = text(other.get(col)).lower()
            if other_value == value:
                same += 1
            else:
                diffs.append(f"{col[len(ATTR_PREFIX):].strip()}: {text(row.get(col))} -> {text(other.get(col)) or '(brak)'}")
        other_title = tokens(text(other.get(ctx.col_title)))
        title_score = len(title & other_title) / len(title | other_title) if title | other_title else 0.0
        score = 0.7 * same / len(values) + 0.3 * title_score if values else title_score
        if score >= SUGGESTION_MIN_SCORE:
            scored.append((round(score, 2), other, diffs))
    scored.sort(key=lambda item: (-item[0], text(item[1].get("id"))))
    return scored[:SUGGESTION_LIMIT]


def plan_removal(header: list[str], rows: list[dict[str, Any]], cats: dict[int, Category], redirects: list[Redirect],
                 site_url: str = SITE_URL, product_base: str = PRODUCT_BASE, category_base: str = CATEGORY_BASE) -> Plan:
    ctx = Context(header, rows, cats, RedirectIndex(redirects), site_url, product_base, category_base)
    if not ctx.col_delete:
        raise SystemExit(f"Eksport nie ma kolumny '{DELETE_COLUMNS[0]}' - dopisz ja i oznacz produkty do usuniecia")
    decisions, suggestions, updates = [], [], []
    seen_sources: dict[str, str] = {}
    for row in rows:
        if not ctx.is_marked(row):
            continue
        decision = decide(row, ctx, seen_sources)
        decisions.append(decision)
        pointing = ctx.index.pointing_to(decision.source_path) if decision.source_path else []
        if pointing:
            updates.extend((rule, decision) for rule in pointing)
            decision.flag(WARNING, f"{len(pointing)} istniejacych przekierowan prowadzi na ten produkt - przestaw je (arkusz 'Istniejace do zmiany')")
        if decision.target_kind == "kategoria" and decision.status != CHECK:
            for rank, (score, other, diffs) in enumerate(suggest(row, ctx), start=1):
                path = ctx.product_path(other)
                suggestions.append(Suggestion(decision.row_id, decision.sku, rank, text(other.get("id")),
                                              text(other.get(ctx.col_sku)), text(other.get(ctx.col_title)),
                                              as_url(path, site_url) if path else "", score, "; ".join(diffs[:6])))
    return Plan(decisions, suggestions, updates, ctx.index.chains())


# --- weryfikacja online -----------------------------------------------------

class _NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *args, **kwargs):  # noqa: D401 - nie podazamy za 301
        return None


def check_online(decisions: list[Decision], site_url: str) -> None:
    opener = urllib.request.build_opener(_NoRedirect)
    cache: dict[str, str] = {}
    for d in decisions:
        if not d.add_redirect:
            continue
        url = d.target_path if d.target_path.startswith("http") else as_url(d.target_path, site_url)
        if url not in cache:
            try:
                request = urllib.request.Request(url, method="HEAD", headers={"User-Agent": "AlkanInator-redirect-check"})
                cache[url] = str(opener.open(request, timeout=15).status)
            except urllib.error.HTTPError as error:
                cache[url] = str(error.code)
            except OSError as error:
                cache[url] = f"blad: {error}"
        code = cache[url]
        if code != "200":
            d.flag(CHECK, f"cel zwraca {code} - nie dodano przekierowania")
            d.add_redirect = False
            d.set_draft = False
        else:
            d.notes.append("cel sprawdzony online: 200")


# --- zapis ------------------------------------------------------------------

def read_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    values = list(load_workbook(path, read_only=True).worksheets[0].iter_rows(values_only=True))
    header = [text(h) for h in values[0]]
    return header, [dict(zip(header, r)) for r in values[1:] if any(v not in (None, "") for v in r)]


def target_url(decision: Decision, site_url: str) -> str:
    return decision.target_path if decision.target_path.startswith("http") else as_url(decision.target_path, site_url)


def write_redirects_csv(path: Path, decisions: list[Decision], site_url: str) -> int:
    rows = [d for d in decisions if d.add_redirect]
    path.parent.mkdir(parents=True, exist_ok=True)
    # Bez BOM: Redirection rozpoznaje naglowek po pierwszej komorce "source".
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["source", "target", "regex", "code"])
        for d in rows:
            writer.writerow([d.source_path + "/", target_url(d, site_url), 0, 301])
    return len(rows)


def write_redirects_json(path: Path, decisions: list[Decision], site_url: str) -> None:
    redirects = []
    for position, d in enumerate((d for d in decisions if d.add_redirect)):
        redirects.append({
            "id": position + 1,
            "url": d.source_path + "/",
            "match_url": d.source_path,
            "match_data": {"source": {"flag_query": "exact", "flag_case": False, "flag_trailing": False, "flag_regex": False}},
            "action_code": 301,
            "action_type": "url",
            "action_data": {"url": target_url(d, site_url)},
            "match_type": "url",
            "title": f"Alkan: usunięty produkt #{d.row_id} {d.sku}".strip(),
            "regex": False,
            "group_id": 1,
            "position": position,
            "enabled": True,
        })
    data = {
        "plugin": {"version": "5.10.1", "date": date.today().isoformat()},
        "groups": [{"id": 1, "name": GROUP_NAME, "module_id": 1, "status": "enabled"}],
        "redirects": redirects,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=4), encoding="utf-8")


def write_draft_copy(input_path: Path, output_path: Path, header: list[str], decisions: list[Decision]) -> int:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    workbook = load_workbook(input_path)
    sheet = workbook.worksheets[0]
    col = {text(sheet.cell(1, c).value): c for c in range(1, sheet.max_column + 1)}
    status_col = col.get(find_column(header, ("Status", "post_status")))
    if not status_col:
        status_col = sheet.max_column + 1
        sheet.cell(1, status_col).value = "Status"
    drafts = {d.row_id for d in decisions if d.set_draft}
    fill = PatternFill("solid", fgColor="F8CBAD")
    changed = 0
    for r in range(2, sheet.max_row + 1):
        if text(sheet.cell(r, col["id"]).value) in drafts:
            sheet.cell(r, status_col).value = "draft"
            for c in range(1, sheet.max_column + 1):
                sheet.cell(r, c).fill = fill
            changed += 1
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return changed


def write_control(path: Path, plan: Plan, site_url: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    fills = {WARNING: "FFF2CC", CHECK: "F8CBAD"}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Do usuniecia"
    sheet.append(["id", "SKU", "Tytul", "Status przed", "Status po", "Przekierowanie z", "Przekierowanie na",
                  "Typ celu", "Cel", "Do importu Redirection", "status", "Uwagi"])
    for d in plan.decisions:
        sheet.append([d.row_id, d.sku, d.title, d.status_before, "draft" if d.set_draft else d.status_before,
                      as_url(d.source_path, site_url) if d.source_path else "",
                      target_url(d, site_url) if d.target_path else "", d.target_kind, d.target_label,
                      "tak" if d.add_redirect else "nie", d.status, "; ".join(d.notes)])
        if d.status in fills:
            for cell in sheet[sheet.max_row]:
                cell.fill = PatternFill("solid", fgColor=fills[d.status])

    sug = workbook.create_sheet("Propozycje zamiennikow")
    sug.append(["id usuwanego", "SKU usuwanego", "Nr", "id kandydata", "SKU kandydata", "Tytul kandydata", "URL",
                "Podobienstwo", "Roznice atrybutow", "status"])
    for s in plan.suggestions:
        sug.append([s.deleted_id, s.deleted_sku, s.rank, s.row_id, s.sku, s.title, s.url, s.score, s.differences, CHECK])

    upd = workbook.create_sheet("Istniejace do zmiany")
    upd.append(["id reguly", "Z adresu", "Obecny cel (usuwany produkt)", "Nowy cel", "SKU usuwanego"])
    for rule, d in plan.updates:
        upd.append([rule.id, rule.source, rule.target, target_url(d, site_url) if d.target_path else "(brak celu)", d.sku])

    chains = workbook.create_sheet("Lancuchy istniejacych")
    chains.append(["id reguly", "Z adresu", "Obecny cel", "Cel koncowy", "Liczba 301 po drodze"])
    for rule, final, hops in plan.chains:
        chains.append([rule.id, rule.source, rule.target, as_url(final, site_url) if final.startswith("/") else final, hops])

    for ws in workbook.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for column in ws.columns:
            width = max(len(text(cell.value)) for cell in column)
            ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 8), 70)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main(argv: list[str] | None = None) -> int:
    today = date.today().isoformat()
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--input", required=True, type=Path, help="Eksport WP All Export z kolumnami 'Do usunięcia' i 'Zamiennik'")
    parser.add_argument("--redirects", required=True, type=Path, help="Eksport JSON wtyczki Redirection z produkcji")
    parser.add_argument("--categories", type=Path, default=None,
                        help="Kategorie z produkcji (XLSX/CSV: ID, nazwa, slug, rodzic); domyslnie zrzut z AlkanLocal")
    parser.add_argument("--site-url", default=SITE_URL)
    parser.add_argument("--verify-online", action="store_true", help="Sprawdz HEAD-em, czy cele zwracaja 200")
    parser.add_argument("--output", type=Path, default=ROOT / "output" / f"Usuwanie-produktow_DRAFT_{today}.xlsx")
    parser.add_argument("--redirects-out", type=Path, default=ROOT / "output" / f"Usuwanie-produktow_redirection_{today}.csv")
    parser.add_argument("--control", type=Path, default=ROOT / "reports" / f"Usuwanie-produktow_kontrola_{today}.xlsx")
    args = parser.parse_args(argv)

    if args.output.resolve() == args.input.resolve():
        raise SystemExit("Plik wynikowy nie moze nadpisac pliku wejsciowego.")

    header, rows = read_rows(args.input)
    category_source = args.categories or DEFAULT_CATEGORY_SNAPSHOT
    cats = categories_from_records(load_table(category_source))
    redirects = load_redirects(args.redirects, hosts_for(args.site_url))
    plan = plan_removal(header, rows, cats, redirects, site_url=args.site_url)
    if args.categories is None and not args.verify_online:
        for d in plan.decisions:
            if d.target_kind == "kategoria" and d.target_path:
                d.flag(WARNING, "kategoria z lokalnej kopii AlkanLocal - podaj --categories z produkcji albo --verify-online")
    if args.verify_online:
        check_online(plan.decisions, args.site_url)

    drafts = write_draft_copy(args.input, args.output, header, plan.decisions)
    count = write_redirects_csv(args.redirects_out, plan.decisions, args.site_url)
    json_out = args.redirects_out.with_suffix(".json")
    write_redirects_json(json_out, plan.decisions, args.site_url)
    write_control(args.control, plan, args.site_url)

    statuses = Counter(d.status for d in plan.decisions)
    print(f"Do usuniecia: {len(plan.decisions)} {dict(statuses)}; draft: {drafts}; nowych przekierowan: {count}; "
          f"propozycji zamiennikow: {len(plan.suggestions)}; istniejacych do zmiany: {len(plan.updates)}; "
          f"lancuchow w istniejacych: {len(plan.chains)}")
    print(f"Import produktow (WP All Import): {args.output}")
    print(f"Przekierowania (Redirection CSV): {args.redirects_out}")
    print(f"Przekierowania (Redirection JSON): {json_out}")
    print(f"Kontrola: {args.control}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
