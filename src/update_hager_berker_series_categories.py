from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DEFAULT_INPUT = Path("input/Hager-Berker-Gniazdka.xlsx")
DEFAULT_CATEGORY_MAP = Path("input/BerkerHagerKategorie.xlsx")
DEFAULT_CATALOG = Path("output/Hager-Berker-Katalog-PDF_BAZA_KOMPATYBILNOSCI.xlsx")
DEFAULT_OUTPUT = Path(
    "output/Hager-Berker-Gniazdka_SERIE_KATEGORIE_POPRAWIONE_2026-08-27.xlsx"
)
DEFAULT_REPORT = Path(
    "reports/Hager-Berker-serie-kategorie-poprawione-2026-08-27.json"
)

STORE_SHEET = "Worksheet"
CATALOG_PRODUCTS_SHEET = "Produkty katalogowe"
CATALOG_HEADER_ROW = 4
STORE_HEADER_ROW = 1
ROOT_CATEGORY = "Gniazdka i Łączniki>Hager / Berker"

TARGET_COLUMNS = {"Atrybut Produktu: Seria", "Kategorie produktów"}

EXPECTED_NEW_LEAVES = [
    "Berker B.3/B.7",
    "Berker K.1/K.5",
    "Berker Q.1/Q.3/Q.7",
    "Berker R.1/R.3/R.8",
    "Berker 1930/Glasserie",
    "Berker W.1",
    "Berker Mechanizmy - one.platform",
    "Hager Lumina",
    "Berker B.Kwadrat/S.1",
    "Berker R.classic",
]

OLD_TO_NEW_LEAF = {
    "Lumina": "Hager Lumina",
    "Hager Lumina": "Hager Lumina",
    "Mechanizmy - one.platform": "Berker Mechanizmy - one.platform",
    "Berker Mechanizmy - one.platform": "Berker Mechanizmy - one.platform",
    "Berker B.Kwadrat": "Berker B.Kwadrat/S.1",
    "Berker B. Kwadrat": "Berker B.Kwadrat/S.1",
    "Berker B.3": "Berker B.3/B.7",
    "Berker B.7": "Berker B.3/B.7",
    "Berker K.1": "Berker K.1/K.5",
    "Berker K.5": "Berker K.1/K.5",
    "Berker Q.": "Berker Q.1/Q.3/Q.7",
    "Berker Q.1": "Berker Q.1/Q.3/Q.7",
    "Berker Q.3": "Berker Q.1/Q.3/Q.7",
    "Berker Q.7": "Berker Q.1/Q.3/Q.7",
    "Berker R.": "Berker R.1/R.3/R.8",
    "Berker R.1": "Berker R.1/R.3/R.8",
    "Berker R.3": "Berker R.1/R.3/R.8",
    "Berker R.8": "Berker R.1/R.3/R.8",
    "Berker S.1": "Berker B.Kwadrat/S.1",
    "Berker Serie 1930": "Berker 1930/Glasserie",
    "Berker Glasserie": "Berker 1930/Glasserie",
    "Berker R.classic": "Berker R.classic",
    "Berker W.1": "Berker W.1",
    "Berker B.3/B.7": "Berker B.3/B.7",
    "Berker K.1/K.5": "Berker K.1/K.5",
    "Berker Q.1/Q.3/Q.7": "Berker Q.1/Q.3/Q.7",
    "Berker R.1/R.3/R.8": "Berker R.1/R.3/R.8",
    "Berker 1930/Glasserie": "Berker 1930/Glasserie",
    "Berker B.Kwadrat/S.1": "Berker B.Kwadrat/S.1",
}

# Ta stara kategoria nie ma następcy w aktualnej liście kategorii.
REMOVED_OLD_LEAVES = {"Berker Integro Flow"}

# Potwierdzone wizualnie w katalogu PDF; parser bazy nie zawsze przenosi oznaczenie
# one.platform do pola "Seria / system".
MANUAL_SERIES_OVERRIDES = {
    "4586/HAG": ["one.platform"],
    "503501/HAG": ["one.platform"],
    "53455111/HAG": ["one.platform"],
}

SERIES_ORDER = {
    "one.platform": 10,
    "B. Kwadrat": 20,
    "B.3": 21,
    "B.7": 22,
    "S.1": 23,
    "K.1": 30,
    "K.5": 31,
    "Q.1": 40,
    "Q.3": 41,
    "Q.7": 42,
    "R.1": 50,
    "R.3": 51,
    "R.8": 52,
    "Serie 1930": 60,
    "Glasserie": 61,
    "R.classic": 70,
    "W.1": 80,
    "Lumina": 90,
    "Lumina soul": 91,
    "Lumina intense": 92,
    "Lumina passion": 93,
    "Serie 1930 Rosenthal": 94,
    "TS": 95,
    "Integro Flow": 100,
    "Inne": 999,
}

SERIES_TO_CATEGORY = {
    "one.platform": "Berker Mechanizmy - one.platform",
    "B. Kwadrat": "Berker B.Kwadrat/S.1",
    "B.3": "Berker B.3/B.7",
    "B.7": "Berker B.3/B.7",
    "S.1": "Berker B.Kwadrat/S.1",
    "K.1": "Berker K.1/K.5",
    "K.5": "Berker K.1/K.5",
    "Q.1": "Berker Q.1/Q.3/Q.7",
    "Q.3": "Berker Q.1/Q.3/Q.7",
    "Q.7": "Berker Q.1/Q.3/Q.7",
    "R.1": "Berker R.1/R.3/R.8",
    "R.3": "Berker R.1/R.3/R.8",
    "R.8": "Berker R.1/R.3/R.8",
    "Serie 1930": "Berker 1930/Glasserie",
    "Glasserie": "Berker 1930/Glasserie",
    "R.classic": "Berker R.classic",
    "W.1": "Berker W.1",
    "Lumina": "Hager Lumina",
    "Lumina soul": "Hager Lumina",
    "Lumina intense": "Hager Lumina",
    "Lumina passion": "Hager Lumina",
    "Serie 1930 Rosenthal": "Berker 1930/Glasserie",
}

CONFIRMED_CATEGORY_REMOVALS = {
    "10439949/HAG": {"Puszki i złączki>Puszka hermetyczna natynkowa"},
    "10516099/HAG": {f"{ROOT_CATEGORY}>Berker Mechanizmy - one.platform"},
}

MOJIBAKE_MARKERS = ("\ufffd", "Ã", "Â", "Ä‚", "Ă", "â€")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def header_map(sheet, row_number: int) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in sheet[row_number]
        if cell.value not in (None, "")
    }


def split_pipe(value: Any) -> list[str]:
    return [part.strip() for part in str(value or "").split("|") if part.strip()]


def unique_in_order(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result


def normalize_series(value: Any) -> list[str]:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if not text:
        return []

    found: list[tuple[int, str]] = []
    lower = text.casefold()
    if "lumina" in lower:
        variants_found = False
        for match in re.finditer(r"\b(soul|intense|passion)\b", text, flags=re.IGNORECASE):
            variant = match.group(1).casefold()
            found.append((match.start(), f"Lumina {variant}"))
            variants_found = True
        if not variants_found:
            position = lower.find("lumina")
            found.append((position, "Lumina"))

    pattern = re.compile(
        r"one\.platform"
        r"|B\.\s*Kwadrat"
        r"|Serie\s+1930\s+Rosenthal"
        r"|Serie\s+1930"
        r"|(?:Serie\s+)?Glasserie"
        r"|(?:Berker\s+)?Serie\s+R\.classic"
        r"|R\.classic"
        r"|Integro\s+Flow"
        r"|B\.[37]"
        r"|S\.1"
        r"|K\.[15]"
        r"|Q\.[137]"
        r"|R\.[138]"
        r"|W\.1"
        r"|Berker\s+TS"
        r"|\bInne\b",
        flags=re.IGNORECASE,
    )
    for match in pattern.finditer(text):
        raw = match.group(0)
        token_lower = raw.casefold()
        if "kwadrat" in token_lower:
            token = "B. Kwadrat"
        elif "rosenthal" in token_lower:
            token = "Serie 1930 Rosenthal"
        elif re.fullmatch(r"serie\s+1930", raw, flags=re.IGNORECASE):
            token = "Serie 1930"
        elif "glas" in token_lower:
            token = "Glasserie"
        elif "r.classic" in token_lower:
            token = "R.classic"
        elif "integro" in token_lower:
            token = "Integro Flow"
        elif "one.platform" in token_lower:
            token = "one.platform"
        elif "berker" in token_lower and "ts" in token_lower:
            token = "TS"
        elif token_lower == "inne":
            token = "Inne"
        else:
            token = raw.upper().replace(" ", "")
        found.append((match.start(), token))

    return unique_in_order(token for _, token in sorted(found, key=lambda item: item[0]))


def sorted_series(tokens: Iterable[str]) -> list[str]:
    return sorted(
        unique_in_order(tokens),
        key=lambda token: (SERIES_ORDER.get(token, 500), token.casefold()),
    )


def normalize_catalog_series(value: Any) -> list[str]:
    """Normalizuje pole katalogowe, gdzie samodzielne `Glas` oznacza Glasserie.

    W tytułach i opisach `Glas` bywa częścią starego oznaczenia B.7 Glas,
    dlatego ogólny parser nie może zamieniać każdego wystąpienia na Glasserie.
    """
    text = re.sub(
        r"(?<!\w)Glas(?!\w)",
        "Glasserie",
        str(value or ""),
        flags=re.IGNORECASE,
    )
    return normalize_series(text)


def load_new_category_leaves(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Arkusz1"]
        leaves = [
            str(sheet.cell(row, 13).value).strip()
            for row in range(11, sheet.max_row + 1)
            if sheet.cell(row, 13).value not in (None, "")
        ]
    finally:
        workbook.close()
    if leaves != EXPECTED_NEW_LEAVES:
        raise ValueError(
            "Lista nowych kategorii nie jest zgodna z oczekiwaną strukturą: "
            f"{leaves!r}"
        )
    return leaves


def load_catalog_series(path: Path) -> tuple[dict[str, list[str]], dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[CATALOG_PRODUCTS_SHEET]
        headers = header_map(sheet, CATALOG_HEADER_ROW)
        required = {
            "Kod katalogowy",
            "Seria / system",
            "Status ekstrakcji",
            "SKU sklepu",
            "Strona PDF",
            "Sekcja strony",
        }
        if not required.issubset(headers):
            raise ValueError(f"Brak kolumn katalogu: {sorted(required - set(headers))}")

        evidence_by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
        direct_codes_by_sku: dict[str, list[str]] = defaultdict(list)
        for row in sheet.iter_rows(min_row=CATALOG_HEADER_ROW + 1, values_only=True):
            get = lambda name: row[headers[name] - 1]
            series = normalize_catalog_series(get("Seria / system"))
            series_source = "Seria / system"
            if not series:
                series = normalize_catalog_series(get("Sekcja strony"))
                series_source = "Sekcja strony"
            code = str(get("Kod katalogowy") or "").strip()
            if not code:
                continue
            if series:
                evidence_by_code[code].append(
                    {
                        "catalog_code": code,
                        "series": series,
                        "status": str(get("Status ekstrakcji") or "").strip(),
                        "page": get("Strona PDF"),
                        "source": series_source,
                    }
                )
            for sku in split_pipe(get("SKU sklepu")):
                direct_codes_by_sku[sku].append(code)

        alias_sheet = workbook["Alias kodów"]
        alias_headers = header_map(alias_sheet, CATALOG_HEADER_ROW)
        aliases: dict[str, list[str]] = defaultdict(list)
        alias_statuses: Counter[str] = Counter()
        for row in alias_sheet.iter_rows(min_row=CATALOG_HEADER_ROW + 1, values_only=True):
            get = lambda name: row[alias_headers[name] - 1]
            sku = str(get("SKU sklepu") or "").strip()
            if not sku:
                continue
            status = str(get("Status") or "").strip()
            alias_statuses[status or "BRAK"] += 1
            if status == "DOPASOWANO":
                code = str(get("Kod katalogowy dopasowany") or "").strip()
                if code:
                    aliases[sku].append(code)
    finally:
        workbook.close()

    result: dict[str, list[str]] = {}
    ambiguous: dict[str, list[list[str]]] = {}
    status_counts: Counter[str] = Counter()
    evidence_sources: Counter[str] = Counter()
    all_skus = set(direct_codes_by_sku) | set(aliases)
    for sku in all_skus:
        codes = unique_in_order(aliases.get(sku, []) or direct_codes_by_sku.get(sku, []))
        rows = [row for code in codes for row in evidence_by_code.get(code, [])]
        if not rows:
            continue
        ok_rows = [row for row in rows if row["status"] == "OK"]
        chosen_rows = ok_rows or rows
        groups = unique_in_order("|".join(row["series"]) for row in chosen_rows)
        if len(groups) > 1:
            ambiguous[sku] = [group.split("|") for group in groups]
        result[sku] = sorted_series(
            token for row in chosen_rows for token in row["series"]
        )
        for row in chosen_rows:
            status_counts[row["status"] or "BRAK"] += 1
            evidence_sources[row["source"]] += 1

    for sku, series in MANUAL_SERIES_OVERRIDES.items():
        result[sku] = series

    diagnostics = {
        "skus_with_series": len(result),
        "ambiguous_skus": ambiguous,
        "chosen_evidence_statuses": dict(status_counts),
        "chosen_evidence_sources": dict(evidence_sources),
        "alias_statuses": dict(alias_statuses),
        "manual_overrides": MANUAL_SERIES_OVERRIDES,
    }
    return result, diagnostics


def desired_category_leaves(series: Iterable[str]) -> list[str]:
    leaves = {SERIES_TO_CATEGORY[token] for token in series if token in SERIES_TO_CATEGORY}
    return [leaf for leaf in EXPECTED_NEW_LEAVES if leaf in leaves]


def normalize_category_path(path: str) -> str | None:
    cleaned = re.sub(r"\s*>\s*", ">", path.strip())
    if not cleaned.startswith(ROOT_CATEGORY + ">"):
        return cleaned
    leaf = cleaned[len(ROOT_CATEGORY) + 1 :]
    if leaf in REMOVED_OLD_LEAVES:
        return None
    mapped = OLD_TO_NEW_LEAF.get(leaf)
    if mapped:
        return f"{ROOT_CATEGORY}>{mapped}"
    return cleaned


def update_categories(value: Any, sku: str, series: list[str]) -> tuple[str, list[str], list[str]]:
    before_paths = split_pipe(value)
    normalized: list[str] = []
    removed: list[str] = []
    explicit_removals = CONFIRMED_CATEGORY_REMOVALS.get(sku, set())
    desired_paths = {
        f"{ROOT_CATEGORY}>{leaf}" for leaf in desired_category_leaves(series)
    }
    for original in before_paths:
        normalized_original = re.sub(r"\s*>\s*", ">", original.strip())
        mapped = normalize_category_path(original)
        if mapped is None or mapped in explicit_removals or normalized_original in explicit_removals:
            removed.append(original)
            continue
        if (
            mapped.startswith(ROOT_CATEGORY + ">")
            and mapped[len(ROOT_CATEGORY) + 1 :] in EXPECTED_NEW_LEAVES
            and mapped not in desired_paths
        ):
            removed.append(original)
            continue
        normalized.append(mapped)

    if not any(
        path == ROOT_CATEGORY or path.startswith(ROOT_CATEGORY + ">")
        for path in normalized
    ):
        normalized.append(ROOT_CATEGORY)
    added: list[str] = []
    for leaf in desired_category_leaves(series):
        path = f"{ROOT_CATEGORY}>{leaf}"
        if path not in normalized:
            normalized.append(path)
            added.append(path)
    return "|".join(unique_in_order(normalized)), added, removed


def scan_mojibake(workbook) -> list[tuple[str, str, str]]:
    bad: list[tuple[str, str, str]] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(marker in cell.value for marker in MOJIBAKE_MARKERS):
                    bad.append((sheet.title, cell.coordinate, cell.value[:120]))
                    if len(bad) >= 20:
                        return bad
    return bad


def collect_workbook_structure(workbook) -> dict[str, Any]:
    result: dict[str, Any] = {"sheetnames": list(workbook.sheetnames), "sheets": {}}
    for sheet in workbook.worksheets:
        result["sheets"][sheet.title] = {
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "merged_ranges": sorted(str(item) for item in sheet.merged_cells.ranges),
            "freeze_panes": str(sheet.freeze_panes or ""),
            "auto_filter": str(sheet.auto_filter.ref or ""),
        }
    return result


def build_output(
    input_path: Path,
    category_map_path: Path,
    catalog_path: Path,
    output_path: Path,
) -> dict[str, Any]:
    source_hashes = {
        "input": sha256(input_path),
        "category_map": sha256(category_map_path),
        "catalog": sha256(catalog_path),
    }
    new_leaves = load_new_category_leaves(category_map_path)
    catalog_series, catalog_diagnostics = load_catalog_series(catalog_path)

    workbook = load_workbook(input_path)
    input_structure = collect_workbook_structure(workbook)
    if workbook.sheetnames != [STORE_SHEET]:
        raise ValueError(f"Nieoczekiwane arkusze pliku wejściowego: {workbook.sheetnames}")
    sheet = workbook[STORE_SHEET]
    headers = header_map(sheet, STORE_HEADER_ROW)
    required = {"SKU", *TARGET_COLUMNS}
    if not required.issubset(headers):
        raise ValueError(f"Brak wymaganych kolumn: {sorted(required - set(headers))}")

    sku_col = headers["SKU"]
    series_col = headers["Atrybut Produktu: Seria"]
    category_col = headers["Kategorie produktów"]

    seen_skus: set[str] = set()
    changed_series: list[dict[str, Any]] = []
    changed_categories: list[dict[str, Any]] = []
    unmatched_skus: list[str] = []
    unknown_series_tokens: Counter[str] = Counter()
    added_category_counts: Counter[str] = Counter()
    removed_category_counts: Counter[str] = Counter()
    styles_before: dict[tuple[int, int], int] = {}

    for row_number in range(2, sheet.max_row + 1):
        sku = str(sheet.cell(row_number, sku_col).value or "").strip()
        if not sku:
            continue
        if sku in seen_skus:
            raise ValueError(f"Powtórzone SKU w pliku wejściowym: {sku}")
        seen_skus.add(sku)

        current_series_value = sheet.cell(row_number, series_col).value
        current_series = normalize_series(current_series_value)
        exact_series = catalog_series.get(sku)
        final_series = exact_series if exact_series else current_series
        if not exact_series:
            unmatched_skus.append(sku)
        for token in final_series:
            if token not in SERIES_ORDER:
                unknown_series_tokens[token] += 1

        before_series_text = str(current_series_value or "").strip()
        final_series_value = (
            "|".join(final_series) if exact_series else before_series_text
        )
        if final_series_value != before_series_text:
            styles_before[(row_number, series_col)] = sheet.cell(row_number, series_col).style_id
            changed_series.append(
                {
                    "sku": sku,
                    "row": row_number,
                    "before": before_series_text,
                    "after": final_series_value,
                    "source": "PDF_PRODUCT_ROW" if sku not in MANUAL_SERIES_OVERRIDES else "PDF_VISUAL_OVERRIDE",
                }
            )
            sheet.cell(row_number, series_col).value = final_series_value or None

        current_category_value = sheet.cell(row_number, category_col).value
        final_category_value, added, removed = update_categories(
            current_category_value, sku, final_series
        )
        before_category_text = str(current_category_value or "").strip()
        if final_category_value != before_category_text:
            styles_before[(row_number, category_col)] = sheet.cell(row_number, category_col).style_id
            changed_categories.append(
                {
                    "sku": sku,
                    "row": row_number,
                    "before": before_category_text,
                    "after": final_category_value,
                    "added": added,
                    "removed": removed,
                }
            )
            sheet.cell(row_number, category_col).value = final_category_value
        for item in added:
            added_category_counts[item] += 1
        for item in removed:
            removed_category_counts[item] += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

    current_hashes = {
        "input": sha256(input_path),
        "category_map": sha256(category_map_path),
        "catalog": sha256(catalog_path),
    }
    if current_hashes != source_hashes:
        raise RuntimeError("Co najmniej jeden plik źródłowy zmienił się podczas operacji")

    source_check = load_workbook(input_path, read_only=False, data_only=False)
    output_check = load_workbook(output_path, read_only=False, data_only=False)
    try:
        output_structure = collect_workbook_structure(output_check)
        if output_structure != input_structure:
            raise ValueError("Struktura skoroszytu wynikowego różni się od wejściowej")
        if scan_mojibake(output_check):
            raise ValueError("W pliku wynikowym wykryto podejrzane uszkodzenie kodowania")

        source_sheet = source_check[STORE_SHEET]
        output_sheet = output_check[STORE_SHEET]
        allowed_columns = {series_col, category_col}
        unexpected_changes: list[str] = []
        formula_count_source = 0
        formula_count_output = 0
        for row_number in range(1, source_sheet.max_row + 1):
            for column_number in range(1, source_sheet.max_column + 1):
                source_cell = source_sheet.cell(row_number, column_number)
                output_cell = output_sheet.cell(row_number, column_number)
                if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
                    formula_count_source += 1
                if isinstance(output_cell.value, str) and output_cell.value.startswith("="):
                    formula_count_output += 1
                if source_cell.value != output_cell.value and column_number not in allowed_columns:
                    unexpected_changes.append(source_cell.coordinate)
                    if len(unexpected_changes) >= 20:
                        break
            if len(unexpected_changes) >= 20:
                break
        if unexpected_changes:
            raise ValueError(f"Zmiany poza dozwolonymi kolumnami: {unexpected_changes}")
        if formula_count_source != formula_count_output:
            raise ValueError("Zmieniła się liczba formuł")
        for (row_number, column_number), style_id in styles_before.items():
            if output_sheet.cell(row_number, column_number).style_id != style_id:
                raise ValueError(
                    f"Zmienił się styl komórki {output_sheet.cell(row_number, column_number).coordinate}"
                )

        final_old_category_paths: Counter[str] = Counter()
        missing_implied_categories: list[dict[str, Any]] = []
        output_headers = header_map(output_sheet, STORE_HEADER_ROW)
        for row_number in range(2, output_sheet.max_row + 1):
            sku = str(output_sheet.cell(row_number, output_headers["SKU"]).value or "").strip()
            if not sku:
                continue
            series = normalize_series(
                output_sheet.cell(row_number, output_headers["Atrybut Produktu: Seria"]).value
            )
            categories = split_pipe(
                output_sheet.cell(row_number, output_headers["Kategorie produktów"]).value
            )
            category_set = set(categories)
            for path in categories:
                if path.startswith(ROOT_CATEGORY + ">"):
                    leaf = path[len(ROOT_CATEGORY) + 1 :]
                    if leaf in OLD_TO_NEW_LEAF and leaf not in new_leaves:
                        final_old_category_paths[path] += 1
                    if leaf in REMOVED_OLD_LEAVES:
                        final_old_category_paths[path] += 1
            for leaf in desired_category_leaves(series):
                path = f"{ROOT_CATEGORY}>{leaf}"
                if path not in category_set:
                    missing_implied_categories.append({"sku": sku, "category": path})
        if final_old_category_paths:
            raise ValueError(f"Pozostały stare kategorie: {dict(final_old_category_paths)}")
        if missing_implied_categories:
            raise ValueError(
                "Brak kategorii wynikających z serii: "
                f"{missing_implied_categories[:20]}"
            )
    finally:
        source_check.close()
        output_check.close()

    return {
        "input": str(input_path.resolve()),
        "category_map": str(category_map_path.resolve()),
        "catalog": str(catalog_path.resolve()),
        "output": str(output_path.resolve()),
        "source_hashes_sha256": source_hashes,
        "rows_including_header": input_structure["sheets"][STORE_SHEET]["max_row"],
        "columns": input_structure["sheets"][STORE_SHEET]["max_column"],
        "products": len(seen_skus),
        "catalog_series_matches": len(seen_skus) - len(unmatched_skus),
        "products_without_catalog_series_match": len(unmatched_skus),
        "unmatched_skus": unmatched_skus,
        "changed_series_count": len(changed_series),
        "changed_category_count": len(changed_categories),
        "changed_series": changed_series,
        "changed_categories": changed_categories,
        "added_category_counts": dict(added_category_counts),
        "removed_category_counts": dict(removed_category_counts),
        "new_category_leaves": new_leaves,
        "unknown_series_tokens": dict(unknown_series_tokens),
        "catalog_diagnostics": catalog_diagnostics,
        "validation": {
            "source_files_unchanged": True,
            "sheet_structure_preserved": True,
            "only_target_columns_changed": True,
            "target_cell_styles_preserved": True,
            "formula_count_preserved": True,
            "no_old_category_leaf_remains": True,
            "all_series_implied_categories_present": True,
            "mojibake_scan_passed": True,
        },
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Aktualizuje serie i połączone kategorie Hager/Berker bez nadpisywania wejścia."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--category-map", type=Path, default=DEFAULT_CATEGORY_MAP)
    parser.add_argument("--catalog", type=Path, default=DEFAULT_CATALOG)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    report = build_output(args.input, args.category_map, args.catalog, args.output)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({key: report[key] for key in (
        "output",
        "products",
        "catalog_series_matches",
        "products_without_catalog_series_match",
        "changed_series_count",
        "changed_category_count",
        "added_category_counts",
        "removed_category_counts",
        "unknown_series_tokens",
        "validation",
    )}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
