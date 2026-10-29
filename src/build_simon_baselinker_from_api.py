"""Buduje CSV Baselinkera dla nowej partii Kontakt Simon z listy kodow i API e-katalogu.

Wejscie to zestawienie z Ekonoma (CSV cp1250, ``;``, kolumny ``Kod towaru``,
``Nazwa towaru``, ``Kod producenta`` = symbol Simon). Dane produktu (EAN, kolor,
seria, IP, montaz Push In, zdjecie) pochodza z API (``src/simon_api.py``,
cache w ``cache/simon_api/``). Nazwy buduje slownik
``dictionaries/simon_product_names.yaml``: rdzen per symbol bazowy + kolor
odmieniony przez rodzaj + seria + kod producenta na koncu.

Wynik to SUROWY CSV - Typ/Podtyp, lista dozwolonych atrybutow i kontrola
wartosci wzgledem terminow Woo robi potem
``src/fix_simon_attributes_to_shop_model.py``. Kolumna ``description`` zostaje
pusta: algorytmiczny ``build_description_html`` pisal o osprzecie jezykiem
oprawy ("laczy uzyteczne swiatlo", "oprawa jest zabezpieczona" w opisie
gniazda), a opisy z FAQ i tak nie przechodza przez Baselinker -> Woo (REST
wycina ``<script>``). Opisy powstaja osobno, recznie per SKU
(skill alkan-product-description) i ida do Woo przez WP All Import.

Uzycie:
    py src/build_simon_baselinker_from_api.py \
        --input "input/Zestawienie SIMON 24.csv" \
        --output output/Simon-24_Baselinker_SUROWY_2026-10-20.csv \
        --control output/Simon-24_nazwy_kontrola_2026-10-20.xlsx \
        --report reports/simon_24_baselinker_2026-10-20.json
"""

from __future__ import annotations

import argparse
import concurrent.futures
import csv
import hashlib
import json
import re
import ssl
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_NAMES = ROOT / "dictionaries" / "simon_product_names.yaml"
BASELINKER_COLUMNS = ["product_id", "name", "sku", "ean", "manufacturer_name", "category", "description", "features", "images_urls"]
SKU_SUFFIX = "/KON"
SYMBOL_RE = re.compile(r"^(?P<base>.+?)/(?P<color>\d{2})(?P<flap>A?)$")
IP_FORMAT = {"IP20": "IP 20", "IP44": "IP 44", "IP40": "IP 40", "IP68": "IP 68", "IP20/IP44": "IP20/IP44"}


@dataclass
class SimonRow:
    lp: str
    sku: str
    symbol: str
    ekonom_name: str
    base: str = ""
    color: str = ""
    name: str = ""
    ean: str = ""
    category: str = ""
    features: dict[str, str] = field(default_factory=dict)
    image: str = ""
    api_name: str = ""
    api_series: str = ""
    api_color: str = ""
    status: str = "OK"
    notes: list[str] = field(default_factory=list)

    def flag(self, status: str, note: str) -> None:
        order = {"OK": 0, "WARNING": 1, "DO_SPRAWDZENIA": 2}
        if order[status] > order[self.status]:
            self.status = status
        self.notes.append(note)


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def load_names(path: Path = DEFAULT_NAMES) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as stream:
        return yaml.safe_load(stream)


def read_ekonom_csv(path: Path) -> list[SimonRow]:
    """Zestawienie z Ekonoma; ``Nazwa towaru`` bywa pocieta znacznikiem `` | `` (limit 40 znakow)."""
    with path.open("r", encoding="cp1250", newline="") as stream:
        reader = csv.DictReader(stream, delimiter=";")
        rows = []
        for record in reader:
            symbol = clean(record.get("Kod producenta"))
            if not symbol:
                continue
            rows.append(SimonRow(
                lp=clean(record.get("LP")),
                sku=clean(record.get("Kod towaru")) or f"{symbol}{SKU_SUFFIX}",
                symbol=symbol,
                ekonom_name=clean(str(record.get("Nazwa towaru") or "").replace(" | ", "")),
            ))
    return rows


def split_symbol(symbol: str) -> tuple[str, str, bool]:
    """``NGZ1BCZ.01/38A`` -> (``NGZ1BCZ.01``, ``38``, True); symbol bez koloru -> (symbol, "", False)."""
    match = SYMBOL_RE.match(symbol)
    if not match:
        return symbol, "", False
    return match.group("base"), match.group("color"), bool(match.group("flap"))


def product_key(symbol: str) -> str:
    base, _, flap = split_symbol(symbol)
    return f"{base}/..A" if flap else base


def color_word(color: dict[str, Any], gender: str | None) -> str:
    name = color["name"]
    if isinstance(name, dict):
        if gender not in name:
            raise ValueError(f"Brak rodzaju '{gender}' dla koloru {name}")
        return name[gender]
    return str(name)


def build_name(symbol: str, config: dict[str, Any]) -> str:
    key = product_key(symbol)
    spec = (config.get("products") or {}).get(key)
    if spec is None:
        raise KeyError(key)
    _, color_code, _ = split_symbol(symbol)
    parts = [spec["stem"]]
    if color_code:
        color = (config.get("colors") or {}).get(color_code)
        if color is None:
            raise KeyError(f"kolor {color_code}")
        parts.append(color_word(color, spec.get("gender")))
    series = spec.get("name_series", spec.get("seria", config.get("default_series")))
    if series:
        parts.append(series)
    parts.append(symbol)
    return clean(" ".join(parts))


def build_features(symbol: str, spec: dict[str, Any], params: dict[str, str], config: dict[str, Any]) -> dict[str, str]:
    features: dict[str, str] = {
        "Dane producenta": config["producer_data"],
        "Producent": config["producent"],
    }
    series = spec.get("seria", config.get("default_series"))
    if series:
        features["Seria"] = series
    if spec.get("typ"):
        features["Typ produktu"] = spec["typ"]
        features["Podtyp produktu"] = spec["podtyp"]
    _, color_code, _ = split_symbol(symbol)
    if color_code:
        color = config["colors"][color_code]
        features["Kolor"] = color["kolor"]
        features["Kolor producenta"] = color["kolor_producenta"]
    if spec["stem"].startswith("Ramka"):
        features["Materiał"] = "Tworzywo sztuczne"
    mounting = spec.get("montaz", "Podtynkowy")
    if mounting:
        features["Sposób montażu"] = mounting
    ip = clean(params.get("Stopień ochrony IP")).replace(" ", "")
    if ip:
        features["Stopień ochrony [IP]"] = IP_FORMAT.get(ip, ip)
    if spec.get("lacznosc"):
        features["Typ łączności"] = spec["lacznosc"]
    return features


def check_against_api(row: SimonRow, spec: dict[str, Any], product: dict | None, params: dict[str, str], config: dict[str, Any]) -> None:
    if product is None:
        row.flag("DO_SPRAWDZENIA", "Brak produktu w API Kontakt-Simon - bez EAN i zdjecia")
        return
    if clean(product.get("symbol")).upper() != row.symbol.upper():
        row.flag("DO_SPRAWDZENIA", f"API zwrocilo inny symbol: {product.get('symbol')}")
    expected_series = spec.get("seria", config.get("default_series"))
    if expected_series and row.api_series and expected_series not in row.api_series:
        row.flag("WARNING", f"Seria w API '{row.api_series}' rozna od przyjetej '{expected_series}'")
    _, color_code, _ = split_symbol(row.symbol)
    if color_code:
        marker = config["colors"][color_code]["api"]
        if marker not in row.api_color.lower():
            row.flag("DO_SPRAWDZENIA", f"Kolor w API '{row.api_color}' nie zgadza sie z kodem {color_code}")
    push_in_api = "push in" in clean(params.get("montaż")).lower() or "push in" in clean(params.get("Sposób mocowania")).lower()
    push_in_name = "Push In" in spec["stem"]
    if push_in_api != push_in_name:
        row.flag("DO_SPRAWDZENIA", f"Push In w nazwie={push_in_name}, w API={push_in_api}")
    if "WKRÓTCE W SPRZEDAŻY" in clean(product.get("name")).upper():
        row.flag("WARNING", "API: WKRÓTCE W SPRZEDAŻY - produkt jeszcze niedostepny u producenta")


def public_image_url(article_id: Any, symbol: str) -> str:
    slug = symbol.lower().replace("/", "_")
    return f"https://www.kontakt-simon.com.pl/assets/catalog/img/product/{article_id},{slug}.png"


# Produkty "WKRÓTCE W SPRZEDAŻY" nie maja zdjecia, a strona Simon zwraca pod ich adresem
# z HTTP 200 grafike "brak zdjecia" (przekreslony aparat, 413x413) - nie jest to zdjecie produktu.
PLACEHOLDER_IMAGE_SHA256 = {"dce9f01b351674251e60ff33f46c2a1f29042f62ee7e24c5cee5b5de49ec9edf"}


def check_image(url: str) -> tuple[str, int, str]:
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urllib.request.urlopen(request, context=ssl._create_unverified_context(), timeout=30) as response:
            body = response.read()
            if hashlib.sha256(body).hexdigest() in PLACEHOLDER_IMAGE_SHA256:
                return url, int(getattr(response, "status", 200)), "placeholder (brak zdjecia)"
            return url, int(getattr(response, "status", 200)), clean(response.headers.get("Content-Type"))
    except Exception as exc:  # noqa: BLE001 - raport, nie przerywamy partii
        return url, 0, type(exc).__name__


def build_rows(rows: list[SimonRow], products: dict[str, dict | None], config: dict[str, Any]) -> None:
    from simon_api import product_params

    specs = config.get("products") or {}
    for row in rows:
        key = product_key(row.symbol)
        row.base, row.color, _ = split_symbol(row.symbol)
        spec = specs.get(key)
        if spec is None:
            row.flag("DO_SPRAWDZENIA", f"Brak rdzenia nazwy dla symbolu bazowego '{key}' w simon_product_names.yaml")
            row.name = row.ekonom_name
            continue
        product = products.get(row.symbol)
        params = product_params(product) if product else {}
        if product:
            row.ean = clean(product.get("ean"))
            row.api_name = clean(product.get("name"))
            row.api_series = clean(params.get("Seria"))
            row.api_color = clean(params.get("Kolor"))
            row.image = public_image_url(product.get("id"), row.symbol)
        row.name = build_name(row.symbol, config)
        row.category = spec.get("category", config["default_category"])
        row.features = build_features(row.symbol, spec, params, config)
        if spec.get("status"):
            row.flag(spec["status"], spec.get("uwaga", ""))
        check_against_api(row, spec, product, params, config)


def validate(rows: list[SimonRow]) -> dict[str, Any]:
    names = [row.name for row in rows]
    duplicates = sorted({name for name in names if names.count(name) > 1})
    eans = [row.ean for row in rows if row.ean]
    for row in rows:
        if not row.name.endswith(row.symbol):
            row.flag("DO_SPRAWDZENIA", "Nazwa nie konczy sie kodem producenta")
        if row.ean and not re.fullmatch(r"\d{13}", row.ean):
            row.flag("DO_SPRAWDZENIA", f"EAN nie ma 13 cyfr: {row.ean}")
        if row.name in duplicates:
            row.flag("DO_SPRAWDZENIA", "Powtorzona nazwa")
    return {
        "rows": len(rows),
        "duplicate_names": duplicates,
        "duplicate_eans": sorted({ean for ean in eans if eans.count(ean) > 1}),
        "missing_ean": [row.symbol for row in rows if not row.ean],
        "missing_image": [row.symbol for row in rows if not row.image],
        "max_name_length": max((len(name) for name in names), default=0),
    }


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=BASELINKER_COLUMNS, delimiter=";", lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        writer.writerows(rows)


def write_control(path: Path, rows: list[SimonRow], image_status: dict[str, tuple[int, str]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    fills = {"WARNING": PatternFill("solid", fgColor="FFF2CC"), "DO_SPRAWDZENIA": PatternFill("solid", fgColor="F8CBAD")}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Nazwy"
    headers = ["LP", "Kod towaru", "Kod producenta", "Nowa nazwa", "Dlugosc", "Nazwa Ekonom", "Nazwa API", "EAN",
               "Seria API", "Kolor API", "Kategoria", "Status", "Uwagi", "Cechy", "Zdjecie", "Zdjecie HTTP"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row.lp, row.sku, row.symbol, row.name, len(row.name), row.ekonom_name, row.api_name, row.ean,
                      row.api_series, row.api_color, row.category, row.status, "; ".join(row.notes),
                      json.dumps(row.features, ensure_ascii=False), row.image, image_status.get(row.image, (0, ""))[0] or ""])
        if row.status in fills:
            for cell in sheet[sheet.max_row]:
                cell.fill = fills[row.status]
    for column, width in zip("ABCDEFGHIJKLMNOP", [5, 22, 18, 80, 8, 60, 60, 15, 22, 16, 50, 15, 60, 60, 40, 8]):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--input", required=True, type=Path, help="Zestawienie z Ekonoma (CSV cp1250, ';').")
    parser.add_argument("--output", required=True, type=Path, help="Surowy CSV Baselinkera (przed fix_simon_attributes_to_shop_model).")
    parser.add_argument("--control", type=Path, help="Skoroszyt kontrolny nazw.")
    parser.add_argument("--report", type=Path, help="Raport JSON.")
    parser.add_argument("--names", type=Path, default=DEFAULT_NAMES)
    parser.add_argument("--refresh", action="store_true", help="Pobierz produkty z API mimo cache.")
    parser.add_argument("--skip-image-check", action="store_true")
    args = parser.parse_args(argv)

    from simon_api import SimonApiClient, is_missing, load_api_key

    config = load_names(args.names)
    rows = read_ekonom_csv(args.input)
    fetched = SimonApiClient(load_api_key()).fetch_products([row.symbol for row in rows], refresh=args.refresh)
    products = {row.symbol: (None if is_missing(fetched.get(row.symbol)) else fetched.get(row.symbol)) for row in rows}
    build_rows(rows, products, config)

    image_status: dict[str, tuple[int, str]] = {}
    if not args.skip_image_check:
        urls = sorted({row.image for row in rows if row.image})
        with concurrent.futures.ThreadPoolExecutor(max_workers=8) as executor:
            for url, status, content_type in executor.map(check_image, urls):
                image_status[url] = (status, content_type)
        for row in rows:
            status, content_type = image_status.get(row.image, (200, "image/"))
            if row.image and (status != 200 or not content_type.startswith("image/")):
                row.flag("WARNING", f"Zdjecie niedostepne ({status} {content_type})")
                row.image = ""

    summary = validate(rows)
    csv_rows = [{
        "product_id": "",
        "name": row.name,
        "sku": row.sku,
        "ean": row.ean,
        "manufacturer_name": config["manufacturer_name"],
        "category": row.category,
        "description": "",
        "features": json.dumps(row.features, ensure_ascii=False, separators=(",", ":")),
        "images_urls": row.image,
    } for row in rows]
    write_csv(args.output, csv_rows)
    if args.control:
        write_control(args.control, rows, image_status)
    summary["status"] = {status: sum(row.status == status for row in rows) for status in ("OK", "WARNING", "DO_SPRAWDZENIA")}
    summary["do_sprawdzenia"] = {row.symbol: row.notes for row in rows if row.status != "OK"}
    summary["categories"] = {category: sum(row.category == category for row in rows) for category in sorted({row.category for row in rows})}
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: summary[key] for key in ("rows", "status", "duplicate_names", "missing_ean", "max_name_length")}, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
