"""Sprawdza (i opcjonalnie poprawia) wartosci atrybutow wzgledem terminow sklepu Woo.

Dziala na kazdym pliku z atrybutami, ktory idzie do sklepu:
- eksport/import WP All Export (XLSX, kolumny ``Atrybut produktu: ...``),
- CSV Baselinkera z kolumna ``features`` (JSON).

Bez ``--output`` tylko raportuje. Z ``--output`` zapisuje KOPIE pliku, w ktorej
wartosci sa przepiete na terminy sklepu (reguly ``woo_attribute_value_rules.yaml``,
wartosc rodzenstwa z ``--shop-export``), a nowe wartosci bez odpowiednika sa
usuwane (``--keep-new-values`` je zostawia). Zmieniane sa tylko komorki
atrybutow. Plik wejsciowy nie jest nadpisywany.

Raport (``reports/``): arkusz "Nowe wartosci" z kolumna decyzji, "Kontrola
wartosci" per SKU, "Terminy do usuniecia w Woo" - wartosci, ktore partia juz
utworzyla w sklepie i ktore po imporcie poprawki zostana bez produktow.

Uzycie (partia juz w sklepie - eksport partii + eksport calej marki):
    py src/audit_shop_attribute_values.py --input input/Simon-Kaszmir.xlsx \
        --shop-export input/Sajmon-caly.xlsx \
        --output output/Simon-Kaszmir_WARTOSCI_POPRAWIONE_2026-10-12.xlsx
Samo sprawdzenie pliku Baselinkera przed importem:
    py src/audit_shop_attribute_values.py --input output/plik_baselinker.csv
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from shop_attribute_values import (
    DEFAULT_RULES,
    DEFAULT_SNAPSHOT,
    KIND_NEW,
    KIND_TYPE_LINK,
    WOO_SEPARATOR,
    ShopExport,
    ShopVocabulary,
    ValueCheck,
    attribute_columns,
    cell_text,
    check_features,
    extend_from_export,
    find_column,
    is_identifier,
    load_rules,
    load_vocabulary,
    row_features,
    sibling_key,
    summary_dict,
    write_check_sheets,
)

ROOT = Path(__file__).resolve().parents[1]
TODAY = date.today().isoformat()


@dataclass
class AuditResult:
    checks: list[ValueCheck] = field(default_factory=list)
    # indeks wiersza danych -> {kolumna albo atrybut: nowa wartosc (None = pusta)}
    changes: dict[int, dict[str, str | None]] = field(default_factory=dict)
    # (atrybut, wartosc) spoza sklepu obecne w pliku przed poprawka -> liczba produktow
    new_before: Counter = field(default_factory=Counter)
    # (atrybut, wartosc) spoza sklepu, ktore zostaja po poprawce
    new_after: Counter = field(default_factory=Counter)
    rows: int = 0


# --------------------------------------------------------------------------- I/O
def read_xlsx(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    header = [cell_text(cell) for cell in next(rows)]
    records = []
    for values in rows:
        record: dict[str, Any] = {}
        for index, name in enumerate(header):
            if name and name not in record:
                record[name] = values[index] if index < len(values) else None
        records.append(record)
    workbook.close()
    return header, records


def write_xlsx_copy(source: Path, target: Path, changes: dict[int, dict[str, str | None]]) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(source)
    sheet = workbook.worksheets[0]
    column_index: dict[str, int] = {}
    for index, cell in enumerate(sheet[1], start=1):
        name = cell_text(cell.value)
        if name and name not in column_index:
            column_index[name] = index
    for row_offset, cells in changes.items():
        for column, value in cells.items():
            # value=None w sheet.cell() nie czysci komorki - trzeba przypisac .value
            sheet.cell(row=row_offset + 2, column=column_index[column]).value = value
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)


def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]], str]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        sample = stream.readline()
        stream.seek(0)
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.DictReader(stream, delimiter=delimiter)
        rows = list(reader)
        return list(reader.fieldnames or []), rows, delimiter


def write_csv(path: Path, header: list[str], rows: list[dict[str, str]], delimiter: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=header, delimiter=delimiter, lineterminator="\r\n")
        writer.writeheader()
        writer.writerows(rows)


# --------------------------------------------------------------------------- logika
def _count_new(counter: Counter, features: dict[str, str], vocabulary: ShopVocabulary, rules: dict[str, Any]) -> None:
    approved = rules.get("approved_new_values") or {}
    for attribute, value in features.items():
        if not vocabulary.has_attribute(attribute) or is_identifier(attribute, rules):
            continue
        allowed = {str(item) for item in (approved.get(attribute) or [])}
        for part in cell_text(value).split(WOO_SEPARATOR):
            part = part.strip()
            if part and part not in allowed and not vocabulary.canonical(attribute, part):
                counter[(attribute, part)] += 1


def audit_feature_rows(
    items: list[tuple[str, str, dict[str, str]]],
    vocabulary: ShopVocabulary,
    rules: dict[str, Any],
    export: ShopExport | None = None,
    keep_new: bool = False,
) -> tuple[AuditResult, list[dict[str, str]]]:
    """``items`` = [(sku, tytul, cechy)]. Zwraca wynik kontroli i poprawione cechy."""
    result = AuditResult(rows=len(items))
    fixed_rows: list[dict[str, str]] = []
    for index, (sku, title, features) in enumerate(items):
        siblings = export.siblings.get(sibling_key(sku, rules)) if export else None
        fixed, checks = check_features(sku, title, features, vocabulary, rules, siblings=siblings, keep_new=keep_new)
        result.checks.extend(checks)
        _count_new(result.new_before, features, vocabulary, rules)
        _count_new(result.new_after, fixed, vocabulary, rules)
        diff = {
            attribute: (cell_text(fixed.get(attribute)) or None)
            for attribute in list(features) + [key for key in fixed if key not in features]
            if cell_text(features.get(attribute)) != cell_text(fixed.get(attribute))
        }
        if diff:
            result.changes[index] = diff
        fixed_rows.append(fixed)
    return result, fixed_rows


def orphaned_terms(result: AuditResult) -> list[dict[str, Any]]:
    """Wartosci spoza slownika, ktore po poprawce znikaja z pliku (w sklepie zostana puste terminy)."""
    rows = []
    for (attribute, value), count in sorted(result.new_before.items()):
        if (attribute, value) in result.new_after:
            continue
        rows.append({"atrybut": attribute, "termin": value, "produktow_w_partii": count})
    return rows


def write_control(path: Path, result: AuditResult, summary: dict[str, Any], orphans: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Podsumowanie"
    for key, value in summary.items():
        if key in ("wartosci", "typ_podtyp"):
            continue
        if isinstance(value, dict):
            overview.append([key])
            for sub_key, sub_value in value.items():
                overview.append(["", sub_key, sub_value])
        else:
            overview.append([key, ", ".join(map(str, value)) if isinstance(value, list) else value])
    overview.column_dimensions["A"].width = 36
    overview.column_dimensions["B"].width = 50
    write_check_sheets(workbook, result.checks)
    orphan_sheet = workbook.create_sheet("Terminy do usuniecia w Woo")
    orphan_sheet.append(["Atrybut", "Termin", "Produktow w partii", "Uwagi"])
    for item in orphans:
        orphan_sheet.append([item["atrybut"], item["termin"], item["produktow_w_partii"],
                             "po imporcie poprawki termin zostanie bez tych produktow - sprawdz licznik w Woo i usun"])
    for cell in orphan_sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCD", [28, 44, 18, 80]):
        orphan_sheet.column_dimensions[column].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def default_paths(input_path: Path) -> tuple[Path, Path]:
    stem = input_path.stem
    return (ROOT / "reports" / f"{stem}_kontrola_wartosci_{TODAY}.xlsx",
            ROOT / "reports" / f"{stem}_kontrola_wartosci_{TODAY}.json")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", required=True, type=Path, help="Eksport WP All Export (XLSX) albo CSV Baselinkera (features)")
    parser.add_argument("--output", type=Path, help="Poprawiona kopia pliku; bez tej opcji tylko raport")
    parser.add_argument("--shop-export", type=Path, action="append", default=[],
                        help="Eksport sklepu (XLSX) dopelniajacy slownik i rodzenstwo; produkty z --input sa pomijane")
    parser.add_argument("--snapshot", type=Path, default=DEFAULT_SNAPSHOT)
    parser.add_argument("--rules", type=Path, default=DEFAULT_RULES)
    parser.add_argument("--control", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--keep-new-values", action="store_true",
                        help="Nowe wartosci bez odpowiednika zostaja w pliku (domyslnie usuwane i raportowane)")
    args = parser.parse_args(argv)

    if args.output and args.output.resolve() == args.input.resolve():
        parser.error("Plik wyjsciowy nie moze nadpisywac wejsciowego.")
    control, report = default_paths(args.input)
    control = args.control or control
    report = args.report or report

    vocabulary = load_vocabulary(args.snapshot)
    rules = load_rules(args.rules)
    is_csv = args.input.suffix.lower() == ".csv"

    if is_csv:
        header, rows, delimiter = read_csv(args.input)
        sku_column = find_column(header, "sku", "SKU") or "sku"
        title_column = find_column(header, "name", "Title") or "name"
        items = [(cell_text(row.get(sku_column)), cell_text(row.get(title_column)), json.loads(row.get("features") or "{}"))
                 for row in rows]
        columns: dict[str, str] = {}
    else:
        header, rows = read_xlsx(args.input)
        columns = attribute_columns(header)
        sku_column = find_column(header, "Sku", "SKU") or "Sku"
        title_column = find_column(header, "Title", "name") or "Title"
        items = [(cell_text(row.get(sku_column)), cell_text(row.get(title_column)), row_features(row, columns)) for row in rows]

    batch_skus = {sku for sku, _, _ in items if sku}
    export = ShopExport()
    for path in args.shop_export:
        export_header, export_rows = read_xlsx(path)
        extend_from_export(vocabulary, export_header, export_rows, rules, exclude_skus=batch_skus, source=str(path), export=export)

    result, fixed_rows = audit_feature_rows(items, vocabulary, rules, export, keep_new=args.keep_new_values)
    orphans = orphaned_terms(result)

    if args.output:
        if is_csv:
            for row, fixed in zip(rows, fixed_rows):
                row["features"] = json.dumps(fixed, ensure_ascii=False)
            write_csv(args.output, header, rows, delimiter)
        else:
            by_attribute = {attribute: column for column, attribute in columns.items()}
            cell_changes = {
                index: {by_attribute[attribute]: value for attribute, value in diff.items() if attribute in by_attribute}
                for index, diff in result.changes.items()
            }
            write_xlsx_copy(args.input, args.output, cell_changes)

    summary = {
        "data": TODAY,
        "plik_wejsciowy": str(args.input),
        "plik_wyjsciowy": str(args.output) if args.output else "(tylko kontrola)",
        "slownik": vocabulary.sources,
        "wiersze": result.rows,
        "wiersze_zmienione": len(result.changes) if args.output else 0,
        "wartosci_spoza_sklepu_przed": sum(result.new_before.values()),
        "wartosci_spoza_sklepu_po": sum(result.new_after.values()),
        "terminy_do_usuniecia_w_woo": len(orphans),
        **summary_dict(result.checks),
        "terminy_do_usuniecia": orphans,
    }
    write_control(control, result, summary, orphans)
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    unresolved = sum(1 for check in result.checks if check.kind == KIND_NEW)
    type_links = sum(1 for check in result.checks if check.kind == KIND_TYPE_LINK)
    kept_new = sum(1 for check in result.checks if check.kind == KIND_NEW and check.new_value)
    print(f"Wiersze: {result.rows}, wartosci spoza sklepu: {sum(result.new_before.values())} -> {sum(result.new_after.values())}, "
          f"nowe bez odpowiednika: {unresolved}, Typ/Podtyp niezgodne: {type_links}, terminy do usuniecia w Woo: {len(orphans)}")
    if args.output:
        print(f"Wynik: {args.output}")
    print(f"Kontrola: {control}")
    blocking = kept_new if args.output else unresolved
    return 1 if blocking or type_links else 0


if __name__ == "__main__":
    raise SystemExit(main())
