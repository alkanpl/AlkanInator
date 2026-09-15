"""Buduje wspólny skoroszyt: priorytety Hager, decyzje kategorii i audyt Woo.

Pliki źródłowe pozostają bez zmian. Wynik trafia do katalogu output/.
"""

from __future__ import annotations

from collections import Counter
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
import yaml


ROOT = Path(__file__).resolve().parents[1]
PRIORITIES_PATH = ROOT / "output" / "Priorytety Hager.xlsx"
AUDIT_PATH = ROOT / "output" / "Hager_audyt_powielen_atrybutow_WooCommerce.xlsx"
TITLE_ANATOMY_PATH = ROOT / "configs" / "hager_title_anatomy.yaml"
OUTPUT_PATH = ROOT / "output" / "Hager_plan_kategorii_atrybutow_i_audyt_uzupelniony.xlsx"


NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
GREEN = "548235"
LIGHT_GREEN = "E2F0D9"
AMBER = "BF8F00"
LIGHT_AMBER = "FFF2CC"
RED = "C00000"
LIGHT_RED = "FCE4D6"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
GRID = "D9E1F2"

THIN_BORDER = Border(
    left=Side(style="thin", color=GRID),
    right=Side(style="thin", color=GRID),
    top=Side(style="thin", color=GRID),
    bottom=Side(style="thin", color=GRID),
)


def copy_worksheet(source, target) -> None:
    """Kopiuje dane i wygląd arkusza bez tworzenia tabel Excela."""
    for row in source.iter_rows():
        for source_cell in row:
            target_cell = target[source_cell.coordinate]
            target_cell.value = source_cell.value
            if source_cell.has_style:
                # Style IDs są lokalne dla skoroszytu, dlatego kopiujemy ich
                # składniki zamiast wewnętrznego obiektu _style.
                target_cell.font = copy(source_cell.font)
                target_cell.fill = copy(source_cell.fill)
                target_cell.border = copy(source_cell.border)
                target_cell.alignment = copy(source_cell.alignment)
                target_cell.number_format = source_cell.number_format
                target_cell.protection = copy(source_cell.protection)
            if source_cell.hyperlink:
                target_cell._hyperlink = copy(source_cell.hyperlink)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)

    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))

    for key, source_dimension in source.column_dimensions.items():
        target_dimension = target.column_dimensions[key]
        target_dimension.width = source_dimension.width
        target_dimension.hidden = source_dimension.hidden
        target_dimension.bestFit = source_dimension.bestFit
        target_dimension.outlineLevel = source_dimension.outlineLevel

    for key, source_dimension in source.row_dimensions.items():
        target_dimension = target.row_dimensions[key]
        target_dimension.height = source_dimension.height
        target_dimension.hidden = source_dimension.hidden
        target_dimension.outlineLevel = source_dimension.outlineLevel

    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.sheet_view.zoomScale = source.sheet_view.zoomScale
    target.auto_filter.ref = source.auto_filter.ref
    target.sheet_format = copy(source.sheet_format)
    target.sheet_properties = copy(source.sheet_properties)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.print_options = copy(source.print_options)
    target.sheet_view.selection = copy(source.sheet_view.selection)

    if source.data_validations:
        target.data_validations = copy(source.data_validations)
    if source.conditional_formatting:
        target.conditional_formatting = copy(source.conditional_formatting)


def set_title(ws, title: str, subtitle: str, last_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 34

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Aptos", size=10, italic=True, color=GRAY)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 34


def style_header(ws, row: int, first_col: int, last_col: int) -> None:
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row, col)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 32


def style_table_body(ws, first_row: int, last_row: int, first_col: int, last_col: int) -> None:
    for row in range(first_row, last_row + 1):
        fill = PatternFill("solid", fgColor=WHITE if row % 2 == 0 else PALE_BLUE)
        for col in range(first_col, last_col + 1):
            cell = ws.cell(row, col)
            cell.font = Font(name="Aptos", size=9, color="000000")
            cell.fill = fill
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_list_validation(ws, cell: str, options: list[str], prompt: str) -> None:
    formula = '"' + ",".join(options) + '"'
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Wybierz jedną z wartości z listy."
    validation.errorTitle = "Nieprawidłowa decyzja"
    validation.prompt = prompt
    validation.promptTitle = "Decyzja"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    ws.add_data_validation(validation)
    validation.add(ws[cell])


def category_counts(priorities_wb) -> Counter[str]:
    ws = priorities_wb["Produkty i kategorie"]
    return Counter(
        str(ws.cell(row, 4).value).strip()
        for row in range(6, ws.max_row + 1)
        if ws.cell(row, 1).value and ws.cell(row, 4).value
    )


def read_categories(priorities_wb) -> list[dict[str, object]]:
    counts = category_counts(priorities_wb)
    ws = priorities_wb["Kategorie"]
    categories: list[dict[str, object]] = []
    for row in range(6, ws.max_row + 1):
        category_id = ws.cell(row, 1).value
        if not category_id:
            continue
        path = str(ws.cell(row, 2).value)
        categories.append(
            {
                "id": str(category_id),
                "path": path,
                "name": ws.cell(row, 3).value,
                "status": ws.cell(row, 4).value,
                "count": counts[path],
            }
        )
    return categories


def read_product_types(priorities_wb) -> dict[str, dict[str, object]]:
    ws = priorities_wb["Produkty i kategorie"]
    product_types: dict[str, dict[str, object]] = {}
    for row in range(6, ws.max_row + 1):
        if not ws.cell(row, 1).value:
            continue
        product_type = str(ws.cell(row, 5).value).strip()
        if product_type not in product_types:
            product_types[product_type] = {
                "count": 0,
                "category": str(ws.cell(row, 4).value).strip(),
                "example": str(ws.cell(row, 3).value).strip(),
            }
        product_types[product_type]["count"] = int(product_types[product_type]["count"]) + 1
    return product_types


def read_title_anatomy() -> dict[str, object]:
    with TITLE_ANATOMY_PATH.open("r", encoding="utf-8") as handle:
        config = yaml.safe_load(handle)
    if not isinstance(config, dict) or not isinstance(config.get("rules"), list):
        raise ValueError(f"Nieprawidłowa konfiguracja anatomii tytułów: {TITLE_ANATOMY_PATH}")
    return config


def build_summary(ws) -> None:
    set_title(
        ws,
        "Hager — plan kategorii, priorytety atrybutów i audyt WooCommerce",
        "Jeden plik roboczy łączący plan 153 produktów, decyzje o strukturze kategorii, audyt 23 nowych atrybutów oraz reguły składania tytułów. Arkusze źródłowe pozostają niezmienione.",
        8,
    )
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A12"

    cards = [
        ("A4", "Produkty Hager", "=COUNTA('Produkty i kategorie'!A6:A500)"),
        ("C4", "Kategorie bazowe", "=COUNTA(Kategorie!A6:A100)"),
        ("E4", "Małe kategorie (<5)", '=COUNTIF(\'Decyzje kategorii\'!F6:F100,"TAK")'),
        ("G4", "Decyzje kluczowe", "=COUNTA('Decyzje kluczowe'!A6:A100)"),
        ("A7", "Atrybuty w audycie", "=COUNTA('Audyt atrybutów'!A6:A100)"),
        ("C7", "Nowe globalne — bazowo", '=COUNTIF(\'Audyt atrybutów\'!B6:B100,"NOWY GLOBALNY")'),
        ("E7", "Wariant z Układem biegunów", "=C8+1"),
        ("G7", "Do ponownego użycia / scalenia", "=A8-C8"),
    ]
    for anchor, label, formula in cards:
        col = ws[anchor].column
        row = ws[anchor].row
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.cell(row, col, label)
        ws.cell(row, col).font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        ws.cell(row, col).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(row, col).alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        ws.cell(row + 1, col, formula)
        ws.cell(row + 1, col).font = Font(name="Aptos Display", size=18, bold=True, color=NAVY)
        ws.cell(row + 1, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row + 1, col).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 23
        ws.row_dimensions[row + 1].height = 30

    ws["A11"] = "Najważniejsze rekomendacje robocze"
    ws["A11"].font = Font(name="Aptos", size=12, bold=True, color=NAVY)
    headers = ["Temat", "Rekomendacja", "Liczba produktów", "Docelowy wynik", "Status"]
    for col, value in enumerate(headers, 1):
        ws.cell(12, col, value)
    ws.merge_cells("E12:H12")
    style_header(ws, 12, 1, 8)

    rows = [
        ["Małe kategorie", "Minimum 5 produktów + wyjątki strategiczne", 14, "Decyzja osobno dla każdej grupy", "DO DECYZJI"],
        ["RCCB i RCBO", "Rozdzielić", 26, "Dwie kategorie i dwa osobne typy produktu", "DO DECYZJI"],
        ["Rozłączniki", "Połączyć", 34, "Rozłączniki i aparatura bezpiecznikowa", "DO DECYZJI"],
        ["Przekaźniki", "Połączyć", 3, "Przekaźniki modułowe", "DO DECYZJI"],
        ["Bieguny", "Zostawić Liczbę i utworzyć Układ biegunów", 137, "17 nowych atrybutów globalnych zamiast 16", "DO DECYZJI"],
        ["Icn a prąd znamionowy", "Nie scalać", 67, "Dwa różne parametry techniczne", "DO DECYZJI"],
    ]
    for row_idx, values in enumerate(rows, 13):
        ws.cell(row_idx, 1, values[0])
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=3)
        ws.cell(row_idx, 2, values[1])
        ws.cell(row_idx, 4, values[2])
        ws.merge_cells(start_row=row_idx, start_column=5, end_row=row_idx, end_column=7)
        ws.cell(row_idx, 5, values[3])
        ws.cell(row_idx, 8, values[4])
        for col in range(1, 9):
            cell = ws.cell(row_idx, col)
            cell.fill = PatternFill("solid", fgColor=WHITE if row_idx % 2 else PALE_BLUE)
            cell.border = THIN_BORDER
            cell.font = Font(name="Aptos", size=9, color="000000")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row_idx, 8).fill = PatternFill("solid", fgColor=LIGHT_AMBER)
        ws.cell(row_idx, 8).font = Font(name="Aptos", size=9, bold=True, color=AMBER)
        ws.row_dimensions[row_idx].height = 34

    ws["A21"] = "Jak korzystać z pliku"
    ws["A21"].font = Font(name="Aptos", size=12, bold=True, color=NAVY)
    notes = [
        "1. W arkuszu „Decyzje kluczowe” wybierz warianty z list rozwijanych.",
        "2. W „Decyzjach kategorii” zatwierdź osobną kategorię albo wskaż kategorię docelową.",
        "3. „Kategorie” i „Kategorie i atrybuty” pokazują plan bazowy sprzed tych rozstrzygnięć.",
        "4. „Audyt atrybutów” rozstrzyga powielenia na podstawie nazw i wartości Woo; wariant z osobnym „Układem biegunów” jest celowo pozostawiony jako decyzja.",
        "5. „Składanie tytułów” pokazuje główne i dodatkowe słowa kluczowe, kolejność atrybutów, szablon oraz przykład dla każdego z 22 typów produktów.",
    ]
    for idx, note in enumerate(notes, 22):
        ws.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=8)
        ws.cell(idx, 1, note)
        ws.cell(idx, 1).font = Font(name="Aptos", size=9, color=GRAY)
        ws.cell(idx, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(idx, 1).fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        ws.row_dimensions[idx].height = 24

    for col, width in enumerate([18, 22, 18, 16, 20, 20, 20, 18], 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def build_title_anatomy_sheet(
    ws,
    title_config: dict[str, object],
    product_types: dict[str, dict[str, object]],
) -> None:
    rules = title_config["rules"]
    configured_types = {str(rule["product_type"]) for rule in rules}
    source_types = set(product_types)
    missing = sorted(source_types - configured_types)
    extra = sorted(configured_types - source_types)
    if missing or extra:
        raise ValueError(
            "Niezgodność typów w anatomii tytułów. "
            f"Brak reguł: {missing or 'brak'}; nadmiarowe reguły: {extra or 'brak'}"
        )

    set_title(
        ws,
        "Składanie tytułów Hager — słowa kluczowe i kolejność atrybutów",
        "22 reguły odpowiadają 22 typom występującym w pliku produktowym. Frazy dodatkowe służą wyszukiwaniu i opisom — nie należy umieszczać wszystkich synonimów w jednym tytule.",
        12,
    )
    headers = [
        "ID reguły",
        "Kategoria bazowa",
        "Typ produktu",
        "Liczba produktów",
        "Główne słowo kluczowe",
        "Dodatkowe frazy SEO / synonimy",
        "Kolejność składania tytułu",
        "Elementy warunkowe",
        "Szablon tytułu",
        "Przykładowy tytuł",
        "Zasady / uwagi",
        "Status",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(5, col, value)
    style_header(ws, 5, 1, len(headers))

    start_row = 6
    for row, rule in enumerate(rules, start_row):
        product_type = str(rule["product_type"])
        source = product_types[product_type]
        ordered = "\n".join(
            f"{number}. {part}" for number, part in enumerate(rule.get("ordered_parts", []), 1)
        )
        conditional = "\n".join(f"• {part}" for part in rule.get("conditional_parts", [])) or "Brak"
        values = [
            rule["rule_id"],
            source["category"],
            product_type,
            source["count"],
            rule["primary_keyword"],
            "; ".join(rule.get("secondary_keywords", [])),
            ordered,
            conditional,
            rule["template"],
            rule.get("example") or source["example"],
            rule.get("notes", ""),
            rule.get("status", "GOTOWE"),
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)

    end_row = start_row + len(rules) - 1
    style_table_body(ws, start_row, end_row, 1, len(headers))
    for row in range(start_row, end_row + 1):
        ws.cell(row, 4).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row, 5).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
        ws.cell(row, 5).font = Font(name="Aptos", size=9, bold=True, color=GREEN)
        ws.cell(row, 7).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row, 7).font = Font(name="Aptos", size=9, bold=True, color=NAVY)
        ws.cell(row, 9).fill = PatternFill("solid", fgColor=LIGHT_AMBER)
        ws.cell(row, 9).font = Font(name="Aptos", size=9, bold=True, color="7F6000")
        ws.cell(row, 12).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
        ws.cell(row, 12).font = Font(name="Aptos", size=9, bold=True, color=GREEN)
        ws.row_dimensions[row].height = 126

    ws.freeze_panes = "E6"
    ws.auto_filter.ref = f"A5:L{end_row}"
    ws.sheet_view.showGridLines = False
    widths = [22, 50, 38, 15, 38, 48, 54, 45, 75, 70, 52, 14]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.print_title_rows = "1:5"


def build_key_decisions(ws) -> None:
    set_title(
        ws,
        "Decyzje kluczowe — Hager",
        "Rekomendacje są robocze. Żółte komórki w kolumnie „Decyzja użytkownika” służą do wyboru ostatecznego wariantu.",
        13,
    )
    headers = [
        "ID",
        "Obszar",
        "Pytanie / decyzja",
        "Produkty objęte",
        "Liczba produktów",
        "Rekomendacja robocza",
        "Jeśli wybierzemy inaczej",
        "Wpływ na sklep",
        "Dostępne opcje",
        "Decyzja użytkownika",
        "Docelowa nazwa / kategoria",
        "Uzasadnienie / uwagi",
        "Status",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(5, col, value)
    style_header(ws, 5, 1, len(headers))

    decisions = [
        {
            "id": "MIN_KATEGORIA",
            "area": "Małe kategorie",
            "question": "Czy tworzymy osobną kategorię, gdy ma mało produktów? Jeżeli nie — dokąd trafiają produkty?",
            "products": "8 grup poniżej 5 produktów",
            "count": 14,
            "recommendation": "Minimum 5 produktów, ale dopuszczać wyjątki dla kategorii o wyraźnym zastosowaniu lub planowanym rozwoju.",
            "alternative": "Przy progu 3 produktów powstaną dodatkowo małe kategorie; przy decyzji indywidualnej każdy przypadek rozstrzygamy w „Decyzjach kategorii”.",
            "impact": "Mniej pustych kategorii i prostsza nawigacja; wyjątki zachowują ważne intencje wyszukiwania.",
            "options": ["MIN. 5 + WYJĄTKI", "MIN. 3 + WYJĄTKI", "INDYWIDUALNIE"],
            "target": "Patrz arkusz „Decyzje kategorii”",
            "notes": "Mała = mniej niż 5 produktów. Kategoria z 5 produktami jest graniczna.",
        },
        {
            "id": "RCCB_RCBO",
            "area": "Wyłączniki różnicowe",
            "question": "Czy rozdzielamy wyłączniki różnicowoprądowe RCCB i różnicowo-nadprądowe RCBO?",
            "products": "13 RCCB + 13 RCBO",
            "count": 26,
            "recommendation": "Rozdzielić.",
            "alternative": "Połączyć w „Wyłączniki różnicowoprądowe”, ale koniecznie dodać filtr „Typ produktu: RCCB / RCBO”.",
            "impact": "RCBO łączy ochronę różnicową i nadprądową, ma charakterystykę B/C i osobną intencję zakupową.",
            "options": ["ROZDZIELIĆ", "POŁĄCZYĆ"],
            "target": "Dwie kategorie: RCCB oraz RCBO",
            "notes": "To dwie różne funkcje techniczne; oba zbiory mają po 13 produktów.",
        },
        {
            "id": "ROZLACZNIKI",
            "area": "Rozłączniki",
            "question": "Czy połączyć rozłączniki izolacyjne oraz rozłączniki i podstawy bezpiecznikowe w jedną kategorię?",
            "products": "29 izolacyjnych + 5 bezpiecznikowych",
            "count": 34,
            "recommendation": "Połączyć w jedną kategorię nadrzędną.",
            "alternative": "Zostawić osobne: „Rozłączniki izolacyjne” oraz „Rozłączniki i podstawy bezpiecznikowe”.",
            "impact": "Jedna mocna kategoria; różnice trzeba pokazać filtrem „Typ produktu” i „Rozmiar wkładki”.",
            "options": ["POŁĄCZYĆ", "ROZDZIELIĆ"],
            "target": "Aparatura modułowa > Rozłączniki i aparatura bezpiecznikowa",
            "notes": "Nazwa robocza obejmuje również podstawy bezpiecznikowe, a nie tylko rozłączniki.",
        },
        {
            "id": "PRZEKAZNIKI",
            "area": "Przekaźniki",
            "question": "Czy przekaźniki bistabilne i instalacyjne trzymać w jednej kategorii?",
            "products": "2 bistabilne + 1 instalacyjny",
            "count": 3,
            "recommendation": "Połączyć.",
            "alternative": "Dwie osobne kategorie po 2 i 1 produkcie.",
            "impact": "Wspólny filtr „Typ produktu” rozdzieli bistabilny i instalacyjny bez tworzenia bardzo małych kategorii.",
            "options": ["POŁĄCZYĆ", "ROZDZIELIĆ"],
            "target": "Aparatura modułowa > Przekaźniki modułowe",
            "notes": "Typ przekaźnika pozostaje P0 i powinien być w nazwie produktu.",
        },
        {
            "id": "STEROWANIE_SYGNALIZACJA",
            "area": "Sterowanie i sygnalizacja",
            "question": "Czy lampki, przyciski i dzwonki modułowe połączyć w jedną kategorię?",
            "products": "5 lampek + 2 przyciski + 1 dzwonek",
            "count": 8,
            "recommendation": "Połączyć.",
            "alternative": "Pozostawić trzy osobne kategorie, w tym dwie bardzo małe.",
            "impact": "Jedna kategoria ma 8 produktów; filtrem „Typ produktu” rozdzielimy lampkę, przycisk i dzwonek.",
            "options": ["POŁĄCZYĆ", "ROZDZIELIĆ"],
            "target": "Aparatura modułowa > Sterowanie i sygnalizacja",
            "notes": "To rekomendacja wynikająca z przyjętego progu dla małych kategorii.",
        },
        {
            "id": "BIEGUNY",
            "area": "Atrybuty biegunów",
            "question": "Czy używać „Liczby biegunów”, „Układu biegunów”, czy obu atrybutów?",
            "products": "137 produktów Hager z układem; 255 produktów Woo z liczbą",
            "count": 137,
            "recommendation": "Zostawić istniejącą „Liczbę biegunów” i utworzyć nowy globalny „Układ biegunów”.",
            "alternative": "Jeden wspólny filtr wymagałby mieszania wartości 1/2/3/4 z 1P, 1P+N, 3P+N; byłby prostszy, lecz mniej precyzyjny.",
            "impact": "Nowy „Układ biegunów” można stosować także do łączników i przełączników; liczba pozostaje uniwersalnym parametrem liczbowym.",
            "options": ["LICZBA + NOWY UKŁAD", "TYLKO LICZBA", "TYLKO UKŁAD"],
            "target": "Liczba biegunów + Układ biegunów",
            "notes": "Po wyborze pierwszej opcji audyt ma 17, a nie 16 nowych atrybutów globalnych. Wartości układu: 1P, 1P+N, 2P, 3P, 3P+N, 4P.",
        },
        {
            "id": "ICN_PRAD",
            "area": "Parametry elektryczne",
            "question": "Czy „Znamionowa zwarciowa zdolność wyłączania Icn [kA]” jest tym samym co „Prąd znamionowy [A]”?",
            "products": "67 produktów z Icn",
            "count": 67,
            "recommendation": "Nie scalać — pozostawić dwa osobne atrybuty.",
            "alternative": "Scalenie byłoby błędem: atrybuty mają inne znaczenie i inne jednostki.",
            "impact": "Prąd znamionowy [A] opisuje obciążenie robocze; Icn [kA] zdolność bezpiecznego wyłączenia zwarcia.",
            "options": ["NIE SCALAĆ", "SCALIĆ"],
            "target": "Prąd znamionowy [A] oraz Znamionowa zwarciowa zdolność wyłączania Icn [kA]",
            "notes": "Poprawiono literówkę „Prąd Zmaminowy” na „Prąd znamionowy”. Rekomendacja techniczna: nie scalać.",
        },
    ]

    start_row = 6
    for idx, decision in enumerate(decisions, start_row):
        values = [
            decision["id"],
            decision["area"],
            decision["question"],
            decision["products"],
            decision["count"],
            decision["recommendation"],
            decision["alternative"],
            decision["impact"],
            "; ".join(decision["options"]),
            None,
            decision["target"],
            decision["notes"],
            f'=IF(J{idx}="","DO DECYZJI","UZGODNIONO")',
        ]
        for col, value in enumerate(values, 1):
            ws.cell(idx, col, value)
        add_list_validation(ws, f"J{idx}", decision["options"], "Wybierz ostateczny wariant.")

    end_row = start_row + len(decisions) - 1
    style_table_body(ws, start_row, end_row, 1, len(headers))
    for row in range(start_row, end_row + 1):
        ws.cell(row, 5).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row, 10).fill = PatternFill("solid", fgColor=LIGHT_AMBER)
        ws.cell(row, 10).font = Font(name="Aptos", size=9, bold=True, color=AMBER)
        ws.cell(row, 13).fill = PatternFill("solid", fgColor=LIGHT_AMBER)
        ws.cell(row, 13).font = Font(name="Aptos", size=9, bold=True, color=AMBER)
        ws.row_dimensions[row].height = 88

    ws.conditional_formatting.add(
        f"M{start_row}:M{end_row}",
        FormulaRule(formula=[f'$M{start_row}="UZGODNIONO"'], fill=PatternFill("solid", fgColor=LIGHT_GREEN)),
    )
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:M{end_row}"
    ws.sheet_view.showGridLines = False
    widths = [22, 24, 45, 35, 15, 46, 46, 46, 34, 30, 46, 52, 18]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def category_recommendations() -> dict[str, dict[str, str]]:
    return {
        "mcb": {
            "decision": "OSOBNA KATEGORIA",
            "target": "Aparatura modułowa > Wyłączniki nadprądowe",
            "fallback": "Aparatura modułowa",
            "reason": "54 produkty i bardzo wyraźna intencja zakupowa.",
        },
        "rccb": {
            "decision": "OSOBNA KATEGORIA",
            "target": "Aparatura modułowa > Wyłączniki różnicowoprądowe RCCB",
            "fallback": "Aparatura modułowa > Wyłączniki różnicowoprądowe",
            "reason": "13 produktów; inna funkcja niż RCBO.",
        },
        "rcbo": {
            "decision": "OSOBNA KATEGORIA",
            "target": "Aparatura modułowa > Wyłączniki różnicowo-nadprądowe RCBO",
            "fallback": "Aparatura modułowa > Wyłączniki różnicowoprądowe",
            "reason": "13 produktów; osobna funkcja i osobna intencja zakupowa.",
        },
        "isolators": {
            "decision": "POŁĄCZYĆ",
            "target": "Aparatura modułowa > Rozłączniki i aparatura bezpiecznikowa",
            "fallback": "Aparatura modułowa > Rozłączniki izolacyjne",
            "reason": "Wariant wskazany do decyzji; wspólna grupa osiągnie 34 produkty.",
        },
        "fuse_holders": {
            "decision": "POŁĄCZYĆ",
            "target": "Aparatura modułowa > Rozłączniki i aparatura bezpiecznikowa",
            "fallback": "Aparatura modułowa > Rozłączniki izolacyjne",
            "reason": "Tylko 5 produktów; rozdzielenie daje cienką kategorię.",
        },
        "contactors": {
            "decision": "OSOBNA KATEGORIA",
            "target": "Aparatura modułowa > Styczniki",
            "fallback": "Aparatura modułowa",
            "reason": "6 produktów i jednoznaczna, techniczna grupa produktowa.",
        },
        "bistable_relays": {
            "decision": "POŁĄCZYĆ",
            "target": "Aparatura modułowa > Przekaźniki modułowe",
            "fallback": "Aparatura modułowa",
            "reason": "2 produkty; wspólnie z przekaźnikiem instalacyjnym tworzą grupę 3 produktów.",
        },
        "installation_relays": {
            "decision": "POŁĄCZYĆ",
            "target": "Aparatura modułowa > Przekaźniki modułowe",
            "fallback": "Aparatura modułowa",
            "reason": "1 produkt; typ przekaźnika pozostaje filtrem i elementem nazwy.",
        },
        "changeover_switches": {
            "decision": "OSOBNA KATEGORIA",
            "target": "Aparatura modułowa > Przełączniki Sieć-Agregat",
            "fallback": "Aparatura modułowa",
            "reason": "7 produktów i odrębne zastosowanie.",
        },
        "spd": {
            "decision": "OSOBNA KATEGORIA",
            "target": "Aparatura modułowa > Ochronniki przepięciowe",
            "fallback": "Aparatura modułowa",
            "reason": "Wyjątek od progu: tylko 3 produkty, ale mocna intencja wyszukiwania i istniejąca kategoria.",
        },
        "busbars": {
            "decision": "OSOBNA KATEGORIA",
            "target": "Aparatura modułowa > Szyny łączeniowe",
            "fallback": "Aparatura modułowa",
            "reason": "7 produktów; do tej kategorii można dołączyć 2 akcesoria.",
        },
        "busbar_accessories": {
            "decision": "POŁĄCZYĆ",
            "target": "Aparatura modułowa > Szyny łączeniowe",
            "fallback": "Aparatura modułowa > Akcesoria do aparatury modułowej",
            "reason": "2 produkty bez samodzielnej masy; kompatybilność i typ produktu rozróżnią akcesoria.",
        },
        "indicators": {
            "decision": "POŁĄCZYĆ",
            "target": "Aparatura modułowa > Sterowanie i sygnalizacja",
            "fallback": "Aparatura modułowa > Lampki sygnalizacyjne",
            "reason": "Razem z przyciskami i dzwonkiem grupa ma 8 produktów.",
        },
        "control_buttons": {
            "decision": "POŁĄCZYĆ",
            "target": "Aparatura modułowa > Sterowanie i sygnalizacja",
            "fallback": "Aparatura modułowa",
            "reason": "2 produkty; typ produktu wystarczy do filtrowania.",
        },
        "acoustic_signaling": {
            "decision": "POŁĄCZYĆ",
            "target": "Aparatura modułowa > Sterowanie i sygnalizacja",
            "fallback": "Aparatura modułowa",
            "reason": "1 produkt; osobna kategoria byłaby zbyt mała.",
        },
        "din_sockets": {
            "decision": "ZOSTAWIĆ W NADRZĘDNEJ",
            "target": "Aparatura modułowa",
            "fallback": "Aparatura modułowa",
            "reason": "1 produkt; nie jest akcesorium i na razie nie uzasadnia osobnej kategorii.",
        },
        "apparatus_accessories": {
            "decision": "OSOBNA KATEGORIA",
            "target": "Aparatura modułowa > Akcesoria do aparatury modułowej",
            "fallback": "Aparatura modułowa",
            "reason": "Wyjątek strategiczny: 2 produkty teraz, ale grupa jest potrzebna do dalszej rozbudowy i cross-sellingu.",
        },
    }


def build_category_decisions(ws, categories: list[dict[str, object]]) -> None:
    set_title(
        ws,
        "Decyzje kategorii — liczebność i miejsca docelowe",
        "Rekomendacja stosuje próg 5 produktów z wyjątkami strategicznymi. Kolumny „Decyzja użytkownika” i „Ścieżka po decyzji” są przeznaczone do zatwierdzenia.",
        13,
    )
    headers = [
        "ID kategorii",
        "Ścieżka bazowa",
        "Nazwa robocza",
        "Status bazowy",
        "Liczba produktów",
        "Mała (<5)?",
        "Rekomendacja",
        "Rekomendowana ścieżka końcowa",
        "Jeżeli nie osobna — przenieść do",
        "Powód",
        "Decyzja użytkownika",
        "Ścieżka po decyzji",
        "Status",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(5, col, value)
    style_header(ws, 5, 1, len(headers))

    recs = category_recommendations()
    start_row = 6
    for row, category in enumerate(categories, start_row):
        rec = recs[category["id"]]
        values = [
            category["id"],
            category["path"],
            category["name"],
            category["status"],
            category["count"],
            "TAK" if int(category["count"]) < 5 else "NIE",
            rec["decision"],
            rec["target"],
            rec["fallback"],
            rec["reason"],
            None,
            None,
            f'=IF(K{row}="","DO DECYZJI","UZGODNIONO")',
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        add_list_validation(
            ws,
            f"K{row}",
            ["OSOBNA KATEGORIA", "POŁĄCZYĆ", "ZOSTAWIĆ W NADRZĘDNEJ"],
            "Wybierz sposób docelowego przypisania kategorii.",
        )

    end_row = start_row + len(categories) - 1
    style_table_body(ws, start_row, end_row, 1, len(headers))
    for row in range(start_row, end_row + 1):
        ws.cell(row, 5).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row, 6).alignment = Alignment(horizontal="center", vertical="top")
        if ws.cell(row, 6).value == "TAK":
            ws.cell(row, 6).fill = PatternFill("solid", fgColor=LIGHT_AMBER)
            ws.cell(row, 6).font = Font(name="Aptos", size=9, bold=True, color=AMBER)
        ws.cell(row, 11).fill = PatternFill("solid", fgColor=LIGHT_AMBER)
        ws.cell(row, 11).font = Font(name="Aptos", size=9, bold=True, color=AMBER)
        ws.cell(row, 12).fill = PatternFill("solid", fgColor=LIGHT_AMBER)
        ws.cell(row, 13).fill = PatternFill("solid", fgColor=LIGHT_AMBER)
        ws.cell(row, 13).font = Font(name="Aptos", size=9, bold=True, color=AMBER)
        ws.row_dimensions[row].height = 54

    ws.conditional_formatting.add(
        f"E{start_row}:E{end_row}",
        CellIsRule(operator="lessThan", formula=["5"], fill=PatternFill("solid", fgColor=LIGHT_AMBER)),
    )
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:M{end_row}"
    ws.sheet_view.showGridLines = False
    widths = [22, 48, 34, 18, 15, 14, 24, 52, 48, 50, 26, 52, 18]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def build_workbook() -> Path:
    priorities_wb = load_workbook(PRIORITIES_PATH, data_only=False)
    audit_wb = load_workbook(AUDIT_PATH, data_only=False)
    categories = read_categories(priorities_wb)
    product_types = read_product_types(priorities_wb)
    title_config = read_title_anatomy()

    output_wb = Workbook()
    output_wb.remove(output_wb.active)

    summary = output_wb.create_sheet("Podsumowanie")
    build_summary(summary)

    key_decisions = output_wb.create_sheet("Decyzje kluczowe")
    build_key_decisions(key_decisions)

    category_decisions = output_wb.create_sheet("Decyzje kategorii")
    build_category_decisions(category_decisions, categories)

    title_anatomy = output_wb.create_sheet("Składanie tytułów")
    build_title_anatomy_sheet(title_anatomy, title_config, product_types)

    for source_name in [
        "Kategorie",
        "Kategorie i atrybuty",
        "Produkty i kategorie",
        "Słownik priorytetów",
    ]:
        target = output_wb.create_sheet(source_name)
        copy_worksheet(priorities_wb[source_name], target)

    audit_sheet_map = {
        "Decyzje": "Audyt atrybutów",
        "Dowody z wartości": "Dowody z wartości",
        "Problemy danych": "Problemy danych",
        "Istniejące Woo": "Istniejące Woo",
        "Metodyka i źródła": "Metodyka i źródła",
    }
    for source_name, target_name in audit_sheet_map.items():
        target = output_wb.create_sheet(target_name)
        copy_worksheet(audit_wb[source_name], target)

    # Kolor zakładek ułatwia odróżnienie części decyzyjnej od danych źródłowych.
    for name in ["Podsumowanie", "Decyzje kluczowe", "Decyzje kategorii", "Składanie tytułów"]:
        output_wb[name].sheet_properties.tabColor = NAVY
    for name in ["Kategorie", "Kategorie i atrybuty", "Produkty i kategorie", "Słownik priorytetów"]:
        output_wb[name].sheet_properties.tabColor = BLUE
    for name in ["Audyt atrybutów", "Dowody z wartości", "Problemy danych", "Istniejące Woo", "Metodyka i źródła"]:
        output_wb[name].sheet_properties.tabColor = AMBER

    output_wb.calculation.fullCalcOnLoad = True
    output_wb.calculation.forceFullCalc = True
    output_wb.calculation.calcMode = "auto"
    output_wb.active = 0
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_workbook())
