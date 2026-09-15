from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from update_hager_berker_series_categories import (
    CATALOG_HEADER_ROW,
    CATALOG_PRODUCTS_SHEET,
    EXPECTED_NEW_LEAVES,
    MANUAL_SERIES_OVERRIDES,
    ROOT_CATEGORY,
    SERIES_ORDER,
    collect_workbook_structure,
    desired_category_leaves,
    header_map,
    load_catalog_series,
    normalize_catalog_series,
    normalize_series,
    scan_mojibake,
    sorted_series,
    split_pipe,
    unique_in_order,
    update_categories,
)
from update_hager_berker_titles_for_series import (
    title_represents_exact_series,
    update_title,
)


DEFAULT_INPUT = Path("input/Hager-Berker-Gniazdka.xlsx")
DEFAULT_CATEGORY_MAP = Path("input/BerkerHagerKategorie.xlsx")
DEFAULT_CATALOG = Path("output/Hager-Berker-Katalog-PDF_BAZA_KOMPATYBILNOSCI.xlsx")
DEFAULT_OUTPUT = Path(
    "output/Hager-Berker-Gniazdka_SERIE_KATEGORIE_OPISY_POPRAWIONE_2026-08-31.xlsx"
)
DEFAULT_JSON_REPORT = Path(
    "reports/Hager-Berker-korekta-serii-kategorii-opisow-2026-08-31.json"
)
DEFAULT_XLSX_REPORT = Path(
    "reports/Hager-Berker-RAPORT_KOREKTY_SERII_KATEGORII_OPISOW_2026-08-31.xlsx"
)

STORE_SHEET = "Worksheet"
HEADER_ROW = 1
TARGET_COLUMNS = {
    "Title",
    "Content",
    "Atrybut Produktu: Seria",
    "Product Tags",
    "Kategorie produktów",
}

SERIES_TOKEN_PATTERNS: list[tuple[str, str]] = [
    ("Serie 1930 Rosenthal", r"(?<!\w)Serie\s+1930\s+Rosenthal(?!\w)"),
    ("Lumina passion", r"(?<!\w)Lumina\s+passion(?!\w)"),
    ("Lumina intense", r"(?<!\w)Lumina\s+intense(?!\w)"),
    ("Lumina soul", r"(?<!\w)Lumina\s+soul(?!\w)"),
    ("B. Kwadrat", r"(?<!\w)B\.\s*Kwadrat(?!\w)"),
    ("R.classic", r"(?<!\w)R\.classic(?!\w)"),
    ("Serie 1930", r"(?<!\w)Serie\s+1930(?!\s+Rosenthal\b)(?!\w)"),
    ("Glasserie", r"(?<!\w)Glasserie(?!\w)"),
    ("one.platform", r"(?<!\w)one\.platform(?!\w)"),
    ("Integro Flow", r"(?<!\w)Integro\s+Flow(?!\w)"),
    ("B.3", r"(?<!\w)B\.3(?!\w)"),
    ("B.7", r"(?<!\w)B\.7(?!\w)"),
    ("S.1", r"(?<!\w)S\.1(?!\w)"),
    ("K.1", r"(?<!\w)K\.1(?!\w)"),
    ("K.5", r"(?<!\w)K\.5(?!\w)"),
    ("Q.1", r"(?<!\w)Q\.1(?!\w)"),
    ("Q.3", r"(?<!\w)Q\.3(?!\w)"),
    ("Q.7", r"(?<!\w)Q\.7(?!\w)"),
    ("R.1", r"(?<!\w)R\.1(?!\w)"),
    ("R.3", r"(?<!\w)R\.3(?!\w)"),
    ("R.8", r"(?<!\w)R\.8(?!\w)"),
    ("W.1", r"(?<!\w)W\.1(?!\w)"),
    ("TS", r"(?<!\w)(?:Berker\s+)?TS(?!\w)"),
    ("Lumina", r"(?<!\w)Lumina(?!\s+(?:soul|intense|passion)\b)(?!\w)"),
]

SERIES_CLUSTER_TOKEN = (
    r"(?:Serie\s+1930\s+Rosenthal|Serie\s+1930|1930|"
    r"Lumina\s+(?:soul|intense|passion)|Lumina|"
    r"B\.\s*Kwadrat|B\.\s*[137]|Q\.\s*[137]|R\.\s*(?:classic|[138])|"
    r"K\.\s*[15]|S\.\s*1|Glasserie|Glas|one\.platform|TS)"
)
SERIES_SLASH_CLUSTER_PATTERN = (
    rf"(?<!\w){SERIES_CLUSTER_TOKEN}(?:\s*/\s*{SERIES_CLUSTER_TOKEN})+(?!\w)"
)

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
WHITE = "FFFFFF"
TEXT = "1F2937"
THIN_GRAY = Side(style="thin", color="D9E1F2")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_direct_catalog_series(
    catalog_path: Path,
) -> tuple[dict[str, list[str]], dict[str, list[dict[str, Any]]]]:
    workbook = load_workbook(catalog_path, read_only=True, data_only=True)
    try:
        sheet = workbook[CATALOG_PRODUCTS_SHEET]
        headers = header_map(sheet, CATALOG_HEADER_ROW)
        evidence_by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in sheet.iter_rows(min_row=CATALOG_HEADER_ROW + 1, values_only=True):
            get = lambda name: row[headers[name] - 1]
            code = str(get("Kod katalogowy") or "").strip()
            if not code:
                continue
            source = "Seria / system"
            raw = str(get(source) or "").strip()
            series = normalize_catalog_series(raw)
            if not series:
                source = "Sekcja strony"
                raw = str(get(source) or "").strip()
                series = normalize_catalog_series(raw)
            if not series:
                continue
            evidence_by_code[code.casefold()].append(
                {
                    "series": series,
                    "status": str(get("Status ekstrakcji") or "").strip(),
                    "page_pdf": get("Strona PDF"),
                    "page_print": get("Strona drukowana"),
                    "source": source,
                    "raw": raw,
                }
            )
    finally:
        workbook.close()

    result: dict[str, list[str]] = {}
    chosen_evidence: dict[str, list[dict[str, Any]]] = {}
    for code, rows in evidence_by_code.items():
        ok_rows = [row for row in rows if row["status"] == "OK"]
        chosen = ok_rows or rows
        result[code] = sorted_series(
            token for row in chosen for token in row["series"]
        )
        chosen_evidence[code] = chosen
    return result, chosen_evidence


def load_current_category_leaves(category_map_path: Path) -> list[str]:
    workbook = load_workbook(category_map_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Arkusz1"]
        marker = next(
            (
                cell
                for row in sheet.iter_rows()
                for cell in row
                if str(cell.value or "").strip() == "Nowe"
            ),
            None,
        )
        if marker is None:
            raise ValueError("Nie znaleziono sekcji 'Nowe' w mapie kategorii")
        leaf_column = marker.column + 2
        leaves = [
            str(sheet.cell(row_number, leaf_column).value or "").strip()
            for row_number in range(marker.row + 1, sheet.max_row + 1)
        ]
        return [value for value in leaves if value]
    finally:
        workbook.close()


def replace_or_remove_pattern(text: str, pattern: str, replacement: str) -> str:
    return re.sub(pattern, replacement, text, flags=re.IGNORECASE)


def tidy_series_text(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\s*/\s*(?:/\s*)+", "/", text)
    text = re.sub(r"(?<=\s)/\s*", "", text)
    text = re.sub(r"\s*/\s*(?=[,.;:<\)])", "", text)
    text = re.sub(r",\s*,+", ",", text)
    text = re.sub(r"\b(?:i|oraz)\s*(?=[,.;:<\)])", "", text, flags=re.IGNORECASE)
    text = re.sub(r"([,:;])\s*(?:i|oraz)\s+", r"\1 ", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+([,.;:])", r"\1", text)
    text = re.sub(r"\(\s*\)", "", text)
    text = re.sub(
        r"(?<!\w)B\.\s*Kwadrat\s*([/,])\s*B\.\s*Kwadrat(?!\w)",
        r"B.Kwadrat\1",
        text,
        flags=re.IGNORECASE,
    )
    return text


def clean_series_mentions(text: str, exact_series: Iterable[str]) -> tuple[str, list[str]]:
    if not text:
        return text, []
    allowed = set(exact_series)
    result = text
    removed: list[str] = []
    canonical = "/".join(
        "B.Kwadrat" if token == "B. Kwadrat" else token for token in exact_series
    )

    def replace_incompatible_cluster(match: re.Match[str]) -> str:
        cluster = match.group(0)
        mentioned = series_mentions(cluster)
        incompatible = mentioned - allowed_mentions(allowed)
        if not incompatible:
            return cluster
        removed.extend(sorted(incompatible))
        return canonical

    result = re.sub(
        SERIES_SLASH_CLUSTER_PATTERN,
        replace_incompatible_cluster,
        result,
        flags=re.IGNORECASE,
    )

    if "Serie 1930" not in allowed and re.search(
        r"(?i)\bseri(?:a|i|ami)\s+1930\b", result
    ):
        result = re.sub(
            r"(?i)(\bseri(?:a|i|ami)\s+)1930\b",
            lambda match: match.group(1) + canonical,
            result,
        )
        removed.append("Serie 1930")

    if re.search(r"(?<!\w)B\.1(?!\w)", result, flags=re.IGNORECASE):
        replacement = "B.Kwadrat" if "B. Kwadrat" in allowed else ""
        result = replace_or_remove_pattern(result, r"(?<!\w)B\.1(?!\w)", replacement)
        removed.append("B.1")

    if re.search(r"(?<!\w)Glas(?!\w)", result, flags=re.IGNORECASE):
        replacement = "Glasserie" if "Glasserie" in allowed else ""
        result = replace_or_remove_pattern(result, r"(?<!\w)Glas(?!\w)", replacement)
        removed.append("Glas")

    for token, pattern in SERIES_TOKEN_PATTERNS:
        permitted = token in allowed
        if token == "Lumina" and any(item.startswith("Lumina ") for item in allowed):
            permitted = True
        if permitted:
            continue
        if re.search(pattern, result, flags=re.IGNORECASE):
            result = replace_or_remove_pattern(result, pattern, "")
            removed.append(token)

    if result == text:
        return text, unique_in_order(removed)
    return tidy_series_text(result), unique_in_order(removed)


def series_mentions(text: str) -> set[str]:
    found: set[str] = set()
    for token, pattern in SERIES_TOKEN_PATTERNS:
        if re.search(pattern, text or "", flags=re.IGNORECASE):
            found.add(token)
    if re.search(r"(?<!\w)B\.1(?!\w)", text or "", flags=re.IGNORECASE):
        found.add("B.1")
    if re.search(r"(?<!\w)Glas(?!\w)", text or "", flags=re.IGNORECASE):
        found.add("Glas")
    if re.search(
        r"(?i)(?:\bseri(?:a|i|ami)\s+|(?<=/))1930\b|\b1930(?=\s*/\s*(?:Glasserie|Glas)\b)",
        text or "",
    ):
        found.add("Serie 1930")
    return found


def allowed_mentions(exact_series: Iterable[str]) -> set[str]:
    allowed = set(exact_series)
    if any(item.startswith("Lumina ") for item in allowed):
        allowed.add("Lumina")
    return allowed


def duplicate_groups(sheet, sku_col: int, title_col: int) -> set[frozenset[str]]:
    titles: dict[str, list[str]] = defaultdict(list)
    for row_number in range(2, sheet.max_row + 1):
        sku = str(sheet.cell(row_number, sku_col).value or "").strip()
        title = str(sheet.cell(row_number, title_col).value or "").strip()
        if title:
            titles[title.casefold()].append(sku)
    return {
        frozenset(skus) for skus in titles.values() if len(skus) > 1
    }


def build_corrected_workbook(
    input_path: Path,
    category_map_path: Path,
    catalog_path: Path,
    output_path: Path,
) -> dict[str, Any]:
    source_hash = sha256(input_path)
    category_map_hash = sha256(category_map_path)
    catalog_hash = sha256(catalog_path)
    current_category_leaves = load_current_category_leaves(category_map_path)
    if current_category_leaves != EXPECTED_NEW_LEAVES:
        raise ValueError(
            "Reguły kategorii nie odpowiadają bieżącej liście w "
            f"{category_map_path}: {current_category_leaves}"
        )
    catalog_series, catalog_diagnostics = load_catalog_series(catalog_path)
    direct_series, direct_evidence = build_direct_catalog_series(catalog_path)

    workbook = load_workbook(input_path)
    input_structure = collect_workbook_structure(workbook)
    if workbook.sheetnames != [STORE_SHEET]:
        raise ValueError(f"Nieoczekiwane arkusze: {workbook.sheetnames}")
    sheet = workbook[STORE_SHEET]
    headers = header_map(sheet, HEADER_ROW)
    required = {"SKU", *TARGET_COLUMNS}
    if not required.issubset(headers):
        raise ValueError(f"Brak kolumn: {sorted(required - set(headers))}")

    sku_col = headers["SKU"]
    title_col = headers["Title"]
    content_col = headers["Content"]
    series_col = headers["Atrybut Produktu: Seria"]
    tags_col = headers["Product Tags"]
    category_col = headers["Kategorie produktów"]
    allowed_column_numbers = {
        title_col,
        content_col,
        series_col,
        tags_col,
        category_col,
    }

    source_duplicates = duplicate_groups(sheet, sku_col, title_col)
    changes: dict[str, list[dict[str, Any]]] = {
        "series": [],
        "categories": [],
        "titles": [],
        "content": [],
        "tags": [],
    }
    unresolved: list[dict[str, Any]] = []
    matched_existing = 0
    matched_direct = 0
    seen_skus: set[str] = set()

    for row_number in range(2, sheet.max_row + 1):
        sku = str(sheet.cell(row_number, sku_col).value or "").strip()
        if not sku:
            continue
        if sku in seen_skus:
            raise ValueError(f"Powtórzone SKU: {sku}")
        seen_skus.add(sku)
        base_code = sku.split("/", 1)[0].strip().casefold()
        current_series_text = str(sheet.cell(row_number, series_col).value or "").strip()
        current_series = normalize_series(current_series_text)

        evidence: list[dict[str, Any]] = []
        if sku in catalog_series:
            final_series = catalog_series[sku]
            series_source = "PDF_PRODUCT_ROW_OR_ALIAS"
            matched_existing += 1
        elif base_code in direct_series:
            final_series = direct_series[base_code]
            evidence = direct_evidence[base_code]
            series_source = "PDF_PRODUCT_ROW_DIRECT_CODE"
            matched_direct += 1
        elif sku in MANUAL_SERIES_OVERRIDES:
            final_series = MANUAL_SERIES_OVERRIDES[sku]
            series_source = "PDF_VISUAL_OVERRIDE"
        else:
            final_series = current_series
            series_source = "CURRENT_ATTRIBUTE_NO_PDF_MATCH"
            unresolved.append(
                {
                    "row": row_number,
                    "sku": sku,
                    "series": current_series_text,
                    "title": str(sheet.cell(row_number, title_col).value or "").strip(),
                    "source": "BRAK_BEZPOSREDNIEGO_DOWODU_PDF",
                    "confidence": 0.0,
                    "status": "DO_SPRAWDZENIA",
                }
            )

        final_series = sorted_series(final_series)
        if set(final_series) == set(current_series):
            final_series = current_series
        final_series_text = "|".join(final_series)
        if final_series_text != current_series_text:
            changes["series"].append(
                {
                    "row": row_number,
                    "sku": sku,
                    "before": current_series_text,
                    "after": final_series_text,
                    "source": series_source,
                    "confidence": 1.0,
                    "status": "OK",
                    "evidence": evidence,
                }
            )
            sheet.cell(row_number, series_col).value = final_series_text or None

        before_category = str(sheet.cell(row_number, category_col).value or "").strip()
        after_category, added, removed = update_categories(
            before_category, sku, final_series
        )
        if after_category != before_category:
            changes["categories"].append(
                {
                    "row": row_number,
                    "sku": sku,
                    "series": final_series_text,
                    "before": before_category,
                    "after": after_category,
                    "added": added,
                    "removed": removed,
                    "source": "ATTRIBUTE_SERIES_TO_CURRENT_CATEGORY_MAP",
                    "confidence": 1.0,
                    "status": "OK",
                }
            )
            sheet.cell(row_number, category_col).value = after_category

        before_title = str(sheet.cell(row_number, title_col).value or "").strip()
        title_missing, title_extra = title_represents_exact_series(
            before_title, final_series
        )
        lumina_without_hager = any(
            token.startswith("Lumina ") for token in final_series
        ) and not re.search(r"\bHager\b", before_title, flags=re.IGNORECASE)
        should_update_title = (
            set(final_series) != set(current_series)
            or bool(title_missing)
            or bool(title_extra)
            or lumina_without_hager
        )
        after_title = (
            update_title(before_title, sku, final_series)
            if should_update_title
            and [token for token in final_series if token != "Inne"]
            else before_title
        )
        if after_title != before_title:
            changes["titles"].append(
                {
                    "row": row_number,
                    "sku": sku,
                    "series": final_series_text,
                    "before": before_title,
                    "after": after_title,
                    "source": "ATTRIBUTE_SERIES_TITLE_STANDARD",
                    "confidence": 1.0,
                    "status": "OK",
                }
            )
            sheet.cell(row_number, title_col).value = after_title

        before_content = str(sheet.cell(row_number, content_col).value or "")
        after_content, removed_content_tokens = clean_series_mentions(
            before_content, final_series
        )
        if after_content != before_content:
            changes["content"].append(
                {
                    "row": row_number,
                    "sku": sku,
                    "series": final_series_text,
                    "removed_or_normalized": removed_content_tokens,
                    "before": before_content,
                    "after": after_content,
                    "source": "ATTRIBUTE_SERIES_CONTENT_NORMALIZATION",
                    "confidence": 1.0,
                    "status": "OK",
                }
            )
            sheet.cell(row_number, content_col).value = after_content

        before_tags = str(sheet.cell(row_number, tags_col).value or "")
        after_tags, removed_tag_tokens = clean_series_mentions(before_tags, final_series)
        if after_tags != before_tags:
            changes["tags"].append(
                {
                    "row": row_number,
                    "sku": sku,
                    "series": final_series_text,
                    "removed_or_normalized": removed_tag_tokens,
                    "before": before_tags,
                    "after": after_tags,
                    "source": "ATTRIBUTE_SERIES_TAG_NORMALIZATION",
                    "confidence": 1.0,
                    "status": "OK",
                }
            )
            sheet.cell(row_number, tags_col).value = after_tags or None

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    if (
        sha256(input_path) != source_hash
        or sha256(category_map_path) != category_map_hash
        or sha256(catalog_path) != catalog_hash
    ):
        raise RuntimeError("Plik źródłowy zmienił się podczas operacji")

    source = load_workbook(input_path, read_only=False, data_only=False)
    result = load_workbook(output_path, read_only=False, data_only=False)
    validation = {
        "source_files_unchanged": True,
        "sheet_structure_preserved": False,
        "only_allowed_columns_changed": False,
        "styles_preserved": False,
        "no_b1_in_product_text": False,
        "no_ambiguous_bare_glas_in_product_text": False,
        "titles_match_exact_series": False,
        "categories_match_exact_series": False,
        "descriptions_do_not_claim_extra_series": False,
        "no_new_duplicate_title_groups": False,
        "mojibake_scan_passed": False,
    }
    try:
        if collect_workbook_structure(result) != input_structure:
            raise ValueError("Struktura skoroszytu nie została zachowana")
        validation["sheet_structure_preserved"] = True
        source_sheet = source[STORE_SHEET]
        result_sheet = result[STORE_SHEET]
        unexpected: list[str] = []
        for row_number in range(1, source_sheet.max_row + 1):
            for column_number in range(1, source_sheet.max_column + 1):
                before = source_sheet.cell(row_number, column_number)
                after = result_sheet.cell(row_number, column_number)
                if before.value != after.value and column_number not in allowed_column_numbers:
                    unexpected.append(before.coordinate)
                if before.style_id != after.style_id:
                    raise ValueError(f"Zmieniono styl komórki {before.coordinate}")
        if unexpected:
            raise ValueError(f"Zmiany poza dozwolonymi kolumnami: {unexpected[:20]}")
        validation["only_allowed_columns_changed"] = True
        validation["styles_preserved"] = True

        bad_b1: list[str] = []
        bad_glas: list[str] = []
        bad_titles: list[dict[str, Any]] = []
        bad_categories: list[dict[str, Any]] = []
        bad_descriptions: list[dict[str, Any]] = []
        for row_number in range(2, result_sheet.max_row + 1):
            sku = str(result_sheet.cell(row_number, sku_col).value or "").strip()
            title = str(result_sheet.cell(row_number, title_col).value or "").strip()
            content = str(result_sheet.cell(row_number, content_col).value or "")
            tags = str(result_sheet.cell(row_number, tags_col).value or "")
            categories = str(result_sheet.cell(row_number, category_col).value or "")
            raw_series = str(
                result_sheet.cell(row_number, series_col).value or ""
            )
            exact_series = normalize_series(raw_series)
            allowed = allowed_mentions(exact_series)

            joined_text = " ".join((title, content, tags, raw_series))
            if re.search(r"(?<!\w)B\.1(?!\w)", joined_text, flags=re.IGNORECASE):
                bad_b1.append(sku)
            if re.search(r"(?<!\w)Glas(?!\w)", joined_text, flags=re.IGNORECASE):
                bad_glas.append(sku)

            missing, extra = title_represents_exact_series(title, exact_series)
            if missing or extra:
                bad_titles.append(
                    {"sku": sku, "missing": missing, "extra": extra, "title": title}
                )

            desired = set(desired_category_leaves(exact_series))
            present = set()
            for path in split_pipe(categories):
                cleaned = re.sub(r"\s*>\s*", ">", path)
                if cleaned.startswith(ROOT_CATEGORY + ">"):
                    leaf = cleaned[len(ROOT_CATEGORY) + 1 :]
                    if leaf in EXPECTED_NEW_LEAVES:
                        present.add(leaf)
            if desired != present:
                bad_categories.append(
                    {
                        "sku": sku,
                        "missing": sorted(desired - present),
                        "extra": sorted(present - desired),
                    }
                )

            description_mentions = series_mentions(content)
            extra_description = sorted(description_mentions - allowed)
            if extra_description:
                bad_descriptions.append(
                    {"sku": sku, "extra": extra_description}
                )

        if bad_b1:
            raise ValueError(f"Pozostało B.1: {bad_b1[:20]}")
        validation["no_b1_in_product_text"] = True
        if bad_glas:
            raise ValueError(f"Pozostało niejednoznaczne Glas: {bad_glas[:20]}")
        validation["no_ambiguous_bare_glas_in_product_text"] = True
        if bad_titles:
            raise ValueError(f"Tytuły niezgodne z serią: {bad_titles[:20]}")
        validation["titles_match_exact_series"] = True
        if bad_categories:
            raise ValueError(f"Kategorie niezgodne z serią: {bad_categories[:20]}")
        validation["categories_match_exact_series"] = True
        if bad_descriptions:
            raise ValueError(f"Opisy deklarują nadmiarowe serie: {bad_descriptions[:20]}")
        validation["descriptions_do_not_claim_extra_series"] = True

        result_duplicates = duplicate_groups(result_sheet, sku_col, title_col)
        new_duplicates = result_duplicates - source_duplicates
        if new_duplicates:
            raise ValueError(
                f"Powstały nowe duplikaty tytułów: {[sorted(item) for item in new_duplicates][:20]}"
            )
        validation["no_new_duplicate_title_groups"] = True
        if scan_mojibake(result):
            raise ValueError("Wykryto podejrzane uszkodzenie kodowania")
        validation["mojibake_scan_passed"] = True
    finally:
        source.close()
        result.close()

    return {
        "input": str(input_path.resolve()),
        "category_map": str(category_map_path.resolve()),
        "catalog": str(catalog_path.resolve()),
        "output": str(output_path.resolve()),
        "input_sha256": source_hash,
        "category_map_sha256": category_map_hash,
        "catalog_sha256": catalog_hash,
        "products": input_structure["sheets"][STORE_SHEET]["max_row"] - 1,
        "catalog_matches_existing": matched_existing,
        "catalog_matches_direct_code": matched_direct,
        "catalog_matches_total": matched_existing + matched_direct,
        "unresolved_count": len(unresolved),
        "unresolved": unresolved,
        "changes": changes,
        "change_counts": {key: len(value) for key, value in changes.items()},
        "catalog_diagnostics": catalog_diagnostics,
        "validation": validation,
    }


def add_table(sheet, start_row: int, end_row: int, end_col: int, name: str) -> None:
    if end_row <= start_row:
        return
    table = Table(
        displayName=name,
        ref=f"A{start_row}:{get_column_letter(end_col)}{end_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def style_title(sheet, title: str, subtitle: str, end_col: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(name="Aptos", size=10, color=TEXT)
    sheet["A2"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 36


def style_header(sheet, row_number: int, end_col: int) -> None:
    for cell in sheet[row_number][:end_col]:
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)
    sheet.row_dimensions[row_number].height = 30


def setup_sheet(sheet, freeze: str) -> None:
    sheet.freeze_panes = freeze
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.auto_filter.ref = None


def set_widths(sheet, widths: dict[int, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width


def build_report(report: dict[str, Any], output_path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Podsumowanie"
    style_title(
        summary,
        "Hager / Berker - korekta serii w całym imporcie",
        "Synchronizacja atrybutu Seria, kategorii, tytułów, opisów i tagów z kontrolną bazą katalogu PDF. B.1 nie jest serią; samodzielne Glas nie jest automatycznie Glasserie.",
        10,
    )
    cards = [
        ("Produkty", report["products"], GREEN),
        ("Zmiany serii", report["change_counts"]["series"], LIGHT_BLUE),
        ("Zmiany kategorii", report["change_counts"]["categories"], LIGHT_BLUE),
        ("Zmiany opisów", report["change_counts"]["content"], LIGHT_BLUE),
        ("DO SPRAWDZENIA", report["unresolved_count"], YELLOW),
    ]
    for index, (label, value, color) in enumerate(cards):
        start_col = 1 + index * 2
        summary.merge_cells(start_row=4, start_column=start_col, end_row=4, end_column=start_col + 1)
        summary.merge_cells(start_row=5, start_column=start_col, end_row=6, end_column=start_col + 1)
        summary.cell(4, start_col, label)
        summary.cell(5, start_col, value)
        summary.cell(4, start_col).font = Font(name="Aptos", bold=True, color=TEXT)
        summary.cell(5, start_col).font = Font(name="Aptos Display", size=20, bold=True, color=NAVY)
        for row in range(4, 7):
            for col in range(start_col, start_col + 2):
                cell = summary.cell(row, col)
                cell.fill = PatternFill("solid", fgColor=color)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)

    summary["A8"] = "Zakres i wynik"
    summary["A8"].font = Font(name="Aptos Display", size=12, bold=True, color=WHITE)
    summary["A8"].fill = PatternFill("solid", fgColor=NAVY)
    summary.merge_cells("A8:J8")
    notes = [
        f"Dopasowano do bazy PDF {report['catalog_matches_total']} z {report['products']} produktów: {report['catalog_matches_existing']} przez alias/powiązanie i {report['catalog_matches_direct_code']} bezpośrednio po kodzie SKU.",
        "Usunięto stare B.1 i niejednoznaczne Glas z opisów oraz tagów. Prawdziwe Glasserie zachowano tam, gdzie wynika z atrybutu produktu.",
        "Kategorie serii są teraz dokładnie wyprowadzone z atrybutu Seria. Szersza nazwa kategorii grupuje serie, ale nie dodaje produktu do niepowiązanych grup.",
        "Pozycje bez bezpośredniego dowodu w bazie PDF pozostawiono z dotychczasowym atrybutem i oznaczono DO_SPRAWDZENIA.",
    ]
    for row_number, note in enumerate(notes, start=9):
        summary.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=10)
        cell = summary.cell(row_number, 1, f"• {note}")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.fill = PatternFill("solid", fgColor=PALE_BLUE if row_number % 2 else WHITE)
        summary.row_dimensions[row_number].height = 34

    summary["A15"] = "Walidacja"
    summary["A15"].font = Font(name="Aptos Display", size=12, bold=True, color=WHITE)
    summary["A15"].fill = PatternFill("solid", fgColor=NAVY)
    summary.merge_cells("A15:J15")
    summary.cell(16, 1, "Test")
    summary.cell(16, 2, "Wynik")
    style_header(summary, 16, 2)
    validation_labels = {
        "source_files_unchanged": "Pliki źródłowe niezmienione",
        "sheet_structure_preserved": "Struktura skoroszytu zachowana",
        "only_allowed_columns_changed": "Zmiany tylko w dozwolonych kolumnach",
        "styles_preserved": "Style komórek zachowane",
        "no_b1_in_product_text": "Brak B.1 w atrybutach, tytułach, opisach i tagach",
        "no_ambiguous_bare_glas_in_product_text": "Brak niejednoznacznego Glas w danych produktu",
        "titles_match_exact_series": "Tytuły zgodne z atrybutem Seria",
        "categories_match_exact_series": "Kategorie zgodne z atrybutem Seria",
        "descriptions_do_not_claim_extra_series": "Opisy nie deklarują nadmiarowych serii",
        "no_new_duplicate_title_groups": "Brak nowych duplikatów tytułów",
        "mojibake_scan_passed": "Kodowanie tekstu poprawne",
    }
    for row_number, (key, passed) in enumerate(report["validation"].items(), start=17):
        summary.cell(row_number, 1, validation_labels.get(key, key))
        summary.cell(row_number, 2, "OK" if passed else "BŁĄD")
        summary.cell(row_number, 2).fill = PatternFill("solid", fgColor=GREEN if passed else RED)
        summary.cell(row_number, 2).font = Font(bold=True, color=TEXT)
        summary.cell(row_number, 2).alignment = Alignment(horizontal="center")
    set_widths(summary, {1: 52, 2: 14, 3: 14, 4: 14, 5: 14, 6: 14, 7: 14, 8: 14, 9: 14, 10: 14})
    setup_sheet(summary, "A4")

    sheet_specs = [
        ("Zmiany serii", "series", ["Lp.", "SKU", "Wiersz", "Przed", "Po", "Źródło", "Pewność", "Status"]),
        ("Zmiany kategorii", "categories", ["Lp.", "SKU", "Wiersz", "Seria", "Przed", "Po", "Dodano", "Usunięto", "Status"]),
        ("Zmiany tytułów", "titles", ["Lp.", "SKU", "Wiersz", "Seria", "Tytuł przed", "Tytuł po", "Status"]),
        ("Zmiany opisów", "content", ["Lp.", "SKU", "Wiersz", "Seria", "Usunięte / znormalizowane", "Opis przed", "Opis po", "Status"]),
        ("Zmiany tagów", "tags", ["Lp.", "SKU", "Wiersz", "Seria", "Usunięte / znormalizowane", "Tagi przed", "Tagi po", "Status"]),
    ]
    for sheet_name, key, headers in sheet_specs:
        sheet = workbook.create_sheet(sheet_name)
        style_title(sheet, sheet_name, "Pełna lista zmian z numerem wiersza w pliku importowym.", len(headers))
        sheet.append([])
        sheet.append(headers)
        style_header(sheet, 4, len(headers))
        for index, change in enumerate(report["changes"][key], start=1):
            if key == "series":
                row = [index, change["sku"], change["row"], change["before"], change["after"], change["source"], change["confidence"], change["status"]]
            elif key == "categories":
                row = [index, change["sku"], change["row"], change["series"], change["before"], change["after"], " | ".join(change["added"]), " | ".join(change["removed"]), change["status"]]
            elif key == "titles":
                row = [index, change["sku"], change["row"], change["series"], change["before"], change["after"], change["status"]]
            else:
                row = [index, change["sku"], change["row"], change["series"], " | ".join(change["removed_or_normalized"]), change["before"], change["after"], change["status"]]
            sheet.append(row)
        add_table(sheet, 4, sheet.max_row, len(headers), f"Table{key.title()}")
        for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column >= 4)
        widths = {1: 7, 2: 22, 3: 12, 4: 38}
        for column in range(5, len(headers) + 1):
            widths[column] = 70 if key in {"content", "categories"} else 45
        set_widths(sheet, widths)
        setup_sheet(sheet, "A5")

    check = workbook.create_sheet("DO_SPRAWDZENIA")
    style_title(
        check,
        "Pozycje bez bezpośredniego dowodu PDF",
        "Nie zmieniano ich atrybutu Seria na podstawie zgadywania. Pozostałe pola uporządkowano względem aktualnego atrybutu.",
        8,
    )
    headers = ["Lp.", "SKU", "Wiersz", "Aktualna seria", "Tytuł", "Źródło", "Pewność", "Status"]
    check.append([])
    check.append(headers)
    style_header(check, 4, len(headers))
    for index, item in enumerate(report["unresolved"], start=1):
        check.append([index, item["sku"], item["row"], item["series"], item["title"], item["source"], item["confidence"], item["status"]])
    add_table(check, 4, check.max_row, len(headers), "TableUnresolved")
    set_widths(check, {1: 7, 2: 22, 3: 12, 4: 42, 5: 100, 6: 34, 7: 12, 8: 20})
    for row in check.iter_rows(min_row=5, max_row=check.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {4, 5, 6})
        row[7].fill = PatternFill("solid", fgColor=YELLOW)
    setup_sheet(check, "A5")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Synchronizuje serie Hager/Berker w całym imporcie i buduje raport zmian."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--category-map", type=Path, default=DEFAULT_CATEGORY_MAP)
    parser.add_argument("--catalog", type=Path, default=DEFAULT_CATALOG)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--json-report", type=Path, default=DEFAULT_JSON_REPORT)
    parser.add_argument("--xlsx-report", type=Path, default=DEFAULT_XLSX_REPORT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    report = build_corrected_workbook(
        args.input, args.category_map, args.catalog, args.output
    )
    args.json_report.parent.mkdir(parents=True, exist_ok=True)
    args.json_report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    build_report(report, args.xlsx_report)
    print(
        json.dumps(
            {
                "output": str(args.output.resolve()),
                "report": str(args.xlsx_report.resolve()),
                "products": report["products"],
                "catalog_matches_total": report["catalog_matches_total"],
                "unresolved_count": report["unresolved_count"],
                "change_counts": report["change_counts"],
                "validation": report["validation"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
