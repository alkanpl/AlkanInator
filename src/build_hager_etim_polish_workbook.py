from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
import shutil
import time
import urllib.parse
import urllib.request
import zipfile
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from copy import copy
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "input" / "BMEcat_Hager_PL_produkty.xlsx"
DEFAULT_OUTPUT = ROOT / "output" / "BMEcat_Hager_PL_produkty_ETIM_PL.xlsx"
DEFAULT_CACHE = ROOT / "output" / "etim_cache"
DEFAULT_REPORT = ROOT / "reports" / "BMEcat_Hager_PL_produkty_ETIM_PL_report.json"

ETIM_PL_CLASSES_URL = "https://etim.org.pl/lista-klas/"
ETIM_VIEWER_CLASS_URL = "https://viewer.etim-international.com/class/{code}?lang=pl-PL"
ETIM10_EN_ZIP_URL = (
    "https://www.etim-na.org/wp-content/uploads/sites/2/2024/12/"
    "ETIM-10.0-ALL-SECTORS-CSV-METRIC-EI-2024-12-05.zip"
)
ETIM10_EN_RELEASE_URL = (
    "https://www.etim-na.org/downloads/"
    "etim-10-0-all-sectors-csv-metric-ei-2024-12-05-zip/"
)
GOOGLE_TRANSLATE_URL = "https://translate.googleapis.com/translate_a/single"
GOOGLE_TRANSLATE_PUBLIC_URL = "https://translate.google.com/"
USER_AGENT = "AlkanInator-ETIM-Translator/1.0 (+local product data workflow)"

STANDARD_CODE_RE = re.compile(r"\b(?:EC|EF|EV|EU)\d{6}\b")
ANY_ETIM_CODE_RE = re.compile(r"\b(?:EC|EF|EV|EU)[A-Z0-9]{6}\b")
CLASS_CODE_RE = re.compile(r"^EC[A-Z0-9]{6}$")
FEATURE_CODE_RE = re.compile(r"^EF[A-Z0-9]{6}$")
VALUE_CODE_RE = re.compile(r"^EV[A-Z0-9]{6}$")
UNIT_CODE_RE = re.compile(r"^EU[A-Z0-9]{6}$")
PARAM_RE = re.compile(r"^(EF[A-Z0-9]{6})=(.*)$")
UNIT_SUFFIX_RE = re.compile(r"\s*\[(EU[A-Z0-9]{6})\]\s*$")

# Corrections for ambiguous, context-sensitive ETIM values where a generic
# machine translation produces an incorrect electrical-installation term.
VALUE_TRANSLATION_OVERRIDES = {
    "EV000038": "Osłona formowana",
    "EV000081": "Zaślepka",
    "EV000139": "Tworzywo sztuczne",
    "EV000154": "Inne",
    "EV000158": "Styk",
    "EV000197": "Widełki",
    "EV000200": "Klawisz kołyskowy",
    "EV000251": "Pokrętło/przycisk",
    "EV000283": "Ocynkowanie ciągłe",
    "EV000302": "Średniozwłoczny",
    "EV000316": "Klawisz jednoczęściowy",
    "EV000317": "Przycisk jednoczęściowy",
    "EV000598": "Suwak z 2 wejściami",
    "EV000623": "1-poziomowy",
    "EV000624": "2-poziomowy",
    "EV000635": "Do 1-poziomowych płyt nośnych",
    "EV000930": "Narożnik wewnętrzny",
    "EV001308": "Złącze szczelinowe IDC",
    "EV002133": "Osłona kabla",
    "EV004143": "Męski",
    "EV004144": "Żeński",
    "EV004173": "Magistrala (gniazdo)",
    "EV004378": "UAE/IAE (ISDN)",
    "EV005307": "aM",
    "EV005523": "Wtykowe",
    "EV006014": "Szafa/szafka",
    "EV006338": "Natynkowa obudowa wskaźnika obecności",
    "EV006357": "Wysoka osłona podświetlenia",
    "EV006375": "Klawisz trzyczęściowy",
    "EV006532": "Klawisz z prowadzeniem cięgna",
    "EV006540": "Klawisz dwuczęściowy",
    "EV007114": "Zatrzask",
    "EV007135": "Kształt puszki",
    "EV007248": "Zawieszany",
    "EV007326": "Do puszek osprzętowych",
    "EV007337": "Puszka osprzętowa",
    "EV007492": "Wypust",
    "EV007855": "Kształt wspornika",
    "EV007869": "Na sucho",
    "EV007965": "Łeb walcowy",
    "EV007981": "Przycisk dwuczęściowy",
    "EV008559": "Styk resetujący",
    "EV008602": "Pierścień dopasowujący",
    "EV008758": "Opadający",
    "EV008828": "Obrót",
    "EV009146": "Ogranicznik drzwi",
    "EV009413": "Blokada drzwi",
    "EV009478": "Mostek",
    "EV010943": "Przekładnik pomiarowy",
    "EV010944": "Moc czynna",
    "EV010946": "Moc czynna i moc bierna",
    "EV011642": "gPV",
    "EV011644": "gTr",
    "EV019070": "gG",
}


@dataclass
class MappingEntry:
    code: str
    kind: str
    label_pl: str
    label_en: str
    source: str
    status: str
    source_url: str
    occurrences: int = 0


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def strip_tags(value: str) -> str:
    return normalize_space(html.unescape(re.sub(r"<[^>]+>", " ", value or "")))


def http_get_bytes(url: str, timeout: int = 45, retries: int = 4) -> bytes:
    last_error: Exception | None = None
    for attempt in range(retries):
        try:
            request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            with urllib.request.urlopen(request, timeout=timeout) as response:
                return response.read()
        except Exception as exc:  # network errors need a retry with backoff
            last_error = exc
            if attempt + 1 < retries:
                time.sleep(0.8 * (attempt + 1))
    raise RuntimeError(f"Could not download {url}: {last_error}")


def http_get_text(url: str, timeout: int = 45, retries: int = 4) -> str:
    return http_get_bytes(url, timeout=timeout, retries=retries).decode("utf-8", errors="replace")


def save_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(path)


def load_json(path: Path, default: object) -> object:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ensure_english_release(cache_dir: Path) -> Path:
    zip_path = cache_dir / "ETIM-10.0-ALL-SECTORS-CSV-METRIC-EI-2024-12-05.zip"
    extract_dir = cache_dir / "etim10_en_csv"
    required = [
        "ETIMARTCLASS.csv",
        "ETIMFEATURE.csv",
        "ETIMVALUE.csv",
        "ETIMUNIT.csv",
    ]
    cache_dir.mkdir(parents=True, exist_ok=True)
    if not zip_path.exists():
        zip_path.write_bytes(http_get_bytes(ETIM10_EN_ZIP_URL, timeout=90))
    if not all((extract_dir / name).exists() for name in required):
        extract_dir.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(zip_path) as archive:
            archive.extractall(extract_dir)
    missing = [name for name in required if not (extract_dir / name).exists()]
    if missing:
        raise RuntimeError(f"English ETIM 10 release is incomplete: {missing}")
    return extract_dir


def load_two_column_csv(path: Path, id_column: str, label_column: str) -> dict[str, str]:
    output: dict[str, str] = {}
    # Official ETIM CSV exports are UTF-16LE, but some releases omit the BOM.
    with path.open("r", encoding="utf-16le", newline="") as handle:
        for row in csv.DictReader(handle, delimiter=";"):
            code = normalize_space(row.get(id_column, ""))
            label = normalize_space(row.get(label_column, ""))
            if code and label:
                output[code] = label
    return output


def load_english_dictionaries(release_dir: Path) -> dict[str, dict[str, str]]:
    return {
        "EC": load_two_column_csv(
            release_dir / "ETIMARTCLASS.csv", "ARTCLASSID", "ARTCLASSDESC"
        ),
        "EF": load_two_column_csv(
            release_dir / "ETIMFEATURE.csv", "FEATUREID", "FEATUREDESC"
        ),
        "EV": load_two_column_csv(
            release_dir / "ETIMVALUE.csv", "VALUEID", "VALUEDESC"
        ),
        "EU": load_two_column_csv(
            release_dir / "ETIMUNIT.csv", "UNITOFMEASID", "UNITDESC"
        ),
    }


def collect_used_codes(
    input_path: Path,
) -> tuple[dict[str, Counter[str]], dict[str, Counter[str]], int, int]:
    counters = {prefix: Counter() for prefix in ("EC", "EF", "EV", "EU")}
    value_feature_contexts: dict[str, Counter[str]] = defaultdict(Counter)
    product_rows = 0
    segment_count = 0
    workbook = load_workbook(input_path, read_only=True, data_only=False)
    worksheet = workbook["Produkty"]
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        product_rows += 1
        class_code = normalize_space(str(row[9] or ""))
        if class_code:
            counters["EC"][class_code] += 1
        params = str(row[10] or "")
        for raw_segment in params.split(";"):
            segment = raw_segment.strip()
            if not segment:
                continue
            segment_count += 1
            match = PARAM_RE.fullmatch(segment)
            if not match:
                continue
            feature_code, raw_value = match.groups()
            counters["EF"][feature_code] += 1
            unit_match = UNIT_SUFFIX_RE.search(raw_value)
            if unit_match:
                counters["EU"][unit_match.group(1)] += 1
                raw_value = UNIT_SUFFIX_RE.sub("", raw_value).strip()
            if VALUE_CODE_RE.fullmatch(raw_value):
                counters["EV"][raw_value] += 1
                value_feature_contexts[raw_value][feature_code] += 1
    workbook.close()
    return counters, value_feature_contexts, product_rows, segment_count


def parse_polish_class_list(page_html: str) -> dict[str, str]:
    rows = re.findall(
        r"<tr[^>]*>\s*<td[^>]*>\s*(EC\d{6})\s*</td>\s*<td[^>]*>(.*?)</td>",
        page_html,
        flags=re.IGNORECASE | re.DOTALL,
    )
    return {code.upper(): strip_tags(label) for code, label in rows if strip_tags(label)}


def parse_viewer_class_page(page_html: str, expected_code: str) -> tuple[str, dict[str, dict[str, str]]]:
    soup = BeautifulSoup(page_html, "html.parser")
    title_text = normalize_space(soup.title.get_text(" ", strip=True)) if soup.title else ""
    title_match = re.match(
        rf"{re.escape(expected_code)}\s*-\s*(.*)$", title_text, flags=re.IGNORECASE
    )
    class_label = normalize_space(title_match.group(1)) if title_match else ""
    features: dict[str, dict[str, str]] = {}
    for italic in soup.find_all("i"):
        descriptor = normalize_space(italic.get_text(" ", strip=True))
        descriptor_match = re.fullmatch(
            r"(EF[A-Z0-9]{6})\s*-\s*(.+)", descriptor, flags=re.IGNORECASE
        )
        if not descriptor_match:
            continue
        code, feature_type = descriptor_match.groups()
        label_node = italic.find_previous_sibling("p")
        label = normalize_space(label_node.get_text(" ", strip=True)) if label_node else ""
        if label:
            features[code.upper()] = {
                "label": label,
                "feature_type": normalize_space(feature_type),
            }
    return class_label, features


def fetch_viewer_data(
    class_codes: Iterable[str], cache_dir: Path, workers: int = 4
) -> dict[str, dict[str, object]]:
    cache_path = cache_dir / "etim_viewer_pl_features.json"
    cached = load_json(cache_path, {})
    if not isinstance(cached, dict):
        cached = {}

    pending = [
        code
        for code in sorted(set(class_codes))
        if code not in cached or cached[code].get("parser_version") != 2
    ]

    def fetch_one(code: str) -> tuple[str, dict[str, object]]:
        url = ETIM_VIEWER_CLASS_URL.format(code=code)
        page = http_get_text(url)
        class_label, features = parse_viewer_class_page(page, code)
        time.sleep(0.12)
        return code, {
            "class_label": class_label,
            "features": features,
            "source_url": url,
            "parser_version": 2,
            "fetched_at": datetime.now(timezone.utc).isoformat(),
        }

    if pending:
        with ThreadPoolExecutor(max_workers=max(1, workers)) as executor:
            futures = {executor.submit(fetch_one, code): code for code in pending}
            completed = 0
            for future in as_completed(futures):
                code = futures[future]
                try:
                    fetched_code, payload = future.result()
                    cached[fetched_code] = payload
                except Exception as exc:
                    cached[code] = {
                        "class_label": "",
                        "features": {},
                        "source_url": ETIM_VIEWER_CLASS_URL.format(code=code),
                        "parser_version": 2,
                        "error": str(exc),
                        "fetched_at": datetime.now(timezone.utc).isoformat(),
                    }
                completed += 1
                if completed % 25 == 0:
                    save_json(cache_path, cached)
        save_json(cache_path, cached)
    return cached


def google_translate_lines(lines: list[str]) -> list[str]:
    if not lines:
        return []
    query = "\n".join(lines)
    params = urllib.parse.urlencode(
        {"client": "gtx", "sl": "en", "tl": "pl", "dt": "t", "q": query}
    )
    response = json.loads(http_get_text(f"{GOOGLE_TRANSLATE_URL}?{params}"))
    translated = "".join(part[0] for part in response[0] if part and part[0])
    result = [normalize_space(item) for item in translated.splitlines()]
    if len(result) != len(lines):
        if len(lines) == 1:
            return [normalize_space(translated)]
        output: list[str] = []
        for line in lines:
            output.extend(google_translate_lines([line]))
            time.sleep(0.08)
        return output
    return result


def translate_missing_labels(labels: Iterable[str], cache_dir: Path) -> dict[str, str]:
    cache_path = cache_dir / "google_translate_en_pl.json"
    cache = load_json(cache_path, {})
    if not isinstance(cache, dict):
        cache = {}
    unique_labels = sorted({normalize_space(label) for label in labels if normalize_space(label)})
    pending = [label for label in unique_labels if label not in cache]

    batches: list[list[str]] = []
    current: list[str] = []
    current_chars = 0
    for label in pending:
        if current and (len(current) >= 35 or current_chars + len(label) + 1 > 1500):
            batches.append(current)
            current = []
            current_chars = 0
        current.append(label)
        current_chars += len(label) + 1
    if current:
        batches.append(current)

    for index, batch in enumerate(batches, start=1):
        translations = google_translate_lines(batch)
        if len(translations) != len(batch):
            raise RuntimeError("Translation batch returned a different number of labels")
        for source, translated in zip(batch, translations):
            cache[source] = translated or source
        if index % 5 == 0 or index == len(batches):
            save_json(cache_path, cache)
        time.sleep(0.2)
    return cache


def build_mapping(
    used: dict[str, Counter[str]],
    value_feature_contexts: dict[str, Counter[str]],
    cache_dir: Path,
    workers: int,
) -> dict[str, MappingEntry]:
    release_dir = ensure_english_release(cache_dir)
    english = load_english_dictionaries(release_dir)

    class_page = http_get_text(ETIM_PL_CLASSES_URL)
    polish_classes = parse_polish_class_list(class_page)
    viewer = fetch_viewer_data(used["EC"].keys(), cache_dir, workers=workers)

    polish_features: dict[str, str] = {}
    feature_sources: dict[str, str] = {}
    feature_types: dict[str, str] = {}
    viewer_classes: dict[str, str] = {}
    for class_code, payload in viewer.items():
        class_label = normalize_space(str(payload.get("class_label", "")))
        if class_label:
            viewer_classes[class_code] = class_label
        source_url = str(payload.get("source_url", ETIM_VIEWER_CLASS_URL.format(code=class_code)))
        features = payload.get("features", {})
        if not isinstance(features, dict):
            continue
        for code, feature_payload in features.items():
            if not isinstance(feature_payload, dict):
                continue
            label = normalize_space(str(feature_payload.get("label", "")))
            if label and code not in polish_features:
                polish_features[code] = label
                feature_sources[code] = source_url
                feature_types[code] = normalize_space(
                    str(feature_payload.get("feature_type", ""))
                )

    labels_to_translate: list[str] = []
    for code in used["EC"]:
        if not polish_classes.get(code) and not viewer_classes.get(code):
            labels_to_translate.append(english["EC"].get(code, ""))
    for code in used["EF"]:
        if not polish_features.get(code):
            labels_to_translate.append(english["EF"].get(code, ""))
    value_translation_queries: dict[str, str] = {}
    for code in used["EV"]:
        label_en = english["EV"].get(code, "")
        contexts = value_feature_contexts.get(code, Counter())
        feature_code = contexts.most_common(1)[0][0] if contexts else ""
        feature_en = english["EF"].get(feature_code, "")
        query = f"{feature_en}: {label_en}" if feature_en and label_en else label_en
        if query:
            value_translation_queries[code] = query
            labels_to_translate.append(query)

    translated = translate_missing_labels(labels_to_translate, cache_dir)
    mapping: dict[str, MappingEntry] = {}

    for prefix, kind in (
        ("EC", "Klasa"),
        ("EF", "Cecha"),
        ("EV", "Wartość"),
        ("EU", "Jednostka"),
    ):
        for code, occurrences in sorted(used[prefix].items()):
            label_en = english[prefix].get(code, "")
            if prefix == "EC":
                label_pl = polish_classes.get(code) or viewer_classes.get(code, "")
                if label_pl:
                    source = "ETIM Polska / ETIM Viewer (PL)"
                    source_url = ETIM_PL_CLASSES_URL if code in polish_classes else ETIM_VIEWER_CLASS_URL.format(code=code)
                    status = "OK"
                elif label_en and translated.get(label_en):
                    label_pl = translated[label_en]
                    source = "ETIM 10 EN + tłumaczenie maszynowe"
                    source_url = ETIM10_EN_RELEASE_URL
                    status = "WARNING"
                else:
                    label_pl = code
                    source = "Nierozpoznany kod"
                    source_url = ""
                    status = "WARNING"
            elif prefix == "EF":
                label_pl = polish_features.get(code, "")
                if label_pl:
                    source = "ETIM Viewer (PL)"
                    source_url = feature_sources.get(code, ETIM_VIEWER_CLASS_URL.format(code=next(iter(used["EC"]))))
                    status = "OK"
                elif label_en and translated.get(label_en):
                    label_pl = translated[label_en]
                    source = "ETIM 10 EN + tłumaczenie maszynowe"
                    source_url = ETIM10_EN_RELEASE_URL
                    status = "WARNING"
                else:
                    label_pl = code
                    source = "Nierozpoznany kod"
                    source_url = ""
                    status = "WARNING"
            elif prefix == "EV":
                translation_query = value_translation_queries.get(code, label_en)
                translated_phrase = translated.get(translation_query, "")
                if ":" in translated_phrase:
                    translated_value = normalize_space(translated_phrase.split(":", 1)[1])
                else:
                    translated_value = normalize_space(translated_phrase)
                if code in VALUE_TRANSLATION_OVERRIDES:
                    translated_value = VALUE_TRANSLATION_OVERRIDES[code]
                elif re.fullmatch(r"[A-Z0-9][A-Z0-9+./_-]{1,15}", label_en or ""):
                    translated_value = label_en
                elif translated_value:
                    translated_value = translated_value[0].upper() + translated_value[1:]
                if label_en and translated_value:
                    label_pl = translated_value
                    source = "ETIM 10 EN + tłumaczenie maszynowe"
                    source_url = ETIM10_EN_RELEASE_URL
                    status = "WARNING"
                else:
                    label_pl = code
                    source = "Nierozpoznany kod"
                    source_url = ""
                    status = "WARNING"
            else:
                if label_en:
                    label_pl = label_en
                    source = "ETIM 10 EN (symbol jednostki)"
                    source_url = ETIM10_EN_RELEASE_URL
                    status = "OK"
                else:
                    label_pl = code
                    source = "Nierozpoznany kod"
                    source_url = ""
                    status = "WARNING"

            mapping[code] = MappingEntry(
                code=code,
                kind=kind,
                label_pl=normalize_space(label_pl) or code,
                label_en=label_en,
                source=source,
                status=status,
                source_url=source_url,
                occurrences=occurrences,
            )

    mapping_cache = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "etim_version": "10.0",
        "entries": {code: asdict(entry) for code, entry in sorted(mapping.items())},
    }
    save_json(cache_dir / "etim10_hager_pl_mapping.json", mapping_cache)
    save_json(cache_dir / "etim10_hager_pl_feature_types.json", feature_types)
    return mapping


def strip_duplicate_unit(feature_label: str, unit_label: str) -> str:
    if not feature_label or not unit_label:
        return feature_label
    suffix = f"({unit_label})"
    if feature_label.rstrip().casefold().endswith(suffix.casefold()):
        return feature_label.rstrip()[: -len(suffix)].rstrip()
    return feature_label


def translate_parameter_string(
    raw_params: str, mapping: dict[str, MappingEntry]
) -> tuple[str, int, list[str]]:
    translated_segments: list[str] = []
    unresolved: list[str] = []
    source_segments = [segment.strip() for segment in str(raw_params or "").split(";") if segment.strip()]
    for segment in source_segments:
        match = PARAM_RE.fullmatch(segment)
        if not match:
            translated_segments.append(segment)
            unresolved.append(segment)
            continue
        feature_code, raw_value = match.groups()
        feature_entry = mapping.get(feature_code)
        feature_label = feature_entry.label_pl if feature_entry else feature_code
        if not feature_entry or feature_entry.label_pl == feature_code:
            unresolved.append(feature_code)

        unit_match = UNIT_SUFFIX_RE.search(raw_value)
        unit_code = unit_match.group(1) if unit_match else ""
        value = UNIT_SUFFIX_RE.sub("", raw_value).strip() if unit_match else raw_value.strip()
        unit_label = ""
        if unit_code:
            unit_entry = mapping.get(unit_code)
            unit_label = unit_entry.label_pl if unit_entry else unit_code
            if not unit_entry or unit_entry.label_pl == unit_code:
                unresolved.append(unit_code)
            feature_label = strip_duplicate_unit(feature_label, unit_label)

        if value == "true":
            translated_value = "Tak"
        elif value == "false":
            translated_value = "Nie"
        elif VALUE_CODE_RE.fullmatch(value):
            value_entry = mapping.get(value)
            translated_value = value_entry.label_pl if value_entry else value
            if not value_entry or value_entry.label_pl == value:
                unresolved.append(value)
        elif "|" in value:
            translated_value = re.sub(r"\s*\|\s*", "–", value)
        else:
            translated_value = value

        if unit_label:
            translated_value = f"{translated_value} {unit_label}"
        translated_segments.append(f"{feature_label} = {translated_value}")
    return "; ".join(translated_segments), len(source_segments), unresolved


def hash_cells(worksheet, min_col: int, max_col: int) -> str:
    digest = hashlib.sha256()
    for row in worksheet.iter_rows(min_row=1, min_col=min_col, max_col=max_col):
        # Empty inline-string cells can be normalized by openpyxl from data_type
        # "str" to "n" even though their value remains None. The user-visible
        # content is unchanged, so compare coordinates and exact values only.
        payload = [(cell.coordinate, cell.value) for cell in row]
        digest.update(json.dumps(payload, ensure_ascii=False, default=str).encode("utf-8"))
    return digest.hexdigest()


def formula_map(worksheet) -> dict[str, str]:
    return {
        cell.coordinate: str(cell.value)
        for row in worksheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    }


def style_audit_sheet(worksheet, widths: dict[str, float]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    header = worksheet[1]
    for cell in header:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
        cell.border = Border(bottom=thin_gray)
    worksheet.row_dimensions[1].height = 24
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def add_dictionary_sheet(workbook, mapping: dict[str, MappingEntry]) -> None:
    if "Słownik ETIM" in workbook.sheetnames:
        del workbook["Słownik ETIM"]
    worksheet = workbook.create_sheet("Słownik ETIM")
    headers = [
        "Typ",
        "Kod",
        "Polska etykieta",
        "Angielska etykieta",
        "Źródło",
        "Status",
        "Liczba wystąpień",
        "URL źródłowy",
    ]
    worksheet.append(headers)
    order = {"Klasa": 0, "Cecha": 1, "Wartość": 2, "Jednostka": 3}
    for entry in sorted(mapping.values(), key=lambda item: (order[item.kind], item.code)):
        worksheet.append(
            [
                entry.kind,
                entry.code,
                entry.label_pl,
                entry.label_en,
                entry.source,
                entry.status,
                entry.occurrences,
                entry.source_url,
            ]
        )
    style_audit_sheet(
        worksheet,
        {"A": 14, "B": 14, "C": 48, "D": 48, "E": 38, "F": 14, "G": 18, "H": 74},
    )
    worksheet.column_dimensions["B"].number_format = "@"
    worksheet.column_dimensions["G"].number_format = "#,##0"
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    worksheet.conditional_formatting.add(
        f"A2:H{worksheet.max_row}",
        FormulaRule(formula=["$F2=\"WARNING\""], fill=warning_fill),
    )


def add_report_sheet(workbook, report: dict[str, object], unresolved_counts: Counter[str]) -> None:
    if "Raport ETIM" in workbook.sheetnames:
        del workbook["Raport ETIM"]
    worksheet = workbook.create_sheet("Raport ETIM")
    worksheet.merge_cells("A1:D1")
    worksheet["A1"] = "Raport tłumaczenia ETIM 10 — Hager"
    worksheet["A1"].font = Font(color="FFFFFF", bold=True, size=15)
    worksheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    worksheet["A1"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 30
    worksheet.sheet_view.showGridLines = False

    rows = [
        ("Wersja ETIM", "10.0"),
        ("Data wygenerowania UTC", report["generated_at"]),
        ("Plik źródłowy", report["input_file"]),
        ("SHA-256 źródła", report["input_sha256"]),
        ("Liczba produktów", report["product_rows"]),
        ("Liczba parametrów", report["source_segments"]),
        ("Przetworzone parametry", report["translated_segments"]),
        ("Formuły w źródle", report["formula_count_before"]),
        ("Najdłuższa komórka parametrów", report["max_parameter_cell_length"]),
        ("Nierozpoznane kody w produktach", sum(unresolved_counts.values())),
        ("Unikalne nierozpoznane kody", len(unresolved_counts)),
    ]
    worksheet["A3"] = "Metryka"
    worksheet["B3"] = "Wartość"
    for row in rows:
        worksheet.append(row)

    source_start = worksheet.max_row + 2
    worksheet.cell(source_start, 1, "Źródło")
    worksheet.cell(source_start, 2, "URL")
    sources = [
        ("Oficjalna lista klas ETIM 10 PL", ETIM_PL_CLASSES_URL),
        ("Oficjalny ETIM Viewer PL", "https://viewer.etim-international.com/"),
        ("Oficjalny angielski eksport ETIM 10", ETIM10_EN_RELEASE_URL),
        ("Fallback tłumaczeniowy", GOOGLE_TRANSLATE_PUBLIC_URL),
    ]
    for label, url in sources:
        worksheet.append((label, url))

    exceptions_start = worksheet.max_row + 2
    worksheet.cell(exceptions_start, 1, "Nierozpoznany kod / segment")
    worksheet.cell(exceptions_start, 2, "Liczba wystąpień")
    if unresolved_counts:
        for code, count in sorted(unresolved_counts.items()):
            worksheet.append((code, count))
    else:
        worksheet.append(("Brak", 0))

    for row_number in (3, source_start, exceptions_start):
        for cell in worksheet[row_number][:2]:
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.font = Font(bold=True, color="1F1F1F")
    worksheet.freeze_panes = "A4"
    worksheet.column_dimensions["A"].width = 42
    worksheet.column_dimensions["B"].width = 92
    worksheet.column_dimensions["C"].width = 18
    worksheet.column_dimensions["D"].width = 18


def build_workbook(
    input_path: Path,
    output_path: Path,
    report_path: Path,
    cache_dir: Path,
    workers: int,
) -> dict[str, object]:
    input_hash_before = sha256_file(input_path)
    used, value_feature_contexts, product_rows, source_segment_count = collect_used_codes(
        input_path
    )
    mapping = build_mapping(
        used, value_feature_contexts, cache_dir, workers=workers
    )

    workbook = load_workbook(input_path, data_only=False, keep_links=True)
    worksheet = workbook["Produkty"]
    hash_a_i_before = hash_cells(worksheet, 1, 9)
    formulas_before = formula_map(worksheet)
    column_dimensions_before = {
        key: (dimension.width, dimension.hidden, dimension.bestFit)
        for key, dimension in worksheet.column_dimensions.items()
    }

    worksheet["J1"] = "Klasa ETIM PL"
    worksheet["K1"] = "Parametry ETIM PL"
    worksheet.column_dimensions["J"].width = max(32, worksheet.column_dimensions["J"].width or 0)

    translated_segments = 0
    unresolved_counts: Counter[str] = Counter()
    max_parameter_cell_length = 0
    for row_number in range(2, worksheet.max_row + 1):
        class_code = normalize_space(str(worksheet.cell(row_number, 10).value or ""))
        if class_code:
            class_entry = mapping.get(class_code)
            if class_entry:
                worksheet.cell(row_number, 10).value = class_entry.label_pl
                if class_entry.label_pl == class_code:
                    unresolved_counts[class_code] += 1
            else:
                unresolved_counts[class_code] += 1

        raw_params = str(worksheet.cell(row_number, 11).value or "")
        translated, segment_count, unresolved = translate_parameter_string(raw_params, mapping)
        worksheet.cell(row_number, 11).value = translated
        translated_segments += segment_count
        unresolved_counts.update(unresolved)
        max_parameter_cell_length = max(max_parameter_cell_length, len(translated))

    if max_parameter_cell_length > 32767:
        raise RuntimeError(
            f"Translated parameter cell exceeds Excel limit: {max_parameter_cell_length}"
        )

    report: dict[str, object] = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "input_file": str(input_path.resolve()),
        "output_file": str(output_path.resolve()),
        "input_sha256": input_hash_before,
        "product_rows": product_rows,
        "source_segments": source_segment_count,
        "translated_segments": translated_segments,
        "formula_count_before": len(formulas_before),
        "max_parameter_cell_length": max_parameter_cell_length,
        "used_unique_codes": {prefix: len(counter) for prefix, counter in used.items()},
        "used_code_occurrences": {
            prefix: sum(counter.values()) for prefix, counter in used.items()
        },
        "mapping_status": dict(Counter(entry.status for entry in mapping.values())),
        "mapping_source": dict(Counter(entry.source for entry in mapping.values())),
        "unresolved_codes": dict(sorted(unresolved_counts.items())),
    }

    add_dictionary_sheet(workbook, mapping)
    add_report_sheet(workbook, report, unresolved_counts)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

    if sha256_file(input_path) != input_hash_before:
        raise RuntimeError("Source workbook changed during processing")

    output_workbook = load_workbook(output_path, data_only=False, keep_links=True)
    output_products = output_workbook["Produkty"]
    formulas_after = formula_map(output_products)
    output_report = output_workbook["Raport ETIM"]
    validation = {
        "sheet_names": output_workbook.sheetnames,
        "rows": output_products.max_row,
        "columns": output_products.max_column,
        "hash_a_i_unchanged": hash_cells(output_products, 1, 9) == hash_a_i_before,
        "formulas_unchanged": formulas_after == formulas_before,
        "formula_count_after": len(formulas_after),
        "column_dimensions_a_i_unchanged": all(
            (
                output_products.column_dimensions[key].width,
                output_products.column_dimensions[key].hidden,
                output_products.column_dimensions[key].bestFit,
            )
            == value
            for key, value in column_dimensions_before.items()
            if key <= "I"
        ),
        "report_sheet_present": output_report.max_row > 1,
    }
    residuals: Counter[str] = Counter()
    segment_count_after = 0
    formula_errors: list[str] = []
    for row_number in range(2, output_products.max_row + 1):
        translated_class = str(output_products.cell(row_number, 10).value or "")
        translated_params = str(output_products.cell(row_number, 11).value or "")
        residuals.update(ANY_ETIM_CODE_RE.findall(f"{translated_class} {translated_params}"))
        segment_count_after += len(
            [segment for segment in translated_params.split(";") if segment.strip()]
        )
        for column in range(1, 12):
            value = output_products.cell(row_number, column).value
            if isinstance(value, str) and re.search(
                r"#REF!|#DIV/0!|#VALUE!|#NAME\?|#N/A", value
            ):
                formula_errors.append(output_products.cell(row_number, column).coordinate)

    validation.update(
        {
            "segments_after": segment_count_after,
            "segments_unchanged": segment_count_after == source_segment_count,
            "residual_codes": dict(sorted(residuals.items())),
            "residual_codes_match_report": dict(sorted(residuals.items()))
            == dict(sorted(unresolved_counts.items())),
            "formula_error_cells": formula_errors,
            "max_parameter_cell_length_after": max(
                len(str(output_products.cell(row_number, 11).value or ""))
                for row_number in range(2, output_products.max_row + 1)
            ),
        }
    )
    output_workbook.close()

    required_checks = [
        validation["rows"] == product_rows + 1,
        validation["columns"] == 11,
        validation["hash_a_i_unchanged"],
        validation["formulas_unchanged"],
        validation["column_dimensions_a_i_unchanged"],
        validation["segments_unchanged"],
        not validation["formula_error_cells"],
        validation["max_parameter_cell_length_after"] <= 32767,
        sha256_file(input_path) == input_hash_before,
    ]
    validation["passed"] = all(required_checks)
    report["validation"] = validation
    report["output_sha256"] = sha256_file(output_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    save_json(report_path, report)
    if not validation["passed"]:
        raise RuntimeError(f"Validation failed: {json.dumps(validation, ensure_ascii=True)}")
    return report


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Translate Hager ETIM 10 codes to Polish in a copied XLSX workbook."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--cache-dir", type=Path, default=DEFAULT_CACHE)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--workers", type=int, default=4)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    report = build_workbook(
        input_path=args.input,
        output_path=args.output,
        report_path=args.report,
        cache_dir=args.cache_dir,
        workers=args.workers,
    )
    print(json.dumps(report, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
