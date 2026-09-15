from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


DEFAULT_CHANGE_DATA = Path(
    "reports/Hager-Berker-serie-kategorie-poprawione-2026-08-27.json"
)
DEFAULT_PRODUCTS = Path(
    "output/Hager-Berker-Gniazdka_SERIE_KATEGORIE_TYTULY_SKROTY_POPRAWIONE_2026-08-27.xlsx"
)
DEFAULT_TITLE_CHANGES = Path(
    "reports/Hager-Berker-zmiany-tytulow-serie-2026-08-27.json"
)
DEFAULT_OUTPUT = Path(
    "reports/Hager-Berker-RAPORT_ZMIAN_SERII_I_KATEGORII_2026-08-27.xlsx"
)

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
WHITE = "FFFFFF"
GRAY = "F2F2F2"
TEXT = "1F2937"
THIN_GRAY = Side(style="thin", color="D9E1F2")


def load_product_lookup(path: Path) -> dict[str, dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Worksheet"]
        headers = {
            str(cell.value): cell.column
            for cell in sheet[1]
            if cell.value not in (None, "")
        }
        required = {
            "SKU",
            "Title",
            "Atrybut Produktu: Seria",
            "Kategorie produktów",
        }
        if not required.issubset(headers):
            raise ValueError(f"Brak kolumn produktowych: {sorted(required - set(headers))}")
        result: dict[str, dict[str, str]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            sku = str(row[headers["SKU"] - 1] or "").strip()
            if not sku:
                continue
            result[sku] = {
                "title": str(row[headers["Title"] - 1] or "").strip(),
                "series": str(
                    row[headers["Atrybut Produktu: Seria"] - 1] or ""
                ).strip(),
                "categories": str(row[headers["Kategorie produktów"] - 1] or "").strip(),
            }
        return result
    finally:
        workbook.close()


def add_table(sheet, start_row: int, end_row: int, end_col: int, name: str) -> None:
    if end_row <= start_row:
        return
    ref = f"A{start_row}:{get_column_letter(end_col)}{end_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def style_title(sheet, title: str, subtitle: str, end_col: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(name="Aptos", size=10, color=TEXT)
    sheet["A2"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 34


def style_header_row(sheet, row_number: int, end_col: int) -> None:
    for cell in sheet.iter_cols(
        min_col=1, max_col=end_col, min_row=row_number, max_row=row_number
    ):
        item = cell[0]
        item.font = Font(name="Aptos", bold=True, color=WHITE)
        item.fill = PatternFill("solid", fgColor=BLUE)
        item.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        item.border = Border(bottom=THIN_GRAY)
    sheet.row_dimensions[row_number].height = 30


def set_widths(sheet, widths: dict[int, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width


def common_sheet_setup(sheet, freeze: str) -> None:
    sheet.freeze_panes = freeze
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = None
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5


def build_summary_sheet(
    workbook: Workbook,
    report: dict[str, Any],
    title_report: dict[str, Any],
    product_lookup: dict[str, dict[str, str]],
) -> None:
    sheet = workbook.active
    sheet.title = "Podsumowanie"
    style_title(
        sheet,
        "Hager / Berker – raport zmian serii i kategorii",
        "Porównanie pliku wejściowego z wersją poprawioną. Seria odpowiada dokładnej grupie produktu przy SKU; kategoria może obejmować kilka serii zgodnie z BerkerHagerKategorie.xlsx.",
        10,
    )

    kpis = [
        ("Produkty w pliku", report["products"], GREEN),
        ("Zmienione serie", report["changed_series_count"], LIGHT_BLUE),
        ("Zmienione kategorie", report["changed_category_count"], LIGHT_BLUE),
        ("Zmienione tytuły", title_report["changed_title_count"], LIGHT_BLUE),
        (
            "Bez jednoznacznego dowodu serii",
            report["products_without_catalog_series_match"],
            YELLOW,
        ),
    ]
    for index, (label, value, color) in enumerate(kpis):
        start_col = 1 + index * 2
        sheet.merge_cells(
            start_row=4, start_column=start_col, end_row=4, end_column=start_col + 1
        )
        sheet.merge_cells(
            start_row=5, start_column=start_col, end_row=6, end_column=start_col + 1
        )
        label_cell = sheet.cell(4, start_col)
        value_cell = sheet.cell(5, start_col)
        label_cell.value = label
        label_cell.font = Font(name="Aptos", bold=True, color=TEXT)
        label_cell.fill = PatternFill("solid", fgColor=color)
        label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        value_cell.value = value
        value_cell.font = Font(name="Aptos Display", size=20, bold=True, color=NAVY)
        value_cell.fill = PatternFill("solid", fgColor=color)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in range(4, 7):
            for col in range(start_col, start_col + 2):
                sheet.cell(row, col).border = Border(
                    left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY
                )

    sheet["A8"] = "Zakres zmian"
    sheet["A8"].font = Font(name="Aptos Display", size=12, bold=True, color=WHITE)
    sheet["A8"].fill = PatternFill("solid", fgColor=NAVY)
    sheet.merge_cells("A8:J8")
    notes = [
        "Zmieniano pola Seria, Kategorie produktów i Title. Pozostałe kolumny nie zostały zmodyfikowane.",
        "Wartość serii pochodzi z dokładnej pozycji produktu/SKU w bazie katalogu PDF. Gdy pozycja produktu zawiera węższą listę niż nagłówek strony, pierwszeństwo ma pozycja produktu.",
        "Kategorie zostały ujednolicone do 10 aktualnych grup z pliku BerkerHagerKategorie.xlsx. Produkt może mieć więcej niż jedną kategorię.",
        "Produkty bez jednoznacznego dowodu katalogowego pozostawiono z dotychczasową serią zamiast przypisywać wartość na podstawie zgadywania.",
    ]
    for row_number, note in enumerate(notes, start=9):
        sheet.merge_cells(
            start_row=row_number, start_column=1, end_row=row_number, end_column=10
        )
        cell = sheet.cell(row_number, 1)
        cell.value = f"• {note}"
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = PatternFill("solid", fgColor=PALE_BLUE if row_number % 2 else WHITE)
        cell.font = Font(name="Aptos", size=10, color=TEXT)
        sheet.row_dimensions[row_number].height = 34

    sheet["A14"] = "Dodane kategorie"
    sheet["A14"].font = Font(name="Aptos Display", size=12, bold=True, color=WHITE)
    sheet["A14"].fill = PatternFill("solid", fgColor=NAVY)
    sheet.merge_cells("A14:D14")
    sheet["F14"] = "Walidacja"
    sheet["F14"].font = Font(name="Aptos Display", size=12, bold=True, color=WHITE)
    sheet["F14"].fill = PatternFill("solid", fgColor=NAVY)
    sheet.merge_cells("F14:J14")

    sheet.append([])
    category_header_row = 15
    sheet.cell(category_header_row, 1, "Kategoria")
    sheet.cell(category_header_row, 2, "Liczba dodań")
    style_header_row(sheet, category_header_row, 2)
    for path, count in report["added_category_counts"].items():
        sheet.append([path.split(">")[-1], count])
    category_end_row = sheet.max_row
    add_table(sheet, category_header_row, category_end_row, 2, "CategoryAdditions")

    validation_names = {
        "source_files_unchanged": "Pliki wejściowe niezmienione",
        "sheet_structure_preserved": "Struktura skoroszytu zachowana",
        "only_target_columns_changed": "Zmiany tylko w dozwolonych kolumnach",
        "target_cell_styles_preserved": "Style komórek zachowane",
        "formula_count_preserved": "Liczba formuł zachowana",
        "no_old_category_leaf_remains": "Brak starych kategorii serii",
        "all_series_implied_categories_present": "Kategorie wynikające z serii obecne",
        "mojibake_scan_passed": "Kodowanie polskich znaków poprawne",
    }
    sheet.cell(15, 6, "Test")
    sheet.cell(15, 7, "Wynik")
    style_header_row(sheet, 15, 7)
    for offset, (key, passed) in enumerate(report["validation"].items(), start=16):
        sheet.cell(offset, 6, validation_names.get(key, key))
        sheet.cell(offset, 7, "OK" if passed else "BŁĄD")
        sheet.cell(offset, 7).fill = PatternFill("solid", fgColor=GREEN if passed else RED)
        sheet.cell(offset, 7).font = Font(bold=True, color=TEXT)
        sheet.cell(offset, 7).alignment = Alignment(horizontal="center")

    title_validation_rows = [
        ("Zmiany tytułów tylko w kolumnie Title", title_report["validation"]["only_title_changed"]),
        ("Dokładne serie poprawnie reprezentowane w tytułach", title_report["validation"]["all_exact_series_present_in_titles"]),
        ("Brak nadmiarowych rozpoznanych serii w tytułach", title_report["validation"]["no_extra_recognized_series_in_titles"]),
        ("Układ: produkt, kolor, marka, seria, kod", title_report["validation"]["product_first_brand_series_code_layout"]),
        ("Każdy tytuł Lumina zawiera markę Hager", title_report["validation"]["all_lumina_titles_include_hager"]),
        ("Brak nowych grup zduplikowanych tytułów", title_report["validation"]["no_new_duplicate_title_groups"]),
    ]
    for offset, (label, passed) in enumerate(title_validation_rows, start=24):
        sheet.cell(offset, 6, label)
        sheet.cell(offset, 7, "OK" if passed else "BŁĄD")
        sheet.cell(offset, 7).fill = PatternFill("solid", fgColor=GREEN if passed else RED)
        sheet.cell(offset, 7).font = Font(bold=True, color=TEXT)
        sheet.cell(offset, 7).alignment = Alignment(horizontal="center")

    sheet["A31"] = "Przykład reguły"
    sheet["A31"].font = Font(name="Aptos Display", size=12, bold=True, color=WHITE)
    sheet["A31"].fill = PatternFill("solid", fgColor=NAVY)
    sheet.merge_cells("A31:J31")
    example = product_lookup.get("10107600/HAG", {})
    sheet.merge_cells("A32:J33")
    sheet["A32"] = (
        "10107600/HAG: seria R.1|R.3 pozostaje dokładną serią produktu, "
        "natomiast kategoria to Berker R.1/R.3/R.8.\n"
        f"Kategorie po zmianie: {example.get('categories', '')}"
    )
    sheet["A32"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet["A32"].fill = PatternFill("solid", fgColor=YELLOW)
    sheet["A32"].font = Font(name="Aptos", size=10, color=TEXT)

    set_widths(
        sheet,
        {1: 34, 2: 14, 3: 14, 4: 14, 5: 14, 6: 52, 7: 14, 8: 14, 9: 14, 10: 14},
    )
    common_sheet_setup(sheet, "A4")


def build_series_sheet(
    workbook: Workbook,
    report: dict[str, Any],
    product_lookup: dict[str, dict[str, str]],
) -> None:
    sheet = workbook.create_sheet("Zmiany serii")
    style_title(
        sheet,
        "Zmiany atrybutu Seria",
        "Każdy wiersz pokazuje wartość przed i po korekcie oraz źródło decyzji.",
        8,
    )
    headers = [
        "Lp.",
        "SKU",
        "Nazwa produktu",
        "Wiersz w imporcie",
        "Seria przed",
        "Seria po",
        "Źródło",
        "Kategorie po zmianie",
    ]
    sheet.append([])
    sheet.append(headers)
    header_row = 4
    style_header_row(sheet, header_row, len(headers))
    source_names = {
        "PDF_PRODUCT_ROW": "Pozycja produktu w bazie PDF",
        "PDF_VISUAL_OVERRIDE": "Kontrola wizualna katalogu PDF",
    }
    for index, change in enumerate(report["changed_series"], start=1):
        product = product_lookup.get(change["sku"], {})
        sheet.append(
            [
                index,
                change["sku"],
                product.get("title", ""),
                change["row"],
                change["before"],
                change["after"],
                source_names.get(change["source"], change["source"]),
                product.get("categories", ""),
            ]
        )
    add_table(sheet, header_row, sheet.max_row, len(headers), "SeriesChanges")
    set_widths(sheet, {1: 7, 2: 22, 3: 60, 4: 17, 5: 38, 6: 38, 7: 31, 8: 85})
    for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {3, 5, 6, 7, 8})
    common_sheet_setup(sheet, "A5")


def category_change_reason(change: dict[str, Any]) -> str:
    added = change.get("added") or []
    removed = change.get("removed") or []
    if added and removed:
        return "Dodano brakującą kategorię i usunięto błędną"
    if added:
        return "Dodano kategorię wynikającą z serii"
    if removed:
        return "Usunięto błędną kategorię"
    return "Zmieniono starą nazwę na aktualną kategorię"


def build_category_sheet(
    workbook: Workbook,
    report: dict[str, Any],
    product_lookup: dict[str, dict[str, str]],
) -> None:
    sheet = workbook.create_sheet("Zmiany kategorii")
    style_title(
        sheet,
        "Zmiany kategorii produktów",
        "Kategorie po zmianie odpowiadają aktualnym grupom z BerkerHagerKategorie.xlsx. Wartości są pokazane jako pełne ścieżki WooCommerce.",
        10,
    )
    headers = [
        "Lp.",
        "SKU",
        "Nazwa produktu",
        "Wiersz w imporcie",
        "Seria po",
        "Kategorie przed",
        "Kategorie po",
        "Dodano",
        "Usunięto",
        "Rodzaj korekty",
    ]
    sheet.append([])
    sheet.append(headers)
    header_row = 4
    style_header_row(sheet, header_row, len(headers))
    for index, change in enumerate(report["changed_categories"], start=1):
        product = product_lookup.get(change["sku"], {})
        sheet.append(
            [
                index,
                change["sku"],
                product.get("title", ""),
                change["row"],
                product.get("series", ""),
                change["before"],
                change["after"],
                "|".join(change.get("added") or []),
                "|".join(change.get("removed") or []),
                category_change_reason(change),
            ]
        )
    add_table(sheet, header_row, sheet.max_row, len(headers), "CategoryChanges")
    set_widths(
        sheet,
        {1: 7, 2: 22, 3: 58, 4: 17, 5: 36, 6: 72, 7: 72, 8: 55, 9: 55, 10: 38},
    )
    for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column not in {1, 2, 4})
    common_sheet_setup(sheet, "A5")


def build_title_sheet(
    workbook: Workbook,
    title_report: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet("Zmiany tytułów")
    style_title(
        sheet,
        "Zmiany tytułów produktów",
        "Standard: nazwa produktu i kolor → marka → seria → kod. Pełne grupy skracane są w tytule do Lumina, B.X, Q.X lub R.X; każda Lumina zawiera markę Hager. Atrybut Seria pozostaje dokładny.",
        6,
    )
    headers = [
        "Lp.",
        "SKU",
        "Wiersz w imporcie",
        "Dokładna seria",
        "Tytuł przed",
        "Tytuł po",
    ]
    sheet.append([])
    sheet.append(headers)
    header_row = 4
    style_header_row(sheet, header_row, len(headers))
    for index, change in enumerate(title_report["changes"], start=1):
        sheet.append(
            [
                index,
                change["sku"],
                change["row"],
                change["series"],
                change["before"],
                change["after"],
            ]
        )
    add_table(sheet, header_row, sheet.max_row, len(headers), "TitleChanges")
    set_widths(sheet, {1: 7, 2: 22, 3: 17, 4: 45, 5: 95, 6: 105})
    for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {4, 5, 6})
    common_sheet_setup(sheet, "A5")


def build_unmatched_sheet(
    workbook: Workbook,
    report: dict[str, Any],
    product_lookup: dict[str, dict[str, str]],
) -> None:
    sheet = workbook.create_sheet("Bez dowodu serii")
    style_title(
        sheet,
        "Produkty bez jednoznacznego dowodu serii",
        "Serii tych produktów nie nadpisano automatycznie. Lista służy do ewentualnej późniejszej kontroli ręcznej lub uzupełnienia bazy katalogowej.",
        7,
    )
    headers = [
        "Lp.",
        "SKU",
        "Nazwa produktu",
        "Pozostawiona seria",
        "Kategorie po zmianie",
        "Status",
        "Zalecenie",
    ]
    sheet.append([])
    sheet.append(headers)
    header_row = 4
    style_header_row(sheet, header_row, len(headers))
    for index, sku in enumerate(report["unmatched_skus"], start=1):
        product = product_lookup.get(sku, {})
        sheet.append(
            [
                index,
                sku,
                product.get("title", ""),
                product.get("series", ""),
                product.get("categories", ""),
                "POZOSTAWIONO BEZ ZMIANY",
                "Kontrola ręczna tylko po uzyskaniu jednoznacznego dowodu przy SKU",
            ]
        )
    add_table(sheet, header_row, sheet.max_row, len(headers), "UnmatchedSeries")
    set_widths(sheet, {1: 7, 2: 22, 3: 62, 4: 40, 5: 85, 6: 25, 7: 55})
    for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row):
        row[5].fill = PatternFill("solid", fgColor=YELLOW)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column not in {1, 2})
    common_sheet_setup(sheet, "A5")


def build_report(
    change_data: Path,
    title_changes_path: Path,
    products_path: Path,
    output_path: Path,
) -> dict[str, Any]:
    report = json.loads(change_data.read_text(encoding="utf-8"))
    title_report = json.loads(title_changes_path.read_text(encoding="utf-8"))
    product_lookup = load_product_lookup(products_path)
    if len(product_lookup) != report["products"]:
        raise ValueError(
            f"Niezgodna liczba produktów: {len(product_lookup)} vs {report['products']}"
        )

    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    build_summary_sheet(workbook, report, title_report, product_lookup)
    build_series_sheet(workbook, report, product_lookup)
    build_category_sheet(workbook, report, product_lookup)
    build_title_sheet(workbook, title_report)
    build_unmatched_sheet(workbook, report, product_lookup)

    for sheet in workbook.worksheets:
        sheet.sheet_view.zoomScale = 90
        sheet.sheet_view.zoomScaleNormal = 90

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

    check = load_workbook(output_path, read_only=False, data_only=False)
    try:
        expected_sheets = [
            "Podsumowanie",
            "Zmiany serii",
            "Zmiany kategorii",
            "Zmiany tytułów",
            "Bez dowodu serii",
        ]
        if check.sheetnames != expected_sheets:
            raise ValueError(f"Nieoczekiwane arkusze raportu: {check.sheetnames}")
        if check["Zmiany serii"].max_row - 4 != report["changed_series_count"]:
            raise ValueError("Niezgodna liczba zmian serii w raporcie")
        if check["Zmiany kategorii"].max_row - 4 != report["changed_category_count"]:
            raise ValueError("Niezgodna liczba zmian kategorii w raporcie")
        if check["Zmiany tytułów"].max_row - 4 != title_report["changed_title_count"]:
            raise ValueError("Niezgodna liczba zmian tytułów w raporcie")
        if check["Bez dowodu serii"].max_row - 4 != report["products_without_catalog_series_match"]:
            raise ValueError("Niezgodna liczba produktów bez dowodu serii")
        formula_errors: list[str] = []
        for sheet in check.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and any(
                        error in cell.value
                        for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
                    ):
                        formula_errors.append(f"{sheet.title}!{cell.coordinate}")
        if formula_errors:
            raise ValueError(f"Błędy formuł: {formula_errors[:20]}")
    finally:
        check.close()

    return {
        "output": str(output_path.resolve()),
        "sheets": 5,
        "series_changes": report["changed_series_count"],
        "category_changes": report["changed_category_count"],
        "title_changes": title_report["changed_title_count"],
        "without_series_evidence": report["products_without_catalog_series_match"],
        "validation": "OK",
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Buduje raport zmian serii i kategorii Hager/Berker.")
    parser.add_argument("--change-data", type=Path, default=DEFAULT_CHANGE_DATA)
    parser.add_argument("--title-changes", type=Path, default=DEFAULT_TITLE_CHANGES)
    parser.add_argument("--products", type=Path, default=DEFAULT_PRODUCTS)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    result = build_report(
        args.change_data, args.title_changes, args.products, args.output
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
