"""Add Hager rows without ``www=simple`` to the reviewed modular workbook."""

from __future__ import annotations

import argparse
import copy
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from build_hager_modular_workbook import (
    SOURCE_FILE,
    base_record,
    make_record,
    php_attribute_serialization,
)


DEFAULT_BASE = Path("output/Aparatura-Modulowa-Hager_nazwy_atrybuty.xlsx")
DEFAULT_SOURCE = Path("input/Hager aparatura na www.xlsx")
DEFAULT_OUTPUT = Path("output/Aparatura-Modulowa-Hager_nazwy_atrybuty_v2.xlsx")

PAGE_BY_CODE = {
    "KZN023": "1020",
    "KDN163B": "1018",
    "KDN363B": "1019",
    "SF463": "1175",
    "EPN520": "1180",
    "KZ059": "brak kodu w katalogu 2023",
    "SA480": "1174",
    "ADC960D": "1077",
    "ADC966D": "1077",
    "ADC970D": "1077",
    "ADC975D": "1077",
    "CDS425D": "1059",
    "ERC216": "1184",
    "ESC227": "1185",
    "ESC440": "1185",
    "ESC463": "1185",
    "NDN325": "1009",
    "NDN350": "1009",
    "SFT240": "1173",
    "SFT340": "1173",
    "SFT440": "1173",
    "SPB413": "1104",
    "SVN121": "1179",
    "SVN122": "1179",
    "SVN127": "1179",
    "SVN129": "1179",
    "SVN221": "1179",
    "SVN311": "1178",
    "SVN452": "1178",
    "ADA910D": "1078",
    "SA463": "1174",
    "LT052": "856",
    "HAE316": "969",
    "HXA004H": "699",
    "LSN503": "882",
    "LR703": "883",
    "SN216": "1193",
    "ACA916D": "1078",
    "ADA906D": "1078",
    "ADH916": "1080",
    "ADM416C": "1074",
    "SFB116": "1173",
    "HAE312": "969",
    "NDN340": "1009",
    "SFT225": "1173",
    "LSN501": "882",
    "ADA966D": "1078",
    "ESC325": "1185",
    "HAB306": "968",
    "KDN463A": "1019",
    "SU213": "1179",
    "EPN510": "1180",
    "KDN180A": "1018",
    "SFT140": "1173",
    "ESC080": "1172",
    "NDN316": "1009",
    "NDN332": "1009",
}

RCBO_DATA = {
    "ADC960D": ("C", 10, 30, "AC", "1P+N", 2, 2),
    "ADC966D": ("C", 16, 30, "AC", "1P+N", 2, 2),
    "ADC970D": ("C", 20, 30, "AC", "1P+N", 2, 2),
    "ADC975D": ("C", 25, 30, "AC", "1P+N", 2, 2),
    "ADA910D": ("B", 10, 30, "A", "1P+N", 2, 2),
    "ACA916D": ("B", 16, 10, "A", "1P+N", 2, 2),
    "ADA906D": ("B", 6, 30, "A", "1P+N", 2, 2),
    "ADH916": ("B", 16, 30, "A-HI", "1P+N", 2, 2),
    "ADM416C": ("B", 16, 30, "A", "3P+N", 4, 4),
    "ADA966D": ("C", 16, 30, "A", "1P+N", 2, 2),
}

BUS_BAR_DATA = {
    "KDN163B": ("1P", 1, 10, 63, 57, 1010, "57 x MCB 1P"),
    "KDN363B": ("3P", 3, 10, 63, 57, 1010, "19 x MCB 3P"),
    "KDN463A": (
        "4P",
        4,
        10,
        63,
        12,
        210,
        "MCB 4P/3P+N, RCCB 3P+N i RCBO 3P+N/4P",
    ),
    "KDN180A": ("1P", 1, 16, 80, 12, 210, "12 x MCB 1P"),
}

CHANGEOVER_DATA = {
    "SFT140": ("1P", 1, 40, 230, "od góry"),
    "SFT225": ("2P", 2, 25, 230, "od góry"),
    "SFT240": ("2P", 2, 40, 230, "od góry"),
    "SFT340": ("3P", 3, 40, 400, "od góry"),
    "SFT440": ("4P", 4, 40, 400, "od góry"),
    "SFB116": ("1P", 1, 16, 230, "od dołu"),
}

CONTACTOR_DATA = {
    "ESC227": (25, "1NO+1NC", 1),
    "ESC325": (25, "3NO", 2),
    "ESC440": (40, "4NO", 3),
    "ESC463": (63, "4NO", 3),
}

INDICATOR_DATA = {
    "SVN121": ("zielona", 1),
    "SVN122": ("czerwona", 1),
    "SVN127": ("3 x czerwona", 3),
    "SVN129": ("czerwona + zielona + pomarańczowa", 3),
    "SVN221": ("3 x zielona", 3),
}


def code_from_sku(sku: object) -> str:
    return str(sku or "").strip().upper().split("/", 1)[0]


def normalize_text(*parts: object) -> str:
    return re.sub(r"\s+", " ", " ".join(str(x or "") for x in parts)).strip()


def add_common_modular(
    record: dict,
    *,
    series: str,
    kind: str,
    product_type: str,
    poles: int | None = None,
    modules: int | float | None = None,
    current: int | float | str | None = None,
    voltage: str | None = None,
    terminals: str | None = None,
) -> None:
    legacy = record["legacy"]
    legacy.update(
        {
            "Rodzaj": kind,
            "Seria": series,
            "Sposób Montażu": "Szyna DIN 35 mm",
            "Typ": product_type,
            "Typ Produktu": product_type,
        }
    )
    if poles is not None:
        legacy["Liczba Biegunów"] = poles
    if modules is not None:
        legacy["Liczba Modułów"] = modules
    if current is not None:
        legacy["Prąd Znamionowy [A]"] = current
    if voltage:
        legacy["Napięcie [V]"] = voltage
    if terminals:
        legacy["Typ Zacisków"] = terminals


def make_addition_record(code: str, raw_title: str, source_id: object, www: object, price: object) -> dict:
    if code.startswith("NDN"):
        r = make_record(code)
    else:
        r = base_record(code)
    r.update(
        page=PAGE_BY_CODE[code],
        source_id=source_id,
        source_www="" if www is None else str(www),
        source_price=float(price) if price not in (None, "") else None,
        added=True,
        raw_source_title=raw_title,
    )
    legacy = r["legacy"]
    tech = r["technical"]

    if code.startswith("NDN"):
        return r

    if code in RCBO_DATA:
        characteristic, current, residual, rcd_type, layout, poles, modules = RCBO_DATA[code]
        r["family"] = "Wyłącznik różnicowo-nadprądowy RCBO"
        r["title"] = (
            f"Wyłącznik różnicowo-nadprądowy RCBO {characteristic}{current} {residual}mA "
            f"typ {rcd_type} {layout} 6kA Hager {code}"
        )
        add_common_modular(
            r,
            series=code[:3],
            kind="Wyłącznik różnicowo-nadprądowy",
            product_type="RCBO",
            poles=poles,
            modules=modules,
            current=current,
            voltage="230 V AC" if poles == 2 else "230/400 V AC",
            terminals="Klatkowe",
        )
        legacy.update(
            {
                "Charakterystyka Wyzwalania": characteristic,
                "Stopień Ochrony [IP]": "IP20",
                "Typ Wyłącznika": "Różnicowo-nadprądowy",
                "Zakres Temperatury Pracy": "-25...+40 °C",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": layout,
                "Atrybut Produktu: Znamionowa Zdolność Zwarciowa Icn [kA]": 6,
                "Atrybut Produktu: Prąd Różnicowy IΔn [mA]": residual,
                "Atrybut Produktu: Typ Prądu Różnicowego": rcd_type,
            }
        )
        if code == "ADH916":
            legacy["Dodatkowe Funkcje"] = "Krótkozwłoczny, podwyższona odporność"
        return r

    if code == "CDS425D":
        r.update(
            family="Wyłącznik różnicowoprądowy RCCB",
            title=(
                "Wyłącznik różnicowoprądowy RCCB 25A 30mA typ A 3P+N "
                "QuickConnect 6kA Hager CDS425D"
            ),
        )
        add_common_modular(
            r,
            series="CDS",
            kind="Wyłącznik różnicowoprądowy",
            product_type="RCCB",
            poles=4,
            modules=4,
            current=25,
            voltage="400 V AC",
            terminals="QuickConnect",
        )
        legacy.update(
            {
                "Stopień Ochrony [IP]": "IP20",
                "Typ Wyłącznika": "Różnicowoprądowy",
                "Zakres Temperatury Pracy": "-25...+40 °C",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": "3P+N",
                "Atrybut Produktu: Warunkowy Prąd Zwarciowy Inc [kA]": 6,
                "Atrybut Produktu: Prąd Różnicowy IΔn [mA]": 30,
                "Atrybut Produktu: Typ Prądu Różnicowego": "A",
            }
        )
        return r

    if code in BUS_BAR_DATA:
        layout, poles, cross_section, current, modules, length, compatibility = BUS_BAR_DATA[code]
        r.update(
            family="Szyna grzebieniowa",
            title=(
                f"Szyna grzebieniowa widełkowa {layout} {cross_section}mm² {current}A "
                f"{modules} modułów Hager {code}"
            ),
        )
        legacy.update(
            {
                "Długość": f"{length} mm",
                "Liczba Biegunów": poles,
                "Liczba Modułów": modules,
                "Napięcie [V]": "415 V AC",
                "Prąd Znamionowy [A]": current,
                "Rodzaj": "Szyna grzebieniowa",
                "Seria": "KDN",
                "Sposób Montażu": "W zaciskach aparatów modułowych",
                "Typ": "Szyna fazowa grzebieniowa",
                "Typ Produktu": "Szyna grzebieniowa",
                "Typ Zacisków": "Widełkowe",
                "Zastosowanie": compatibility,
                "Znamionowy Przekrój Poprzeczny": f"{cross_section} mm²",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": layout,
                "Atrybut Produktu: Typ Szyny": "Grzebieniowa, widełkowa",
                "Atrybut Produktu: Przekrój Szyny [mm²]": cross_section,
                "Atrybut Produktu: Kompatybilność": compatibility,
            }
        )
        return r

    if code == "KZN023":
        r.update(
            family="Akcesoria do szyn grzebieniowych",
            title=(
                "Osłony końcowe do szyn grzebieniowych 3P 10mm² / 2P 16mm² "
                "Hager KZN023"
            ),
        )
        legacy.update(
            {
                "Rodzaj": "Akcesoria do szyn grzebieniowych",
                "Rodzaj Akcesoria": "Osłony końcowe",
                "Seria": "KZN",
                "Typ": "Osłony końcowe",
                "Typ Produktu": "Akcesoria do szyn grzebieniowych",
                "Zastosowanie": "Szyny 3P 10 mm² oraz 2P 16 mm²",
                "Znamionowy Przekrój Poprzeczny": "10 / 16 mm²",
            }
        )
        tech["Atrybut Produktu: Kompatybilność"] = "Szyny 3P 10 mm² oraz 2P 16 mm²"
        return r

    if code == "KZ059":
        r.update(
            family="Akcesoria do szyn grzebieniowych",
            title="Osłona ochronna przed dotykiem do szyn grzebieniowych 5M Hager KZ059",
            confidence="średnia",
            status="DO_SPRAWDZENIA",
            notes=(
                "Kod nie został odnaleziony w polskim katalogu Hager 2023. "
                "Nazwa i parametry pochodzą z pliku Hager aparatura na www.xlsx."
            ),
            url="https://hager.com/pl/szkolenia-i-wsparcie/do-pobrania",
        )
        legacy.update(
            {
                "Liczba Modułów": 5,
                "Rodzaj": "Akcesoria do szyn grzebieniowych",
                "Rodzaj Akcesoria": "Osłona ochronna przed dotykiem",
                "Seria": "KZ",
                "Typ": "Osłona ochronna",
                "Typ Produktu": "Akcesoria do szyn grzebieniowych",
                "Zastosowanie": "Szyny grzebieniowe kołkowe i widełkowe",
            }
        )
        tech["Atrybut Produktu: Kompatybilność"] = "Szyny grzebieniowe kołkowe i widełkowe"
        return r

    if code in CHANGEOVER_DATA:
        layout, poles, current, voltage, common = CHANGEOVER_DATA[code]
        r.update(
            family="Przełącznik instalacyjny modułowy",
            title=(
                f"Przełącznik instalacyjny modułowy I-0-II {layout} {current}A "
                f"{voltage}V AC punkt wspólny {common} Hager {code}"
            ),
        )
        add_common_modular(
            r,
            series=code[:3],
            kind="Przełącznik instalacyjny",
            product_type="Przełącznik I-0-II",
            poles=poles,
            modules=poles,
            current=current,
            voltage=f"{voltage} V AC",
            terminals="Śrubowe",
        )
        legacy["Zastosowanie"] = f"Przełączanie źródła lub obwodu; punkt wspólny {common}"
        tech["Atrybut Produktu: Układ Biegunów"] = layout
        return r

    if code == "SF463":
        r.update(
            family="Przełącznik zasilania modułowy",
            title="Przełącznik zasilania modułowy I-0-II 3P+N 63A 400V AC Hager SF463",
        )
        add_common_modular(
            r,
            series="SF",
            kind="Przełącznik zasilania",
            product_type="Przełącznik I-0-II",
            poles=4,
            modules=8,
            current=63,
            voltage="400 V AC",
            terminals="Śrubowe",
        )
        legacy["Zastosowanie"] = "Ręczne przełączanie źródła zasilania"
        tech["Atrybut Produktu: Układ Biegunów"] = "3P+N"
        return r

    if code in {"SA463", "SA480"}:
        current = 63 if code == "SA463" else 80
        r.update(
            family="Rozłącznik izolacyjny modułowy",
            title=(
                f"Rozłącznik izolacyjny modułowy z wyzwalaniem 4P {current}A "
                f"400V AC ze stykiem pomocniczym Hager {code}"
            ),
        )
        add_common_modular(
            r,
            series="SA",
            kind="Rozłącznik izolacyjny z wyzwalaniem",
            product_type="Rozłącznik izolacyjny modułowy",
            poles=4,
            modules=4.5,
            current=current,
            voltage="400 V AC",
            terminals="Klatkowe",
        )
        legacy["Dodatkowe Funkcje"] = "Możliwość wyzwalania; styk pomocniczy"
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": "4P",
                "Atrybut Produktu: Konfiguracja Styków": "Styk pomocniczy",
            }
        )
        return r

    if code in {"HAE312", "HAE316", "HAB306"}:
        current = {"HAE312": 125, "HAE316": 160, "HAB306": 63}[code]
        size = 4 if code.startswith("HAE") else 1
        visible_break = code.startswith("HAE")
        visible_text = " z widoczną przerwą" if visible_break else ""
        r.update(
            family="Rozłącznik izolacyjny obrotowy",
            title=(
                f"Rozłącznik izolacyjny obrotowy 3P {current}A{visible_text} "
                f"rozmiar {size} Hager {code}"
            ),
        )
        legacy.update(
            {
                "Liczba Biegunów": 3,
                "Prąd Znamionowy [A]": current,
                "Rodzaj": "Rozłącznik izolacyjny obrotowy",
                "Rozmiar": size,
                "Seria": code[:3],
                "Typ": "Rozłącznik izolacyjny obrotowy",
                "Typ Produktu": "Rozłącznik izolacyjny",
                "Zastosowanie": "Rozłączanie i izolowanie obwodów",
            }
        )
        if visible_break:
            legacy["Dodatkowe Funkcje"] = "Widoczna przerwa izolacyjna"
        tech["Atrybut Produktu: Układ Biegunów"] = "3P"
        return r

    if code == "LT052":
        r.update(
            family="Rozłącznik bezpiecznikowy",
            title=(
                "Rozłącznik bezpiecznikowy NH00 3P 160A 690V AC "
                "na płytę lub szynę DIN Hager LT052"
            ),
        )
        legacy.update(
            {
                "Liczba Biegunów": 3,
                "Napięcie [V]": "690 V AC",
                "Prąd Znamionowy [A]": 160,
                "Rodzaj": "Rozłącznik bezpiecznikowy",
                "Rozmiar": "NH00",
                "Seria": "LT",
                "Sposób Montażu": "Płyta montażowa / szyna DIN 35 mm",
                "Typ": "NH00",
                "Typ Produktu": "Rozłącznik bezpiecznikowy",
                "Typ Zacisków": "Śruba M8",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": "3P",
                "Atrybut Produktu: Rodzaj Wkładki": "NH00",
            }
        )
        return r

    if code in {"LSN501", "LSN503", "LR703"}:
        if code.startswith("LSN"):
            poles = 1 if code == "LSN501" else 3
            fuse_size, current = "10x38 mm", 32
        else:
            poles, fuse_size, current = 3, "22x58 mm", 125
        r.update(
            family="Podstawa bezpiecznikowa modułowa",
            title=(
                f"Podstawa bezpiecznikowa modułowa {poles}P {fuse_size} "
                f"{current}A 690V AC Hager {code}"
            ),
        )
        add_common_modular(
            r,
            series=code[:3],
            kind="Podstawa bezpiecznikowa",
            product_type="Podstawa bezpiecznikowa modułowa",
            poles=poles,
            modules=poles,
            current=current,
            voltage="690 V AC",
            terminals="Klatkowe",
        )
        legacy["Rozmiar"] = fuse_size
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": f"{poles}P",
                "Atrybut Produktu: Rodzaj Wkładki": fuse_size,
            }
        )
        return r

    if code in CONTACTOR_DATA:
        current, contacts, modules = CONTACTOR_DATA[code]
        poles = sum(int(x) if x else 0 for x in re.findall(r"(\d+)(?:NO|NC)", contacts))
        r.update(
            family="Stycznik modułowy",
            title=(
                f"Stycznik modułowy {current}A 230V AC {contacts} "
                f"AC-7a/b Hager {code}"
            ),
        )
        add_common_modular(
            r,
            series="ESC",
            kind="Stycznik modułowy",
            product_type="Stycznik instalacyjny",
            poles=poles,
            modules=modules,
            current=current,
            voltage="230 V AC",
            terminals="Śrubowe",
        )
        legacy.update(
            {
                "Ilość Styków Zwiernych/rozwiernych": contacts,
                "Zakres Temperatury Pracy": "-10...+50 °C",
                "Zastosowanie": "Sterowanie odbiornikami w kategorii AC-7a/AC-7b",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": f"{poles}P",
                "Atrybut Produktu: Napięcie Cewki [V]": 230,
                "Atrybut Produktu: Konfiguracja Styków": contacts,
                "Atrybut Produktu: Kategoria Użytkowania": "AC-7a / AC-7b",
            }
        )
        return r

    if code in {"EPN510", "EPN520"}:
        contacts = "1NO" if code == "EPN510" else "2NO"
        r.update(
            family="Przekaźnik bistabilny modułowy",
            title=(
                f"Przekaźnik bistabilny modułowy 16A {contacts} "
                f"230V AC / 110V DC Hager {code}"
            ),
        )
        add_common_modular(
            r,
            series="EPN",
            kind="Przekaźnik bistabilny",
            product_type="Przekaźnik bistabilny modułowy",
            modules=1,
            current=16,
            voltage="230 V AC / 110 V DC",
            terminals="Śrubowe",
        )
        legacy["Ilość Styków Zwiernych/rozwiernych"] = contacts
        tech.update(
            {
                "Atrybut Produktu: Napięcie Cewki [V]": 230,
                "Atrybut Produktu: Konfiguracja Styków": contacts,
            }
        )
        return r

    if code == "ERC216":
        r.update(
            family="Przekaźnik instalacyjny modułowy",
            title="Przekaźnik instalacyjny modułowy 16A 2NO 230V AC Hager ERC216",
        )
        add_common_modular(
            r,
            series="ERC",
            kind="Przekaźnik instalacyjny",
            product_type="Przekaźnik instalacyjny modułowy",
            modules=1,
            current=16,
            voltage="230 V AC",
            terminals="Śrubowe",
        )
        legacy["Ilość Styków Zwiernych/rozwiernych"] = "2NO"
        tech.update(
            {
                "Atrybut Produktu: Napięcie Cewki [V]": 230,
                "Atrybut Produktu: Konfiguracja Styków": "2NO",
            }
        )
        return r

    if code == "ESC080":
        r.update(
            family="Styk pomocniczy",
            title="Styk pomocniczy 1NO+1NC 6A 250V AC do SBN i styczników Hager ESC080",
        )
        add_common_modular(
            r,
            series="ESC",
            kind="Styk pomocniczy",
            product_type="Akcesoria do aparatury modułowej",
            modules=0.5,
            current=6,
            voltage="250 V AC",
            terminals="Śrubowe",
        )
        legacy.update(
            {
                "Ilość Styków Zwiernych/rozwiernych": "1NO+1NC",
                "Rodzaj Akcesoria": "Styk pomocniczy",
                "Zastosowanie": "Rozłączniki SBN i styczniki modułowe",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Konfiguracja Styków": "1NO+1NC",
                "Atrybut Produktu: Kompatybilność": "Rozłączniki SBN i styczniki modułowe",
            }
        )
        return r

    if code == "SPB413":
        r.update(
            family="Ogranicznik przepięć SPD",
            title=(
                "Ogranicznik przepięć SPD T2 4P 3+1 TN-S/TT In 20kA "
                "Up≤1,35kV Uc 275V Hager SPB413"
            ),
        )
        add_common_modular(
            r,
            series="SPB",
            kind="Ogranicznik przepięć",
            product_type="SPD",
            poles=4,
            modules=4,
            voltage="230/400 V AC",
            terminals="Śrubowe",
        )
        legacy.update(
            {
                "Stopień Ochrony [IP]": "IP20",
                "Zastosowanie": "Ochrona instalacji TN-S/TT przed przepięciami",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": "3+1 / 4P",
                "Atrybut Produktu: Typ SPD": "T2",
                "Atrybut Produktu: Technologia SPD": "MOV",
                "Atrybut Produktu: Znamionowy Prąd Wyładowczy In [kA]": 20,
                "Atrybut Produktu: Poziom Ochrony Up [kV]": 1.35,
                "Atrybut Produktu: Maksymalne Napięcie Ciągłej Pracy Uc [V]": 275,
                "Atrybut Produktu: Typ Sieci": "TN-S / TT",
                "Atrybut Produktu: Styk Sygnalizacyjny": "Nie",
            }
        )
        return r

    if code in INDICATOR_DATA:
        color, lamp_count = INDICATOR_DATA[code]
        r.update(
            family="Lampka sygnalizacyjna modułowa",
            title=f"Lampka sygnalizacyjna LED {color} 230V AC Hager {code}",
        )
        add_common_modular(
            r,
            series="SVN",
            kind="Lampka sygnalizacyjna",
            product_type="Lampka sygnalizacyjna modułowa",
            modules=1,
            voltage="230 V AC",
            terminals="Śrubowe",
        )
        legacy["Barwa Lampek"] = color
        legacy["Ilość Lampek"] = lamp_count
        return r

    if code in {"SVN311", "SVN452"}:
        contacts = "1NO" if code == "SVN311" else "1NO+1NC"
        led = "" if code == "SVN311" else " LED czerwona"
        r.update(
            family="Przycisk sterowniczy modułowy",
            title=(
                f"Przycisk sterowniczy modułowy {contacts}{led} "
                f"16A 230V AC Hager {code}"
            ),
        )
        add_common_modular(
            r,
            series="SVN",
            kind="Przycisk sterowniczy",
            product_type="Przycisk sterowniczy modułowy",
            modules=1,
            current=16,
            voltage="230 V AC",
            terminals="Śrubowe",
        )
        legacy["Ilość Styków Zwiernych/rozwiernych"] = contacts
        if code == "SVN452":
            legacy["Barwa Lampek"] = "czerwona"
            legacy["Ilość Lampek"] = 1
        tech["Atrybut Produktu: Konfiguracja Styków"] = contacts
        return r

    if code == "SN216":
        r.update(
            family="Gniazdo modułowe na szynę DIN",
            title="Gniazdo modułowe na szynę DIN z uziemieniem 230V 10/16A Hager SN216",
        )
        add_common_modular(
            r,
            series="SN",
            kind="Gniazdo modułowe",
            product_type="Gniazdo na szynę DIN",
            modules=2.5,
            current="10/16",
            voltage="230 V AC",
            terminals="Śrubowe",
        )
        legacy["Uziemienie"] = "Tak"
        return r

    if code == "SU213":
        r.update(
            family="Dzwonek modułowy",
            title="Dzwonek modułowy 230V AC 0,03A 85dB Hager SU213",
        )
        add_common_modular(
            r,
            series="SU",
            kind="Dzwonek modułowy",
            product_type="Dzwonek modułowy",
            modules=1,
            current=0.03,
            voltage="230 V AC",
            terminals="Śrubowe",
        )
        legacy["Głośność [dB]"] = 85
        return r

    if code == "HXA004H":
        r.update(
            family="Wyzwalacz wzrostowy",
            title="Wyzwalacz wzrostowy 200-240V AC do x160-P630 Hager HXA004H",
        )
        legacy.update(
            {
                "Napięcie [V]": "200-240 V AC",
                "Rodzaj": "Wyzwalacz wzrostowy",
                "Rodzaj Akcesoria": "Wyzwalacz wzrostowy",
                "Seria": "HXA",
                "Typ": "Wyzwalacz wzrostowy",
                "Typ Produktu": "Akcesoria do rozłączników i wyłączników",
                "Zastosowanie": "x160, P160, x250, P250, x630 i P630",
            }
        )
        tech["Atrybut Produktu: Kompatybilność"] = "x160 / P160 / x250 / P250 / x630 / P630"
        return r

    raise ValueError(f"Brak reguły dla nowego kodu {code}: {raw_title}")


def source_rows(source_path: Path) -> list[dict]:
    wb = load_workbook(source_path, read_only=True, data_only=True)
    ws = wb["Arkusz1"]
    result = []
    for row in range(2, ws.max_row + 1):
        sku = str(ws.cell(row, 2).value or "").strip().upper()
        www = str(ws.cell(row, 5).value or "").strip()
        if not sku or www.lower() == "simple":
            continue
        result.append(
            {
                "source_row": row,
                "source_id": ws.cell(row, 1).value,
                "sku": sku,
                "www": www,
                "raw_title": normalize_text(ws.cell(row, 6).value, ws.cell(row, 7).value),
                "price": ws.cell(row, 8).value,
            }
        )
    return result


def copy_row_style(ws, source_row: int, target_row: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        target.font = copy.copy(source.font)
        target.fill = copy.copy(source.fill)
        target.border = copy.copy(source.border)
        target.alignment = copy.copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy.copy(source.protection)


def existing_old_titles(control_ws) -> dict[str, str]:
    headers = {str(control_ws.cell(5, c).value): c for c in range(1, control_ws.max_column + 1)}
    sku_col = headers.get("SKU")
    old_col = headers.get("Nazwa pierwotna")
    if not sku_col or not old_col:
        return {}
    return {
        str(control_ws.cell(row, sku_col).value or "").strip().upper(): str(
            control_ws.cell(row, old_col).value or ""
        )
        for row in range(6, control_ws.max_row + 1)
        if control_ws.cell(row, sku_col).value
    }


def key_parameters(record: dict) -> str:
    technical = record["technical"]
    pairs = [
        ("Icn", "Atrybut Produktu: Znamionowa Zdolność Zwarciowa Icn [kA]", "kA"),
        ("Icu", "Atrybut Produktu: Graniczna Zdolność Zwarciowa Icu [kA]", "kA"),
        ("Inc", "Atrybut Produktu: Warunkowy Prąd Zwarciowy Inc [kA]", "kA"),
        ("IΔn", "Atrybut Produktu: Prąd Różnicowy IΔn [mA]", "mA"),
        ("In", "Atrybut Produktu: Znamionowy Prąd Wyładowczy In [kA]", "kA"),
        ("Iimp", "Atrybut Produktu: Prąd Udarowy Iimp [kA, łącznie]", "kA"),
        ("Up", "Atrybut Produktu: Poziom Ochrony Up [kV]", "kV"),
    ]
    values = []
    for label, key, unit in pairs:
        value = technical.get(key)
        if value not in (None, ""):
            values.append(f"{label} {str(value).replace('.', ',')} {unit}")
    contacts = technical.get("Atrybut Produktu: Konfiguracja Styków")
    if contacts:
        values.append(str(contacts))
    return "; ".join(values)


def make_control_sheet(wb, records: list[tuple[int, str, str, dict]], added_count: int) -> None:
    if "Kontrola" in wb.sheetnames:
        del wb["Kontrola"]
    ws = wb.create_sheet("Kontrola")
    ws.sheet_view.showGridLines = False

    dark_blue = "17365D"
    light_blue = "DCE6F1"
    pale_blue = "EEF5FB"
    green = "E2F0D9"
    yellow = "FFF2CC"
    red = "F4CCCC"
    thin_gray = Side(style="thin", color="D9E1F2")

    ws.merge_cells("A1:U1")
    ws["A1"] = "Kontrola nazw i atrybutów - aparatura Hager"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=dark_blue)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26

    counts = Counter(record[3]["status"] for record in records)
    ws.merge_cells("A2:U2")
    ws["A2"] = (
        f"Produkty: {len(records)} | dodano: {added_count} | OK: {counts.get('OK', 0)} | "
        f"Do sprawdzenia: {counts.get('DO_SPRAWDZENIA', 0)} | Pliki źródłowe nie zostały nadpisane."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=light_blue)
    ws["A2"].font = Font(italic=True, color="17365D")

    ws.merge_cells("A3:U3")
    ws["A3"] = (
        "Źródła: Hager aparatura na www.xlsx; polski katalog Hager "
        f"„Rozdział energii 2023” ({SOURCE_FILE}) oraz karty produktów Hager PL."
    )
    ws["A3"].font = Font(size=9, color="666666")

    headers = [
        "Lp.",
        "SKU",
        "Kod",
        "ID źródłowe",
        "www źródłowe",
        "Cena internetowa netto [PLN]",
        "Nazwa pierwotna",
        "Nazwa po zmianie",
        "Rodzina",
        "Układ biegunów",
        "Prąd [A]",
        "Napięcie",
        "Moduły",
        "Charakterystyka / typ",
        "Kluczowe parametry",
        "Źródło",
        "Strona katalogu",
        "Adres źródła",
        "Pewność",
        "Status",
        "Uwagi",
    ]
    header_row = 5
    for col, value in enumerate(headers, 1):
        cell = ws.cell(header_row, col, value)
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EADBF"))
    ws.row_dimensions[header_row].height = 34

    for out_row, (_, sku, old_title, record) in enumerate(records, header_row + 1):
        legacy = record["legacy"]
        technical = record["technical"]
        characteristic = (
            legacy.get("Charakterystyka Wyzwalania")
            or technical.get("Atrybut Produktu: Typ Prądu Różnicowego")
            or technical.get("Atrybut Produktu: Typ SPD")
            or legacy.get("Typ", "")
        )
        source_name = "Hager aparatura na www.xlsx; katalog Hager 2023" if record.get("added") else "Hager, katalog Rozdział energii 2023; karta produktu Hager PL"
        row_values = [
            out_row - header_row,
            sku,
            record["code"],
            record.get("source_id", ""),
            record.get("source_www", ""),
            record.get("source_price", ""),
            old_title,
            record["title"],
            record["family"],
            technical.get("Atrybut Produktu: Układ Biegunów", ""),
            legacy.get("Prąd Znamionowy [A]", ""),
            legacy.get("Napięcie [V]", ""),
            legacy.get("Liczba Modułów", ""),
            characteristic,
            key_parameters(record),
            source_name,
            record["page"],
            record["url"],
            record["confidence"],
            record["status"],
            record["notes"],
        ]
        for col, value in enumerate(row_values, 1):
            cell = ws.cell(out_row, col, value)
            if out_row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=pale_blue)
            cell.border = Border(bottom=thin_gray)
            cell.alignment = Alignment(vertical="top", wrap_text=col in {7, 8, 15, 16, 21})
        ws.cell(out_row, 6).number_format = "#,##0.00"
        url_cell = ws.cell(out_row, 18)
        url_cell.hyperlink = record["url"]
        url_cell.style = "Hyperlink"

    last_row = header_row + len(records)
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:U{last_row}"
    widths = [
        7,
        19,
        14,
        14,
        14,
        18,
        43,
        49,
        32,
        15,
        11,
        20,
        10,
        20,
        31,
        35,
        22,
        38,
        12,
        19,
        55,
    ]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.conditional_formatting.add(
        f"T6:T{last_row}",
        FormulaRule(formula=['T6="OK"'], fill=PatternFill("solid", fgColor=green)),
    )
    ws.conditional_formatting.add(
        f"T6:T{last_row}",
        FormulaRule(formula=['T6="DO_SPRAWDZENIA"'], fill=PatternFill("solid", fgColor=red)),
    )
    ws.conditional_formatting.add(
        f"S6:S{last_row}",
        FormulaRule(formula=['S6="średnia"'], fill=PatternFill("solid", fgColor=yellow)),
    )


def validate(records: list[tuple[int, str, str, dict]], added_skus: list[str]) -> None:
    assert len(records) == 153, len(records)
    assert len(added_skus) == 57, len(added_skus)
    all_skus = [sku for _, sku, _, _ in records]
    titles = [record["title"] for _, _, _, record in records]
    assert len(set(all_skus)) == len(all_skus), "Powielone SKU"
    assert len(set(titles)) == len(titles), "Powielone nazwy"
    assert not (set(added_skus) & set(all_skus[:-57])), "Nowy SKU już istniał w pliku bazowym"
    for _, sku, _, record in records:
        assert record["code"] in record["title"], sku
        assert "Hager" in record["title"], sku
        assert record["page"], sku
        assert record["legacy"].get("Producent") == "Hager", sku
    statuses = Counter(record["status"] for _, _, _, record in records)
    assert statuses == Counter({"OK": 151, "DO_SPRAWDZENIA": 2}), statuses
    review_codes = sorted(record["code"] for _, _, _, record in records if record["status"] != "OK")
    assert review_codes == ["CDA440D", "KZ059"], review_codes


def build(base_path: Path, source_path: Path, output_path: Path) -> None:
    if output_path.resolve() in {base_path.resolve(), source_path.resolve()}:
        raise ValueError("Plik wynikowy nie może nadpisywać pliku źródłowego")
    additions = source_rows(source_path)
    assert len(additions) == 57, f"Oczekiwano 57 pozycji bez www=simple, jest {len(additions)}"

    wb = load_workbook(base_path)
    ws = wb["Worksheet"]
    headers = {str(cell.value): cell.column for cell in ws[1] if cell.value is not None}
    existing_skus = {
        str(ws.cell(row, headers["SKU"]).value or "").strip().upper()
        for row in range(2, ws.max_row + 1)
        if ws.cell(row, headers["SKU"]).value
    }
    added_skus = [item["sku"] for item in additions]
    duplicates = sorted(set(added_skus) & existing_skus)
    if duplicates:
        raise ValueError(f"Nowe pozycje zawierają istniejące SKU: {duplicates}")

    old_titles = existing_old_titles(wb["Kontrola"])
    records: list[tuple[int, str, str, dict]] = []
    for row in range(2, ws.max_row + 1):
        sku = str(ws.cell(row, headers["SKU"]).value or "").strip().upper()
        if not sku:
            continue
        code = code_from_sku(sku)
        record = make_record(code)
        records.append((row, sku, old_titles.get(sku, ws.cell(row, headers["Title"]).value), record))

    template_row = ws.max_row
    manufacturer_data = ws.cell(2, headers["Atrybut Produktu: Dane Producenta"]).value
    for item in additions:
        code = code_from_sku(item["sku"])
        record = make_addition_record(
            code,
            item["raw_title"],
            item["source_id"],
            item["www"],
            item["price"],
        )
        row = ws.max_row + 1
        copy_row_style(ws, template_row, row)
        for col in range(1, ws.max_column + 1):
            ws.cell(row, col).value = None
        ws.cell(row, headers["id"], None)
        ws.cell(row, headers["Title"], record["title"])
        ws.cell(row, headers["Content"], None)
        ws.cell(row, headers["SKU"], item["sku"])
        ws.cell(row, headers["Product Type"], "simple")
        ws.cell(row, headers["Parent Product ID"], "0")
        record["legacy"]["Dane Producenta"] = manufacturer_data
        for attribute, value in record["legacy"].items():
            header = f"Atrybut Produktu: {attribute}"
            if header in headers and value not in (None, ""):
                ws.cell(row, headers[header], value)
        for header, value in record["technical"].items():
            if header in headers and value not in (None, ""):
                ws.cell(row, headers[header], value)

        if code == "KZ059":
            source_text = "Hager aparatura na www.xlsx"
            source_url = "https://hager.com/pl/szkolenia-i-wsparcie/do-pobrania"
        else:
            source_text = (
                f"Hager aparatura na www.xlsx; Hager, katalog Rozdział energii 2023 ({SOURCE_FILE}); "
                "karta produktu Hager PL"
            )
            source_url = record["url"]
        ws.cell(row, headers["Źródło"], source_text)
        ws.cell(row, headers["Strona katalogu"], record["page"])
        url_cell = ws.cell(row, headers["Adres źródła"], source_url)
        url_cell.hyperlink = source_url
        url_cell.style = "Hyperlink"
        ws.cell(row, headers["Pewność"], record["confidence"])
        ws.cell(row, headers["Status"], record["status"])
        note = record["notes"]
        price_text = (
            f"{float(item['price']):.2f}".replace(".", ",")
            if item["price"] not in (None, "")
            else "brak"
        )
        source_note = f"ID źródłowe: {item['source_id']}; cena netto: {price_text} PLN."
        ws.cell(row, headers["Uwagi"], f"{note} {source_note}".strip())
        ws.cell(row, headers["Product Attributes"], php_attribute_serialization(headers, ws, row))
        status_cell = ws.cell(row, headers["Status"])
        status_cell.fill = PatternFill(
            "solid", fgColor="E2F0D9" if record["status"] == "OK" else "F4CCCC"
        )
        records.append((row, item["sku"], item["raw_title"], record))

    validate(records, added_skus)
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    make_control_sheet(wb, records, added_count=len(additions))
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    print(f"Zapisano: {output_path}")
    print(f"Produkty łącznie: {len(records)}")
    print(f"Dodano: {len(additions)}")
    print("Statusy: OK=151, DO_SPRAWDZENIA=2 (CDA440D, KZ059)")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", type=Path, default=DEFAULT_BASE)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build(args.base, args.source, args.output)
