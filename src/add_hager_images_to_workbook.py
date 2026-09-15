from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import xml.etree.ElementTree as ET
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XML = ROOT / "input" / "BMEcat_PL 17.04.2026.xml"
DEFAULT_INPUT = ROOT / "output" / "BMEcat_Hager_PL_produkty_ETIM_PL.xlsx"
DEFAULT_OUTPUT = ROOT / "output" / "BMEcat_Hager_PL_produkty_ETIM_PL_ZDJECIA.xlsx"
DEFAULT_REPORT = ROOT / "reports" / "BMEcat_Hager_PL_produkty_ETIM_PL_ZDJECIA_report.json"

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
LINK_FONT = Font(color="0563C1", underline="single")


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_product_images(xml_path: Path) -> dict[str, dict[str, list[str]]]:
    products: dict[str, dict[str, list[str]]] = {}
    for _event, element in ET.iterparse(xml_path, events=("end",)):
        if local_name(element.tag) != "PRODUCT":
            continue

        product_code = ""
        for child in element:
            if local_name(child.tag) == "SUPPLIER_PID":
                product_code = (child.text or "").strip()
                break

        image_groups = {"MD01": [], "MD25": []}
        for mime in element.iter():
            if local_name(mime.tag) != "UDX.EDXF.MIME":
                continue
            values = {
                local_name(child.tag): (child.text or "").strip()
                for child in mime
            }
            mime_code = values.get("UDX.EDXF.MIME_CODE", "")
            source = values.get("UDX.EDXF.MIME_SOURCE", "")
            if mime_code in image_groups and source and source not in image_groups[mime_code]:
                image_groups[mime_code].append(source)

        if product_code:
            if product_code in products:
                raise ValueError(f"Duplicate SUPPLIER_PID in XML: {product_code}")
            products[product_code] = image_groups
        element.clear()
    return products


def copy_header_style(source_cell, target_cell) -> None:
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.number_format = source_cell.number_format
    target_cell.protection = copy(source_cell.protection)


def add_report_sheet(workbook, metrics: dict, missing_rows: list[tuple[str, str]]) -> None:
    title = "Raport zdjęć"
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"

    sheet.merge_cells("A1:D1")
    sheet["A1"] = "Raport zdjęć produktowych — Hager BMEcat"
    sheet["A1"].fill = PatternFill("solid", fgColor=BLUE)
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28

    rows = [
        ("Metryka", "Wartość", None, None),
        ("Plik źródłowy XML", metrics["xml_file"], None, None),
        ("SHA-256 XML", metrics["xml_sha256"], None, None),
        ("Liczba produktów", metrics["product_rows"], None, None),
        ("Dopasowane kody XML ↔ XLSX", metrics["matched_products"], None, None),
        ("Produkty ze zdjęciem głównym", metrics["products_with_primary_image"], None, None),
        ("Produkty ze zdjęciem dodatkowym", metrics["products_with_additional_image"], None, None),
        ("Produkty bez zdjęcia produktowego", metrics["products_without_product_image"], None, None),
        ("Typy wykorzystanych zasobów", "MD01 Product picture; MD25 Product picture detailed view", None, None),
        ("Sposób zapisu", "Klikalne adresy URL; bez osadzania bitmap", None, None),
        (None, None, None, None),
        ("Braki w XML", None, None, None),
        ("Kod produktu", "Nazwa", "Status", "Uwagi"),
    ]
    for row in rows:
        sheet.append(row)
    for product_code, product_name in missing_rows:
        sheet.append((product_code, product_name, "BRAK W XML", "Brak MD01 i MD25"))

    for row_number in (2, 13):
        for cell in sheet[row_number]:
            cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            cell.font = Font(bold=True, color="1F1F1F")
    for cell in sheet[14]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(bold=True, color="FFFFFF")
    for row in sheet.iter_rows(min_row=15, max_col=4):
        row[2].fill = WARNING_FILL

    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 62
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 24
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_workbook(input_path: Path, xml_path: Path, output_path: Path, report_path: Path) -> dict:
    input_hash = sha256(input_path)
    xml_hash = sha256(xml_path)
    images = extract_product_images(xml_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(input_path, output_path)
    workbook = load_workbook(output_path, data_only=False)
    products_sheet = workbook["Produkty"]

    original_rows = products_sheet.max_row
    original_columns = products_sheet.max_column
    if original_rows != 14543 or original_columns != 11:
        raise ValueError(
            f"Unexpected Produkty dimensions: {original_rows} rows, {original_columns} columns"
        )

    original_formula_count = sum(
        1
        for row in products_sheet.iter_rows(min_row=2, min_col=1, max_col=11)
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    before_values = [
        tuple(cell.value for cell in row)
        for row in products_sheet.iter_rows(min_row=1, max_col=11)
    ]

    header_source = products_sheet["K1"]
    headers = ("Zdjęcie główne URL", "Zdjęcie dodatkowe URL", "Status zdjęć")
    for column, header in zip(range(12, 15), headers):
        cell = products_sheet.cell(1, column, header)
        copy_header_style(header_source, cell)

    matched = 0
    with_primary = 0
    with_additional = 0
    missing_rows: list[tuple[str, str]] = []
    for row_number in range(2, original_rows + 1):
        product_code = str(products_sheet.cell(row_number, 1).value or "").strip()
        product_name = str(products_sheet.cell(row_number, 3).value or "").strip()
        groups = images.get(product_code)
        if groups is not None:
            matched += 1
        else:
            groups = {"MD01": [], "MD25": []}

        primary_candidates = list(groups["MD01"])
        detail_candidates = list(groups["MD25"])
        if primary_candidates:
            primary_url = primary_candidates[0]
            additional_urls = primary_candidates[1:] + detail_candidates
        elif detail_candidates:
            primary_url = detail_candidates[0]
            additional_urls = detail_candidates[1:]
        else:
            primary_url = ""
            additional_urls = []

        primary_cell = products_sheet.cell(row_number, 12, primary_url or None)
        additional_cell = products_sheet.cell(
            row_number, 13, " | ".join(additional_urls) or None
        )
        status_cell = products_sheet.cell(
            row_number, 14, "OK" if primary_url else "BRAK W XML"
        )

        if primary_url:
            with_primary += 1
            primary_cell.hyperlink = primary_url
            primary_cell.font = copy(LINK_FONT)
        else:
            missing_rows.append((product_code, product_name))
            status_cell.fill = copy(WARNING_FILL)
        if additional_urls:
            with_additional += 1
            additional_cell.hyperlink = additional_urls[0]
            additional_cell.font = copy(LINK_FONT)

        for cell in (primary_cell, additional_cell, status_cell):
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            cell.number_format = "@"

    products_sheet.column_dimensions["L"].width = 72
    products_sheet.column_dimensions["M"].width = 72
    products_sheet.column_dimensions["N"].width = 16
    if products_sheet.auto_filter.ref:
        products_sheet.auto_filter.ref = f"A1:N{original_rows}"

    metrics = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "input_file": str(input_path),
        "input_sha256": input_hash,
        "xml_file": str(xml_path),
        "xml_sha256": xml_hash,
        "output_file": str(output_path),
        "product_rows": original_rows - 1,
        "matched_products": matched,
        "products_with_primary_image": with_primary,
        "products_with_additional_image": with_additional,
        "products_without_product_image": len(missing_rows),
        "missing_product_codes": [code for code, _name in missing_rows],
        "formula_count_before": original_formula_count,
    }
    add_report_sheet(workbook, metrics, missing_rows)
    workbook.save(output_path)
    workbook.close()

    validation_book = load_workbook(output_path, read_only=True, data_only=False)
    validation_sheet = validation_book["Produkty"]
    after_values = [
        tuple(cell.value for cell in row)
        for row in validation_sheet.iter_rows(min_row=1, max_col=11)
    ]
    formula_count_after = sum(
        1
        for row in validation_sheet.iter_rows(min_row=2, min_col=1, max_col=11)
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    primary_links = sum(
        1
        for (value,) in validation_sheet.iter_rows(
            min_row=2, min_col=12, max_col=12, values_only=True
        )
        if value
    )
    missing_statuses = sum(
        1
        for (value,) in validation_sheet.iter_rows(
            min_row=2, min_col=14, max_col=14, values_only=True
        )
        if value == "BRAK W XML"
    )
    validation = {
        "rows": validation_sheet.max_row,
        "columns": validation_sheet.max_column,
        "original_a_k_unchanged": before_values == after_values,
        "formula_count_after": formula_count_after,
        "formulas_unchanged": formula_count_after == original_formula_count,
        "primary_image_urls_after": primary_links,
        "missing_statuses_after": missing_statuses,
        "image_report_present": "Raport zdjęć" in validation_book.sheetnames,
    }
    validation["passed"] = all(
        (
            validation["rows"] == 14543,
            validation["columns"] == 14,
            validation["original_a_k_unchanged"],
            validation["formulas_unchanged"],
            validation["primary_image_urls_after"] == with_primary,
            validation["missing_statuses_after"] == len(missing_rows),
            validation["image_report_present"],
        )
    )
    validation_book.close()
    if not validation["passed"]:
        raise AssertionError(json.dumps(validation, ensure_ascii=False, indent=2))

    metrics["validation"] = validation
    metrics["output_sha256"] = sha256(output_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")
    return metrics


def main() -> None:
    parser = argparse.ArgumentParser(description="Add Hager BMEcat product image URLs to XLSX")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--xml", type=Path, default=DEFAULT_XML)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    metrics = build_workbook(args.input, args.xml, args.output, args.report)
    print(json.dumps(metrics, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
