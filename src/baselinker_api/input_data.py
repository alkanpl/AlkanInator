from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass, field
from io import StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from utils import normalize_header, valid_ean


ALIASES = {
    "product_id": ["product_id", "baselinker id", "id baselinker"],
    "name": ["name", "nazwa", "new_title", "tytul"],
    "sku": ["sku", "kod producenta", "kod produktu"],
    "ean": ["ean", "gtin", "ean (gtin)"],
    "manufacturer_name": ["manufacturer_name", "producent", "marka"],
    "category": ["category", "kategoria", "proponowana_kategoria_1"],
    "description": ["description", "opis html", "opis"],
    "features": ["features", "cechy", "parametry"],
    "images_urls": ["images_urls", "zdjęcia url", "zdjecia url", "obrazki"],
    "images_paths": ["images_paths", "ścieżki zdjęć", "sciezki zdjec"],
}


@dataclass
class ProductRecord:
    row_number: int
    product_id: str = ""
    name: str = ""
    sku: str = ""
    ean: str = ""
    manufacturer_name: str = ""
    category: str = ""
    description: str = ""
    features: dict[str, str] | None = None
    images_urls: list[str] = field(default_factory=list)
    images_paths: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def _string(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _headers_map(headers: list[str]) -> dict[str, str]:
    seen: dict[str, str] = {}
    duplicates: list[str] = []
    for header in headers:
        normalized = normalize_header(header)
        if not normalized:
            continue
        if normalized in seen:
            duplicates.append(header)
        else:
            seen[normalized] = header
    if duplicates:
        raise ValueError(f"Powtórzone nagłówki kolumn: {', '.join(duplicates)}")
    result: dict[str, str] = {}
    for target, candidates in ALIASES.items():
        for candidate in candidates:
            source = seen.get(normalize_header(candidate))
            if source is not None:
                result[target] = source
                break
    return result


def _detect_delimiter(text: str) -> str:
    first = next((line for line in text.splitlines() if line.strip()), "")
    counts = {delimiter: first.count(delimiter) for delimiter in (";", "\t", ",")}
    return max(counts, key=counts.get) if max(counts.values(), default=0) else ";"


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    raw = path.read_bytes()
    text = ""
    for encoding in ("utf-8-sig", "utf-8", "cp1250"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        raise ValueError("Nie można odczytać kodowania CSV (obsługiwane: UTF-8 i CP1250).")
    reader = csv.reader(StringIO(text), delimiter=_detect_delimiter(text))
    rows = list(reader)
    if not rows:
        return [], []
    headers = [_string(value) for value in rows[0]]
    _headers_map(headers)
    output: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(_string(value) for value in row):
            continue
        if len(row) > len(headers):
            raise ValueError(f"Wiersz {row_number} ma więcej pól niż nagłówek CSV.")
        output.append(dict(zip(headers, row + [""] * (len(headers) - len(row)))))
    return headers, output


def _read_xlsx(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        values = worksheet.iter_rows(values_only=True)
        header_values = next((row for row in values if any(_string(v) for v in row)), None)
        if header_values is None:
            return [], []
        headers = [_string(value) for value in header_values]
        _headers_map(headers)
        rows: list[dict[str, Any]] = []
        for row in values:
            if not any(_string(v) for v in row):
                continue
            padded = list(row) + [None] * (len(headers) - len(row))
            rows.append(dict(zip(headers, padded)))
        return headers, rows
    finally:
        workbook.close()


def _read_json(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(payload, dict):
        payload = payload.get("products", [])
    if not isinstance(payload, list) or not all(isinstance(row, dict) for row in payload):
        raise ValueError("JSON musi być listą obiektów albo obiektem z listą w polu 'products'.")
    headers: list[str] = []
    for row in payload:
        for key in row:
            key_text = str(key)
            if key_text not in headers:
                headers.append(key_text)
    _headers_map(headers)
    return headers, [dict(row) for row in payload]


def _parse_features(value: Any) -> tuple[dict[str, str] | None, str | None]:
    if value is None or _string(value) == "":
        return None, None
    parsed = value
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError as exc:
            return None, f"Niepoprawny JSON w features: {exc.msg}"
    if not isinstance(parsed, dict):
        return None, "Pole features musi być obiektem JSON."
    features: dict[str, str] = {}
    for key, raw_value in parsed.items():
        name = " ".join(str(key).split())
        feature_value = _string(raw_value)
        if name and feature_value:
            features[name] = feature_value
    return features, None


def _split_multi(value: Any) -> list[str]:
    if isinstance(value, list):
        values = [_string(item) for item in value]
    else:
        values = [_string(item) for item in re.split(r"[|\n]", _string(value))]
    return [item for item in values if item]


def load_product_records(path: str | Path) -> list[ProductRecord]:
    input_path = Path(path)
    suffix = input_path.suffix.lower()
    if suffix == ".csv":
        headers, raw_rows = _read_csv(input_path)
    elif suffix in {".xlsx", ".xlsm"}:
        headers, raw_rows = _read_xlsx(input_path)
    elif suffix == ".json":
        headers, raw_rows = _read_json(input_path)
    else:
        raise ValueError("Obsługiwane pliki wejściowe: CSV, XLSX, XLSM i JSON.")

    mapping = _headers_map(headers)
    records: list[ProductRecord] = []
    for index, row in enumerate(raw_rows, start=2):
        get = lambda key: row.get(mapping[key], "") if key in mapping else ""
        features, feature_error = _parse_features(get("features"))
        record = ProductRecord(
            row_number=index,
            product_id=_string(get("product_id")),
            name=_string(get("name")),
            sku=_string(get("sku")),
            ean=re.sub(r"\D", "", _string(get("ean"))),
            manufacturer_name=_string(get("manufacturer_name")),
            category=_string(get("category")),
            description=_string(get("description")),
            features=features,
            images_urls=_split_multi(get("images_urls")),
            images_paths=_split_multi(get("images_paths")),
        )
        if feature_error:
            record.errors.append(feature_error)
        if record.ean and not valid_ean(record.ean):
            record.warnings.append(f"EAN/GTIN nie przechodzi sumy kontrolnej: {record.ean}")
        if not record.product_id and not record.sku and not record.ean:
            record.errors.append("Brak product_id, SKU i EAN do dopasowania.")
        records.append(record)

    _mark_duplicates(records, "product_id")
    _mark_duplicates(records, "sku", casefold=True)
    _mark_duplicates(records, "ean")
    return records


def _mark_duplicates(records: list[ProductRecord], field_name: str, *, casefold: bool = False) -> None:
    grouped: dict[str, list[ProductRecord]] = {}
    for record in records:
        value = str(getattr(record, field_name, "")).strip()
        if not value:
            continue
        normalized = value.casefold() if casefold else value
        grouped.setdefault(normalized, []).append(record)
    for value, matches in grouped.items():
        if len(matches) < 2:
            continue
        for record in matches:
            record.errors.append(f"Powtórzona wartość {field_name} w wejściu: {value}")
