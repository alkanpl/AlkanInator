from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = (
    ROOT
    / "output"
    / "Hager-Berker-Gniazdka_WOO_ALL_IMPORT_CROSS_UP_SELL_PO_AUDYCIE_PDF.xlsx"
)
DEFAULT_OUTPUT = (
    ROOT
    / "output"
    / "Hager-Berker-Gniazdka_WOO_ALL_IMPORT_CROSS_UP_SELL_PO_AUDYCIE_PDF_TYTULY_STARYCH_SKU.xlsx"
)

BERKER_ALIAS_RE = re.compile(r"\(\s*Berker\s+([A-Z0-9._/-]+)\s*\)", re.IGNORECASE)
MOJIBAKE_MARKERS = ("\ufffd", "Ã", "Â", "Ä", "Å", "â€")

# The Polish replacement-reference PDF describes WTW1000PBG as white, but the
# source catalogue page for Berker 16476500 and Hager product data identify the
# porcelain Rosenthal part as black, glossy. Keep the old product factually
# correct while applying the new Hager/Berker title convention.
LEGACY_TITLE_OVERRIDES = {
    "16476500": (
        "Płytka czołowa z pokrętłem do łącznika obrotowego, porcelana Rosenthal, "
        "czarny, połysk Hager Serie 1930 WTW1000PBG (Berker 16476500)"
    ),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalized_code(value: object) -> str:
    return str(value or "").strip().upper().removesuffix("/HAG")


def update_legacy_titles(input_path: Path, output_path: Path) -> dict[str, object]:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("Plik wyjściowy musi być inną kopią niż plik wejściowy.")
    if output_path.exists():
        raise FileExistsError(f"Plik wyjściowy już istnieje: {output_path}")

    source_hash_before = sha256(input_path)
    workbook = load_workbook(input_path)
    sheet = workbook.active

    headers = {
        str(cell.value).strip(): cell.column
        for cell in sheet[1]
        if cell.value is not None and str(cell.value).strip()
    }
    missing_headers = {"SKU", "Title"} - set(headers)
    if missing_headers:
        raise KeyError(f"Brak wymaganych kolumn: {sorted(missing_headers)}")

    sku_column = headers["SKU"]
    title_column = headers["Title"]
    rows_by_code: dict[str, int] = {}
    for row in range(2, sheet.max_row + 1):
        code = normalized_code(sheet.cell(row, sku_column).value)
        if not code:
            continue
        if code in rows_by_code:
            raise ValueError(f"Powtórzone SKU w arkuszu: {code}/HAG")
        rows_by_code[code] = row

    pairs_by_old_code: dict[str, tuple[int, int]] = {}
    for new_code, new_row in rows_by_code.items():
        new_title = str(sheet.cell(new_row, title_column).value or "").strip()
        match = BERKER_ALIAS_RE.search(new_title)
        if not match:
            continue
        old_code = normalized_code(match.group(1))
        old_row = rows_by_code.get(old_code)
        if old_row is None or old_code == new_code:
            continue
        previous = pairs_by_old_code.get(old_code)
        if previous and previous != (old_row, new_row):
            raise ValueError(f"Niejednoznaczny kod zastępujący dla Berker {old_code}")
        pairs_by_old_code[old_code] = (old_row, new_row)

    if not pairs_by_old_code:
        raise ValueError("Nie znaleziono żadnych jednoznacznych par stare SKU ↔ nowe SKU.")

    changes: list[dict[str, object]] = []
    for old_code, (old_row, new_row) in sorted(pairs_by_old_code.items()):
        new_code = normalized_code(sheet.cell(new_row, sku_column).value)
        new_title = str(sheet.cell(new_row, title_column).value or "").strip()
        target_title = LEGACY_TITLE_OVERRIDES.get(old_code, new_title)
        old_cell = sheet.cell(old_row, title_column)
        old_title = str(old_cell.value or "")

        expected_alias = f"(Berker {old_code})"
        if new_code not in target_title.upper() or expected_alias.lower() not in target_title.lower():
            raise ValueError(
                f"Tytuł docelowy nie zawiera obu kodów dla pary {old_code} → {new_code}: "
                f"{target_title}"
            )

        old_cell.value = target_title
        changes.append(
            {
                "old_sku": f"{old_code}/HAG",
                "new_sku": f"{new_code}/HAG",
                "row": old_row,
                "old_title": old_title,
                "new_title": target_title,
                "used_catalog_correction": old_code in LEGACY_TITLE_OVERRIDES,
            }
        )

    suspicious_cells: list[str] = []
    for row in sheet.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            if any(marker in cell.value for marker in MOJIBAKE_MARKERS):
                suspicious_cells.append(cell.coordinate)
                if len(suspicious_cells) >= 20:
                    break
        if len(suspicious_cells) >= 20:
            break
    if suspicious_cells:
        raise ValueError(f"Wykryto możliwe uszkodzenie kodowania w komórkach: {suspicious_cells}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

    source_hash_after = sha256(input_path)
    if source_hash_after != source_hash_before:
        raise RuntimeError("Plik wejściowy został nieoczekiwanie zmieniony.")

    return {
        "input": str(input_path),
        "output": str(output_path),
        "changed_title_count": len(changes),
        "source_unchanged": True,
        "source_sha256": source_hash_after,
        "changes": changes,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Kopiuje nowe tytuły Hager z aliasem '(Berker ...)' do odpowiadających "
            "im starych SKU, bez zmiany SKU ani pozostałych pól."
        )
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    result = update_legacy_titles(args.input.resolve(), args.output.resolve())
    print(json.dumps(result, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
