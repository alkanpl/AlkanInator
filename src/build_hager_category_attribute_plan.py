"""Build a category and attribute-priority plan for the prepared Hager workbook."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


DEFAULT_INPUT = Path("output/Aparatura-Modulowa-Hager_nazwy_atrybuty_v2.xlsx")
DEFAULT_CONFIG = Path("configs/hager_category_attribute_priorities.yaml")
DEFAULT_OUTPUT = Path("output/Hager_aparatura_kategorie_atrybuty_priorytety.xlsx")

DARK_BLUE = "17365D"
MID_BLUE = "366092"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "EEF5FB"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="D9E1F2")


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def is_filled(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def title_band(ws, title: str, subtitle: str, last_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, title)
    cell.font = Font(size=16, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    cell = ws.cell(2, 1, subtitle)
    cell.font = Font(size=9, italic=True, color="4F4F4F")
    cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30


def table_header(ws, row: int, headers: list[str]) -> None:
    for col, value in enumerate(headers, 1):
        cell = ws.cell(row, col, value)
        cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EADBF"))
    ws.row_dimensions[row].height = 34


def style_data_rows(ws, start_row: int, end_row: int, end_col: int, wrap_cols: set[int]) -> None:
    for row in range(start_row, end_row + 1):
        for col in range(1, end_col + 1):
            cell = ws.cell(row, col)
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
            cell.border = Border(bottom=THIN_GRAY)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=col in wrap_cols,
                horizontal="right" if isinstance(cell.value, (int, float)) else "left",
            )


def add_excel_table(ws, ref: str, name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def load_products(input_path: Path) -> tuple[list[dict], dict[str, int]]:
    wb = load_workbook(input_path, read_only=False, data_only=False)
    main = wb["Worksheet"]
    control = wb["Kontrola"]
    main_headers = {
        text(cell.value): cell.column
        for cell in main[1]
        if cell.value is not None
    }
    control_headers = {
        text(control.cell(5, col).value): col
        for col in range(1, control.max_column + 1)
        if control.cell(5, col).value is not None
    }
    control_by_sku = {}
    for row in range(6, control.max_row + 1):
        sku = text(control.cell(row, control_headers["SKU"]).value).upper()
        if sku:
            control_by_sku[sku] = {
                "family": text(control.cell(row, control_headers["Rodzina"]).value),
                "old_title": text(control.cell(row, control_headers["Nazwa pierwotna"]).value),
            }
    products = []
    for row in range(2, main.max_row + 1):
        sku = text(main.cell(row, main_headers["SKU"]).value).upper()
        if not sku:
            continue
        products.append(
            {
                "row": row,
                "sku": sku,
                "code": sku.split("/", 1)[0],
                "title": text(main.cell(row, main_headers["Title"]).value),
                "family": control_by_sku[sku]["family"],
                "source_status": text(main.cell(row, main_headers["Status"]).value),
                "source": text(main.cell(row, main_headers["Źródło"]).value),
                "source_page": text(main.cell(row, main_headers["Strona katalogu"]).value),
                "source_url": text(main.cell(row, main_headers["Adres źródła"]).value),
                "values": {
                    header.removeprefix("Atrybut Produktu: "): main.cell(row, col).value
                    for header, col in main_headers.items()
                    if header.startswith("Atrybut Produktu: ")
                },
            }
        )
    return products, main_headers


def build_category_maps(config: dict) -> tuple[dict[str, dict], dict[str, dict]]:
    by_id = {item["id"]: item for item in config["categories"]}
    by_family = {}
    for category in config["categories"]:
        for family in category["families"]:
            if family in by_family:
                raise ValueError(f"Rodzina przypisana do dwóch kategorii: {family}")
            by_family[family] = category
    return by_id, by_family


def expanded_attributes(config: dict, category: dict) -> list[dict]:
    return [dict(item) for item in config["common_attributes"]] + [
        dict(item) for item in category["attributes"]
    ]


def attribute_required_for_product(category: dict, attribute: dict, product: dict) -> bool:
    if attribute["priority"] != "P0":
        return False
    name = attribute["attribute"]
    if category["id"] == "spd":
        spd_type = text(product["values"].get("Typ SPD"))
        if name == "Prąd Udarowy Iimp [kA, łącznie]":
            return "T1" in spd_type
        if name == "Znamionowy Prąd Wyładowczy In [kA]":
            return "T2" in spd_type
    return True


def examples(values: list[object], limit: int = 5) -> str:
    unique = []
    for value in values:
        candidate = text(value)
        if candidate and candidate not in unique:
            unique.append(candidate)
    if len(unique) > limit:
        return "; ".join(unique[:limit]) + f"; +{len(unique) - limit}"
    return "; ".join(unique)


def build_summary_sheet(wb: Workbook, config: dict, categories: list[dict]) -> None:
    ws = wb.create_sheet("Podsumowanie")
    ws.sheet_view.showGridLines = False
    title_band(
        ws,
        "Hager - plan kategorii i priorytetów atrybutów",
        "Źródło danych: Aparatura-Modulowa-Hager_nazwy_atrybuty_v2.xlsx. Priorytety P0/P1/P2 służą do nazw, filtrów i opisów produktów.",
        8,
    )

    cards = [
        ("A4:B5", "Produkty", "=COUNTA('Produkty i kategorie'!$A$6:$A$158)", "0"),
        ("C4:D5", "Kategorie", "=COUNTA('Kategorie'!$A$6:$A$22)", "0"),
        (
            "E4:F5",
            "Kompletność P0",
            "=COUNTIF('Produkty i kategorie'!$J$6:$J$158,\"OK\")/COUNTA('Produkty i kategorie'!$A$6:$A$158)",
            "0.0%",
        ),
        (
            "G4:H5",
            "Kategorie do utworzenia",
            "=COUNTIF('Kategorie'!$D$6:$D$22,\"DO_UTWORZENIA\")",
            "0",
        ),
    ]
    for merged, label, formula, number_format in cards:
        ws.merge_cells(merged)
        top_left = ws[merged.split(":")[0]]
        top_left.value = f"{label}\n"
        top_left.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        top_left.font = Font(bold=True, color=DARK_BLUE, size=11)
        top_left.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        top_left.border = Border(
            left=Side(style="medium", color=MID_BLUE),
            right=Side(style="medium", color=MID_BLUE),
            top=Side(style="medium", color=MID_BLUE),
            bottom=Side(style="medium", color=MID_BLUE),
        )
        formula_cell = ws.cell(top_left.row + 1, top_left.column)
        # Merged 2x2 cards use a separate centered result overlay in the lower-left cell.
        ws.unmerge_cells(merged)
        ws.merge_cells(
            start_row=top_left.row,
            start_column=top_left.column,
            end_row=top_left.row,
            end_column=top_left.column + 1,
        )
        ws.merge_cells(
            start_row=top_left.row + 1,
            start_column=top_left.column,
            end_row=top_left.row + 1,
            end_column=top_left.column + 1,
        )
        label_cell = ws.cell(top_left.row, top_left.column)
        label_cell.value = label
        label_cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        label_cell.font = Font(bold=True, color=DARK_BLUE)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell = ws.cell(top_left.row + 1, top_left.column)
        value_cell.value = formula
        value_cell.number_format = number_format
        value_cell.fill = PatternFill("solid", fgColor=WHITE)
        value_cell.font = Font(bold=True, color=DARK_BLUE, size=15)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in range(top_left.row, top_left.row + 2):
            for col in range(top_left.column, top_left.column + 2):
                ws.cell(row, col).border = Border(
                    left=Side(style="thin", color=MID_BLUE),
                    right=Side(style="thin", color=MID_BLUE),
                    top=Side(style="thin", color=MID_BLUE),
                    bottom=Side(style="thin", color=MID_BLUE),
                )

    headers = [
        "Kategoria",
        "Status kategorii",
        "Liczba produktów",
        "Atrybuty P0",
        "Średnia kompletność P0",
        "Produkty z brakami P0",
        "Rodziny źródłowe",
        "Uwagi",
    ]
    table_header(ws, 8, headers)
    for index, category in enumerate(categories, 9):
        source_last = 158
        matrix_last = 400
        ws.cell(index, 1, category["path"])
        ws.cell(index, 2, category["taxonomy_status"])
        ws.cell(
            index,
            3,
            f'=COUNTIF(\'Produkty i kategorie\'!$D$6:$D${source_last},A{index})',
        )
        ws.cell(
            index,
            4,
            f'=COUNTIFS(\'Kategorie i atrybuty\'!$B$6:$B${matrix_last},A{index},\'Kategorie i atrybuty\'!$E$6:$E${matrix_last},"P0")',
        )
        ws.cell(
            index,
            5,
            f'=IFERROR(AVERAGEIF(\'Produkty i kategorie\'!$D$6:$D${source_last},A{index},\'Produkty i kategorie\'!$I$6:$I${source_last}),0)',
        )
        ws.cell(
            index,
            6,
            f'=COUNTIFS(\'Produkty i kategorie\'!$D$6:$D${source_last},A{index},\'Produkty i kategorie\'!$J$6:$J${source_last},"DO_UZUPELNIENIA")',
        )
        ws.cell(index, 7, "; ".join(category["families"]))
        ws.cell(
            index,
            8,
            "Ścieżka istnieje w lokalnej taksonomii Alkan."
            if category["taxonomy_status"] == "ISTNIEJE"
            else "Rekomendowana osobna kategoria - wymaga akceptacji lub utworzenia w sklepie.",
        )
        ws.cell(index, 5).number_format = "0.0%"
    last_row = 8 + len(categories)
    style_data_rows(ws, 9, last_row, 8, {1, 7, 8})
    ws.auto_filter.ref = f"A8:H{last_row}"
    ws.freeze_panes = "A9"
    widths = [48, 19, 16, 13, 20, 20, 45, 58]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = width
    ws.conditional_formatting.add(
        f"B9:B{last_row}",
        FormulaRule(formula=['B9="ISTNIEJE"'], fill=PatternFill("solid", fgColor=GREEN)),
    )
    ws.conditional_formatting.add(
        f"B9:B{last_row}",
        FormulaRule(formula=['B9="DO_UTWORZENIA"'], fill=PatternFill("solid", fgColor=YELLOW)),
    )
    ws.conditional_formatting.add(
        f"E9:E{last_row}",
        CellIsRule(operator="lessThan", formula=["1"], fill=PatternFill("solid", fgColor=RED)),
    )


def build_categories_sheet(wb: Workbook, categories: list[dict]) -> None:
    ws = wb.create_sheet("Kategorie")
    ws.sheet_view.showGridLines = False
    title_band(
        ws,
        "Kategorie docelowe Hager",
        "Status ISTNIEJE pochodzi z lokalnej taksonomii Alkan. DO_UTWORZENIA oznacza rekomendowaną, bardziej precyzyjną kategorię.",
        10,
    )
    headers = [
        "ID kategorii",
        "Ścieżka kategorii",
        "Nazwa robocza",
        "Status",
        "Rodziny źródłowe",
        "Liczba produktów",
        "Liczba P0",
        "Liczba P1",
        "Liczba P2",
        "Decyzja / uwagi",
    ]
    table_header(ws, 5, headers)
    for row, category in enumerate(categories, 6):
        ws.cell(row, 1, category["id"])
        ws.cell(row, 2, category["path"])
        ws.cell(row, 3, category["label"])
        ws.cell(row, 4, category["taxonomy_status"])
        ws.cell(row, 5, "; ".join(category["families"]))
        ws.cell(
            row,
            6,
            f'=COUNTIF(\'Produkty i kategorie\'!$D$6:$D$158,B{row})',
        )
        for col, priority in [(7, "P0"), (8, "P1"), (9, "P2")]:
            ws.cell(
                row,
                col,
                f'=COUNTIFS(\'Kategorie i atrybuty\'!$A$6:$A$400,A{row},\'Kategorie i atrybuty\'!$E$6:$E$400,"{priority}")',
            )
        ws.cell(
            row,
            10,
            "Zaakceptować istniejącą ścieżkę."
            if category["taxonomy_status"] == "ISTNIEJE"
            else "Potwierdzić nazwę i utworzyć kategorię albo przypisać do nadrzędnej Aparatura modułowa.",
        )
    last_row = 5 + len(categories)
    style_data_rows(ws, 6, last_row, 10, {2, 3, 5, 10})
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:J{last_row}"
    widths = [20, 52, 38, 19, 48, 16, 12, 12, 12, 65]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = width
    ws.conditional_formatting.add(
        f"D6:D{last_row}",
        FormulaRule(formula=['D6="ISTNIEJE"'], fill=PatternFill("solid", fgColor=GREEN)),
    )
    ws.conditional_formatting.add(
        f"D6:D{last_row}",
        FormulaRule(formula=['D6="DO_UTWORZENIA"'], fill=PatternFill("solid", fgColor=YELLOW)),
    )


def build_matrix_sheet(
    wb: Workbook,
    config: dict,
    categories: list[dict],
    products_by_category: dict[str, list[dict]],
) -> int:
    ws = wb.create_sheet("Kategorie i atrybuty")
    ws.sheet_view.showGridLines = False
    title_band(
        ws,
        "Matryca kategorii, atrybutów i priorytetów",
        "P0 = obowiązkowy, P1 = ważny, P2 = uzupełniający. Pokrycie liczone jest na podstawie 153 produktów w aktualnym pliku Hager.",
        15,
    )
    headers = [
        "ID kategorii",
        "Ścieżka kategorii",
        "Status kategorii",
        "Atrybut",
        "Priorytet",
        "Wymagane",
        "W nazwie",
        "Rola",
        "Uzasadnienie",
        "Produkty z wartością",
        "Produkty w kategorii",
        "Pokrycie",
        "Przykładowe wartości",
        "Źródło reguły",
        "Status pokrycia",
    ]
    table_header(ws, 5, headers)
    row = 6
    for category in categories:
        category_products = products_by_category[category["id"]]
        for attribute in expanded_attributes(config, category):
            name = attribute["attribute"]
            values = [product["values"].get(name) for product in category_products]
            filled = sum(is_filled(value) for value in values)
            conditional = (
                category["id"] == "spd"
                and name
                in {
                    "Prąd Udarowy Iimp [kA, łącznie]",
                    "Znamionowy Prąd Wyładowczy In [kA]",
                }
            )
            ws.cell(row, 1, category["id"])
            ws.cell(row, 2, category["path"])
            ws.cell(row, 3, category["taxonomy_status"])
            ws.cell(row, 4, name)
            ws.cell(row, 5, attribute["priority"])
            ws.cell(
                row,
                6,
                "Warunkowo" if conditional else ("Tak" if attribute["priority"] == "P0" else "Nie"),
            )
            ws.cell(row, 7, attribute["in_title"])
            ws.cell(row, 8, attribute["role"])
            ws.cell(row, 9, attribute["reason"])
            ws.cell(row, 10, filled)
            ws.cell(row, 11, len(category_products))
            ws.cell(row, 12, filled / len(category_products) if category_products else 0)
            ws.cell(row, 12).number_format = "0.0%"
            ws.cell(row, 13, examples(values))
            ws.cell(
                row,
                14,
                "Plan Hager + lokalna taksonomia Alkan + katalog Hager 2023",
            )
            ws.cell(
                row,
                15,
                f'=IF(AND(E{row}="P0",F{row}="Tak",L{row}<1),"BRAKI_P0",IF(AND(E{row}="P1",L{row}<0.7),"NISKIE_POKRYCIE_P1","OK"))',
            )
            row += 1
    last_row = row - 1
    style_data_rows(ws, 6, last_row, 15, {2, 4, 8, 9, 13, 14})
    ws.freeze_panes = "D6"
    ws.auto_filter.ref = f"A5:O{last_row}"
    widths = [18, 52, 19, 43, 11, 13, 13, 23, 58, 18, 19, 13, 52, 46, 16]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col) if col <= 26 else "A"].width = width
    validation = DataValidation(type="list", formula1='"P0,P1,P2"', allow_blank=False)
    ws.add_data_validation(validation)
    validation.add(f"E6:E{last_row}")
    for priority, color in [("P0", RED), ("P1", YELLOW), ("P2", GREEN)]:
        ws.conditional_formatting.add(
            f"E6:E{last_row}",
            FormulaRule(formula=[f'E6="{priority}"'], fill=PatternFill("solid", fgColor=color)),
        )
    ws.conditional_formatting.add(
        f"O6:O{last_row}",
        FormulaRule(formula=['O6="BRAKI_P0"'], fill=PatternFill("solid", fgColor=RED)),
    )
    ws.conditional_formatting.add(
        f"O6:O{last_row}",
        FormulaRule(formula=['O6="NISKIE_POKRYCIE_P1"'], fill=PatternFill("solid", fgColor=YELLOW)),
    )
    ws.conditional_formatting.add(
        f"O6:O{last_row}",
        FormulaRule(formula=['O6="OK"'], fill=PatternFill("solid", fgColor=GREEN)),
    )
    return last_row


def build_products_sheet(
    wb: Workbook,
    config: dict,
    products: list[dict],
    by_family: dict[str, dict],
) -> None:
    ws = wb.create_sheet("Produkty i kategorie")
    ws.sheet_view.showGridLines = False
    title_band(
        ws,
        "Przypisanie produktów Hager do kategorii",
        "Kompletność P0 uwzględnia atrybuty wspólne i wymagane atrybuty kategorii. Dla SPD Iimp jest wymagany dla T1, a In dla T2.",
        13,
    )
    headers = [
        "SKU",
        "Kod",
        "Nazwa produktu",
        "Kategoria docelowa",
        "Rodzina źródłowa",
        "Status produktu źródłowego",
        "Wypełnione P0",
        "Wymagane P0",
        "Kompletność P0",
        "Status P0",
        "Brakujące P0",
        "Źródło / strona",
        "Adres źródła",
    ]
    table_header(ws, 5, headers)
    for row, product in enumerate(products, 6):
        category = by_family[product["family"]]
        attributes = expanded_attributes(config, category)
        required = [
            attribute
            for attribute in attributes
            if attribute_required_for_product(category, attribute, product)
        ]
        missing = [
            attribute["attribute"]
            for attribute in required
            if not is_filled(product["values"].get(attribute["attribute"]))
        ]
        filled = len(required) - len(missing)
        ws.cell(row, 1, product["sku"])
        ws.cell(row, 2, product["code"])
        ws.cell(row, 3, product["title"])
        ws.cell(row, 4, category["path"])
        ws.cell(row, 5, product["family"])
        ws.cell(row, 6, product["source_status"])
        ws.cell(row, 7, filled)
        ws.cell(row, 8, len(required))
        ws.cell(row, 9, f"=IF(H{row}=0,1,G{row}/H{row})")
        ws.cell(row, 9).number_format = "0.0%"
        ws.cell(row, 10, f'=IF(I{row}=1,"OK","DO_UZUPELNIENIA")')
        ws.cell(row, 11, "; ".join(missing))
        ws.cell(row, 12, f"{product['source']} | str. {product['source_page']}")
        ws.cell(row, 13, product["source_url"])
        ws.cell(row, 13).hyperlink = product["source_url"]
        ws.cell(row, 13).style = "Hyperlink"
    last_row = 5 + len(products)
    style_data_rows(ws, 6, last_row, 13, {3, 4, 5, 11, 12})
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:M{last_row}"
    widths = [19, 14, 55, 52, 38, 22, 15, 15, 17, 20, 50, 60, 46]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = width
    ws.conditional_formatting.add(
        f"J6:J{last_row}",
        FormulaRule(formula=['J6="OK"'], fill=PatternFill("solid", fgColor=GREEN)),
    )
    ws.conditional_formatting.add(
        f"J6:J{last_row}",
        FormulaRule(formula=['J6="DO_UZUPELNIENIA"'], fill=PatternFill("solid", fgColor=RED)),
    )
    ws.conditional_formatting.add(
        f"F6:F{last_row}",
        FormulaRule(formula=['F6="DO_SPRAWDZENIA"'], fill=PatternFill("solid", fgColor=YELLOW)),
    )


def build_dictionary_sheet(wb: Workbook, config: dict) -> None:
    ws = wb.create_sheet("Słownik priorytetów")
    ws.sheet_view.showGridLines = False
    title_band(
        ws,
        "Słownik priorytetów i zasady stosowania",
        "Ten arkusz definiuje znaczenie priorytetów oraz pola wspólne dla każdej kategorii.",
        6,
    )
    table_header(ws, 5, ["Priorytet", "Nazwa", "Znaczenie", "Minimalne zastosowanie", "Kolor", "Decyzja"])
    uses = {
        "P0": "Nazwa, podstawowe filtry, przypisanie produktu",
        "P1": "Porównanie produktów, opis techniczny, dodatkowe filtry",
        "P2": "Rozszerzony opis i informacje montażowe",
    }
    colors = {"P0": "czerwony", "P1": "żółty", "P2": "zielony"}
    for row, priority in enumerate(["P0", "P1", "P2"], 6):
        definition = config["priorities"][priority]
        ws.cell(row, 1, priority)
        ws.cell(row, 2, definition["label"])
        ws.cell(row, 3, definition["definition"])
        ws.cell(row, 4, uses[priority])
        ws.cell(row, 5, colors[priority])
        ws.cell(row, 6, "Utrzymać jako regułę bazową")
        ws.cell(row, 1).fill = PatternFill(
            "solid", fgColor={"P0": RED, "P1": YELLOW, "P2": GREEN}[priority]
        )
    style_data_rows(ws, 6, 8, 6, {3, 4, 6})

    table_header(ws, 11, ["Atrybut wspólny", "Priorytet", "W nazwie", "Rola", "Uzasadnienie", "Zakres"])
    for row, attribute in enumerate(config["common_attributes"], 12):
        ws.cell(row, 1, attribute["attribute"])
        ws.cell(row, 2, attribute["priority"])
        ws.cell(row, 3, attribute["in_title"])
        ws.cell(row, 4, attribute["role"])
        ws.cell(row, 5, attribute["reason"])
        ws.cell(row, 6, "Wszystkie kategorie Hager")
    common_last = 11 + len(config["common_attributes"])
    style_data_rows(ws, 12, common_last, 6, {1, 4, 5, 6})

    table_header(ws, common_last + 3, ["Źródło", "Zastosowanie", "Ścieżka / adres", "Status", "Uwagi", "Właściciel"])
    source_rows = [
        (
            "Plik produktowy Hager",
            "Pokrycie atrybutów i lista SKU",
            "output/Aparatura-Modulowa-Hager_nazwy_atrybuty_v2.xlsx",
            "AKTUALNE",
            "153 produkty",
            "Alkan",
        ),
        (
            "Lokalna taksonomia Alkan",
            "Status istniejących kategorii",
            "dictionaries/woocommerce_catalog_knowledge.yaml; dictionaries/learned_product_taxonomy.yaml",
            "AKTUALNE",
            "Ścieżki z lokalnego katalogu WooCommerce",
            "Alkan",
        ),
        (
            "Polski katalog Hager 2023",
            "Parametry techniczne produktów",
            "input/23PL013_KATALOG ROZDZIAL ENERGII_2023.pdf",
            "ŹRÓDŁO",
            "Strony źródłowe są zapisane w pliku produktowym",
            "Hager",
        ),
    ]
    source_start = common_last + 4
    for row, values in enumerate(source_rows, source_start):
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
    style_data_rows(ws, source_start, source_start + len(source_rows) - 1, 6, {1, 2, 3, 5})
    ws.freeze_panes = "A5"
    widths = [30, 23, 62, 19, 65, 18]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = width


def validate(products: list[dict], categories: list[dict], by_family: dict[str, dict]) -> None:
    assert len(products) == 153, len(products)
    assert len(categories) == 17, len(categories)
    assert len({product["sku"] for product in products}) == 153
    missing_families = sorted({product["family"] for product in products} - set(by_family))
    assert not missing_families, missing_families
    counts = Counter(by_family[product["family"]]["id"] for product in products)
    assert sum(counts.values()) == 153
    assert all(counts[category["id"]] > 0 for category in categories)


def build(input_path: Path, config_path: Path, output_path: Path) -> None:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("Plik wynikowy nie może nadpisywać pliku produktowego")
    with config_path.open("r", encoding="utf-8-sig") as handle:
        config = yaml.safe_load(handle)
    products, _ = load_products(input_path)
    _, by_family = build_category_maps(config)
    categories = config["categories"]
    validate(products, categories, by_family)

    products_by_category: dict[str, list[dict]] = defaultdict(list)
    for product in products:
        products_by_category[by_family[product["family"]]["id"]].append(product)

    wb = Workbook()
    wb.remove(wb.active)
    build_summary_sheet(wb, config, categories)
    build_categories_sheet(wb, categories)
    matrix_last = build_matrix_sheet(wb, config, categories, products_by_category)
    build_products_sheet(wb, config, products, by_family)
    build_dictionary_sheet(wb, config)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    print(f"output={output_path}")
    print(f"products={len(products)} categories={len(categories)} matrix_rows={matrix_last - 5}")
    print(
        "category_statuses="
        + str(dict(Counter(category["taxonomy_status"] for category in categories)))
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build(args.input, args.config, args.output)
