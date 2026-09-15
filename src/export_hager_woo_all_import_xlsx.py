from __future__ import annotations

import argparse
import json
import shutil
from copy import copy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from export_hager_woo_variant_updates import PhpUnserializer, clean, sha256


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "input" / "Hager-Berker-Gniazdka.xlsx"
DEFAULT_ANALYSIS = ROOT / "output" / "Hager-Berker-Gniazdka_WARIANTY_ATRYBUTY_I_TAGI.xlsx"
DEFAULT_OUTPUT = ROOT / "output" / "Hager-Berker-Gniazdka_WOO_ALL_IMPORT_ATRYBUTY_I_TAGI.xlsx"
DEFAULT_REPORT = ROOT / "reports" / "Hager-Berker-Gniazdka_WOO_ALL_IMPORT_ATRYBUTY_I_TAGI_report.json"

ORIENTATION_HEADER = "Atrybut Produktu: Orientacja"
EDITED_HEADERS = (
    "Atrybut Produktu: Kolor",
    "Atrybut Produktu: Kolor Producenta",
    "Atrybut Produktu: Materiał",
    "Atrybut Produktu: Krotność",
    "Atrybut Produktu: Kategoria Kabli Sieciowych",
    "Atrybut Produktu: Ekranowany",
    "Atrybut Produktu: Uziemienie",
    "Atrybut Produktu: Typ Zacisków",
    "Atrybut Produktu: Napięcie [V]",
    "Atrybut Produktu: Prąd Znamionowy [A]",
    "Atrybut Produktu: Liczba Biegunów",
    "Atrybut Produktu: Przesłony Torów Prądowych",
    "Product Tags",
)
ATTRIBUTE_SLUGS = {
    "Atrybut Produktu: Kolor": "pa_kolor",
    "Atrybut Produktu: Kolor Producenta": "pa_kolor-producenta",
    "Atrybut Produktu: Materiał": "pa_material",
    "Atrybut Produktu: Krotność": "pa_krotnosc",
    "Atrybut Produktu: Kategoria Kabli Sieciowych": "pa_kategoria-kabli-sieciowych",
    "Atrybut Produktu: Ekranowany": "pa_ekranowany",
    "Atrybut Produktu: Uziemienie": "pa_uziemienie",
    "Atrybut Produktu: Typ Zacisków": "pa_typ-zaciskow",
    "Atrybut Produktu: Napięcie [V]": "pa_napiecie-v",
    "Atrybut Produktu: Prąd Znamionowy [A]": "pa_prad-znamionowy-a",
    "Atrybut Produktu: Liczba Biegunów": "pa_liczba-biegunow",
    "Atrybut Produktu: Przesłony Torów Prądowych": "pa_przeslony-torow-pradowych",
    ORIENTATION_HEADER: "pa_orientacja",
}


def php_serialize(value: Any) -> str:
    if value is None:
        return "N;"
    if isinstance(value, bool):
        return f"b:{1 if value else 0};"
    if isinstance(value, int):
        return f"i:{value};"
    if isinstance(value, float):
        return f"d:{value};"
    if isinstance(value, str):
        return f's:{len(value.encode("utf-8"))}:"{value}";'
    if isinstance(value, dict):
        body = "".join(php_serialize(key) + php_serialize(item) for key, item in value.items())
        return f"a:{len(value)}:{{{body}}}"
    if isinstance(value, (list, tuple)):
        payload = {index: item for index, item in enumerate(value)}
        return php_serialize(payload)
    raise TypeError(f"Nieobsługiwany typ serializacji PHP: {type(value).__name__}")


def parse_product_attributes(value: Any) -> dict[Any, Any]:
    serialized = clean(value)
    if not serialized:
        return {}
    payload = PhpUnserializer(serialized).parse()
    if not isinstance(payload, dict):
        raise ValueError("Product Attributes nie zawiera tablicy PHP")
    return payload


def attribute_is_declared(payload: dict[Any, Any], slug: str) -> bool:
    for key, details in payload.items():
        if clean(key) == slug:
            return True
        if isinstance(details, dict) and clean(details.get("name")) == slug:
            return True
    return False


def add_attribute_declaration(payload: dict[Any, Any], slug: str) -> bool:
    if attribute_is_declared(payload, slug):
        return False
    positions = [
        int(details.get("position", 0))
        for details in payload.values()
        if isinstance(details, dict)
    ]
    payload[slug] = {
        "name": slug,
        "value": "",
        "position": max(positions, default=-1) + 1,
        "is_visible": 1,
        "is_variation": 0,
        "is_taxonomy": 1,
    }
    return True


def copy_cell_style(source_cell, target_cell) -> None:
    target_cell._style = copy(source_cell._style)
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = source_cell.number_format
        target_cell.protection = copy(source_cell.protection)


def build_file(source_path: Path, analysis_path: Path, output_path: Path) -> dict[str, Any]:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, output_path)

    source_book = load_workbook(output_path, data_only=False)
    analysis_book = load_workbook(analysis_path, read_only=True, data_only=False)
    try:
        if len(source_book.sheetnames) != 1:
            raise ValueError(f"Eksport źródłowy powinien mieć jeden arkusz: {source_book.sheetnames}")
        source_sheet = source_book.active
        analysis_sheet = analysis_book["Produkty warianty"]

        source_headers = [cell.value for cell in source_sheet[1]]
        analysis_headers = [cell.value for cell in analysis_sheet[1]]
        source_columns = {value: index + 1 for index, value in enumerate(source_headers)}
        analysis_columns = {value: index for index, value in enumerate(analysis_headers)}
        original_column_count = source_sheet.max_column
        original_row_count = source_sheet.max_row

        for required in ("SKU", "Product Attributes", *EDITED_HEADERS):
            if required not in source_columns:
                raise KeyError(f"Brak kolumny w eksporcie źródłowym: {required}")
        for required in ("SKU", "Wariant: tag", "Wariant: decyzja", ORIENTATION_HEADER, *EDITED_HEADERS):
            if required not in analysis_columns:
                raise KeyError(f"Brak kolumny w pliku analitycznym: {required}")

        source_rows: dict[str, int] = {}
        for row in range(2, source_sheet.max_row + 1):
            sku = clean(source_sheet.cell(row, source_columns["SKU"]).value)
            if sku:
                if sku in source_rows:
                    raise ValueError(f"Powtórzony SKU w eksporcie źródłowym: {sku}")
                source_rows[sku] = row

        orientation_column = original_column_count + 1
        orientation_header_cell = source_sheet.cell(1, orientation_column, ORIENTATION_HEADER)
        style_source_column = source_columns["Atrybut Produktu: Krotność"]
        copy_cell_style(source_sheet.cell(1, style_source_column), orientation_header_cell)
        source_sheet.column_dimensions[get_column_letter(orientation_column)].width = source_sheet.column_dimensions[
            get_column_letter(style_source_column)
        ].width
        for row in range(2, source_sheet.max_row + 1):
            copy_cell_style(source_sheet.cell(row, style_source_column), source_sheet.cell(row, orientation_column))

        tagged_products = 0
        changed_cells = 0
        product_attribute_payload_changes = 0
        declarations_added = {slug: 0 for slug in ATTRIBUTE_SLUGS.values()}
        missing_skus: list[str] = []
        active_skus: set[str] = set()

        for values in analysis_sheet.iter_rows(min_row=2, values_only=True):
            variant_tag = clean(values[analysis_columns["Wariant: tag"]])
            if not variant_tag:
                continue
            sku = clean(values[analysis_columns["SKU"]])
            source_row = source_rows.get(sku)
            if not source_row:
                missing_skus.append(sku)
                continue
            if clean(values[analysis_columns["Wariant: decyzja"]]).startswith("POMINIĘTY"):
                raise ValueError(f"Stary duplikat otrzymał tag wariantowy: {sku}")
            active_skus.add(sku)
            tagged_products += 1

            for header in EDITED_HEADERS:
                new_value = values[analysis_columns[header]]
                cell = source_sheet.cell(source_row, source_columns[header])
                if cell.value != new_value:
                    cell.value = new_value
                    changed_cells += 1

            orientation_value = values[analysis_columns[ORIENTATION_HEADER]]
            if clean(orientation_value):
                source_sheet.cell(source_row, orientation_column).value = orientation_value
                changed_cells += 1

            product_attributes_cell = source_sheet.cell(source_row, source_columns["Product Attributes"])
            payload = parse_product_attributes(product_attributes_cell.value)
            payload_changed = False
            for header, slug in ATTRIBUTE_SLUGS.items():
                if header == ORIENTATION_HEADER:
                    value = orientation_value
                else:
                    value = values[analysis_columns[header]]
                if clean(value) and add_attribute_declaration(payload, slug):
                    declarations_added[slug] += 1
                    payload_changed = True
            if payload_changed:
                product_attributes_cell.value = php_serialize(payload)
                product_attribute_payload_changes += 1

        if missing_skus:
            raise ValueError(f"Brakujące SKU w eksporcie źródłowym ({len(missing_skus)}): {missing_skus[:20]}")

        source_book.save(output_path)
    finally:
        source_book.close()
        analysis_book.close()

    return {
        "source_sheet": source_book.sheetnames[0],
        "original_rows": original_row_count,
        "original_columns": original_column_count,
        "output_columns": original_column_count + 1,
        "orientation_column": orientation_column,
        "tagged_products": tagged_products,
        "changed_cells": changed_cells,
        "product_attribute_payload_changes": product_attribute_payload_changes,
        "declarations_added": declarations_added,
        "active_skus": sorted(active_skus),
    }


def validate_file(
    source_path: Path,
    analysis_path: Path,
    output_path: Path,
    metrics: dict[str, Any],
) -> dict[str, Any]:
    source_book = load_workbook(source_path, read_only=True, data_only=False)
    output_book = load_workbook(output_path, read_only=True, data_only=False)
    analysis_book = load_workbook(analysis_path, read_only=True, data_only=False)
    try:
        source_sheet = source_book.active
        output_sheet = output_book.active
        analysis_sheet = analysis_book["Produkty warianty"]
        source_headers = [cell.value for cell in source_sheet[1]]
        output_headers = [cell.value for cell in output_sheet[1]]
        analysis_headers = [cell.value for cell in analysis_sheet[1]]
        source_columns = {value: index for index, value in enumerate(source_headers)}
        output_columns = {value: index for index, value in enumerate(output_headers)}
        analysis_columns = {value: index for index, value in enumerate(analysis_headers)}

        analysis_rows = {
            clean(row[analysis_columns["SKU"]]): row
            for row in analysis_sheet.iter_rows(min_row=2, values_only=True)
        }
        active_skus = set(metrics["active_skus"])
        allowed_original_headers = set(EDITED_HEADERS) | {"Product Attributes"}
        unexpected_changes: list[dict[str, Any]] = []
        active_values_match = True
        declarations_complete = True
        old_duplicates_unchanged = True
        formula_errors: list[dict[str, str]] = []

        for source_row, output_row in zip(
            source_sheet.iter_rows(min_row=2, values_only=True),
            output_sheet.iter_rows(min_row=2, values_only=True),
        ):
            sku = clean(source_row[source_columns["SKU"]])
            for index, (before, after) in enumerate(zip(source_row, output_row[: len(source_headers)])):
                header = source_headers[index]
                if before != after and header not in allowed_original_headers:
                    unexpected_changes.append({"sku": sku, "header": header, "before": before, "after": after})
            if sku in active_skus:
                analysis_row = analysis_rows[sku]
                for header in EDITED_HEADERS:
                    active_values_match = active_values_match and (
                        output_row[output_columns[header]] == analysis_row[analysis_columns[header]]
                    )
                active_values_match = active_values_match and (
                    output_row[output_columns[ORIENTATION_HEADER]]
                    == analysis_row[analysis_columns[ORIENTATION_HEADER]]
                )
                payload = parse_product_attributes(output_row[output_columns["Product Attributes"]])
                for header, slug in ATTRIBUTE_SLUGS.items():
                    value = output_row[output_columns[header]]
                    if clean(value):
                        declarations_complete = declarations_complete and attribute_is_declared(payload, slug)
            else:
                old_duplicates_unchanged = old_duplicates_unchanged and (
                    output_row[: len(source_headers)] == source_row
                    and not clean(output_row[output_columns[ORIENTATION_HEADER]])
                )
            for value in output_row:
                if isinstance(value, str) and any(error in value for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")):
                    formula_errors.append({"sku": sku, "value": value})

        validation = {
            "single_sheet_preserved": output_book.sheetnames == source_book.sheetnames,
            "row_count_preserved": output_sheet.max_row == source_sheet.max_row,
            "original_headers_preserved": output_headers[: len(source_headers)] == source_headers,
            "only_orientation_header_added": output_headers[len(source_headers) :] == [ORIENTATION_HEADER],
            "no_unexpected_original_column_changes": not unexpected_changes,
            "all_active_values_match_analysis": active_values_match,
            "all_nonempty_attributes_declared": declarations_complete,
            "all_nonvariant_products_unchanged": old_duplicates_unchanged,
            "up_sells_unchanged": all(
                source_row[source_columns["Up-Sells"]] == output_row[output_columns["Up-Sells"]]
                for source_row, output_row in zip(
                    source_sheet.iter_rows(min_row=2, values_only=True),
                    output_sheet.iter_rows(min_row=2, values_only=True),
                )
            ),
            "cross_sells_unchanged": all(
                source_row[source_columns["Cross-Sells"]] == output_row[output_columns["Cross-Sells"]]
                for source_row, output_row in zip(
                    source_sheet.iter_rows(min_row=2, values_only=True),
                    output_sheet.iter_rows(min_row=2, values_only=True),
                )
            ),
            "no_formula_error_strings": not formula_errors,
        }
        validation["passed"] = all(validation.values())
        if not validation["passed"]:
            raise AssertionError(json.dumps(validation, ensure_ascii=False, indent=2))
        return validation
    finally:
        source_book.close()
        output_book.close()
        analysis_book.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Buduje XLSX WP All Import na bazie pełnego eksportu WP All Export")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--analysis", type=Path, default=DEFAULT_ANALYSIS)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()

    metrics = build_file(args.source, args.analysis, args.output)
    validation = validate_file(args.source, args.analysis, args.output, metrics)
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "workflow": "WP All Export -> pełny XLSX -> WP All Import",
        "source_file": str(args.source.resolve()),
        "source_sha256": sha256(args.source),
        "analysis_file": str(args.analysis.resolve()),
        "analysis_sha256": sha256(args.analysis),
        "output_file": str(args.output.resolve()),
        "output_sha256": sha256(args.output),
        **{key: value for key, value in metrics.items() if key != "active_skus"},
        "validation": validation,
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
