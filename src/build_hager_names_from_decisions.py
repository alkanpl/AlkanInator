"""Build a Hager naming workbook from the reviewed decision workbook.

The archived decision workbook and the prepared Hager product workbook are
treated as immutable sources. The result is a new control workbook in output/.
"""

from __future__ import annotations

from collections import Counter
from copy import copy
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DECISIONS_PATH = (
    ROOT
    / "archived_input_files"
    / "input_2026-08-07"
    / "Hager_plan_kategorii_atrybutow_i_audyt_uzupelniony_decyzje.xlsx"
)
HAGER_PATH = ROOT / "output" / "Aparatura-Modulowa-Hager_nazwy_atrybuty_v2.xlsx"
OUTPUT_PATH = ROOT / "output" / "Hager_nazwy_po_decyzjach_2026-08-07.xlsx"
REPORT_PATH = ROOT / "reports" / "hager_names_decisions" / "validation.txt"


NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
GREEN = "548235"
LIGHT_GREEN = "E2F0D9"
AMBER = "BF8F00"
LIGHT_AMBER = "FFF2CC"
RED = "C00000"
LIGHT_RED = "FCE4D6"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
GRID = "D9E1F2"

THIN_BORDER = Border(
    left=Side(style="thin", color=GRID),
    right=Side(style="thin", color=GRID),
    top=Side(style="thin", color=GRID),
    bottom=Side(style="thin", color=GRID),
)


def text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def sheet_starting_with(workbook, prefix: str):
    return workbook[next(name for name in workbook.sheetnames if name.startswith(prefix))]


def headers_by_name(ws, row: int) -> dict[str, int]:
    return {
        text(ws.cell(row, column).value): column
        for column in range(1, ws.max_column + 1)
        if ws.cell(row, column).value not in (None, "")
    }


def set_title(ws, title: str, subtitle: str, last_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, title)
    ws.cell(1, 1).font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 34

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).font = Font(name="Aptos", size=10, italic=True, color=GRAY)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 34


def style_header(ws, row: int, first_col: int, last_col: int) -> None:
    for column in range(first_col, last_col + 1):
        cell = ws.cell(row, column)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 34


def style_body(ws, first_row: int, last_row: int, first_col: int, last_col: int) -> None:
    for row in range(first_row, last_row + 1):
        fill = PatternFill("solid", fgColor=WHITE if row % 2 == 0 else PALE_BLUE)
        for column in range(first_col, last_col + 1):
            cell = ws.cell(row, column)
            cell.font = Font(name="Aptos", size=9, color="000000")
            cell.fill = fill
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def phase_from_layout(layout: object) -> int | None:
    value = text(layout).upper().replace(" ", "")
    if not value:
        return None
    if value in {"1P", "1P+N", "2P"}:
        return 1
    if value in {"3P", "3P+N", "4P", "3+1/4P", "4P/3P"}:
        return 3
    if value.startswith("1P") or value.startswith("2P"):
        return 1
    if value.startswith("3P") or value.startswith("4P") or value.startswith("3+1"):
        return 3
    raise ValueError(f"Nieznany układ biegunów do mapowania liczby faz: {layout!r}")


def add_network_generator_phrase(title: str) -> str:
    if "sieć-agregat" in title.lower():
        return title
    if "I-0-II" not in title:
        raise ValueError(f"Brak I-0-II w tytule przełącznika: {title}")
    return title.replace("I-0-II", "I-0-II (sieć-agregat)", 1)


def read_decisions(decisions_wb) -> tuple[dict[str, dict[str, str]], dict[str, dict[str, str]], list[dict[str, str]]]:
    key_ws = sheet_starting_with(decisions_wb, "Decyzje kluczowe")
    key_decisions: dict[str, dict[str, str]] = {}
    for row in range(6, key_ws.max_row + 1):
        decision_id = text(key_ws.cell(row, 1).value)
        if not decision_id:
            continue
        key_decisions[decision_id] = {
            "choice": text(key_ws.cell(row, 10).value),
            "target": text(key_ws.cell(row, 11).value),
            "notes": text(key_ws.cell(row, 12).value),
        }

    category_ws = sheet_starting_with(decisions_wb, "Decyzje kategorii")
    categories: dict[str, dict[str, str]] = {}
    for row in range(6, category_ws.max_row + 1):
        category_id = text(category_ws.cell(row, 1).value)
        if not category_id:
            continue
        base_path = text(category_ws.cell(row, 2).value)
        recommendation = text(category_ws.cell(row, 7).value)
        recommended_path = text(category_ws.cell(row, 8).value)
        user_choice = text(category_ws.cell(row, 11).value)
        user_path = text(category_ws.cell(row, 12).value)
        choice = user_choice or recommendation
        if choice == "OSOBNA KATEGORIA":
            resolved_path = user_path or base_path
        elif choice in {"POŁĄCZYĆ", "ZOSTAWIĆ W NADRZĘDNEJ"}:
            resolved_path = user_path or recommended_path or base_path
        else:
            resolved_path = user_path or recommended_path or base_path
        categories[category_id] = {
            "base_path": base_path,
            "choice": choice,
            "resolved_path": resolved_path,
            "source": "Decyzje kategorii — wybór użytkownika" if user_choice else "Decyzje kategorii — rekomendacja bez wyboru",
        }

    conflicts: list[dict[str, str]] = []

    def override_group(decision_id: str, category_ids: list[str], preferred_target_id: str) -> None:
        decision = key_decisions.get(decision_id, {})
        if decision.get("choice") != "POŁĄCZYĆ":
            return
        preferred = categories[preferred_target_id]["resolved_path"]
        target = preferred or decision.get("target", "")
        previous = {category_id: categories[category_id]["resolved_path"] for category_id in category_ids}
        if len(set(previous.values())) > 1:
            conflicts.append(
                {
                    "area": "Kategoria",
                    "scope": ", ".join(category_ids),
                    "conflict": "Decyzja grupowa mówi POŁĄCZYĆ, ale decyzje wierszowe prowadziły do różnych kategorii.",
                    "resolution": target,
                    "source": f"Decyzje kluczowe: {decision_id} + jawna ścieżka z Decyzji kategorii",
                    "status": "ROZWIĄZANY",
                }
            )
        for category_id in category_ids:
            categories[category_id]["resolved_path"] = target
            categories[category_id]["source"] = f"Decyzja grupowa {decision_id} ma pierwszeństwo"

    override_group("RCCB_RCBO", ["rccb", "rcbo"], "rcbo")
    override_group("ROZLACZNIKI", ["isolators", "fuse_holders"], "fuse_holders")
    override_group("PRZEKAZNIKI", ["bistable_relays", "installation_relays"], "installation_relays")
    override_group(
        "STEROWANIE_SYGNALIZACJA",
        ["indicators", "control_buttons", "acoustic_signaling"],
        "control_buttons",
    )

    return key_decisions, categories, conflicts


def read_product_category_ids(decisions_wb) -> dict[str, str]:
    category_ws = decisions_wb["Kategorie"]
    path_to_id = {
        text(category_ws.cell(row, 2).value): text(category_ws.cell(row, 1).value)
        for row in range(6, category_ws.max_row + 1)
        if category_ws.cell(row, 1).value
    }
    products_ws = decisions_wb["Produkty i kategorie"]
    result: dict[str, str] = {}
    for row in range(6, products_ws.max_row + 1):
        sku = text(products_ws.cell(row, 1).value)
        path = text(products_ws.cell(row, 4).value)
        if sku:
            result[sku] = path_to_id[path]
    return result


def build_summary_sheet(ws, record_count: int, changed_names: int, phase_count: int, final_categories: int, conflict_count: int, check_count: int) -> None:
    set_title(
        ws,
        "Hager — nazwy po decyzjach",
        "Plik kontrolny z nazwami, kategoriami i korektami konfliktów. Źródła nie zostały nadpisane; arkusz nie jest jeszcze plikiem importowym Baselinker/WooCommerce.",
        8,
    )
    ws.sheet_view.showGridLines = False

    cards = [
        ("A4", "Produkty", record_count),
        ("C4", "Zmienione nazwy", changed_names),
        ("E4", "Uzupełniona liczba faz", phase_count),
        ("G4", "Kategorie końcowe", final_categories),
        ("A7", "Konflikty kategorii", conflict_count),
        ("C7", "Do sprawdzenia", check_count),
        ("E7", "Brakujące nazwy", 0),
        ("G7", "Duplikaty nazw", 0),
    ]
    for anchor, label, value in cards:
        column = ws[anchor].column
        row = ws[anchor].row
        ws.merge_cells(start_row=row, start_column=column, end_row=row, end_column=column + 1)
        ws.cell(row, column, label)
        ws.cell(row, column).font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        ws.cell(row, column).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(row, column).alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=row + 1, start_column=column, end_row=row + 1, end_column=column + 1)
        ws.cell(row + 1, column, value)
        ws.cell(row + 1, column).font = Font(name="Aptos Display", size=18, bold=True, color=NAVY)
        ws.cell(row + 1, column).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row + 1, column).alignment = Alignment(horizontal="center", vertical="center")

    ws["A11"] = "Zastosowane decyzje"
    ws["A11"].font = Font(name="Aptos", size=12, bold=True, color=NAVY)
    headers = ["Obszar", "Decyzja", "Skutek"]
    for column, value in enumerate(headers, 1):
        ws.cell(12, column, value)
    ws.merge_cells("C12:H12")
    style_header(ws, 12, 1, 8)
    rows = [
        ("Małe kategorie", "Indywidualnie", "Zastosowano wybory z arkusza kategorii; brak wyboru = rekomendacja robocza."),
        ("RCCB + RCBO", "Połączyć", "Wspólna kategoria „Wyłączniki różnicowoprądowe”."),
        ("Rozłączniki", "Połączyć", "Wspólna kategoria „Rozłączniki i aparatura bezpiecznikowa”."),
        ("Przekaźniki", "Połączyć", "Wspólna kategoria „Przekaźniki modułowe”."),
        ("Sterowanie i sygnalizacja", "Połączyć", "Lampki, przyciski i dzwonek w jednej kategorii."),
        ("Bieguny / fazy", "Zmienić na Liczba faz", "Wartości 1 albo 3 wyprowadzone z układu Hager."),
        ("Przełączniki I-0-II", "Dodać sieć-agregat", "Fraza dodana do 7 nazw przełączników; nie została dodana do SPD."),
        ("Icn a prąd znamionowy", "Nie scalać", "Pozostają dwoma różnymi parametrami i jednostkami."),
    ]
    for row, values in enumerate(rows, 13):
        ws.cell(row, 1, values[0])
        ws.cell(row, 2, values[1])
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        ws.cell(row, 3, values[2])
        for column in range(1, 9):
            cell = ws.cell(row, column)
            cell.fill = PatternFill("solid", fgColor=WHITE if row % 2 else PALE_BLUE)
            cell.border = THIN_BORDER
            cell.font = Font(name="Aptos", size=9, color="000000")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30

    ws["A23"] = "Ważne"
    ws["A23"].font = Font(name="Aptos", size=12, bold=True, color=NAVY)
    notes = [
        "• Nazwy zachowują techniczny układ biegunów z pliku Hager (np. 1P, 3P+N), natomiast osobny atrybut „Liczba faz” ma wartości wyłącznie 1 albo 3.",
        "• Korekty z arkusza „Problemy danych” zastosowano w pełnym arkuszu produktowym oraz opisano w „Konfliktach rozwiązanych”.",
        "• Dwa produkty oznaczone wcześniej jako DO_SPRAWDZENIA pozostają do ręcznej kontroli; nie blokują przygotowania pozostałych nazw.",
    ]
    for row, note in enumerate(notes, 24):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        ws.cell(row, 1, note)
        ws.cell(row, 1).font = Font(name="Aptos", size=9, color=GRAY)
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 28

    for column, width in enumerate([20, 24, 20, 18, 22, 22, 20, 20], 1):
        ws.column_dimensions[get_column_letter(column)].width = width


def build_names_sheet(ws, records: list[dict[str, object]]) -> None:
    set_title(
        ws,
        "Nazwy Hager po decyzjach",
        "153 produkty. Zielony kolor oznacza nazwę zmienioną decyzją użytkownika; żółty status wymaga ręcznej kontroli.",
        17,
    )
    headers = [
        "Lp.",
        "SKU",
        "Kod",
        "Nazwa pierwotna",
        "Nazwa Hager przed decyzjami",
        "Nazwa po decyzjach",
        "Zmiana nazwy",
        "Typ produktu",
        "Kategoria źródłowa",
        "Kategoria po decyzjach",
        "Źródło decyzji kategorii",
        "Układ biegunów Hager",
        "Liczba faz",
        "Korekty zastosowane",
        "Status końcowy",
        "Źródło / strona",
        "Adres źródła",
    ]
    for column, value in enumerate(headers, 1):
        ws.cell(5, column, value)
    style_header(ws, 5, 1, len(headers))
    for row, record in enumerate(records, 6):
        values = [
            record["lp"],
            record["sku"],
            record["code"],
            record["original_title"],
            record["hager_title"],
            record["final_title"],
            record["title_change"],
            record["family"],
            record["source_category"],
            record["final_category"],
            record["category_source"],
            record["source_layout"],
            record["phase_count"],
            record["corrections"],
            record["final_status"],
            record["source_page"],
            record["url"],
        ]
        for column, value in enumerate(values, 1):
            ws.cell(row, column, value)
    end_row = 5 + len(records)
    style_body(ws, 6, end_row, 1, len(headers))
    for row in range(6, end_row + 1):
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row, 13).alignment = Alignment(horizontal="center", vertical="top")
        if ws.cell(row, 7).value == "ZMIENIONA":
            ws.cell(row, 6).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
            ws.cell(row, 6).font = Font(name="Aptos", size=9, bold=True, color=GREEN)
            ws.cell(row, 7).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
            ws.cell(row, 7).font = Font(name="Aptos", size=9, bold=True, color=GREEN)
        if ws.cell(row, 15).value != "OK":
            ws.cell(row, 15).fill = PatternFill("solid", fgColor=LIGHT_AMBER)
            ws.cell(row, 15).font = Font(name="Aptos", size=9, bold=True, color=AMBER)
        ws.row_dimensions[row].height = 42
    ws.freeze_panes = "D6"
    ws.auto_filter.ref = f"A5:Q{end_row}"
    ws.sheet_view.showGridLines = False
    widths = [8, 18, 16, 48, 56, 60, 17, 38, 48, 52, 34, 20, 12, 52, 18, 48, 44]
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width


def build_conflicts_sheet(ws, conflicts: list[dict[str, str]]) -> None:
    set_title(
        ws,
        "Konflikty rozwiązane",
        "Rozbieżności decyzji kategorii i problemy wartości atrybutów rozwiązane na podstawie decyzji użytkownika oraz przygotowanego pliku Hager.",
        6,
    )
    headers = ["Obszar", "Zakres / SKU", "Konflikt", "Zastosowane rozwiązanie", "Źródło", "Status"]
    for column, value in enumerate(headers, 1):
        ws.cell(5, column, value)
    style_header(ws, 5, 1, len(headers))
    for row, conflict in enumerate(conflicts, 6):
        values = [
            conflict["area"],
            conflict["scope"],
            conflict["conflict"],
            conflict["resolution"],
            conflict["source"],
            conflict["status"],
        ]
        for column, value in enumerate(values, 1):
            ws.cell(row, column, value)
    end_row = 5 + len(conflicts)
    style_body(ws, 6, end_row, 1, len(headers))
    for row in range(6, end_row + 1):
        ws.cell(row, 4).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
        ws.cell(row, 4).font = Font(name="Aptos", size=9, bold=True, color=GREEN)
        ws.cell(row, 6).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
        ws.cell(row, 6).font = Font(name="Aptos", size=9, bold=True, color=GREEN)
        ws.row_dimensions[row].height = 64
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:F{end_row}"
    ws.sheet_view.showGridLines = False
    for column, width in enumerate([24, 32, 58, 58, 50, 18], 1):
        ws.column_dimensions[get_column_letter(column)].width = width


def copy_header_style(ws, source_column: int, target_column: int) -> None:
    source = ws.cell(1, source_column)
    target = ws.cell(1, target_column)
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)
    target.number_format = source.number_format


def build_workbook() -> Path:
    decisions_wb = load_workbook(DECISIONS_PATH, data_only=False)
    hager_wb = load_workbook(HAGER_PATH, data_only=False)
    key_decisions, categories, category_conflicts = read_decisions(decisions_wb)
    product_category_ids = read_product_category_ids(decisions_wb)

    control_ws = hager_wb["Kontrola"]
    control_headers = headers_by_name(control_ws, 5)
    product_ws = hager_wb["Worksheet"]
    product_headers = headers_by_name(product_ws, 1)

    sku_column = product_headers["SKU"]
    product_rows = {
        text(product_ws.cell(row, sku_column).value): row
        for row in range(2, product_ws.max_row + 1)
        if product_ws.cell(row, sku_column).value
    }

    required_columns = [
        "Title",
        "Atrybut Produktu: Liczba Faz",
        "Atrybut Produktu: Układ Biegunów",
        "Atrybut Produktu: Maksymalne Napięcie Ciągłej Pracy Uc [V]",
        "Atrybut Produktu: Napięcie Cewki [V]",
        "Atrybut Produktu: Konfiguracja Styków",
        "Atrybut Produktu: Typ Produktu",
        "Atrybut Produktu: Typ Wyłącznika",
        "Atrybut Produktu: Typ Prądu Różnicowego",
    ]
    missing_columns = [column for column in required_columns if column not in product_headers]
    if missing_columns:
        raise KeyError(f"Brak kolumn w pliku Hager: {missing_columns}")

    append_headers = [
        "Kategoria po decyzjach",
        "Źródło decyzji kategorii",
        "Styk pomocniczy po korekcie",
        "Korekty zastosowane",
    ]
    append_start = product_ws.max_column + 1
    for offset, header in enumerate(append_headers):
        column = append_start + offset
        product_ws.cell(1, column, header)
        copy_header_style(product_ws, append_start - 1, column)
        product_ws.column_dimensions[get_column_letter(column)].width = [52, 34, 24, 58][offset]
    append_indices = {header: append_start + index for index, header in enumerate(append_headers)}

    source_category_by_sku: dict[str, str] = {}
    source_products_ws = decisions_wb["Produkty i kategorie"]
    for row in range(6, source_products_ws.max_row + 1):
        sku = text(source_products_ws.cell(row, 1).value)
        if sku:
            source_category_by_sku[sku] = text(source_products_ws.cell(row, 4).value)

    records: list[dict[str, object]] = []
    all_conflicts = list(category_conflicts)

    data_conflicts = [
        {
            "area": "Atrybut",
            "scope": "SPA931/HAG",
            "conflict": "Uc miało wartość 255 V.",
            "resolution": "Maksymalne napięcie ciągłej pracy Uc = 335 V.",
            "source": "Problemy danych + karta Hager wskazana w pliku decyzji",
            "status": "ROZWIĄZANY",
        },
        {
            "area": "Atrybut",
            "scope": "SPA931/HAG",
            "conflict": "Układ biegunów 4P opisywał liczbę biegunów, nie układ faz i N.",
            "resolution": "Układ biegunów = 3P+N; Liczba faz = 3.",
            "source": "Problemy danych + przygotowany plik Hager",
            "status": "ROZWIĄZANY",
        },
        {
            "area": "Atrybut",
            "scope": "SPB413/HAG",
            "conflict": "Wartość 3+1 / 4P mieszała schemat SPD i liczbę biegunów.",
            "resolution": "Układ biegunów = 3P+N; Liczba faz = 3.",
            "source": "Problemy danych + przygotowany plik Hager",
            "status": "ROZWIĄZANY",
        },
        {
            "area": "Atrybut",
            "scope": "EPN510/HAG; EPN520/HAG",
            "conflict": "Napięcie cewki 230 V pomijało obsługę 110 V DC.",
            "resolution": "Napięcie cewki = 230 V AC / 110 V DC.",
            "source": "Problemy danych + przygotowany plik Hager",
            "status": "ROZWIĄZANY",
        },
        {
            "area": "Atrybut",
            "scope": "SA463/HAG; SA480/HAG",
            "conflict": "„Styk pomocniczy” był wpisany jako konfiguracja styków.",
            "resolution": "Konfiguracja styków wyczyszczona; Styk pomocniczy = Tak.",
            "source": "Problemy danych + przygotowany plik Hager",
            "status": "ROZWIĄZANY",
        },
        {
            "area": "Model atrybutów",
            "scope": "Produkty z Typem wyłącznika opisującym rodzinę",
            "conflict": "Typ wyłącznika mieszał rodzinę produktu z typem prądu RCD.",
            "resolution": "Rodzina pozostaje w Typie produktu; Typ wyłącznika zawiera wyłącznie AC/A/A-HI albo jest pusty.",
            "source": "Problemy danych + wartości Hager",
            "status": "ROZWIĄZANY",
        },
        {
            "area": "Nazwa",
            "scope": "7 przełączników I-0-II",
            "conflict": "W uwagach użytkownik wymaga frazy „sieć-agregat”; identyczna dopiska przy SPD jest sprzeczna z typem produktu.",
            "resolution": "Dodano „(sieć-agregat)” wyłącznie do przełączników I-0-II; tytuły SPD zachowano zgodnie z plikiem Hager.",
            "source": "Składanie tytułów — uwagi użytkownika + Typ produktu Hager",
            "status": "ROZWIĄZANY",
        },
    ]
    all_conflicts.extend(data_conflicts)

    # Apply model-wide attribute cleanup before building the records.
    type_column = product_headers["Atrybut Produktu: Typ Produktu"]
    old_switch_type_column = product_headers["Atrybut Produktu: Typ Wyłącznika"]
    rcd_type_column = product_headers["Atrybut Produktu: Typ Prądu Różnicowego"]
    for row in range(2, product_ws.max_row + 1):
        old_value = text(product_ws.cell(row, old_switch_type_column).value)
        rcd_type = text(product_ws.cell(row, rcd_type_column).value)
        if rcd_type:
            product_ws.cell(row, old_switch_type_column, rcd_type)
        elif old_value in {"Nadprądowy", "Różnicowoprądowy", "Różnicowo-nadprądowy"}:
            product_ws.cell(row, old_switch_type_column).value = None

    for row in range(6, control_ws.max_row + 1):
        sku = text(control_ws.cell(row, control_headers["SKU"]).value)
        if not sku:
            continue
        if sku not in product_rows:
            raise KeyError(f"Brak SKU z Kontroli w arkuszu Worksheet: {sku}")
        product_row = product_rows[sku]
        family = text(control_ws.cell(row, control_headers["Rodzina"]).value)
        hager_title = text(control_ws.cell(row, control_headers["Nazwa po zmianie"]).value)
        final_title = hager_title
        corrections: list[str] = []

        if family in {"Przełącznik zasilania modułowy", "Przełącznik instalacyjny modułowy"}:
            final_title = add_network_generator_phrase(final_title)
            corrections.append("Dodano frazę „sieć-agregat” do tytułu")

        source_layout = text(product_ws.cell(product_row, product_headers["Atrybut Produktu: Układ Biegunów"]).value)
        corrected_layout = source_layout
        if sku == "SPA931/HAG":
            product_ws.cell(
                product_row,
                product_headers["Atrybut Produktu: Maksymalne Napięcie Ciągłej Pracy Uc [V]"],
                335,
            )
            corrected_layout = "3P+N"
            corrections.extend(["Uc: 255 → 335 V", "Układ biegunów: 4P → 3P+N"])
        elif sku == "SPB413/HAG":
            corrected_layout = "3P+N"
            corrections.append("Układ biegunów: 3+1 / 4P → 3P+N")

        if corrected_layout != source_layout:
            product_ws.cell(product_row, product_headers["Atrybut Produktu: Układ Biegunów"], corrected_layout)

        phase_count = phase_from_layout(corrected_layout)
        if phase_count is not None:
            product_ws.cell(product_row, product_headers["Atrybut Produktu: Liczba Faz"], phase_count)

        if sku in {"EPN510/HAG", "EPN520/HAG"}:
            product_ws.cell(
                product_row,
                product_headers["Atrybut Produktu: Napięcie Cewki [V]"],
                "230 V AC / 110 V DC",
            )
            corrections.append("Uzupełniono napięcie cewki: 230 V AC / 110 V DC")

        auxiliary_contact = None
        if sku in {"SA463/HAG", "SA480/HAG"}:
            product_ws.cell(
                product_row,
                product_headers["Atrybut Produktu: Konfiguracja Styków"],
            ).value = None
            auxiliary_contact = "Tak"
            corrections.append("Styk pomocniczy przeniesiono poza Konfigurację styków")

        category_id = product_category_ids[sku]
        final_category = categories[category_id]["resolved_path"]
        category_source = categories[category_id]["source"]
        product_ws.cell(product_row, product_headers["Title"], final_title)
        product_ws.cell(product_row, type_column, family)
        product_ws.cell(product_row, append_indices["Kategoria po decyzjach"], final_category)
        product_ws.cell(product_row, append_indices["Źródło decyzji kategorii"], category_source)
        product_ws.cell(product_row, append_indices["Styk pomocniczy po korekcie"], auxiliary_contact)
        product_ws.cell(product_row, append_indices["Korekty zastosowane"], "; ".join(corrections) or "Brak")

        control_ws.cell(row, control_headers["Nazwa po zmianie"], final_title)
        if corrected_layout:
            control_ws.cell(row, control_headers["Układ biegunów"], corrected_layout)

        original_status = text(control_ws.cell(row, control_headers["Status"]).value)
        final_status = original_status if original_status != "OK" else "OK"
        records.append(
            {
                "lp": control_ws.cell(row, control_headers["Lp."]).value,
                "sku": sku,
                "code": text(control_ws.cell(row, control_headers["Kod"]).value),
                "original_title": text(control_ws.cell(row, control_headers["Nazwa pierwotna"]).value),
                "hager_title": hager_title,
                "final_title": final_title,
                "title_change": "ZMIENIONA" if final_title != hager_title else "BEZ ZMIANY",
                "family": family,
                "source_category": source_category_by_sku[sku],
                "final_category": final_category,
                "category_source": category_source,
                "source_layout": source_layout,
                "phase_count": phase_count,
                "corrections": "; ".join(corrections) or "Brak",
                "final_status": final_status,
                "source_page": text(control_ws.cell(row, control_headers["Źródło"]).value)
                + (" | str. " + text(control_ws.cell(row, control_headers["Strona katalogu"]).value) if control_ws.cell(row, control_headers["Strona katalogu"]).value else ""),
                "url": text(control_ws.cell(row, control_headers["Adres źródła"]).value),
            }
        )

    # Extend the existing control sheet with the final decisions.
    control_append_headers = ["Kategoria po decyzjach", "Liczba faz", "Korekty zastosowane"]
    control_start = control_ws.max_column + 1
    for offset, header in enumerate(control_append_headers):
        column = control_start + offset
        control_ws.cell(5, column, header)
        source_cell = control_ws.cell(5, control_start - 1)
        target_cell = control_ws.cell(5, column)
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        control_ws.column_dimensions[get_column_letter(column)].width = [52, 14, 58][offset]
    record_by_sku = {str(record["sku"]): record for record in records}
    for row in range(6, control_ws.max_row + 1):
        sku = text(control_ws.cell(row, control_headers["SKU"]).value)
        if not sku:
            continue
        record = record_by_sku[sku]
        control_ws.cell(row, control_start, record["final_category"])
        control_ws.cell(row, control_start + 1, record["phase_count"])
        control_ws.cell(row, control_start + 2, record["corrections"])
        for column in range(control_start, control_start + 3):
            control_ws.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)

    if len(records) != 153:
        raise AssertionError(f"Oczekiwano 153 produktów, otrzymano {len(records)}")
    titles = [str(record["final_title"]) for record in records]
    if any(not title for title in titles):
        raise AssertionError("Wykryto pustą nazwę")
    if len(set(titles)) != len(titles):
        duplicates = [title for title, count in Counter(titles).items() if count > 1]
        raise AssertionError(f"Powielone nazwy: {duplicates}")
    for record in records:
        if str(record["code"]) not in str(record["final_title"]):
            raise AssertionError(f"Brak kodu w nazwie: {record['sku']}")
        if "Hager" not in str(record["final_title"]):
            raise AssertionError(f"Brak producenta w nazwie: {record['sku']}")
        if not record["final_category"]:
            raise AssertionError(f"Brak kategorii po decyzjach: {record['sku']}")

    changed_names = sum(record["title_change"] == "ZMIENIONA" for record in records)
    phase_count = sum(record["phase_count"] is not None for record in records)
    final_category_count = len({str(record["final_category"]) for record in records})
    check_count = sum(record["final_status"] != "OK" for record in records)

    summary_ws = hager_wb.create_sheet("Podsumowanie decyzji", 0)
    names_ws = hager_wb.create_sheet("Nazwy po decyzjach", 1)
    conflicts_ws = hager_wb.create_sheet("Konflikty rozwiązane", 2)
    build_summary_sheet(
        summary_ws,
        len(records),
        changed_names,
        phase_count,
        final_category_count,
        len(category_conflicts),
        check_count,
    )
    build_names_sheet(names_ws, records)
    build_conflicts_sheet(conflicts_ws, all_conflicts)
    for ws in [summary_ws, names_ws, conflicts_ws]:
        ws.sheet_properties.tabColor = NAVY
    control_ws.sheet_properties.tabColor = BLUE
    product_ws.sheet_properties.tabColor = BLUE

    hager_wb.calculation.fullCalcOnLoad = True
    hager_wb.calculation.forceFullCalc = True
    hager_wb.calculation.calcMode = "auto"
    hager_wb.active = 0
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    hager_wb.save(OUTPUT_PATH)

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(
        "\n".join(
            [
                f"input_decisions={DECISIONS_PATH}",
                f"input_hager={HAGER_PATH}",
                f"output={OUTPUT_PATH}",
                f"products={len(records)}",
                f"changed_names={changed_names}",
                f"phase_values={phase_count}",
                f"final_categories={final_category_count}",
                f"category_conflicts_resolved={len(category_conflicts)}",
                f"products_to_check={check_count}",
                "blank_names=0",
                "duplicate_names=0",
                "missing_codes_in_names=0",
                "missing_manufacturer_in_names=0",
                "status=OK",
            ]
        ),
        encoding="utf-8",
    )
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_workbook())
