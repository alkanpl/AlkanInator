from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from update_hager_berker_series_categories import (
    collect_workbook_structure,
    header_map,
    normalize_series,
    scan_mojibake,
)


DEFAULT_INPUT = Path(
    "output/Hager-Berker-Gniazdka_SERIE_KATEGORIE_POPRAWIONE_2026-08-27.xlsx"
)
DEFAULT_OUTPUT = Path(
    "output/Hager-Berker-Gniazdka_SERIE_KATEGORIE_TYTULY_SKROTY_POPRAWIONE_2026-08-27.xlsx"
)
DEFAULT_REPORT = Path(
    "reports/Hager-Berker-zmiany-tytulow-serie-2026-08-27.json"
)

STORE_SHEET = "Worksheet"
HEADER_ROW = 1
TARGET_COLUMN = "Title"

SERIES_PATTERNS = [
    r"(?<!\w)Lumina\s*(?:2|Soul|Intense|Passion)?(?!\w)",
    r"(?<!\w)Serie\s+1930\s+Rosenthal(?!\w)",
    r"(?<!\w)Serie\s+1930(?!\w)",
    r"(?<!\w)(?:Berker\s+)?Serie\s+R\.classic(?!\w)",
    r"(?<!\w)R\.classic(?!\w)",
    r"(?<!\w)(?:Berker\s+)?Serie\s+Glas(?!\w)",
    r"(?<!\w)Glasserie(?!\w)",
    r"(?<!\w)Glas(?!\w)",
    r"(?<!\w)one\.platform(?!\w)",
    r"(?<!\w)Integro\s+Flow(?!\w)",
    r"(?<!\w)B\.\s*Kwadrat(?!\w)",
    r"(?<!\w)B\.X(?!\w)",
    r"(?<!\w)Q\.X(?!\w)",
    r"(?<!\w)R\.X(?!\w)",
    r"(?<!\w)(?:B\.[137]|S\.1|K\.[15]|Q\.[137]|R\.[138]|W\.1)(?!\w)",
    r"(?<!\w)R\.x(?!\w)",
    r"(?<!\w)TS(?!\w)",
]

FULL_LUMINA_GROUP = {"Lumina soul", "Lumina intense", "Lumina passion"}
FULL_B_GROUP = {"B. Kwadrat", "B.3", "B.7"}
FULL_Q_GROUP = {"Q.1", "Q.3", "Q.7"}
FULL_R_GROUP = {"R.1", "R.3", "R.8"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def display_series_token(token: str) -> str:
    replacements = {
        "B. Kwadrat": "B.Kwadrat",
        "Lumina soul": "Lumina Soul",
        "Lumina intense": "Lumina Intense",
        "Lumina passion": "Lumina Passion",
    }
    return replacements.get(token, token)


def title_series_fragments(exact_series: list[str]) -> list[str]:
    series_set = set(exact_series)
    used: set[str] = set()
    fragments: list[str] = []
    for token in exact_series:
        if token in used or token == "Inne":
            continue
        if token in FULL_LUMINA_GROUP and FULL_LUMINA_GROUP.issubset(series_set):
            fragments.append("Lumina")
            used.update(FULL_LUMINA_GROUP)
            continue
        if token in FULL_B_GROUP and FULL_B_GROUP.issubset(series_set):
            fragments.append("B.X")
            used.update(FULL_B_GROUP)
            continue
        if token in FULL_Q_GROUP and FULL_Q_GROUP.issubset(series_set):
            fragments.append("Q.X")
            used.update(FULL_Q_GROUP)
            continue
        if token in FULL_R_GROUP and FULL_R_GROUP.issubset(series_set):
            fragments.append("R.X")
            used.update(FULL_R_GROUP)
            continue
        fragments.append(display_series_token(token))
        used.add(token)
    return fragments


def title_represents_exact_series(
    title: str, exact_series: list[str]
) -> tuple[list[str], list[str]]:
    expected = [token for token in exact_series if token != "Inne"]
    actual = normalize_series(title)
    expected_set = set(expected)
    allowed_actual = set(expected)
    covered: set[str] = set(actual)

    if FULL_LUMINA_GROUP.issubset(expected_set) and re.search(
        r"(?<!\w)Lumina(?!\s+(?:Soul|Intense|Passion)\b)",
        title,
        flags=re.IGNORECASE,
    ):
        covered.update(FULL_LUMINA_GROUP)
        allowed_actual.add("Lumina")
    if FULL_B_GROUP.issubset(expected_set) and re.search(
        r"(?<!\w)B\.X(?!\w)", title, flags=re.IGNORECASE
    ):
        covered.update(FULL_B_GROUP)
    if FULL_Q_GROUP.issubset(expected_set) and re.search(
        r"(?<!\w)Q\.X(?!\w)", title, flags=re.IGNORECASE
    ):
        covered.update(FULL_Q_GROUP)
    if FULL_R_GROUP.issubset(expected_set) and re.search(
        r"(?<!\w)R\.X(?!\w)", title, flags=re.IGNORECASE
    ):
        covered.update(FULL_R_GROUP)

    missing = [token for token in expected if token not in covered]
    extra = [token for token in actual if token not in allowed_actual]
    return missing, extra


def clean_title_after_removal(title: str) -> str:
    title = re.sub(r"(?<!\S)/(?!\S)", " ", title)
    title = re.sub(r"\s+", " ", title)
    title = re.sub(r"\s+([,.;:)])", r"\1", title)
    title = re.sub(r"([(])\s+", r"\1", title)
    title = re.sub(r"\(\s*\)", "", title)
    title = re.sub(r"\s+", " ", title)
    return title.strip(" ,;-")


def remove_series_mentions(title: str) -> str:
    result = title
    for pattern in SERIES_PATTERNS:
        result = re.sub(pattern, " ", result, flags=re.IGNORECASE)
    return clean_title_after_removal(result)


def is_inside_parentheses(text: str, position: int) -> bool:
    return text.rfind("(", 0, position) > text.rfind(")", 0, position)


def outside_matches(pattern: str, text: str) -> list[re.Match[str]]:
    return [
        match
        for match in re.finditer(pattern, text, flags=re.IGNORECASE)
        if not is_inside_parentheses(text, match.start())
    ]


def remove_external_brands(title: str, brands: set[str]) -> str:
    result = title
    for brand in brands:
        matches = outside_matches(rf"\b{re.escape(brand)}\b", result)
        for match in reversed(matches):
            result = f"{result[:match.start()]} {result[match.end():]}"
    return clean_title_after_removal(result)


def find_product_code_anchor(title: str, sku: str) -> re.Match[str] | None:
    code = sku.split("/", 1)[0].strip()
    if code:
        matches = outside_matches(re.escape(code), title)
        if matches:
            return matches[-1]

    excluded = {"ip20", "ip21", "ip44", "ip55", "ip66", "knx", "usb"}
    candidates = outside_matches(
        r"(?<![\w.-])(?=[A-Z0-9-]{4,}(?![\w.-]))(?=[A-Z0-9-]*\d)[A-Z0-9-]+",
        title,
    )
    candidates = [
        match
        for match in candidates
        if match.group(0).casefold() not in excluded
        and not re.fullmatch(
            r"\d+(?:V|W|A|MA|MM)", match.group(0), flags=re.IGNORECASE
        )
    ]
    return candidates[-1] if candidates else None


def capitalize_title(title: str) -> str:
    for index, character in enumerate(title):
        if character.isalpha():
            return f"{title[:index]}{character.upper()}{title[index + 1:]}"
    return title


def insert_series_fragment(
    title: str, fragment: str, sku: str, exact_series: list[str]
) -> str:
    # W tytułach Hager/Berker seria jest najczytelniejsza bezpośrednio po marce.
    # Jeżeli występują Hager i referencja Berker w nawiasie, wybieramy Hager.
    has_lumina = any(token.casefold().startswith("lumina") for token in exact_series)
    external_hager = outside_matches(r"\bHager\b", title)
    external_berker = outside_matches(r"\bBerker\b", title)

    if has_lumina:
        brand = "Hager"
        base_title = remove_external_brands(title, {"Hager", "Berker"})
    elif external_hager:
        brand = "Hager"
        base_title = remove_external_brands(title, {"Hager"})
    elif external_berker:
        brand = "Berker"
        base_title = remove_external_brands(title, {"Berker"})
    else:
        brand = ""
        base_title = title

    brand_and_series = " ".join(part for part in (brand, fragment) if part)
    anchor = find_product_code_anchor(base_title, sku)
    if anchor:
        result = clean_title_after_removal(
            f"{base_title[:anchor.start()]}{brand_and_series} {base_title[anchor.start():]}"
        )
    else:
        result = clean_title_after_removal(f"{base_title} {brand_and_series}")
    return capitalize_title(result)


def update_title(title: str, sku: str, exact_series: list[str]) -> str:
    title_without_series = remove_series_mentions(title)
    fragments = title_series_fragments(exact_series)
    if not fragments:
        return title
    fragment = " ".join(fragments)
    return insert_series_fragment(title_without_series, fragment, sku, exact_series)


def build_output(input_path: Path, output_path: Path) -> dict[str, Any]:
    input_hash = sha256(input_path)
    workbook = load_workbook(input_path)
    input_structure = collect_workbook_structure(workbook)
    if workbook.sheetnames != [STORE_SHEET]:
        raise ValueError(f"Nieoczekiwane arkusze: {workbook.sheetnames}")
    sheet = workbook[STORE_SHEET]
    headers = header_map(sheet, HEADER_ROW)
    required = {"SKU", "Title", "Atrybut Produktu: Seria"}
    if not required.issubset(headers):
        raise ValueError(f"Brak kolumn: {sorted(required - set(headers))}")

    sku_col = headers["SKU"]
    title_col = headers["Title"]
    series_col = headers["Atrybut Produktu: Seria"]
    changes: list[dict[str, Any]] = []
    products_with_series = 0

    for row_number in range(2, sheet.max_row + 1):
        sku = str(sheet.cell(row_number, sku_col).value or "").strip()
        title = str(sheet.cell(row_number, title_col).value or "").strip()
        series_value = str(sheet.cell(row_number, series_col).value or "").strip()
        exact_series = normalize_series(series_value)
        expected_for_title = [token for token in exact_series if token != "Inne"]
        if not sku or not title or not expected_for_title:
            continue
        products_with_series += 1
        new_title = update_title(title, sku, exact_series)
        if new_title == title:
            continue
        sheet.cell(row_number, title_col).value = new_title
        changes.append(
            {
                "sku": sku,
                "row": row_number,
                "series": series_value,
                "before": title,
                "after": new_title,
            }
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    if sha256(input_path) != input_hash:
        raise RuntimeError("Plik wejściowy zmienił się podczas operacji")

    source = load_workbook(input_path, read_only=False, data_only=False)
    result = load_workbook(output_path, read_only=False, data_only=False)
    try:
        if collect_workbook_structure(result) != input_structure:
            raise ValueError("Struktura skoroszytu nie została zachowana")
        if scan_mojibake(result):
            raise ValueError("Wykryto podejrzane uszkodzenie kodowania")
        source_sheet = source[STORE_SHEET]
        result_sheet = result[STORE_SHEET]
        unexpected: list[str] = []
        missing_series_in_title: list[dict[str, Any]] = []
        extra_series_in_title: list[dict[str, Any]] = []
        malformed_titles: list[dict[str, str]] = []
        title_layout_violations: list[dict[str, str]] = []
        lumina_without_hager: list[dict[str, str]] = []
        invalid_abbreviations: list[dict[str, str]] = []
        duplicate_titles: dict[str, list[str]] = {}
        title_to_skus: dict[str, list[str]] = {}
        source_title_to_skus: dict[str, list[str]] = {}
        abbreviation_counts = {"Lumina": 0, "B.X": 0, "Q.X": 0, "R.X": 0}
        max_title_length = 0
        changed_skus = {change["sku"] for change in changes}
        for row_number in range(2, source_sheet.max_row + 1):
            source_sku = str(source_sheet.cell(row_number, sku_col).value or "").strip()
            source_title = str(source_sheet.cell(row_number, title_col).value or "").strip()
            if source_title:
                source_title_to_skus.setdefault(source_title.casefold(), []).append(source_sku)
        for row_number in range(1, source_sheet.max_row + 1):
            for column_number in range(1, source_sheet.max_column + 1):
                before = source_sheet.cell(row_number, column_number)
                after = result_sheet.cell(row_number, column_number)
                if before.value != after.value and column_number != title_col:
                    unexpected.append(before.coordinate)
                    if len(unexpected) >= 20:
                        break
                if before.style_id != after.style_id:
                    raise ValueError(f"Zmieniono styl komórki {before.coordinate}")
            if len(unexpected) >= 20:
                break
        if unexpected:
            raise ValueError(f"Zmiany poza kolumną Title: {unexpected}")

        for row_number in range(2, result_sheet.max_row + 1):
            sku = str(result_sheet.cell(row_number, sku_col).value or "").strip()
            title = str(result_sheet.cell(row_number, title_col).value or "").strip()
            series_value = str(result_sheet.cell(row_number, series_col).value or "").strip()
            expected = [token for token in normalize_series(series_value) if token != "Inne"]
            expected_set = set(expected)
            missing, extra = title_represents_exact_series(title, expected)
            if missing:
                missing_series_in_title.append({"sku": sku, "missing": missing, "title": title})
            if expected and extra:
                extra_series_in_title.append({"sku": sku, "extra": extra, "title": title})
            if sku in changed_skus and re.search(r"/(?:\s*/)+|\s{2,}|\(\s*\)", title):
                malformed_titles.append({"sku": sku, "title": title})
            if expected and re.match(r"(?i)(?:Hager|Berker)\b", title):
                title_layout_violations.append({"sku": sku, "title": title})
            if any(token.casefold().startswith("lumina") for token in expected):
                if not re.search(r"\bHager\b", title, flags=re.IGNORECASE):
                    lumina_without_hager.append({"sku": sku, "title": title})

            abbreviations = (
                ("Lumina", FULL_LUMINA_GROUP, r"(?<!\w)Lumina(?!\s+(?:Soul|Intense|Passion)\b)"),
                ("B.X", FULL_B_GROUP, r"(?<!\w)B\.X(?!\w)"),
                ("Q.X", FULL_Q_GROUP, r"(?<!\w)Q\.X(?!\w)"),
                ("R.X", FULL_R_GROUP, r"(?<!\w)R\.X(?!\w)"),
            )
            for abbreviation, full_group, pattern in abbreviations:
                present = bool(re.search(pattern, title, flags=re.IGNORECASE))
                complete = full_group.issubset(expected_set)
                if present:
                    abbreviation_counts[abbreviation] += 1
                if present and not complete:
                    invalid_abbreviations.append(
                        {"sku": sku, "abbreviation": abbreviation, "title": title}
                    )
            if title:
                title_to_skus.setdefault(title.casefold(), []).append(sku)
                max_title_length = max(max_title_length, len(title))
        if missing_series_in_title:
            raise ValueError(f"Brak serii w tytułach: {missing_series_in_title[:20]}")
        if extra_series_in_title:
            raise ValueError(f"Nadmiarowe serie w tytułach: {extra_series_in_title[:20]}")
        if malformed_titles:
            raise ValueError(f"Nieprawidłowe separatory w tytułach: {malformed_titles[:20]}")
        if title_layout_violations:
            raise ValueError(f"Marka na poczatku tytulu: {title_layout_violations[:20]}")
        if lumina_without_hager:
            raise ValueError(f"Lumina bez marki Hager: {lumina_without_hager[:20]}")
        if invalid_abbreviations:
            raise ValueError(f"Skrot bez pelnej grupy serii: {invalid_abbreviations[:20]}")
        duplicate_titles = {
            key: skus for key, skus in title_to_skus.items() if len(skus) > 1
        }
        source_duplicate_groups = {
            frozenset(skus)
            for skus in source_title_to_skus.values()
            if len(skus) > 1
        }
        new_duplicate_groups = [
            skus
            for skus in duplicate_titles.values()
            if frozenset(skus) not in source_duplicate_groups
        ]
        if new_duplicate_groups:
            raise ValueError(f"Nowe duplikaty tytulow: {new_duplicate_groups[:20]}")
    finally:
        source.close()
        result.close()

    return {
        "input": str(input_path.resolve()),
        "output": str(output_path.resolve()),
        "input_sha256": input_hash,
        "products": input_structure["sheets"][STORE_SHEET]["max_row"] - 1,
        "products_with_series_for_title": products_with_series,
        "changed_title_count": len(changes),
        "changes": changes,
        "duplicate_title_groups": list(duplicate_titles.values()),
        "abbreviation_counts": abbreviation_counts,
        "max_title_length": max_title_length,
        "validation": {
            "source_file_unchanged": True,
            "sheet_structure_preserved": True,
            "only_title_changed": True,
            "styles_preserved": True,
            "all_exact_series_present_in_titles": True,
            "title_abbreviations_applied": True,
            "product_first_brand_series_code_layout": True,
            "all_lumina_titles_include_hager": True,
            "no_new_duplicate_title_groups": True,
            "no_extra_recognized_series_in_titles": True,
            "mojibake_scan_passed": True,
        },
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Uzupełnia tytuły Hager/Berker pełną, dokładną listą serii."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    report = build_output(args.input, args.output)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                key: report[key]
                for key in (
                    "output",
                    "products",
                    "products_with_series_for_title",
                    "changed_title_count",
                    "duplicate_title_groups",
                    "abbreviation_counts",
                    "max_title_length",
                    "validation",
                )
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
