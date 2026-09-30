"""Buduje pole "Produkty kompatybilne" dla katalogu Hager/Berker.

Sklep (wtyczka alkan-compatible-products) trzyma powiazania w polu meta
`_alkan_compatible_product_ids` i pokazuje je na karcie produktu przed sekcja
Up-Sell. To NIE jest cross-sell: tu trafiaja tylko powiazania techniczne
potwierdzone przez producenta, a nie reguly sprzedazowe.

Zrodla, od najmocniejszego:
  1. katalog PDF - bloki "Wspolpracuje z" (arkusz "Kompatybilnosc" bazy
     katalogowej),
  2. BMEcat - relacje `mandatory` (elementy obowiazkowe),
  3. BMEcat - relacje `accessories`.
Kazda relacja dziala w obie strony: jezeli ramka wymienia klawisz jako
akcesorium, to na karcie klawisza ta ramka tez jest produktem kompatybilnym.

Punktacja, limity i slowa obnizajace ranking sa w
`dictionaries/hager_berker_compatible_products.yaml`.
Plik wejsciowy nie jest nadpisywany.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Iterable

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from build_hager_berker_cross_upsell import (
    BLUE,
    WHITE,
    CatalogProduct,
    same_variant_group,
    series_compatible,
    sha256,
)
from build_hager_berker_variants import (
    Record,
    alias_berker_code,
    find_headers,
    infer_primary_series,
    infer_series_key,
    is_frame,
    load_records,
    mark_old_duplicate_products,
    mojibake_markers,
    norm,
    sku_base,
)


ROOT = Path(__file__).resolve().parents[1]
TODAY = date.today().isoformat()
DEFAULT_INPUT = ROOT / "input" / "BerkerHager" / "Hager-Berker-Gniazdka.xlsx"
DEFAULT_PDF_BASE = ROOT / "output" / "Hager-Berker-Katalog-PDF_BAZA_KOMPATYBILNOSCI.xlsx"
DEFAULT_BMECAT = ROOT / "output" / "BMEcat_Hager_PL_produkty_ETIM_PL_ZDJECIA_RELACJE.xlsx"
DEFAULT_UPSELL_SOURCE = (
    ROOT / "output" / "Hager-Berker-Gniazdka_WOO_ALL_IMPORT_CROSS_UP_SELL_PO_AUDYCIE_PDF_TYTULY_STARYCH_SKU.xlsx"
)
DEFAULT_RULES = ROOT / "dictionaries" / "hager_berker_compatible_products.yaml"
DEFAULT_OUTPUT = ROOT / "output" / f"Hager-Berker_PRODUKTY_KOMPATYBILNE_{TODAY}.xlsx"
DEFAULT_CSV = ROOT / "output" / f"Hager-Berker_PRODUKTY_KOMPATYBILNE_WP_ALL_IMPORT_{TODAY}.csv"
DEFAULT_FILL_OUTPUT = ROOT / "output" / f"Kompatybilne-Berker-Hager_UZUPELNIONE_{TODAY}.xlsx"
DEFAULT_REPORT = ROOT / "reports" / f"Hager-Berker_produkty_kompatybilne_{TODAY}.json"

META_KEY = "_alkan_compatible_product_ids"
PDF_SHEET = "Kompatybilność"
BMECAT_SHEET = "Relacje produktowe"

SOURCE_LABELS = {
    "pdf": "Katalog PDF: Współpracuje z",
    "mandatory": "BMEcat: element obowiązkowy",
    "accessories": "BMEcat: akcesorium",
    "series_rule": "Reguła sklepu: ramka 1-krotna tej samej serii i koloru",
}
STATUS_OK = "OK"
STATUS_WARNING = "WARNING"
STATUS_REVIEW = "DO_SPRAWDZENIA"


@dataclass(frozen=True)
class Relation:
    source_sku: str
    target_sku: str
    source_key: str
    reverse: bool
    evidence: str


@dataclass(frozen=True)
class Candidate:
    target_sku: str
    score: int
    source_label: str
    reason: str
    evidence: str
    status: str = STATUS_OK


def load_rules(path: Path = DEFAULT_RULES) -> dict[str, Any]:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def general_colors(product: CatalogProduct) -> frozenset[str]:
    return frozenset(norm(item) for item in str(product.record.color or "").split("|") if norm(item))


def is_colorless(product: CatalogProduct, rules: dict[str, Any]) -> bool:
    return (
        product.role in rules["colorless_roles"]
        or not product.color
        or product.color in rules["neutral_colors"]
    )


def color_bonus(source: CatalogProduct, target: CatalogProduct, rules: dict[str, Any]) -> int | None:
    """Zwraca premie za kolor albo None, gdy kolory sie wykluczaja."""
    if is_colorless(source, rules) or is_colorless(target, rules):
        return 0
    if source.color == target.color:
        return rules["bonuses"]["same_manufacturer_color"]
    left, right = general_colors(source), general_colors(target)
    shared = left & right
    if not shared:
        return None
    return 1 + round(3 * len(shared) / len(left | right))


def is_niche(product: CatalogProduct, rules: dict[str, Any]) -> bool:
    title = norm(product.title)
    return any(token in title for token in rules["niche_title_tokens"])


def build_code_index(products: Iterable[CatalogProduct]) -> dict[str, list[str]]:
    """Kod katalogowy (Hager albo alias Berker) -> SKU sklepu, bez starych duplikatow."""
    index: dict[str, list[str]] = defaultdict(list)
    for product in products:
        if product.record.is_old_duplicate:
            continue
        for code in {product.record.code, product.record.alias_code}:
            if code and product.sku not in index[code.upper()]:
                index[code.upper()].append(product.sku)
    return dict(index)


def split_cell(value: object) -> list[str]:
    return [item.strip() for item in str(value or "").split("|") if item.strip()]


def load_pdf_relations(path: Path, known_skus: set[str]) -> list[Relation]:
    workbook = load_workbook(path, read_only=True)
    try:
        sheet = workbook[PDF_SHEET]
        rows = sheet.iter_rows(min_row=4, values_only=True)
        header = [str(value or "") for value in next(rows)]
        column = {name: header.index(name) for name in (
            "SKU źródłowe sklepu", "SKU docelowe sklepu", "Status", "Dowód katalogowy",
        )}
        relations: list[Relation] = []
        for row in rows:
            if str(row[column["Status"]] or "") != "OK":
                continue
            evidence = str(row[column["Dowód katalogowy"]] or "")
            for source in split_cell(row[column["SKU źródłowe sklepu"]]):
                for target in split_cell(row[column["SKU docelowe sklepu"]]):
                    if source in known_skus and target in known_skus and source != target:
                        relations.append(Relation(source, target, "pdf", False, evidence))
        return relations
    finally:
        workbook.close()


def load_bmecat_relations(path: Path, code_index: dict[str, list[str]]) -> list[Relation]:
    workbook = load_workbook(path, read_only=True)
    try:
        sheet = workbook[BMECAT_SHEET]
        rows = sheet.iter_rows(values_only=True)
        header = [str(value or "") for value in next(rows)]
        code_column = header.index("Kod produktu")
        columns = {"mandatory": header.index("Elementy obowiązkowe"), "accessories": header.index("Akcesoria")}
        relations: list[Relation] = []
        for row in rows:
            source_code = str(row[code_column] or "").strip().upper()
            source_skus = code_index.get(source_code)
            if not source_skus:
                continue
            for source_key, index in columns.items():
                for target_code in split_cell(row[index]):
                    for target in code_index.get(target_code.upper(), []):
                        for source in source_skus:
                            if source != target:
                                relations.append(Relation(
                                    source, target, source_key, False,
                                    f"BMEcat: {source_code} -> {target_code.upper()} ({source_key})",
                                ))
        return relations
    finally:
        workbook.close()


def header_index(header: list[str], name: str) -> int | None:
    """Eksporty WP All Export pisza raz `SKU`, raz `Sku` - szukamy bez rozrozniania wielkosci liter."""
    lowered = [item.strip().lower() for item in header]
    return lowered.index(name.lower()) if name.lower() in lowered else None


def load_upsells(path: Path | None) -> dict[str, set[str]]:
    """Up-selle sa na tej samej karcie produktu - nie powtarzamy ich w kompatybilnych."""
    if not path or not path.exists():
        return {}
    workbook = load_workbook(path, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [str(value or "") for value in next(rows)]
        sku_column, up_column = header_index(header, "SKU"), header_index(header, "Up-Sells")
        if sku_column is None or up_column is None:
            return {}
        return {
            str(row[sku_column]).strip(): set(split_cell(row[up_column]))
            for row in rows
            if row[sku_column] and row[up_column]
        }
    finally:
        workbook.close()


def minimal_record(excel_row: int, product_id: str, title: str, sku: str) -> Record:
    """Produkt z eksportu sklepu, ktorego nie ma w pliku atrybutow - znamy tylko tytul."""
    primary = infer_primary_series("", title)
    code = sku_base(sku)
    record = Record(
        excel_row=excel_row, product_id=product_id, title=title, sku=sku, code=code,
        product_type="", series=primary, primary_series=primary, color="",
        manufacturer_color_before="", manufacturer_color="", material="", multiplicity="",
        tags_before="", core=title, family="inne", frame=is_frame(title, ""),
        color_detail="", material_detail="", finish_detail="", orientation="",
        network_category="", screened="", grounding="", terminal_type="", voltage="",
        nominal_current="", pole_count="", shutters="", detection_range="", output_count="",
        alias_code=alias_berker_code(title),
    )
    record.series_key = infer_series_key(primary, title, code)
    return record


def read_fill_rows(path: Path) -> list[tuple[int, str, str, str]]:
    """(wiersz, id, SKU, tytul) z eksportu sklepu, ktory mamy uzupelnic."""
    workbook = load_workbook(path, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [str(value or "") for value in next(rows)]
        columns = {name: header_index(header, name) for name in ("id", "SKU", "Title", META_KEY)}
        missing = [name for name, index in columns.items() if index is None]
        if missing:
            raise ValueError(f"W pliku {path.name} brakuje kolumn: {', '.join(missing)}")
        return [
            (number, str(row[columns["id"]]).strip(), str(row[columns["SKU"]]).strip(), str(row[columns["Title"]] or ""))
            for number, row in enumerate(rows, 2)
            if row[columns["SKU"]]
        ]
    finally:
        workbook.close()


def parse_existing_ids(value: object) -> list[str]:
    """Czyta recznie wpisane ID: serializowana tablica PHP (`a:1:{i:0;i:20648;}`) albo lista po przecinku."""
    text = str(value or "").strip()
    if text.startswith("a:"):
        return re.findall(r'i:\d+;(?:i:|s:\d+:")(\d+)', text)
    return re.findall(r"\d+", text)


def fill_workbook(
    source: Path,
    target: Path,
    products: list[CatalogProduct],
    selected: dict[str, list[Candidate]],
    max_items: int,
) -> dict[str, object]:
    """Zapisuje kopie eksportu sklepu z wypelniona kolumna pola meta; reszta komorek bez zmian."""
    by_sku = {product.sku: product for product in products}
    workbook = load_workbook(source)
    sheet = workbook.worksheets[0]
    header = [str(cell.value or "") for cell in sheet[1]]
    sku_column, meta_column = header_index(header, "SKU"), header_index(header, META_KEY)
    before = [[cell.value for cell in row] for row in sheet.iter_rows()]
    filled = 0
    manual: list[str] = []
    for row in sheet.iter_rows(min_row=2):
        sku = str(row[sku_column].value or "").strip()
        ids = [by_sku[item.target_sku].record.product_id for item in selected.get(sku, [])]
        existing = parse_existing_ids(row[meta_column].value)
        if existing:
            # Reczne wpisy ze sklepu maja pierwszenstwo - automat tylko dopelnia liste.
            manual.append(sku)
            ids = (existing + [item for item in ids if item not in existing])[: max(max_items, len(existing))]
        if ids:
            row[meta_column].value = ",".join(ids)
            filled += 1
    after = [[cell.value for cell in row] for row in sheet.iter_rows()]
    untouched = len(before) == len(after) and all(
        old == new
        for old_row, new_row in zip(before, after)
        for index, (old, new) in enumerate(zip(old_row, new_row))
        if index != meta_column
    )
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return {
        "fill_source": str(source),
        "fill_output": str(target),
        "fill_rows": len(before) - 1,
        "fill_rows_filled": filled,
        "fill_rows_with_manual_value_kept": manual,
        "fill_other_cells_unchanged": untouched,
    }


def with_reverse(relations: Iterable[Relation]) -> list[Relation]:
    result: list[Relation] = []
    for relation in relations:
        result.append(relation)
        result.append(Relation(
            relation.target_sku, relation.source_sku, relation.source_key, True, relation.evidence,
        ))
    return result


def build_candidate(
    source: CatalogProduct,
    target: CatalogProduct,
    relation: Relation,
    rules: dict[str, Any],
) -> Candidate | None:
    if source.sku == target.sku or target.record.is_old_duplicate:
        return None
    # Warianty tego samego produktu przelacza juz wtyczka wariantow.
    if same_variant_group(source, target):
        return None
    score = rules["source_scores"][relation.source_key]["reverse" if relation.reverse else "direct"]
    label = SOURCE_LABELS[relation.source_key] + (" (relacja odwrotna)" if relation.reverse else "")
    reasons = []
    status = STATUS_OK
    bonus = color_bonus(source, target, rules)
    if bonus is None:
        score -= rules["color_mismatch_penalty"]
        status = STATUS_WARNING
        reasons.append("inny kolor - producent nie ma tego elementu w kolorze produktu")
    elif bonus >= rules["bonuses"]["same_manufacturer_color"]:
        score += bonus
        reasons.append("ten sam kolor producenta")
    elif bonus:
        score += bonus
        reasons.append("zgodny kolor ogólny")
    if series_compatible(source, target):
        score += rules["bonuses"]["shared_series"]
        reasons.append("wspólna seria")
    if target.role == "frame" and target.record.multiplicity == "1x":
        score += rules["bonuses"]["single_frame"]
        reasons.append("ramka 1-krotna")
    if source.role == "frame":
        score += rules["function_priority"].get(target.function, 0)
    if is_niche(target, rules) and not is_niche(source, rules):
        score -= rules["niche_penalty"]
    return Candidate(
        target.sku, score, label, ", ".join(reasons) or "relacja producenta", relation.evidence, status,
    )


def diversity_key(source: CatalogProduct, target: CatalogProduct, rules: dict[str, Any]) -> tuple[str, ...]:
    series = target.record.series_key or norm(target.record.primary_series)
    if target.role == "frame":
        # Wklad B.X pasuje do ramek B.Kwadrat, B.3 i B.7 - pokazujemy przekroj serii.
        return ("frame", series)
    role = f"{target.role}:{target.function}"
    if is_colorless(source, rules):
        # Mechanizm pasuje do klawiszy wielu serii - nie 8 kolorow jednej serii.
        return (role, series)
    return (role,)


def pick_diverse(
    source: CatalogProduct,
    candidates: Iterable[Candidate],
    by_sku: dict[str, CatalogProduct],
    color_popularity: Counter,
    rules: dict[str, Any],
    max_items: int,
) -> list[Candidate]:
    ordered = sorted(
        candidates,
        key=lambda item: (-item.score, -color_popularity[by_sku[item.target_sku].color], item.target_sku),
    )
    chosen: list[Candidate] = []
    used: Counter = Counter()
    # Pierwszy przebieg: po jednym z kazdego rodzaju, kolejne dopelniaja do limitu.
    for limit in range(1, rules["max_per_kind"] + 1):
        for candidate in ordered:
            if len(chosen) >= max_items:
                break
            if candidate in chosen:
                continue
            key = diversity_key(source, by_sku[candidate.target_sku], rules)
            if used[key] >= limit:
                continue
            chosen.append(candidate)
            used[key] += 1
    return sorted(chosen, key=lambda item: (-item.score, item.target_sku))


def select_compatible(
    products: list[CatalogProduct],
    relations: Iterable[Relation],
    rules: dict[str, Any],
    upsells: dict[str, set[str]] | None = None,
    max_items: int | None = None,
) -> dict[str, list[Candidate]]:
    by_sku = {product.sku: product for product in products}
    upsells = upsells or {}
    max_items = max_items or rules["max_items"]
    color_popularity = Counter(product.color for product in products)

    best: dict[str, dict[str, Candidate]] = defaultdict(dict)
    for relation in relations:
        source, target = by_sku.get(relation.source_sku), by_sku.get(relation.target_sku)
        if not source or not target or source.record.is_old_duplicate:
            continue
        if target.sku in upsells.get(source.sku, set()):
            continue
        candidate = build_candidate(source, target, relation, rules)
        if not candidate:
            continue
        current = best[source.sku].get(target.sku)
        if not current or candidate.score > current.score:
            best[source.sku][target.sku] = candidate

    selected: dict[str, list[Candidate]] = {}
    for source_sku, candidates in best.items():
        matching = [item for item in candidates.values() if item.status == STATUS_OK]
        # Relacje w innym kolorze tylko wtedy, gdy nic nie pasuje kolorem.
        chosen = pick_diverse(
            by_sku[source_sku], matching or list(candidates.values()), by_sku, color_popularity, rules, max_items,
        )
        if chosen:
            selected[source_sku] = chosen
    return selected


def series_fallback(
    products: list[CatalogProduct],
    covered: set[str],
    rules: dict[str, Any],
    upsells: dict[str, set[str]] | None = None,
) -> dict[str, list[Candidate]]:
    """Propozycje dla produktow bez relacji producenta: ramka 1-krotna tej samej serii i koloru."""
    config = rules["series_fallback"]
    upsells = upsells or {}
    frames = [
        product for product in products
        if product.role == "frame" and product.record.multiplicity == "1x" and not product.record.is_old_duplicate
    ]
    proposals: dict[str, list[Candidate]] = {}
    for source in products:
        if source.sku in covered or source.record.is_old_duplicate or source.role not in config["source_roles"]:
            continue
        title = norm(source.title)
        if not source.color or any(token in title for token in config["skip_title_tokens"]):
            continue
        matches = [
            Candidate(
                frame.sku, config["score"], SOURCE_LABELS["series_rule"],
                "brak relacji producenta; wspólna seria i ten sam kolor producenta", "", STATUS_REVIEW,
            )
            for frame in frames
            if series_compatible(source, frame)
            and source.color == frame.color
            and not same_variant_group(source, frame)
            and frame.sku not in upsells.get(source.sku, set())
        ]
        if matches:
            proposals[source.sku] = sorted(matches, key=lambda item: item.target_sku)[: rules["max_per_kind"] + 1]
    return proposals


def validate(
    products: list[CatalogProduct],
    selected: dict[str, list[Candidate]],
    upsells: dict[str, set[str]],
    max_items: int,
) -> dict[str, bool]:
    by_sku = {product.sku: product for product in products}
    pairs = [(source, item.target_sku) for source, items in selected.items() for item in items]
    return {
        "all_targets_exist": all(target in by_sku for _, target in pairs),
        "all_targets_have_id": all(by_sku[target].record.product_id for _, target in pairs),
        "no_self_reference": all(source != target for source, target in pairs),
        "no_old_duplicates": not any(
            by_sku[sku].record.is_old_duplicate for pair in pairs for sku in pair
        ),
        "no_same_variant_group": not any(
            same_variant_group(by_sku[source], by_sku[target]) for source, target in pairs
        ),
        "no_upsell_overlap": not any(target in upsells.get(source, set()) for source, target in pairs),
        "no_duplicate_targets": all(
            len({item.target_sku for item in items}) == len(items) for items in selected.values()
        ),
        "limit_respected": all(len(items) <= max_items for items in selected.values()),
        "no_mojibake": not any(mojibake_markers(item.reason) for items in selected.values() for item in items),
    }


def style_header(sheet, row: int = 1) -> None:
    for cell in sheet[row]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)


LINK_HEADER = [
    "SKU źródłowe", "Tytuł źródłowy", "Lp.", "SKU kompatybilne", "ID kompatybilne", "Tytuł kompatybilny",
    "Źródło relacji", "Punkty", "Status", "Uzasadnienie", "Dowód",
]
LINK_WIDTHS = (22, 70, 6, 22, 10, 70, 44, 8, 16, 44, 80)


def write_link_sheet(sheet, products: list[CatalogProduct], links: dict[str, list[Candidate]]) -> None:
    by_sku = {product.sku: product for product in products}
    sheet.append(LINK_HEADER)
    for product in products:
        for position, item in enumerate(links.get(product.sku, []), 1):
            target = by_sku[item.target_sku]
            sheet.append([
                product.sku, product.title, position, target.sku, target.record.product_id, target.title,
                item.source_label, item.score, item.status, item.reason, item.evidence,
            ])
    style_header(sheet)
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(LINK_WIDTHS):
        sheet.column_dimensions["ABCDEFGHIJK"[index]].width = width


def write_outputs(
    products: list[CatalogProduct],
    selected: dict[str, list[Candidate]],
    proposals: dict[str, list[Candidate]],
    metrics: dict[str, object],
    output: Path,
    csv_path: Path,
) -> None:
    by_sku = {product.sku: product for product in products}
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Podsumowanie"
    summary.append(["Hager / Berker - produkty kompatybilne"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([f"Pole meta sklepu: {META_KEY} (ID po przecinku). To nie jest cross-sell ani up-sell."])
    summary.append([])
    for key, value in metrics.items():
        if isinstance(value, dict):
            if key != "validation":
                for name, count in value.items():
                    summary.append([f"{key}: {name}", count])
        elif isinstance(value, list):
            summary.append([key, ", ".join(str(item) for item in value)])
        else:
            summary.append([key, value])
    summary.append([])
    summary.append(["Walidacja", "Wynik"])
    for key, value in metrics["validation"].items():
        summary.append([key, "OK" if value else "BŁĄD"])
    summary.column_dimensions["A"].width = 70
    summary.column_dimensions["B"].width = 90

    sheet = workbook.create_sheet("Import")
    sheet.append(["id", "SKU", "Title", META_KEY, "Produkty kompatybilne (SKU)", "Liczba", "Status"])
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, delimiter=";")
        writer.writerow(["id", "SKU", META_KEY])
        for product in products:
            items = selected.get(product.sku)
            if not items:
                continue
            ids = ",".join(by_sku[item.target_sku].record.product_id for item in items)
            statuses = {item.status for item in items}
            status = next(name for name in (STATUS_REVIEW, STATUS_WARNING, STATUS_OK) if name in statuses)
            sheet.append([
                product.record.product_id, product.sku, product.title, ids,
                "|".join(item.target_sku for item in items), len(items), status,
            ])
            writer.writerow([product.record.product_id, product.sku, ids])
    style_header(sheet)
    sheet.auto_filter.ref = sheet.dimensions
    for letter, width in zip("ABCDEFG", (10, 22, 80, 50, 80, 10, 16)):
        sheet.column_dimensions[letter].width = width

    write_link_sheet(workbook.create_sheet("Powiązania"), products, selected)
    write_link_sheet(workbook.create_sheet("Propozycje bez relacji"), products, proposals)

    sheet = workbook.create_sheet("Bez powiązań")
    sheet.append(["SKU", "id", "Tytuł", "Typ produktu", "Seria", "Rola", "Powód"])
    for product in products:
        if product.sku in selected:
            continue
        if product.record.is_old_duplicate:
            reason = "stary duplikat - pomijany"
        elif product.sku in proposals:
            reason = "brak relacji producenta - jest propozycja w arkuszu 'Propozycje bez relacji'"
        else:
            reason = "brak relacji producenta do produktów obecnych w sklepie (PDF i BMEcat)"
        sheet.append([
            product.sku, product.record.product_id, product.title, product.record.product_type,
            product.record.series, product.role, reason,
        ])
    style_header(sheet)
    sheet.auto_filter.ref = sheet.dimensions
    for letter, width in zip("ABCDEFG", (22, 10, 80, 30, 30, 14, 70)):
        sheet.column_dimensions[letter].width = width

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def build(
    input_path: Path,
    pdf_base: Path,
    bmecat: Path,
    upsell_source: Path | None,
    output: Path,
    csv_path: Path,
    report: Path,
    rules_path: Path = DEFAULT_RULES,
    max_items: int | None = None,
    include_series_fallback: bool = False,
    fill_source: Path | None = None,
    fill_output: Path | None = None,
) -> dict[str, object]:
    protected = {input_path.resolve()} | ({fill_source.resolve()} if fill_source else set())
    for path in (output, csv_path, fill_output):
        if path and path.resolve() in protected:
            raise ValueError("Plik wynikowy nie moze nadpisac pliku wejsciowego.")

    rules = load_rules(rules_path)
    max_items = max_items or rules["max_items"]
    workbook = load_workbook(input_path, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        records = load_records(sheet, find_headers(sheet))
    finally:
        workbook.close()
    id_mismatches: list[str] = []
    added_from_fill: list[str] = []
    if fill_source:
        by_sku = {record.sku: record for record in records}
        for number, product_id, sku, title in read_fill_rows(fill_source):
            if sku not in by_sku:
                records.append(minimal_record(number, product_id, title, sku))
                added_from_fill.append(sku)
            elif by_sku[sku].product_id != product_id:
                # ID ida do pola meta, wiec musza pochodzic z uzupelnianego eksportu.
                id_mismatches.append(sku)
                by_sku[sku].product_id = product_id
        # Aktualne Up-Sells sa w uzupelnianym eksporcie - to one decyduja o powtorkach na karcie.
        upsell_source = fill_source
    old_duplicates = mark_old_duplicate_products(records)
    products = [CatalogProduct(record) for record in records]
    known_skus = {product.sku for product in products}

    direct = load_pdf_relations(pdf_base, known_skus)
    direct += load_bmecat_relations(bmecat, build_code_index(products))
    upsells = load_upsells(upsell_source)
    selected = select_compatible(products, with_reverse(direct), rules, upsells, max_items)
    proposals = series_fallback(products, set(selected), rules, upsells)
    if include_series_fallback:
        selected.update(proposals)
        proposals = {}

    counts = [len(items) for items in selected.values()]
    active = len(products) - old_duplicates
    metrics: dict[str, object] = {
        "generated": TODAY,
        "input_file": str(input_path),
        "input_sha256": sha256(input_path),
        "rules_file": str(rules_path),
        "products": len(products),
        "old_duplicates_skipped": old_duplicates,
        "direct_relations_in_shop": dict(Counter(relation.source_key for relation in direct)),
        "products_with_compatible": len(selected),
        "products_without_compatible": active - len(selected),
        "coverage_percent": round(100 * len(selected) / active, 1) if active else 0,
        "compatible_links": sum(counts),
        "average_per_product": round(sum(counts) / len(counts), 1) if counts else 0,
        "links_by_source": dict(Counter(
            item.source_label for items in selected.values() for item in items
        )),
        "links_by_status": dict(Counter(item.status for items in selected.values() for item in items)),
        "series_fallback_included": include_series_fallback,
        "series_fallback_proposals_products": len(proposals),
        "without_compatible_by_role": dict(Counter(
            product.role for product in products
            if product.sku not in selected and not product.record.is_old_duplicate
        )),
        "max_items": max_items,
        "validation": validate(products, selected, upsells, max_items),
    }
    if fill_source:
        metrics["products_added_from_fill_file"] = added_from_fill
        metrics["id_mismatches_fixed_from_fill_file"] = id_mismatches
        fill_metrics = fill_workbook(
            fill_source, fill_output or DEFAULT_FILL_OUTPUT, products, selected, max_items,
        )
        metrics.update(fill_metrics)
        metrics["validation"]["fill_other_cells_unchanged"] = fill_metrics["fill_other_cells_unchanged"]
    metrics["passed"] = all(metrics["validation"].values())

    write_outputs(products, selected, proposals, metrics, output, csv_path)
    metrics["output_file"] = str(output)
    metrics["csv_file"] = str(csv_path)
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")
    return metrics


def main() -> None:
    parser = argparse.ArgumentParser(description="Buduje pole 'Produkty kompatybilne' dla Hager/Berker.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--pdf-base", type=Path, default=DEFAULT_PDF_BASE)
    parser.add_argument("--bmecat", type=Path, default=DEFAULT_BMECAT)
    parser.add_argument("--upsell-source", type=Path, default=DEFAULT_UPSELL_SOURCE)
    parser.add_argument("--rules", type=Path, default=DEFAULT_RULES)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--csv", type=Path, default=DEFAULT_CSV)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--max-items", type=int, default=None)
    parser.add_argument(
        "--include-series-fallback", action="store_true",
        help="Dolacz do importu propozycje z reguly sklepu (status DO_SPRAWDZENIA).",
    )
    parser.add_argument(
        "--fill", type=Path, default=None,
        help="Eksport sklepu z kolumna _alkan_compatible_product_ids do uzupelnienia (kopia w --fill-output).",
    )
    parser.add_argument("--fill-output", type=Path, default=DEFAULT_FILL_OUTPUT)
    args = parser.parse_args()

    metrics = build(
        args.input, args.pdf_base, args.bmecat, args.upsell_source, args.output, args.csv, args.report,
        args.rules, args.max_items, args.include_series_fallback, args.fill, args.fill_output,
    )
    # Konsola Windows (cp1250) - tylko ASCII; pelny raport jest w JSON.
    print(f"produkty z kompatybilnymi: {metrics['products_with_compatible']} / {metrics['products']}")
    print(f"powiazania: {metrics['compatible_links']}, walidacja: {'OK' if metrics['passed'] else 'BLAD'}")
    print(f"propozycje bez relacji producenta: {metrics['series_fallback_proposals_products']}")
    if args.fill:
        print(f"uzupelniony eksport: {metrics['fill_rows_filled']} / {metrics['fill_rows']} wierszy")
    print(f"raport: {args.report}")


if __name__ == "__main__":
    main()
