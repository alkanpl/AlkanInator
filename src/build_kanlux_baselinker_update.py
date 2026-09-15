from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path
from typing import Any

import pandas as pd

from export_to_baselinker_csv import (
    BASELINKER_COLUMNS,
    build_baselinker_rows,
    build_feature_value_normalizer,
    build_manufacturer_data_by_producer,
    feature_value_lookup_keys,
    hermetic_fluorescent_title,
    load_producer_suffixes,
    load_supplier_feature_policy,
    is_accessory_product_type,
    match_attribute_category_rule,
    product_type_from_title,
    product_types_agree,
    normalize_feature_name,
    normalize_feature_value,
    write_baselinker_csv,
    write_feature_review_file,
)
from extract_attributes import extract_attributes_from_text
from optimize_titles import normalize_seo_title_terms, normalize_title_uniqueness_key
from utils import color_stem, compact_spaces, ensure_dir, load_yaml, normalize_header, read_products


DEFAULT_REFERENCE = "archived_input_files/Kanlux (4).xlsx"
DEFAULT_WOO = "archived_input_files/KanluxWoo.xlsx"
DEFAULT_ENRICHED = "output/alkan_kanlux_2026-06-11_enriched.xlsx"
DEFAULT_ACCEPTED = "output/alkan_kanlux_titles_no_duplicates_categories_descriptions.xlsx"
DEFAULT_CONTROL = "output/kanlux_baselinker_update_2026-06-11_control.xlsx"
DEFAULT_CSV = "output/kanlux_baselinker_update_2026-06-11.csv"
DEFAULT_ATTRIBUTE_CSV = "output/kanlux_baselinker_attributes_update_2026-06-11.csv"
DEFAULT_MISSING_XLSX = "output/kanlux_braki_wymaganych_2026-06-11.xlsx"
DEFAULT_MANUAL_FILLS = "archived_input_files/kanlux_braki_wymaganych_uzupelnione.xlsx"
DEFAULT_REPORTS = "reports/kanlux_baselinker_update_2026-06-11"
WOO_ATTRIBUTE_PREFIX = "Atrybut Produktu: "
ATTRIBUTE_UPDATE_COLUMNS = ["product_id", "sku", "name", "features"]

UNMATCHED_OVERRIDES = {
    "36312": {
        "type": "Ogranicznik przepięć",
        "attributes": {"seria": "KSPD-T2"},
        "category": "Aparatura modułowa > Ochronniki przepięciowe",
        "description": (
            "<p>Ogranicznik przepięć KSPD-T2 275/160 Kanlux 36312 jest przeznaczony do ochrony instalacji "
            "elektrycznej przed skutkami przepięć. Układ 3P+N pozwala zastosować go w instalacji trójfazowej "
            "z przewodem neutralnym.</p><h3>Najważniejsze informacje</h3><ul><li>Producent: Kanlux</li>"
            "<li>Typ: ogranicznik przepięć</li><li>Układ: 3P+N</li><li>Seria: KSPD-T2</li>"
            "<li>Kod producenta: 36312</li></ul>"
        ),
    },
    "45930": {
        "type": "Lampa ogrodowa",
        "attributes": {
            "seria": "STONO",
            "gwint": "E27",
            "srednica": "200mm",
            "ip": "IP65",
            "kolor": "biały",
            "zrodlo_swiatla": "Wymienne",
        },
        "category": "Oświetlenie > Oświetlenie zewnętrzne > Lampy ogrodowe",
        "description": (
            "<p>Lampa ogrodowa kula STONO Kanlux 45930 jest oprawą zewnętrzną na źródło światła z trzonkiem "
            "E27. Średnica 200 mm i biała obudowa pozwalają wykorzystać ją przy ścieżkach, tarasach i zieleni. "
            "Stopień ochrony IP65 zabezpiecza oprawę przed pyłem i strumieniami wody.</p>"
            "<h3>Najważniejsze informacje</h3><ul><li>Producent: Kanlux</li><li>Seria: STONO</li>"
            "<li>Trzonek: E27</li><li>Średnica: 200 mm</li><li>Stopień ochrony: IP65</li>"
            "<li>Kolor: biały</li><li>Kod producenta: 45930</li></ul>"
        ),
    },
    "28997": {
        "type": "Oprawa elewacyjna solarna LED",
        "attributes": {
            "seria": "REKA",
            "moc": "6W",
            "strumien": "600lm",
            "barwa": "3000-6500K",
            "ip": "IP65",
            "czujnik": "z czujnikiem mikrofalowym",
            "kolor": "grafitowy",
            "zrodlo_swiatla": "LED",
        },
        "fields": {"Barwa - kategoria": "Zmienna"},
        "category": "Oświetlenie > Oświetlenie zewnętrzne > Lampy ogrodowe > Kinkiety zewnętrzne",
        "description": (
            "<p>Oprawa elewacyjna solarna LED REKA Kanlux 28997 zapewnia światło bez doprowadzania przewodu "
            "zasilającego. Moc 6 W, strumień 600 lm i regulowana temperatura barwowa 3000-6500 K pozwalają "
            "dopasować oświetlenie elewacji lub wejścia. Czujnik mikrofalowy automatyzuje włączanie światła, "
            "a IP65 wspiera pracę na zewnątrz.</p><h3>Najważniejsze informacje</h3><ul>"
            "<li>Producent: Kanlux</li><li>Seria: REKA</li><li>Moc: 6 W</li><li>Strumień: 600 lm</li>"
            "<li>Temperatura barwowa: 3000-6500 K</li><li>Stopień ochrony: IP65</li>"
            "<li>Czujnik: mikrofalowy</li><li>Kolor: grafitowy</li><li>Kod producenta: 28997</li></ul>"
        ),
    },
}


def normalize_sku(value: Any) -> str:
    text = compact_spaces(str(value or ""))
    text = re.sub(r"/KANL?$", "", text, flags=re.IGNORECASE)
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def normalize_ean(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def indexed_rows(df: pd.DataFrame, sku_columns: list[str], ean_columns: list[str]) -> tuple[dict[str, pd.Series], dict[str, pd.Series]]:
    by_sku: dict[str, pd.Series] = {}
    by_ean: dict[str, pd.Series] = {}
    for _, row in df.iterrows():
        for column in sku_columns:
            if column in row:
                sku = normalize_sku(row.get(column, ""))
                if sku:
                    by_sku.setdefault(sku, row)
                    break
        for column in ean_columns:
            if column in row:
                ean = normalize_ean(row.get(column, ""))
                if ean:
                    by_ean.setdefault(ean, row)
                    break
    return by_sku, by_ean


def reference_value(row: pd.Series, column: str) -> str:
    return compact_spaces(str(row.get(column, "")))


def build_attribute_aliases(attribute_knowledge: dict[str, Any]) -> dict[str, str]:
    aliases: dict[str, str] = {}
    for mapping in attribute_knowledge.get("attribute_mappings") or []:
        source_name = compact_spaces(str(mapping.get("source_name", "")))
        canonical_name = compact_spaces(str(mapping.get("canonical_name", "")))
        if source_name and canonical_name:
            aliases[normalize_header(source_name)] = canonical_name
            aliases[normalize_header(canonical_name)] = canonical_name
    return aliases


def woo_features(row: pd.Series, attribute_aliases: dict[str, str]) -> dict[str, str]:
    features: dict[str, str] = {}
    for column, raw_value in row.items():
        column_name = str(column)
        if not column_name.startswith(WOO_ATTRIBUTE_PREFIX):
            continue
        if pd.isna(raw_value):
            continue
        value = compact_spaces(str(raw_value or ""))
        if not value or value.lower() in {"nan", "none", "null"}:
            continue
        raw_name = compact_spaces(column_name.removeprefix(WOO_ATTRIBUTE_PREFIX))
        feature_name = attribute_aliases.get(normalize_header(raw_name), normalize_feature_name(raw_name))
        if feature_name == "Dostępność":
            continue
        features[feature_name] = value
    return features


def add_woo_feature_fallbacks(
    row: dict[str, Any],
    woo_row: pd.Series | None,
    attribute_aliases: dict[str, str],
) -> None:
    if woo_row is None:
        row["woo_match_status"] = "MISSING"
        row["woo_attribute_count"] = 0
        row["woo_features_json"] = "{}"
        return
    features = woo_features(woo_row, attribute_aliases)
    row["woo_match_status"] = "MATCHED"
    row["woo_product_id"] = reference_value(woo_row, "id")
    row["woo_title"] = reference_value(woo_row, "Title")
    row["woo_attribute_count"] = len(features)
    row["woo_features_json"] = json.dumps(features, ensure_ascii=False, separators=(",", ":"))
    for feature_name, value in features.items():
        row[f"Parametr: {feature_name}"] = value


def add_woo_values_to_feature_normalizer(
    normalizer: dict[str, dict[str, str]],
    woo: pd.DataFrame,
    attribute_aliases: dict[str, str],
) -> None:
    for _, woo_row in woo.iterrows():
        for feature_name, value in woo_features(woo_row, attribute_aliases).items():
            canonical_value = normalize_feature_value(value, feature_name) or value
            known_values = normalizer.setdefault(feature_name, {})
            for key in feature_value_lookup_keys(value):
                known_values.setdefault(key, canonical_value)
            # Wartosci wielokrotne (np. "ABS|PE") maja byc znane takze per czesc.
            for part in canonical_value.split("|"):
                for key in feature_value_lookup_keys(part):
                    known_values.setdefault(key, part)


def unmatched_product_row(reference: pd.Series, code: str) -> dict[str, Any]:
    override = UNMATCHED_OVERRIDES[code]
    title = reference_value(reference, "Title")
    extracted = extract_attributes_from_text(title, {"default_values": {"producent": "Kanlux"}}, "Kanlux")
    row: dict[str, Any] = {
        "product_id": reference_value(reference, "id"),
        "new_title": title,
        "SKU": reference_value(reference, "Sku"),
        "EAN": normalize_ean(reference.get("hwp_product_gtin", "")),
        "Producent": "Kanlux",
        "attr_producent": "Kanlux",
        "attr_typ": override["type"],
        "proponowana_kategoria_1": override["category"],
        "description_html": override["description"],
        "source_match_status": "UNMATCHED_MANUAL",
    }
    for attribute, value in extracted["attributes"].items():
        row[f"attr_{attribute}"] = value
    for attribute, value in override.get("attributes", {}).items():
        row[f"attr_{attribute}"] = value
    row.update(override.get("fields", {}))
    return row


def build_control_rows(
    reference: pd.DataFrame,
    enriched: pd.DataFrame,
    accepted: pd.DataFrame,
    woo: pd.DataFrame,
    attribute_knowledge: dict[str, Any],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    enriched_by_sku, enriched_by_ean = indexed_rows(enriched, ["Kod", "SKU"], ["EAN"])
    accepted_by_sku, accepted_by_ean = indexed_rows(accepted, ["Kod", "SKU"], ["EAN"])
    woo_by_sku, _ = indexed_rows(woo, ["SKU"], [])
    attribute_aliases = build_attribute_aliases(attribute_knowledge)
    rows: list[dict[str, Any]] = []
    unmatched: list[dict[str, str]] = []

    added_codes: set[str] = set()
    for _, reference_row in reference.iterrows():
        code = normalize_sku(reference_row.get("Sku", ""))
        ean = normalize_ean(reference_row.get("hwp_product_gtin", ""))
        source = enriched_by_sku.get(code)
        if source is None:
            source = enriched_by_ean.get(ean)
        approved = accepted_by_sku.get(code)
        if approved is None:
            approved = accepted_by_ean.get(ean)
        woo_row = woo_by_sku.get(code)
        if woo_row is None:
            continue
        if source is None:
            row = unmatched_product_row(reference_row, code)
            add_woo_feature_fallbacks(row, woo_row, attribute_aliases)
            unmatched.append(
                {
                    "product_id": reference_value(reference_row, "id"),
                    "sku": reference_value(reference_row, "Sku"),
                    "ean": ean,
                    "name": reference_value(reference_row, "Title"),
                    "resolution": "manual_override",
                }
            )
            rows.append(row)
            added_codes.add(code)
            continue

        row = source.to_dict()
        row["pipeline_title"] = compact_spaces(str(source.get("new_title", "")))
        row.update(
            {
                "product_id": reference_value(reference_row, "id"),
                "new_title": reference_value(reference_row, "Title"),
                "SKU": reference_value(reference_row, "Sku"),
                "EAN": ean,
                "Producent": "Kanlux",
                "attr_producent": "Kanlux",
                "source_match_status": "MATCHED_MASTER",
            }
        )
        if approved is not None:
            row["proponowana_kategoria_1"] = compact_spaces(str(approved.get("proponowana_kategoria_1", "")))
            row["description_html"] = str(approved.get("description_html", "")).strip()
        add_woo_feature_fallbacks(row, woo_row, attribute_aliases)
        rows.append(row)
        added_codes.add(code)

    # Produkty istniejace w WooCommerce + w master (enriched), ktorych NIE ma na liscie
    # referencyjnej Baselinkera - wzbogacamy je gotowymi danymi do pliku Woo (nie ida do
    # importu Baselinkera, bo nie maja ID Baselinkera - patrz source_match_status).
    for code, woo_row in woo_by_sku.items():
        if code in added_codes:
            continue
        source = enriched_by_sku.get(code)
        if source is None:
            continue
        ean = normalize_ean(source.get("EAN", "")) or normalize_ean(woo_row.get("EAN", ""))
        approved = accepted_by_sku.get(code)
        if approved is None:
            approved = accepted_by_ean.get(ean)
        row = source.to_dict()
        row["pipeline_title"] = compact_spaces(str(source.get("new_title", "")))
        row.update(
            {
                "product_id": reference_value(woo_row, "id"),
                "new_title": reference_value(woo_row, "Title"),
                "SKU": reference_value(woo_row, "SKU") or code,
                "EAN": ean,
                "Producent": "Kanlux",
                "attr_producent": "Kanlux",
                "source_match_status": "MATCHED_OUTSIDE_REFERENCE",
            }
        )
        if approved is not None:
            row["proponowana_kategoria_1"] = compact_spaces(str(approved.get("proponowana_kategoria_1", "")))
            row["description_html"] = str(approved.get("description_html", "")).strip()
        add_woo_feature_fallbacks(row, woo_row, attribute_aliases)
        rows.append(row)
        added_codes.add(code)
    return pd.DataFrame(rows), pd.DataFrame(unmatched)


def add_export_columns(control: pd.DataFrame, exported: list[dict[str, str]]) -> pd.DataFrame:
    result = control.copy()
    export_df = pd.DataFrame(exported)
    for column in BASELINKER_COLUMNS:
        result[f"baselinker_{column}"] = export_df[column].values
    result["attribute_changes"] = [
        json.dumps(compare_feature_sets(old, new), ensure_ascii=False, separators=(",", ":"))
        for old, new in zip(result["woo_features_json"], export_df["features"])
    ]
    return result


def compare_feature_sets(old_features_json: str, new_features_json: str) -> dict[str, Any]:
    old_features = json.loads(old_features_json or "{}")
    new_features = json.loads(new_features_json or "{}")
    return {
        "added": sorted(name for name in new_features if name not in old_features),
        "changed": sorted(
            name
            for name in new_features.keys() & old_features.keys()
            if compact_spaces(str(new_features[name])) != compact_spaces(str(old_features[name]))
        ),
        "removed": sorted(name for name in old_features if name not in new_features),
    }


def write_attribute_update_csv(rows: list[dict[str, str]], output: str | Path) -> None:
    output_path = Path(output)
    ensure_dir(output_path.parent)
    with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=ATTRIBUTE_UPDATE_COLUMNS, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "") for column in ATTRIBUTE_UPDATE_COLUMNS})


def write_control_workbook(
    path: Path,
    control: pd.DataFrame,
    exported: list[dict[str, str]],
    unmatched: pd.DataFrame,
    summary_rows: list[dict[str, Any]],
) -> None:
    ensure_dir(path.parent)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        control.to_excel(writer, sheet_name="Kontrola", index=False)
        pd.DataFrame(exported, columns=BASELINKER_COLUMNS).to_excel(writer, sheet_name="Baselinker", index=False)
        pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Podsumowanie", index=False)
        unmatched.to_excel(writer, sheet_name="Reczne dopasowania", index=False)
        for sheet_name, worksheet in writer.sheets.items():
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.style = "Headline 4"
            worksheet.column_dimensions["A"].width = 18
            for column in worksheet.columns:
                letter = column[0].column_letter
                if letter == "A":
                    continue
                max_length = max((len(str(cell.value or "")) for cell in column[:200]), default=0)
                worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 60)


def validate_export(exported: list[dict[str, str]], expected_rows: int, review_rows: list[dict[str, str]]) -> tuple[dict[str, int], list[dict[str, str]]]:
    metrics = {
        "rows": len(exported),
        "expected_rows": expected_rows,
        "empty_required": 0,
        "invalid_features_json": 0,
        "duplicate_product_id": 0,
        "duplicate_sku": 0,
        "missing_product_id": 0,
        "missing_ean": 0,
        "missing_images": 0,
        "feature_review_rows": len(review_rows),
    }
    issues: list[dict[str, str]] = []
    required = ["product_id", "name", "sku", "ean", "manufacturer_name", "category", "description", "features"]
    product_ids: set[str] = set()
    skus: set[str] = set()
    for index, row in enumerate(exported, start=2):
        for field in required:
            if not compact_spaces(row.get(field, "")):
                metrics["empty_required"] += 1
                issues.append({"row": str(index), "sku": row.get("sku", ""), "issue": f"EMPTY_{field.upper()}"})
        try:
            json.loads(row.get("features", ""))
        except json.JSONDecodeError:
            metrics["invalid_features_json"] += 1
            issues.append({"row": str(index), "sku": row.get("sku", ""), "issue": "INVALID_FEATURES_JSON"})
        product_id = row.get("product_id", "")
        if not product_id:
            metrics["missing_product_id"] += 1
        elif product_id in product_ids:
            metrics["duplicate_product_id"] += 1
        product_ids.add(product_id)
        sku = row.get("sku", "")
        if sku in skus:
            metrics["duplicate_sku"] += 1
        skus.add(sku)
        if not row.get("ean", ""):
            metrics["missing_ean"] += 1
        if not row.get("images_urls", ""):
            metrics["missing_images"] += 1
    return metrics, issues


def write_reports(
    reports_dir: Path,
    metrics: dict[str, int],
    issues: list[dict[str, str]],
    unmatched: pd.DataFrame,
    review_rows: list[dict[str, str]],
    changes: list[dict[str, Any]],
    missing_required: list[dict[str, Any]] | None = None,
    power_voltage_audit: list[dict[str, Any]] | None = None,
    manual_fills_applied: list[dict[str, Any]] | None = None,
    duplicate_disambiguated: list[dict[str, Any]] | None = None,
    remaining_duplicates: list[dict[str, Any]] | None = None,
) -> None:
    ensure_dir(reports_dir)
    (reports_dir / "summary.json").write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")
    pd.DataFrame([{"metric": key, "value": value} for key, value in metrics.items()]).to_csv(
        reports_dir / "validation_summary.csv", index=False, encoding="utf-8-sig"
    )
    pd.DataFrame(issues, columns=["row", "sku", "issue"]).to_csv(
        reports_dir / "validation_issues.csv", index=False, encoding="utf-8-sig"
    )
    unmatched.to_csv(reports_dir / "manual_matches.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(changes).to_csv(reports_dir / "attribute_changes.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(
        manual_fills_applied or [],
        columns=["sku", "attribute", "old_value", "new_value"],
    ).to_csv(reports_dir / "manual_fills_applied.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(
        missing_required or [],
        columns=["product_id", "sku", "name", "category", "rule_category", "status", "missing_required"],
    ).to_csv(reports_dir / "missing_required_attributes.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(
        power_voltage_audit or [],
        columns=[
            "product_id",
            "sku",
            "name",
            "moc_eksport",
            "moc_max_zarowki",
            "moc_woo",
            "moc_zrodlo",
            "napiecie_eksport",
            "napiecie_woo",
            "napiecie_zrodlo",
            "trzonek",
            "zrodlo_swiatla",
            "flagi",
        ],
    ).to_csv(reports_dir / "power_voltage_audit.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(
        duplicate_disambiguated or [],
        columns=["sku", "old_name", "new_name", "group_key"],
    ).to_csv(reports_dir / "duplicate_disambiguated.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(
        remaining_duplicates or [],
        columns=["group_key", "sku", "name", "count"],
    ).to_csv(reports_dir / "duplicate_names_remaining.csv", index=False, encoding="utf-8-sig")
    write_feature_review_file(review_rows, reports_dir / "feature_review_required.xlsx")


# Atrybuty, ktorych brak nie jest brakiem: pusty czujnik ruchu znaczy "nie ma",
# a liczba zrodel swiatla dotyczy tylko opraw na wymienne zrodla.
REQUIRED_ATTRIBUTES_OPTIONAL_WHEN_ABSENT = {"Czujnik ruchu", "Liczba gniazd", "Liczba źródeł światła"}

# Przy wymiennym zrodle swiatla moc, barwa i strumien zaleza od wkreconej zarowki,
# wiec nie sa atrybutami wymaganymi produktu.
REQUIRED_ATTRIBUTES_DEPENDENT_ON_LIGHT_SOURCE = {
    "Moc [W]",
    "Barwa światła",
    "Strumień świetlny [lm]",
    "Temperatura barwowa [K]",
}

# Przy zintegrowanym zrodle nie ma trzonka, liczby wymiennych zrodel ani doboru kierunku.
REQUIRED_ATTRIBUTES_NOT_FOR_INTEGRATED_SOURCE = {
    "Trzonek",
    "Liczba źródeł światła",
    "Liczba gniazd",
    "Kierunek świecenia",
}


def build_missing_required_rows(
    control: pd.DataFrame,
    exported: list[dict[str, str]],
    attribute_knowledge: dict[str, Any],
) -> list[dict[str, Any]]:
    """Produkty, ktorym brakuje atrybutow obowiazkowych dla ich typu (macierz Poprawione atrybuty)."""
    rows: list[dict[str, Any]] = []
    for (_, control_row), exported_row in zip(control.iterrows(), exported):
        features = json.loads(exported_row.get("features", "{}") or "{}")
        if is_accessory_product_type(str(features.get("Typ produktu", ""))):
            continue
        rule = match_attribute_category_rule(control_row, exported_row.get("category", ""), attribute_knowledge, features)
        if not rule:
            rows.append(
                {
                    "product_id": exported_row["product_id"],
                    "sku": exported_row["sku"],
                    "name": exported_row["name"],
                    "category": exported_row.get("category", ""),
                    "rule_category": "",
                    "status": "BRAK_REGULY",
                    "missing_required": "",
                }
            )
            continue
        present = {normalize_header(name) for name in features}
        row_skip = row_skip_for_features(features)
        missing = [
            str(name)
            for name in rule.get("required_description") or []
            if normalize_header(str(name)) not in present and normalize_header(str(name)) not in row_skip
        ]
        if missing:
            rows.append(
                {
                    "product_id": exported_row["product_id"],
                    "sku": exported_row["sku"],
                    "name": exported_row["name"],
                    "category": exported_row.get("category", ""),
                    "rule_category": str(rule.get("category", "")),
                    "status": "BRAKI_WYMAGANYCH",
                    "missing_required": " | ".join(missing),
                }
            )
    return rows


def apply_pipeline_titles_for_mismatched_names(
    control: pd.DataFrame,
    exported: list[dict[str, str]],
) -> list[dict[str, str]]:
    """Gdy nazwa referencyjna kloci sie z atrybutami, bierzemy nazwe z pipeline'u.

    Przyklad: stare Woo nazywa panel "Zasilacz BLINGO AIO 40W" - pipeline nazw
    wygenerowal poprawna nazwe panelowa i to ona idzie do importu. To samo dotyczy
    starych parametrow w tytule, np. IP12 przy finalnym atrybucie IP 54|IP 20.
    """
    renamed: list[dict[str, str]] = []
    for (_, control_row), exported_row in zip(control.iterrows(), exported):
        features = json.loads(exported_row.get("features", "{}") or "{}")
        product_type = compact_spaces(str(features.get("Typ produktu", "")))
        name_type = product_type_from_title(exported_row.get("name", ""))
        if not product_type or not name_type:
            continue
        pipeline_title = compact_spaces(str(control_row.get("pipeline_title", "") or ""))
        if pipeline_title.lower() in {"nan", "none"}:
            pipeline_title = ""
        pipeline_type = product_type_from_title(pipeline_title)
        if not pipeline_title or not pipeline_type or not product_types_agree(pipeline_type, product_type):
            continue
        reasons: list[str] = []
        if not product_types_agree(name_type, product_type):
            reasons.append("TYPE_MISMATCH")
        if title_has_feature_conflict(exported_row.get("name", ""), pipeline_title, features):
            reasons.append("FEATURE_MISMATCH")
        if reasons:
            renamed.append(
                {
                    "sku": exported_row.get("sku", ""),
                    "old_name": exported_row.get("name", ""),
                    "new_name": pipeline_title,
                    "reason": "|".join(reasons),
                }
            )
            exported_row["name"] = pipeline_title
    return renamed


def strip_product_code(name: str) -> str:
    """Ucina koncowy kod produktu ("... Kanlux 36504"), zostawiajac sama nazwe."""
    return re.sub(r"\s*Kanlux\s+[\w/.\-]+\s*$", "", compact_spaces(name), flags=re.IGNORECASE).strip()


def duplicate_name_key(name: str) -> str:
    return normalize_title_uniqueness_key(strip_product_code(name))


def apply_pipeline_titles_for_duplicate_names(
    control: pd.DataFrame,
    exported: list[dict[str, str]],
) -> list[dict[str, str]]:
    """Gdy kilka produktow ma identyczna nazwe (po odcieciu kodu), bierzemy
    rozrozniajaca nazwe z pipeline'u.

    Stary tytul legacy ze starego Woo gubil wymiary/warianty/czujnik, ktore
    pipeline juz policzyl (np. "Plafon ERTON E27 IP44 bialy" -> "...252mm..."),
    ale chain trzymal sie nazwy referencyjnej, bo nie bylo konfliktu typu/atrybutu.
    """
    pipeline_titles: list[str] = []
    for _, control_row in control.iterrows():
        title = compact_spaces(str(control_row.get("pipeline_title", "") or ""))
        if title.lower() in {"nan", "none"}:
            title = ""
        pipeline_titles.append(title)

    groups: dict[str, list[int]] = {}
    for index, exported_row in enumerate(exported):
        groups.setdefault(duplicate_name_key(exported_row.get("name", "")), []).append(index)

    renamed: list[dict[str, str]] = []
    for group_key, indexes in groups.items():
        if not group_key or len(indexes) < 2:
            continue
        for index in indexes:
            exported_row = exported[index]
            pipeline_title = pipeline_titles[index]
            if not pipeline_title:
                continue
            current = exported_row.get("name", "")
            name_type = product_type_from_title(current)
            pipeline_type = product_type_from_title(pipeline_title)
            if name_type and pipeline_type and not product_types_agree(name_type, pipeline_type):
                continue
            # Pipeline nie rozroznia tego produktu - nie podmieniamy (zostaje duplikatem).
            if duplicate_name_key(pipeline_title) == group_key:
                continue
            features = json.loads(exported_row.get("features", "{}") or "{}")
            new_name = hermetic_fluorescent_title(normalize_seo_title_terms(pipeline_title), features)
            if compact_spaces(new_name) == compact_spaces(current):
                continue
            renamed.append(
                {
                    "sku": exported_row.get("sku", ""),
                    "old_name": current,
                    "new_name": new_name,
                    "group_key": group_key,
                }
            )
            exported_row["name"] = new_name
    return renamed


def remaining_duplicate_name_rows(exported: list[dict[str, str]]) -> list[dict[str, str]]:
    """Produkty, ktore po dysambiguacji nadal maja identyczna nazwe (do scalenia w sklepie)."""
    groups: dict[str, list[dict[str, str]]] = {}
    for exported_row in exported:
        groups.setdefault(duplicate_name_key(exported_row.get("name", "")), []).append(exported_row)
    rows: list[dict[str, str]] = []
    for group_key, members in groups.items():
        if not group_key or len(members) < 2:
            continue
        for member in members:
            rows.append(
                {
                    "group_key": group_key,
                    "sku": member.get("sku", ""),
                    "name": member.get("name", ""),
                    "count": str(len(members)),
                }
            )
    return rows


def title_has_feature_conflict(title: str, pipeline_title: str, features: dict[str, Any]) -> bool:
    return (
        title_mounting_conflicts_with_feature(title, pipeline_title, features)
        or title_missing_pipeline_identity(title, pipeline_title)
        or title_word_feature_missing(title, pipeline_title, features, "ksztalt")
        or title_word_feature_missing(title, pipeline_title, features, "kolor")
        or title_word_feature_missing(title, pipeline_title, features, "trzonek")
        or title_word_feature_missing(title, pipeline_title, features, "czujnik ruchu")
        or any(
        title_token_conflicts_with_feature(title, pipeline_title, feature_value, pattern)
        for feature_value, pattern in [
            (feature_by_normalized_name(features, "stopien ochrony ip"), r"\bIP\s*(\d{2})\b"),
            (feature_by_normalized_name(features, "moc w"), r"\b(\d+(?:[,.]\d+)?)\s*W\b"),
            (feature_by_normalized_name(features, "temperatura barwowa k"), r"\b(\d{4})\s*K\b"),
            (feature_by_normalized_name(features, "strumien swietlny lm"), r"\b(\d+(?:-\d+)?)\s*lm\b"),
            (feature_by_normalized_name(features, "wymiary mm"), r"\b(\d+(?:[xX]\d+){1,2})\s*(?:mm|cm)?\b"),
        ]
    )
    )


def title_missing_pipeline_identity(title: str, pipeline_title: str) -> bool:
    current = normalize_header(title)
    pipeline = normalize_header(pipeline_title)
    if not current or not pipeline:
        return False
    for term in ["asymetryczny", "symetryczny"]:
        if term in pipeline and term not in current:
            return True
    protected_suffixes = ["sky v", "sky", "pt", "pro", "hi", "aio"]
    protected_prefixes = ["adtr", "fl agor", "varso", "daba"]
    for prefix in protected_prefixes:
        for suffix in protected_suffixes:
            phrase = f"{prefix} {suffix}"
            if phrase in pipeline and phrase not in current:
                return True
    return False


def title_mounting_conflicts_with_feature(title: str, pipeline_title: str, features: dict[str, Any]) -> bool:
    current = mounting_from_title(title)
    expected = normalize_header(feature_by_normalized_name(features, "sposob montazu"))
    pipeline = mounting_from_title(pipeline_title)
    return bool(current and expected and current != expected and pipeline == expected)


def mounting_from_title(title: str) -> str:
    normalized = normalize_header(title)
    if "podtynkowy" in normalized:
        return "podtynkowy"
    if "natynkowy" in normalized:
        return "natynkowy"
    return ""


def title_word_feature_missing(title: str, pipeline_title: str, features: dict[str, Any], normalized_name: str) -> bool:
    feature_value = feature_by_normalized_name(features, normalized_name)
    if normalized_name == "kolor":
        # Kolor porownujemy po rdzeniu (szary/szara/szare -> szar), zeby wykryc realny
        # konflikt koloru obudowy niezaleznie od formy gramatycznej.
        stem = color_stem(feature_value)
        if stem:
            return stem in normalize_header(pipeline_title) and stem not in normalize_header(title)
    expected = normalized_title_words(feature_value)
    if not expected:
        return False
    current = normalize_header(title)
    pipeline = normalize_header(pipeline_title)
    return any(word in pipeline and word not in current for word in expected)


def normalized_title_words(value: str) -> set[str]:
    ignored = {"z", "ze", "do", "bez", "nie", "tak", "na", "w", "i"}
    return {
        word
        for word in re.split(r"\s+", normalize_header(value))
        if len(word) > 1 and word not in ignored
    }


def feature_by_normalized_name(features: dict[str, Any], normalized_name: str) -> str:
    for name, value in features.items():
        if normalize_header(str(name)) == normalized_name:
            return compact_spaces(str(value))
    return ""


def title_token_conflicts_with_feature(title: str, pipeline_title: str, feature_value: str, pattern: str) -> bool:
    current = normalized_title_tokens(title, pattern)
    expected = normalized_title_tokens(feature_value, pattern)
    if not expected:
        return False
    pipeline = normalized_title_tokens(pipeline_title, pattern)
    if not expected <= pipeline:
        return False
    if not current:
        return True
    return not expected <= current


def normalized_title_tokens(text: str, pattern: str) -> set[str]:
    tokens: set[str] = set()
    for match in re.findall(pattern, compact_spaces(str(text)), flags=re.IGNORECASE):
        value = match[0] if isinstance(match, tuple) else match
        tokens.add(value.replace(",", ".").replace(" ", "").upper())
    return tokens


def align_panel_mounting_in_names(exported: list[dict[str, str]]) -> int:
    """Nazwa panelu ma zgadzac sie z (recznie zatwierdzonym) sposobem montazu."""
    fixed = 0
    for exported_row in exported:
        features = json.loads(exported_row.get("features", "{}") or "{}")
        if features.get("Typ produktu") != "Panel LED":
            continue
        mounting = features.get("Sposób montażu", "")
        name = exported_row.get("name", "")
        if mounting == "Natynkowy" and re.search(r"\bpodtynkowy\b", name):
            exported_row["name"] = re.sub(r"\bpodtynkowy\b", "natynkowy", name)
            fixed += 1
        elif mounting == "Podtynkowy" and re.search(r"\bnatynkowy\b", name):
            exported_row["name"] = re.sub(r"\bnatynkowy\b", "podtynkowy", name)
            fixed += 1
    return fixed


def load_manual_attribute_fills(path: str | Path) -> dict[str, dict[str, str]]:
    """Wczytuje recznie uzupelniony skoroszyt brakow (arkusz per kategoria).

    Puste komorki i "n/d" sa pomijane; reszta to zatwierdzone przez czlowieka wartosci.
    """
    fills: dict[str, dict[str, str]] = {}
    fill_path = Path(path)
    if not fill_path.exists():
        return fills
    from openpyxl import load_workbook

    workbook = load_workbook(fill_path, read_only=True)
    for sheet in workbook.worksheets:
        headers = [compact_spaces(str(cell.value or "")) for cell in sheet[1]]
        if headers[:4] != ["SKU", "EAN", "Nazwa", "Kategoria sklepu"]:
            continue
        attributes = headers[4:]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            sku = compact_spaces(str(row[0] or ""))
            if not sku:
                continue
            for offset, attribute in enumerate(attributes):
                value = compact_spaces(str(row[4 + offset] if 4 + offset < len(row) and row[4 + offset] is not None else ""))
                if not value or value.lower() in {"n/d", "nd", "nan", "none"}:
                    continue
                fills.setdefault(sku, {})[attribute] = value
    workbook.close()
    return fills


def apply_manual_attribute_fills(
    exported: list[dict[str, str]],
    fills: dict[str, dict[str, str]],
) -> list[dict[str, Any]]:
    """Reczne uzupelnienia wygrywaja z danymi automatycznymi."""
    applied: list[dict[str, Any]] = []
    for exported_row in exported:
        sku_fills = fills.get(exported_row.get("sku", ""))
        if not sku_fills:
            continue
        features = json.loads(exported_row.get("features", "{}") or "{}")
        changed = False
        for attribute, raw_value in sku_fills.items():
            value = normalize_feature_value(raw_value, attribute) or raw_value
            if compact_spaces(str(features.get(attribute, ""))) == value:
                continue
            applied.append(
                {
                    "sku": exported_row.get("sku", ""),
                    "attribute": attribute,
                    "old_value": features.get(attribute, ""),
                    "new_value": value,
                }
            )
            features[attribute] = value
            changed = True
        if changed:
            exported_row["features"] = json.dumps(features, ensure_ascii=False, separators=(",", ":"))
    return applied


def row_skip_for_features(features: dict[str, Any]) -> set[str]:
    """Atrybuty wymagane, ktorych nie liczymy jako brak dla danego produktu."""
    skip = {normalize_header(name) for name in REQUIRED_ATTRIBUTES_OPTIONAL_WHEN_ABSENT}
    light_source = compact_spaces(str(features.get("Źródło światła", ""))).lower()
    replaceable_source = (
        "Maksymalna moc źródła światła" in features
        or light_source in {"wymienne", "nie zintegrowane", "niezintegrowane"}
    )
    if replaceable_source:
        skip.update(normalize_header(name) for name in REQUIRED_ATTRIBUTES_DEPENDENT_ON_LIGHT_SOURCE)
    if light_source == "zintegrowane":
        skip.update(normalize_header(name) for name in REQUIRED_ATTRIBUTES_NOT_FOR_INTEGRATED_SOURCE)
    return skip


def write_missing_required_workbook(
    path: Path,
    control: pd.DataFrame,
    exported: list[dict[str, str]],
    attribute_knowledge: dict[str, Any],
) -> int:
    """Skoroszyt do recznego uzupelnienia brakow: arkusz per kategoria, braki na zolto."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    missing_fill = PatternFill("solid", fgColor="FFF3B0")
    not_applicable_font = Font(color="9AA0A6", italic=True)

    grouped: dict[str, dict[str, Any]] = {}
    for (_, control_row), exported_row in zip(control.iterrows(), exported):
        features = json.loads(exported_row.get("features", "{}") or "{}")
        if is_accessory_product_type(str(features.get("Typ produktu", ""))):
            continue
        rule = match_attribute_category_rule(control_row, exported_row.get("category", ""), attribute_knowledge, features)
        if not rule:
            continue
        required = [str(name) for name in rule.get("required_description") or []]
        if not required:
            continue
        by_normalized = {normalize_header(name): value for name, value in features.items()}
        skip = row_skip_for_features(features)
        missing = [
            name for name in required
            if normalize_header(name) not in by_normalized and normalize_header(name) not in skip
        ]
        if not missing:
            continue
        group = grouped.setdefault(
            str(rule.get("category", "")),
            {"required": required, "products": []},
        )
        group["products"].append((exported_row, by_normalized, skip))

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Podsumowanie"
    summary.append(["Kategoria (reguła)", "Produkty z brakami", "Atrybuty wymagane"])
    for category, group in sorted(grouped.items(), key=lambda item: -len(item[1]["products"])):
        summary.append([category, len(group["products"]), ", ".join(group["required"])])
    summary.append([])
    summary.append(["Legenda:", "żółte pole = brak do uzupełnienia", "„n/d” = zależne od wkręcanego źródła światła"])

    total = 0
    for category, group in sorted(grouped.items(), key=lambda item: -len(item[1]["products"])):
        required = group["required"]
        sheet_title = re.sub(r"[\\/*?\[\]:]", " ", category)[:31].strip() or "Kategoria"
        sheet = workbook.create_sheet(sheet_title)
        sheet.append(["SKU", "EAN", "Nazwa", "Kategoria sklepu", *required])
        for exported_row, by_normalized, skip in group["products"]:
            total += 1
            row_index = sheet.max_row + 1
            sheet.cell(row_index, 1, exported_row.get("sku", ""))
            sheet.cell(row_index, 2, exported_row.get("ean", ""))
            sheet.cell(row_index, 3, exported_row.get("name", ""))
            sheet.cell(row_index, 4, exported_row.get("category", ""))
            for offset, attribute in enumerate(required):
                cell = sheet.cell(row_index, 5 + offset)
                normalized = normalize_header(attribute)
                if normalized in by_normalized:
                    cell.value = str(by_normalized[normalized])
                elif normalized in skip:
                    cell.value = "n/d"
                    cell.font = not_applicable_font
                else:
                    cell.fill = missing_fill

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column_index in range(1, sheet.max_column + 1):
            letter = get_column_letter(column_index)
            width = max(
                (len(str(cell.value or "")) for cell in sheet[letter][:200]),
                default=0,
            )
            sheet.column_dimensions[letter].width = min(max(width + 2, 12), 55)

    ensure_dir(path.parent)
    workbook.save(path)
    return total


def build_power_voltage_audit_rows(
    control: pd.DataFrame,
    exported: list[dict[str, str]],
) -> list[dict[str, Any]]:
    """Audyt mocy i napiecia: flaguje wartosci, ktore moga pochodzic z podobnych, ale innych parametrow."""
    rows: list[dict[str, Any]] = []
    for (_, control_row), exported_row in zip(control.iterrows(), exported):
        features = json.loads(exported_row.get("features", "{}") or "{}")
        moc = compact_spaces(str(features.get("Moc [W]", "")))
        moc_max = compact_spaces(str(features.get("Maksymalna moc źródła światła", "")))
        napiecie = compact_spaces(str(features.get("Napięcie [V]", "")))
        trzonek = compact_spaces(str(features.get("Trzonek", "")))
        zrodlo = compact_spaces(str(features.get("Źródło światła", "")))
        try:
            sources = json.loads(str(control_row.get("attribute_sources", "{}") or "{}"))
        except (json.JSONDecodeError, TypeError):
            sources = {}

        def clean_cell(value: Any) -> str:
            text = compact_spaces(str(value if value is not None else ""))
            return "" if text.lower() in {"nan", "none", "null"} else text

        woo_moc = clean_cell(control_row.get("Parametr: Moc [W]", ""))
        woo_napiecie = clean_cell(control_row.get("Parametr: Napięcie [V]", ""))

        flags: list[str] = []
        if moc and re.search(r"max", moc, flags=re.IGNORECASE):
            flags.append("MOC_Z_PRZEDROSTKIEM_MAX")
        if moc and trzonek and not moc_max and "led" not in zrodlo.lower():
            flags.append("MOC_PRZY_WYMIENNYM_ZRODLE_DO_WERYFIKACJI")
        # Roznica mocy wzgledem starego Woo nie jest flagowana: parametry producenta
        # sa zrodlem prawdy (kolumna moc_woo zostaje w raporcie tylko dla kontekstu).
        if napiecie and not re.fullmatch(r"\d+(?:[,.]\d+)?(?:-\d+(?:[,.]\d+)?)?(?:\s(?:AC|DC))?", napiecie):
            flags.append("NAPIECIE_NIETYPOWE")
        if not napiecie and woo_napiecie:
            flags.append("NAPIECIE_TYLKO_W_WOO")
        if not flags:
            continue
        rows.append(
            {
                "product_id": exported_row["product_id"],
                "sku": exported_row["sku"],
                "name": exported_row["name"],
                "moc_eksport": moc,
                "moc_max_zarowki": moc_max,
                "moc_woo": woo_moc,
                "moc_zrodlo": str(sources.get("moc", "")),
                "napiecie_eksport": napiecie,
                "napiecie_woo": woo_napiecie,
                "napiecie_zrodlo": str(sources.get("napiecie", "")),
                "trzonek": trzonek,
                "zrodlo_swiatla": zrodlo,
                "flagi": " | ".join(flags),
            }
        )
    return rows


def build_change_rows(control: pd.DataFrame, exported: list[dict[str, str]]) -> list[dict[str, Any]]:
    changes: list[dict[str, Any]] = []
    for (_, row), exported_row in zip(control.iterrows(), exported):
        comparison = compare_feature_sets(str(row.get("woo_features_json", "{}")), exported_row["features"])
        changes.append(
            {
                "product_id": exported_row["product_id"],
                "sku": exported_row["sku"],
                "woo_product_id": row.get("woo_product_id", ""),
                "added_count": len(comparison["added"]),
                "changed_count": len(comparison["changed"]),
                "removed_count": len(comparison["removed"]),
                "added": " | ".join(comparison["added"]),
                "changed": " | ".join(comparison["changed"]),
                "removed": " | ".join(comparison["removed"]),
            }
        )
    return changes


def main() -> None:
    parser = argparse.ArgumentParser(description="Buduje aktualizacyjny import Kanlux do Baselinkera.")
    parser.add_argument("--reference", default=DEFAULT_REFERENCE)
    parser.add_argument("--woo", default=DEFAULT_WOO)
    parser.add_argument("--enriched", default=DEFAULT_ENRICHED)
    parser.add_argument("--accepted", default=DEFAULT_ACCEPTED)
    parser.add_argument("--control-output", default=DEFAULT_CONTROL)
    parser.add_argument("--csv-output", default=DEFAULT_CSV)
    parser.add_argument("--attribute-csv-output", default=DEFAULT_ATTRIBUTE_CSV)
    parser.add_argument("--missing-required-xlsx", default=DEFAULT_MISSING_XLSX)
    parser.add_argument("--manual-fills", default=DEFAULT_MANUAL_FILLS)
    parser.add_argument("--reports-dir", default=DEFAULT_REPORTS)
    parser.add_argument("--catalog-knowledge", default="dictionaries/woocommerce_catalog_knowledge.yaml")
    parser.add_argument("--supplier-knowledge", default="dictionaries/supplier_global_knowledge.yaml")
    parser.add_argument("--attribute-knowledge", default="dictionaries/lighting_attribute_knowledge.yaml")
    parser.add_argument("--taxonomy", default="dictionaries/learned_product_taxonomy.yaml")
    args = parser.parse_args()

    reference = read_products(args.reference)
    woo = read_products(args.woo)
    enriched = read_products(args.enriched)
    accepted = read_products(args.accepted)
    attribute_knowledge = load_yaml(args.attribute_knowledge)
    control, unmatched = build_control_rows(reference, enriched, accepted, woo, attribute_knowledge)

    catalog = load_yaml(args.catalog_knowledge)
    attribute_aliases = build_attribute_aliases(attribute_knowledge)
    feature_value_normalizer = build_feature_value_normalizer(catalog)
    add_woo_values_to_feature_normalizer(feature_value_normalizer, woo, attribute_aliases)
    review_rows: list[dict[str, str]] = []
    exported = build_baselinker_rows(
        control,
        producer_suffixes=load_producer_suffixes(args.taxonomy, args.catalog_knowledge),
        manufacturer_data_by_producer=build_manufacturer_data_by_producer(catalog),
        feature_value_normalizer=feature_value_normalizer,
        feature_review_rows=review_rows,
        feature_policy=load_supplier_feature_policy(args.supplier_knowledge),
        attribute_knowledge=attribute_knowledge,
    )
    renamed_from_pipeline = apply_pipeline_titles_for_mismatched_names(control, exported)
    for exported_row in exported:
        # Normalizacja SEO (np. "Plafon drewniana" -> "Plafon drewniany") dotyczy
        # takze tytulow legacy zachowanych ze starego Woo, nie tylko przebudowanych.
        name = normalize_seo_title_terms(exported_row["name"])
        exported_row["name"] = hermetic_fluorescent_title(
            name, json.loads(exported_row["features"] or "{}")
        )
    manual_fills_applied = apply_manual_attribute_fills(exported, load_manual_attribute_fills(args.manual_fills))
    metrics, issues = validate_export(exported, len(control), review_rows)
    metrics["manual_fills_applied"] = len(manual_fills_applied)
    metrics["renamed_from_pipeline"] = len(renamed_from_pipeline)
    metrics["panel_mounting_names_fixed"] = align_panel_mounting_in_names(exported)
    duplicate_disambiguated = apply_pipeline_titles_for_duplicate_names(control, exported)
    metrics["duplicate_disambiguated"] = len(duplicate_disambiguated)
    remaining_duplicates = remaining_duplicate_name_rows(exported)
    metrics["duplicate_names_remaining_groups"] = len({row["group_key"] for row in remaining_duplicates})
    metrics["woo_matches"] = int((control["woo_match_status"] == "MATCHED").sum())
    metrics["woo_missing_matches"] = int((control["woo_match_status"] != "MATCHED").sum())
    control_with_export = add_export_columns(control, exported)
    changes = build_change_rows(control, exported)
    metrics["products_with_added_attributes"] = sum(row["added_count"] > 0 for row in changes)
    metrics["products_with_changed_attributes"] = sum(row["changed_count"] > 0 for row in changes)
    missing_required = build_missing_required_rows(control, exported, attribute_knowledge)
    metrics["missing_required_workbook_rows"] = write_missing_required_workbook(
        Path(args.missing_required_xlsx), control, exported, attribute_knowledge
    )
    power_voltage_audit = build_power_voltage_audit_rows(control, exported)
    metrics["products_missing_required_attributes"] = sum(
        row["status"] == "BRAKI_WYMAGANYCH" for row in missing_required
    )
    metrics["products_without_required_rule"] = sum(row["status"] == "BRAK_REGULY" for row in missing_required)
    metrics["power_voltage_audit_rows"] = len(power_voltage_audit)
    summary_rows = [{"metric": key, "value": value} for key, value in metrics.items()]

    # Import Baselinkera tylko dla produktow z listy referencyjnej (maja ID Baselinkera).
    # Produkty spoza referencji ida wylacznie do pliku Woo (przez atrybutowy CSV).
    baselinker_export = [row for row in exported if not row.get("_outside_reference")]
    metrics["outside_reference_woo_only"] = len(exported) - len(baselinker_export)
    for row in exported:
        row.pop("_outside_reference", None)
    write_baselinker_csv(baselinker_export, args.csv_output)
    write_attribute_update_csv(exported, args.attribute_csv_output)
    write_control_workbook(Path(args.control_output), control_with_export, exported, unmatched, summary_rows)
    write_reports(
        Path(args.reports_dir),
        metrics,
        issues,
        unmatched,
        review_rows,
        changes,
        missing_required,
        power_voltage_audit,
        manual_fills_applied,
        duplicate_disambiguated,
        remaining_duplicates,
    )
    print(json.dumps(metrics, ensure_ascii=True, sort_keys=True))


if __name__ == "__main__":
    main()
