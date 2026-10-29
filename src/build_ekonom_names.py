"""Krotkie nazwy produktow do Ekonoma (dla sprzedawcow, nie SEO).

Wejscie: CSV z Ekonoma (cp1250, ';': LP; Kod towaru; Nazwa towaru; Nazwa dodatkowa;
Kod producenta; ...) i katalog PDF-ow dostawcy. Parametry (moc, barwa, gwint,
dlugosc...) sa czytane z nazwy, lumeny z tabel PDF po EAN (`Kod producenta`).
Reguly typow: dictionaries/ekonom_name_rules.yaml.

Wynik:
- CSV dla Ekonoma w tym samym formacie co wejscie, z nowa nazwa podzielona na
  `Nazwa towaru` (max 40 znakow, ciecie na spacji) i `Nazwa dodatkowa` (reszta),
- skoroszyt kontrolny: wszystkie nazwy ze statusem, arkusz "Powtorzenia" z grupami
  identycznych nazw i tym, czym roznia sie ich nazwy oryginalne.

Inne zrodla parametrow: feed Ceneo (`--feed`, Kobi) albo pliki Parametry*.xlsx producenta
(`--params-dir`, Kanlux - format dlugi kod/atrybut/wartosc; parametry nadpisuja nazwe).
Katalog PDF bez EAN (GTV) jest czytany po kodzie producenta (`pdf_match: code` w YAML).

Przyklad:
    py src/build_ekonom_names.py --supplier ledvance \
        --input "input/Nazwy do ekonom/Ledvance/Ledvance.csv" \
        --pdf-dir "input/Nazwy do ekonom/Ledvance"
    py src/build_ekonom_names.py --supplier kanlux \
        --input "input/Nazwy do ekonom/Kanlux/Kanlux.csv" \
        --params-dir "input/Nazwy do ekonom/Kanlux"
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import unicodedata
from collections import Counter, defaultdict
from itertools import combinations
from datetime import date
from pathlib import Path
from typing import Any

from utils import ensure_dir, load_yaml

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RULES = ROOT / "dictionaries" / "ekonom_name_rules.yaml"
NAME_LIMIT = 40

EAN_RE = re.compile(r"^\d{13}$")
# 16 / 17: Vintage 1906 "816 Blue" / "817 Smoke" = 1600 / 1700 K (2026-10-21: bez barwy w nazwie)
VALID_CCT = {"16", "17", "18", "20", "22", "24", "25", "27", "30", "35", "40", "50", "57", "60", "65", "80"}
LIFETIMES = {6000, 8000, 10000, 12000, 15000, 20000, 25000, 30000, 35000, 40000, 50000, 60000, 70000, 100000}

SOCKET_RE = re.compile(
    r"(?<![A-Z])("
    r"GX24Q\.?-?[1-5]|GX24D\.?-?[1-3]|G24Q\.?-?[1-4]|G24D\.?-?[1-3]|GX24Q|GX24D|G24Q|G24D|"
    r"2GX13|2G11|2G10|2G7|GU5[.,]3|GX5[.,]3|GY6[.,]35|GX6[.,]35|GY9[.,]5|G5[.,]3|GU10|GU11|GU4|GX53|GR10Q|GR8|"
    r"S14D|S14S|S15S|E27|E14|E40|E10|B22D|B15D|BA15D|R7S|K12S|G13|G23|G53|G12|G9|G4|G5"
    r")(?![A-Z0-9.])"
)
# Gwint sklejony z kodem OSRAM: "FRE144X1" (E14 + 4X1), "B22DFS1", "BKE274X1".
SOCKET_GLUED_RE = re.compile(r"(?:(?<=[A-Z])|(?<=\s))(E27|E14|E40|B22D|GU10)(?=\d+X\d|FS\d|\b)")


# ---------------------------------------------------------------- tekst

def ascii_upper(value: str) -> str:
    value = value.replace("ł", "l").replace("Ł", "L")
    norm = unicodedata.normalize("NFKD", value)
    return "".join(ch for ch in norm if not unicodedata.combining(ch)).upper()


def fmt_num(value: float) -> str:
    text = f"{value:.2f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")


def to_float(text: str) -> float:
    return float(text.replace(",", "."))


PL_MONTHS = ["sty", "lut", "mar", "kwi", "maj", "cze", "lip", "sie", "wrz", "paź", "lis", "gru"]


def undo_excel_date(cell: str) -> str:
    """Komorka zamieniona przez Excela na date: "7/5" -> "05-lip", "27/5" -> "27-maj", "40/5" -> "maj-40".

    Excel czyta "a/b" najpierw jako miesiac/dzien, potem dzien/miesiac, a na koncu rok/miesiac - koncowki kodow tasm
    ("LS PFM-1000/COB/92" + "7/5" = 927/5) wracaja w ten sam sposob (przeglad 2026-10-21).
    """
    m = re.fullmatch(r"(\d{2})-([a-zź]{3})", cell.strip())
    if m and m.group(2) in PL_MONTHS:
        day, month = int(m.group(1)), PL_MONTHS.index(m.group(2)) + 1
        return f"{month}/{day}" if day <= 12 else f"{day}/{month}"
    m = re.fullmatch(r"([a-zź]{3})-(\d{2})", cell.strip())
    if m and m.group(1) in PL_MONTHS:
        return f"{int(m.group(2))}/{PL_MONTHS.index(m.group(1)) + 1}"
    return cell


def full_name(row: dict[str, str]) -> str:
    # Ekonom tnie nazwe na 40. znaku i reszte wpisuje do `Nazwa dodatkowa` bez separatora.
    return " ".join((row.get("Nazwa towaru", "") + undo_excel_date(row.get("Nazwa dodatkowa", ""))).split())


def split_for_ekonom(name: str, limit: int = NAME_LIMIT) -> tuple[str, str]:
    if len(name) <= limit:
        return name, ""
    cut = name.rfind(" ", 0, limit + 1)
    if cut <= 0:
        return name[:limit], name[limit:]
    # Spacja zostaje na poczatku drugiej czesci, jak w eksporcie Ekonoma ("Clear" + " E27").
    return name[:cut], name[cut:]


# ---------------------------------------------------------------- parametry z nazwy

def parse_power(u: str) -> tuple[str, float | None]:
    """Zwraca (tekst do nazwy, moc pojedyncza do weryfikacji lm albo None)."""
    # kod barwy przyklejony do mocy: "13.1W/15.6W830LEDV" - inaczej druga moc nie byla czytana (2026-10-21)
    u = re.sub(r"W(?=[6-9]\d{2}(?!\d))", "W ", u)
    m = re.search(r"(?<![\d.,])(\d+)\s?X\s?(\d+(?:[.,]\d+)?)\s?W(?![A-Z])", u)
    if m:
        return f"{m.group(1)}x{fmt_num(to_float(m.group(2)))}W", None
    # moc przelaczana: "900/1200W", "3/5/7W", "20-30-40W", "30W/40W/60W"
    m = re.search(r"(?<![\d.,])(\d+(?:[.,]\d+)?(?:W?[/-]\d+(?:[.,]\d+)?)+)\s?W(?![A-Z\d])", u)
    if m:
        steps = [fmt_num(to_float(x)) for x in re.findall(r"\d+(?:[.,]\d+)?", m.group(1))]
        sep = "-" if "-" in m.group(1) else "/"
        return f"{sep.join(steps)}W", None
    # najwyzej 4 cyfry: kod zamowieniowy z karty GTIN "LVPR16501204,5W" (PAR16 50 120 4,5W) to nie 16501204,5 W (2026-10-26)
    m = re.search(r"(?<![\d.,/])(\d{1,4}(?:[.,]\d+)?)\s?W(?=[^A-Z]|$|CCT|TW|WT|PS|BK|GR|DALI|RGB)", u)
    if m:
        value = to_float(m.group(1))
        if value > 0:
            return f"{fmt_num(value)}W", value
    return "", None


def parse_cct(u: str) -> str:
    # tylko Smart+ / Matter - w dekorach "Decor Straw Multicolor" to kolorowe szklo, nie RGBW (2026-10-21)
    if re.search(r"MULTICOLOU?R", u) and re.search(r"SMART|MATTER|WIFI|\bZB\b|\bBT\b", u):
        return "RGBW"  # Smart+ Matter "827...865 Multicolor" = RGB + barwa przestrajana (2026-10-21)
    # "TWFRE14" / "TWFRB22D": TW sklejone z kloszem w kodach SMARTWIFI (2026-10-21: dawalo samo "2700K")
    if re.search(r"\bTW\b|TUNABLE|/TW\b|TW(?=[A-Z]{2}\b)|TW(?=FR)|RGBTW", u):
        base = "TW"
    else:
        base = ""
    kelvins: list[str] = []
    m = re.search(r"(?<!\d)(\d{4})\s?(?:-|…|\.\.\.)\s?(\d{4})\s?K\b", u)
    if m:
        kelvins.append(f"{m.group(1)}-{m.group(2)}")
        if int(m.group(1)) <= 3000 and int(m.group(2)) >= 5000:
            base = "TW"  # karta Smart+ "4.9 W/2700…6500 K" bez slowa Tunable (2026-10-21)
    else:
        for k in re.findall(r"(?<![\d.,])([12-6]\d{3})\s?K\b", u):
            if k not in kelvins and 1500 <= int(k) <= 6500:
                kelvins.append(k)
    if not kelvins and re.search(r"VINTAGE|\b1906\b", u):
        # Vintage 1906 "618 Copper" = CRI 60, 1800 K
        m = re.search(r"(?<![\d.,])6(1[68])(?![\d.,])", u)
        if m:
            kelvins.append(m.group(1) + "00")
    if not kelvins:
        # kod barwy przyklejony do mocy: "9W830", "4,9W/827"
        for m in re.finditer(r"\dW/?([789])(\d{2})(?!\d)", u):
            if m.group(2) in VALID_CCT and m.group(2) + "00" not in kelvins:
                kelvins.append(m.group(2) + "00")
    if not kelvins:
        # Kody barwy: 840, 930/940, MS827, 830-840, 827/840/865, multi lumen 83040, 8405065.
        for m in re.finditer(r"(?<![A-Z0-9.,-])(?<!\d/)(?:MS|ML)?([789])(\d{2})((?:[/-][789]\d{2})*)(?![\d,.]|\s?MM|\s?LM|LED(?!V))", u):  # "840LED" to liczba diod ("830 LEDV" to barwa)
            temps = [m.group(2)] + [t[2:] for t in re.findall(r"[/-][789]\d{2}", m.group(3))]
            if m.group(1) == "7" and m.group(2) == "20":
                continue  # "T8 EM V 720 7W 830" - 720 to dlugosc tuby, nie barwa (2026-10-21)
            if all(t in VALID_CCT for t in temps):
                for t in temps:
                    if t + "00" not in kelvins:
                        kelvins.append(t + "00")
        for m in re.finditer(r"(?<![A-Z0-9.,/-])(?:ML)?([789])((?:\d{2}){2,3})(?![\d,.]|\s?LM)", u):
            pairs = [m.group(2)[i:i + 2] for i in range(0, len(m.group(2)), 2)]
            if all(t in VALID_CCT for t in pairs):
                for t in pairs:
                    if t + "00" not in kelvins:
                        kelvins.append(t + "00")
    if not kelvins:
        # moduly: "1800-840" (lm-kod barwy)
        m = re.search(r"\d{3,5}-([789])(\d{2})\b", u)
        if m and m.group(2) in VALID_CCT:
            kelvins.append(m.group(2) + "00")
    parts = []
    if kelvins:
        parts.append("/".join(kelvins) + "K")
    elif base:
        parts.append(base)
    elif re.search(r"\b\d?CCT\b|CCT\b|CCT SELECT", u):
        parts.append("CCT")
    if re.search(r"RGBW|RGB ?\+ ?W", u):
        parts.append("RGBW")
    elif re.search(r"\bRGB", u):
        parts.append("RGB")
    if base and kelvins:
        if len(kelvins) == 1 and "-" not in kelvins[0] and "/" not in kelvins[0]:
            # "TW 9W/827" - barwa startowa nic nie mowi przy barwie przestrajanej (2026-10-21: "TW 2700K")
            parts = [p for p in parts if not p.endswith("K")]
        parts.insert(0, "TW")
    return " ".join(parts)


def parse_socket(u: str) -> str:
    m = SOCKET_RE.search(u) or SOCKET_GLUED_RE.search(u)
    if not m:
        return ""
    s = m.group(1).replace(",", ".").replace(".-", "-")
    s = re.sub(r"^(GX?24)([DQ])-?(\d)$", lambda x: f"{x.group(1)}{x.group(2).lower()}-{x.group(3)}", s)
    s = re.sub(r"^(GX?24)([DQ])$", lambda x: f"{x.group(1)}{x.group(2).lower()}", s)
    return {"B22D": "B22d", "B15D": "B15d", "BA15D": "BA15d", "R7S": "R7s", "K12S": "K12s",
            "GR10Q": "GR10q", "S14D": "S14d", "S14S": "S14s", "S15S": "S15s"}.get(s, s)


def card_ml_cct(name: str) -> str:
    """Multi Select z karty GTIN: "13WML 93040" / "ML 8304057" -> "3000/4000K" / "3000/4000/5700K"."""
    m = re.search(r"ML\s?[89]((?:\d{2}){2,3})(?!\d)", ascii_upper(name))
    if not m:
        return ""
    temps = [m.group(1)[i:i + 2] for i in range(0, len(m.group(1)), 2)]
    return "/".join(t + "00" for t in temps) + "K" if all(t in VALID_CCT for t in temps) else ""


def card_cct(cu: str) -> str:
    """Barwa z nazwy karty GTIN: "2700…6500 K" = TW (Smart+), "822…827" = 2200/2700K (GlowDim)."""
    found = parse_cct(re.sub(r"(?<=\d)\s?(?:…|\.\.\.)\s?(?=\d)", "-", cu))
    return "TW" if found == "2700-6500K" else found


def parse_ip(u: str) -> str:
    m = re.search(r"\bIP\s?(\d{2})\b", u)
    return f"IP{m.group(1)}" if m else ""


def parse_control(u: str) -> str:
    parts = []
    # w TruSys "DA" to optyka Double Asymmetric, nie DALI (2026-10-21: "TRUSYS LED 53W DA" i "... DALI 53W DA" = dwa DALI)
    da = r"DALI(?!O)|DAVR" if "TRUSYS" in u else r"DALI(?!O)|\bDA\b|DAVR|\bDA(?=EUE|LEDV)|AL?DA(?=EUE)"
    if re.search(da, u):
        parts.append("DALI")
    if re.search(r"ZIGBEE|\bZB\b|ZBVR|\bZB(?=[A-Z0-9])", u):
        parts.append("Zigbee")
    # "SMART WFA30D", "SMART BTA60D", "SMART BTHKA60" - protokol sklejony z ksztaltem w kodach OSRAM
    if re.search(r"WIFI|WI-FI|SMARTWF|SMART WF\b|SMART WF(?=[ABP]\d{2})", u):
        parts.append("WiFi")
    if re.search(r"\bBT\b|BLUETOOTH|\bBT(?=[ABP]\d{2}|HK)", u):
        parts.append("BT")
    if re.search(r"\bMATTER\b", u):
        parts.append("Matter")
    return " ".join(parts)


def parse_sensor(u: str) -> str:
    """Czujnik z rodzajem, gdy nazwa go podaje (2026-10-21: "z czujnikiem z czujnikiem ruchu i zmierzchu").

    "STAR+ CL A MSFR" / "LEDSCLA60MS" = ruchu, "DSFR" / DAYLIGHT = zmierzchu; samo SENSOR zostaje ogolne.
    """
    if not re.search(r"SENSOR|\bSEN\b|CZUJNIK|MOTION|\bLX\b|\bPIR\b|\b[MD]SFR\b|LEDSCL[ABP]\d+MS", u):
        return ""
    motion = re.search(r"MOTION|\bPIR\b|\bMSFR\b|LEDSCL[ABP]\d+MS|CZUJNIK(?:IEM)? RUCHU", u)
    day = re.search(r"DAYLIGHT|\bDSFR\b|ZMIERZCH", u)
    if motion and day:
        return "z czujnikiem ruchu i zmierzchu"
    if motion:
        return "z czujnikiem ruchu"
    if day:
        return "z czujnikiem zmierzchu"
    return "z czujnikiem"


def parse_emergency(u: str) -> str:
    return "awaryjna" if re.search(r"EMERGENCY|AWARYJN|\bEM(?=BK|WT|\b)", u) else ""


def parse_tube_type(u: str) -> str:
    # tez sklejone "LEDTUBE T8EM" i kody kart "ST8E-AC" / "ST8V-EM" (SubstiTUBE T8), 2026-10-21
    m = re.search(r"(?<![A-Z0-9])S?(T5C|T9C|T8|T5|T9|T12)(?=\b|EM\b|HF\b|AC\b|[EVPS]-)", u)
    if m:
        return m.group(1)
    if re.search(r"\bFC\b", u):
        return "T5 FC"
    return ""


def parse_tube_supply(u: str) -> str:
    # DULUX LED: CCG = statecznik konwencjonalny (EM), ECG = elektroniczny (HF); "CCG, AC mains" dziala tez z 230V
    em = bool(re.search(r"\bEM\b|\bCCG\b", u))
    hf = bool(re.search(r"\bHF\b|\bECG\b", u))
    ac = bool(re.search(r"AC MAINS|\bAC\b|230 ?V|\bMAINS\b", u))
    if re.search(r"\bUN\b|UNIVERSAL", u):
        return "EM/HF/230V"
    if em and ac:
        return "EM/230V"
    if re.search(r"EXTERNAL|\bEXT\b", u):
        return "zasilacz zewn."
    if hf and ac and re.search(r"\bECG\b", u):
        return "HF/230V"  # "DULUX L LED FR 36 ECG, AC mains" (2026-10-21: bylo samo 230V)
    if hf:
        return "HF"
    if em:
        return "EM"
    if ac:
        return "230V"
    return ""


def mm_to_cm(mm: float) -> str:
    return f"{round(mm / 10):d}cm"


def parse_tube_length(u: str, tube_type: str, power: float | None, rules: dict[str, Any]) -> str:
    m = re.search(r"(?<![\d.,])(\d{3,4})\s?MM\b", u)
    if m:
        return mm_to_cm(int(m.group(1)))
    m = re.search(r"(?<![\d.,])(\d{2,3})\s?CM\b", u)
    if m:
        return f"{int(m.group(1))}cm"  # Kobi: "LED T8 9W 60CM"
    for tok in re.findall(r"(?<![\d.,/])(\d{3,4})(?![\d.,/])", u):
        # 720 = "LEDTUBE T8 EM V 720 7W" (tuba 72 cm, 2026-10-21)
        if int(tok) in {438, 449, 517, 549, 563, 588, 600, 720, 849, 895, 900, 1050, 1149, 1200, 1449, 1500, 1513, 1800}:
            return mm_to_cm(int(tok))
    m = re.search(r"\bH[OE] ?(\d{2})\b", u)
    table = rules.get("fluorescent_lengths", {})
    if m and str(int(m.group(1))) and "T5" in table:
        cm = table["T5"].get(int(m.group(1)))
        if cm:
            return f"{cm}cm"
    if power and tube_type in table:
        cm = table[tube_type].get(int(power)) if float(power).is_integer() else None
        if cm:
            return f"{cm}cm"
    return ""


def parse_length(u: str) -> str:
    m = re.search(r"(?<![\d.,])(\d{2,3})\s?CM\b", u)
    if m:
        return f"{int(m.group(1))}cm"
    m = re.search(r"(?<![\d.,])(\d{3,4})\s?MM\b", u)
    if m and int(m.group(1)) >= 300:
        return mm_to_cm(int(m.group(1)))
    m = re.search(r"(?<![\d.,])([0-2])[.,](\d)\s?M\b", u)  # "SUBMARINE Integrated 0.6M / 1.2M" (2026-10-20)
    if m:
        return f"{int(m.group(1)) * 100 + int(m.group(2)) * 10}cm"
    for tok in re.findall(r"(?<![\d.,/X])(\d{3,4})(?![\d.,/X])", u):
        if int(tok) in {600, 900, 1000, 1200, 1500, 1800}:
            return mm_to_cm(int(tok))
    return ""


def parse_panel_size(u: str) -> str:
    m = re.search(r"(?<!\d)(\d{3,4})\s?X\s?(\d{3,4})(?:\s?MM)?\b", u)
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        return f"{fmt_num(a / 10)}x{fmt_num(b / 10)}"
    m = re.search(r"(?<![\d.,])(30|60|62|120)\s?X\s?(30|60|62|120)(?:\s?CM)?\b", u)  # Kobi: "60x60" w cm
    if m:
        return f"{m.group(1)}x{m.group(2)}"
    for tok in re.findall(r"(?<![\d.,/X])(\d{3,4})(?![\d.,/X])", u):
        n = int(tok)
        if n == 600:
            return "60x60"
        if n == 625:
            return "62,5x62,5"
        if n == 1200:
            return "120x30"
    m = re.search(r"ROUND (\d{3})\s?MM|ROUND.*?(\d{3})\s?MM", u)
    if m:
        return f"fi{int(m.group(1) or m.group(2)) // 10}cm"
    return ""


def parse_diameter(u: str, bare: bool = False) -> str:
    # oprawa kwadratowa: wymiar boku w mm jak srednica ("SURFACE SQUARE 330" -> "330x330 mm", "500 SQ" -> "500x500 mm",
    # "SQUARE DALI 500", "Square 200mm"); sama liczba bez jednostki - "SQUARE 16W" to moc, a "DULUX LED SQ" /
    # "Cardboard Square" bez liczby nic nie daja
    m = re.search(r"\bSQ(?:UARE)?\s+(?:(?:S|DALI|DA|SENSOR|SEN|SN|ZB)\s+)?(\d{2,4})(?=MM\b|\b)"
                  r"(?![.,]\d|\s?(?:W|K|LM|V|CM)\b)", u) \
        or re.search(r"(?<![\d.,X])(\d{3,4})\s+SQ\b", u)
    if m:
        return f"{int(m.group(1))}x{int(m.group(1))} mm"
    m = re.search(r"(?<![\d.,X])(\d{2,4})\s?X\s?(\d{2,4})\s?MM\b", u)  # "Orbis Alicia 292X292MM", "430x430mm"
    if m:
        return f"{int(m.group(1))}x{int(m.group(2))} mm"
    m = re.search(r"(?<![\d.,X])(\d{2})\s?X\s?(\d{2})\s?CM\b", u)  # "Orbis Disc 30X30CM"
    if m:
        return f"{int(m.group(1)) * 10}x{int(m.group(2)) * 10} mm"
    m = re.search(r"(?<![\d.,X])(\d{2})\s?CM\b", u)  # "Orbis Disc 30CM"
    if m and 10 <= int(m.group(1)) <= 90:
        return f"fi{int(m.group(1)) * 10}"
    # bez przeskakiwania mocy: w "AR SPOT1 RD V 6W 830" liczba za moca to kod barwy, nie srednica
    m = re.search(r"\b(?:DN|D|RD)\s?(\d{2,3})\b", u) or re.search(
        r"\b(?:ROUND|CIRCULAR|CIRC|RD)\s+(?:(?!\d+(?:[.,]\d+)?W\b)[A-Z0-9]+\s+){0,2}?(\d{3})\b", u) \
        or re.search(r"(?<![\d.,X])(\d{3})\s+R\b", u)  # "SF FLAT 500 R" (okragly) obok "500 SQ"
    if m:
        return f"fi{int(m.group(1))}"
    m = re.search(r"(?<![\d.,X])(\d{2,3})\s?MM\b", u)  # "TWIST PRO 70mm", "RECESS SLIM 85mm"
    if m and int(m.group(1)) >= 40:
        return f"fi{int(m.group(1))}"
    if bare:
        # plafony: sama liczba 200-699 ("SURFACE BULKHEAD 250 10W", "SF DISC 300 V 21W"); 7xx-9xx to kody barwy
        # "SF DISC 220 V 14W": V za liczba to klasa Value, nie napiecie (napiecie: "230V", "230 V" bez mocy dalej)
        m = re.search(r"(?<![\d.,X/-])([2-6]\d{2})(?![\d.,/X-])(?!V\b)(?!\s?(?:W|K|LM|MM|CM|MA|H)\b)"
                      r"(?!\s+V\b(?!\s+\d))", u)
        if m:
            return f"fi{int(m.group(1))}"
    return ""


def parse_strip(u: str) -> str:
    # "SUP-2000/TW/927-965/5" - zakres barwy TW przed dlugoscia (2026-10-21: bylo "927m")
    m = re.search(r"-(\d{3,4})/(?:(RGBW|RGB|COB|TW)/)?(?:(\d{3}(?:-\d{3})?|RGBW)/)?(\d+(?:[.,]\d+)?)(?=/|\b)", u)
    if not m:
        return ""
    lm, kind, code, length = m.group(1), m.group(2), m.group(3), m.group(4)
    parts = [f"{fmt_num(to_float(length))}m", f"{lm}lm/m"]
    if code and code.isdigit() and code[1:] in VALID_CCT:
        parts.append(f"{code[1:]}00K")
    if kind:
        parts.append(kind)
    elif code == "RGBW":
        parts.append("RGBW")
    return " ".join(parts)


def parse_driver(u: str) -> str:
    m = re.search(r"(\d+(?:[.,]\d+)?)/([\d.…]+(?:-|\.\.\.|…)[\d.]+)/(\d+A\d|\d+(?:-\d+)?)", u)
    if not m:
        m2 = re.search(r"(\d+)\s?W\b", u) or re.search(r"\b(\d{2,4})/2[0-4]0(?:-2[0-4]0)?\b", u)
        return f"{m2.group(1)}W" if m2 else ""  # tez "OT SLIM 100/220-240" (moc / napiecie wejsciowe), 2026-10-21
    power, middle, out = m.group(1), m.group(2), m.group(3)
    bounds = [b for b in re.split(r"-|\.\.\.|…", middle) if b]
    parts = [f"{fmt_num(to_float(power))}W"]
    is_input_voltage = (len(bounds) == 2 and bounds[0].isdigit() and bounds[1].isdigit()
                        and int(bounds[0]) in (100, 110, 120, 198, 200, 220, 230)
                        and int(bounds[1]) in (230, 240, 250, 254, 264, 277))
    if is_input_voltage:
        # wejscie 220-240 V -> wyjscie napieciowe (24) albo pradowe (700, 1A0)
        if "A" in out:
            amps = out.replace("A", ".")
            parts.append(f"{int(round(float(amps) * 1000))}mA")
        elif out.isdigit() and int(out) <= 48:
            parts.append(f"{out}V")
        else:
            parts.append(f"{out}mA")
    else:
        parts.append(f"{middle.replace('…', '-').replace('...', '-')}mA")
    return " ".join(parts)


def parse_ballast(u: str) -> str:
    # "FIT8 3x18/4x18" - jeden statecznik do 3 albo 4 swietlowek (2026-10-26: dawalo "3x18W")
    m = re.search(r"\b(\d)\s?X\s?(\d+)/(\d)\s?X\s?\2\b", u)
    if m:
        return f"{m.group(1)}/{m.group(3)}x{m.group(2)}W"
    m = re.search(r"\b(\d)\s?X\s?(\d+(?:-\d+)?(?:[ /]\d{2}(?!\d))*)", u)
    if m:
        cfg = re.sub(r"[ ]", "/", m.group(2))
        return f"{m.group(1)}x{cfg}W" + (" DIM" if re.search(r"\bDIM\b", u) and "DALI" not in u else "")
    m = re.search(r"\b(\d{2,4})/2\d0", u)
    if m:
        return f"{m.group(1)}W" + (" DIM" if "DIM" in u else "")
    return ""


SHAPES: list[tuple[str, Any]] = [
    # zamienniki lamp wyladowczych - rodzaj lampy mowi sprzedawcy wiecej niz ksztalt (2026-10-21: "HQL LED 1800
    # 14,5W" dawalo sama "Zarowke LED E27 14,5W"); HID LED HB = zamiennik do opraw high bay
    (r"\bHQL\b", "HQL"),
    (r"\bNAV\b", "NAV"),
    (r"\bHQI\b", "HQI"),
    (r"\bHID LED\b", "HID"),
    (r"HIGHBAY BLADE", "High Bay Blade"),
    (r"\bCIRCOLUX\b", "Circolux"),
    (r"\bLEDINESTRA\b(?:\D*?(\d{2})\s?CM\b)?", lambda m: f"LEDinestra {m.group(1)}cm" if m.group(1) else "LEDinestra"),
    # lista kilku ksztaltow w nazwie z Ekonoma ("CLASSIC A,B,P SLIM", "CLASSIC A, Globe and EDISON") - ksztalt tylko
    # z karty GTIN (2026-10-21: dawalo "A60" swieczkom E14)
    (r"\bA,\s?B,\s?P\b|\bA,\s?GLOBE AND EDISON\b|\bCLASSIC SHAPES\b", ""),
    # rozmiar globe zawsze w ksztalcie ("globe G125"), tez sklejony "1906LGL125" i Vintage "GOLD 95" / "SMOKE 80"
    # tylko srednice banki; "Big Globe 15" / "Globe 28" na kartach Vintage to zamiennik mocy, nie rozmiar
    (r"\bGLOBE ?(45|60|80|95|100|125|130|150|165|183|200|220)\b(?![.,]\d)(?!\s?W\b)|\bG(60|80|95|125)\b"
     r"|LGL(80|95|125)D?\b|\b(?:GOLD|SMOKE|SMOEK) ?(80|95|125|200)\b",
     lambda m: f"globe G{next(g for g in m.groups() if g)}"),
    # Vintage "PINECONE" / karta "Big Globe 40 Filament Pine" - szyszka, nie globe (2026-10-26)
    (r"\bPINE(?:CONE)?\b", "Pinecone"),
    (r"\bGLOBE\b", "globe"),
    (r"\bPAR ?(11|16|20|30|38)\b", lambda m: f"PAR{m.group(1)}"),
    (r"\bMR ?(11|16)\b", lambda m: f"MR{m.group(1)}"),
    (r"\bAR ?111\b", "AR111"),
    (r"\bR ?(50|63|80)\b", lambda m: f"R{m.group(1)}"),
    # "ENERGY EFFICIENCY REFLECTOR 2W 827 GU10" - karta: PAR16 (2026-10-21)
    (r"\bREFLECTOR\b(?=.*\bGU10\b)", "PAR16"),
    # "SMART+ Spot GU10 Dimmable 40 45°" / karta "SMART+ BT SPOT 50 45 ° 4.9 W GU10" - lampa PAR16 (2026-10-21)
    (r"\bSPOT\b(?=.*\bGU10\b)", "PAR16"),
    (r"\bT ?26\b", "T26"),
    # "CLASSIC ST P" / "ST45" / "E60 D" (Smart+ BT) - ksztalt Edison wg karty GTIN
    (r"\bEDISON\b|\bST(?:45|64)\b|\bCL(?:ASSIC)? ST\b|\bE(?:55|60)D?\b", "Edison"),
    (r"\bSTICK\b|\bT SLIM\b|\bSLIM T\b", "sztyft"),
    (r"\bOVAL\b", "owal"),
    # nazwy ksztaltu z kart Smart+ ("SMART+ BT Candle 40", "Mini bulb 40")
    (r"\bCANDLE\b", "świeczka"),
    (r"\bMINI BULB\b", "kulka"),
    # sklejone z mocowym zamiennikiem: "Classic BA40", "CL B40", "Classic P25", "LEDSCLB40", "VALCLA100" (2026-10-21)
    (r"\b(?:CL(?:ASSIC)? ?)?BA(?: ?\d{2,3})?\b", "płomyk"),
    # BW to wariant plomyka (wiatr), sprzedawcy wystarczy "plomyk" - decyzja uzytkownika 2026-10-21
    (r"\b(?:CL(?:ASSIC)? ?)?BW(?: ?\d{2,3})?\b", "płomyk"),
    (r"\bCL(?:ASSIC)? ?B(?: ?\d{2,3})?\b|\bB(25|35|40|60|75)\b|\bCLB\b|(?:LEDS?|VAL)CLB\d", "świeczka"),
    (r"\bCL(?:ASSIC)? ?P(?: ?\d{2,3})?\b|\bP(25|40|45|60)\b|\bCLP\b|(?:LEDS?|VAL)CLP\d", "kulka"),
    (r"\bA(60|65|67)\b", lambda m: f"A{m.group(1)}"),
    (r"\bA ?(125|160)\b", lambda m: f"A{m.group(1)}"),  # Vintage 1906 "CL A125" / "A160" - duze banki
    # "A40" / "A75" / "A100" / "A150" to zamiennik mocy (75 W), nie rozmiar banki - ten sam produkt co "CLASSIC A
    # V 75" = A60 (2026-10-21: obok siebie stalo "A75" i "A60"); Vintage A125 / A160 to rzeczywiste rozmiary
    (r"\bA(?:15|25|40|45|55|75|100|120|150|200)\b|CLASSIC ?A\d{2,3}\b", "A60"),
    (r"\bCL(?:ASSIC)? ?A\b|\bCLA\b|(?:LEDS?|VAL)CLA\d", "A60"),
    (r"\bPIN\b", "kapsułka"),
    (r"\bFLAT\b", "płaska"),
]


def parse_shape(u: str) -> str:
    for pattern, render in SHAPES:
        m = re.search(pattern, u)
        if m:
            return render(m) if callable(render) else render
    return ""


BULB_LINES: list[tuple[str, str]] = [
    (r"SUPERSTAR ?\+|SUPERSTAR PLUS", "Superstar+"),
    (r"\bSUPERSTAR\b", "Superstar"),
    (r"\bSUPERIOR\b", "Superior"),
    (r"\bPERFORMANCE\b", "Performance"),
    (r"\bSTAR ?\+", "Star+"),
    (r"\bSTAR\b", "Star"),
    (r"\bVALUE\b", "Value"),
    # samodzielne slowo oprawy ("DOWNLIGHT PFM DN195"); w kodach zasilaczy i akcesoriow ("DR DS-PFM-20", "LS AY PFM -CSD",
    # "PFM and VAL") zostaje kod modelu, ktory odroznia je od innych serii
    (r"(?<![\w-])PFM(?![\w-])(?!\s+-)(?!\s+AND\b)", "Performance"),
    (r"\bENTRY\b", "Entry"),  # tuby "T8 ENTRY EM"
    (r"\bADV(?:ANCED)?\b", "Advanced"),  # tuby "T8 ADV Ultra Output"
    (r"\bESSENTIAL\b", "Essential"),
    (r"\bLED BASE\b|\bBASE (?:CLASSIC|CL)\b", "Base"),  # "rail and base" to podstawa szyny
    (r"\bPARATHOM (?:DIM )?PRO\b", "Parathom Pro"),  # tez "PARATHOM DIM PRO Classic A60"
    (r"\bPARATHOM\b", "Parathom"),
    (r"\bPRO\b", "Pro"),  # "TWIST PRO", "HALOLINE PRO", "Quicktronic PRO", "HQL LED PRO"
    (r"\bBELLALUX\b", "Bellalux"),
    (r"\bVINTAGE\b|\b1906\b", "Vintage 1906"),
    (r"SUN@HOME", "Sun@Home"),
    (r"\bSMART(?!LUX)", "Smart+"),  # SMARTLUX to swietlowka, nie Smart+
    (r"\bPCR\b", "PCR"),
    (r"ANTYBAKTERYJN", "antybakteryjna"),
]


# jednowyrazowe linie - `move_brand_tail` przestawia je do marki, gdy przyszly jako wyroznik
LINE_LABELS = {label for _, label in BULB_LINES if " " not in label} - {"antybakteryjna"}


def parse_bulb_line(u: str) -> str:
    """Linia produktu; tuby i liniowe niosa dodatkowo strumien "Ultra Output" ("UO P" = Performance Ultra Output)."""
    output = ("Ultra Output" if re.search(r"\bULTRA OUTPUT\b|\bUO\b", u)
              else "High Output" if re.search(r"\bHIGH OUTPUT\b", u) else "")
    # tylko linia zapisana slowem; sama litera ("MR16 DIM S", "EM P 600") nie jest dowodem - 2026-10-19 wylaczone,
    # bo bywa czyms innym: "URBAN LITE S 15W" to rozmiar S, nie Superior
    label = next((label for pattern, label in BULB_LINES if re.search(pattern, u)), "")
    if label or output:
        return f"{label} {output}".strip()
    # klasa energetyczna jako "linia": "ENERGY EFFICIENCY C" to klasa C, nie A
    m = re.search(r"ENERGY (?:EFFICIENCY|CLASS)\s+(?:CLASS\s+)?([A-G])\b", u)
    return f"klasa {m.group(1)}" if m else ("klasa A" if re.search(r"\bEECA\b", u) else "")


def parse_family(u: str, rules: dict[str, Any]) -> str:
    """Rodzina opraw z `families` ("URBAN LITE S" -> "Urban Lite S"); litery rozmiaru zostaja wielkie."""
    for item in rules.get("families", []):
        m = re.search(item["match"], u)
        if m:
            return item.get("label") or " ".join(w if len(w) <= 2 else w.capitalize() for w in m.group(0).split())
    return ""


def card_family_name(family: str, cfg: dict[str, Any]) -> str:
    """Nazwa wlasna z rodziny na karcie GTIN (`family_from_card`): "DAMP PROOF MAX EMERGENCY SENSOR" -> "Damp Proof Max".

    Pusta, gdy rodzina to zaslepka ("Phase out") albo sama nazwa kategorii ("PANEL PERFORMANCE 600" -> "Panel").
    """
    if not family or any(re.search(p, family.strip()) for p in cfg.get("placeholders", [])):
        return ""
    text = family
    for item in cfg.get("rename", []):
        text = re.sub(item["match"], item["repl"], text)
    for pattern in cfg.get("remove", []):
        text = re.sub(pattern, " ", text)
    words = re.sub(r"\s+", " ", text).strip(" -/").split()
    while words and words[-1].lower() in {"and", "with", "for", "of", "&", "-", "/"}:
        words.pop()  # "LED Color and White" bez koloru -> "Color", nie "Color and"
    generic = {w.lower() for w in cfg.get("generic_words", [])}
    if not words or (len(words) == 1 and words[0].lower() in generic):
        return ""  # samo slowo kategorii ("Floodlights", "Spot"); "Damp Proof", "Panel Compact" to nazwy wlasne
    spelling = cfg.get("spelling", {})

    def case(w: str) -> str:
        if w.lower() in {"and", "with", "for", "of"}:
            return w.lower()
        if re.fullmatch(r"[A-Z]{1,2}|XL", w):  # litery rozmiaru / klasy (Lite S, MR)
            return w
        if w.isupper() or w.islower():
            w = w.capitalize()
        return spelling.get(w, w)  # mieszana pisownia z karty zostaje (IndiviLED, TruSys)

    return " ".join("/".join(case(p) for p in w.split("/")) for w in words)  # "Direct/Indirect"


def card_families(web: dict[str, dict[str, Any]], cfg: dict[str, Any]) -> dict[str, str]:
    """EAN -> nazwa wlasna z rodziny na karcie GTIN (tylko karty dokladnie tego EAN).

    Pusty napis = karta ma rodzine, ale to sama kategoria ("HIGH BAY V") - karta decyduje, bez szukania dalej.
    Brak klucza = brak karty albo rodzina-zaslepka ("Phase out") - wtedy `family_phrases`.
    """
    out = {}
    for ean, card in web.items():
        family = card.get("family", "")
        placeholder = any(re.search(p, family.strip()) for p in cfg.get("placeholders", []))
        if card.get("exact") and placeholder and card.get("name") and not re.search(r"\d", card["name"]):
            # "Phase out" + nazwa karty bez parametrow ("DAMP PROOF ECO VERY WIDE") - nazwa karty to nazwa rodziny
            family, placeholder = card["name"], False
        if card.get("name") and any(re.search(p, family.strip()) for p in cfg.get("refine_from_name", [])):
            family = card["name"]  # "Orbis Click Sensor" -> "Orbis Square Click Sensor 430x430mm 50W"
        if not card.get("exact") or not family or placeholder:
            continue
        out[ean] = card_family_name(family, cfg)
    return out


def family_phrases(web: dict[str, dict[str, Any]], cfg: dict[str, Any]) -> list[str]:
    """Nazwy wlasne ze wszystkich kart (i rodzin z podpowiedzi wyszukiwarki), od najdluzszej - do szukania w nazwie z Ekonoma.

    Tylko nazwy z co najmniej dwoch slow: "Damp Proof", "Panel Compact", "Streetlight Area" (pojedyncze "Click", "Spot"
    trafialyby przypadkiem).
    """
    names = {card_family_name(card.get("family", ""), cfg) for card in web.values() if card.get("family")}
    return sorted((n for n in names if len(n.split()) >= 2), key=lambda n: (-len(n.split()), n))


def family_from_name(u: str, phrases: list[str]) -> str:
    """Najdluzsza nazwa wlasna z kart, ktora stoi doslownie w nazwie z Ekonoma ("OPRAWA DAMP PROOF HE DA 1500" -> "Damp Proof")."""
    # "PLANON™" -> "PLANONTM" po ascii_upper - znak towarowy nie jest czescia nazwy
    text = f" {re.sub(r'[^A-Z0-9+/]+', ' ', re.sub(r'(?<=[A-Z])TM\b', '', ascii_upper(u.replace('™', ''))))} "
    for phrase in phrases:
        if f" {re.sub(r'[^A-Z0-9+/]+', ' ', ascii_upper(phrase))} " in text:
            return phrase
    return ""


def parse_glass(u: str) -> str:
    # "CL" w nazwach Ledvance to CLASSIC, nie "clear" - nie jest brane pod uwage
    # klosz lustrzany ("Mirror Silver 50", "Mirror Gold") mowi wiecej niz "przezroczysta" pod spodem (2026-10-21)
    m = re.search(r"\bMIRROR(?:\s+(SILVER|GOLD))?\b", u)
    if m:
        return {"SILVER": "lustrzana srebrna", "GOLD": "lustrzana złota"}.get(m.group(1) or "", "lustrzana")
    # barwione szklo klosza po angielsku, jak na opakowaniu (decyzja uzytkownika 2026-10-21): Gold / Smoke / Amber
    # zastepuje "przezroczysta" (barwione szklo jest przezroczyste), "matowa" zostaje obok
    m = re.search(r"\b(GOLD|SMOKE|AMBER)\b", u)
    if m:
        tint = m.group(1).capitalize()
        return f"matowa {tint}" if re.search(r"FROST|\bFR\b|MATOW", u) else tint
    if re.search(r"\bCLEAR\b|PRZEZROCZYST", u):
        return "przezroczysta"
    # GLFR = szklo matowe (Superior Class "CLASSIC P GLFR 40")
    if re.search(r"FROST|\bFR\b|MATOW|MLECZN|DSFR|MSFR|\bFR(?=E\d)|\bGLFR\b|GLFR(?=\s|E\d)", u):
        return "matowa"
    return ""


def parse_dim(u: str) -> str:
    # "non.-dim" z kropka (Parathom PIN, Superstar+) to tez nieściemnialna (2026-10-21: dawalo "ściemnialna")
    if re.search(r"NON\.?[- ]?DIM", u):
        return ""
    # sposob sciemniania nazwa z opakowania (decyzja uzytkownika 2026-10-21): "3 STEP DIM" = 3-stopniowa (wylacznikiem),
    # "DIM Duo Click" = Duo Click (wylacznikiem, bez sciemniacza), GLOWdim = sciemnialna ze zmiana barwy
    if re.search(r"\b3[ -]?STEP\b", u):
        return "3-stopniowa"
    if re.search(r"\bDUO ?CLICK\b", u):
        return "Duo Click"
    if re.search(r"GLOWDIM", u):
        return "ściemnialna GlowDim"
    return "ściemnialna" if re.search(r"\bDIM\b|DIMMABLE|SCIEMNIALN|GLOWDIM|\bDIM(?=FR)", u) else ""


LETTER_OPTICS = {"VN": "opt. VN", "VW": "opt. VW", "WW": "opt. WW", "DS": "opt. DS", "OV": "opt. OV",
                 "NFL": "opt. NFL", "N": "opt. N", "W": "opt. W", "NB": "opt. NB", "MB": "opt. MB", "WB": "opt. WB"}


def parse_optics(u: str) -> str:
    """Kat swiecenia albo kod optyki (wersja), gdy katalog nie podaje kata."""
    m = re.search(r"\bA(?:SYM(?:METRIC)?)?\s?(\d{2,3})\s?X\s?(\d{2,3})\b", u)
    if m:
        return f"asym. {m.group(1)}x{m.group(2)}°"
    # "R45 WAL" (Floodlight Max) - WAL to kolor, nie moc (2026-10-21: kat ginal)
    m = re.search(r"\bSYM(?:METRIC)?\s?(\d{2,3})\b|\bR(\d{2})\b(?!\s?(?:W|MM)\b)", u)
    if m:
        return f"{int(m.group(1) or m.group(2))}°"
    # kod optyki ulicznej Ledvance zawsze jako kod (decyzja uzytkownika 2026-10-21): RV20ST i RW20ST to rozne
    # rozsyly, "20°" je zlewal; przed regula Kobi "36ST", tez sklejony z barwa ("740RW20STNWAC2...")
    m = re.search(r"(?:\b|(?<=\d))((?:R[UVWN]|PW)\d{2}L?ST)(?=\b|[A-Z]{2}(?:AC)?\d)", u)
    if m:
        return f"opt. {m.group(1)}"
    # Kobi: "36ST", "10-60ST", "45/80ST" (stopnie)
    m = re.search(r"(?<![\d.,])(\d{1,3}(?:[-/]\d{1,3})?)\s?ST\b", u)
    if m:
        return f"{m.group(1)}°"
    # bez liczby sklejonej z moca: "9W830 D BK" to barwa 830 + Dimmable, nie 830° (2026-10-21)
    m = re.search(r"(?<![\d.,W])(\d{1,3})\s?(?:°|DEG\b|D\b)", u)
    if m and int(m.group(1)) <= 360:
        return f"{int(m.group(1))}°"
    # High Bay FLEX: kat bez jednostki miedzy moca a barwa ("240W 110 840", "240W 35X80 840")
    m = re.search(r"\dW\s+(\d{2,3})(?:X(\d{2,3}))?\s+(?:ML\s+)?[789]\d{2}\b", u)
    if m:
        return f"{m.group(1)}x{m.group(2)}°" if m.group(2) else f"{int(m.group(1))}°"
    if re.search(r"\bASYM(?:METRIC)?\b", u):
        return "asym."
    # litery optyki TruSys/Low Bay/Damp Proof - bez mocy, zeby "36 W" nie dalo "W"
    # "140W 840 W": 840 to kod barwy, a W za nim - optyka (wczesniej "840 W" wypadalo jak moc, 2026-10-20)
    rest = re.sub(r"(?<![\d.,])(?![6-9]\d\d\s+W\b)\d+(?:[.,]\d+)?\s?W\b|\b\d+\s?W/", " ", u)
    for tok in re.findall(r"\b(VN|VW|WW|DS|OV|NFL|NB|MB|WB|N|W)\b", rest):
        return LETTER_OPTICS[tok]
    return ""


REFLECTOR_SHAPE_RE = re.compile(r"^(PAR|MR|R\d|AR111|ES111)")
BULB_BEAMS = {10, 15, 20, 24, 25, 30, 36, 38, 40, 45, 60, 90, 100, 110, 120}


def parse_beam(u: str, shape: str) -> str:
    """Kat swiecenia zarowki reflektorowej (PAR16, R50, MR16...) - decyzja uzytkownika 2026-10-14."""
    if not REFLECTOR_SHAPE_RE.match(shape):
        return ""
    m = re.search(r"(?<![\d.,])(\d{1,3})\s?(?:°|DEG\b)", u)
    if m:
        return f"{int(m.group(1))}°"
    # "PAR16 80 36 6,9W", "R80 100 36": zamiennik W, potem kat bez jednostki
    m = re.search(r"\b(?:PAR ?\d{2}|MR ?1[16]|R ?(?:50|63|80)|AR ?111)\s+(?:(?:GL|GLASS|WHITE|BLACK)\s+)?\d{2,3}\s+(\d{2,3})\b(?!\s?(?:W|V)\b|[.,])", u)
    if m and int(m.group(1)) in BULB_BEAMS:
        return f"{int(m.group(1))}°"
    return ""


def parse_color(u: str, gender: str, colors: dict[str, Any], two_tone: bool = False) -> str:
    if two_tone:
        # "Black and Blue" -> "czarno-niebieski", "White and Wood Decor" -> "biały/drewno" (Ledvance Orbis, 2026-10-20)
        # kazde wystapienie koloru, nie tylko pierwsze: "Decor Wood table TW bk and wood" - "Wood" z nazwy serii stoi
        # przed para "bk and wood" (2026-10-21)
        found = sorted(((m.start(), m.end(), color) for color in colors.values()
                        for p in color["match"] for m in re.finditer(p, u)), key=lambda x: x[0])
        for (s1, e1, c1), (s2, e2, c2) in zip(found, found[1:]):
            if re.fullmatch(r"\s*(?:AND|&|/|\+)\s*", u[e1:s2]) and c1 is not c2:
                second = c2.get(gender, c2.get("f", ""))
                if c1.get("stem") and c2.get("stem"):
                    return f"{c1['stem']}-{second}"
                return f"{c1.get(gender, c1.get('f', ''))}/{second}"
    for color in colors.values():
        if any(re.search(p, u) for p in color["match"]):
            return color.get(gender, color.get("f", ""))
    return ""


def parse_ugr(u: str) -> str:
    # tez UGR16 (Panel Louver "U16", 2026-10-21: gubiony) i UGR22
    m = re.search(r"UGR\s?[<-]?\s?(1[69]|22)\b|\bU(1[69]|22)\b|S90U19|UGR(1[69]|22)", u)
    return f"UGR{next((g for g in m.groups() if g), '19')}" if m else ""


def parse_voltage(u: str) -> str:
    m = re.search(r"(?<![\d.,])(6|12|24)\s?V\b", u)
    return f"{m.group(1)}V" if m else ""


def parse_dulux_socket(u: str, power: float | None, rules: dict[str, Any]) -> str:
    m = re.search(r"DULUX (?:LED )?(D/E|T/E|S/E|D|T|S|L|F)\b", u)
    if not m or not power:
        return ""
    table = rules.get("dulux_sockets", {}).get(m.group(1), {})
    return table.get(int(power), "") if float(power).is_integer() else ""


def parse_name_lumen(u: str, rules: dict[str, Any]) -> float | None:
    for pattern in rules.get("lumen_patterns", []):
        m = re.search(pattern, u)
        if m:
            value = float(re.sub(r"[ .]", "", m.group(1)))
            if value >= 10:
                return value
    return None


# ---------------------------------------------------------------- lumeny z PDF

def pdf_texts(pdf_dir: Path, cache_dir: Path) -> dict[str, list[str]]:
    ensure_dir(cache_dir)
    raw: dict[str, str] = {}
    for pdf in sorted(pdf_dir.glob("*.pdf")):
        stat = pdf.stat()
        cached = cache_dir / f"{pdf.stem}.{stat.st_size}.txt"
        if cached.exists():
            text = cached.read_text(encoding="utf-8")
        else:
            import fitz  # PyMuPDF - tylko gdy trzeba odczytac nowy PDF

            doc = fitz.open(pdf)
            text = "\n".join(page.get_text() for page in doc)
            cached.write_text(text, encoding="utf-8")
        raw[pdf.name] = text

    # Katalogi PDF zajmuja setki MB, wiec moga zniknac z input/ - ich tekst zostaje
    # w cache ("<nazwa>.<rozmiar>.txt"). Przy kilku wersjach jednego PDF wygrywa najnowsza.
    cached_only: dict[str, Path] = {}
    for cached in sorted(cache_dir.glob("*.txt"), key=lambda path: path.stat().st_mtime):
        name = cached.name.rsplit(".", 2)[0] + ".pdf"
        if name not in raw:
            cached_only[name] = cached
    for name, cached in cached_only.items():
        raw[name] = cached.read_text(encoding="utf-8")

    texts: dict[str, list[str]] = {}
    # Kolejnosc jak przy sortowaniu plikow w katalogu - od niej zalezy, ktore dopasowanie wygrywa.
    for name in sorted(raw, key=lambda item: pdf_dir / item):
        texts[name] = [line.strip() for line in raw[name].split("\n")]
    return texts


def pdf_number(text: str) -> float | None:
    text = text.replace(" ", " ").replace("\xa0", " ").strip()
    text = re.sub(r"(?<=\d) (?=\d{3}$)", "", text)
    if re.fullmatch(r"\d+(?:[.,]\d+)?", text):
        return to_float(text)
    return None


def lumen_from_window(window: list[str], power: float, cct_k: set[int], skip_cct: bool = False) -> float | None:
    values = [v for v in (pdf_number(x) for x in window) if v is not None]
    power_pos = [i for i, v in enumerate(values) if abs(v - power) <= 0.051]
    if not power_pos:
        return None  # bez mocy w wierszu nie wiadomo, czy to ten produkt

    def plausible(v: float) -> bool:
        return v >= 10 and abs(v - power) > 0.05 and 35 <= v / power <= 260

    cands = sorted({v for v in values if plausible(v)})
    if not cands:
        return None
    # Tabele katalogow stawiaja lm tuz obok mocy (W|lm albo lm|W); inne liczby w wierszu
    # (zamiennik 40 W, CRI 80, wymiary, trwalosc) bywaja przypadkiem "wiarygodne".
    adjacent = {values[j] for i in power_pos for j in (i - 1, i + 1)
                if 0 <= j < len(values) and plausible(values[j]) and values[j] not in LIFETIMES
                and not (skip_cct and int(values[j]) in cct_k)}  # karty Bemko: "4000 | 100 | 16000" - 4000 to barwa
    if len(adjacent) == 1:
        return adjacent.pop()
    confirmed = [c for c in cands if any(abs(c / power - e) <= max(2.0, 0.03 * e) for e in values if e != c and 50 <= e <= 260)]
    if len(set(confirmed)) == 1:
        return confirmed[0]
    rest = [c for c in cands if c not in LIFETIMES and int(c) not in cct_k]
    if len(set(rest)) == 1:
        return rest[0]
    return None


def lumens_from_pdfs(texts: dict[str, list[str]], products: dict[str, tuple[float, set[int]]],
                     by_code: bool = False) -> dict[str, dict[str, Any]]:
    """Wiersz tabeli zaczyna sie od EAN albo (`by_code`, karty Bemko: kod | EAN | opak. | kolor | K | W | lm) od kodu."""
    found: dict[str, dict[str, set]] = defaultdict(lambda: {"lm": set(), "src": set(), "beam": set()})
    for fname, lines in texts.items():
        for i, line in enumerate(lines):
            if line not in products or not (by_code or EAN_RE.match(line)):
                continue
            power, cct_k = products[line]
            j = i + 1
            while j < len(lines) and (EAN_RE.match(lines[j]) or not lines[j]):
                j += 1
            window = []
            while j < len(lines) and len(window) < 18 and not EAN_RE.match(lines[j]):
                window.append(lines[j])
                j += 1
            lm = lumen_from_window(window, power, cct_k, skip_cct=by_code)
            if lm:
                found[line]["lm"].add(lm)
                found[line]["src"].add(fname)
                # kat swiecenia z tego samego wiersza tabeli ("36°", "36 °")
                for cell in window:
                    m = re.fullmatch(r"(\d{1,3})\s?°", cell.replace("\xa0", " "))
                    if m and int(m.group(1)) in BULB_BEAMS:
                        found[line]["beam"].add(int(m.group(1)))
    out = {}
    for ean, data in found.items():
        out[ean] = {"lm": sorted(data["lm"]), "src": sorted(data["src"]), "beam": sorted(data["beam"])}
    return out


# ---------------------------------------------------------------- opakowanie z PDF (marka, sztuki, barwa)

ROW_EAN_RE = re.compile(r"^(?:(.*\S)\s+)?(\d{13})$")


def package_from_pdfs(texts: dict[str, list[str]], columns: dict[str, str],
                      trusted: list[str] | None = None,
                      preferred: list[str] | None = None) -> dict[str, dict[str, Any]]:
    """EAN -> wiersz katalogu: nazwa, marka kolumny EAN, liczby z wiersza, sztuk w opakowaniu.

    Polski katalog Ledvance ma wiersz "Nazwa | OSRAM EAN | LEDVANCE EAN | BELLALUX EAN" - kolejne linie
    z samym EAN to ta sama lampa pod inna marka. Katalog retail konczy wiersz liczbami
    "dlugosc | srednica | sztuk w opakowaniu | sztuk w kartonie". Sztuki i nazwa wiersza (do kontroli
    barwy) tylko z katalogow `trusted` - w innych ta kolumna to karton albo moc, a oprawy maja tabele,
    w ktorych nazwa nie stoi przy swoim EAN. Nazwa wiersza pochodzi z pierwszego katalogu, ktory zna
    EAN: najpierw `trusted`, potem `preferred` (katalogi lamp z nazwa przy EAN), potem reszta.
    """
    trusted = trusted or []
    preferred = preferred or []
    out: dict[str, dict[str, Any]] = {}

    def rank(f: str) -> tuple[int, str]:
        return (0 if any(t in f for t in trusted) else 1 if any(p in f for p in preferred) else 2, f)

    for fname in sorted(texts, key=rank):
        lines = texts[fname]
        is_trusted = any(t in fname for t in trusted)
        header: list[str] = []
        last_name = ""
        i = 0
        while i < len(lines):
            line = lines[i]
            if line.startswith("Nazwa produktu"):
                header = [w.strip().upper() for w in lines[i + 1:i + 10] if w.strip().upper() in columns]
            m = ROW_EAN_RE.match(line)
            if not m:
                if re.search(r"[A-Za-z]{2}", line) and len(line) > 6:
                    last_name = line
                i += 1
                continue
            name = m.group(1) or last_name
            run = [m.group(2)]
            j = i + 1
            while j < len(lines) and EAN_RE.match(lines[j]):
                run.append(lines[j])
                j += 1
            k = j
            while k < len(lines) and not ROW_EAN_RE.match(lines[k]):
                k += 1
            cells = lines[j:k - 1]  # ostatnia linia przed nastepnym EAN to nazwa nastepnego wiersza
            # po trwalosci (15 000 h) ida: dlugosc | srednica | sztuk w opak. | sztuk w kartonie;
            # dalej bywa numer strony i przypisy "1 2 3 4" - nie naleza do wiersza
            pack = None
            for c, cell in enumerate(cells[:20]):
                if pdf_number(re.sub(r"[\u2009\u202f]", " ", cell)) in LIFETIMES and not re.fullmatch(r"\d+", cell):
                    tail = []
                    for nxt in cells[c + 1:c + 5]:
                        if not re.fullmatch(r"\d{1,4}", nxt):
                            break
                        tail.append(int(nxt))
                    pack = tail[2] if len(tail) == 4 else None
                    break
            numbers = [v for v in (pdf_number(c) for c in cells[:12]) if v is not None]
            for pos, ean in enumerate(run):
                rec = out.setdefault(ean, {"name": name, "src": fname, "numbers": numbers, "brand": "",
                                           "pack": None, "trusted": is_trusted})
                if len(run) > 1 and pos < len(header) and not rec["brand"]:
                    rec["brand"] = columns[header[pos]]  # marka z innego katalogu niz nazwa wiersza
                if is_trusted and pack and 1 <= pack <= 20 and rec["pack"] is None and rec["src"] == fname:
                    rec["pack"] = pack
            i = j
    return out


def catalog_cct(name: str) -> set[int]:
    """Barwy z kodu katalogowego ("GL 50 840" -> {4000}); zakres GLOWdim "822…827" pomijamy."""
    if "…" in name or "..." in name:
        return set()
    codes = re.findall(r"(?<![\d/.,])[89](2[247]|3\d|4\d|5\d|6[05])(?![\d.,])(?!\s*mm)", name, re.I)
    return {int(c) * 100 for c in codes}


def catalog_power_mismatch(name: str, power: float | None) -> bool:
    """Nazwa wiersza z jawna inna moca ("... 46W 865") - wiersz nie nalezy do tego produktu."""
    found = [to_float(x) for x in re.findall(r"(\d+(?:[.,]\d+)?)\s?W\b", name)]
    return bool(power and found and all(abs(v - power) > 0.051 for v in found))


def apply_package_brand(results: list[dict[str, Any]], rules: dict[str, Any],
                        pkg: dict[str, dict[str, Any]]) -> None:
    """Marka z opakowania (decyzja uzytkownika 2026-10-15): Osram/Bellalux zamiast domyslnej Ledvance."""
    cfg = rules.get("package_brand")
    if not cfg:
        return
    for res in results:
        ean = (res["row"].get("Kod producenta") or "").strip()
        from_pdf = pkg.get(ean, {}).get("brand", "")
        from_name = next((item["brand"] for item in cfg.get("name_rules", [])
                          if re.search(item["match"], ascii_upper(res["stara"]))), "")
        if from_pdf and from_name and from_pdf != from_name:
            res["uwagi"].append(f"marka: katalog {from_pdf}, kod w Ekonomie {from_name} - zostaje katalog")
        brand = from_pdf or from_name
        if not brand or brand == res.get("marka") or not res.get("marka"):
            continue
        name = re.sub(rf"(?<!\S){re.escape(res['marka'])}(?!\S)", brand, res["nazwa"], count=1)
        res["nazwa"] = re.sub(rf"(?<!\S){re.escape(brand)} {re.escape(brand)}(?!\S)", brand, name)
        res["uwagi"].append(f"marka z opakowania: {brand} ({'katalog' if from_pdf else 'kod w Ekonomie'})")
        res["marka"] = brand


def strip_brand_from_names(results: list[dict[str, Any]], rules: dict[str, Any]) -> None:
    """Jedna marka w nazwie (`brand_words`) albo nazwa bez marki (`brand_in_name: false`).

    Z marka (Ledvance, decyzja 2026-10-19 "jednak chcemy"): "Ledvance Bellalux" -> jedna marka (Osram /
    Bellalux wygrywa z domyslna Ledvance) w miejscu pierwszego slowa marki.
    Bez marki: gdy nazwy bez marki bylyby identyczne, a marki rozne (ta sama lampa pod marka Osram /
    Ledvance / Bellalux, osobne EAN), marka zostaje (jedna, w miejscu pierwszej).
    """
    words = [w for w in rules.get("brand_words") or [] if w]
    if not words:
        return
    pattern = re.compile(rf"(?<!\S)(?:{'|'.join(re.escape(w) for w in words)})(?!\S)", re.I)
    canon = {w.upper(): w for w in words}
    default = (rules.get("brand") or "").upper()

    def effective_brand(res: dict[str, Any]) -> str:
        found = [canon.get(m.upper(), m) for m in pattern.findall(res["nazwa"])]
        other = [b for b in found if b.upper() != default and b.upper() in {"OSRAM", "BELLALUX"}]
        return other[-1] if other else res.get("marka") or rules.get("brand", "")

    def one_brand(name: str, brand: str) -> str:
        marked = pattern.sub("\x00", name).replace("\x00", brand, 1).replace("\x00", "")
        return re.sub(r"\s{2,}", " ", marked).strip()

    if rules.get("brand_in_name", True):
        for res in results:
            if len(pattern.findall(res["nazwa"])) > 1:
                res["nazwa"] = one_brand(res["nazwa"], effective_brand(res))
        return

    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for res in results:
        res["_brand"] = effective_brand(res)
        res["_bez_marki"] = re.sub(r"\s{2,}", " ", pattern.sub("", res["nazwa"])).strip()
        groups[res["_bez_marki"]].append(res)
    for members in groups.values():
        keep = len({m["_brand"] for m in members}) > 1
        for res in members:
            if keep and pattern.search(res["nazwa"]):
                res["nazwa"] = one_brand(res["nazwa"], res["_brand"])
                res["uwagi"].append(f"marka {res['_brand']} zostaje - bez niej nazwa jak u innej marki")
            else:
                res["nazwa"] = res["_bez_marki"]
            del res["_brand"], res["_bez_marki"]


def diameter_in_mm(name: str) -> str:
    """Srednica "fi220" -> "220 mm", "fi30cm" -> "300 mm" (`diameter_mm`, Ledvance 2026-10-19)."""
    name = re.sub(r"(?<!\S)fi(\d+(?:,\d+)?)cm(?!\S)", lambda m: f"{fmt_num(to_float(m.group(1)) * 10)} mm", name)
    return re.sub(r"(?<!\S)fi(\d+(?:,\d+)?)(?!\S)", r"\1 mm", name)


def move_brand_tail(results: list[dict[str, Any]], rules: dict[str, Any]) -> None:
    """Marka z linia na koncu nazwy (`brand_last`, Ledvance 2026-10-19), przed "op. N szt.".

    Wersja i wyroznik sa dopisywane na koncu po zbudowaniu nazwy, wiec trafiaja za marke
    ("Zasilacz LED 150W 24V Ledvance DR-PFM-150 Performance") - przenosimy marke i slowa, ktore przy niej
    staly przy budowie (`po_marce`: linia, protokol Smart+), za nie.
    """
    if not rules.get("brand_last"):
        return
    words_cfg = [w for w in rules.get("brand_words") or [] if w]
    pattern = re.compile(rf"(?:{'|'.join(re.escape(w) for w in words_cfg)})", re.I) if words_cfg else None
    for res in results:
        head, pack = re.match(r"^(.*?)((?: op\. \d+ szt\.)?)$", res["nazwa"]).groups()
        words = head.split()
        # marka wielowyrazowa ("Kobi Pro", "GTV G-TECH") - szukamy calego ciagu; Ledvance: tez `brand_words`
        brand = [w.upper() for w in (res.get("marka") or "").split()]
        upper = [w.upper() for w in words]
        idx = next((i for i in range(len(words) - len(brand) + 1) if brand and upper[i:i + len(brand)] == brand), None)
        size = len(brand)
        if idx is None and pattern:
            idx, size = next((i for i, w in enumerate(words) if pattern.fullmatch(w)), None), 1
        if idx is None:
            continue
        after = res.get("po_marce", [])
        k = 0
        while k < len(after) and idx + size + k < len(words) and words[idx + size + k] == after[k]:
            k += 1
        tail = words[idx:idx + size + k]
        rest = words[:idx] + words[idx + size + k:]
        if res.get("linia") and k == 0:
            # linia dopisana pozniej jako wyroznik ("świeczka PCR Ledvance") tez stoi przy marce
            moved = [w for w in rest[1:] if w in LINE_LABELS]
            rest = rest[:1] + [w for w in rest[1:] if w not in LINE_LABELS]
            tail += moved
        res["nazwa"] = " ".join(rest + tail) + pack


def check_catalog_cct(results: list[dict[str, Any]], pkg: dict[str, dict[str, Any]]) -> None:
    """Barwa w nazwie inna niz w katalogu (dla tego EAN): moc zgodna -> barwa z katalogu, inaczej DO_SPRAWDZENIA."""
    for res in results:
        rec = pkg.get((res["row"].get("Kod producenta") or "").strip())
        m = re.search(r"\b(\d{4})K\b", res["nazwa"])
        if not rec or not rec.get("trusted") or not m:
            continue
        named = {int(k) for k in re.findall(r"(\d{4})(?=[/K])", res["nazwa"])}  # "3000/4000K" = obie barwy
        cct = catalog_cct(rec["name"])
        if len(cct) != 1 or cct & named:
            continue
        cct_k = cct.pop()
        _, power = parse_power(parse_view(res.get("zrodlo_stara") or res["stara"]))
        if catalog_power_mismatch(rec["name"], power):
            continue
        if power and any(abs(v - power) <= 0.051 for v in rec["numbers"]):
            res["nazwa"] = res["nazwa"].replace(m.group(0), f"{cct_k}K", 1)
            res["uwagi"].append(f"barwa z katalogu {cct_k}K (w Ekonomie {m.group(1)}K): {rec['name']}")
            res["status"] = "WARNING" if res["status"] == "OK" else res["status"]
        else:
            res["uwagi"].append(f"katalog opisuje inny produkt dla tego EAN: {rec['name']} "
                                f"({' | '.join(fmt_num(v) for v in rec['numbers'][:6])}) - nazwa bez zmian")
            res["status"] = "DO_SPRAWDZENIA"


def apply_pack_size(results: list[dict[str, Any]], rules: dict[str, Any], pkg: dict[str, dict[str, Any]],
                    web: dict[str, dict[str, Any]] | None = None) -> None:
    """Multipak (decyzja uzytkownika 2026-10-15): liczba sztuk w opakowaniu na koncu nazwy, gdy > 1.

    Zrodla: nazwa w Ekonomie, karta produktu ledvance.com/en/gtin (dokladnie ten EAN), katalog PDF.
    """
    cfg = rules.get("pack_size")
    if not cfg:
        return
    for res in results:
        site = (web or {}).get((res["row"].get("Kod producenta") or "").strip(), {})
        from_web = site.get("pieces") if site.get("exact") else None
        u = ascii_upper(res["stara"])
        from_name = next((int(m.group(1)) for p in cfg.get("name_patterns", []) for m in [re.search(p, u)] if m), None)
        if from_name is None:
            from_name = next((n for w, n in cfg.get("word_packs", {}).items() if re.search(rf"\b{w}\s*PACK\b", u)), None)
        if from_name is None:
            from_name = next((n for w, n in cfg.get("fixed_words", {}).items() if re.search(rf"\b{w}\b", u)), None)
        rec = pkg.get((res["row"].get("Kod producenta") or "").strip(), {})
        from_pdf = None
        if cfg.get("from_catalog") and rec.get("pack"):
            _, power = parse_power(parse_view(res["stara"]))
            if power and any(abs(v - power) <= 0.051 for v in rec["numbers"]):
                from_pdf = rec["pack"]  # tylko gdy wiersz katalogu ma moc z nazwy - to na pewno ten produkt
        if from_name and from_pdf and from_name != from_pdf:
            res["uwagi"].append(f"sztuk w opakowaniu: nazwa {from_name}, katalog {from_pdf} - zostaje nazwa")
        if from_web and (from_name or from_pdf) and from_web != (from_name or from_pdf):
            res["uwagi"].append(f"sztuk w opakowaniu: ledvance.com {from_web}, "
                                f"{'nazwa' if from_name else 'katalog'} {from_name or from_pdf} - do sprawdzenia")
        n = from_name or from_web or from_pdf
        src = "nazwa w Ekonomie" if from_name else ("ledvance.com" if from_web else "katalog")
        if n and 1 < n <= 20:
            # "LF-2TERM 10 PACK", "DOT-IT 3" (z "3 pack") - liczba sztuk tylko w dopisku
            name = re.sub(rf"(?i)\s+{n}\s*-?\s*PACK\b", "", res["nazwa"])
            if re.search(rf"\b{n}\s*-?\s*PACK\b", u):
                name = re.sub(rf"\s+{n}$", "", name)
            res["nazwa"] = f"{name} {cfg['suffix'].format(n=n)}"
            res["uwagi"].append(f"multipak {n} szt. ({src})")


CATALOG_SOCKET_FOOTNOTE = re.compile(r"\b(E27|E14|GU10|G53|GU5\.3|G9|G4|B22d|R7s)\d\b")
CCT_CODE = re.compile(r"[89](18|22|24|25|27|30|35|40|50|57|60|65|80)")  # kod barwy Ledvance: 8xx/9xx = CRI 80/90
MODEL_WORD = re.compile(r"[A-Z][A-Za-z+]*|[A-Z0-9+-]*[A-Z][A-Z0-9+-]*")  # "Globe", "HALOLINE", "T26"; nie "matowe"


def name_tokens(text: str, letters: dict[str, str] | None = None) -> list[tuple[str, str, str, str]]:
    """Kandydaci na wyroznik z nazwy (katalogowej, z Ekonoma albo ze strony): (klucz, oryginal, slowo przed, rodzaj).

    Rodzaj: "word" (slowo), "num" (liczba: moc "20W", barwa "3000K", dlugosc "3m", liczba ze slowem
    modelu przed nia "Globe 80", kod "64695", zakres "1200-840"), "volt" (napiecie - wyroznik tylko,
    gdy podaja je wszyscy w grupie). Moc i barwa odpadaja pozniej, jesli sa juz w nazwie. Pomija kody
    barwy 8xx/9xx, luzne litery (poza klasami V/P/S/E z `letters`) i sklejone kody zamowieniowe
    ("740RW20STTWAC2M7D4EUELEDV", "230VTWFRB22DFS1LEDV").
    """
    text = CATALOG_SOCKET_FOOTNOTE.sub(r"\1", text)  # "E271" = E27 + przypis 1
    text = re.sub(r"(?<=[A-Za-z])1$|,\s*\d$", "", text.strip())  # "Non-DIM1", "E27, 5"
    text = re.sub(r"(?<=\d)\s+(W|V|K|lm|mm|cm)\b", r"\1", text, flags=re.I)  # "35 W 240 V" -> "35W 240V"
    text = re.sub(r"\b(\d)\s+pins?\b", r"\1-pin", text, flags=re.I)  # CFL SQUARE "2 pins" / "4 pins"
    toks = [t.strip(",.()") for t in text.replace("/", " ").split()]
    out = []
    for i, tok in enumerate(toks):
        key = ascii_upper(tok).replace(",", ".")
        if not key:
            continue
        if SOCKET_RE.fullmatch(key):
            out.append((key, tok, "", "word"))  # gwint - wyroznik tylko, gdy nie ma go w nazwie (E14 / E27)
            continue
        if re.fullmatch(r"\d+(?:\.\d+)?(?:M|CM)", key):
            out.append((key, tok.lower().replace(".", ","), "", "num"))  # dlugosc tasmy/szyny: "3m", "0,8m"
            continue
        if re.fullmatch(r"\d+(?:\.\d+)?W|\d{4}K|\d{3,5}LM", key):
            label = re.sub(r"(?i)lm$", "lm", tok.replace(".", ","))
            out.append((key, label, "", "num"))  # moc / barwa / strumien - gdy nie ma ich w nazwie
            continue
        if len(key) == 1:
            if key in (letters or {}) and ascii_upper(toks[i - 1] if i else "") not in ("CLASSIC", "CL"):  # "CLASSIC P" = kulka
                out.append((key, letters[key], "", "word"))
            continue
        if re.search(r"\d", key) and (re.search(r"(EUE|LEDV|OSRAM)$", key) or (len(key) > 10 and "-" not in key)):
            continue  # sklejony kod zamowieniowy
        if re.fullmatch(r"[\d.]+(?:-[\d.]+)?V(?:AC|DC)?", key):
            out.append((key, tok, "", "volt"))
            continue
        if re.fullmatch(r"\d{2,5}-\d{2,5}", key):
            out.append((key, tok, "", "num"))  # "1200-840" = strumien-barwa modulu
            continue
        if re.fullmatch(r"[89]\d{4}", key) and CCT_CODE.fullmatch(key[:3]):
            if key[0] == "9":  # "93040" = barwa przelaczana 3000/4000K, CRI 90 ("83040" - standard)
                out.append(("CRI90", "CRI90", "", "word"))
            continue
        if CCT_CODE.fullmatch(key) and not re.search(r"\d-\d", toks[i - 1] if i else ""):
            # kod barwy (827, 940): barwa, gdy nazwa jej nie ma (tasma 830/840/865), a 9xx to CRI 90
            # (8xx to standard - bez dopisku); nie "220-240/925" - tam 925 to prad zasilacza
            out.append((f"{int(key[1:]) * 100}K", f"{int(key[1:]) * 100}K", "", "num"))
            if key[0] == "9":
                out.append(("CRI90", "CRI90", "", "word"))
            continue
        if re.fullmatch(r"[789]\d\d|[\d.]+(?:K|MM|MA|A|M)|[\d.]+", key):
            prev = toks[i - 1] if i else ""
            if re.fullmatch(r"\d{2,5}", key) and not re.fullmatch(r"[789]\d\d", key):
                if (len(prev) > 1 and MODEL_WORD.fullmatch(prev) and not SOCKET_RE.fullmatch(ascii_upper(prev))
                        and not re.fullmatch(r"(?i)W|V|K|LM|MM|CM|DIM", prev)):
                    out.append((key, tok, prev, "num"))
                elif len(key) == 5:
                    out.append((key, tok, "", "num"))  # numer katalogowy halogenu "64695"
            continue
        out.append((key, tok, "", "word"))
    return out


def distinguish_label(text: str) -> str:
    """"GLOBE95" -> "Globe 95", "STRAW" -> "Straw"; "T26", "Spiral" bez zmian."""
    text = re.sub(r"(?<=[A-Za-z]{3})(?=\d)", " ", text)
    return " ".join(w.capitalize() if w.isupper() and len(w) > 3 and w.isalpha() else w for w in text.split())


def minimal_keys(own: list[str], numeric: set[str], others: list[set[str]], limit: int = 3) -> list[str]:
    """Najmniejszy zestaw slow, ktorego zaden inny produkt z grupy nie ma w calosci.

    Slowa przed liczbami ("Spiral" przed "80"), potem mniejszy zestaw, potem kolejnosc z nazwy.
    Pusta lista = produktu nie da sie odroznic tymi slowami.
    """
    cands = []
    for size in range(1, min(limit, len(own)) + 1):
        for combo in combinations(range(len(own)), size):
            keys = {own[c] for c in combo}
            if all(not keys <= other for other in others):
                cands.append((sum(own[c] in numeric for c in combo), size, combo))
    if not cands:
        return []
    return [own[c] for c in min(cands)[2]]


def color_label(key: str, gender: str, colors: dict[str, Any]) -> str:
    """Kod koloru ze starej nazwy ("BK", "GY") jako kolor w rodzaju nazwy ("czarny" zasilacz)."""
    for color in colors.values():
        if any(re.search(p, key) for p in color.get("match", [])):
            return color.get(gender, color.get("f", ""))
    return ""


def add_distinguishers(results: list[dict[str, Any]], rules: dict[str, Any], source: str,
                       text_of: Any, polish: bool = False) -> None:
    """Wyroznik dla nazw, ktore nadal sie powtarzaja (decyzja uzytkownika 2026-10-15: tyle, zeby
    sie roznily, nie wiecej). `text_of(res)` zwraca (tekst zrodla, styl) albo None; porownujemy tylko
    produkty opisane tym samym stylem - "LED STAR CLASSIC A FR 40" (retail) i "LED Classic A 40 ...
    Frosted" (broszura) to ten sam produkt, nie roznica. Wyroznik stoi przed "op. N szt.".
    """
    move_brand_tail(results, rules)  # "G9 Ledvance 35W" i "G9 35W Ledvance" to ta sama nazwa
    cfg = rules.get("distinguish", {})
    words = {ascii_upper(k): v for k, v in cfg.get("words", {}).items()}
    stop = {ascii_upper(w) for w in cfg.get("stop", [])}
    stop_re = [re.compile(p) for p in cfg.get("stop_patterns", [])]
    conflicts = [set(g) for g in cfg.get("conflicts", [])]
    colors = rules.get("housing_colors", {})
    letters = rules.get("version_letters", {})
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for res in results:
        groups[res["nazwa"]].append(res)
    for name, group in groups.items():
        if len(group) < 2:
            continue
        by_style: dict[str, list[tuple[dict[str, Any], str]]] = defaultdict(list)
        for m in group:
            got = text_of(m)
            if got and got[0]:
                by_style[got[1]].append((m, got[0]))
        head, pack = re.match(r"^(.*?)((?: op\. \d+ szt\.)?)$", name).groups()
        present = set(ascii_upper(head).replace("+", "").replace(",", ".").split())  # "6,5W" = "6.5W"
        present |= {t[1:] for t in present if re.fullmatch(r"G\d{2,3}", t)}  # "globe G125" = "Globe 125"
        last = ascii_upper(head.split()[-1])
        first = head.split()[0].lower()  # rodzaj do koloru: "Zasilacz" czarny, "Oprawa" czarna
        gender = "f" if first.endswith("a") else ("n" if first.endswith(("o", "e")) else "m")

        def usable(x: tuple[str, str, str, str]) -> bool:
            key, orig, prev = x[0], x[1], x[2]
            # litera klasy "V" ma oryginal "Value" - sprawdzamy oba, inaczej kolejna runda dopisze "Value Value"
            if key in stop or key in present or ascii_upper(orig) in present or any(p.fullmatch(key) for p in stop_re):
                return False
            if prev and (ascii_upper(prev) in stop or any(p.fullmatch(ascii_upper(prev)) for p in stop_re)):
                return False  # "CLP 40", "FR 40", "T8 36" - zamiennik mocy po kodzie nic nie mowi (2026-10-21)
            color = color_label(key, gender, colors)
            if color and ascii_upper(color) in ascii_upper(head):
                return False  # kolor jest juz w nazwie ("lustrzana srebrna" + SILVER)
            if key in words:  # null = slowo bez znaczenia; tlumaczenie juz w nazwie ("przezroczysta")
                return words[key] is not None and ascii_upper(words[key]) not in ascii_upper(head)
            # polskie slowa opisowe ze starej nazwy ("możliwością"); w nazwach angielskich "pins" itp. zostaja
            return not (polish and re.fullmatch(r"[a-ząćęłńóśźż-]+", orig))

        for pairs in by_style.values():
            if len(pairs) < 2:
                continue
            toks = [[x for x in name_tokens(text, letters) if usable(x)] for _, text in pairs]
            if not all(any(x[3] == "volt" for x in t) for t in toks):
                toks = [[x for x in t if x[3] != "volt"] for t in toks]  # napiecie tylko gdy podaja je wszyscy
            sets = [{x[0] for x in t} for t in toks]
            common = set.intersection(*sets)
            for idx, ((m, text), t) in enumerate(zip(pairs, toks)):
                own = list(dict.fromkeys(x[0] for x in t if x[0] not in common))
                numeric = {x[0] for x in t if x[3] != "word"}
                # bliznaki (te same slowa) i tak zostana razem - nie blokuja oddzielenia od reszty grupy
                chosen = set(minimal_keys(own, numeric, [s for j, s in enumerate(sets) if j != idx and s != sets[idx]]))
                if not chosen:
                    continue
                labels: list[str] = []
                for key, orig, prev, _kind in t:
                    if key not in chosen:
                        continue
                    chosen.discard(key)
                    color = color_label(key, gender, colors)
                    if key in words:
                        labels.append(words[key])
                    elif color:
                        labels.append(color)
                    elif prev and ascii_upper(prev) != last:
                        labels.append(distinguish_label(f"{prev} {orig}"))  # "Globe 80"
                    else:
                        labels.append(distinguish_label(orig))
                clash = next((f"{lab} (w nazwie: {other})" for lab in labels for grp in conflicts if lab in grp
                              for other in grp if other != lab and re.search(rf"(?<!\w){other}(?!\w)", head)), "")
                if clash:  # zrodlo przeczy nazwie z Ekonoma ("matowa" przy "przezroczysta") - nie zgadujemy
                    m["uwagi"].append(f"{source} podaje {clash} - do sprawdzenia: {text}")
                    m["status"] = "DO_SPRAWDZENIA"
                    continue
                suffix = " ".join(labels)
                if head.endswith(" z czujnikiem") and suffix.startswith("z czujnikiem "):
                    suffix = suffix[len("z czujnikiem "):]  # "... z czujnikiem" + "zmierzchu"
                m["nazwa"] = f"{head} {suffix}{pack}"
                m["uwagi"].append(f"wyroznik ({source}): {suffix} <- {text}")


def catalog_style(rec: dict[str, Any], rules: dict[str, Any]) -> str:
    """Styl zapisu nazw katalogu PDF (`distinguish.catalogs`: fragment nazwy pliku -> styl)."""
    return next((f for c, f in rules.get("distinguish", {}).get("catalogs", {}).items()
                 if c in rec.get("src", "")), "")


def load_web_names(cache_dir: Path) -> dict[str, dict[str, Any]]:
    """EAN -> {title, exact, pieces} z ledvance.com (cache `scripts/fetch_ledvance_web.py`).

    Karta `/en/gtin/<EAN>` to dokladnie ten produkt (`exact`, z liczba sztuk w opakowaniu); podpowiedz
    wyszukiwarki tylko gdy karty nie ma i trafienie jest jedno.
    """
    out: dict[str, dict[str, Any]] = {}
    for path in cache_dir.glob("*.json"):
        rec = json.loads(path.read_text(encoding="utf-8"))
        page = rec.get("gtin_page")
        if page:
            title = f"{page['name']} {page.get('order_reference', '')}".strip()
            out[rec["ean"]] = {"title": title, "exact": True, "pieces": page.get("pieces"),
                               "name": page["name"], "url": page.get("url", ""), "cri": page.get("cri"),
                               "lumen": page.get("lumen"), "power": page.get("power"), "cct": page.get("cct"),
                               "lumens": page.get("lumens") or [], "powers": page.get("powers") or [],
                               "family": page.get("family", ""), "beam": page.get("beam")}
        elif len(rec.get("products", [])) == 1:  # kilka trafien = EAN niejednoznaczny, pomijamy
            out[rec["ean"]] = {"title": rec["products"][0]["title"], "exact": False, "pieces": None,
                               "family": rec["products"][0].get("description", "")}
    return out


def add_all_distinguishers(results: list[dict[str, Any]], rules: dict[str, Any],
                           pkg: dict[str, dict[str, Any]], web: dict[str, dict[str, Any]]) -> None:
    """Kolejnosc zrodel: katalog PDF, stara nazwa w Ekonomie, strona producenta."""
    if not rules.get("distinguish"):
        return

    def ean(m: dict[str, Any]) -> str:
        return (m["row"].get("Kod producenta") or "").strip()

    def from_pdf(m: dict[str, Any]) -> tuple[str, str] | None:
        rec = pkg.get(ean(m), {})
        style = catalog_style(rec, rules)
        return (rec["name"], style) if style else None

    def from_web(m: dict[str, Any]) -> tuple[str, str] | None:
        title = web.get(ean(m), {}).get("title")
        if not title:
            return None
        clash = web_mismatch(m.get("zrodlo_stara") or m["stara"], title)
        if clash:  # strona opisuje pod tym EAN cos innego niz Ekonom - nie dopisujemy, do sprawdzenia
            note = f"ledvance.com pod tym EAN: {title} ({clash}) - do sprawdzenia"
            if note not in m["uwagi"]:
                m["uwagi"].append(note)
            return None
        return title, "web"

    add_card_cri(results, rules, web)
    # kilka rund: "-150" nie odrozni sie od "-150 OEM" w jednej rundzie, ale po oddzieleniu wersji OEM tak
    for _ in range(3):
        before = [m["nazwa"] for m in results]
        add_distinguishers(results, rules, "katalog", from_pdf)
        add_distinguishers(results, rules, "nazwa w Ekonomie",
                           lambda m: (m.get("zrodlo_stara") or m["stara"], "ekonom"), polish=True)
        add_distinguishers(results, rules, "ledvance.com", from_web)
        if before == [m["nazwa"] for m in results]:
            break


CARD_LINE_WORDS = {"VALUE": "Value", "PERFORMANCE": "Performance", "SUPERIOR": "Superior"}


def card_family_lines(web: dict[str, dict[str, Any]], placeholders: list[str] | None = None) -> dict[str, str]:
    """EAN -> linia jego rodziny z kart GTIN (`line_from_card_family`, decyzja uzytkownika 2026-10-20).

    Rodzina ("LED CLASSIC B V") dostaje linie tylko wtedy, gdy ktoras karta tej rodziny zapisuje ja slowem
    ("VALUE CLAS B 60 6.8 W") i zadna karta jej nie przeczy - wtedy producent potwierdza znaczenie litery.
    """
    found: dict[str, set[str]] = defaultdict(set)
    # "Phase out" / "Coming soon" to nie rodzina, tylko status - laczyl tasmy, zasilacze i swietlowki (2026-10-20)
    web = {ean: card for ean, card in web.items()
           if not any(re.search(p, (card.get("family") or "").strip()) for p in placeholders or [])}
    for card in web.values():
        if card.get("exact") and card.get("family"):
            name = ascii_upper(card.get("name", ""))
            found[card["family"]] |= {line for word, line in CARD_LINE_WORDS.items() if re.search(rf"\b{word}\b", name)}
    lines = {fam: next(iter(ls)) for fam, ls in found.items() if len(ls) == 1}
    return {ean: lines[card["family"]] for ean, card in web.items()
            if card.get("exact") and card.get("family") in lines}


def card_lumen(card: dict[str, Any], power: float | None, cct: str, result: dict[str, Any],
               power_text: str = "") -> float | str | None:
    """Brakujace lm z karty ledvance.com/en/gtin/<EAN> (`lumen_from_gtin_card`, 2026-10-20).

    Karta to dokladnie ten EAN, ale kartoteka w Ekonomie bywa kopia sasiedniej - lm bierzemy tylko, gdy moc
    (i barwa, gdy jest jedna) z karty zgadza sie z nazwa; inaczej uwaga i lm zostaje puste.
    """
    if not card.get("exact") or not card.get("lumen"):
        return None
    where = card.get("url") or "ledvance.com"
    pairs = list(zip(card.get("powers") or [], card.get("lumens") or []))
    watts = [to_float(w) for w in re.findall(r"\d+(?:,\d+)?", power_text)] if power is None else []
    if len(watts) > 1 and len(pairs) > 1:
        # nazwa z kilkoma mocami ("11,3/14W", "900/1200W", "4-7-9W") - lm kazdej z karty, ta sama kolejnosc i separator
        found = [next((lm for w, lm in pairs if abs(w - watt) <= 0.051), None) for watt in watts]
        named = re.fullmatch(r"(\d{4})K", cct or "")
        if all(found) and not (named and card.get("cct") and int(named.group(1)) != card["cct"]):
            result["lm_zrodlo"] = f"karta {card.get('url') or 'ledvance.com'}"
            return ("/" if "/" in power_text else "-").join(fmt_num(lm) for lm in found)
    if len(pairs) > 1 and power is not None:
        # Power Select: "55 W / 49 W / 43 W" + "6650 lm / 5900 lm / 5250 lm" - lm stopnia o mocy z Ekonoma
        hit = next(((w, lm) for w, lm in pairs if abs(w - power) <= 0.051), None)
        top = max(pairs)
        if not hit and abs(top[0] - power) <= 0.05 * power:
            # Multi Lumen: w Ekonomie moc zaokraglona do najwyzszego stopnia ("40W" = 39,3 W na karcie)
            hit = top
            result["uwagi"].append(f"lm najwyzszego stopnia z karty: {fmt_num(top[0])}W / {fmt_num(top[1])}lm")
        if hit:
            card = {**card, "power": power, "lumen": hit[1]}  # stopien dopasowany do mocy z Ekonoma
    if power and card.get("power") and len(pairs) <= 1:
        sets = card["power"] / power
        if round(sets) >= 2 and abs(sets - round(sets)) < 0.01:
            # komplet ("PANEL ... 33W SET6"): karta podaje moc i lm calego zestawu - lm na sztuke
            k = round(sets)
            card = {**card, "power": power, "lumen": round(card["lumen"] / k)}
            result["uwagi"].append(f"karta podaje komplet {k} szt. - lm na sztuke")
    if power is None or card.get("power") is None or abs(card["power"] - power) > 0.051:
        result["uwagi"].append(f"karta {where} podaje {fmt_num(card['lumen'])}lm przy "
                               f"{fmt_num(card['power']) if card.get('power') else '?'}W - moc w Ekonomie inna, lm pominiete")
        return None
    named = re.fullmatch(r"(\d{4})K", cct or "")
    if named and card.get("cct") and int(named.group(1)) != card["cct"]:
        result["uwagi"].append(f"karta {where} podaje {fmt_num(card['lumen'])}lm przy {card['cct']}K"
                               " - barwa w Ekonomie inna, lm pominiete")
        return None
    result["lm_zrodlo"] = f"karta {where}"
    return card["lumen"]


def add_card_cri(results: list[dict[str, Any]], rules: dict[str, Any], web: dict[str, dict[str, Any]]) -> None:
    """Powtorzone nazwy rozniace sie CRI na karcie GTIN dostaja "CRI 97" / "CRI 90" z karty (2026-10-20).

    "LED MR16 DIM S" i "DIM P" maja te sama moc, barwe i lm - rozni je CRI 97 / 90. Litera S / P tego nie
    dowodzi, karta producenta pod dokladnie tym EAN tak. Karta opisujaca inny produkt niz Ekonom jest pomijana.
    """
    move_brand_tail(results, rules)
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for res in results:
        groups[res["nazwa"]].append(res)
    for members in groups.values():
        if len(members) < 2:
            continue
        cards = [web.get((m["row"].get("Kod producenta") or "").strip(), {}) for m in members]
        if not all(c.get("exact") and c.get("cri") for c in cards) or len({c["cri"] for c in cards}) < 2:
            continue
        if min(c["cri"] for c in cards) < 90:
            continue  # tylko lampy wysokiego CRI (97 / 90); "CRI 80" to standard, nic nie mowi sprzedawcy
        if any(web_mismatch(m.get("zrodlo_stara") or m["stara"], c["title"]) for m, c in zip(members, cards)):
            continue
        for m, c in zip(members, cards):
            head, pack = re.match(r"^(.*?)((?: op\. \d+ szt\.)?)$", m["nazwa"]).groups()
            m["nazwa"] = f"{head} CRI {c['cri']}{pack}"
            m["uwagi"].append(f"wyroznik (karta {c.get('url') or 'ledvance.com'}): CRI {c['cri']}")


def web_mismatch(old: str, title: str) -> str:
    """Rozbieznosc mocy, kodu barwy albo gwintu miedzy nazwa w Ekonomie a nazwa ze strony producenta."""
    u_old, u_web = parse_view(old), parse_view(title)
    _, p_old = parse_power(u_old)
    _, p_web = parse_power(u_web)
    if p_old and p_web and abs(p_old - p_web) > 0.051:
        return f"moc {fmt_num(p_old)}W / {fmt_num(p_web)}W"
    k_old = {k[1:] for k in re.findall(r"(?<!\d)[89]\d\d(?!\d)", u_old)} | {k[:2] for k in re.findall(r"(\d{4}) ?K\b", u_old)}
    k_web = {k[1:] for k in re.findall(r"(?<!\d)[89]\d\d(?!\d)", u_web)} | {k[:2] for k in re.findall(r"(\d{4}) ?K\b", u_web)}
    if k_old and k_web and not k_old & k_web:
        return "inna barwa"
    s_old = {s.upper() for s in SOCKET_RE.findall(u_old)}
    s_web = {s.upper() for s in SOCKET_RE.findall(u_web)}
    if s_old and s_web and not s_old & s_web:
        return "inny gwint"
    return ""


def web_name_mismatch(old: str, page: str, code: str, rules: dict[str, Any]) -> str:
    """Nazwa w Ekonomie opisuje inny produkt niz karta producenta pod tym kodem (Bemko: kopia sasiedniej kartoteki).

    Oprocz mocy/barwy/gwintu (`web_mismatch`) porownuje prady (63A, 100mA) i kontrole kodu (`code_checks`):
    kod A03-V7-2-63-100 z nazwa "25A 300mA" przeczy Ekonomowi, a z nazwa karty "63A 100mA" - nie.
    """
    why = web_mismatch(old, page)
    if why:
        return why
    u_old, u_web = parse_view(old), parse_view(page)
    for unit, label in (("MA", "prad roznicowy"), ("A", "prad")):
        a_old = set(re.findall(rf"(?<![\d.,])(\d+)\s?{unit}\b", u_old))
        a_web = set(re.findall(rf"(?<![\d.,])(\d+)\s?{unit}\b", u_web))
        if a_old and a_web and a_old != a_web:
            return f"{label} {'/'.join(sorted(a_old))} / {'/'.join(sorted(a_web))}"
    if code_mismatches(code, u_old, rules) and not code_mismatches(code, u_web, rules):
        return "kod przeczy nazwie w Ekonomie, zgadza sie z karta"
    return ""


def gtin_other_product(old_u: str, page_u: str, s_old: set[str], s_page: set[str], p_old: float | None,
                       p_page: float | None, rules: dict[str, Any]) -> str:
    """Czym karta GTIN przeczy nazwie z Ekonoma tak, ze to inny produkt, a nie literowka (pusty napis = ten sam).

    Inny trzonek ("E27" / karta "B22d"), inny ksztalt banki ("A60" / "swieczka") albo inna dlugosc tuby ("600" / 120 cm)
    - decyzja uzytkownika 2026-10-21. Sama inna moc albo barwa nie wystarcza.
    """
    def norm(sockets: set[str]) -> set[str]:
        return {re.sub(r"[^A-Z0-9]", "", s.upper()) for s in sockets}  # "GX24Q.-3" w Ekonomie = "GX24Q-3" na karcie

    if s_old and s_page and not norm(s_old) & norm(s_page):
        return f"trzonek {'/'.join(sorted(s_old))} / karta {'/'.join(sorted(s_page))}"
    shape_old, shape_page = parse_shape(old_u), parse_shape(page_u)
    if shape_old and shape_page and shape_old.split()[0].lower() != shape_page.split()[0].lower():
        return f"kształt {shape_old} / karta {shape_page}"
    t_old, t_page = parse_tube_type(old_u), parse_tube_type(page_u)
    if t_old or t_page:
        len_old = parse_tube_length(old_u, t_old, p_old, rules)
        len_page = parse_tube_length(page_u, t_page or t_old, p_page, rules)
        if len_old and len_page and len_old != len_page:
            return f"długość tuby {len_old} / karta {len_page}"
    return ""


def apply_web_card(res: dict[str, Any], row: dict[str, str], rules: dict[str, Any],
                   pdf_lm: dict[str, dict[str, Any]], feed: dict[str, dict[str, str]] | None) -> dict[str, Any]:
    """Karta producenta pod dokladnie tym kodem (`fix_from_web_name`, Bemko).

    Nazwa w Ekonomie opisuje inny produkt (`web_name_mismatch`) -> nazwa budowana z karty, WARNING ze stara
    nazwa w uwagach. Nazwa bez marki -> marka z pola karty (`web_brand_key`) zamiast zgadywania z rodzenstwa.
    """
    ean = (row.get("Kod producenta") or "").strip()
    page = rules.get("_web_params", {}).get(ean, {}) if rules.get("fix_from_web_name") else {}
    brand = page.get(rules.get("web_brand_key", ""), "")
    why = web_name_mismatch(res["stara"], page["_name"], ean, rules) if page.get("_name") else ""
    if not why and not (brand and not res.get("marka_jawna")):
        return res
    if why:
        fixed = {**row, "Nazwa towaru": f"{page['_name']} {brand}".strip(), "Nazwa dodatkowa": ""}
    else:
        fixed = {**row, "Nazwa dodatkowa": f"{row.get('Nazwa dodatkowa', '')} {brand}".strip()}
    new = build_name(fixed, rules, pdf_lm, feed)
    new.update(row=row, stara=res["stara"])
    if why:
        new["uwagi"].append(f"nazwa z karty {page.get('_url') or 'producenta'} ({why}; w Ekonomie: {res['stara']})")
        new["status"] = "WARNING" if new["status"] == "OK" else new["status"]
    elif brand != res.get("marka"):
        new["uwagi"].append(f"marka {brand} z karty producenta (nazwa w Ekonomie bez marki)")
    return new


def nice_case(text: str, keep: set[str]) -> str:
    """WIELKIE LITERY GTV -> male, z wyjatkiem skrotow (USB, HDMI, RJ45) i tokenow z cyframi."""
    out = []
    for tok in text.split():
        core = re.sub(r"[^\wĄĆĘŁŃÓŚŹŻąćęłńóśźż+]", "", tok)
        if not core or core.upper() in keep or re.search(r"\d", core) or not tok.isupper():
            out.append(tok)
        else:
            out.append(tok.lower())
    return " ".join(out)


def split_segments(text: str) -> list[str]:
    """Czlony nazwy GTV po przecinkach; przecinek miedzy cyframi ("3X1,0mm2", "1,5m") nie dzieli."""
    return re.split(r"(?<!\d),|,(?!\d)", text)


def parse_segments(original: str, rules: dict[str, Any]) -> str:
    """Czlony nazwy po pierwszym przecinku bez szumu (zasilanie, PF, Ra, gwarancja) - dla typow bez schematu."""
    parts = []
    for seg in split_segments(original)[1:]:
        seg = " ".join(seg.split())
        if not seg or any(re.search(p, ascii_upper(seg)) for p in rules.get("segment_drop", [])):
            continue
        parts.append(nice_case(seg, set(rules.get("keep_upper", []))))
    return " ".join(parts)


PDF_CODE_RE = re.compile(r"^[A-Z]{1,4}\d?-[A-Z0-9][A-Z0-9.,/+-]*$")


def lumens_from_pdf_codes(texts: dict[str, list[str]], codes: set[str]) -> dict[str, dict[str, Any]]:
    """Katalog z tabelami po kodzie (GTV): wiersz od kodu do nastepnego kodu, wartosci z jednostkami.

    Bierze tylko komorki z jednostka ("1000 lm", "650 lm/m", "3200 K"), bez zgadywania liczb.
    """
    found: dict[str, dict[str, set]] = defaultdict(lambda: {"lm": set(), "per_m": set(), "k": set(), "src": set()})
    for fname, lines in texts.items():
        for i, line in enumerate(lines):
            if line not in codes:
                continue
            for cell in lines[i + 1:i + 16]:
                if cell in codes or PDF_CODE_RE.match(cell):
                    break
                cell = cell.replace("\xa0", " ").strip()
                m = re.fullmatch(r"(\d{1,3}(?: \d{3})?|\d{2,6})\s?lm(/m)?", cell, re.IGNORECASE)
                if m:
                    found[line]["lm"].add(to_float(m.group(1).replace(" ", "")))
                    found[line]["per_m"].add(bool(m.group(2)))
                    found[line]["src"].add(fname)
                m = re.fullmatch(r"(\d{4})\s?K", cell)
                if m and 1700 <= int(m.group(1)) <= 10000:
                    found[line]["k"].add(int(m.group(1)))
    return {code: {"lm": sorted(d["lm"]), "per_m": any(d["per_m"]), "k": sorted(d["k"]),
                   "src": sorted(d["src"]), "beam": []} for code, d in found.items()}


# ---------------------------------------------------------------- produkty bez parametrow swietlnych

LS_ACCESSORY_NOUNS = [
    (r"/SMB\b|^UCHWYT", "Uchwyt do taśm LED", "m"),
    (r"/SE\b|SIL-?SEAL|SILICONE ?SEAL|^ZASLEPKA", "Zaślepka do taśm LED", "f"),
    (r"^(ZLACZ|PASEK|ZASILACZ)", "Złączka do taśm LED", "f"),
    (r"^ZASLEPKA", "Zaślepka do profilu LED", "f"),
    (r"^UCHWYT", "Uchwyt do profilu LED", "m"),
    (r"^(AKCESORIA|ACCESSORIES)", "Akcesorium do profilu LED", "n"),
]


def model_prefix_len(words: list[str]) -> int:
    """Ile slow WIELKIMI na poczatku to model, po ktorym idzie rzeczownik ("EDGAR II spodnie" -> 2)."""
    n = 0
    while n < len(words) and words[n].isupper() and re.search(r"[A-Z]", words[n]):
        n += 1
    return n if 0 < n < len(words) and not words[n].isupper() and re.match(r"[A-Za-ząćęłńóśźżĄĆĘŁŃÓŚŹŻ]", words[n]) else 0


def resolve_noun(rule: dict[str, Any], u: str, original: str, rules: dict[str, Any]) -> tuple[str, str]:
    """Rzeczownik nazwy i jego rodzaj (m/f/n) - staly z reguly albo wyliczany z nazwy katalogowej."""
    source = rule.get("noun_from")
    noun, gender = rule.get("noun", ""), rule.get("gender", "f")
    if source in ("mount", "device", "accessory"):
        noun, gender = rule.get("default_noun", ""), "f"
        for item in rules.get(f"{source}_nouns", []):
            if re.search(item["match"], u):
                noun, gender = item["noun"], item.get("gender", "f")
                break
        if not noun and source == "accessory":
            return resolve_noun({"noun_from": "first_word"}, u, original, rules)
        if source == "mount" and re.search(rules.get("outdoor_pattern", "$^"), u) \
                and not re.search(r"ogrodow|Kamera|Taśma|Lampka|Panel solarny|Transformator|Kotwy|Tuleja", noun):
            noun = f"{noun} {rules.get('outdoor_adjective', {}).get(gender, '')}".strip()
    elif source == "ls_accessory":
        noun, gender = "Akcesorium do taśm LED", "n"
        for pattern, label, g in LS_ACCESSORY_NOUNS:
            if re.search(pattern, u):
                noun, gender = label, g
                break
    elif source == "model_first":
        words = original.split()
        n_model = model_prefix_len(words)
        if n_model:
            lead = [words[n_model]]
            for w in words[n_model + 1:]:
                if re.fullmatch(r"[a-ząćęłńóśźż/-]+", w):
                    lead.append(w)
                else:
                    break
            phrase = " ".join(lead)
            return phrase[:1].upper() + phrase[1:], "f"
        return resolve_noun({"noun_from": "first_word"}, u, original, rules)
    elif source == "first_word":
        words = original.split()
        if words and words[0].isupper() and len(words[0]) > 1:
            # stara nazwa WIELKIMI LITERAMI - rzeczownik ze slownika `caps_nouns`
            for item in rules.get("caps_nouns", []):
                if re.search(item["match"], u):
                    return item["noun"], item.get("gender", "f")
            return rule.get("default_noun", words[0].capitalize()), "f"
        lead = []
        stop = {w.lower() for w in rules.get("noun_stop_words", [])}
        keep_after = {w.lower() for w in rules.get("noun_keep_after", [])}
        for i, w in enumerate(words[:rule.get("noun_words", len(words))]):
            nxt = words[i + 1].lower() if i + 1 < len(words) else ""
            if lead and ((w.lower() in stop and nxt not in keep_after) or (rules.get("noun_stop_digits")
                                                and re.search(r"\d", w) and not re.fullmatch(r"\d+-[a-ząćęłńóśźż]+", w))):
                break  # Kobi: marka przed "do ..." ("Rozeta sufitowa Kobi Pro do KOLINE K2") i przed parametrami
            # polskie slowo: pierwsze, pisane malymi albo z ogonkiem ("Sterownik z czujnikiem ruchu i światła")
            if not lead or (w.islower() and not (rules.get("noun_no_digits") and re.search(r"\d", w))) \
                    or (re.search(r"[ąćęłńóśźż]", w) and not w.isupper()):
                lead.append(w)
            else:
                break
            if len(lead) == 6:
                break
        noun = " ".join(lead).rstrip(",") if rules.get("noun_no_digits") else " ".join(lead)
        gender = "f" if lead and lead[0].lower().endswith("a") else "m"
    return noun, gender


def noun_end(original: str, noun: str, u: str, rules: dict[str, Any]) -> int:
    """Gdzie w nazwie konczy sie rzeczownik `first_word`.

    W starej nazwie WIELKIMI rzeczownik pochodzi ze slownika `caps_nouns` i ma inna dlugosc niz tekst
    ("AKUM. 18650" -> "Akumulator"), wiec ciecie po len(noun) zjadalo poczatek modelu.
    """
    words = original.split()
    if not (words and words[0].isupper() and len(words[0]) > 1):
        return len(noun)
    for item in rules.get("caps_nouns", []):
        m = re.search(item["match"], u)
        if m and m.start() == 0 and len(u) == len(original):
            end = m.end()
            while end < len(original) and not original[end].isspace():
                end += 1  # "^OPRAWA DO NA" konczy sie w srodku slowa NABUDOWANIA
            return end
    return len(words[0])


def strip_words(original: str, rule: dict[str, Any], rules: dict[str, Any]) -> str:
    text = original
    if rule.get("strip_prefix"):
        text = re.sub(rule["strip_prefix"], "", text)  # wielkosc liter wg wzorca: (?i) w YAML
    text = re.sub(r"(?<=[A-Za-z0-9])(EUE)?LEDV(ANCE)?", " ", text)
    for pattern in rules.get("noise_patterns", []):
        text = re.sub(pattern, " ", text)
    return text


POLISH_LOWER_EXTRA = {"LEWY", "PRAWY", "LEWE", "PRAWE", "LEWA", "PRAWA", "MAT", "SZARY", "BIALY", "BIALE", "BIALA"}


def is_polish_word(tok: str) -> bool:
    """Polskie slowo opisowe (nie nazwa modelu): ogonki, typowe koncowki albo przyimek."""
    t = tok.upper().strip(",.()")
    if re.search(r"[ĄĆĘŁŃÓŚŹŻ]", t) or t in {"DO", "Z", "ZE", "W", "NA", "I", "BEZ", "POD", "DLA", "OD", "PRZY", "NAD"}:
        return True
    return bool(re.search(r"(OWY|OWA|OWE|OWYCH|NY|NA|NE|NYCH|SKI|SKA|CJA|NIE|IEM|AMI|ACH|EGO|EJ|YCH|ANY|ANA|ANE)$", t)) \
        and len(t) > 4


def tidy_tokens(text: str, limit: int = 6, split_slash: bool = True, keep_case: bool = False,
                keep_upper: frozenset[str] = frozenset(), polish_lower: bool = False, lower_re: str = "",
                vocab: frozenset[str] = frozenset(), models: frozenset[str] = frozenset(), dedupe: bool = True) -> str:
    """Bez polskich slow pisanych malymi literami, bez powtorzen, WIELKIE SLOWA -> Wielkie.

    `dedupe=False` (opis akcesorium): powtorzone slowo niesie tresc ("uchwytem FI60mm na slupach z FI60mm").
    """
    out, seen = [], set()
    if keep_case:
        # po wycieciu parametrow z kodu "AL-SL-NW-MAT-W-NT" zostaja wiszace myslniki: "AL-SL -MAT -NT"
        text = re.sub(r"-+", "-", re.sub(r"\s+-(?=[A-Za-z0-9])", "-", text))
    for tok in (text.replace("/", " / ") if split_slash else text).split():
        tok = tok.strip(",.+-" if keep_case else ",.+")
        if not tok or tok == "/" or tok == "-":
            continue
        if re.fullmatch(r"[a-ząćęłńóśźż]+", tok):
            continue
        key = ascii_upper(tok)
        if (dedupe and key in seen) or re.fullmatch(r"\d+X\d+", key):
            continue
        seen.add(key)
        if polish_lower and tok.isupper() and key not in keep_upper and key.strip("()") not in models \
                and (is_polish_word(tok) or key.strip("()") in vocab or (lower_re and re.fullmatch(lower_re, key))):
            tok = tok.lower()  # GTV: "OPRAWA ... MONTAŻOWY" -> "montażowy", model "ARALIA" -> "Aralia"
        elif tok.isalpha() and tok.isupper() and len(tok) > 2 and not keep_case and key not in keep_upper:
            tok = tok.capitalize()
        out.append(tok)
    return " ".join(out[:limit])


def parse_series(original: str, rule: dict[str, Any], rules: dict[str, Any]) -> str:
    text = strip_words(original, rule, rules)
    for pattern in [] if rule.get("series_raw") else rules.get("series_drop", []):
        text = re.sub(pattern, " ", text)
    if rules.get("series_drop_polish") and not rule.get("series_raw"):
        # Bemko: "SKALER KWADRATOWA ... CZ. MIKR." -> "SKALER"; nazwy modeli z polska koncowka chroni `proper_words`
        proper = set(rules.get("proper_words", []))
        text = " ".join(t for t in text.split() if ascii_upper(t) in proper
                        or not (is_polish_word(t) or re.fullmatch(r"[^\W\d]{1,6}\.", t)))
    if rules.get("proper_words"):
        text = " ".join(t.capitalize() if ascii_upper(t) in set(rules["proper_words"]) else t for t in text.split())
    return tidy_tokens(text, rule.get("series_limit", 5), split_slash=not rules.get("keep_case"),
                       keep_case=bool(rules.get("keep_case")),
                       keep_upper=frozenset(rules.get("keep_upper", [])), polish_lower=bool(rules.get("polish_lower")),
                       lower_re=rules.get("polish_lower_re", "") if rule.get("series_raw") else "",
                       vocab=rules.get("_vocab", frozenset()), models=rules.get("_models", frozenset()),
                       dedupe=not rule.get("series_all_segments"))


def parse_deco_size(u: str, noun: str, bare_cm: bool = False) -> str:
    # panel okragly tez ze srednica ("Planon Frameless Round 300 mm" -> "300 mm", nie "30cm"), 2026-10-21
    round_noun = noun.startswith(("Plafon", "Lampa wisząca")) or bool(re.search(r"\b(?:ROUND|CIRCULAR)\b", u))
    m = re.search(r"(?<![\d.,])(\d{2,4})\s?X\s?(\d{1,4}(?:[.,]\d+)?)(?![\d.,]*\s?W\b)\s?(MM|CM)?\b", u)
    if m:
        a, b = to_float(m.group(1)), to_float(m.group(2))
        # bez jednostki: mm, gdy pierwszy wymiar >= 100 (GTV "SAKURA 120x85" = 12x8,5cm); `deco_size_bare_cm`
        # (Ledvance) - mm tylko, gdy oba >= 100, bo "PLANON 120X30" / "BL100x25" to cm (2026-10-21: bylo "12x3cm")
        if m.group(3) == "MM" or (m.group(3) is None and a >= 100 and (b >= 100 or not bare_cm)):
            a, b = a / 10, b / 10
        return f"{fmt_num(a)}x{fmt_num(b)}cm"
    m = re.search(r"(?<![\d.,])(\d{2,4})\s?MM\b", u)
    if m:
        n = int(m.group(1))
        return f"fi{n}" if round_noun else f"{fmt_num(n / 10)}cm"
    m = re.search(r"(?<![\d.,])(\d{2,3})\s?CM\b", u)
    if m:
        return f"{int(m.group(1))}cm"
    m = re.search(r"(?<![\d.,])(\d+(?:[.,]\d+)?)\s?M\b", u)
    if m:
        return f"{fmt_num(to_float(m.group(1)))}m"
    m = re.search(r"\b(?:ROUND|RD|CIRCULAR)\s+(\d{3})\b", u)
    if m:
        return f"fi{int(m.group(1))}"
    if round_noun:
        # liczba tuz za moca to kod barwy ("Cylinder Pendant 28W 830"), nie srednica
        # ani zakres napiecia ("AC220-240V" -> nie "220 mm"), ani kat ("110°", "360ST")
        m = re.search(r"(?<![\d.,X])(?<!\dW )(?<!AC)(?<!AC )(?<!-)(\d{3})(?![\d.,X])(?!\s?(?:W|K|LM|V|ST|DEG)\b)"
                      r"(?!\s?-\s?\d)(?!\s?[°º])", u)
        if m:
            return f"fi{int(m.group(1))}"
    return ""


def parse_sockets(u: str) -> str:
    m = re.search(r"\b(\d)\s?X\s?(E27|E14|G9|GU10|GX53|R7S|G4)\b", u)
    if m:
        return f"{m.group(1)}x{parse_socket(m.group(2))}"
    m = re.search(r"\b(GU?5[.,]3)/(GU10)\b", u)
    if m:
        return f"{parse_socket(m.group(1))}/GU10"  # Kobi: "Downlight AQUA G5,3/GU10" - oprawka na oba trzonki
    socket = parse_socket(u)
    if socket in {"E27", "E14", "G9", "GU10", "GX53", "G4", "R7s"}:
        return socket
    return "MR16" if re.search(r"\bMR ?16\b", u) else ""


def parse_system(u: str, rules: dict[str, Any]) -> str:
    for item in rules.get("accessory_systems", []):
        if re.search(item["match"], u):
            return f"do {item['label']}"
    return ""


def parse_acc_size(u: str) -> str:
    m = re.search(r"\bSQUARE\s+(\d{3})\b", u)
    if m:
        return f"{m.group(1)}x{m.group(1)} mm"
    dia = re.search(r"\b(?:DN|D|RD)\s?(\d{2,3})\b", u)
    box = re.search(r"(?<![\d.,])(\d{3,4})\s?X\s?(\d{3,4})\b", u)
    if dia or box:
        # ramka wielokrotna: "SPOT FRAME D68 100X300" -> "fi68 100x300"
        return " ".join(p for p in (f"fi{int(dia.group(1))}" if dia else "",
                                    f"{box.group(1)}x{box.group(2)}" if box else "") if p)
    if re.search(r"\d\s?W\b", u):
        return ""  # "HIGH BAY REFLECTOR GEN 4 147 190 210W" - liczby to moce opraw, nie wymiar (2026-10-21)
    for tok in re.findall(r"(?<![\d.,/X-])(\d{3,4})(?![\d.,/XW])", u):
        n = int(tok)
        if n >= 600:
            return mm_to_cm(n)
        if n >= 100:
            return f"fi{n}"
    return ""


def parse_poles(u: str) -> str:
    m = re.search(r"\b(\d)P\b|\b(\d)\s?X\s?2[.,]5\b|(\d)-POLE", u)
    return f"{m.group(1) or m.group(2) or m.group(3)}P" if m else ""


def parse_housing_tubes(u: str, words: tuple[str, str] = ("świetlówkę", "świetlówki")) -> str:
    m = re.search(r"\b(\d)\s?X\s?LAMP", u)
    count = int(m.group(1)) if m else 1
    length = parse_length(u)
    k = re.search(r"\b(\d)\s?X\s?(60|120|150)\b", u)  # Kobi: "HERMIC 2X120"
    if k:
        count, length = int(k.group(1)), f"{k.group(2)}cm"
    t = re.search(r"\b(\d)\s?X\s?T8\b", u)  # "2XT8 18W"
    if t:
        count = int(t.group(1))
        p = re.search(r"(?<![\d.,])(\d+)\s?W\b", u)
        length = {9: "60cm", 18: "120cm", 22: "150cm", 24: "150cm"}.get(int(p.group(1)), "") if p else ""
    word = words[0] if count == 1 else words[1]  # Ledvance od 2026-10-19: "tubę" / "tuby" (`tube_words`)
    return f"pod {word} LED {count}x{length}" if length else f"pod {word} LED {count}x"


def parse_ls_profile(u: str) -> str:
    m = re.search(r"-(P[A-Z]\d{2})(?:/([A-Z]+))?/(\d+(?:[.,]\d+)?)X(\d+(?:[.,]\d+)?)/\d+/(\d+(?:[.,]\d+)?)?", u)
    if not m:
        return parse_ls_code(u)
    code = m.group(1) + (f"/{m.group(2)}" if m.group(2) else "")
    length = f" {fmt_num(to_float(m.group(5)))}m" if m.group(5) else ""
    return f"{m.group(3)}x{m.group(4)}mm{length} {code}"


def parse_ls_cover(u: str) -> str:
    m = re.search(r"PC/([A-Z]\d{2})/([CDR])/(\d+(?:[.,]\d+)?)", u)
    if not m:
        return parse_ls_code(u)
    kind = {"C": "przezroczysta", "D": "mleczna"}.get(m.group(2), m.group(2))
    return f"{kind} {fmt_num(to_float(m.group(3)))}m PC/{m.group(1)}"


def parse_ls_code(u: str) -> str:
    for pattern in (r"\bAY\b\s*-?\s*(.+)$", r"PROFILES\s*-?\s*(.+)$", r"^\S+(?:\s+DO\s+\S+)?\s+(.+)$"):
        m = re.search(pattern, u)
        if m:
            code = re.sub(r"\bLEDV\b|\b\d+X\d+\b", " ", m.group(1))
            return " ".join(code.split())
    return ""


def parse_module_family(u: str) -> str:
    # "PANEL-CN111" w Ekonomie to PL-CN111 = PrevaLED Coin (zamiana PL -> PANEL w Ekonomie, cofana w `name_cleanup`)
    for pattern, label in ((r"BACKLED", "BackLED"), (r"\bPL-CN", "PrevaLED Coin"), (r"\bPL-CUBE", "PrevaLED Cube"),
                           (r"PANEL-CN", "PrevaLED Coin"), (r"PREVALED", "PrevaLED"),
                           (r"LISOEV", "IndiviLED")):
        if re.search(pattern, u):
            return label
    return ""


def parse_model_code(original: str, rule: dict[str, Any], rules: dict[str, Any]) -> str:
    text = strip_words(original, rule, rules)
    text = re.sub(r"(?i)\b(sensors?|czujnik|coupler|wt|bk|gy|dg|si|white|black|\d+x\d+)\b", " ", text)
    return tidy_tokens(text, 6, split_slash=False)


def parse_hid_model(u: str) -> str:
    if re.search(r"\bPLANTASTAR\b", u):
        return "Plantastar"
    m = re.search(r"\b(HQI|NAV|HCI|HMI)[- ]?(TS|BT|T|E|R)?(?=[\s\d/]|$)", u)
    if not m:
        return ""
    return f"{m.group(1)}-{m.group(2)}" if m.group(2) else m.group(1)


def parse_hid_power(u: str) -> str:
    text, _ = parse_power(u)
    if text:
        return text
    m = re.search(r"\b(?:HQI|NAV|HCI|HMI|PLANTASTAR)[- ]?(?:TS|BT|T|E|R)?\s?(\d{2,4})(?=[/\s]|$)", u)
    return f"{m.group(1)}W" if m else ""


def parse_device_extra(u: str, gender: str, rules: dict[str, Any]) -> str:
    """Cechy urzadzen Smart+: zewnetrzne, pomiar energii, wtyczka UK, solarne."""
    parts = []
    if re.search(r"OUTDOOR|ZEWNETRZN", u):
        parts.append(rules.get("outdoor_adjective", {}).get(gender, ""))
    if re.search(r"\b3 WAY\b", u):
        parts.append("potrójne")
    if re.search(r"LICZNIK|ENERGY|ENERGII", u):
        parts.append("z pomiarem energii")
    if re.search(r"SOLAR|SLONECZN", u):
        parts.append("solarny")
    if re.search(r"\bUK\b|WIELKA BRYTANIA", u):
        parts.append("UK")
    return " ".join(p for p in parts if p)


def parse_hid_color(u: str) -> str:
    # HQI: N / NDL = neutralna 4000K, WDL = 3000K, D / DL = dzienna (2026-10-21: "HQI-T 1000W N" bylo bez barwy,
    # "N/L", "D HF" zostawaly jako surowe litery)
    if re.search(r"\bNDL\b|\dW?\s?N(?=/|\s|$)", u):
        return "4000K"
    if re.search(r"\bWDL\b", u):
        return "3000K"
    if re.search(r"\bDL\b|\dW?\s?D(?=/|\s|$)", u):
        return "barwa dzienna"
    return ""


# ---------------------------------------------------------------- budowa nazwy

def match_type(u: str, types: list[dict[str, Any]], has_power: bool, code: str = "") -> dict[str, Any] | None:
    for t in types:
        if t.get("code_match") and not re.search(t["code_match"], code):
            continue
        if not any(re.search(p, u) for p in t.get("match", ["."])):
            continue
        if any(re.search(p, u) for p in t.get("exclude", [])):
            continue
        if t.get("requires_no_power") and has_power:
            continue
        if t.get("requires_power") and not has_power:
            continue
        if t.get("require_pattern") and not re.search(t["require_pattern"], u):
            continue
        return t
    return None


def clean_model(name: str, rules: dict[str, Any], strip_prefix: str = "") -> str:
    text = name
    if strip_prefix:
        text = re.sub(strip_prefix, "", text, flags=re.IGNORECASE)
    for pattern in rules.get("noise_patterns", []):
        text = re.sub(pattern, " ", text)
    return " ".join(text.split())


def parse_view(name: str) -> str:
    """Nazwa do parsowania: wielkie litery bez ogonkow, przyklejone LEDV/EUE odklejone."""
    u = ascii_upper(name)
    u = re.sub(r"(?<=[A-Z0-9])(EUE)?LEDV(ANCE)?", " LEDV", u)
    u = re.sub(r"(?<=[A-Z0-9])EUE\b", " EUE", u)
    # kody OSRAM bez spacji: "4,9W/827230VDIMFRE144X3" -> "4,9W/827 230V DIMFRE144X3"
    u = re.sub(r"(?<=\d{3})(2[34]0V)", r" \1", u)
    u = re.sub(r"\b(2[34]0V)(?=[A-Z])", r"\1 ", u)
    return u


# ---------------------------------------------------------------- feed XML dostawcy (Ceneo)

def load_feed(path: Path) -> dict[str, dict[str, str]]:
    """Feed Ceneo: kod producenta (bez zer wiodacych) -> atrybuty <a name=...> + kategoria i nazwa."""
    import xml.etree.ElementTree as ET

    feed: dict[str, dict[str, str]] = {}
    for offer in ET.parse(path).getroot().iter("o"):
        attrs = {a.get("name", ""): (a.text or "").strip() for a in offer.iter("a")}
        attrs["_cat"] = offer.findtext("cat") or ""
        attrs["_name"] = offer.findtext("name") or ""
        code = attrs.get("Kod producenta", "").lstrip("0")
        if code:
            feed[code] = attrs
    return feed


def feed_key(row: dict[str, str]) -> str:
    return (row.get("Kod towaru") or "").split("/")[0].strip().lstrip("0")


def resolve_brand(u: str, feed_rec: dict[str, str], rules: dict[str, Any]) -> str:
    """Marka/linia: z feedu (np. Kobi Premium, LED2B), potem reguly po nazwie, potem domyslna."""
    mapped = rules.get("brand_map", {})
    value = feed_rec.get(rules.get("feed_attrs", {}).get("brand", ""), "")
    if value:
        return mapped.get(value, value)
    for item in rules.get("brand_rules", []):
        if re.search(item["match"], u):
            return item["brand"]
    return rules.get("brand", "")


def feed_mismatch(name: str, feed_rec: dict[str, str]) -> str:
    """Opis rozbieznosci, gdy moc albo barwa z nazwy Ekonoma nie zgadza sie z feedem."""
    if not feed_rec:
        return ""
    u = ascii_upper(name)
    f = ascii_upper(feed_rec.get("_name", ""))
    for pattern in (r"(?<![\d.,])(\d+(?:[.,]\d+)?)\s?W\b", r"(?<!\d)(\d{4})\s?K\b"):
        mine = set(re.findall(pattern, u))
        theirs = set(re.findall(pattern, f))
        if mine and theirs and not mine & theirs:
            return f"feed ma inny produkt pod tym kodem: {feed_rec.get('_name', '')}"
    return ""


def feed_number(value: str) -> float | None:
    value = value.replace("\xa0", " ").strip()
    if re.fullmatch(r"\d{1,3}(?: \d{3})+|\d+(?:[.,]\d+)?", value):
        return to_float(value.replace(" ", ""))
    return None


def apply_feed(values: dict[str, str], rule: dict[str, Any], feed_rec: dict[str, str], rules: dict[str, Any],
               gender: str, result: dict[str, Any]) -> None:
    """Uzupelnia puste atrybuty z feedu. Nazwa ma pierwszenstwo, poza czujnikiem (feed jest dokladniejszy)."""
    if not feed_rec:
        return
    names = {k: v for k, v in rules.get("feed_attrs", {}).items() if k not in rule.get("no_feed", [])}
    attrs = rule.get("attrs", [])

    def get(key: str) -> str:
        return feed_rec.get(names.get(key, ""), "").strip()

    if "power" in attrs and not values.get("power"):
        n = feed_number(get("power"))
        if n:
            values["power"] = f"{fmt_num(n)}W"
    if "lumen" in attrs and not values.get("lumen"):
        n = feed_number(get("lumen"))
        if n and n >= 10:
            values["lumen"] = f"{fmt_num(n)}lm"
            result["lm_zrodlo"] = "feed XML"
    if "cct" in attrs and not values.get("cct"):
        v = get("cct")
        if re.fullmatch(r"\d{4}", v):
            values["cct"] = f"{v}K"
        elif v.upper() in ("CCT", "RGB", "RGBW"):
            values["cct"] = v.upper()
    if "gwint" in attrs and not values.get("gwint"):
        v = get("gwint")
        values["gwint"] = parse_socket(ascii_upper(v)) or v
    if "ip" in attrs and not values.get("ip"):
        m = re.search(r"(\d{2})", get("ip"))
        if m and int(m.group(1)) >= 44:  # IP20/IP40 z feedu nic sprzedawcy nie mowi
            values["ip"] = f"IP{m.group(1)}"
    if "color" in attrs and not values.get("color"):
        values["color"] = parse_color(ascii_upper(get("color")), gender, rules.get("housing_colors", {}))
    if "sensor" in attrs:
        v = get("sensor").lower()
        if "ruchu" in v and "zmierzch" in v:
            values["sensor"] = "z czujnikiem ruchu i zmierzchu"
        elif "ruchu" in v:
            values["sensor"] = "z czujnikiem ruchu"
        elif "zmierzch" in v:
            values["sensor"] = "z czujnikiem zmierzchu"
    if rule.get("feed_optics"):
        for key in ("optics", "beam"):
            if key in attrs and not values.get(key):
                v = get("beam")
                if re.fullmatch(r"\d{1,3}(?:[-x]\d{1,3})?", v):
                    values[key] = f"{v}°"
    if "length" in attrs and not values.get("length"):
        n = feed_number(get("length"))
        if n and n >= 30:
            values["length"] = f"{fmt_num(n)}cm"


# ---------------------------------------------------------------- parametry producenta (Kanlux)

def load_params(path: Path) -> dict[str, dict[str, str]]:
    """Pliki Parametry*.xlsx w formacie dlugim Kanlux: Kod Kanlux; Nazwa; ID atrybutu; Nazwa atrybutu; Wartosc.

    Zwraca kod (bez zer wiodacych) -> {nazwa atrybutu: wartosc, "_name": nazwa katalogowa}.
    """
    import openpyxl

    params: dict[str, dict[str, str]] = defaultdict(dict)
    files = sorted(path.glob("Parametry*.xlsx")) if path.is_dir() else [path]
    for f in files:
        wb = openpyxl.load_workbook(f, read_only=True)
        for row in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 5 or row[0] is None or row[3] is None:
                continue
            rec = params[str(row[0]).strip().lstrip("0")]
            rec.setdefault("_name", " ".join(str(row[1] or "").split()))
            key = " ".join(str(row[3]).split())  # "Całkowity znamionowy strumień świetlny [lm]\n"
            value = "" if row[4] is None else str(row[4]).strip()
            old = rec.get(key)
            # ten sam atrybut dwa razy (Φuse: "806" i "w kuli (360°)") - wygrywa wartosc z liczba
            if old is None or (not re.search(r"\d", old) and re.search(r"\d", value)):
                rec[key] = value
        wb.close()
    return dict(params)


def load_params_web(cache_dir: Path) -> tuple[dict[str, dict[str, str]], set[str]]:
    """Cache kart kanlux.com: kod -> parametry jak w Parametry.xlsx; drugi wynik - kody spoza oferty."""
    params: dict[str, dict[str, str]] = {}
    withdrawn: set[str] = set()
    for f in sorted(cache_dir.glob("*.json")) if cache_dir.is_dir() else []:
        rec = json.loads(f.read_text(encoding="utf-8"))
        if rec.get("found"):
            params[rec["code"]] = {**rec.get("params", {}), "_name": rec.get("name", ""), "_src": "kanlux.com"}
        else:
            withdrawn.add(rec["code"])
    return params, withdrawn


def load_web_params(cache_dir: Path) -> tuple[dict[str, dict[str, str]], set[str], frozenset[str], frozenset[str]]:
    """Cache kart producenta (GTV: scripts/fetch_gtv_web.py): kod -> parametry, kody spoza strony, slownik.

    Slownik to slowa pisane na stronie producenta malymi literami ("przewodem", "paneli") - w starej nazwie
    WIELKIMI takie slowo jest polskim opisem, a nie nazwa modelu. Czwarty wynik - nazwy modeli (slowa pisane
    na stronie WIELKIMI w nazwie mieszanej: NOELIA, VILANA), ktorych nie wolno zmienic na male litery.
    """
    params: dict[str, dict[str, str]] = {}
    withdrawn: set[str] = set()
    lower_n: Counter[str] = Counter()
    upper_n: Counter[str] = Counter()
    for f in sorted(cache_dir.glob("*.json")) if cache_dir.is_dir() else []:
        rec = json.loads(f.read_text(encoding="utf-8"))
        if not rec.get("found"):
            withdrawn.add(rec["code"])
            continue
        params[rec["code"]] = {**rec.get("params", {}), "_name": rec.get("name", ""), "_url": rec.get("url", "")}
        name = rec.get("name", "")
        if name.isupper():
            continue  # Bemko: nazwy kart WIELKIMI LITERAMI - nic nie mowia o pisowni slow
        for tok in re.findall(r"[A-Za-zĄĆĘŁŃÓŚŹŻąćęłńóśźż]{2,}", name):
            if tok.islower():
                lower_n[ascii_upper(tok)] += 1
            elif tok.isupper() and len(tok) >= 3 and re.search(r"[a-ząćęłńóśźż]{3}", name):
                upper_n[ascii_upper(tok)] += 1
    # zwykle slowo czasem bywa WIELKIMI ("TYPU HIGH BAY"), model czasem malymi - decyduje czestszy zapis
    vocab = {w for w, n in lower_n.items() if n > upper_n[w]}
    models = {w for w, n in upper_n.items() if n > lower_n[w]}
    return params, withdrawn, frozenset(vocab), frozenset(models)


def apply_web_params(values: dict[str, str], rule: dict[str, Any], rec: dict[str, str], rules: dict[str, Any],
                     result: dict[str, Any]) -> None:
    """Karta producenta uzupelnia tylko braki - parametry z nazwy w Ekonomie maja pierwszenstwo."""
    keys = rules.get("web_params_keys", {})
    attrs = rule.get("attrs", [])
    lm = feed_number(rec.get(keys.get("lumen", ""), ""))
    if "lumen" in attrs and not values.get("lumen") and lm and lm >= 1:
        values["lumen"] = f"{fmt_num(lm)}lm"
        result["lm_zrodlo"] = rules.get("web_params_label", "strona producenta")
    cct = kx_cct(rec.get(keys.get("cct", ""), ""))
    if "cct" in attrs and cct and (not values.get("cct") or not re.search(r"\d", values["cct"])):
        values["cct"] = cct  # K ze strony zamiast braku albo slow "ciepła biała"
    m = re.search(r"(\d{2})", rec.get(keys.get("ip", ""), ""))
    if "ip" in attrs and not values.get("ip") and m and int(m.group(1)) >= 44:
        values["ip"] = f"IP{m.group(1)}"


POLISH_WORD = r"[A-ZĄĆĘŁŃÓŚŹŻ][a-ząćęłńóśźż]{2,}"


def split_head_tail(name: str) -> tuple[str, str]:
    """Nazwa Kanlux = kod modelu + polski typ: ("XLED GU10 4,9W-WW", "Źródło światła LED")."""
    # typ przyklejony do kodu po obcieciu na 40 znakach: "SMDE27-WWŹródło", "czw.Gniazdo"
    text = re.sub(r"(?<=[A-Za-z0-9.)])(?=[A-ZĄĆĘŁŃÓŚŹŻ][a-ząćęłńóśźż]{3,})", " ", name)
    m = re.search(rf"(?:^|\s)({POLISH_WORD}\b.*)$", text)
    if not m:
        return text.strip(), ""
    head, tail = text[:m.start(1)].strip(), m.group(1).strip()
    if not head:
        # nazwa przerobiona na opis, a na koncu nazwa katalogowa: "Rozdzielnica modułowa ... kolor szary
        # KDB-IP65-S09T MK2 Rozdzielnica serii" -> kod i typ z konca
        words = text.split()
        for i in range(len(words) - 1, 0, -1):
            if not re.fullmatch(POLISH_WORD + r"\W?", words[i]):
                continue
            code = []
            for w in reversed(words[:i]):
                if re.search(r"[a-ząćęłńóśźż]{2,}", w) or not re.search(r"[A-Z0-9]", w):
                    break
                code.insert(0, w)
            if code:
                return " ".join(dict.fromkeys(code)), " ".join(words[i:])
    if not head:
        # nazwa zaczyna sie od typu ("Oprawa liniowa LED AL-..."): typ = slowa do pierwszego kodu
        words = tail.split()
        lead = [words[0]]
        for w in words[1:]:
            if w.islower() or (w == "LED" and len(lead) < 4) or re.fullmatch(r"[a-ząćęłńóśźż-]+", w):
                lead.append(w)
            else:
                break
        head, tail = " ".join(words[len(lead):]), " ".join(lead)
    return head, tail


def tail_noun(tail: str, rules: dict[str, Any]) -> tuple[str, str]:
    """Rzeczownik z polskiego typu Kanlux, z podmianami z `tail_nouns` i bez szumu z `tail_drop`."""
    tu = ascii_upper(tail)
    for item in rules.get("tail_nouns", []):
        if re.search(item["match"], tu):
            return item["noun"], item.get("gender", "f")
    noun = tail
    for pattern in rules.get("tail_drop", []):
        noun = re.sub(pattern, "", noun)
    noun = " ".join(noun.replace(" ,", ",").split()).strip(" ,")
    first = noun.split()[0].lower() if noun else ""
    gender = "f" if first.endswith("a") else ("n" if first.endswith(("o", "e")) else "m")
    return noun, gender


def kx_numbers(value: str) -> list[float]:
    return [to_float(x) for x in re.findall(r"\d+(?:[.,]\d+)?", value.replace(" ", ""))]


def kx_power(rec: dict[str, str]) -> str:
    """Moc z parametrow; przy wymiennym zrodle swiatla "Moc maksymalna" to limit zarowki, nie moc."""
    value = rec.get("Moc znamionowa [W]", "")
    socket = rec.get("Trzonek", "")
    replaceable = (rec.get("Wymienne źródło światła") == "tak" or rec.get("Źródło światła w komplecie") == "nie"
                   or (socket and not re.search(r"(?i)przewod|wolne", socket)
                       and rec.get("Zintegrowane źródło światła LED") != "1"))
    if not value and not replaceable:
        value = rec.get("Moc maksymalna [W]", "")
    if re.fullmatch(r"\d+(?:[.,]\d+)?\s?-\s?\d+(?:[.,]\d+)?", value) and rec.get("Moc maksymalna [W]"):
        value = rec["Moc maksymalna [W]"]  # zasilacz TRETO: "16-40" to zakres obciazenia, moc to 40W
    if not value or (value.startswith("max") and replaceable):
        return ""  # "max 7,3" przy zintegrowanym LED (PREDA) to moc lampki, przy trzonku - limit zarowki
    per_m = bool(re.match(r"1\s?m\s?=", value))
    nums = kx_numbers(re.sub(r"^1\s?m\s?=", "", value))
    nums = [n for n in nums if n > 0]
    if not nums:
        return ""
    text = "/".join(fmt_num(n) for n in dict.fromkeys(nums)) + "W"
    return text + "/m" if per_m else text


def kx_lumen(rec: dict[str, str], keys: list[str]) -> str:
    for key in keys:
        value = rec.get(key, "")
        if not re.search(r"\d", value):
            continue
        per_m = bool(re.match(r"1\s?m\s?=", value))
        value = re.sub(r"\([^)]*\)", " ", value)  # "2280 (600mA) / 3040 (850mA)" - prad to nie lm
        nums = sorted({n for n in kx_numbers(re.sub(r"^1\s?m\s?=", "", value)) if n >= 1})
        if not nums:
            continue
        # "2090(WW) /2310(WN)", "37000/29600/22200" (moc przelaczana) -> zakres
        text = fmt_num(nums[0]) if len(nums) == 1 else f"{fmt_num(nums[0])}-{fmt_num(nums[-1])}"
        return f"{text}lm/m" if per_m else f"{text}lm"
    return ""


REPLACEABLE_SOCKETS = {"G13", "G5", "GU10", "E27", "E14", "G9", "GX53", "R7s", "GU5.3", "GX5.3", "G4", "2G11", "G23",
                       "G24d-1", "G24d-2", "G24d-3", "2G7", "B22d", "E40", "GY6.35"}


def kx_replaceable_sockets(rec: dict[str, str]) -> str:
    """Oprawa na wymienne zrodlo: "4xG13" (NOTUS 418, moc "4 x 18"), "3xGU10", "E27"; pusty, gdy LED zintegrowany."""
    if rec.get("Zintegrowane źródło światła LED") == "1" or rec.get("Źródło światła w komplecie") == "tak":
        return ""
    sockets = [parse_socket(ascii_upper(x)) or x for x in rec.get("Trzonek", "").replace(" ", "").split("/") if x]
    if not sockets or not all(x in REPLACEABLE_SOCKETS for x in sockets):
        return ""
    m = re.match(r"\s*(\d+)\s*x", rec.get("Moc maksymalna [W]", ""))
    count = int(m.group(1)) if m else 1
    text = "/".join(dict.fromkeys(sockets))
    return f"{count}x{text}" if count > 1 else text


def kx_cct(value: str) -> str:
    value = value.replace(" ", "")
    kel = [int(x) for x in re.findall(r"\d{4}", value) if 1700 <= int(x) <= 10000]
    parts = []
    if kel:
        if "-" in value and len(kel) == 2:
            parts.append(f"{kel[0]}-{kel[1]}K")
        elif len(kel) > 3:
            parts.append(f"CCT {min(kel)}-{max(kel)}K")
        else:
            parts.append("/".join(str(k) for k in dict.fromkeys(kel)) + "K")
    if "RGBW" in value.upper():
        parts.append("RGBW")
    elif "RGB" in value.upper():
        parts.append("RGB")
    return " ".join(parts)


def kx_dims(rec: dict[str, str]) -> tuple[float, float]:
    """(dluzszy, krotszy) bok z Dlugosc/Szerokosc [mm] - tuby T8 maja dlugosc w polu Szerokosc."""
    dims = []
    for key in ("Długość [mm]", "Szerokość [mm]"):
        nums = kx_numbers(rec.get(key, ""))
        if nums:
            dims.append(nums[0])
    dims.sort(reverse=True)
    return (dims[0] if dims else 0, dims[1] if len(dims) > 1 else 0)


def cm_round(mm: float) -> str:
    return str(int(mm / 10 + 0.5))


def parse_breaker(head_u: str) -> str:
    """Aparatura Kanlux z kodu: KMB10-B16/1, KMB6-P-16/3, KRD6-2/25/100-A, KCP-09-230, KFL-D02-25, KMCCB-RT-125A-4P."""
    m = re.search(r"\bKMB(\d+)-([BCD])(\d+)/(\d)", head_u)
    if m:
        return f"{m.group(2)}{m.group(3)} {m.group(4)}P {m.group(1)}kA"
    m = re.search(r"\bKMB\d+-P-(\d+)/(\d)", head_u)
    if m:
        return f"{m.group(1)}A {m.group(2)}P"
    m = re.search(r"\bKRD(\d+)-(\d)/(\d+)/(\d+)(?:-(AC|A|B|F))?", head_u)
    if m:
        return f"{m.group(2)}P {m.group(3)}A {m.group(4)}mA typ {m.group(5) or 'AC'} {m.group(1)}kA"
    m = re.search(r"\bKCPM?-(\d+)-(\d+)\b", head_u)
    if m:
        return f"{int(m.group(1))}A cewka {m.group(2)}V"
    m = re.search(r"\bKFL-(D0?2)-(\d+)", head_u)
    if m:
        return f"D02 {m.group(2)}A"
    m = re.search(r"\bKMP-(\d+(?:,\d+)?)\b", head_u)
    if m:
        return f"do {m.group(1)}A"
    m = re.search(r"\bKMCCB-RT-(\d+)A(?:-(\d)P)?", head_u)
    if m:
        return f"{m.group(1)}A {m.group(2) or 3}P"
    return ""


def apply_params(values: dict[str, str], rule: dict[str, Any], rec: dict[str, str], rules: dict[str, Any],
                 gender: str, result: dict[str, Any]) -> None:
    """Parametry producenta (Kanlux) - zrodlo prawdy: nadpisuja wartosci odczytane z nazwy."""
    if not rec:
        return
    attrs = rule.get("attrs", [])

    def put(key: str, value: str) -> None:
        if key in attrs and value:
            values[key] = value

    put("power", kx_power(rec))
    lumen = kx_lumen(rec, rules.get("params_lumen_keys", []))
    if "lumen" in attrs and lumen:
        values["lumen"] = lumen
        result["lm_zrodlo"] = rec.get("_src", "Parametry.xlsx")
    cct = kx_cct(rec.get("Temperatura barwowa [K]", ""))
    if "cct" in attrs and cct:
        values["cct"] = cct
        result["uwagi"] = [u for u in result["uwagi"] if not u.startswith("barwa z kodu")]
    light = rec.get("Barwa światła", "")
    if "cct" in attrs and not values.get("cct") and not values.get("light_color") and re.fullmatch(r"(?i)czerwona|zielona|niebieska|żółta|pomarańczowa|różowa|fioletowa|bursztynowa", light):
        values["cct"] = light.lower()  # tasma kolorowa (LCOB ...-R) nie ma K - barwa swiatla to kolor
    sockets = kx_replaceable_sockets(rec)
    if sockets and not values.get("lumen"):
        result["wymienne_zrodlo"] = sockets  # moc/lm/K zalezy od zarowki - nie sa wymagane
        if not rec.get("Moc znamionowa [W]") and values.get("power"):
            values["power"] = ""  # "LARATA EL-10 W", "1xGU10 W" - liczba z kodu modelu, a nie moc oprawy
    trzonek = rec.get("Trzonek", "")
    if trzonek and not re.search(r"(?i)przewod|wolne", trzonek):
        put("gwint", "/".join(parse_socket(ascii_upper(s)) or s for s in trzonek.replace(" ", "").split("/")))
        if sockets and values.get("gwint") == sockets.split("x")[-1]:
            values["gwint"] = sockets  # "EVALO FL 3xGU10" -> 3xGU10, nie samo GU10
    m = re.match(r"(\d{2})", rec.get("Stopień IP", ""))
    if m and int(m.group(1)) >= 44:
        put("ip", f"IP{m.group(1)}")
    color = rec.get("Kolor", "")
    if color and not re.search(r"(?i)transparent", color):
        put("color", parse_color(ascii_upper(color), gender, rules.get("housing_colors", {})) or color.lower())
    longer, shorter = kx_dims(rec)
    if longer >= 300:
        put("length", mm_to_cm(longer))
        nominal = [n for n in (45, 60, 90, 120, 150) if abs(longer / 10 - n) <= 4]
        put("tube_length", f"{nominal[0]}cm" if nominal else mm_to_cm(longer))
    if longer >= 1000:
        put("strip_length", f"{fmt_num(longer / 1000)}m")
    if longer >= 150 and shorter >= 150:
        put("panel_size", f"{cm_round(longer)}x{cm_round(shorter)}")
    dia = kx_numbers(rec.get("Średnica [mm]", ""))
    if dia and dia[0] >= 40:
        put("diameter", f"fi{fmt_num(dia[0])}")
    ugr = rec.get("UGR", "").replace(" ", "")
    if re.fullmatch(r"[<≤]\d{2}", ugr):
        put("ugr", f"UGR{ugr}")
    dimmer = rec.get("Możliwość współpracy ze ściemniaczem", "")
    if dimmer in ("DALI", "1-10V", "0-10V"):
        put("control", dimmer)
    if dimmer == "tak" or rec.get("Lampa jest przeznaczona do ściemniania") == "1":
        put("dim", "ściemnialna")
    sensor = rec.get("Rodzaj czujnika", "")
    if sensor:
        put("sensor", {"PIR": "z czujnikiem PIR", "mikrofalowy": "z czujnikiem mikrofalowym",
                       "zmierzchu": "z czujnikiem zmierzchu"}.get(sensor, "z czujnikiem"))
    elif rec.get("Wykrywanie ruchu") == "tak":
        put("sensor", "z czujnikiem ruchu")
    shape_src = rec.get("Źródło światła", "")
    if "shape" in attrs and shape_src and not values.get("shape"):
        put("shape", parse_shape_code(ascii_upper(shape_src), rules) or parse_shape(ascii_upper(shape_src)))
    if "beam" in attrs and REFLECTOR_SHAPE_RE.match(values.get("shape", "")):
        beam = kx_numbers(rec.get("Kąt promieniowania [°]", "") or rec.get("Znamionowy kąt promieniowania [°]", ""))
        if beam:
            values["beam"] = f"{fmt_num(beam[0])}°"
    volt = re.match(r"(12|24|48)\s?(?:DC|AC)?$", rec.get("Napięcie znamionowe [V]", ""))
    if volt:
        put("voltage", f"{volt.group(1)}V")
    em = kx_numbers(rec.get("Czas pracy awaryjnej [h]", ""))
    if em:
        put("em_time", "/".join(fmt_num(x) for x in em) + "h")


def parse_features(u: str, gender: str, rules: dict[str, Any], items: list[dict[str, Any]] | None = None,
                   code: str = "") -> str:
    """Cechy z nazwy (`match`) albo z kodu producenta (`code_match`); z jednej `group` wygrywa pierwsza trafiona."""
    out, groups = [], set()
    for item in rules.get("features", []) if items is None else items:
        if item.get("group") and item["group"] in groups:
            continue
        m = re.search(item["code_match"], code) if "code_match" in item else re.search(item["match"], u)
        if m:
            label = item["label"]
            label = label.get(gender, label.get("m", "")) if isinstance(label, dict) else label
            if "{" in label:
                label = label.format(*(g or "" for g in m.groups()))
                # wartosci z nazwy WIELKIMI: "10X38" -> "10x38", "1.5" -> "1,5"
                label = re.sub(r"(?<=\d)X(?=\d)", "x", re.sub(r"(?<=\d)\.(?=\d)", ",", label))
            out.append(label)
            if item.get("group"):
                groups.add(item["group"])
    return " ".join(dict.fromkeys(out))


def parse_shape_code(u: str, rules: dict[str, Any]) -> str:
    for item in rules.get("shape_codes", []):
        if re.search(item["match"], u):
            return item["label"]
    return ""


def parse_cri(u: str) -> str:
    m = re.search(r"\bCRI?\s?(\d{2})\b", u)
    return f"CRI{m.group(1)}" if m else ""


def parse_psu_kind(u: str) -> str:
    if re.search(r"DESKTOP", u):
        return "desktop"
    if re.search(r"PUSZKOW", u):
        return "puszkowy"  # Kobi: inny produkt niz "wodoodporny" o tych samych V/W/IP
    if re.search(r"INSTALACYJN|WODOODPORN|IP6[5-8]", u):
        return "hermetyczny"
    if re.search(r"MONTAZOW", u):
        return "montażowy"
    return ""


def parse_socket_spec(u: str) -> str:
    """Przedluzacze/rozgalezniki Kobi: 3gn/5m/Zu+W, 4gn+USB/1,5m/Zu+W/1,5mm, 2gn+2USB/1USB-C."""
    m = re.search(r"\b(\d)GN((?:\+\d?USB(?:-C)?)*(?:/\d?USB-C)?)(?:/(\d+(?:,\d+)?)M)?(?:/(ZU(?:\+W)?|BU))?"
                  r"(?:/(\d(?:,\d+)?(?:X\d(?:,\d+)?)?)MM)?", u)
    if not m:
        return ""
    parts = [f"{m.group(1)}gn{m.group(2)}"]
    if m.group(3):
        parts.append(f"{m.group(3)}m")
    if m.group(4):
        # Zu = z uziemieniem, Bu = bez uziemienia ("LINEA G 1gn/20m/Bu/2x1mm")
        parts.append(m.group(4).replace("ZU", "Zu").replace("BU", "Bu"))
    if m.group(5):
        parts.append(f"{m.group(5).replace('X', 'x')}mm2")
    return " ".join(parts)


def parse_code(u: str, rule: dict[str, Any]) -> str:
    pattern = rule.get("code_regex")
    if not pattern:
        return ""
    m = re.search(pattern, u)
    return (m.group(1) if m.groups() else m.group(0)) if m else ""


def code_mismatches(code: str, raw_u: str, rules: dict[str, Any]) -> list[str]:
    """Kod producenta przeczy nazwie z Ekonoma (Bemko: kod -6K, a w nazwie 4000K) - nazwy nie poprawiamy, zglaszamy."""
    out = []
    for item in rules.get("code_checks", []):
        m = re.search(item["code"], code)
        if m and not re.search(item["name"].format(*m.groups()), raw_u):
            out.append(f"kod {code} wskazuje {item['label'].format(*m.groups())}, a nazwa w Ekonomie mowi inaczej")
    return out


def propagate_family_brand(results: list[dict[str, Any]], rules: dict[str, Any]) -> None:
    """Marka z rodzenstwa: w rodzinie kodu (C66-LFC) tylko czesc nazw podaje marke (Luxmena) - reszta ja dziedziczy.

    Dziala tylko, gdy cala rodzina wymienia jedna marke; nazwa z jawna marka nigdy nie jest zmieniana.
    """
    pattern = rules.get("brand_family")
    if not pattern:
        return
    families: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for res in results:
        m = re.search(pattern, (res["row"].get("Kod producenta") or "").strip())
        if m:
            families[m.group(1)].append(res)
    for key, members in families.items():
        brands = {r["marka"] for r in members if r.get("marka_jawna")}
        if len(brands) != 1:
            continue
        brand = brands.pop()
        for res in members:
            if res.get("marka_jawna") or res.get("marka") == brand or not res.get("marka"):
                continue
            res["nazwa"] = re.sub(rf"(?<!\S){re.escape(res['marka'])}(?!\S)", brand, res["nazwa"], count=1)
            res["uwagi"].append(f"marka {brand} z rodzenstwa {key} (nazwa w Ekonomie bez marki)")
            res["marka"] = brand


def build_name(row: dict[str, str], rules: dict[str, Any], pdf_lm: dict[str, dict[str, Any]],
               feed: dict[str, dict[str, str]] | None = None) -> dict[str, Any]:
    feed_rec = (feed or {}).get(feed_key(row), {})
    original = full_name(row)
    raw_u = parse_view(original)  # przed czyszczeniem - Bemko: marka na koncu nazwy jest wycinana w `name_cleanup`
    for item in rules.get("name_cleanup", []):
        # dopiski marketingowe w starych nazwach ("... 3 lat Gwar. LCOBB 10W/M ...") - zostaje nazwa katalogowa;
        # pozycja {match, repl} normalizuje zapis ("1*E27" -> "1xE27", "360ST" -> "360°")
        pattern, repl = (item["match"], item["repl"]) if isinstance(item, dict) else (item, " ")
        original = " ".join(re.sub(pattern, repl, original).split())
    feed_note = feed_mismatch(original, feed_rec)
    if feed_note:
        feed_rec = {}  # kod w Ekonomie wskazuje inny produkt niz w feedzie - feed pomijamy
    polish_start = bool(re.match(POLISH_WORD + r"\b", original))
    if feed_rec.get("_name") and ((rules.get("prefer_feed_name") == "polish_start" and polish_start)
                                  or (rules.get("prefer_feed_name") is True and original == original.upper())):
        # stara nazwa WIELKIMI LITERAMI -> nowa nazwa dostawcy z feedu ("Naświetlacz LED DELYO MHN 10W ...")
        original = " ".join(feed_rec["_name"].split())
    head, tail = split_head_tail(original) if rules.get("split_tail") else (original, "")
    for item in rules.get("socket_rules", []):
        if re.search(item["match"], parse_view(head)):
            head = f"{head} {item['socket']}"  # gwint, ktorego nazwa nie podaje ("J-118MM" = R7s)
            original = f"{head} {tail}".strip()
    head_u = parse_view(head)
    u = parse_view(original)
    ean = (row.get("Kod producenta") or "").strip()
    # Kanlux: gwint przed kodem koloru ("GU10 W") nie jest moca 10W
    # "3xGU10 W-SR": gwint sklejony z krotnoscia (SOCKET_RE wymaga litery przed) - inaczej "10 W" to moc
    power_text, power = parse_power(SOCKET_RE.sub(" ", re.sub(r"(?<=\dX)(?=[A-Z])", " ", head_u))
                                    if rules.get("split_tail") else u)
    has_power = bool(power_text) or bool(rules.get("params") and kx_power(feed_rec))
    code = (row.get("Kod towaru") or "").strip()
    rule = match_type(u, rules["types"], has_power, code)
    brand = resolve_brand(raw_u if rules.get("brand_from_raw") else u, feed_rec, rules)
    for item in rules.get("brand_code_rules", []):
        if re.search(item["match"], code):
            brand = item["brand"]  # GTV: HT... = narzedzia Högert, GT-... = linia G-TECH
            break
    result: dict[str, Any] = {"typ": "", "lm_zrodlo": "", "braki": [], "uwagi": [feed_note] if feed_note else [],
                              "marka": brand,
                              "marka_jawna": bool(re.search(rules["brand_explicit"], raw_u))
                              if rules.get("brand_explicit") else True,
                              # wyroznik wersji z nazwy po oczyszczeniu i rozklejeniu ("NWPanel" -> "NW Panel")
                              "zrodlo_nazwy": f"{head} {tail}".strip() if rules.get("split_tail")
                              else (original if rules.get("version_from_clean") else "")}
    if rules.get("params") and not feed_rec:
        result["uwagi"].append("brak w ofercie kanlux.com (wycofany) - dane tylko z nazwy"
                               if feed_key(row) in rules.get("_withdrawn", ()) else
                               "brak w plikach Parametry - dane tylko z nazwy")
    if rule is None:
        # marka na koncu nazwy dostawcy ("... czarny Kobi") idzie za rzeczownik, a nie drugi raz na koniec
        words = (re.sub(rf"\s+{re.escape(brand)}$", "", original) if brand else original).split()
        if words and re.search(r"[a-ząćęłńóśźż]", words[0]):
            # polski rzeczownik na poczatku ("Ramka", "Moduł") - marka za nim
            result["nazwa"] = " ".join(words[:1] + [brand] + words[1:])
        else:
            result["nazwa"] = " ".join([brand] + words)
        result["typ"] = "brak reguly"
        result["status"] = "DO_SPRAWDZENIA"
        result["uwagi"].append("brak reguly typu - nazwa oryginalna z marka")
        return result

    result["typ"] = rule["key"]
    result["typ_reguly"] = rule["key"]  # `noun_overrides` zmienia "typ" w raporcie, wersje licza sie po regule
    lead_len = 0
    if rules.get("type_words"):
        # GTV: "OPRAWA SUFITOWA ARALIA, AC220-240V, ..." - slowa typu z poczatku, dalej model
        type_words = set(rules["type_words"])
        first_words = split_segments(original)[0].split()
        prep_stop = {"DO", "Z", "ZE", "NA", "DLA", "W", "POD"} if rule.get("brand_before_prep") else set()
        for w in first_words:
            if ascii_upper(w) not in type_words or (lead_len and ascii_upper(w) in prep_stop):
                break  # Bemko: marka przed "do ..." ("Uchwyt regulowany Bemko do opraw ...")
            lead_len += 1
        while lead_len and ascii_upper(first_words[lead_len - 1]) in {"Z", "ZE", "I", "NA", "DO", "W"}:
            lead_len -= 1  # "GNIAZDO MEBLOWE WPUSZCZANE Z PRZESŁONĄ" - "z" zostaje przy modelu
    if rule.get("noun_from") == "tail":
        noun, gender = tail_noun(tail, rules) if tail else (rule.get("default_noun", ""), "f")
        gender = rule.get("gender", gender)
    elif rule.get("noun_from") == "type_words":
        lead = split_segments(original)[0].split()[:max(lead_len, 1)]
        phrase = nice_case(" ".join(lead), set(rules.get("keep_upper", [])))
        noun = phrase[:1].upper() + phrase[1:]
        first = lead[0].lower()
        gender = rule.get("gender") or ("f" if first.endswith("a") else ("n" if first.endswith(("o", "e")) else "m"))
        if rule.get("noun_add_led") and "LED" not in noun.upper().split():
            noun += " LED"
    else:
        noun, gender = resolve_noun(rule, u, original, rules)
    # tez rodzina z karty GTIN: "Orbis Flame" w Ekonomie to na karcie "Orbis Flame Wall" albo "... Ceiling" (2026-10-21)
    noun_src = f"{u} {parse_view(rules.get('_card_family_raw', {}).get(ean, ''))}"
    for item in rules.get("noun_overrides", []):
        # rodzaj oprawy z kategorii producenta zamiast ogolnego "Oprawa LED" (2026-10-19: "Oprawa LED nic nie mowi")
        if (not item.get("types") or rule["key"] in item["types"]) and re.search(item["match"], noun_src):
            noun, gender = item["noun"], item.get("gender", gender)
            result["typ"] = item.get("typ", result["typ"])
            break
    color_src = head_u if rules.get("split_tail") else u
    for pattern in rules.get("color_ignore", []):
        color_src = re.sub(pattern, " ", color_src)
    values: dict[str, str] = {}
    # karta GTIN tego EAN uzupelnia to, czego nazwa w Ekonomie nie podaje (zarowki / tuby, 2026-10-21): moc
    # ("SMART+ WiFi Classic A60 DIM 2700K" bez W), ksztalt, barwe ("2700…6500 K" = TW), protokol Smart+, T8
    card = (rules.get("_gtin_cards") or {}).get(ean, {})
    card_fill = card if (card.get("exact") and rule["key"] in rules.get("card_fill_types", [])
                         and not web_mismatch(original, card.get("title", ""))) else {}
    if card_fill and not power_text and card_fill.get("power") and "power" in rule.get("attrs", []):
        power = float(card_fill["power"])
        power_text = f"{fmt_num(power)}W"
        result["uwagi"].append(f"moc {power_text} z karty {card_fill.get('url') or 'ledvance.com'} (brak w Ekonomie)")
    tube_type = parse_tube_type(u)
    attrs0 = list(rule.get("attrs", []))
    # Ledvance (2026-10-19): linia "jesli jest" w kazdym typie; typy z seria / modelem ("Vintage 1906 PenduLum")
    # maja linie w serii
    add_line = rules.get("line_all_types") and not {"line", "series", "model", "model_code"} & set(attrs0)
    if rules.get("brand_last") and "brand" not in attrs0:
        # marka (+ linia) na koncu, jak w zarowkach: "Downlight LED fi200 18W 4000K bialy Ledvance Value"
        # rodzina opraw za marka, przed linia: "... Ledvance Urban Lite S" (`families`, 2026-10-19)
        tail_attrs = ["brand"] + (["family"] if rules.get("families") and not rules.get("family_after_noun") else [])             + (["line"] if add_line else [])
        if rules.get("families") and rules.get("family_after_noun"):
            # nazwa wlasna zaraz za rzeczownikiem (2026-10-20): "Naswietlacz LED Floodlight Compact 50W ... Ledvance"
            attrs0 = ["family"] + attrs0
        if rules.get("series_after_brand") and "series" in attrs0 and {"power", "lumen", "cct"} & set(attrs0):
            # seria oprawy z parametrami jest jak linia: "Naswietlacz LED 10W 800lm 4000K IP65 czarny Bemko Kastel 2";
            # opis z poczatku serii ("z akumulatorem", "puszkowy") zostaje na jej miejscu jako `series_desc`
            attrs0 = [("series_desc" if a == "series" else a) for a in attrs0]
            tail_attrs.append("series")
        rule = {**rule, "attrs": attrs0 + tail_attrs}
    elif add_line:
        rule = {**rule, "attrs": ([] if "brand" in attrs0 else ["brand"]) + ["line"] + attrs0}
    for attr in rule.get("attrs", []):
        if attr == "gwint":
            values[attr] = parse_socket(u) or (parse_dulux_socket(u, power, rules) if "DULUX" in u else "")
        elif attr == "tube_type":
            values[attr] = tube_type
        elif attr == "tube_length":
            values[attr] = parse_tube_length(u, tube_type, power, {**rules, **{k: v for k, v in rule.items()
                                                                                 if k == "fluorescent_lengths"}})
        elif attr == "length":
            values[attr] = parse_length(u)
        elif attr == "panel_size":
            values[attr] = parse_panel_size(head_u if rules.get("split_tail") else u)
            m = re.search(r"(?<!\d)(120|60|62|30)(30|60|62|120)(?!\d)", head_u) if rules.get("split_tail") else None
            if not values[attr] and m:
                values[attr] = f"{m.group(1)}x{m.group(2)}"  # Kanlux: "6060", "12030", "6262"
        elif attr == "diameter":
            values[attr] = parse_diameter(u, bare=bool(rule.get("diameter_bare")))
        elif attr == "power":
            values[attr] = power_text
            if power_text and (rules.get("split_tail") or rules.get("per_meter_power")) \
                    and re.search(r"\dW/M\b", head_u):
                values[attr] += "/m"
        elif attr == "lumen":
            lm = None
            info = pdf_lm.get(ean)
            name_lm = parse_name_lumen(u, rules)
            if rules.get("name_lumen_first") and name_lm:
                if info and len(info["lm"]) == 1 and abs(info["lm"][0] - name_lm) > 0.1 * name_lm:
                    result["uwagi"].append(f"katalog PDF podaje {fmt_num(info['lm'][0])}lm, nazwa {fmt_num(name_lm)}lm")
                info = None  # GTV: nazwa ma lm sklepu, katalog tylko uzupelnia braki
            if info and len(info["lm"]) == 1:
                lm = info["lm"][0]
                result["lm_zrodlo"] = "PDF: " + ", ".join(info["src"])
                if info.get("per_m"):
                    values[attr] = f"{fmt_num(lm)}lm/m"
                    continue
            elif info and len(info["lm"]) > 1:
                result["uwagi"].append("PDF-y podaja rozne lm: " + "/".join(fmt_num(x) for x in info["lm"]))
            if lm is None:
                lm = parse_name_lumen(u, rules)
                if lm:
                    result["lm_zrodlo"] = "nazwa"
            eff = re.search(r"(\d{2,3})\s?LM/W\b", u) if lm is None and rules.get("lumen_from_efficacy") else None
            if eff:
                # GTV: "150W-120W-90W, 150lm/W" -> strumien dla najwiekszej mocy
                watts = [to_float(w) for w in re.findall(r"(?<![\d.,])(\d+(?:[.,]\d+)?)\s?W\b", u)]
                if watts:
                    lm = round(max(watts) * int(eff.group(1)), -1)
                    result["lm_zrodlo"] = "moc x lm/W z nazwy"
                    result["uwagi"].append(f"lm wyliczone: {fmt_num(max(watts))}W x {eff.group(1)}lm/W")
            if lm is None and rules.get("_gtin_cards") and "tasma" not in rule["key"]:
                lm = card_lumen(rules["_gtin_cards"].get(ean, {}), power, parse_cct(u), result, power_text)
            if isinstance(lm, str):  # kilka mocy w nazwie ("11,3/14W") -> lm w tej samej kolejnosci ("1900/2300lm")
                values[attr] = f"{lm}lm"
                continue
            per_m = bool(lm) and bool(re.search(r"\d\s?LM/M\b", u))
            values[attr] = (f"{fmt_num(lm)}lm/m" if per_m else f"{fmt_num(lm)}lm") if lm else ""
        elif attr == "cct":
            values[attr] = parse_cct(u)
            kel = (pdf_lm.get(ean) or {}).get("k") or []
            if not values[attr] and len(kel) == 1:
                values[attr] = f"{kel[0]}K"  # barwa z katalogu PDF
            for item in rules.get("cct_words", []) if not values[attr] else []:
                if re.search(item["match"], u):
                    values[attr] = item["label"]  # "CIEPLY BIALY" bez K w nazwie i katalogu
                    break
            codes = rules.get("cct_codes", {})
            if not values[attr] and codes:
                # kody barwy Kanlux w kodzie modelu: "7W-WW", "AL-SL-NW-MAT", "36W 6262 NW"
                found = [codes[t] for t in re.findall(r"(?<![A-Z])(" + "|".join(codes) + r")(?![A-Z])", head_u)]
                found += [f"{m}00K" for m in re.findall(r"-[89](\d{2})\b", head_u) if m in VALID_CCT]
                if found:
                    values[attr] = found[0]
                    result["uwagi"].append(f"barwa z kodu w nazwie ({found[0]})")
        elif attr == "ip":
            values[attr] = parse_ip(u)
        elif attr == "control":
            values[attr] = parse_control(u)
        elif attr == "sensor":
            values[attr] = parse_sensor(u)
            if not values[attr] and rule.get("sensor_ms") and re.search(r"\bMS\b", u):
                values[attr] = "z czujnikiem ruchu"  # tuby "LEDTUBE T8 EM MS P 1200" = MOTION SENSOR (karta GTIN)
        elif attr == "emergency":
            values[attr] = parse_emergency(u)
        elif attr == "tube_supply":
            values[attr] = parse_tube_supply(u)
        elif attr == "strip":
            values[attr] = parse_strip(u)
        elif attr == "driver":
            values[attr] = parse_driver(u)
        elif attr == "ballast":
            values[attr] = parse_ballast(u)
        elif attr == "shape":
            values[attr] = parse_shape_code(u, rules) or parse_shape(u)
        elif attr.startswith("features"):
            # "features" - lista typu albo dostawcy; "features_pre" itp. - druga lista typu (inna pozycja w nazwie)
            items = rule.get(attr) if attr in rule else (rules.get("features", []) if attr == "features" else [])
            # pozycja-napis to nazwa wspolnej listy z `feature_sets` (czujniki, ksztalt...)
            items = [x for item in items for x in (rules.get("feature_sets", {}).get(item, [])
                                                   if isinstance(item, str) else [item])]
            values[attr] = parse_features(u, gender, rules, items, ean)
        elif attr == "cri":
            values[attr] = parse_cri(u)
        elif attr == "psu_kind":
            values[attr] = parse_psu_kind(u)
        elif attr == "socket_spec":
            values[attr] = parse_socket_spec(u)
        elif attr == "code":
            values[attr] = parse_code(u, rule)
        elif attr == "beam":
            shape = values.get("shape") or parse_shape(u)
            values[attr] = parse_beam(u, shape)
            beams = (pdf_lm.get(ean) or {}).get("beam") or []
            if not values[attr] and len(beams) == 1 and REFLECTOR_SHAPE_RE.match(shape):
                values[attr] = f"{beams[0]}°"
        elif attr == "line":
            values[attr] = parse_bulb_line(u)
            fam_line = rules.get("_family_lines", {}).get(ean, "")
            if fam_line and values[attr] in ("", "Ultra Output", "High Output"):
                # linia rodziny z kart GTIN (`line_from_card_family`): "LED CLASSIC B V" = Value, bo karta tej rodziny
                # pisze "VALUE CLAS B 60" - sama litera w Ekonomie bywa bledna ("CLASSIC P P" w rodzinie "... P V")
                values[attr] = f"{fam_line} {values[attr]}".strip()
                result["uwagi"].append(f"linia {fam_line} z rodziny na karcie ledvance.com")
        elif attr == "family":
            values[attr] = parse_family(u, rules)
            card_cfg = rules.get("family_from_card") or {}
            if not values[attr] and rule["key"] in card_cfg.get("types", []):
                # nazwa wlasna z rodziny na karcie GTIN ("... Ledvance Damp Proof Max"), 2026-10-20
                cards = rules.get("_card_families", {})
                found = cards.get(ean)
                if found is None:
                    # bez karty albo karta "Phase out": nazwa wlasna innych kart, jesli stoi doslownie w nazwie z Ekonoma
                    found = family_from_name(raw_u, rules.get("_family_phrases", []))
                    if found:
                        result["uwagi"].append(f"nazwa wlasna {found} z kart ledvance.com innych produktow (jest w nazwie)")
                words = found.split()
                noun_words = {w.lower() for w in noun.split()}
                while words and words[0].lower() in noun_words:
                    # slowo z poczatku nazwy nie wraca (decyzja 2026-10-20): "Panel LED ... Ledvance Flex",
                    # nie "... Ledvance Panel Flex"; "Oprawa High Bay LED ... Ledvance Compact Brackets"
                    words.pop(0)
                while len(words) > 1 and words[-1].lower() in noun_words:
                    words.pop()  # "Panel LED ... Ledvance Biolux HCL Panel" -> "... Biolux HCL" (2026-10-21)
                values[attr] = " ".join(words)
        elif attr == "click":
            # Orbis: sterowanie wylacznikiem - Click-DIM (sciemnianie) / Click-CCT (zmiana barwy), 2026-10-20
            # Remote-CCT / Remote-DIM - pilot (blizniak Click-CCT), samo "REMOTE" bez CCT/DIM nie jest cecha
            m = re.search(r"\b(CLICK|REMOTE)(?:[- ](CCT|DIM))?\b", u)
            if m and m.group(1) == "REMOTE" and not m.group(2):
                m = re.search(r"\bCLICK(?:[- ](CCT|DIM))?\b", u)
                m = m and re.match(r"(CLICK)(?:[- ](CCT|DIM))?", m.group(0))
            values[attr] = (f"{m.group(1).capitalize()}-{m.group(2)}" if m.group(2) else "Click") if m else ""
        elif attr == "remote":
            # dekor z pilotem w komplecie: "COB STRIPS CCT WITH REMOTE CONTROL 1.5M TW + RC" (2026-10-20)
            # samo "RC" tez: "LINEAR EDGE 30CM RGBW DIM RC BK" (karta: Remote Control), 2026-10-21
            values[attr] = "z pilotem" if re.search(r"\bREMOTE CONTROL\b|\+\s?RC\b|\bRC\b", u) else ""
        elif attr == "round":
            # oprawa okragla (2026-10-20: "jak jest round to fajnie by bylo dodac, ze plafon jest okragly")
            values[attr] = ({"m": "okrągły", "f": "okrągła", "n": "okrągłe"}.get(gender, "okrągły")
                            if re.search(r"\bROUND\b|\bRD\b|OKRAGL|\bCIRC(?:ULAR)?\b", u) else "")
        elif attr == "glass":
            values[attr] = parse_glass(u)
        elif attr == "dim":
            values[attr] = parse_dim(u)
        elif attr == "optics":
            values[attr] = parse_optics(u)
            beam = (rules.get("_gtin_cards") or {}).get(ean, {}).get("beam")
            if beam and values[attr] in LETTER_OPTICS.values():
                # kod optyki -> kat z karty tego EAN (2026-10-20: "opt. N" nic nie mowi): Low Bay Flex N = 60°,
                # W = 90°, VW = 120°, OV = 40x90°
                result["uwagi"].append(f"kat {beam} z karty ledvance.com zamiast {values[attr]}")
                values[attr] = beam
        elif attr == "color":
            values[attr] = parse_color(color_src, gender, rules.get("housing_colors", {}),
                                       two_tone=bool(rules.get("two_tone_colors")))
        elif attr == "series":
            src = head if rules.get("split_tail") else original
            if rules.get("series_first_segment"):
                # GTV: "OPRAWA SUFITOWA ARALIA, AC220-240V, ..." -> model z 1. czlonu, bez slow typu
                segs = split_segments(src)
                if rule.get("series_all_segments"):
                    segs = [" ".join(segs)]  # Bemko akcesoria: "UCHWYT ... LOP1 30W, JAS1 30/50W" - caly opis
                src = " ".join(segs[0].split()[lead_len:])
                if not src.strip() and len(segs) > 1 and re.fullmatch(r"\s*[A-Z][A-Z0-9 .+-]{1,20}", segs[1]) \
                        and not re.search(r"\d\s?(W|V|HZ|LM|K|MM|CM)\b|^\s*(IP|AC|DC|MAX|G\d|GU|E\d)", segs[1]):
                    src = segs[1]  # "OPRAWA NA SZYNOPRZEWÓD 1-FAZOWY, LAGOS II, ..." -> model z 2. czlonu
            values[attr] = parse_series(src, rule, rules)
        elif attr == "segments":
            values[attr] = parse_segments(original, rules)
        elif attr == "light_color":
            found = [c for code, c in rules.get("light_colors", {}).items() if re.search(rf"-{code}\b", head_u)]
            values[attr] = found[0] if found else ""
        elif attr == "breaker":
            values[attr] = parse_breaker(head_u)
        elif attr == "osp_code":
            m = re.search(r"\b\d{2}-\d{4}-\d{3}\b", head)
            values[attr] = m.group(0) if m else ""
        elif attr == "strip_length":
            m = re.search(r"(?<![\d.,])(\d+(?:[.,]\d+)?)\s?M\b(?!/)", head_u)
            values[attr] = f"{fmt_num(to_float(m.group(1)))}m" if m else ""
        elif attr == "head":
            values[attr] = " ".join(head.split())
        elif attr == "size":
            values[attr] = parse_deco_size(u, noun, bool(rules.get("deco_size_bare_cm")))
        elif attr == "sockets":
            values[attr] = parse_sockets(u)
        elif attr == "system":
            values[attr] = parse_system(u, rules)
        elif attr == "acc_size":
            values[attr] = parse_acc_size(u)
        elif attr == "poles":
            values[attr] = parse_poles(u)
        elif attr == "housing_tubes":
            values[attr] = parse_housing_tubes(u, tuple(rules.get("tube_words", ("świetlówkę", "świetlówki"))))
        elif attr == "ls_profile":
            values[attr] = parse_ls_profile(u)
        elif attr == "ls_cover":
            values[attr] = parse_ls_cover(u)
        elif attr == "ls_code":
            values[attr] = parse_ls_code(u)
        elif attr == "module_family":
            values[attr] = parse_module_family(u)
        elif attr == "model_code":
            values[attr] = parse_model_code(original, rule, rules)
        elif attr == "hid_model":
            values[attr] = parse_hid_model(u)
        elif attr == "hid_power":
            values[attr] = parse_hid_power(u)
        elif attr == "hid_color":
            values[attr] = parse_hid_color(u)
        elif attr == "ugr":
            values[attr] = parse_ugr(u)
        elif attr == "voltage":
            values[attr] = parse_voltage(u)
        elif attr == "model":
            if rule.get("noun_from") == "model_first" and noun:
                words = original.split()
                n_lead = len(noun.split())
                n_model = model_prefix_len(words)
                values[attr] = " ".join(words[:n_model] + words[n_model + n_lead:]) if n_model \
                    else clean_model(original[len(noun):], rules).lstrip(" ,")
            elif rule.get("noun_from") == "first_word":
                values[attr] = clean_model(original[noun_end(original, noun, u, rules):], rules).lstrip(" ,")
                if rules.get("model_polish_lower") and original.isupper():
                    # stara nazwa WIELKIMI: "Z WYSIĘGNIKIEM", "CZARNY", "DO NINA HB" -> polskie slowa malymi
                    values[attr] = " ".join("WiFi" if w == "WIFI" else w.lower()
                                            if is_polish_word(w) or ascii_upper(w) in POLISH_LOWER_EXTRA
                                            else w for w in values[attr].split())
            else:
                values[attr] = clean_model(original, rules, rule.get("strip_prefix", r"^\S+\s+"))
        elif attr == "device_extra":
            values[attr] = parse_device_extra(u, gender, rules)
    ml = card_ml_cct(card.get("name", "")) if card.get("exact") and "cct" in values else ""
    if ml and (not values["cct"] or values["cct"].rstrip("K") in ml.rstrip("K").split("/")):
        # Ekonom "13W MS 940" / "10W CPS" = Multi Select, karta: "ML 93040" - dwie barwy do wyboru (przeglad 2026-10-21)
        if values["cct"] != ml:
            result["uwagi"].append(f"barwa {ml} z karty ledvance.com (Multi Select, w Ekonomie: {values['cct'] or 'brak'})")
        values["cct"] = ml
    if card_fill:
        cu = parse_view(re.sub(r"\bCLAS\b", "CLASSIC", card_fill.get("name", "")))
        if "shape" in values and values["shape"] in ("", "globe"):
            # "CLASSIC LAMPS FROSTED" / "A,B,P SLIM" bez ksztaltu -> z karty; "globe" dostaje rozmiar ("globe G125")
            found = parse_shape_code(cu, rules) or parse_shape(cu)
            if found and (not values["shape"] or found.startswith("globe")):
                values["shape"] = found
        if "cct" in values and not values["cct"] and rule.get("cct_from_card") and card_fill.get("cct"):
            # karta podaje barwe tylko liczba ("Color temperature 3000 K"), nie w nazwie - kamery Smart+ (2026-10-21)
            values["cct"] = f"{card_fill['cct']}K"
        if "cct" in values and (not values["cct"] or values["cct"] == "2700K"):
            found = card_cct(cu)
            if found and (not values["cct"] or found == "TW"):
                values["cct"] = found
        if "control" in values and not values["control"] and re.search(r"SMART", u):
            values["control"] = parse_control(cu)  # "SMART+ Filament Classic Dimmable" - protokol z karty (BT)
        if "tube_type" in values and not values["tube_type"]:
            values["tube_type"] = parse_tube_type(cu)  # "Tuba LED Entry EM 16 W" - karta "ST8E-EM"
        if "gwint" in values and not values["gwint"]:
            values["gwint"] = parse_socket(cu)
        if "features" in values and not values["features"] and rule.get("features"):
            # "HQL LED Special Accessory W/" - strumien lampy, do ktorej jest klosz, tylko na karcie ("... 2000LM LAMP SHADE")
            values["features"] = parse_features(cu, gender, rules, rule["features"], ean)
    if rules.get("hid_cct_from_card") and values.get("hid_color") == "barwa dzienna" and card.get("exact") \
            and card.get("cct"):
        # HQI D / DL: barwa w K z karty tego EAN zamiast "barwa dzienna" (decyzja uzytkownika 2026-10-21); karta ponizej
        # 5000 K przeczy D / DL ("HQI TS 1000W DL/S" = karta "HQI-TS 1000 W/NDL" 4400 K) - do sprawdzenia
        values["hid_color"] = f"{card['cct']}K"
        result["uwagi"].append(f"barwa {card['cct']}K z karty {card.get('url') or 'ledvance.com'} (w Ekonomie: D/DL)")
        if card["cct"] < 5000:
            result["uwagi"].append(f"karta podaje {card['cct']}K ({card.get('name', '')}), Ekonom D/DL - do sprawdzenia")
            result["_card_name_only"] = True
    if rules.get("params"):
        apply_params(values, rule, feed_rec, rules, gender, result)
    else:
        apply_feed(values, rule, feed_rec, rules, gender, result)
    if rules.get("_web_params") is not None:
        web_rec = rules["_web_params"].get(ean)
        if web_rec:
            apply_web_params(values, rule, web_rec, rules, result)
        elif ean in rules.get("_withdrawn", ()):
            result["uwagi"].append(rules.get("web_withdrawn_note", "brak na stronie producenta (wycofany)"))
    missing = [a for a in rule.get("required", []) if not values.get(a)
               and not re.search(rule.get("required_unless", {}).get(a, r"(?!x)x"), u)]
    if values.get("sockets") and not values.get("power"):
        # oprawa na wymienne zrodlo (1xGU10, G5,3/GU10) nie ma wlasnego strumienia
        missing = [a for a in missing if a != "lumen"]
    if result.get("wymienne_zrodlo"):
        # Kanlux: trzonek z parametrow (NOTUS 418 -> 4xG13) - moc, lm i barwa zaleza od wlozonej zarowki
        missing = [a for a in missing if a not in ("power", "lumen", "cct")]
    if values.get("gwint") and values["gwint"] in values.get("shape", "").split():
        values["gwint"] = ""  # "R7s 118mm" - trzonek jest juz w ksztalcie
    if values.get("line") == "antybakteryjna" and "dim" in rule.get("attrs", []):
        # to cecha, nie linia - stoi z cechami przed marka: "... matowa antybakteryjna Ledvance"
        values["line"], values["dim"] = "", f"{values.get('dim', '')} antybakteryjna".strip()
    if "series_desc" in rule.get("attrs", []) and values.get("series"):
        # GTV: seria bywa opisem pisanym malymi ("z akumulatorem Alledo", "puszkowy", "specjalna do Maszyn Szycia") -
        # opis zostaje przed parametrami, za marke idzie tylko model ("Alledo"); opis konczacy sie przyimkiem
        # ("specjalna do") niesie dalsza czesc ze soba
        words = values["series"].split()
        k = 0
        while k < len(words) and words[k][:1].islower():
            k += 1
        if k and k < len(words) and ascii_upper(words[k - 1]) in {"DO", "Z", "ZE", "NA", "DLA", "POD", "W", "I", "BEZ"}:
            k = len(words)
        values["series_desc"], values["series"] = " ".join(words[:k]), " ".join(words[k:])
    if rules.get("onoff_twins") and "control" in values and not values["control"] and ONOFF_RE.search(u):
        # "DAMP PROOF ECO 3TH ON-OFF" - zostaje tylko przy blizniaku DALI / 1DIM (`drop_lone_onoff`, 2026-10-21)
        values["control"] = "ON-OFF"
    for word, (key, full) in (rules.get("feature_abbrev") or {}).items():
        # skroty Ledvance w nazwie z Ekonoma: "DP HE EL 1500", "DP HE SN 1500" - EL = awaryjna, SN = czujnik (2026-10-20)
        parser = {"sensor": parse_sensor, "emergency": parse_emergency, "control": parse_control}[key]
        if key in rule.get("attrs", []) and not values.get(key) and re.search(rf"\b{word}\b", u):
            values[key] = parser(full)
    card_cfg = rules.get("family_from_card") or {}
    raw_family = rules.get("_card_family_raw", {}).get(ean, "") if rule["key"] in card_cfg.get("types", []) else ""
    if raw_family:
        # rodzina na karcie tego EAN mowi wiecej niz Ekonom ("SURFACE SLIM IK10 RD 200 SENSOR" przy "SURFACE SLIM RD S 200"):
        # czujnik / awaryjna / DALI z karty, gdy nazwa ich nie podaje (2026-10-20)
        # tez nazwa produktu na karcie: rodzina "Orbis IP44", nazwa "Orbis IP44 Sensor 300mm 15.5W" (2026-10-20)
        fu = parse_view(f"{raw_family} {(rules.get('_gtin_cards') or {}).get(ean, {}).get('name', '')}")
        for key, parser in (("sensor", parse_sensor), ("emergency", parse_emergency), ("control", parse_control)):
            if key in rule.get("attrs", []) and not values.get(key):
                found = parser(fu)
                if found:
                    values[key] = found
                    if parser(parse_view(raw_family)) or key == "control":
                        result["uwagi"].append(f"{found} z rodziny na karcie ledvance.com ({raw_family})")
                    else:
                        # tylko nazwa karty, rodzina nie potwierdza: "Essentials Floodlights with EU Plug" 20W ma na
                        # karcie "Floodlight Essential Sensor", a 10W / 30W tej rodziny nie - do sprawdzenia
                        result["uwagi"].append(f"{found} tylko z nazwy na karcie ledvance.com (rodzina: {raw_family})"
                                               " - do sprawdzenia")
                        result["_card_name_only"] = True
    # Ekonom + rodzina i nazwa z karty: "DAMP PROOF LED ECO 600 TH" to na karcie "ECO CLASS DAMP PROOF"
    card_name = (rules.get("_gtin_cards") or {}).get(ean, {}).get("name", "") if raw_family else ""
    extra_src = f"{u} {parse_view(raw_family)} {parse_view(card_name)}" if raw_family else u
    for item in rules.get("family_strip", []):
        # slowa nazwy wlasnej, ktore mowi juz kat z karty ("Damp Proof Eco Narrow ... 90x60°" -> "Damp Proof Eco")
        if rule["key"] in item.get("types", []) and values.get("family"):
            values["family"] = " ".join(re.sub(item["match"], " ", values["family"]).split())
    for item in rules.get("family_extra", []):
        # dopisek do nazwy wlasnej tylko w wybranych typach (2026-10-20): "DAMP PROOF HE" -> "Damp Proof High Efficiency",
        # `prepend` - przed nazwa ("Eco Class Damp Proof", jak u Ledvance); `{0}` = grupa z `match`, `family` = dopisek
        # tylko do tej nazwy wlasnej ("PRINT SKY1" -> "Print Sky 1", 2026-10-21)
        m_extra = re.search(item["match"], extra_src) if rule["key"] in item.get("types", []) else None
        if not m_extra or (item.get("family") and not re.search(item["family"], values.get("family", ""))):
            continue
        add = item["add"].format(*(g or "" for g in m_extra.groups()))
        if add.lower() not in values.get("family", "").lower():
            parts_f = [add, values.get("family", "")] if item.get("prepend") else [values.get("family", ""), add]
            values["family"] = " ".join(p for p in parts_f if p)
    if values.get("round") and re.search(r"(?i)\b(?:ROUND|RD|CIRCULAR)\b", values.get("family", "")):
        values["round"] = ""  # "Ceiling Round" juz mowi, ze okragly (2026-10-20)
    if "gwint" in values and values["gwint"] == "GU5.3" and rules.get("gu53_voltage"):
        # GU5.3 to zwykle 12V (MR16) - sprzedawca ma to widziec (2026-10-20); 230V tylko gdy nazwa / karta tak mowi
        card_name = ascii_upper((rules.get("_gtin_cards") or {}).get(ean, {}).get("name", ""))
        values["gwint"] = "GU5.3 230V" if re.search(r"230\s?V", f"{u} {card_name}") else "GU5.3 12V"
    if "click" in values and values["click"]:
        if values.get("cct") == "CCT" and values["click"].endswith("-CCT"):
            values["cct"] = ""  # "CCT Click-CCT" - przelaczana barwa wystarczy raz
        fam_words = values.get("family", "").split()
        if fam_words and fam_words[-1] == "Click":
            # "Orbis Ultra Slim Click" + "Click-DIM" -> "Orbis Ultra Slim ... Click-DIM"; samo "Click" juz jest w rodzinie
            values["family"] = " ".join(fam_words[:-1]) if values["click"] != "Click" else values["family"]
            if values["click"] == "Click":
                values["click"] = ""
    if values.get("family") and values.get("housing_tubes", "").startswith("pod ") and rules.get("family_after_noun"):
        # "Oprawa hermetyczna pod tube LED Damp Proof Housing 1x120cm", nie "Oprawa hermetyczna Damp Proof Housing pod tube"
        values["housing_tubes"] = re.sub(r"^(pod \S+ LED)", lambda m: f"{m.group(1)} {values['family']}", values["housing_tubes"])
        values["family"] = ""
    for key in ("family", "line"):
        found = values.get(key, "")
        if found and any(re.search(rf"(?i)(?<!\S){re.escape(found)}(?!\S)", str(v))
                         for v in [noun, rule.get("noun_suffix", "")] + [x for k, x in values.items() if k != key]):
            values[key] = ""  # rodzina / linia jest juz w serii / modelu / "Smart+" rzeczownika
        elif key == "family" and found:
            # wszystkie slowa rodziny sa juz w modelu ("UVC Linear Housing" przy "do UVC 440mm LINEAR HOUSING")
            others = {w.lower() for v in [noun] + [x for k, x in values.items() if k != key] for w in str(v).split()}
            if {w.lower() for w in found.split()} <= others:
                values[key] = ""
    attrs = rule.get("attrs", [])
    if "brand" in attrs:
        # marka na pozycji z `attrs` (zarowki Ledvance: "... 36° Osram Value", jak "... 36ST Philips CorePro")
        parts = [noun, rule.get("noun_suffix", "")] + [brand if a == "brand" else values.get(a, "") for a in attrs]
    else:
        parts = [noun, rule.get("noun_suffix", ""), brand] + [values[a] for a in attrs if values.get(a)]
    sock = result.get("wymienne_zrodlo", "")
    if sock and sock.split("x")[-1].upper() not in " ".join(parts).upper():
        at = parts.index(values["color"]) if values.get("color") in parts else len(parts)
        parts.insert(at, sock)
    result["nazwa"] = " ".join(p for p in parts if p)
    if "brand" in attrs:  # slowa za marka (linia, protokol) - `move_brand_tail` trzyma je przy marce
        result["po_marce"] = " ".join(values.get(a, "") for a in attrs[attrs.index("brand") + 1:]).split()
        result["linia"] = "line" in attrs
    if rules.get("dedupe_words"):
        # GTV: cecha z nazwy i ten sam czlon w modelu ("kwadratowa Noelia ... kwadratowa"), pusty nawias po
        # wycieciu mocy ("VILANA ( )")
        colors = {w for c in rules.get("housing_colors", {}).values() for g in ("m", "f", "n") for w in c.get(g, "").split()}
        words, seen = [], set()
        for w in re.sub(r"\(\s*\)", " ", result["nazwa"]).split():
            key = w.lower()
            # tylko przymiotniki cech; kolor obudowy po barwie ("neutralna biała ... biała") i rzeczowniki zostaja
            if key in seen and key not in colors and (re.fullmatch(r"[a-ząćęłńóśźż]{3,}(owa|owy|owe|na|ny|ne|ła|ły|łe)", key)
                                                      or w in ("RGB", "RGBW", "CCT") or re.fullmatch(r"\d+LED", w)):
                continue
            seen.add(key)
            words.append(w)
        result["nazwa"] = " ".join(words)
    result["braki"] = missing
    result["atrybuty"] = values
    mismatch = code_mismatches(ean, raw_u, rules)
    result["uwagi"] += mismatch
    if missing or mismatch:
        result["status"] = "DO_SPRAWDZENIA"
    elif not rule.get("trusted") and ("model" in rule.get("attrs", []) or "segments" in rule.get("attrs", [])
                                      or (rule.get("no_schema") and not values.get("breaker"))):
        result["status"] = "WARNING"
        result["uwagi"].append("typ bez schematu - nazwa z modelu katalogowego")
    else:
        # cecha tylko z nazwy na karcie, rodzina jej nie potwierdza - do sprawdzenia (2026-10-20)
        result["status"] = "WARNING" if result.pop("_card_name_only", False) else "OK"
    result.pop("_card_name_only", None)
    return result


ONOFF_RE = re.compile(r"\bON[- ]?OFF\b|\bON/OFF\b")
CONTROL_TWIN_RE = re.compile(r"(?<!\S)(?:ON-OFF|DALI(?:-2)?|1DIM|2DIM)(?!\S)")


def drop_lone_onoff(results: list[dict[str, Any]]) -> None:
    """"ON-OFF" w nazwie tylko wtedy, gdy ten sam produkt jest tez w wersji DALI / 1DIM (decyzja uzytkownika 2026-10-21).

    Blizniak = ta sama nazwa po zdjeciu slow sterowania. Bez blizniaka "ON-OFF" nic nie mowi sprzedawcy (oprawa bez
    DALI jest wlaczana zwyklym wylacznikiem) i wypada - takze, gdy przyszlo jako wyroznik ze strony ledvance.com.
    """
    def base(name: str) -> str:
        return " ".join(CONTROL_TWIN_RE.sub(" ", name).split())

    dimmable = {base(r["nazwa"]) for r in results if re.search(r"(?<!\S)(?:DALI(?:-2)?|1DIM|2DIM)(?!\S)", r["nazwa"])}
    for res in results:
        if re.search(r"(?<!\S)ON-OFF(?!\S)", res["nazwa"]) and base(res["nazwa"]) not in dimmable:
            res["nazwa"] = " ".join(re.sub(r"(?<!\S)ON-OFF(?!\S)", " ", res["nazwa"]).split())
            res["uwagi"].append("bez ON-OFF - brak blizniaka DALI / 1DIM")


# ---------------------------------------------------------------- powtorzenia

DIFF_CATEGORIES = [
    ("klosz", r"^(CLEAR|FROSTED|FR|CL|MATOWE|MATOWA|PRZEZROCZYSTE|SZKLO|PLASTIK|GLASS|OPAL|OP|MS|PS|DSFR)$"),
    ("sciemnialnosc", r"^(DIM|NON-DIM|NON|DIMMABLE|SCIEMNIALNA|GLOWDIM|1DIM|2DIM)$"),
    ("kolor obudowy", r"^(BK|WT|GY|GR|SI|DG|BN|CR|WAL|ALW|WA|BLACK|WHITE|GRAY|GREY|SILVER|CHROME|GOLD|SMOKE|AMBER|BG|ST)$"),
    ("ksztalt", r"^(A|B|P|BA|BW|R50|R63|R80|GLOBE|GLOBE95|GLOBE125|STICK|EDISON|PAR16|PAR20|PAR30|PAR38|MR16|MR11|A60|A40|A75|B40|P40|P45|T26|PIN|CLA|CLB|CLP|SPOT|REFLECTOR|FLAT)$"),
    ("optyka / kat", r"^(SYM|ASYM|ASYMMETRIC|SYMMETRIC|R\d{2}|A\d{2}X\d{2,3}|SYM\d{2,3}|\d{2,3}DEG|\d{2,3}°|N|W|VW|WW|DS|OV|NFL|FL|SPOT|WIDE|NARROW|NB|MB|WB|RU\d{2}ST|RV\d{2}ST)$"),
    ("linia produktu", r"^(VALUE|PERFORMANCE|SUPERIOR|STAR|SUPERSTAR|BASE|CLASSIC|PRO|COMFORT|ECO|ESSENTIAL|ESSENTIALS|ENTRY|PLUS|SPECIAL|RETROFIT|PCR|PARATHOM|BELLALUX|FILAMENT|VINTAGE|1906|ENERGY|EFFICIENCY|CLASS|V|S|E)$"),
    ("sterowanie", r"^(DALI|ZIGBEE|ZB|WIFI|BT|DALIVR|DAVR|DA|ON|ONEUE|PS)$"),
    ("opakowanie", r"^(\d+X\d+|\d?PACK|FS\d|\d+PACK)$"),
]


VERSION_STOP = {"LED", "LEDV", "LEDVANCE", "OSRAM", "EUE", "LAMPA", "OPRAWA", "W", "K", "MM", "CM", "LM",
                "DAMP", "PROOF", "HOUSING", "XLAMP", "BK", "WT", "GY", "GR", "SI", "DG", "BN", "CR", "WH", "BLACK", "WHITE", "GRAY", "GREY",
                "DALI", "IP65", "IP44", "IP66", "IP20", "IP54", "IP67", "IP69", "EM", "EMERGENCY", "SENSOR",
                "UGR19", "U19", "UGR", "TUBA", "TUBY", "TUBE", "LEDTUBE", "LINIOWA", "SCIENNA", "SUFITOWA", "ZASILACZ",
                "ZAROWKA", "HALOGENOWA", "STATECZNIK", "ELEKTRONICZNY", "PASEK", "STEROWNIK", "DRIVER", "LS",
                "PLUG", "OUTDOOR", "OUTD", "ZB", "BT", "WIFI", "SMART", "SMART+", "FS1", "FS2", "FS3", "FS5", "FS10",
                "UO", "ADV", "PFM"}  # skroty linii (Ultra Output, Advanced, Performance) - linia jest juz w nazwie


def version_tokens(name: str, letters: dict[str, str] | None = None,
                   numbers: bool = False, with_word: bool = False) -> list[tuple[str, str]]:
    """Slowa nazwy katalogowej, ktore moga wyroznic wersje: (klucz, oryginal).

    `with_word` (Kobi): sam numer albo litera ("SOLARIS 2", "LAITILINE L") nic sprzedawcy nie mowi -
    wyroznik dostaje poprzedzajace slowo modelu.
    """
    name = re.sub(r"(?<=[A-Za-z0-9])(EUE)?(LEDV(ANCE)?|OSRAM)", " ", name)
    out = []
    prev = ""
    for tok in name.replace("/", " ").split():
        key = ascii_upper(tok).strip(",.")
        word = prev
        prev = tok.strip(",.") if re.fullmatch(r"[A-Z][A-Z-]{1,}", key) and key not in VERSION_STOP else ""
        if not key or key in VERSION_STOP:
            continue
        sized = re.fullmatch(r"(\d{2,3})-[A-Z]{1,2}", key) if with_word else None
        if sized:
            # Kanlux: rozmiar sklejony z kodem koloru ("STOBI DLP 350-W", "MERSA 300-W/M") - w nazwie jest juz kolor
            if numbers:
                out.append((sized.group(1), f"{word} {sized.group(1)}" if word else sized.group(1)))
            continue
        if not re.search(r"[A-Z]", key):
            # numer modelu ("QUAZAR 17") - tylko w drugim podejsciu, bez mocy/barwy/wymiarow
            if numbers and re.fullmatch(r"\d{1,3}(?:[.,]\d{1,2})?", key):
                out.append((key, f"{word} {tok.strip(',')}" if with_word and word else tok.strip(",")))
            continue
        if with_word and word and re.fullmatch(r"[A-Z]\d?", key):
            out.append((key, f"{word} {tok.strip(',.')}"))
            continue
        if re.fullmatch(r"[a-ząćęłńóśźż]+|[A-ZĄĆĘŁŃÓŚŹŻ][a-ząćęłńóśźż]*[ąćęłńóśźż][a-ząćęłńóśźż]*", tok.strip(",.")):
            continue  # polskie slowa opisowe ("wykrywający", "Połączenie") nie sa wersja
        if re.fullmatch(r"[\d.,X-]+(W|K|MM|CM|LM|V|M)?", key):
            continue
        if SOCKET_RE.fullmatch(key):
            continue  # gwint jest juz w nazwie
        if len(key) > 10 and re.search(r"\d", key) and "-" not in key:
            continue  # sklejony kod zamowieniowy, np. 740RW20STTWAC2M7D4 (kody z myslnikami zostaja)
        tok = (letters or {}).get(key, tok.strip(","))
        out.append((key, tok))
    return out


def add_version_suffixes(results: list[dict[str, Any]], rules: dict[str, Any]) -> None:
    """Drugi przebieg: gdy nazwy nadal sie powtarzaja, dopisz wyroznik wersji z nazwy katalogowej.

    Tylko dla typow z `version_suffix: true` (decyzja uzytkownika 2026-10-14: najpierw kat, a gdy
    nie wystarcza - wersje). Identyczne nazwy w Ekonomie zostaja powtorzeniem.
    """
    allowed = {t["key"] for t in rules["types"] if t.get("version_suffix")}
    for numbers in (False, True):  # najpierw slowa, potem tez numery modeli
        groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for res in results:
            if res.get("typ_reguly", res["typ"]) in allowed:
                groups[res["nazwa"]].append(res)
        _suffix_groups(groups, rules, numbers)


def _suffix_groups(groups: dict[str, list[dict[str, Any]]], rules: dict[str, Any], numbers: bool) -> None:
    for name, members in groups.items():
        if len(members) < 2 or len({m["stara"] for m in members}) < 2:
            continue
        # w lampach wyladowczych N/E to kody barwy i trzonka, nie klasy Value/Eco
        letters = rules.get("version_letters", {})
        toks = [version_tokens(m.get("zrodlo_nazwy") or m["stara"], None if m["typ"] in ("sodowa", "metalohalogenkowa") else letters,
                               numbers, bool(rules.get("version_number_with_word")))
                for m in members]
        common = set.intersection(*({k for k, _ in t} for t in toks))
        common |= set(ascii_upper(name).replace("+", "").split())  # slowa juz obecne w nazwie
        for m, t in zip(members, toks):
            extra, seen = [], set()
            for key, orig in t:
                if any(re.fullmatch(p, key) for p in rules.get("version_skip", [])):
                    continue  # GTV: "MA", "25SZT.", "AC180-250V" nie sa wersja
                if m["typ"] == "metalohalogenkowa" and key in {"N", "NDL", "WDL", "D", "DL"}:
                    continue  # barwa HQI jest juz w nazwie (4000K / 3000K / barwa dzienna), 2026-10-21
                if key.isdigit() and " " not in orig and (re.fullmatch(r"[789]\d\d", key) or key in ("220", "230", "240")
                                      or f"{key}W" in ascii_upper(name).replace(",", ".")):
                    continue  # kod barwy, napiecie albo moc, ktora juz jest w nazwie
                # litera klasy ("E" -> "Eco") porownywana tez po rozwinieciu: bez "ECO Eco"
                if not {key, ascii_upper(orig)} & (common | seen):
                    extra.append(orig)
                    seen |= {key, ascii_upper(orig)}
            if extra:
                # "SOLARIS" + "SOLARIS 2" -> raz; wielkosc liter bez znaczenia ("ECO" / "Eco")
                uniq: dict[str, str] = {}
                for w in " ".join(extra).split():
                    uniq.setdefault(ascii_upper(w), w)
                extra = list(uniq.values())
                note = " ".join(extra)
                words = name.split()
                for i, w in enumerate(words):
                    key = ascii_upper(w)
                    if key in extra[:-1] and re.fullmatch(r"[A-Z]\d?|\d{1,3}(?:[.,]\d{1,2})?",
                                                          extra[extra.index(key) + 1]):
                        # seria jest juz w nazwie ("Kobi Quazar ...") - numer idzie za nia: "Kobi Quazar 15 1xGU10"
                        j = extra.index(key)
                        words.insert(i + 1, extra[j + 1])
                        extra = extra[:j] + extra[j + 2:]
                        break
                m["nazwa"] = " ".join(words + extra)
                m["uwagi"].append("dopisana wersja z nazwy katalogowej: " + note)


def diff_tokens(names: list[str]) -> list[set[str]]:
    toks = [set(re.findall(r"[A-Z0-9°+/.,-]+", ascii_upper(n).replace("/", " "))) for n in names]
    common = set.intersection(*toks) if toks else set()
    return [t - common for t in toks]


def diff_categories(diffs: list[set[str]]) -> list[str]:
    cats = []
    allt = set().union(*diffs) if diffs else set()
    for cat, pattern in DIFF_CATEGORIES:
        if any(re.match(pattern, t) for t in allt):
            cats.append(cat)
    return cats


# ---------------------------------------------------------------- main

JOINED_NAME_SEP = " | "


def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(encoding="cp1250", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=";")
        fields, rows = list(reader.fieldnames or []), list(reader)
    if "Nazwa dodatkowa" not in fields and "Nazwa towaru" in fields:
        # "Zestawienie" z Ekonoma (2026-10-26): obie czesci w jednej kolumnie, "<40 znakow> | <reszta>"
        for row in rows:
            head, _, tail = row["Nazwa towaru"].partition(JOINED_NAME_SEP)
            row["Nazwa towaru"], row["Nazwa dodatkowa"] = head, tail
    return fields, rows


# Uklad pliku do Ekonoma (jak eksport Ledvance.csv). Zestawienie ma inny uklad (nazwa "<40> | <reszta>" w jednej
# kolumnie, ilosci i stany) - wynik zawsze idzie w tym ukladzie (decyzja uzytkownika 2026-10-27).
EKONOM_FIELDS = ["LP", "Kod towaru", "Nazwa towaru", "Nazwa dodatkowa", "Kod producenta", "JM", "Cena netto", "Waluta"]
ZESTAWIENIE_FIELD_SOURCES = {"JM": "J.M.", "Cena netto": "Cena C3 - Cena Katalogowa"}


def to_ekonom_layout(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    out = []
    for row in rows:
        new = {f: row.get(ZESTAWIENIE_FIELD_SOURCES.get(f, f)) or "" for f in EKONOM_FIELDS}
        new["Waluta"] = new["Waluta"] or "PLN"
        out.append(new)
    return out


def write_ekonom_csv(path: Path, fields: list[str], rows: list[dict[str, str]]) -> None:
    if "Nazwa dodatkowa" not in fields:
        fields, rows = EKONOM_FIELDS, to_ekonom_layout(rows)
    with path.open("w", encoding="cp1250", newline="", errors="replace") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, delimiter=";", quoting=csv.QUOTE_NONNUMERIC, lineterminator="\r\n")
        fh.write(";".join(fields) + "\r\n")  # naglowek bez cudzyslowow, jak w eksporcie Ekonoma
        for row in rows:
            out = dict(row)
            for key in ("LP",):
                if out.get(key, "").isdigit():
                    out[key] = int(out[key])
            if "Cena netto" in out:
                out["Cena netto"] = _raw_number(out["Cena netto"])
            writer.writerow(out)


class _raw_number(float):
    """Liczba zapisywana bez cudzyslowu, ale z oryginalnym przecinkiem z Ekonoma."""

    def __new__(cls, text: str):
        obj = super().__new__(cls, float(text.replace(",", ".") or 0))
        obj.text = text
        return obj

    def __repr__(self) -> str:
        return self.text

    __str__ = __repr__


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--supplier", required=True)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--pdf-dir", type=Path)
    parser.add_argument("--feed", type=Path, help="feed XML dostawcy (format Ceneo) z atrybutami produktow")
    parser.add_argument("--params-dir", type=Path,
                        help="katalog z Parametry*.xlsx (format dlugi Kanlux) - parametry producenta po kodzie")
    parser.add_argument("--rules", type=Path, default=DEFAULT_RULES)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    rules = load_yaml(args.rules)["suppliers"][args.supplier]
    fields, rows = read_csv(args.input)
    stamp = date.today().isoformat()
    label = args.supplier.capitalize()
    output = args.output or ROOT / "output" / f"Ekonom-nazwy-{label}_{stamp}.csv"
    report = args.report or ROOT / "reports" / f"Ekonom-nazwy-{label}_kontrola_{stamp}.xlsx"
    if output.resolve() == args.input.resolve():
        raise SystemExit("Plik wynikowy nie moze nadpisac pliku wejsciowego.")

    pdf_lm: dict[str, dict[str, Any]] = {}
    pkg: dict[str, dict[str, Any]] = {}
    if args.pdf_dir:
        products = {}
        for row in rows:
            u = parse_view(full_name(row))
            _, power = parse_power(u)
            if power:
                kel = {int(k[:4]) for k in re.findall(r"\d{4}", parse_cct(u))}
                products[(row.get("Kod producenta") or "").strip()] = (power, kel)
        texts = pdf_texts(args.pdf_dir, ROOT / "cache" / "ekonom_pdf" / args.supplier)
        if rules.get("pdf_match") == "code":
            # katalog bez EAN - wiersze tabel zaczynaja sie od kodu producenta (GTV: LD-SMBD45-10)
            pdf_lm = lumens_from_pdf_codes(texts, {(r.get("Kod producenta") or "").strip() for r in rows} - {""})
        else:
            pdf_lm = lumens_from_pdfs(texts, products, by_code=rules.get("pdf_match") == "code_table")
        if rules.get("package_from_pdf"):
            pkg = package_from_pdfs(texts, rules.get("package_brand", {}).get("catalog_columns", {}),
                                    rules.get("package_catalogs", []),
                                    list(rules.get("distinguish", {}).get("catalogs", {})))
        ensure_dir(ROOT / "cache" / "ekonom_pdf")
        (ROOT / "cache" / "ekonom_pdf" / f"{args.supplier}_lumeny.json").write_text(
            json.dumps(pdf_lm, ensure_ascii=False, indent=1), encoding="utf-8")

    feed = load_feed(args.feed) if args.feed else None
    if args.params_dir:
        feed = load_params(args.params_dir)
        rules = {**rules, "params": True}
        if rules.get("params_web_cache"):
            # karty z kanlux.com (scripts/fetch_kanlux_web.py) - dla kodow spoza plikow Parametry
            web_params, withdrawn = load_params_web(ROOT / rules["params_web_cache"])
            feed = {**web_params, **feed}
            rules["_withdrawn"] = withdrawn
    if feed is not None:
        hits = sum(1 for row in rows if feed_key(row) in feed)
        print(f"Feed: {len(feed)} produktow, dopasowane do CSV: {hits}/{len(rows)}")

    if rules.get("web_params_cache"):
        web_params, withdrawn, vocab, models = load_web_params(ROOT / rules["web_params_cache"])
        rules = {**rules, "_web_params": web_params, "_withdrawn": withdrawn, "_vocab": vocab, "_models": models}
        print(f"Strona producenta: {len(web_params)} kart, {len(withdrawn)} kodow spoza strony")

    web_cache = rules.get("distinguish", {}).get("web_cache")
    web = load_web_names(ROOT / web_cache) if web_cache else {}  # scripts/fetch_ledvance_web.py
    if rules.get("lumen_from_gtin_card"):
        rules = {**rules, "_gtin_cards": web}
    if rules.get("line_from_card_family"):
        rules = {**rules, "_family_lines": card_family_lines(
            web, (rules.get("family_from_card") or {}).get("placeholders", []))}
    if rules.get("family_from_card"):
        rules = {**rules, "_card_families": card_families(web, rules["family_from_card"]),
                 "_card_family_raw": {e: c["family"] for e, c in web.items() if c.get("exact") and c.get("family")},
                 "_family_phrases": family_phrases(web, rules["family_from_card"])}

    results = []
    for row in rows:
        res = build_name(row, rules, pdf_lm, feed)
        res["row"] = row
        res["stara"] = full_name(row)
        res = apply_web_card(res, row, rules, pdf_lm, feed)
        site = web.get((row.get("Kod producenta") or "").strip(), {})
        if rules.get("distinguish", {}).get("fix_from_gtin_page") and site.get("exact") \
                and web_mismatch(res["stara"], site["title"]):
            # karta ledvance.com/en/gtin to dokladnie ten EAN (nadruk na opakowaniu) - nazwa w Ekonomie
            # opisuje inny produkt (zwykle kopia sasiedniej kartoteki), nazwe budujemy z karty
            # rzeczownik z Ekonoma ("Lampa", "Tuba", "Oprawa") pozwala dobrac typ; "CLAS A" = "CLASSIC A"
            lead = res["stara"].split()[0] if re.match(POLISH_WORD, res["stara"]) else ""
            if res["typ"] == "zarowka_led":
                # "Oprawa SMART+ WiFi Spot R50" to lampa, a karta "SMART+ WIFI SPOT 40 R50 ..." bez "Lampa" trafiala
                # do downlightow (2026-10-21)
                lead = "Lampa"
            page_name = re.sub(r"\bCLAS\b", "CLASSIC", site["name"])
            # ten sam trzonek = ta sama lampa z literowka w Ekonomie ("PARATHOM PRO DIM AR111 1,7W" przy 11,7 W na karcie):
            # linia, marka z opakowania i sciemnianie z Ekonoma zostaja, karta ich nie powtarza ("PPRO AR111 75 24°")
            # - 2026-10-21: znikalo Parathom / Superior / Bellalux / antybakteryjna. Inny trzonek = inny produkt.
            old_u, page_u = parse_view(res["stara"]), parse_view(page_name)
            s_old = {s.upper() for s in SOCKET_RE.findall(old_u)}
            s_page = {s.upper() for s in SOCKET_RE.findall(page_u)}
            carry = ""
            _, p_old = parse_power(old_u)
            _, p_page = parse_power(page_u)
            # literowka w mocy: "0,5W" przy 10,5 W, "1,7W" przy 11,7 W, "4,9W" przy 24,9 W - albo moc bliska (6W / 5,5W);
            # "5W" przy 17 W to inny produkt (2026-10-21)
            typo = (not p_old or not p_page or max(p_old, p_page) / min(p_old, p_page) <= 1.25
                    or fmt_num(p_page).endswith(fmt_num(p_old)))
            other = gtin_other_product(old_u, page_u, s_old, s_page, p_old, p_page, rules) \
                if rules["distinguish"].get("gtin_other_product") else ""
            if other:
                # karta opisuje fizycznie inny produkt (inny trzonek / ksztalt / dlugosc tuby) - nie budujemy nazwy z karty,
                # zostaje nazwa z Ekonoma do sprawdzenia (decyzja uzytkownika 2026-10-21)
                res["uwagi"].append(f"karta producenta {site['url']} opisuje inny produkt ({other}: {site['name']})"
                                    " - nazwa z Ekonoma do sprawdzenia")
                res["status"] = "DO_SPRAWDZENIA"
                results.append(res)
                continue
            if typo and not (s_old and s_page and not s_old & s_page):
                carry = " ".join(w for w in dict.fromkeys(re.findall(
                    r"PARATHOM(?: PRO)?|SUPERIOR|SUPERSTAR(?: PLUS)?|ANTYBAKTERYJN\w*|BELLALUX|OSRAM|"
                    r"(?<!NON-)(?<!NON\.-)(?<!NON )\bDIM\b|\bPCR\b|"
                    # ksztalt i barwa Smart+ ("SMART+ ZB PAR16 TW 5W" przy karcie "SMART+ ZB SPOT 4.9 W ... GU10")
                    r"\bPAR1[16]\b|\bR(?:50|63|80)\b|\bTW\b|\bRGBW\b|"
                    # rozmiar banki Vintage: "CL GLOBE200" przy karcie "Vintage 1906 Big Globe 28" (2026-10-21)
                    r"\bGLOBE ?(?:45|60|80|95|100|125|130|150|165|183|200|220)\b", old_u))
                    if not re.search(rf"\b{re.escape(w)}\b", page_u))
            fixed = {**row, "Nazwa towaru": f"{lead} {page_name} {carry}".strip(), "Nazwa dodatkowa": ""}
            new = build_name(fixed, rules, pdf_lm, feed)
            new.update(row=row, stara=res["stara"], zrodlo_stara=site["name"])
            new["uwagi"].append(f"nazwa z karty producenta {site['url']} ({web_mismatch(res['stara'], site['title'])};"
                                f" w Ekonomie: {res['stara']})")
            new["status"] = "WARNING" if new["status"] == "OK" else new["status"]
            if new["typ"] == "brak reguly" and res["typ"] != "brak reguly":
                # nazwa karty nie pasuje do zadnego typu ("LED TUBE T5 AC HE14 ...") - zostaje nazwa z Ekonoma,
                # rozbieznosc z karta tylko w uwagach (2026-10-20)
                res["uwagi"].append(f"karta producenta {site['url']} przeczy nazwie "
                                    f"({web_mismatch(res['stara'], site['title'])}: {site['name']}) - do sprawdzenia")
                res["status"] = "WARNING" if res["status"] == "OK" else res["status"]
            else:
                res = new
        results.append(res)

    if rules.get("catalog_cct_check"):
        check_catalog_cct(results, pkg)
    apply_package_brand(results, rules, pkg)
    propagate_family_brand(results, rules)
    strip_brand_from_names(results, rules)  # jedna marka albo bez marki - przed wyroznikami
    add_version_suffixes(results, rules)
    move_brand_tail(results, rules)  # wersja przed marka, zanim policzymy multipaki i wyrozniki
    apply_pack_size(results, rules, pkg, web)  # po wersjach - "op. 3 szt." zawsze na koncu nazwy
    add_all_distinguishers(results, rules, pkg, web)  # wyroznik przed "op. N szt."
    move_brand_tail(results, rules)  # marka + linia na koncu, za wersja i wyroznikiem
    if rules.get("diameter_mm"):
        # srednica jako "220 mm" zamiast "fi220" (decyzja uzytkownika 2026-10-19); "fi30cm" -> "300 mm"
        for res in results:
            res["nazwa"] = diameter_in_mm(res["nazwa"])
    for item in rules.get("name_replace", []):
        # koncowe zamiany slow ("Oprawa świetlówkowa" -> "Oprawa pod tuby", 2026-10-19)
        for res in results:
            res["nazwa"] = re.sub(item["match"], item["repl"], res["nazwa"])
    if rules.get("dedupe_adjacent"):
        # to samo slowo dwa razy obok siebie: "globe Globe 95", "Haloline Haloline 1500", "SENSOR SENSOR" z Ekonoma
        for res in results:
            res["nazwa"] = re.sub(r"(?i)(?<!\S)(\S+) (?=\1(?!\S))", "", res["nazwa"])
    if rules.get("onoff_twins"):
        drop_lone_onoff(results)

    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for res in results:
        groups[res["nazwa"]].append(res)
    dup_rows = []
    for gid, (name, members) in enumerate(sorted(((n, m) for n, m in groups.items() if len(m) > 1), key=lambda x: (-len(x[1]), x[0])), 1):
        olds = [m["stara"] for m in members]
        diffs = diff_tokens(olds)
        same_old = len(set(olds)) == 1
        cats = ["identyczna nazwa w Ekonomie"] if same_old else (diff_categories(diffs) or ["inne"])
        for m, d in zip(members, diffs):
            m["powtorzenie"] = gid
            if m["status"] == "OK":
                m["status"] = "WARNING"
            dup_rows.append({
                "Grupa": gid, "Ile w grupie": len(members), "Nowa nazwa": name, "Typ": m["typ"],
                "Kod towaru": m["row"].get("Kod towaru", ""), "EAN": m["row"].get("Kod producenta", ""),
                "Stara nazwa": m["stara"], "Czym rozni sie stara nazwa": " ".join(sorted(d)),
                "Kategorie roznic w grupie": ", ".join(cats),
            })

    out_rows = []
    for res in results:
        row = dict(res["row"])
        row["Nazwa towaru"], row["Nazwa dodatkowa"] = split_for_ekonom(res["nazwa"])
        out_rows.append(row)
    ensure_dir(output.parent)
    write_ekonom_csv(output, fields, out_rows)

    write_report(report, results, dup_rows)
    stats = Counter(r["status"] for r in results)
    missing = Counter(b for r in results for b in r["braki"])
    print(f"Produkty: {len(results)}  statusy: {dict(stats)}")
    print(f"Braki: {dict(missing)}")
    print(f"Grupy powtorzen: {len({d['Grupa'] for d in dup_rows})} ({len(dup_rows)} produktow)")
    print(f"CSV Ekonom: {output}")
    print(f"Raport: {report}")


def write_report(path: Path, results: list[dict[str, Any]], dup_rows: list[dict[str, Any]]) -> None:
    import pandas as pd

    ensure_dir(path.parent)
    names = pd.DataFrame([{
        "LP": r["row"].get("LP", ""), "Kod towaru": r["row"].get("Kod towaru", ""),
        "EAN": r["row"].get("Kod producenta", ""), "Stara nazwa": r["stara"], "Nowa nazwa": r["nazwa"],
        "Dlugosc": len(r["nazwa"]), "Typ": r["typ"], "Status": r["status"],
        "Braki": ", ".join(r["braki"]), "Zrodlo lm": r["lm_zrodlo"],
        "Grupa powtorzen": r.get("powtorzenie", ""), "Uwagi": "; ".join(r["uwagi"]),
    } for r in results])
    dups = pd.DataFrame(dup_rows)
    summary_rows = []
    for (typ, status), n in Counter((r["typ"], r["status"]) for r in results).most_common():
        summary_rows.append({"Typ": typ, "Status": status, "Sztuk": n})
    cat_rows = []
    if not dups.empty:
        per_group = dups.drop_duplicates("Grupa")
        for _, g in per_group.iterrows():
            for cat in g["Kategorie roznic w grupie"].split(", "):
                cat_rows.append({"Typ": g["Typ"], "Kategoria roznicy": cat, "Grup": 1, "Produktow": g["Ile w grupie"]})
    cats = pd.DataFrame(cat_rows)
    if not cats.empty:
        cats = cats.groupby(["Typ", "Kategoria roznicy"], as_index=False).sum().sort_values(["Produktow"], ascending=False)
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        names.to_excel(xl, sheet_name="Nazwy", index=False)
        dups.to_excel(xl, sheet_name="Powtorzenia", index=False)
        cats.to_excel(xl, sheet_name="Roznice w powtorzeniach", index=False)
        pd.DataFrame(summary_rows).to_excel(xl, sheet_name="Podsumowanie typow", index=False)
        for ws in xl.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                width = max(len(str(c.value or "")) for c in col[:300])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 70)


if __name__ == "__main__":
    main()
