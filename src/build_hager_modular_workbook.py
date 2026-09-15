"""Build a reviewed Hager modular-apparatus workbook without touching input data."""

from __future__ import annotations

import argparse
import copy
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_INPUT = Path("input/Aparatura-Modulowa-Hager.xlsx")
DEFAULT_OUTPUT = Path("output/Aparatura-Modulowa-Hager_nazwy_atrybuty.xlsx")

SOURCE_FILE = "23PL013_KATALOG ROZDZIAL ENERGII_2023.pdf"
PRODUCT_URL = "https://hager.com/pl/katalog/catalog/product/download/id/{code}"

TECHNICAL_HEADERS = [
    "Atrybut Produktu: Układ Biegunów",
    "Atrybut Produktu: Znamionowa Zdolność Zwarciowa Icn [kA]",
    "Atrybut Produktu: Graniczna Zdolność Zwarciowa Icu [kA]",
    "Atrybut Produktu: Warunkowy Prąd Zwarciowy Inc [kA]",
    "Atrybut Produktu: Klasa Ograniczenia Energii",
    "Atrybut Produktu: Prąd Różnicowy IΔn [mA]",
    "Atrybut Produktu: Typ Prądu Różnicowego",
    "Atrybut Produktu: Napięcie Cewki [V]",
    "Atrybut Produktu: Konfiguracja Styków",
    "Atrybut Produktu: Kategoria Użytkowania",
    "Atrybut Produktu: Typ SPD",
    "Atrybut Produktu: Technologia SPD",
    "Atrybut Produktu: Prąd Udarowy Iimp [kA, łącznie]",
    "Atrybut Produktu: Znamionowy Prąd Wyładowczy In [kA]",
    "Atrybut Produktu: Poziom Ochrony Up [kV]",
    "Atrybut Produktu: Maksymalne Napięcie Ciągłej Pracy Uc [V]",
    "Atrybut Produktu: Typ Sieci",
    "Atrybut Produktu: Styk Sygnalizacyjny",
    "Atrybut Produktu: Zdolność Zwarciowa Icc [kA]",
    "Atrybut Produktu: Rodzaj Wkładki",
    "Atrybut Produktu: Typ Szyny",
    "Atrybut Produktu: Przekrój Szyny [mm²]",
    "Atrybut Produktu: Kompatybilność",
    "Źródło",
    "Strona katalogu",
    "Adres źródła",
    "Pewność",
    "Status",
    "Uwagi",
]

ATTRIBUTE_SLUGS = {
    "Barwa Lampek": "pa_barwa-lampek",
    "Dane Producenta": "pa_dane-producenta",
    "Dodatkowe Funkcje": "pa_dodatkowe-funkcje",
    "Głośność [dB]": "pa_glosnosc-db",
    "Ilość Lampek": "pa_ilosc-lampek",
    "Producent": "pa_producent",
    "Rodzaj": "pa_rodzaj",
    "Rodzaj Akcesoria": "pa_rodzaj-akcesoria",
    "Symbol": "pa_symbol",
    "Charakterystyka Wyzwalania": "pa_charakterystyka-wyzwalania",
    "Liczba Biegunów": "pa_liczba-biegunow",
    "Liczba Modułów": "pa_liczba-modulow",
    "Napięcie [V]": "pa_napiecie-v",
    "Prąd Znamionowy [A]": "pa_prad-znamionowy-a",
    "Rozmiar": "pa_rozmiar",
    "Seria": "pa_seria",
    "Sposób Montażu": "pa_sposob-montazu",
    "Stopień Ochrony [IP]": "pa_stopien-ochrony-ip",
    "Typ": "pa_typ",
    "Typ Produktu": "pa_typ-produktu",
    "Typ Wyłącznika": "pa_typ-wylacznika",
    "Typ Zacisków": "pa_typ-zaciskow",
    "Uziemienie": "pa_uziemienie",
    "Zakres Temperatury Pracy": "pa_zakres-temperatury-pracy",
    "Długość": "pa_dlugosc",
    "Ilość Styków Zwiernych/rozwiernych": "pa_ilosc-stykow-zwiernych-rozwiernych",
    "Zastosowanie": "pa_zastosowanie",
    "Znamionowy Przekrój Poprzeczny": "pa_znamionowy-przekroj-poprzeczny",
}


def code_from_sku(sku: object) -> str:
    return str(sku or "").split("/", 1)[0].strip().upper()


def current_from_code(code: str) -> int:
    match = re.search(r"^[A-Z]{3}\d(\d{2,3})[A-Z]?$", code)
    if not match:
        raise ValueError(f"Nie można odczytać prądu znamionowego z kodu {code}")
    raw = match.group(1)
    return int(raw)


def base_record(code: str) -> dict:
    return {
        "code": code,
        "family": "",
        "title": "",
        "page": "",
        "url": PRODUCT_URL.format(code=code),
        "confidence": "wysoka",
        "status": "OK",
        "notes": "",
        "legacy": {
            "Producent": "Hager",
            "Symbol": code,
        },
        "technical": {},
    }


def make_record(code: str) -> dict:
    r = base_record(code)
    legacy = r["legacy"]
    tech = r["technical"]

    if code.startswith(("MBN", "MCN", "NDN")):
        series = code[:3]
        characteristic = {"MBN": "B", "MCN": "C", "NDN": "D"}[series]
        current = current_from_code(code)
        poles = int(code[3])
        layout = f"{poles}P"
        icn = 10 if series == "NDN" else 6
        r.update(
            family="Wyłącznik nadprądowy MCB",
            title=(
                f"Wyłącznik nadprądowy {characteristic}{current} {layout} "
                f"{icn}kA Hager {code}"
            ),
            page="1009" if series == "NDN" else ("996" if poles == 1 else "997")
            if series == "MBN"
            else ("999" if poles == 1 else "1000"),
        )
        legacy.update(
            {
                "Charakterystyka Wyzwalania": characteristic,
                "Liczba Biegunów": poles,
                "Liczba Modułów": poles,
                "Napięcie [V]": "230/400 V AC" if poles == 1 else "400 V AC",
                "Prąd Znamionowy [A]": current,
                "Rodzaj": "Wyłącznik nadprądowy",
                "Seria": series,
                "Sposób Montażu": "Szyna DIN 35 mm",
                "Stopień Ochrony [IP]": "IP2X (zaciski) / IP40 (w rozdzielnicy)",
                "Typ": "MCB",
                "Typ Produktu": "Wyłącznik nadprądowy",
                "Typ Wyłącznika": "Nadprądowy",
                "Typ Zacisków": "Klatkowe / Bi-Connect",
                "Zakres Temperatury Pracy": "-25...+60 °C",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": layout,
                "Atrybut Produktu: Znamionowa Zdolność Zwarciowa Icn [kA]": icn,
                "Atrybut Produktu: Klasa Ograniczenia Energii": 3,
            }
        )
        if series == "NDN":
            tech["Atrybut Produktu: Graniczna Zdolność Zwarciowa Icu [kA]"] = 15
        return r

    if code.startswith(("CDC", "CDA")):
        series = code[:3]
        current = int(code[4:6])
        poles = int(code[3])
        layout = "1P+N" if poles == 2 else "3P+N"
        rcd_type = "AC" if series == "CDC" else "A"
        r.update(
            family="Wyłącznik różnicowoprądowy RCCB",
            title=(
                f"Wyłącznik różnicowoprądowy RCCB {current}A 30mA typ {rcd_type} "
                f"{layout} 6kA Hager {code}"
            ),
            page="1056" if series == "CDC" else "1058",
        )
        legacy.update(
            {
                "Liczba Biegunów": poles,
                "Liczba Modułów": poles,
                "Napięcie [V]": "230/400 V AC" if poles == 2 else "400 V AC",
                "Prąd Znamionowy [A]": current,
                "Rodzaj": "Wyłącznik różnicowoprądowy",
                "Seria": series,
                "Sposób Montażu": "Szyna DIN 35 mm",
                "Stopień Ochrony [IP]": "IP2X (zaciski) / IP40 (w rozdzielnicy)",
                "Typ": "RCCB",
                "Typ Produktu": "Wyłącznik różnicowoprądowy",
                "Typ Wyłącznika": "Różnicowoprądowy",
                "Typ Zacisków": "Klatkowe / Bi-Connect",
                "Zakres Temperatury Pracy": "-25...+40 °C",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": layout,
                "Atrybut Produktu: Warunkowy Prąd Zwarciowy Inc [kA]": 6,
                "Atrybut Produktu: Prąd Różnicowy IΔn [mA]": 30,
                "Atrybut Produktu: Typ Prądu Różnicowego": rcd_type,
            }
        )
        if code == "CDA440D":
            r.update(
                page="464 (tabela strat mocy); odpowiednik bieżący CDA440J: 1058",
                confidence="średnia",
                status="DO_SPRAWDZENIA",
                notes=(
                    "Kod CDA440D występuje w katalogu tylko w tabeli technicznej. "
                    "W bieżącej tabeli produktowej widnieje CDA440J. Zachowano SKU CDA440D."
                ),
            )
        return r

    if code.startswith(("ADC", "ADA")):
        series = code[:3]
        rcd_type = "AC" if series == "ADC" else "A"
        current = int(code[4:6])
        characteristic = "B"
        r.update(
            family="Wyłącznik różnicowo-nadprądowy RCBO",
            title=(
                f"Wyłącznik różnicowo-nadprądowy RCBO {characteristic}{current} "
                f"30mA typ {rcd_type} 1P+N 6kA Hager {code}"
            ),
            page="1077" if series == "ADC" else "1078",
        )
        legacy.update(
            {
                "Charakterystyka Wyzwalania": characteristic,
                "Liczba Biegunów": 2,
                "Liczba Modułów": 2,
                "Napięcie [V]": "230 V AC",
                "Prąd Znamionowy [A]": current,
                "Rodzaj": "Wyłącznik różnicowo-nadprądowy",
                "Seria": series,
                "Sposób Montażu": "Szyna DIN 35 mm",
                "Stopień Ochrony [IP]": "IP2X (zaciski) / IP40 (w rozdzielnicy)",
                "Typ": "RCBO",
                "Typ Produktu": "Wyłącznik różnicowo-nadprądowy",
                "Typ Wyłącznika": "Różnicowo-nadprądowy",
                "Typ Zacisków": "Klatkowe / Bi-Connect",
                "Zakres Temperatury Pracy": "-25...+40 °C",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": "1P+N",
                "Atrybut Produktu: Znamionowa Zdolność Zwarciowa Icn [kA]": 6,
                "Atrybut Produktu: Prąd Różnicowy IΔn [mA]": 30,
                "Atrybut Produktu: Typ Prądu Różnicowego": rcd_type,
            }
        )
        return r

    if code.startswith("SBN"):
        poles = int(code[3])
        current = int(code[4:])
        modules = 1 if poles == 1 else (2 if current <= 32 else poles)
        voltage = "230 V AC" if poles == 1 else "400 V AC"
        r.update(
            family="Rozłącznik izolacyjny modułowy",
            title=(
                f"Rozłącznik izolacyjny modułowy {poles}P {current}A "
                f"{voltage} Hager {code}"
            ),
            page="1170",
        )
        legacy.update(
            {
                "Liczba Biegunów": poles,
                "Liczba Modułów": modules,
                "Napięcie [V]": voltage,
                "Prąd Znamionowy [A]": current,
                "Rodzaj": "Rozłącznik izolacyjny",
                "Seria": "SBN",
                "Sposób Montażu": "Szyna DIN 35 mm",
                "Stopień Ochrony [IP]": "IP20",
                "Typ": "Rozłącznik izolacyjny",
                "Typ Produktu": "Rozłącznik izolacyjny modułowy",
                "Typ Zacisków": "Klatkowe / Bi-Connect",
                "Zakres Temperatury Pracy": "-25...+50 °C",
                "Zastosowanie": "Rozłączanie i izolowanie obwodów",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": f"{poles}P",
                "Atrybut Produktu: Kategoria Użytkowania": "AC-22A",
            }
        )
        return r

    if code == "L71M":
        r.update(
            family="Rozłącznik bezpiecznikowy modułowy",
            title="Rozłącznik bezpiecznikowy modułowy D02 1P 63A 400V AC Hager L71M",
            page="874; dane techniczne także 469",
        )
        legacy.update(
            {
                "Liczba Biegunów": 1,
                "Napięcie [V]": "400 V AC / 110-220 V DC",
                "Prąd Znamionowy [A]": 63,
                "Rodzaj": "Rozłącznik bezpiecznikowy",
                "Rozmiar": "D02",
                "Seria": "L71M",
                "Sposób Montażu": "Szyna DIN 35 mm",
                "Typ": "D02",
                "Typ Produktu": "Rozłącznik bezpiecznikowy modułowy",
                "Typ Zacisków": "Klatkowe",
                "Zakres Temperatury Pracy": "-25...+40 °C",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": "1P",
                "Atrybut Produktu: Zdolność Zwarciowa Icc [kA]": 50,
                "Atrybut Produktu: Rodzaj Wkładki": "D02",
            }
        )
        return r

    if code in {"ESC225", "ESC425"}:
        contacts = "2NO" if code == "ESC225" else "4NO"
        poles = 2 if code == "ESC225" else 4
        modules = 1 if code == "ESC225" else 2
        r.update(
            family="Stycznik modułowy",
            title=f"Stycznik modułowy 25A 230V AC {contacts} AC-7a/b Hager {code}",
            page="1185",
        )
        legacy.update(
            {
                "Ilość Styków Zwiernych/rozwiernych": contacts,
                "Liczba Biegunów": poles,
                "Liczba Modułów": modules,
                "Napięcie [V]": "230 V AC",
                "Prąd Znamionowy [A]": 25,
                "Rodzaj": "Stycznik modułowy",
                "Seria": "ESC",
                "Sposób Montażu": "Szyna DIN 35 mm",
                "Typ": "Stycznik instalacyjny",
                "Typ Produktu": "Stycznik modułowy",
                "Typ Zacisków": "Śrubowe",
                "Zakres Temperatury Pracy": "-10...+50 °C",
                "Zastosowanie": "Sterowanie odbiornikami w kategorii AC-7a/AC-7b",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": f"{poles}P",
                "Atrybut Produktu: Napięcie Cewki [V]": 230,
                "Atrybut Produktu: Konfiguracja Styków": contacts,
                "Atrybut Produktu: Kategoria Użytkowania": "AC-7a / AC-7b",
            }
        )
        return r

    if code in {"SPA931", "SPA400"}:
        if code == "SPA931":
            title = (
                "Ogranicznik przepięć SPD T1+T2 4P TN-S/TT Iimp 50kA "
                "In 50kA Up≤1,2kV Hager SPA931"
            )
            values = {
                "poles": 4,
                "type": "T1+T2",
                "technology": "MOV",
                "iimp": 50,
                "in": 50,
                "up": 1.2,
                "network": "TN-S / TT",
                "contact": "Tak",
                "page": "1103",
            }
        else:
            title = (
                "Ogranicznik przepięć SPD T1 3P TN-C Iimp 37,5kA "
                "Up≤1,5kV Hager SPA400"
            )
            values = {
                "poles": 3,
                "type": "T1",
                "technology": "Kombinowany, iskiernikowy",
                "iimp": 37.5,
                "in": "",
                "up": 1.5,
                "network": "TN-C",
                "contact": "Nie",
                "page": "1102",
            }
        r.update(family="Ogranicznik przepięć SPD", title=title, page=values["page"])
        legacy.update(
            {
                "Liczba Biegunów": values["poles"],
                "Liczba Modułów": 4,
                "Napięcie [V]": "230/400 V AC",
                "Rodzaj": "Ogranicznik przepięć",
                "Seria": "SPA",
                "Sposób Montażu": "Szyna DIN 35 mm",
                "Stopień Ochrony [IP]": "IP20",
                "Typ": "SPD",
                "Typ Produktu": "Ogranicznik przepięć",
                "Typ Zacisków": "Śrubowe",
                "Zakres Temperatury Pracy": "-40...+80 °C",
                "Zastosowanie": f"Ochrona instalacji {values['network']} przed przepięciami",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": f"{values['poles']}P",
                "Atrybut Produktu: Typ SPD": values["type"],
                "Atrybut Produktu: Technologia SPD": values["technology"],
                "Atrybut Produktu: Prąd Udarowy Iimp [kA, łącznie]": values["iimp"],
                "Atrybut Produktu: Znamionowy Prąd Wyładowczy In [kA]": values["in"],
                "Atrybut Produktu: Poziom Ochrony Up [kV]": values["up"],
                "Atrybut Produktu: Maksymalne Napięcie Ciągłej Pracy Uc [V]": 255,
                "Atrybut Produktu: Typ Sieci": values["network"],
                "Atrybut Produktu: Styk Sygnalizacyjny": values["contact"],
            }
        )
        return r

    if code in {"KDN163A", "KDN363A", "KDN363F"}:
        title_layout = "1P" if code == "KDN163A" else ("3P" if code == "KDN363A" else "4P/3P")
        layout = "1P" if code == "KDN163A" else "3P"
        poles = 1 if code == "KDN163A" else 3
        voltage = "415 V AC" if code != "KDN363F" else "230-400 V AC"
        compatibility = {
            "KDN163A": "Aparatura modułowa 1P z zaciskiem widełkowym",
            "KDN363A": "Aparatura modułowa 3P z zaciskiem widełkowym",
            "KDN363F": "RCCB 4P + MCB lub SPD 4P + MCB 3P",
        }[code]
        r.update(
            family="Szyna grzebieniowa",
            title=(
                f"Szyna grzebieniowa widełkowa {title_layout} 10mm² 63A "
                f"12 modułów Hager {code}"
            ),
            page="1018" if code == "KDN163A" else "1019",
        )
        legacy.update(
            {
                "Długość": "210 mm",
                "Liczba Biegunów": poles,
                "Liczba Modułów": 12,
                "Napięcie [V]": voltage,
                "Prąd Znamionowy [A]": 63,
                "Rodzaj": "Szyna grzebieniowa",
                "Seria": "KDN",
                "Sposób Montażu": "W zaciskach aparatów modułowych",
                "Typ": "Szyna fazowa grzebieniowa",
                "Typ Produktu": "Szyna grzebieniowa",
                "Typ Zacisków": "Widełkowe",
                "Zastosowanie": compatibility,
                "Znamionowy Przekrój Poprzeczny": "10 mm²",
            }
        )
        tech.update(
            {
                "Atrybut Produktu: Układ Biegunów": layout,
                "Atrybut Produktu: Typ Szyny": "Grzebieniowa, widełkowa",
                "Atrybut Produktu: Przekrój Szyny [mm²]": 10,
                "Atrybut Produktu: Kompatybilność": compatibility,
            }
        )
        return r

    raise ValueError(f"Brak reguły dla kodu: {code}")


def php_attribute_serialization(headers: dict[str, int], ws, row: int) -> str:
    keys = []
    for attribute, slug in ATTRIBUTE_SLUGS.items():
        column = headers.get(f"Atrybut Produktu: {attribute}")
        if column and ws.cell(row, column).value not in (None, ""):
            keys.append((slug, 0 if attribute == "Symbol" else 1))
    chunks = [f"a:{len(keys)}:{{"]
    for position, (slug, is_variation) in enumerate(keys):
        chunks.append(
            f's:{len(slug)}:"{slug}";a:6:{{'
            f's:4:"name";s:{len(slug)}:"{slug}";'
            's:5:"value";s:0:"";'
            f's:8:"position";i:{position};'
            's:10:"is_visible";i:1;'
            f's:12:"is_variation";i:{is_variation};'
            's:11:"is_taxonomy";i:1;}'
        )
    chunks.append("}")
    return "".join(chunks)


def append_headers(ws) -> dict[str, int]:
    headers = {str(cell.value): cell.column for cell in ws[1] if cell.value is not None}
    template = ws.cell(1, ws.max_column)
    for name in TECHNICAL_HEADERS:
        if name in headers:
            continue
        col = ws.max_column + 1
        cell = ws.cell(1, col, name)
        cell.font = copy.copy(template.font)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.border = copy.copy(template.border)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.number_format = copy.copy(template.number_format)
        cell.protection = copy.copy(template.protection)
        headers[name] = col
        ws.column_dimensions[get_column_letter(col)].width = 18
    return headers


def make_control_sheet(wb, records: list[tuple[int, str, str, dict]]) -> None:
    if "Kontrola" in wb.sheetnames:
        del wb["Kontrola"]
    ws = wb.create_sheet("Kontrola")
    ws.sheet_view.showGridLines = False

    dark_blue = "17365D"
    light_blue = "DCE6F1"
    pale_blue = "EEF5FB"
    green = "E2F0D9"
    yellow = "FFF2CC"
    red = "F4CCCC"
    thin_gray = Side(style="thin", color="D9E1F2")

    ws.merge_cells("A1:R1")
    ws["A1"] = "Kontrola nazw i atrybutów — aparatura modułowa Hager"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=dark_blue)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26

    counts = Counter(record[3]["status"] for record in records)
    ws.merge_cells("A2:R2")
    ws["A2"] = (
        f"Produkty: {len(records)} | OK: {counts.get('OK', 0)} | "
        f"Do sprawdzenia: {counts.get('DO_SPRAWDZENIA', 0)} | "
        "Plik wejściowy nie został nadpisany."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=light_blue)
    ws["A2"].font = Font(italic=True, color="17365D")
    ws["A2"].alignment = Alignment(vertical="center")

    ws.merge_cells("A3:R3")
    ws["A3"] = (
        "Źródło podstawowe: polski katalog Hager „Rozdział energii 2023” "
        f"({SOURCE_FILE}) oraz karty produktów Hager PL."
    )
    ws["A3"].font = Font(size=9, color="666666")

    headers = [
        "Lp.",
        "SKU",
        "Kod",
        "Nazwa pierwotna",
        "Nazwa po zmianie",
        "Rodzina",
        "Układ biegunów",
        "Prąd [A]",
        "Napięcie",
        "Moduły",
        "Charakterystyka / typ",
        "Kluczowe parametry",
        "Źródło",
        "Strona katalogu",
        "Adres źródła",
        "Pewność",
        "Status",
        "Uwagi",
    ]
    header_row = 5
    for col, value in enumerate(headers, 1):
        cell = ws.cell(header_row, col, value)
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9EADBF"))
    ws.row_dimensions[header_row].height = 34

    for out_row, (_, sku, old_title, record) in enumerate(records, header_row + 1):
        legacy = record["legacy"]
        tech = record["technical"]
        characteristic = legacy.get("Charakterystyka Wyzwalania") or tech.get(
            "Atrybut Produktu: Typ Prądu Różnicowego"
        ) or tech.get("Atrybut Produktu: Typ SPD") or legacy.get("Typ", "")
        key_params = []
        for label, key, suffix in [
            ("Icn", "Atrybut Produktu: Znamionowa Zdolność Zwarciowa Icn [kA]", " kA"),
            ("Icu", "Atrybut Produktu: Graniczna Zdolność Zwarciowa Icu [kA]", " kA"),
            ("Inc", "Atrybut Produktu: Warunkowy Prąd Zwarciowy Inc [kA]", " kA"),
            ("IΔn", "Atrybut Produktu: Prąd Różnicowy IΔn [mA]", " mA"),
            ("Iimp", "Atrybut Produktu: Prąd Udarowy Iimp [kA, łącznie]", " kA"),
            ("Up", "Atrybut Produktu: Poziom Ochrony Up [kV]", " kV"),
        ]:
            value = tech.get(key)
            if value not in (None, ""):
                key_params.append(f"{label} {str(value).replace('.', ',')}{suffix}")
        row_values = [
            out_row - header_row,
            sku,
            record["code"],
            old_title,
            record["title"],
            record["family"],
            tech.get("Atrybut Produktu: Układ Biegunów", ""),
            legacy.get("Prąd Znamionowy [A]", ""),
            legacy.get("Napięcie [V]", ""),
            legacy.get("Liczba Modułów", ""),
            characteristic,
            "; ".join(key_params),
            f"Hager, katalog Rozdział energii 2023; karta produktu Hager PL",
            record["page"],
            record["url"],
            record["confidence"],
            record["status"],
            record["notes"],
        ]
        for col, value in enumerate(row_values, 1):
            cell = ws.cell(out_row, col, value)
            if out_row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=pale_blue)
            cell.border = Border(bottom=thin_gray)
            cell.alignment = Alignment(vertical="top", wrap_text=col in {4, 5, 12, 13, 18})
        url_cell = ws.cell(out_row, 15)
        url_cell.hyperlink = record["url"]
        url_cell.style = "Hyperlink"

    last_row = header_row + len(records)
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:R{last_row}"
    widths = [7, 19, 14, 43, 49, 32, 15, 11, 20, 10, 20, 31, 35, 22, 38, 12, 19, 55]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.conditional_formatting.add(
        f"Q6:Q{last_row}",
        FormulaRule(formula=['Q6="OK"'], fill=PatternFill("solid", fgColor=green)),
    )
    ws.conditional_formatting.add(
        f"Q6:Q{last_row}",
        FormulaRule(formula=['Q6="DO_SPRAWDZENIA"'], fill=PatternFill("solid", fgColor=red)),
    )
    ws.conditional_formatting.add(
        f"P6:P{last_row}",
        FormulaRule(formula=['P6="średnia"'], fill=PatternFill("solid", fgColor=yellow)),
    )


def validate_records(records: list[tuple[int, str, str, dict]]) -> None:
    assert len(records) == 96, f"Oczekiwano 96 produktów, jest {len(records)}"
    skus = [sku for _, sku, _, _ in records]
    titles = [record["title"] for _, _, _, record in records]
    assert len(set(skus)) == len(skus), "SKU nie są unikalne"
    assert len(set(titles)) == len(titles), "Nowe nazwy nie są unikalne"
    for _, sku, _, record in records:
        assert record["code"] in record["title"], f"Brak kodu w nazwie: {sku}"
        assert "Hager" in record["title"], f"Brak marki w nazwie: {sku}"
        assert record["page"], f"Brak strony źródłowej: {sku}"
        assert record["status"] in {"OK", "DO_SPRAWDZENIA"}
    statuses = Counter(record["status"] for _, _, _, record in records)
    assert statuses == Counter({"OK": 95, "DO_SPRAWDZENIA": 1}), statuses
    review_codes = [record["code"] for _, _, _, record in records if record["status"] != "OK"]
    assert review_codes == ["CDA440D"], review_codes


def build(input_path: Path, output_path: Path) -> None:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("Plik wynikowy nie może nadpisywać pliku wejściowego")
    wb = load_workbook(input_path)
    ws = wb["Worksheet"]
    headers = append_headers(ws)
    required = ["Title", "SKU", "Product Attributes"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise KeyError(f"Brak wymaganych kolumn: {', '.join(missing)}")

    records: list[tuple[int, str, str, dict]] = []
    for row in range(2, ws.max_row + 1):
        sku = str(ws.cell(row, headers["SKU"]).value or "").strip()
        if not sku:
            continue
        code = code_from_sku(sku)
        old_title = str(ws.cell(row, headers["Title"]).value or "")
        record = make_record(code)
        ws.cell(row, headers["Title"], record["title"])

        for attribute, value in record["legacy"].items():
            header = f"Atrybut Produktu: {attribute}"
            if header in headers and value not in (None, ""):
                ws.cell(row, headers[header], value)
        for header, value in record["technical"].items():
            if value not in (None, ""):
                ws.cell(row, headers[header], value)

        source = f"Hager, katalog Rozdział energii 2023 ({SOURCE_FILE}); karta produktu Hager PL"
        ws.cell(row, headers["Źródło"], source)
        ws.cell(row, headers["Strona katalogu"], record["page"])
        url_cell = ws.cell(row, headers["Adres źródła"], record["url"])
        url_cell.hyperlink = record["url"]
        url_cell.style = "Hyperlink"
        ws.cell(row, headers["Pewność"], record["confidence"])
        ws.cell(row, headers["Status"], record["status"])
        ws.cell(row, headers["Uwagi"], record["notes"])

        ws.cell(
            row,
            headers["Product Attributes"],
            php_attribute_serialization(headers, ws, row),
        )
        records.append((row, sku, old_title, record))

    validate_records(records)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 15, 32)
    ws.sheet_view.showGridLines = True
    for row, _, _, record in records:
        status_cell = ws.cell(row, headers["Status"])
        status_cell.fill = PatternFill(
            "solid", fgColor="E2F0D9" if record["status"] == "OK" else "F4CCCC"
        )

    make_control_sheet(wb, records)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    counts = Counter(record[3]["family"] for record in records)
    print(f"Zapisano: {output_path}")
    print(f"Produkty: {len(records)}")
    print(f"Rodziny: {dict(counts)}")
    print("Statusy: OK=95, DO_SPRAWDZENIA=1 (CDA440D)")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build(args.input, args.output)
