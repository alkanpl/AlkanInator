from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import fitz
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "input" / "Hager-Berker-Gniazdka.xlsx"
DEFAULT_PDF = ROOT / "input" / "21PL001_Katalog_Osprzet_Hager_2021.pdf"
DEFAULT_OUTPUT = ROOT / "output" / "Hager-Berker-Gniazdka_WARIANTY_ATRYBUTY_I_TAGI.xlsx"
DEFAULT_REPORT = ROOT / "reports" / "Hager-Berker-Gniazdka_WARIANTY_ATRYBUTY_I_TAGI_report.json"
DEFAULT_REVIEW_SOURCE = ROOT / "output" / "Hager-Berker-Gniazdka_WARIANTY_BEZ_STARYCH_DUPLIKATOW.xlsx"

BLUE = "1F4E78"
MID_BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="D9E1F2")

MOJIBAKE_MARKERS = (
    "Ã³", "Ä…", "Ä‡", "Ä™", "Å‚", "Å", "Å„", "Å›", "Åš", "Åº", "Å¹", "Å¼", "Å»",
    "Ĺ‚", "Ĺ›", "Ĺš", "Ĺş", "Ĺź", "ĹĽ", "Ĺ»",
    "â€“", "â€”", "â€ž", "â€ť", "�",
)

SPECIAL_TRANSLATION = str.maketrans(
    {
        "ł": "l",
        "Ł": "L",
        "đ": "d",
        "Đ": "D",
        "ø": "o",
        "Ø": "O",
    }
)

COLOR_WORDS = {
    "bialy", "biala", "biale", "czarny", "czarna", "czarne", "szary", "szara", "szare",
    "srebrny", "srebrna", "srebrne", "bezowy", "bezowa", "bezowe", "kremowy", "kremowa",
    "kremowe", "antracyt", "antracytowy", "antracytowa", "grafitowy", "grafitowa", "grafitowe",
    "brazowy", "brazowa", "brazowe", "czerwony", "czerwona", "czerwone", "zielony", "zielona",
    "zielone", "zolty", "zolta", "zolte", "pomaranczowy", "pomaranczowa", "pomaranczowe", "zloty", "zlota", "zlote",
    "transparentny", "transparentna", "transparentne", "przezroczysty", "przezroczysta", "przezroczyste",
    "jasny", "jasna", "jasne", "polarnej", "bieli", "alu",
}
MATERIAL_WORDS = {
    "aluminium", "aluminiowy", "aluminiowa", "aluminiowe", "szklo", "szklany", "szklana",
    "szklane", "szkliwiony", "szkliwiona", "metal", "metalowy", "metalowa", "metalowe", "stal",
    "stalowy", "stalowa", "stalowe", "szlachetny", "szlachetna", "nierdzewny", "nierdzewna",
    "mosiadz", "mosiezny", "mosiezna", "porcelana", "porcelanowy", "porcelanowa", "beton",
    "betonowy", "betonowa", "lupek", "lupkowy", "lupkowa", "drewno", "drewniany", "drewniana",
    "skora", "skorzany", "skorzana", "kamien", "kamienny", "kamienna", "akryl", "buk", "dab",
    "termoplastyczne", "tworzywo", "tworzywa",
}
FINISH_WORDS = {
    "mat", "matowy", "matowa", "matowe", "polysk", "blyszczacy", "blyszczaca", "blyszczace",
    "lakierowany", "lakierowana", "lakierowane", "aksamit", "aksamitny", "aksamitna", "satyna",
    "satynowy", "satynowa", "chrom", "chromowany", "chromowana", "nikiel", "niklowany", "niklowana",
}
AXIS_WORDS = COLOR_WORDS | MATERIAL_WORDS | FINISH_WORDS

CATALOG_OVERRIDES = {
    "10116020": {
        "material_detail": "beton strukturalny",
        "manufacturer_color": "Beton strukturalny",
        "catalog_note": "Katalog, strona drukowana 215: beton strukturalny. Tytuł źródłowy błędnie podaje łupek.",
    },
    "10116030": {
        "material_detail": "łupek naturalny",
        "manufacturer_color": "Łupek naturalny",
        "catalog_note": "Katalog, strona drukowana 215: łupek naturalny.",
    },
    "10126020": {
        "material_detail": "beton strukturalny",
        "manufacturer_color": "Beton strukturalny",
        "catalog_note": "Katalog, strona drukowana 215: beton strukturalny. Tytuł źródłowy błędnie podaje łupek.",
    },
    "10126030": {
        "material_detail": "łupek naturalny",
        "manufacturer_color": "Łupek naturalny",
        "catalog_note": "Katalog, strona drukowana 215: łupek naturalny.",
    },
    "10136020": {
        "material_detail": "beton strukturalny",
        "manufacturer_color": "Beton strukturalny",
        "catalog_note": "Katalog, strona drukowana 215: beton strukturalny. Tytuł źródłowy błędnie podaje łupek.",
    },
    "10136030": {
        "material_detail": "łupek naturalny",
        "manufacturer_color": "Łupek naturalny",
        "catalog_note": "Katalog, strona drukowana 215: łupek naturalny.",
    },
    "3315396089": {
        "color_detail": "biały aksamit",
        "manufacturer_color": "Biały aksamit",
        "catalog_note": "Katalog, strona drukowana 131: wykonanie białe aksamitne. Tytuł źródłowy błędnie podaje antracyt.",
    },
    "85145183": {
        "color_detail": "aluminium mat",
        "material_detail": "aluminium",
        "finish_detail": "mat",
        "manufacturer_color": "Aluminium mat, lakierowane",
        "catalog_note": "Katalog, strona drukowana 507: aluminium mat, lakierowane. Tytuł i atrybut koloru błędnie podają biały.",
    },
    "85145188": {
        "color_detail": "biały mat",
        "finish_detail": "mat",
        "manufacturer_color": "Biały mat",
        "catalog_note": "Katalog, strona drukowana 507: biały mat.",
    },
    "10122384": {
        "multiplicity": "2x",
        "catalog_note": "Tytuł źródłowy oznacza ramkę 2-krotną, mimo braku pełnego słowa „krotna”.",
    },
    "5316238992": {
        "multiplicity": "2x",
        "catalog_note": "Kod rodziny 531623... odpowiada klawiszowi podwójnemu; tytuł źródłowy pomija krotność.",
    },
}

SERIES_PATTERNS = (
    (r"\bb\s*kwadrat\b", "B. Kwadrat"),
    (r"\bb\s*1\b", "B.1"), (r"\bb\s*3\b", "B.3"), (r"\bb\s*7\b", "B.7"),
    (r"\bq\s*1\b", "Q.1"), (r"\bq\s*3\b", "Q.3"), (r"\bq\s*7\b", "Q.7"),
    (r"\bk\s*1\b", "K.1"), (r"\bk\s*5\b", "K.5"),
    (r"\br\s*classic\b", "R.classic"), (r"\br\s*x\b", "R.x"),
    (r"\br\s*1\b", "R.1"), (r"\br\s*3\b", "R.3"), (r"\br\s*8\b", "R.8"),
    (r"\bs\s*1\b", "S.1"), (r"\bw\s*1\b", "W.1"),
    (r"\bglasserie\b", "Glasserie"), (r"\bserie\s*1930\b", "Serie 1930"),
    (r"\bone\s*platform\b", "one.platform"), (r"\bintegro\s*flow\b", "Integro Flow"),
    (r"\blumina\b", "Lumina"),
)


def norm(value: Any) -> str:
    text = str(value or "").translate(SPECIAL_TRANSLATION)
    text = "".join(
        char for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    ).lower()
    text = text.replace("b/u", "bez uziemienia").replace("z/u", "z uziemieniem")
    return " ".join(re.findall(r"[a-z0-9]+", text))


def compact(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip(" ,;|/")


def mojibake_markers(value: Any) -> list[str]:
    text = str(value or "")
    return [marker for marker in MOJIBAKE_MARKERS if marker in text]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sku_base(value: Any) -> str:
    return compact(value).split("/", 1)[0].upper()


def split_tags(value: Any) -> list[str]:
    text = compact(value)
    if not text:
        return []
    return [compact(item) for item in re.split(r"\s*[,;|]\s*", text) if compact(item)]


def append_tag(value: Any, tag: str) -> str:
    tags = split_tags(value)
    known = {norm(item) for item in tags}
    if norm(tag) not in known:
        tags.append(tag)
    return ", ".join(tags)


def infer_primary_series(series_value: Any, title: Any) -> str:
    series = compact(series_value)
    if series:
        return compact(series.split("|", 1)[0])
    title_normalized = norm(title)
    for pattern, label in SERIES_PATTERNS:
        if re.search(pattern, title_normalized):
            return label
    return ""


def remove_series_phrases(text: str) -> str:
    result = text
    for pattern, _label in SERIES_PATTERNS:
        result = re.sub(pattern, " ", result)
    return compact(result)


def title_without_codes(title: Any, sku: Any) -> str:
    value = str(title or "")
    code = re.escape(sku_base(sku))
    if code:
        value = re.sub(rf"(?<![A-Za-z0-9]){code}(?![A-Za-z0-9])", " ", value, flags=re.I)
    value = re.sub(r"\(\s*Berker\s+[A-Z0-9]+\s*\)", " ", value, flags=re.I)
    value = re.sub(r"\b(?:Hager|Berker)\b", " ", value, flags=re.I)
    value = re.sub(r"\b\d{5,14}\b", " ", value)
    return compact(value)


def is_frame(title: Any, product_type: Any) -> bool:
    return "ramk" in norm(product_type) or norm(title).startswith("ramka ")


MEASURE_UNIT_TOKENS = {"m", "mm", "cm"}
COUNT_NOUN_PREFIXES = ("wyjsciow", "wejsciow", "strzal")


def title_core(title: Any, sku: Any, frame: bool) -> str:
    value = norm(title_without_codes(title, sku))
    value = remove_series_phrases(value)
    # norm() rozbija "1,1 m" na tokeny "1", "1", "m"; sklej je z powrotem, zeby
    # zasieg czujnika nie zostal potraktowany jak krotnosc i wyciety z rdzenia.
    value = re.sub(r"\b([0-9])\s+([0-9])\s+(mm|cm|m)\b", r"\1,\2 \3", value)
    tokens = value.split()
    result: list[str] = []
    for position, token in enumerate(tokens):
        if token in AXIS_WORDS:
            continue
        if token in {
            "krotna", "krotny", "krotne", "krotnego", "krotnych", "kr",
            "pionowa", "pionowy", "pionowe", "pozioma", "poziomy", "poziome",
        }:
            continue
        if re.fullmatch(r"[1-5]x?", token):
            following = tokens[position + 1] if position + 1 < len(tokens) else ""
            # cyfra przed jednostka miary lub przed "wyjsciowy/wejsciowy" to
            # parametr techniczny, a nie krotnosc - zostaje w rdzeniu
            if not (following in MEASURE_UNIT_TOKENS or following.startswith(COUNT_NOUN_PREFIXES)):
                continue
        result.append(token)
    replacements = {
        "gn": "gniazdo", "gniazda": "gniazdo", "plytki": "plytka", "klawisze": "klawisz",
        "ramki": "ramka", "pojedyncza": "pojedynczy", "pojedyncze": "pojedynczy",
        "podwojna": "podwojny", "podwojne": "podwojny", "potrojna": "potrojny",
        "potrojne": "potrojny", "poczworna": "poczworny", "poczworne": "poczworny",
    }
    result = [replacements.get(token, token) for token in result]
    return compact(" ".join(result))


def category_family(core: str, product_type: Any, frame: bool) -> str:
    if frame:
        return "ramki"
    first = core.split(" ", 1)[0] if core else ""
    mapping = {
        "klawisz": "klawisze", "plytka": "plytki", "element": "pokrywy", "nasadka": "pokrywy",
        "zaslepka": "pokrywy", "gniazdo": "gniazda", "zlacze": "gniazda", "lacznik": "sterowanie",
        "przycisk": "sterowanie", "sterownik": "sterowanie", "sciemniacz": "sterowanie",
        "czujnik": "sterowanie", "regulator": "sterowanie", "termostat": "sterowanie",
        "puszka": "montaz", "adapter": "montaz", "pierscien": "montaz", "ramka": "ramki",
    }
    return mapping.get(first, norm(product_type) or "inne")


def extracted_axis(title: Any, words: set[str]) -> str:
    tokens = norm(title).split()
    return " ".join(token for token in tokens if token in words)


def multiplicity_axis(title: Any, value: Any, frame: bool) -> str:
    if compact(value):
        return compact(value)
    text = norm(title)
    match = re.search(r"\b([1-5])\s*(?:x|kr\b|krotn)", text)
    if match:
        return f"{match.group(1)}x"
    word_map = {"pojedyncz": "1x", "podwojn": "2x", "potrojn": "3x", "poczworn": "4x", "pieciokrotn": "5x"}
    for prefix, label in word_map.items():
        if prefix in text:
            return label
    return ""


def orientation_axis(title: Any, frame: bool, multiplicity: str = "") -> str:
    if not frame:
        return ""
    text = norm(title)
    if "poziom" in text:
        return "Pozioma"
    if "pionow" in text:
        return "Pionowa"
    if multiplicity == "1x":
        return "Bez orientacji"
    return "Uniwersalna"


def label_from_title(title: Any, sku: Any, series: str, frame: bool) -> str:
    source = title_without_codes(title, sku)
    source = re.sub(r"\b(?:B\.?\s*Kwadrat|[BQKRSW]\.?\s*[13758]|R\.?\s*classic|R\.?\s*x|Glasserie|Serie\s*1930|one\.?platform|Integro\s*Flow|Lumina)\b", " ", source, flags=re.I)
    axis_pattern = r"\b(?:bia\w*|czarn\w*|szar\w*|srebr\w*|beż\w*|krem\w*|antracyt\w*|grafit\w*|brąz\w*|czerwon\w*|zielon\w*|pomarańcz\w*|złot\w*|transparent\w*|aluminium|aluminiow\w*|alu|szkł\w*|metal\w*|stal\w*|szlachetn\w*|nierdzewn\w*|mosiądz\w*|porcelan\w*|beton\w*|łupek\w*|drewn\w*|skór\w*|kamien\w*|akryl\w*|buk|dąb|mat\w*|połysk\w*|lakierowan\w*|aksamit\w*|satyn\w*|chrom\w*|nikl\w*)\b"
    source = re.sub(axis_pattern, " ", source, flags=re.I)
    source = re.sub(
        r"\b[1-5]\s*[- ]?(?:kr\b|krotn\w*)|\b(?:pojedyncz\w*|podwĂłjn\w*|potrĂłjn\w*|poczwĂłrn\w*|poziom\w*|pionow\w*)\b",
        " ", source, flags=re.I,
    )
    source = re.sub(
        r"\bkat(?:egoria)?\.?\s*[56](?:e|a)?\b|\b(?:s/?ftp|ftp|utp|stp)\b|"
        r"\b(?:bez|z) uziemieni\w*\b|\b(?:zaciski? )?(?:samozacisk\w*|śrubowo[- ]window\w*|śrubow\w*)\b|"
        r"\b\d+(?:[.,]\d+)?\s*(?:v|a)\b|\b[1-4]\s*[- ]?biegun\w*\b|"
        r"\b(?:bez|z) (?:przesłon\w*|podwyższon\w* ochron\w* styk\w*)\b",
        " ", source, flags=re.I,
    )
    source = compact(re.sub(r"\s*[/;,()]\s*", " ", source))
    if not source:
        source = "Ramka" if frame else "Produkt"
    return source[0].upper() + source[1:]


def alias_berker_code(title: Any) -> str:
    match = re.search(r"\(\s*Berker\s+([A-Z0-9]+)\s*\)", str(title or ""), flags=re.I)
    return match.group(1).upper() if match else ""


def infer_series_key(primary_series: str, title: str, code: str) -> str:
    """Return a grouping series key without changing the source series attribute."""
    title_normalized = norm(title)
    if code.startswith(("WDA4527", "WDA1527", "WDA4628", "WDA4528")):
        return "akcesoria uniwersalne"
    if "ip44" in title_normalized or code.startswith(("WWD", "WWS", "WWW", "WWA")):
        return "system ip44"
    if code.startswith(("475171", "475172", "475271", "475272")):
        return "k.1 k.5"
    if re.match(r"^(?:104|105)[123]70(?:03|04|06|09)$", code):
        return "k.1 k.5"
    if re.match(r"^(?:53104[123]99(?:02|09)|104[123]9949)$", code):
        return "b kwadrat shared box"
    if re.match(r"^104[123]89(?:35|89)$", code):
        return "s.1 r.3 shared box"
    if code.startswith(("8514", "85648", "85655", "85656")):
        suffix = code[-2:]
        if suffix in {"82", "83", "85", "88", "89"}:
            return "b s system"
        if suffix in {"73", "75", "77", "79"}:
            return "k.1 k.5"
        if suffix in {"24", "26", "29"}:
            return "q system"
        if suffix in {"31", "39"}:
            return "r system"
    if code.startswith("113572"):
        return "k.1 k.5"
    if code.startswith("104570"):
        return "k.1 k.5"
    if code in {"11371404", "11371606", "11371909"}:
        return "b s dimmer"
    if code.startswith("113789"):
        return "b s system"
    if code.startswith(("4123", "4143", "4723", "4743")):
        if primary_series in {"B. Kwadrat", "B.1", "B.3", "B.7", "S.1", "Glasserie"}:
            return "b s system"
        if primary_series in {"Q.1", "Q.3", "Q.7"}:
            return "q system"
        if primary_series in {"R.1", "R.3", "R.8"}:
            return "r.1 r.3 r.8"
    return norm(primary_series)


def apply_product_specifics(record: "Record") -> None:
    """Add catalog-confirmed construction details omitted or wrong in source titles."""
    code = record.code
    override = CATALOG_OVERRIDES.get(code, {})
    if override:
        record.color_detail = override.get("color_detail", record.color_detail)
        record.material_detail = override.get("material_detail", record.material_detail)
        record.finish_detail = override.get("finish_detail", record.finish_detail)
        record.multiplicity = override.get("multiplicity", record.multiplicity)
        if override.get("multiplicity"):
            record.orientation = orientation_axis(record.title, record.frame, record.multiplicity)
        record.catalog_note = override.get("catalog_note", "")

    if record.series_key == "q system" and code.startswith(("4123", "4143", "4723", "4743")):
        record.core = "gniazdo schuko"
    elif code.startswith("104570"):
        record.core = "zaslepka z plytka"
    elif code.startswith("471570"):
        record.core = "gniazdo schuko bez podwyzszonej ochrony stykow"
    elif code.startswith("473570"):
        record.core = "gniazdo schuko z podwyzszona ochrona stykow"

    def set_detail(feature: str, value: str, evidence: str, page: int | None = None) -> None:
        record.distinguishing_feature = feature
        record.distinguishing_value = value
        record.audit_evidence = evidence
        if page is not None:
            record.catalog_printed_page = str(page)
        record.construction_key = norm(f"{feature} {value}")

    if code == "16201404":
        set_detail("Soczewka", "bez soczewki", "Katalog: zwykły klawisz.", 29)
    elif code == "16211404":
        set_detail("Soczewka", "z soczewką", "Katalog: klawisz z soczewką.", 30)
    elif code.startswith("471570"):
        set_detail("Podwyższona ochrona styków", "nie", "Hager: wykonanie bez podwyższonej ochrony styków.")
    elif code.startswith("473570"):
        set_detail("Podwyższona ochrona styków", "tak", "Hager: wykonanie z podwyższoną ochroną styków.")
    elif code.startswith("4143"):
        set_detail("Przyłącze / ochrona styków", "zaciski śrubowo-windowe; bez przesłon styków", "Katalog rozdziela wykonania według rodzaju zacisków i przesłon.", 166)
    elif code.startswith("4123"):
        set_detail("Przyłącze / ochrona styków", "zaciski śrubowo-windowe; z przesłonami styków", "Katalog rozdziela wykonania według rodzaju zacisków i przesłon.", 167)
    elif code.startswith("4743"):
        set_detail("Przyłącze / ochrona styków", "samozaciski; bez przesłon styków", "Katalog rozdziela wykonania według rodzaju zacisków i przesłon.", 166)
    elif code.startswith("4723"):
        set_detail("Przyłącze / ochrona styków", "samozaciski; z przesłonami styków", "Katalog rozdziela wykonania według rodzaju zacisków i przesłon.", 166)
    elif code.startswith("475171"):
        set_detail("Pokrywa", "blokowana w pozycji otwartej; bez pola opisowego", "Katalog: pokrywa z blokadą w pozycji otwartej.", 277)
    elif code.startswith("475172"):
        set_detail("Pokrywa", "samozamykająca; bez pola opisowego", "Katalog: pokrywa samozamykająca, opcja IP44.", 277)
    elif code.startswith("475271"):
        set_detail("Pokrywa", "blokowana w pozycji otwartej; z polem opisowym", "Katalog: pokrywa z polem opisowym i blokadą w pozycji otwartej.", 277)
    elif code.startswith("475272"):
        set_detail("Pokrywa", "samozamykająca; z polem opisowym", "Katalog: pokrywa samozamykająca z polem opisowym, opcja IP44.", 278)
    elif code.startswith("104570"):
        set_detail("Sposób montażu", "bez cokołu i pazurków rozporowych", "Katalog: zaślepka bez pazurków rozporowych.", 299)
    elif code.startswith("67104570"):
        set_detail("Sposób montażu", "z cokołem i pazurkami", "Katalog: komplet z cokołem i pazurkami.", 299)
    elif code in {"11707003", "WLF1058WG", "WLF1058BG"}:
        set_detail("Wykonanie płytki", "pojedyncza", "Katalog: do pojedynczej płytki montażowej.", 114)
    elif code in {"11827003", "WLF1059WG", "WLF1059BG"}:
        set_detail("Wykonanie płytki", "podwójna", "Katalog: do podwójnej płytki montażowej.", 116)
    elif code == "WDE5101":
        set_detail("Zestyki zgłoszeniowe", "bez oddzielnych zestyków zgłoszeniowych", "Katalog: łącznik i sygnalizator E10, zestyk zwierny.", 56)
    elif code == "WDE510110":
        set_detail("Zestyki zgłoszeniowe", "z oddzielnymi zestykami zgłoszeniowymi", "Katalog wprost wskazuje oddzielne zestyki zgłoszeniowe.", 56)
    elif code.startswith("WL021"):
        set_detail("Rodzaj łącznika", "1-biegunowy z podświetleniem", "Katalog: 10 AX, lampka neonowa, bez nadruku 0/1.", 425)
    elif code.startswith("WL061"):
        set_detail("Rodzaj łącznika", "1-biegunowy z podświetleniem kontrolnym i nadrukiem 0/1", "Katalog: 16 AX, podświetlenie kontrolne, nadruk 0/1.", 425)
    elif code == "WDF459210":
        set_detail("Liczba wyjść", "RTV — 2 wyjścia", "Katalog: gniazdo końcowe RTV.", 119)
    elif code == "WDF459315":
        set_detail("Liczba wyjść", "RTV + 2×SAT — 4 wyjścia", "Katalog: dwa niezależne wyjścia SAT.", 119)
    elif code == "85145183":
        record.distinguishing_feature = "Kolor katalogowy"
        record.distinguishing_value = "aluminium mat"
        record.audit_evidence = "Katalog wskazuje aluminium mat; błąd w tytule i atrybucie źródłowym."
        record.catalog_printed_page = "507"
    elif code == "85145188":
        record.distinguishing_feature = "Kolor katalogowy"
        record.distinguishing_value = "biały mat"
        record.audit_evidence = "Katalog wskazuje biały mat."
        record.catalog_printed_page = "507"
    elif code == "3315396089":
        record.distinguishing_feature = "Kolor katalogowy"
        record.distinguishing_value = "biały aksamit"
        record.audit_evidence = "Katalog potwierdza białe wykonanie; tytuł źródłowy błędnie podaje antracyt."
        record.catalog_printed_page = "131"

    if record.construction_key:
        record.core = compact(f"{record.core} [{record.construction_key}]")


@dataclass
class Record:
    excel_row: int
    product_id: str
    title: str
    sku: str
    code: str
    product_type: str
    series: str
    primary_series: str
    color: str
    manufacturer_color_before: str
    manufacturer_color: str
    material: str
    multiplicity: str
    tags_before: str
    core: str
    family: str
    frame: bool
    color_detail: str
    material_detail: str
    finish_detail: str
    orientation: str
    network_category: str
    screened: str
    grounding: str
    terminal_type: str
    voltage: str
    nominal_current: str
    pole_count: str
    shutters: str
    detection_range: str
    output_count: str
    alias_code: str
    series_key: str = ""
    construction_key: str = ""
    distinguishing_feature: str = ""
    distinguishing_value: str = ""
    audit_evidence: str = ""
    catalog_printed_page: str = ""
    catalog_note: str = ""
    is_old_duplicate: bool = False
    replacement_code: str = ""
    catalog_match: str = ""
    catalog_reference: str = ""
    catalog_pages: str = ""
    variant_group_id: str = ""
    variant_group_key: str = ""
    variant_group_name: str = ""
    variant_group_size: int = 0
    variant_tag: str = ""
    variant_axes: str = ""
    variant_values: str = ""
    decision: str = "BRAK GRUPY"
    reason: str = "brak drugiego produktu o tej samej funkcji i serii"
    review_reasons: list[str] = field(default_factory=list)

    @property
    def axis_signature(self) -> tuple[str, ...]:
        return (
            norm(self.manufacturer_color),
            norm(self.multiplicity),
            norm(self.orientation),
            norm(self.network_category),
            norm(self.screened),
            norm(self.grounding),
            norm(self.terminal_type),
            norm(self.voltage),
            norm(self.nominal_current),
            norm(self.pole_count),
            norm(self.shutters),
            norm(self.detection_range),
            norm(self.output_count),
        )


AXIS_SAFE_CONSTRUCTION_FEATURES = {
    "wykonanie plytki",
    "podwyzszona ochrona stykow",
    "przylacze ochrona stykow",
    "liczba wyjsc",
}

VARIANT_AXIS_FIELDS = (
    ("Kolor producenta", "manufacturer_color"),
    ("Krotność", "multiplicity"),
    ("Orientacja", "orientation"),
    ("Kategoria kabli sieciowych", "network_category"),
    ("Ekranowany", "screened"),
    ("Uziemienie", "grounding"),
    ("Typ zacisków", "terminal_type"),
    ("Napięcie [V]", "voltage"),
    ("Prąd znamionowy [A]", "nominal_current"),
    ("Liczba biegunów", "pole_count"),
    ("Przesłony torów prądowych", "shutters"),
    ("Zasięg detekcji", "detection_range"),
    ("Liczba wyjść", "output_count"),
)
VARIANT_AXIS_FIELD_BY_LABEL = dict(VARIANT_AXIS_FIELDS)


def infer_variant_axes(record: Record) -> None:
    text = norm(f"{record.title} {record.distinguishing_value}")
    record.multiplicity = multiplicity_axis(
        f"{record.title} {record.distinguishing_value}", record.multiplicity, record.frame,
    )

    if record.frame:
        record.orientation = orientation_axis(record.title, True, record.multiplicity)
    elif "puszka" in text and ("natyn" in text or " n t " in f" {text} "):
        if record.code.startswith("104") and record.series in {"K.1", "K.5"} and record.multiplicity in {"2x", "3x"}:
            record.orientation = "Pionowa"
        elif record.code.startswith("105") and record.series in {"K.1", "K.5"}:
            record.orientation = "Pozioma"
        elif record.multiplicity == "1x":
            record.orientation = "Bez orientacji"
        else:
            record.orientation = "Uniwersalna"

    if not record.network_category:
        match = re.search(r"\bkat(?:egoria)?\s*([56](?:e|a)?)\b", text)
        if match:
            record.network_category = f"Kat. {match.group(1).upper()}"

    if not record.screened:
        if re.search(r"\b(?:s ftp|ftp|stp|ekranowan)", text):
            record.screened = "Tak"
        elif re.search(r"\butp\b|nieekranowan", text):
            record.screened = "Nie"

    if not record.grounding:
        if "bez uziemienia" in text:
            record.grounding = "Nie"
        elif "z uziemieniem" in text:
            record.grounding = "Tak"

    if not record.terminal_type:
        if "lsa" in text:
            record.terminal_type = "LSA+"
        elif "samozacisk" in text:
            record.terminal_type = "Samozaciski"
        elif "srubowo window" in text:
            record.terminal_type = "Śrubowo-windowe"
        elif "zacisk" in text and "srub" in text:
            record.terminal_type = "Śrubowe"

    if not record.voltage:
        match = re.search(r"\b(\d+(?:[.,]\d+)?)\s*v\b", text)
        if match:
            record.voltage = match.group(1).replace(".", ",")
    if record.code.startswith("2031"):
        record.voltage = "24"
    elif record.code.startswith("2030") and record.primary_series not in {"R.1", "R.3", "R.8"}:
        record.voltage = "230"

    if not record.nominal_current:
        match = re.search(r"\b(\d+(?:[.,]\d+)?)\s*a\b", text)
        if match:
            record.nominal_current = match.group(1).replace(".", ",")

    if not record.pole_count:
        match = re.search(r"\b([1-4])\s*biegunow", text)
        if match:
            record.pole_count = match.group(1)

    if not record.detection_range and re.search(r"czujnik|nasadka ir|ruchu", text):
        match = re.search(r"(\d+(?:[.,]\d+)?)\s*m\b", record.title)
        if match:
            record.detection_range = f"{match.group(1).replace('.', ',')} m"

    if not record.output_count:
        match = re.search(r"\b(\d+)\s*(?:wyjsciow|wejsciow)", text)
        if match:
            record.output_count = match.group(1)

    if not record.shutters:
        # dane producenta uzywaja obu pisowni: "przeslony" i "przyslony"
        if re.search(r"bez (?:prze|przy)slon|bez podwyzszonej ochrony stykow", text):
            record.shutters = "Nie"
        elif re.search(r"\b(?:prze|przy)slon\w*|z podwyzszona ochrona stykow", text):
            record.shutters = "Tak"


def normalize_variant_core(record: Record) -> None:
    value = record.core
    if norm(record.distinguishing_feature) in AXIS_SAFE_CONSTRUCTION_FEATURES:
        value = re.sub(r"\s*\[[^\]]+\]\s*", " ", value)

    patterns = (
        r"\b[1-9]\s*(?:kr\b|krotn\w*\b)",
        r"\b[1-9]x\b",
        r"\b(?:jednokrotn\w*|dwukrotn\w*|trzykrotn\w*|czterokrotn\w*|pieciokrotn\w*|pojedyncz\w*|podwojn\w*|potrojn\w*|poczworn\w*|pionow\w*|poziom\w*|snieznobial\w*|krem\w*)\b",
        r"\bkat(?:egoria)?\s*[56](?:e|a)?\b",
        r"\b(?:s?ftp|utp|stp|ekranowan\w*|nieekranowan\w*)\b",
        r"\b(?:bez|z) uziemieni\w*\b",
        r"\b(?:zaciski? )?(?:samozacisk\w*|srubowo window\w*|srubow\w*)\b",
        r"\b\d+(?:[.,]\d+)?\s*v\b",
        r"\b\d+(?:[.,]\d+)?\s*a\b",
        r"\b[1-4]\s*biegunow\w*\b",
        r"\b(?:bez|z) (?:prze|przy)slon\w*(?: stykow)?\b",
        r"\b(?:bez|z) podwyzszon\w* ochron\w* styk\w*\b",
    )
    for pattern in patterns:
        value = re.sub(pattern, " ", value)
    value = re.sub(r"\bpuszka n t\b", "puszka natynkowa", value)
    value = re.sub(r"\bgniazdo komputerowego\b", "gniazdo komputerowe", value)
    value = re.sub(r"\bgniazdo glosnikowego\b", "gniazdo glosnikowe", value)
    record.core = compact(value)


COLOR_TERM_INFO = {
    "bialy": ("Biały", "Biały"), "biala": ("Biały", "Biały"), "biale": ("Biały", "Biały"),
    "snieznobialy": ("Śnieżnobiały", "Biały"), "snieznobiala": ("Śnieżnobiały", "Biały"),
    "czarny": ("Czarny", "Czarny"), "czarna": ("Czarny", "Czarny"), "czarne": ("Czarny", "Czarny"),
    "antracyt": ("Antracyt", "Czarny"), "grafitowy": ("Grafitowy", "Grafitowy"),
    "szary": ("Szary", "Szary"), "szara": ("Szary", "Szary"), "szare": ("Szary", "Szary"),
    "srebrny": ("Srebrny", "Srebrny"), "srebrna": ("Srebrny", "Srebrny"), "srebrne": ("Srebrny", "Srebrny"),
    "kremowy": ("Kremowy", "Beżowy"), "kremowa": ("Kremowy", "Beżowy"), "krem": ("Kremowy", "Beżowy"),
    "bezowy": ("Beżowy", "Beżowy"), "bezowa": ("Beżowy", "Beżowy"),
    "czerwony": ("Czerwony", "Czerwony"), "czerwona": ("Czerwony", "Czerwony"),
    "zielony": ("Zielony", "Zielony"), "zielona": ("Zielony", "Zielony"),
    "zolty": ("Żółty", "Żółty"), "zolta": ("Żółty", "Żółty"),
    "pomaranczowy": ("Pomarańczowy", "Pomarańczowy"), "pomaranczowa": ("Pomarańczowy", "Pomarańczowy"),
    "brazowy": ("Brązowy", "Brązowy"), "brazowa": ("Brązowy", "Brązowy"),
    "zloty": ("Złoty", "Złoty"), "zlota": ("Złoty", "Złoty"),
    "alu": ("Aluminium", "Szary"),
    "przezroczysty": ("Przezroczysty", "Transparentny"),
    "przezroczysta": ("Przezroczysty", "Transparentny"),
    "transparentny": ("Przezroczysty", "Transparentny"),
    "jasny": ("Jasny", "Transparentny"),
}

GENERAL_COLOR_ORDER = [
    "Biały", "Czarny", "Grafitowy", "Szary", "Srebrny", "Beżowy", "Brązowy",
    "Czerwony", "Zielony", "Żółty", "Pomarańczowy", "Złoty", "Transparentny",
]


def canonical_material_detail(value: str) -> str:
    text = norm(value)
    if not text:
        return ""
    if "beton strukturalny" in text:
        return "Beton strukturalny"
    if "lupek naturalny" in text:
        return "Łupek naturalny"
    if "stal szlachetna nierdzewna" in text:
        return "Stal szlachetna nierdzewna"
    if "stal szlachetna" in text:
        return "Stal szlachetna"
    if "aluminium mosiadz" in text:
        return "Aluminium mosiężne"
    if ("szklo" in text or "szklana" in text) and ("aluminium" in text or "aluminiowa" in text):
        return "Szkło aluminium"
    if "aluminium" in text or "aluminiowa" in text or re.search(r"\balu\b", text):
        return "Aluminium"
    if "szklo" in text or "szklana" in text:
        return "Szkło"
    if "porcelan" in text:
        return "Porcelana"
    if "akryl" in text:
        return "Akryl"
    if "beton" in text:
        return "Beton"
    if "lupek" in text:
        return "Łupek"
    if "drewno buk" in text or "buk" in text:
        return "Drewno bukowe"
    if "drewno dab" in text or "dab" in text:
        return "Dąb"
    if "drewno" in text:
        return "Drewno"
    if "skora" in text:
        return "Skóra"
    if "mosiadz" in text:
        return "Mosiądz"
    if "kamien" in text:
        return "Kamień"
    if "stal" in text:
        return "Stal"
    return ""


def material_attribute_value(value: str) -> str:
    detail = canonical_material_detail(value)
    if detail in {"Aluminium", "Aluminium mosiężne", "Stal", "Stal szlachetna", "Stal szlachetna nierdzewna", "Mosiądz"}:
        return "Metal"
    if detail.startswith("Szkło"):
        return "Szkło"
    if detail in {"Drewno", "Drewno bukowe", "Dąb"}:
        return "Drewno"
    if detail:
        return detail
    text = norm(value)
    if "tworzyw" in text or "termoplast" in text:
        return "Tworzywo sztuczne"
    return ""


def canonical_finish(value: str) -> tuple[str, bool]:
    text = norm(value)
    primary = ""
    if "aksamit" in text:
        primary = "aksamit"
    elif "mat" in text or "matow" in text:
        primary = "mat"
    elif "polysk" in text:
        primary = "połysk"
    elif "satyn" in text:
        primary = "satyna"
    elif "chrom" in text:
        primary = "chromowany"
    elif "nikl" in text:
        primary = "niklowany"
    return primary, "lakier" in text


def broad_color_matches(expected: str, source_values: set[str]) -> bool:
    if expected in source_values:
        return True
    metallic = {"Szary", "Srebrny"}
    if expected in metallic and source_values & metallic:
        return True
    if expected == "Czarny" and source_values & {"Czarny", "Grafitowy"}:
        return True
    if expected == "Grafitowy" and source_values & {"Czarny", "Grafitowy", "Szary"}:
        return True
    return False


def derive_manufacturer_color(record: Record) -> str:
    override = CATALOG_OVERRIDES.get(record.code, {})
    if override.get("manufacturer_color"):
        return compact(override["manufacturer_color"])

    source_general = {compact(item) for item in str(record.color or "").split("|") if compact(item)}
    color_terms: list[tuple[str, str]] = []
    for token in norm(record.color_detail).split():
        info = COLOR_TERM_INFO.get(token)
        if not info or info in color_terms:
            continue
        display, broad = info
        color_terms.append((display, broad))

    material_display = canonical_material_detail(record.material_detail)
    color_displays = [display for display, _broad in color_terms]
    if material_display == "Szkło" and color_displays == ["Biały"]:
        base = "Szkło białe"
    elif material_display == "Szkło" and color_displays == ["Czarny"]:
        base = "Szkło czarne"
    elif material_display and color_displays:
        filtered = [value for value in color_displays if norm(value) != norm(material_display)]
        base = material_display if not filtered else f"{material_display} / {' / '.join(filtered)}"
    elif material_display:
        base = material_display
    elif color_displays:
        base = " / ".join(color_displays)
    else:
        base = compact(record.manufacturer_color_before)
        if not base and len(source_general) == 1:
            base = next(iter(source_general))

    finish, lacquered = canonical_finish(record.finish_detail)
    if finish and norm(finish) not in norm(base):
        base = compact(f"{base} {finish}")
    if lacquered and "lakier" not in norm(base):
        base = compact(f"{base}, lakierowany")
    return compact(base)


def general_color_from_manufacturer(value: str, fallback: str = "") -> str:
    text = norm(value)
    found: set[str] = set()
    checks = (
        ("Biały", ("bial", "snieznobial")),
        ("Czarny", ("czarn", "antracyt")),
        ("Grafitowy", ("grafit",)),
        ("Szary", ("szar", "aluminium", "alu", "stal", "beton", "lupek")),
        ("Srebrny", ("srebr", "chrom", "nikiel")),
        ("Beżowy", ("krem", "bez")),
        ("Brązowy", ("braz", "dab", "buk", "drewno", "skora")),
        ("Czerwony", ("czerwon",)),
        ("Zielony", ("zielon",)),
        ("Żółty", ("zolt",)),
        ("Pomarańczowy", ("pomarancz",)),
        ("Złoty", ("zlot", "mosiadz")),
        ("Transparentny", ("przezrocz", "transparent", "jasny")),
    )
    for label, needles in checks:
        if any(needle in text for needle in needles):
            found.add(label)
    if not found:
        return compact(fallback)
    return "|".join(label for label in GENERAL_COLOR_ORDER if label in found)


def choose_general_color(current: str, manufacturer_color: str) -> str:
    expected = general_color_from_manufacturer(manufacturer_color, current)
    current_values = {compact(item) for item in str(current or "").split("|") if compact(item)}
    expected_values = {compact(item) for item in str(expected or "").split("|") if compact(item)}
    if current_values and expected_values and all(
        broad_color_matches(value, current_values) for value in expected_values
    ):
        return compact(current)
    return expected


def find_headers(sheet) -> dict[str, int]:
    headers = [cell.value for cell in sheet[1]]
    exact = {str(value): index for index, value in enumerate(headers)}

    def find(text: str) -> int:
        target = norm(text)
        for index, value in enumerate(headers):
            if norm(value) == target:
                return index
        raise KeyError(text)

    return {
        "id": exact["id"], "title": exact["Title"], "sku": exact["SKU"],
        "product_type": exact["Product Type"], "parent": exact["Parent Product ID"],
        "color": find("Atrybut Produktu: Kolor"),
        "manufacturer_color": find("Atrybut Produktu: Kolor Producenta"),
        "material": find("Atrybut Produktu: Materiał"),
        "multiplicity": find("Atrybut Produktu: Krotność"), "series": find("Atrybut Produktu: Seria"),
        "network_category": find("Atrybut Produktu: Kategoria Kabli Sieciowych"),
        "screened": find("Atrybut Produktu: Ekranowany"),
        "grounding": find("Atrybut Produktu: Uziemienie"),
        "terminal_type": find("Atrybut Produktu: Typ Zacisków"),
        "voltage": find("Atrybut Produktu: Napięcie [V]"),
        "nominal_current": find("Atrybut Produktu: Prąd Znamionowy [A]"),
        "pole_count": find("Atrybut Produktu: Liczba Biegunów"),
        "shutters": find("Atrybut Produktu: Przesłony Torów Prądowych"),
        "detection_range": find("Atrybut Produktu: Zasięg Detekcji"),
        "output_count": find("Atrybut Produktu: Liczba Wyjść"),
        "type": find("Atrybut Produktu: Typ Produktu"), "tags": exact["Product Tags"],
    }


def load_records(sheet, columns: dict[str, int]) -> list[Record]:
    records: list[Record] = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        title = compact(values[columns["title"]])
        sku = compact(values[columns["sku"]])
        product_type = compact(values[columns["type"]])
        series = compact(values[columns["series"]])
        frame = is_frame(title, product_type)
        core = title_core(title, sku, frame)
        primary = infer_primary_series(series, title)
        multiplicity = multiplicity_axis(title, values[columns["multiplicity"]], frame)
        manufacturer_color_before = compact(values[columns["manufacturer_color"]])
        record = Record(
            excel_row=excel_row,
            product_id=compact(values[columns["id"]]),
            title=title,
            sku=sku,
            code=sku_base(sku),
            product_type=product_type,
            series=series,
            primary_series=primary,
            color=compact(values[columns["color"]]),
            manufacturer_color_before=manufacturer_color_before,
            manufacturer_color=manufacturer_color_before,
            material=compact(values[columns["material"]]),
            multiplicity=multiplicity,
            tags_before=compact(values[columns["tags"]]),
            core=core,
            family=category_family(core, product_type, frame),
            frame=frame,
            color_detail=extracted_axis(title, COLOR_WORDS),
            material_detail=extracted_axis(title, MATERIAL_WORDS),
            finish_detail=extracted_axis(title, FINISH_WORDS),
            orientation=orientation_axis(title, frame, multiplicity),
            network_category=compact(values[columns["network_category"]]),
            screened=compact(values[columns["screened"]]),
            grounding=compact(values[columns["grounding"]]),
            terminal_type=compact(values[columns["terminal_type"]]),
            voltage=compact(values[columns["voltage"]]),
            nominal_current=compact(values[columns["nominal_current"]]),
            pole_count=compact(values[columns["pole_count"]]),
            shutters=compact(values[columns["shutters"]]),
            detection_range=compact(values[columns["detection_range"]]),
            output_count=compact(values[columns["output_count"]]),
            alias_code=alias_berker_code(title),
        )
        record.series_key = infer_series_key(record.primary_series, record.title, record.code)
        apply_product_specifics(record)
        record.manufacturer_color = derive_manufacturer_color(record)
        infer_variant_axes(record)
        normalize_variant_core(record)
        records.append(record)
    return records


def build_pdf_index(pdf_path: Path, target_codes: set[str]) -> tuple[dict[str, list[int]], int]:
    pages_by_code: dict[str, list[int]] = defaultdict(list)
    doc = fitz.open(pdf_path)
    try:
        for page_number, page in enumerate(doc, 1):
            text = page.get_text("text").upper()
            tokens = set(re.findall(r"(?<![A-Z0-9])[A-Z0-9]{4,14}(?![A-Z0-9])", text))
            for code in tokens & target_codes:
                pages_by_code[code].append(page_number)
        return dict(pages_by_code), len(doc)
    finally:
        doc.close()


def assign_catalog_matches(records: list[Record], pages_by_code: dict[str, list[int]]) -> None:
    for record in records:
        code_pages = pages_by_code.get(record.code, [])
        alias_pages = pages_by_code.get(record.alias_code, []) if record.alias_code else []
        if code_pages:
            record.catalog_match = "DOPASOWANIE BEZPOŚREDNIE"
            record.catalog_reference = record.code
            record.catalog_pages = ", ".join(str(page) for page in code_pages[:12])
        elif alias_pages:
            record.catalog_match = "DOPASOWANIE PO KODZIE BERKER"
            record.catalog_reference = record.alias_code
            record.catalog_pages = ", ".join(str(page) for page in alias_pages[:12])
        else:
            record.catalog_match = "BRAK DOPASOWANIA"


def differing_axes(records: list[Record]) -> list[str]:
    axes: list[str] = []
    for label, field_name in VARIANT_AXIS_FIELDS:
        values = {
            norm(getattr(item, field_name))
            for item in records
            if norm(getattr(item, field_name))
        }
        if len(values) > 1:
            axes.append(label)
    return axes


def variant_axis_values(record: Record, axes: list[str]) -> tuple[str, ...]:
    return tuple(norm(getattr(record, VARIANT_AXIS_FIELD_BY_LABEL[axis])) for axis in axes)


def differs_on_exactly_one_axis(left: Record, right: Record, axes: list[str]) -> bool:
    left_values = variant_axis_values(left, axes)
    right_values = variant_axis_values(right, axes)
    return sum(left_value != right_value for left_value, right_value in zip(left_values, right_values)) == 1


def isolated_variant_records(records: list[Record], axes: list[str]) -> list[Record]:
    """Products without a one-axis transition cannot be switched by the variant plugin."""
    return [
        record
        for record in records
        if not any(
            other is not record and differs_on_exactly_one_axis(record, other, axes)
            for other in records
        )
    ]


def axis_summary(record: Record) -> str:
    values = []
    for label, value in (
        ("kolor producenta", record.manufacturer_color),
        ("materiał", record.material_detail or record.material),
        ("krotność", record.multiplicity),
        ("orientacja", record.orientation),
        ("kategoria kabli", record.network_category),
        ("ekranowany", record.screened),
        ("uziemienie", record.grounding),
        ("typ zacisków", record.terminal_type),
        ("napięcie [V]", record.voltage),
        ("prąd znamionowy [A]", record.nominal_current),
        ("liczba biegunów", record.pole_count),
        ("przesłony torów prądowych", record.shutters),
        ("zasięg detekcji", record.detection_range),
        ("liczba wyjść", record.output_count),
    ):
        if compact(value):
            values.append(f"{label}: {compact(value)}")
    return "; ".join(values)


def unsafe_technical_group(records: list[Record]) -> bool:
    text = " ".join(norm(item.title) for item in records)
    return bool(re.search(r"\bmechanizm\b|\bwklad\b|\bwkladka\b", text))


def assign_variant_groups(records: list[Record]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    candidates: dict[tuple[str, str, str], list[Record]] = defaultdict(list)
    missing_series_groups: dict[tuple[str, str], list[Record]] = defaultdict(list)
    for record in records:
        if record.is_old_duplicate:
            continue
        if not record.core:
            continue
        if record.series_key and record.series_key != "inne":
            candidates[(record.series_key, record.family, record.core)].append(record)
        else:
            missing_series_groups[(record.family, record.core)].append(record)

    accepted_groups: list[dict[str, Any]] = []
    rejected_groups: list[dict[str, Any]] = []
    accepted_pending: list[tuple[tuple[str, str, str], list[Record], list[str], list[Record]]] = []

    for key, group in sorted(candidates.items()):
        if len(group) < 2:
            continue
        reasons: list[str] = []

        signature_groups: dict[tuple[str, ...], list[Record]] = defaultdict(list)
        for record in group:
            signature_groups[record.axis_signature].append(record)
        ambiguous = [item for items in signature_groups.values() if len(items) > 1 for item in items]
        accepted = [items[0] for items in signature_groups.values() if len(items) == 1]
        axes = differing_axes(accepted)
        if len(accepted) < 2 or not axes:
            reasons.append("brak co najmniej dwóch unikalnych kombinacji bezpiecznych osi wariantu")

        if reasons:
            reason = "; ".join(dict.fromkeys(reasons))
            for record in group:
                record.decision = "DO WERYFIKACJI"
                record.reason = reason
                record.review_reasons.append(reason)
            rejected_groups.append({"key": key, "records": group, "reason": reason})
            continue

        for record in ambiguous:
            reason = "powtórzona kombinacja osi wariantu — możliwy duplikat lub inna cecha techniczna"
            record.decision = "DO WERYFIKACJI"
            record.reason = reason
            record.review_reasons.append(reason)

        if ambiguous:
            accepted = [record for record in accepted if record not in ambiguous]
            axes = differing_axes(accepted)
        if len(accepted) < 2 or not axes:
            reason = "po odrzuceniu powtórzonych kombinacji pozostało za mało pewnych wariantów"
            for record in group:
                if record.decision != "DO WERYFIKACJI":
                    record.decision = "DO WERYFIKACJI"
                    record.reason = reason
                    record.review_reasons.append(reason)
            rejected_groups.append({"key": key, "records": group, "reason": reason})
            continue

        if len(accepted) == 2 and len(axes) > 1:
            reason = (
                "grupa dwuproduktowa różni się na więcej niż jednej osi; "
                "wtyczka wariantowa nie może zmienić jednej cechy przy zachowaniu pozostałych"
            )
            for record in group:
                record.decision = "DO WERYFIKACJI"
                record.reason = reason
                record.review_reasons.append(reason)
            rejected_groups.append({"key": key, "records": group, "reason": reason})
            continue

        isolated = isolated_variant_records(accepted, axes)
        if isolated:
            reason = (
                "brak wariantu różniącego się dokładnie jedną osią; "
                "produkt byłby nieprzełączalny we wtyczce wariantowej"
            )
            for record in isolated:
                record.decision = "DO WERYFIKACJI"
                record.reason = reason
                record.review_reasons.append(reason)
            accepted = [record for record in accepted if record not in isolated]
            axes = differing_axes(accepted)

        if len(accepted) < 2 or not axes:
            reason = "po odrzuceniu nieprzełączalnych produktów pozostało za mało pewnych wariantów"
            for record in group:
                if record.decision != "DO WERYFIKACJI":
                    record.decision = "DO WERYFIKACJI"
                    record.reason = reason
                    record.review_reasons.append(reason)
            rejected_groups.append({"key": key, "records": group, "reason": reason})
            continue
        accepted_pending.append((key, accepted, axes, ambiguous))

    used_tags: set[str] = set()
    for group_number, (key, group, axes, ambiguous) in enumerate(sorted(accepted_pending, key=lambda item: item[0]), 1):
        group_id = f"HBV-{group_number:04d}"
        first = group[0]
        label = label_from_title(first.title, first.sku, first.primary_series, first.frame)
        series_labels = {
            "akcesoria uniwersalne": "Akcesoria uniwersalne",
            "system ip44": "System IP44",
            "k.1 k.5": "K.1 / K.5",
            "b s system": "B. Kwadrat / B.1 / B.3 / B.7 / S.1",
            "q system": "Q.1 / Q.3 / Q.7",
            "r.1 r.3 r.8": "R.1 / R.3 / R.8",
            "r system": "R.1 / R.3 / R.classic / Serie 1930",
            "b kwadrat shared box": "B. Kwadrat / Q.1 / Q.3",
            "s.1 r.3 shared box": "S.1 / R.3",
        }
        series_label = series_labels.get(key[0], first.primary_series or key[0])
        construction_label = first.distinguishing_value if first.construction_key else ""
        name_detail = f" ({construction_label})" if construction_label else ""
        group_name = compact(f"{label}{name_detail} — {series_label}")
        tag_detail = f" | {construction_label}" if construction_label else ""
        tag = compact(f"Wariant - {label} | {series_label}{tag_detail}")[:240]
        if norm(tag) in used_tags:
            tag = compact(f"{tag} | {group_id}")[:240]
        used_tags.add(norm(tag))
        raw_key = f"hager-berker:{key[0]}:{key[1]}:{key[2]}"
        for record in group:
            record.variant_group_id = group_id
            record.variant_group_key = raw_key
            record.variant_group_name = group_name
            record.variant_group_size = len(group)
            record.variant_tag = tag
            record.variant_axes = ", ".join(axes)
            record.variant_values = axis_summary(record)
            record.decision = "WARIANT — WYSOKA PEWNOŚĆ"
            record.reason = "wspólna funkcja i seria; różnice wyłącznie na bezpiecznych osiach wariantu"
        accepted_groups.append(
            {
                "id": group_id,
                "key": raw_key,
                "name": group_name,
                "tag": tag,
                "axes": axes,
                "records": group,
                "ambiguous": ambiguous,
            }
        )

    # Products without a reliable series and without a confirmed peer remain
    # standalone. Missing series alone is not evidence of an unresolved variant.

    return accepted_groups, rejected_groups


def mark_old_duplicate_products(records: list[Record]) -> int:
    by_code = {record.code: record for record in records if record.code}
    old_codes: set[str] = set()
    for new_record in records:
        if not new_record.alias_code or new_record.alias_code not in by_code:
            continue
        old_record = by_code[new_record.alias_code]
        old_record.is_old_duplicate = True
        old_record.replacement_code = new_record.code
        old_record.variant_group_id = ""
        old_record.variant_group_key = ""
        old_record.variant_group_name = ""
        old_record.variant_group_size = 0
        old_record.variant_tag = ""
        old_record.variant_axes = ""
        old_record.variant_values = ""
        old_record.decision = "POMINIĘTY — STARY PRODUKT DUPLIKATOWY"
        old_record.reason = f"stary numer Berker; aktualny odpowiednik Hager: {new_record.code}"
        old_record.review_reasons.append(old_record.reason)
        old_codes.add(old_record.code)
    return len(old_codes)


def duplicate_reference_rows(records: list[Record]) -> list[dict[str, Any]]:
    by_code = {record.code: record for record in records if record.code}
    rows: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for record in records:
        if not record.alias_code or record.alias_code not in by_code:
            continue
        old = by_code[record.alias_code]
        pair = tuple(sorted((record.code, old.code)))
        if pair in seen:
            continue
        seen.add(pair)
        rows.append(
            {
                "reference": record.alias_code,
                "old_sku": old.sku,
                "old_title": old.title,
                "new_sku": record.sku,
                "new_title": record.title,
                "same_primary_series": "TAK" if norm(old.primary_series) == norm(record.primary_series) else "NIE",
                "note": "Stary SKU pominięto w wariantach; nowy kod Hager pozostaje aktywnym odpowiednikiem.",
            }
        )
    return rows


def load_review_audit_source(path: Path | None) -> dict[str, list[str]]:
    """Load the exact SKU set the user asked to verify from the prior workbook."""
    if path is None or not path.exists():
        return {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Do weryfikacji" not in workbook.sheetnames:
            return {}
        sheet = workbook["Do weryfikacji"]
        reasons: dict[str, list[str]] = defaultdict(list)
        for values in sheet.iter_rows(min_row=2, values_only=True):
            reason = compact(values[0])
            sku = compact(values[1])
            if sku and reason and reason not in reasons[sku]:
                reasons[sku].append(reason)
        return dict(reasons)
    finally:
        workbook.close()


def audit_outcome(record: Record) -> tuple[str, str, str, str]:
    """Return outcome, distinguishing feature, value and evidence for audit sheet."""
    if record.is_old_duplicate:
        return (
            "STARY DUPLIKAT — POMINIĘTY",
            "Numer referencyjny",
            f"stary kod; aktualny odpowiednik {record.replacement_code}",
            "Powiązanie starego numeru Berker z aktualnym kodem Hager w nazwie produktu.",
        )
    if record.decision == "WARIANT — WYSOKA PEWNOŚĆ":
        feature = record.distinguishing_feature or record.variant_axes or "Bezpieczna oś wariantu"
        value = record.distinguishing_value or record.variant_values or axis_summary(record)
        evidence = record.audit_evidence or "Wspólna funkcja i seria; unikalna kombinacja koloru, materiału, wykończenia, krotności lub orientacji."
        return "WARIANT POTWIERDZONY", feature, value, evidence
    if record.decision == "DO WERYFIKACJI":
        feature = record.distinguishing_feature or "Nierozstrzygnięta cecha"
        value = record.distinguishing_value or record.reason
        evidence = record.audit_evidence or "Dane źródłowe i katalog nie pozwalają bezpiecznie rozstrzygnąć."
        return "NADAL DO WERYFIKACJI", feature, value, evidence
    feature = record.distinguishing_feature or "Funkcja / wykonanie techniczne"
    value = record.distinguishing_value or record.core
    evidence = record.audit_evidence or "Brak drugiego aktywnego SKU o tej samej funkcji i wykonaniu; produkt pozostaje samodzielny."
    return "PRODUKT SAMODZIELNY", feature, value, evidence


def style_header(sheet, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor=BLUE)
    for cell in sheet[row]:
        cell.fill = fill
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)
    sheet.row_dimensions[row].height = 36


def style_table_sheet(sheet, widths: dict[str, float], filter_ref: str | None = None) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    style_header(sheet)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            cell.border = Border(bottom=THIN_GRAY)
    if filter_ref:
        sheet.auto_filter.ref = filter_ref
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"


def add_report_sheet(workbook, metrics: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Raport")
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "Hager / Berker — warianty produktów"
    sheet["A1"].font = Font(size=18, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=BLUE)
    sheet.merge_cells("A1:D1")
    sheet["A3"] = "Zakres"
    sheet["B3"] = "Tylko warianty. Up-sell i cross-sell pozostawiono bez zmian. Produkty zachowują typ simple."
    rows = [
        ("Produkty źródłowe", metrics["products"]),
        ("Grupy wysokiej pewności", metrics["variant_groups"]),
        ("Produkty otagowane jako wariant", metrics["variant_products"]),
        ("Produkty do weryfikacji", metrics["review_products"]),
        ("Produkty objęte pełnym audytem", metrics["audited_products"]),
        ("Z audytu: warianty potwierdzone", metrics["audited_variants_confirmed"]),
        ("Z audytu: produkty samodzielne", metrics["audited_standalone"]),
        ("Z audytu: stare duplikaty", metrics["audited_old_duplicates"]),
        ("Łączne zmiany atrybutów wariantowych", metrics["attribute_changes"]),
        ("Zmiany: Kolor Producenta", metrics["manufacturer_color_changes"]),
        ("Zmiany: Kolor", metrics["general_color_changes"]),
        ("Zmiany: Materiał", metrics["material_changes"]),
        ("Zmiany: Krotność", metrics["multiplicity_changes"]),
        ("Zmiany: Orientacja", metrics["orientation_changes"]),
        ("Stare duplikaty pominięte w wariantach", metrics["old_duplicate_products_omitted"]),
        ("Dopasowanie PDF bezpośrednie", metrics["pdf_direct_matches"]),
        ("Dopasowanie PDF po kodzie Berker", metrics["pdf_alias_matches"]),
        ("Brak dopasowania w PDF", metrics["pdf_unmatched"]),
        ("Pary stary/nowy kod", metrics["duplicate_reference_pairs"]),
    ]
    sheet["A5"] = "Metryka"
    sheet["B5"] = "Wartość"
    for row in rows:
        sheet.append(row)
    style_header(sheet, 5)
    instruction_header_row = sheet.max_row + 2
    sheet.cell(instruction_header_row, 1, "Jak użyć pliku")
    sheet.cell(instruction_header_row, 1).font = Font(bold=True, color=WHITE)
    sheet.cell(instruction_header_row, 1).fill = PatternFill("solid", fgColor=BLUE)
    instructions = [
        "1. Arkusz „Produkty warianty” zawiera komplet danych źródłowych i dopisane kolumny kontrolne.",
        "2. Wspólny tag dopisano tylko rekordom oznaczonym „WARIANT — WYSOKA PEWNOŚĆ”.",
        "3. Kolor Producenta jest osią wariantu; ogólny Kolor pozostaje filtrem klienta.",
        "4. Uzupełniono istniejące atrybuty Kolor Producenta, Kolor, Materiał i Krotność oraz dodano Orientację.",
        "5. Arkusz „Uzupełnione atrybuty” pokazuje każdą zmianę wartości przed/po wraz ze źródłem reguły.",
        "6. Arkusz „Grupy wariantów” pokazuje gotowe grupy, osie wariantów i wszystkie SKU.",
        "7. Arkusz „Audyt weryfikacji” dokumentuje wynik kontroli każdej pozycji z poprzedniego arkusza „Do weryfikacji”.",
        "8. Arkusz „Do weryfikacji” zawiera wyłącznie pozycje, których nadal nie dało się bezpiecznie rozstrzygnąć.",
        "9. Nie zmieniono kolumn Product Type, Parent Product ID, Up-Sells ani Cross-Sells.",
        "10. Stare produkty duplikatowe pozostają w danych, ale nie mają tagu ani uzupełnionych atrybutów wariantowych.",
    ]
    for offset, instruction in enumerate(instructions, instruction_header_row + 1):
        sheet.cell(offset, 1, instruction)
        sheet.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=4)
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 95
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 18
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A5"


def write_workbook(
    input_path: Path,
    pdf_path: Path,
    output_path: Path,
    report_path: Path,
    review_source: Path | None = None,
) -> dict[str, Any]:
    input_hash = sha256(input_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(input_path, output_path)
    workbook = load_workbook(output_path, data_only=False)
    products = workbook.active
    products.title = "Produkty warianty"
    columns = find_headers(products)
    records = load_records(products, columns)
    audit_source = load_review_audit_source(review_source)

    target_codes = {record.code for record in records if len(record.code) >= 4}
    target_codes.update(record.alias_code for record in records if len(record.alias_code) >= 4)
    pages_by_code, pdf_pages = build_pdf_index(pdf_path, target_codes)
    assign_catalog_matches(records, pages_by_code)
    old_duplicate_count = mark_old_duplicate_products(records)
    accepted_groups, rejected_groups = assign_variant_groups(records)
    duplicate_rows = duplicate_reference_rows(records)

    orientation_col = products.max_column + 1
    products.cell(1, orientation_col, "Atrybut Produktu: Orientacja")
    attribute_changes: list[dict[str, str]] = []

    def set_attribute(record: Record, column: int, attribute: str, new_value: str, source: str) -> None:
        new_value = compact(new_value)
        if not new_value:
            return
        cell = products.cell(record.excel_row, column)
        old_value = compact(cell.value)
        if old_value == new_value:
            return
        cell.value = new_value
        attribute_changes.append(
            {
                "sku": record.sku,
                "title": record.title,
                "attribute": attribute,
                "before": old_value,
                "after": new_value,
                "change_type": "UZUPEŁNIONO" if not old_value else "POPRAWIONO",
                "source": source,
                "variant_group_id": record.variant_group_id,
                "variant_tag": record.variant_tag,
            }
        )

    for record in records:
        if not record.variant_tag or record.is_old_duplicate:
            continue
        set_attribute(
            record, columns["manufacturer_color"] + 1, "Kolor Producenta",
            record.manufacturer_color, "Nazwa produktu / katalog producenta",
        )
        set_attribute(
            record, columns["color"] + 1, "Kolor",
            choose_general_color(record.color, record.manufacturer_color),
            "Mapowanie ogólnego koloru z Koloru Producenta",
        )
        material_value = (
            material_attribute_value(record.material_detail)
            or material_attribute_value(record.manufacturer_color)
        )
        if material_value and not compact(products.cell(record.excel_row, columns["material"] + 1).value):
            set_attribute(
                record, columns["material"] + 1, "Materiał", material_value,
                "Materiał jawnie wskazany w nazwie lub katalogu",
            )
        if record.multiplicity:
            set_attribute(
                record, columns["multiplicity"] + 1, "Krotność", record.multiplicity,
                "Krotność ramki / elementu wskazana w nazwie",
            )
        if record.orientation:
            set_attribute(
                record, orientation_col, "Orientacja", record.orientation,
                "Orientacja wskazana w nazwie; brak kierunku oznacza wykonanie uniwersalne",
            )
        for column_key, label, value in (
            ("network_category", "Kategoria Kabli Sieciowych", record.network_category),
            ("screened", "Ekranowany", record.screened),
            ("grounding", "Uziemienie", record.grounding),
            ("terminal_type", "Typ Zacisków", record.terminal_type),
            ("voltage", "Napięcie [V]", record.voltage),
            ("nominal_current", "Prąd Znamionowy [A]", record.nominal_current),
            ("pole_count", "Liczba Biegunów", record.pole_count),
            ("shutters", "Przesłony Torów Prądowych", record.shutters),
            ("detection_range", "Zasięg Detekcji", record.detection_range),
            ("output_count", "Liczba Wyjść", record.output_count),
        ):
            if value:
                set_attribute(
                    record, columns[column_key] + 1, label, value,
                    "Nazwa produktu / katalog producenta; oś techniczna wariantu",
                )

    helper_headers = [
        "PDF: status dopasowania", "PDF: referencja", "PDF: strony pliku", "Wariant: decyzja",
        "Wariant: ID grupy", "Wariant: klucz grupy", "Wariant: nazwa grupy", "Wariant: liczba produktów",
        "Wariant: tag", "Wariant: osie", "Wariant: wartości", "Wariant: seria główna", "Wariant: rdzeń funkcji",
        "Wariant: uzasadnienie", "Product Tags — przed zmianą",
        "PDF: korekta danych katalogowych", "Audyt: cecha rozróżniająca",
        "Audyt: wartość cechy", "Audyt: strona drukowana katalogu",
    ]
    first_helper_col = products.max_column + 1
    for offset, header in enumerate(helper_headers):
        products.cell(1, first_helper_col + offset, header)

    for record in records:
        values = [
            record.catalog_match, record.catalog_reference, record.catalog_pages, record.decision,
            record.variant_group_id, record.variant_group_key, record.variant_group_name,
            record.variant_group_size or None, record.variant_tag, record.variant_axes, record.variant_values,
            record.primary_series, record.core, record.reason, record.tags_before,
            record.catalog_note, record.distinguishing_feature, record.distinguishing_value,
            record.catalog_printed_page,
        ]
        for offset, value in enumerate(values):
            products.cell(record.excel_row, first_helper_col + offset, value)
        if record.variant_tag:
            products.cell(record.excel_row, columns["tags"] + 1, append_tag(record.tags_before, record.variant_tag))

    products.freeze_panes = "D2"
    products.sheet_view.showGridLines = False
    products.auto_filter.ref = f"A1:{get_column_letter(products.max_column)}{products.max_row}"
    style_header(products)
    products.row_dimensions[1].height = 46
    for col in range(first_helper_col, products.max_column + 1):
        products.column_dimensions[get_column_letter(col)].width = 24
    products.column_dimensions[get_column_letter(first_helper_col + 5)].width = 46
    products.column_dimensions[get_column_letter(first_helper_col + 6)].width = 42
    products.column_dimensions[get_column_letter(first_helper_col + 10)].width = 48
    products.column_dimensions[get_column_letter(first_helper_col + 13)].width = 58
    products.conditional_formatting.add(
        f"{get_column_letter(first_helper_col + 3)}2:{get_column_letter(first_helper_col + 3)}{products.max_row}",
        FormulaRule(formula=[f'{get_column_letter(first_helper_col + 3)}2="WARIANT — WYSOKA PEWNOŚĆ"'], fill=PatternFill("solid", fgColor=GREEN)),
    )
    products.conditional_formatting.add(
        f"{get_column_letter(first_helper_col + 3)}2:{get_column_letter(first_helper_col + 3)}{products.max_row}",
        FormulaRule(formula=[f'{get_column_letter(first_helper_col + 3)}2="DO WERYFIKACJI"'], fill=PatternFill("solid", fgColor=YELLOW)),
    )

    groups_sheet = workbook.create_sheet("Grupy wariantów")
    groups_headers = [
        "ID grupy", "Nazwa grupy", "Wspólny tag", "Liczba produktów", "Osie wariantów",
        "SKU", "Tytuły", "Wartości wariantów", "Status", "Uzasadnienie",
    ]
    groups_sheet.append(groups_headers)
    for group in accepted_groups:
        items: list[Record] = group["records"]
        groups_sheet.append(
            [
                group["id"], group["name"], group["tag"], len(items), ", ".join(group["axes"]),
                " | ".join(item.sku for item in items),
                " | ".join(item.title for item in items),
                " | ".join(f"{item.sku}: {axis_summary(item)}" for item in items),
                "WYSOKA PEWNOŚĆ",
                "Wspólna funkcja i seria; każda pozycja ma unikalną kombinację bezpiecznych osi wariantu.",
            ]
        )
    style_table_sheet(
        groups_sheet,
        {"A": 13, "B": 42, "C": 55, "D": 16, "E": 28, "F": 75, "G": 90, "H": 90, "I": 20, "J": 70},
        f"A1:J{groups_sheet.max_row}",
    )
    for row in groups_sheet.iter_rows(min_row=2):
        for index in (5, 6, 7, 9):
            row[index].alignment = Alignment(vertical="top", wrap_text=True)

    changes_sheet = workbook.create_sheet("Uzupełnione atrybuty")
    changes_sheet.append(
        [
            "SKU", "Tytuł", "Atrybut", "Wartość przed", "Wartość po", "Rodzaj zmiany",
            "Źródło / reguła", "ID grupy", "Wspólny tag wariantu",
        ]
    )
    for change in attribute_changes:
        changes_sheet.append(list(change.values()))
    style_table_sheet(
        changes_sheet,
        {"A": 22, "B": 72, "C": 24, "D": 30, "E": 36, "F": 18, "G": 58, "H": 14, "I": 70},
        f"A1:I{changes_sheet.max_row}",
    )
    for row in changes_sheet.iter_rows(min_row=2):
        row[5].fill = PatternFill("solid", fgColor=GREEN if row[5].value == "UZUPEŁNIONO" else YELLOW)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    review_sheet = workbook.create_sheet("Do weryfikacji")
    review_headers = [
        "Powód", "SKU", "Tytuł", "Seria główna", "Rdzeń funkcji", "Typ produktu",
        "Kolor / materiał / wykończenie", "Krotność / orientacja", "Status PDF", "Strony PDF",
        "Kandydat do grupy", "Decyzja",
    ]
    review_sheet.append(review_headers)
    review_seen: set[tuple[str, str]] = set()
    for record in records:
        if record.decision != "DO WERYFIKACJI" or record.is_old_duplicate:
            continue
        reasons = list(dict.fromkeys(record.review_reasons or [record.reason]))
        for reason in dict.fromkeys(reasons):
            key = (record.sku, reason)
            if key in review_seen:
                continue
            review_seen.add(key)
            review_sheet.append(
                [
                    reason, record.sku, record.title, record.primary_series, record.core, record.product_type,
                    "; ".join(filter(None, [record.color_detail or record.color, record.material_detail or record.material, record.finish_detail])),
                    "; ".join(filter(None, [record.multiplicity, record.orientation])),
                    record.catalog_match, record.catalog_pages,
                    compact(f"{record.family} | {record.core} | {record.primary_series}"), record.decision,
                ]
            )
    style_table_sheet(
        review_sheet,
        {"A": 62, "B": 20, "C": 72, "D": 20, "E": 46, "F": 28, "G": 42, "H": 28, "I": 28, "J": 16, "K": 62, "L": 26},
        f"A1:L{review_sheet.max_row}",
    )
    for row in review_sheet.iter_rows(min_row=2):
        row[0].fill = PatternFill("solid", fgColor=YELLOW)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    audit_sheet = workbook.create_sheet("Audyt weryfikacji")
    audit_headers = [
        "SKU", "Tytuł", "Pierwotne powody weryfikacji", "Wynik audytu",
        "Cecha rozróżniająca", "Wartość cechy", "Dowód / uzasadnienie",
        "Strona drukowana katalogu", "Status PDF", "Finalna decyzja",
        "ID grupy", "Wspólny tag",
    ]
    audit_sheet.append(audit_headers)
    by_sku = {record.sku: record for record in records}
    for sku in sorted(audit_source, key=lambda value: norm(value)):
        record = by_sku.get(sku)
        if record is None:
            audit_sheet.append([sku, "", " | ".join(audit_source[sku]), "BRAK W ŹRÓDLE"])
            continue
        outcome, feature, value, evidence = audit_outcome(record)
        audit_sheet.append(
            [
                record.sku, record.title, " | ".join(audit_source[sku]), outcome,
                feature, value, evidence, record.catalog_printed_page,
                record.catalog_match, record.decision, record.variant_group_id, record.variant_tag,
            ]
        )
    style_table_sheet(
        audit_sheet,
        {"A": 22, "B": 72, "C": 58, "D": 28, "E": 34, "F": 50, "G": 72, "H": 18, "I": 28, "J": 28, "K": 14, "L": 55},
        f"A1:L{audit_sheet.max_row}",
    )
    for row in audit_sheet.iter_rows(min_row=2):
        color = GREEN if row[3].value in {"WARIANT POTWIERDZONY", "PRODUKT SAMODZIELNY"} else GRAY
        if row[3].value == "NADAL DO WERYFIKACJI":
            color = YELLOW
        row[3].fill = PatternFill("solid", fgColor=color)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    duplicates_sheet = workbook.create_sheet("Duplikaty referencji")
    duplicates_sheet.append(["Referencja Berker", "Stary SKU", "Stary tytuł", "Nowy SKU Hager", "Nowy tytuł", "Ta sama seria główna", "Uwagi"])
    for row in duplicate_rows:
        duplicates_sheet.append(list(row.values()))
    style_table_sheet(
        duplicates_sheet,
        {"A": 20, "B": 22, "C": 72, "D": 22, "E": 72, "F": 22, "G": 72},
        f"A1:G{duplicates_sheet.max_row}",
    )
    for row in duplicates_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    changes_by_attribute = Counter(change["attribute"] for change in attribute_changes)
    metrics = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "input_file": str(input_path),
        "input_sha256": input_hash,
        "pdf_file": str(pdf_path),
        "pdf_sha256": sha256(pdf_path),
        "pdf_pages": pdf_pages,
        "output_file": str(output_path),
        "products": len(records),
        "variant_groups": len(accepted_groups),
        "variant_products": sum(1 for item in records if item.variant_tag),
        "review_products": sum(1 for item in records if item.decision == "DO WERYFIKACJI" and not item.is_old_duplicate),
        "audited_products": len(audit_source),
        "audited_variants_confirmed": sum(1 for item in records if item.sku in audit_source and item.decision == "WARIANT — WYSOKA PEWNOŚĆ"),
        "audited_standalone": sum(1 for item in records if item.sku in audit_source and item.decision == "BRAK GRUPY"),
        "audited_old_duplicates": sum(1 for item in records if item.sku in audit_source and item.is_old_duplicate),
        "attribute_changes": len(attribute_changes),
        "manufacturer_color_changes": changes_by_attribute["Kolor Producenta"],
        "general_color_changes": changes_by_attribute["Kolor"],
        "material_changes": changes_by_attribute["Materiał"],
        "multiplicity_changes": changes_by_attribute["Krotność"],
        "orientation_changes": changes_by_attribute["Orientacja"],
        "network_category_changes": changes_by_attribute["Kategoria Kabli Sieciowych"],
        "screened_changes": changes_by_attribute["Ekranowany"],
        "grounding_changes": changes_by_attribute["Uziemienie"],
        "terminal_type_changes": changes_by_attribute["Typ Zacisków"],
        "voltage_changes": changes_by_attribute["Napięcie [V]"],
        "nominal_current_changes": changes_by_attribute["Prąd Znamionowy [A]"],
        "pole_count_changes": changes_by_attribute["Liczba Biegunów"],
        "shutters_changes": changes_by_attribute["Przesłony Torów Prądowych"],
        "detection_range_changes": changes_by_attribute["Zasięg Detekcji"],
        "output_count_changes": changes_by_attribute["Liczba Wyjść"],
        "rejected_candidate_groups": len(rejected_groups),
        "pdf_direct_matches": sum(1 for item in records if item.catalog_match == "DOPASOWANIE BEZPOŚREDNIE"),
        "pdf_alias_matches": sum(1 for item in records if item.catalog_match == "DOPASOWANIE PO KODZIE BERKER"),
        "pdf_unmatched": sum(1 for item in records if item.catalog_match == "BRAK DOPASOWANIA"),
        "duplicate_reference_pairs": len(duplicate_rows),
        "old_duplicate_products_omitted": old_duplicate_count,
    }
    add_report_sheet(workbook, metrics)
    workbook._sheets = [
        workbook["Raport"], workbook["Produkty warianty"], workbook["Grupy wariantów"],
        workbook["Uzupełnione atrybuty"], workbook["Audyt weryfikacji"],
        workbook["Do weryfikacji"], workbook["Duplikaty referencji"],
    ]
    workbook.active = 0
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output_path)
    workbook.close()

    if sha256(input_path) != input_hash:
        raise AssertionError("Plik źródłowy został zmieniony")

    validation_book = load_workbook(output_path, read_only=True, data_only=False)
    try:
        expected_sheets = [
            "Raport", "Produkty warianty", "Grupy wariantów", "Uzupełnione atrybuty",
            "Audyt weryfikacji", "Do weryfikacji", "Duplikaty referencji",
        ]
        products_sheet = validation_book["Produkty warianty"]
        group_sheet = validation_book["Grupy wariantów"]
        decision_col = first_helper_col + 3
        tagged_rows = 0
        for row in products_sheet.iter_rows(min_row=2, min_col=decision_col, max_col=decision_col, values_only=True):
            if row[0] == "WARIANT — WYSOKA PEWNOŚĆ":
                tagged_rows += 1
        product_rows = list(products_sheet.iter_rows(min_row=2, values_only=True))
        group_signature_counts: dict[str, Counter[tuple[str, ...]]] = defaultdict(Counter)
        group_id_col = first_helper_col + 4
        for values in product_rows:
            group_id = compact(values[group_id_col - 1])
            if not group_id:
                continue
            signature = (
                norm(values[columns["manufacturer_color"]]),
                norm(values[columns["multiplicity"]]),
                norm(values[orientation_col - 1]),
                norm(values[columns["network_category"]]),
                norm(values[columns["screened"]]),
                norm(values[columns["grounding"]]),
                norm(values[columns["terminal_type"]]),
                norm(values[columns["voltage"]]),
                norm(values[columns["nominal_current"]]),
                norm(values[columns["pole_count"]]),
                norm(values[columns["shutters"]]),
                norm(values[columns["detection_range"]]),
                norm(values[columns["output_count"]]),
            )
            group_signature_counts[group_id][signature] += 1
        axis_value_getters = (
            ("Kolor producenta", lambda item: item.manufacturer_color),
            ("Krotność", lambda item: item.multiplicity),
            ("Orientacja", lambda item: item.orientation),
            ("Kategoria kabli sieciowych", lambda item: item.network_category),
            ("Ekranowany", lambda item: item.screened),
            ("Uziemienie", lambda item: item.grounding),
            ("Typ zacisków", lambda item: item.terminal_type),
            ("Napięcie [V]", lambda item: item.voltage),
            ("Prąd znamionowy [A]", lambda item: item.nominal_current),
            ("Liczba biegunów", lambda item: item.pole_count),
            ("Przesłony torów prądowych", lambda item: item.shutters),
            ("Zasięg detekcji", lambda item: item.detection_range),
            ("Liczba wyjść", lambda item: item.output_count),
        )
        axis_value_violations = [
            {"sku": item.sku, "axis": label, "axes": item.variant_axes}
            for item in records if item.variant_tag
            for label, getter in axis_value_getters
            if label in item.variant_axes and not compact(getter(item))
        ]
        switchability_violations = [
            {
                "group_id": group["id"],
                "sku": item.sku,
                "axes": group["axes"],
            }
            for group in accepted_groups
            for item in isolated_variant_records(group["records"], group["axes"])
        ]
        two_product_multi_axis_violations = [
            {
                "group_id": group["id"],
                "skus": [item.sku for item in group["records"]],
                "axes": group["axes"],
            }
            for group in accepted_groups
            if len(group["records"]) == 2 and len(group["axes"]) > 1
        ]
        encoding_violations: list[dict[str, Any]] = []
        for sheet_name in expected_sheets:
            for row in validation_book[sheet_name].iter_rows():
                for cell in row:
                    markers = mojibake_markers(cell.value)
                    if markers:
                        encoding_violations.append({
                            "sheet": sheet_name,
                            "cell": cell.coordinate,
                            "markers": markers,
                            "value": str(cell.value)[:240],
                        })
                        if len(encoding_violations) >= 50:
                            break
                if len(encoding_violations) >= 50:
                    break
            if len(encoding_violations) >= 50:
                break
        validation = {
            "sheet_names_match": validation_book.sheetnames == expected_sheets,
            "product_rows_match": products_sheet.max_row == len(records) + 1,
            "group_rows_match": group_sheet.max_row == len(accepted_groups) + 1,
            "tagged_rows_match": tagged_rows == metrics["variant_products"],
            "all_products_remain_simple": all(
                value[0] == "simple"
                for value in products_sheet.iter_rows(min_row=2, min_col=columns["product_type"] + 1, max_col=columns["product_type"] + 1, values_only=True)
            ),
            "parent_ids_unchanged_zero": all(
                value[0] in (0, "0", None)
                for value in products_sheet.iter_rows(min_row=2, min_col=columns["parent"] + 1, max_col=columns["parent"] + 1, values_only=True)
            ),
            "old_duplicates_have_no_variant_tag": all(
                not item.variant_tag and item.decision == "POMINIĘTY — STARY PRODUKT DUPLIKATOWY"
                for item in records if item.is_old_duplicate
            ),
            "all_requested_review_products_audited": validation_book["Audyt weryfikacji"].max_row == len(audit_source) + 1,
            "attribute_change_rows_match": validation_book["Uzupełnione atrybuty"].max_row == len(attribute_changes) + 1,
            "variant_attribute_combinations_unique": all(
                count == 1
                for signatures in group_signature_counts.values()
                for count in signatures.values()
            ),
            "variant_axis_values_complete": not axis_value_violations,
            "variant_products_have_single_axis_transition": not switchability_violations,
            "two_product_groups_use_one_axis_only": not two_product_multi_axis_violations,
            "polish_text_encoding_intact": not encoding_violations,
        }
        validation["passed"] = all(validation.values())
        if not validation["passed"]:
            raise AssertionError(json.dumps(
                {
                    "validation": validation,
                    "axis_value_violations": axis_value_violations,
                    "switchability_violations": switchability_violations,
                    "two_product_multi_axis_violations": two_product_multi_axis_violations,
                    "encoding_violations": encoding_violations,
                },
                ensure_ascii=False, indent=2,
            ))
    finally:
        validation_book.close()

    metrics["validation"] = validation
    metrics["output_sha256"] = sha256(output_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")
    return metrics


def main() -> None:
    parser = argparse.ArgumentParser(description="Buduje konserwatywne grupy wariantów Hager/Berker")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--pdf", type=Path, default=DEFAULT_PDF)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--review-source", type=Path, default=DEFAULT_REVIEW_SOURCE)
    args = parser.parse_args()
    metrics = write_workbook(args.input, args.pdf, args.output, args.report, args.review_source)
    print(json.dumps(metrics, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
