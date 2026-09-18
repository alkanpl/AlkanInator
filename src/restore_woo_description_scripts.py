"""Przywraca w eksporcie Woo opisy, z ktorych REST API wycielo <script> JSON-LD.

Baselinker zapisuje produkty przez REST API WooCommerce, a kontroler produktow
zawsze robi ``wp_filter_post_kses`` na opisie - znacznik
``<script type="application/ld+json">`` znika i schema FAQ wyswietla sie jako
goly tekst. WP All Import dziala jako administrator i znacznik zachowuje, wiec
poprawny opis wgrywamy plikiem importu.

Opis jest podmieniany TYLKO wtedy, gdy tresc w Woo jest identyczna z nasza po
pominieciu znacznikow <script> - reczne zmiany w sklepie nie sa nadpisywane.

Uzycie:
    py src/restore_woo_description_scripts.py \
        --woo-export output/Simon-55-i-55-Go-i-54_TAGI_KASZMIR_v2.xlsx \
        --descriptions output/Simon_54_55_kaszmir_98_Baselinker_opisy_2026-09-16.csv \
        --output output/Simon-55-i-55-Go-i-54_TAGI_I_OPISY_KASZMIR_v3.xlsx \
        --report reports/simon_opisy_script_restore_report.json
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path
from typing import Any

SKU_COLUMN = "Sku"
CONTENT_COLUMN = "Content"
SCRIPT_TAG = re.compile(r"</?script[^>]*>", flags=re.IGNORECASE)
LD_JSON_BLOCK = re.compile(r"<script[^>]*application/ld\+json[^>]*>(.*?)</script>", flags=re.IGNORECASE | re.DOTALL)


def normalize_without_scripts(html: str | None) -> str:
    return re.sub(r"\s+", " ", SCRIPT_TAG.sub("", html or "")).strip()


def has_valid_ld_json(html: str) -> bool:
    blocks = LD_JSON_BLOCK.findall(html or "")
    if not blocks:
        return False
    for block in blocks:
        try:
            json.loads(block)
        except json.JSONDecodeError:
            return False
    return True


def decide(woo_content: str | None, our_description: str | None) -> str:
    if not our_description:
        return "brak naszego opisu"
    if not has_valid_ld_json(our_description):
        return "nasz opis nie ma poprawnego JSON-LD"
    if has_valid_ld_json(woo_content or ""):
        return "Woo ma juz poprawny script - bez zmian"
    if normalize_without_scripts(woo_content) != normalize_without_scripts(our_description):
        return "tresc w Woo rozni sie od naszej - NIE nadpisano"
    return "przywrocono"


def read_descriptions(path: Path) -> dict[str, str]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return {row["sku"]: row.get("description", "") for row in csv.DictReader(stream, delimiter=";")}


def main(argv: list[str] | None = None) -> int:
    from openpyxl import load_workbook

    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--woo-export", required=True, type=Path)
    parser.add_argument("--descriptions", required=True, type=Path, help="CSV Baselinkera z kolumnami sku;description.")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args(argv)
    if args.output.resolve() == args.woo_export.resolve():
        parser.error("Plik wyjsciowy nie moze nadpisywac wejsciowego.")

    descriptions = read_descriptions(args.descriptions)
    workbook = load_workbook(args.woo_export)
    sheet = workbook.worksheets[0]
    header = [cell.value for cell in sheet[1]]
    sku_index = header.index(SKU_COLUMN)
    content_index = header.index(CONTENT_COLUMN)  # pierwsza kolumna Content = opis produktu

    results: list[dict[str, Any]] = []
    for cells in sheet.iter_rows(min_row=2):
        sku = str(cells[sku_index].value or "")
        if sku not in descriptions:
            continue
        outcome = decide(cells[content_index].value, descriptions[sku])
        if outcome == "przywrocono":
            cells[content_index].value = descriptions[sku]
        results.append({"sku": sku, "wynik": outcome})
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)

    check = load_workbook(args.output, read_only=True).worksheets[0]
    restored = {item["sku"] for item in results if item["wynik"] == "przywrocono"}
    verified = sum(
        1
        for row in check.iter_rows(min_row=2, values_only=True)
        if str(row[sku_index] or "") in restored and has_valid_ld_json(row[content_index] or "")
    )
    counts: dict[str, int] = {}
    for item in results:
        counts[item["wynik"]] = counts.get(item["wynik"], 0) + 1
    report = {
        "woo_export": str(args.woo_export),
        "output": str(args.output),
        "matched_products": len(results),
        "results": counts,
        "verified_valid_ld_json_in_output": verified,
        "passed": verified == len(restored),
        "not_restored": [item for item in results if item["wynik"] != "przywrocono"],
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Dopasowane: {len(results)}, wyniki: {counts}, zweryfikowany JSON-LD: {verified} -> {args.output}")
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
