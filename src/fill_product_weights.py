"""Uzupelnia wage (Woo "Weight", kg) i proponowana klase wysylkowa w eksporcie Woo.

Wagi: dictionaries/product_shipping_weights_verified.yaml (karty producentow;
pierwszenstwo ma brutto, potem waga kartonu zbiorczego / szt., na koncu netto).
Istniejaca waga > 0 jest zostawiana, chyba ze --overwrite.

Klasa: z --gabaryt-class (domyslnie "Gabaryt") dostaja produkty, ktore nie
mieszcza sie w skrytce paczkomatu (64x38x41 cm): krotszy bok > 380 mm
(60x60, okragle 60 cm), dluzszy bok > 640 mm (120x30) albo kod z --force-sku
(ramki ADTR 60x30 pakowane po 2 szt. w karton 123 cm). Akcesoria (linki,
klipsy, uchwyty, zasilacze) i mini panele bez zmian. Pusta komorka "Shipping
Class" = nie zmieniaj (w WP All Import ustaw aktualizacje tylko niepustych).
Wymiary wysylkowe (--ship-dims, domyslnie dictionaries/product_shipping_dims_verified.yaml):
Woo Length/Width/Height = miejsce, jakie 1 sztuka zajmuje w kartonie zbiorczym
producenta (karton a x b x c/N). Nadpisuja wymiary produktu, bo z L/W/H
Flexible Shipping liczy wage gabarytowa (progi 30/60/90 kg = liczba paczek).
Panele w klasie Gabaryt dostaja wage wysylkowa wedlug producenta
(--packing, domyslnie dictionaries/panel_shipping_packing.yaml): Philips/Signify/
PILA 8 szt. w paczce (3,6 kg), Ledvance 12 szt. (2,4 kg), pozostali 10 szt. (2,88 kg).
Gdy waga gabarytowa (GLS, dzielnik 5000) nie miesci tylu sztuk w 30 kg, a produkt
nie ma zweryfikowanego kartonu, wysokosc wysylkowa jest obnizana (WARNING);
przy zweryfikowanym kartonie rozbieznosc jest tylko raportowana.
Plik wejsciowy nie jest nadpisywany.
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import math
import re
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import PatternFill

COL_DIM = "Atrybut produktu: Wymiary [mm]"
COL_D = "Atrybut produktu: Średnica"
PANEL_RE = re.compile(r"^(Panel LED|Ramka oświetleniowa LED)\b", re.I)
ACCESSORY_RE = re.compile(r"^(Linka|Klips|Uchwyt|Zasilacz)\b", re.I)


def mm_values(v):
    return [float(x.replace(",", ".")) for x in re.findall(r"\d+(?:[.,]\d+)?", str(v or ""))]


def packing_rule(title: str, sku: str, packing: dict) -> tuple[str, dict]:
    """Regula pakowania producenta: dopasowanie po marce w nazwie, potem po sufiksie SKU."""
    suffix = sku.rsplit("/", 1)[-1].upper() if "/" in sku else ""
    for name, rule in (packing.get("brands") or {}).items():
        if any(re.search(rf"\b{re.escape(m)}\b", title, re.I) for m in rule.get("match", [])):
            return name, rule
        if suffix and suffix in {x.upper() for x in rule.get("sku_suffixes", [])}:
            return name, rule
    return "default", packing["default"]


def shipping_weight(per_package: int, packing: dict) -> float:
    """Waga 1 szt. tak, by per_package szt. bylo < progu, a per_package+1 >= progu."""
    limit, fill_kg = packing["package_limit_kg"], packing["fill_kg"]
    w = round(fill_kg / per_package, 2)
    if not (per_package * w < limit <= (per_package + 1) * w):
        raise ValueError(f"waga {w} kg nie daje {per_package} szt. w paczce {limit} kg")
    return w


def pieces_per_package(unit: float, limit: float) -> int:
    """Ile sztuk miesci sie ponizej progu (Flexible Shipping: przedzial 30-60 zaczyna sie od 30)."""
    if unit <= 0:
        return 999
    n = math.floor(limit / unit)
    return n - 1 if n * unit >= limit else n


def max_height_cm(length: float, width: float, per_package: int, ratio: float, limit: float) -> float:
    """Najwieksza wysokosc (co 0,1 cm), przy ktorej per_package szt. ma wage gabarytowa < progu."""
    h = math.floor(limit * ratio / (per_package * length * width) * 10) / 10
    while h > 0 and per_package * length * width * h / ratio >= limit:
        h = round(h - 0.1, 1)
    return h


def main() -> None:
    today = dt.date.today().isoformat()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default=f"output/Panele-Led_WYMIARY_UZUPELNIONE_{today}.xlsx")
    ap.add_argument("--weights", default="dictionaries/product_shipping_weights_verified.yaml")
    ap.add_argument("--ship-dims", default="dictionaries/product_shipping_dims_verified.yaml",
                    help="wymiary kartonow; pusty string = nie zmieniaj Length/Width/Height")
    ap.add_argument("--output", default=f"output/Panele-Led_WYSYLKA_KARTONY_{today}.xlsx")
    ap.add_argument("--report", default=f"reports/Panele-Led_WYSYLKA_KARTONY_KONTROLA_{today}.xlsx")
    ap.add_argument("--gabaryt-class", default="Gabaryt")
    ap.add_argument("--force-sku", default="29840/KAN,29841/KAN,29842/KAN")
    ap.add_argument("--no-class", action="store_true", help="nie dodawaj kolumny Shipping Class")
    ap.add_argument("--overwrite", action="store_true")
    ap.add_argument("--packing", default="dictionaries/panel_shipping_packing.yaml",
                    help="ile paneli w paczce wg producenta; pusty string = prawdziwa waga paneli")
    args = ap.parse_args()

    weights = yaml.safe_load(Path(args.weights).read_text(encoding="utf-8"))["products"]
    ship = (yaml.safe_load(Path(args.ship_dims).read_text(encoding="utf-8")) or {}).get("products", {}) if args.ship_dims else {}
    force = {s.strip() for s in args.force_sku.split(",") if s.strip()}
    packing = yaml.safe_load(Path(args.packing).read_text(encoding="utf-8")) if args.packing else None
    wb = openpyxl.load_workbook(args.input)
    ws = wb.active
    header = [c.value for c in ws[1]]
    if not args.no_class and "Shipping Class" not in header:
        header.append("Shipping Class")
        ws.cell(row=1, column=len(header), value="Shipping Class")
    col = {n: i + 1 for i, n in enumerate(header)}
    fill = PatternFill("solid", fgColor="FFF2CC")
    report = []
    for r in range(2, ws.max_row + 1):
        get = lambda n: ws.cell(r, col[n]).value if n in col else None
        sku, title = str(get("Sku") or ""), str(get("Title") or "")
        old_w = get("Weight")
        w = weights.get(sku)
        new_w, status, note = old_w, "OK", ""
        if w:
            if old_w not in (None, "", 0, "0") and not args.overwrite:
                note = f"zostawiona istniejaca waga {old_w} kg (karta: {w['kg']} kg)"
            else:
                new_w = w["kg"]
            note = "; ".join(x for x in (note, w.get("note", "")) if x)
            if w["kind"] == "netto":
                status = "WARNING"
        elif old_w in (None, "", 0, "0"):
            status, note = "DO_SPRAWDZENIA", "brak wagi w zrodlach producenta"
        if new_w != old_w:
            ws.cell(r, col["Weight"]).value = new_w
            ws.cell(r, col["Weight"]).fill = fill

        dims_before = "/".join(str(get(k)) for k in ("Length", "Width", "Height"))
        sd = ship.get(sku)
        if sd:
            for k, v in zip(("Length", "Width", "Height"), sd["piece"]):
                v = int(v) if float(v).is_integer() else v
                if get(k) != v:
                    ws.cell(r, col[k]).value = v
                    ws.cell(r, col[k]).fill = fill
            if sd.get("status") == "WARNING":
                status = "WARNING" if status == "OK" else status
                note = "; ".join(x for x in (note, sd.get("note", "")) if x)
        dims_after = "/".join(str(get(k)) for k in ("Length", "Width", "Height"))

        cls = None
        if not args.no_class and not ACCESSORY_RE.match(title):
            dims = mm_values(get(COL_DIM))[:2] or mm_values(get(COL_D))[:1] * 2
            if sku in force or (len(dims) == 2 and (min(dims) > 380 or max(dims) > 640)):
                cls = args.gabaryt_class
                ws.cell(r, col["Shipping Class"]).value = cls
                ws.cell(r, col["Shipping Class"]).fill = fill
        real_w = new_w
        brand = per_package = fit_w = fit_dim = None
        if cls and packing and PANEL_RE.match(title):
            brand, prule = packing_rule(title, sku, packing)
            per_package = prule["per_package"]
            new_w = shipping_weight(per_package, packing)
            ws.cell(r, col["Weight"]).value = new_w
            ws.cell(r, col["Weight"]).fill = fill
            note = "; ".join(x for x in (note, f"waga wysylkowa {new_w} kg (max {per_package} szt./paczke, "
                                                 f"{brand}), prawdziwa {real_w} kg") if x)
            limit, ratio = packing["package_limit_kg"], packing["dim_weight_ratio"]
            lwh = [float(get(k) or 0) for k in ("Length", "Width", "Height")]
            fit_w = pieces_per_package(new_w, limit)
            fit_dim = pieces_per_package(lwh[0] * lwh[1] * lwh[2] / ratio, limit)
            if fit_dim < per_package:
                status = "WARNING" if status == "OK" else status
                h = max_height_cm(lwh[0], lwh[1], per_package, ratio, limit)
                if sd or not h:
                    note += (f"; waga gabarytowa (dzielnik {ratio:g}) miesci tylko {fit_dim} szt. w paczce "
                             f"- karton zweryfikowany, nie zmieniane")
                else:
                    ws.cell(r, col["Height"]).value = int(h) if float(h).is_integer() else h
                    ws.cell(r, col["Height"]).fill = fill
                    note += (f"; wysokosc wysylkowa {lwh[2]:g} -> {h:g} cm, zeby waga gabarytowa (dzielnik "
                             f"{ratio:g}) miescila {per_package} szt. w paczce (brak danych kartonu)")
                    fit_dim = pieces_per_package(lwh[0] * lwh[1] * h / ratio, limit)
            dims_after = "/".join(str(get(k)) for k in ("Length", "Width", "Height"))
        report.append({"id": get("id"), "Sku": sku, "Title": title, "Weight przed": old_w, "Weight po": new_w,
                       "waga producenta": real_w,
                       "rodzaj wagi": (w or {}).get("kind"), "source": (w or {}).get("source"),
                       "L/W/H przed [cm]": dims_before, "L/W/H po [cm]": dims_after,
                       "karton zbiorczy": f"{sd['carton']} / {sd['qty']} szt." if sd else None,
                       "Shipping Class (propozycja)": cls, "producent (pakowanie)": brand,
                       "szt. w paczce (cel)": per_package, "szt. w paczce wg wagi": fit_w,
                       "szt. w paczce wg gabarytu": fit_dim, "status": status, "uwagi": note})
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    rb = openpyxl.Workbook()
    rs = rb.active
    rs.title = "Waga i klasa"
    rs.append(list(report[0]))
    for row in report:
        rs.append(list(row.values()))
    rs.freeze_panes = "D2"
    rs.auto_filter.ref = rs.dimensions
    rb.save(args.report)
    c = collections.Counter(x["status"] for x in report)
    print(f"rows={len(report)} {dict(c)} gabaryt={sum(1 for x in report if x['Shipping Class (propozycja)'])} -> {args.output} | {args.report}")


if __name__ == "__main__":
    main()
