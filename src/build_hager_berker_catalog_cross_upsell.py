from __future__ import annotations

import argparse
import json
import re
import shutil
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import fitz
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from build_hager_berker_cross_upsell import (
    BLUE,
    KNOWN_TEXT_REPAIRS,
    WHITE,
    CatalogProduct,
    Recommendation,
    compare_preserved_cells,
    colors_compatible,
    cross_candidate,
    diversity_key,
    same_design_family,
    same_variant_group,
    series_compatible,
    sha256,
    up_sell_candidate,
    validate_recommendations,
)
from build_hager_berker_variants import (
    compact,
    find_headers,
    load_records,
    mark_old_duplicate_products,
    mojibake_markers,
    norm,
)


DEFAULT_INPUT = Path("output/Hager-Berker-Gniazdka_WOO_ALL_IMPORT_OSIE_POPRAWIONE.xlsx")
DEFAULT_PDF = Path("input/21PL001_Katalog_Osprzet_Hager_2021.pdf")
DEFAULT_OUTPUT = Path("output/Hager-Berker-Gniazdka_WOO_ALL_IMPORT_CROSS_UP_SELL_KATALOG.xlsx")
DEFAULT_REPORT = Path("reports/hager_berker_cross_up_sell_katalog/cross_up_sell_katalog_report.xlsx")

INDEX_PAGE_MARKER = "spis indeksow katalogowych"


@dataclass(frozen=True)
class CatalogEvidence:
    catalog_code: str
    printed_page: int
    pdf_page: int
    context: str
    page_text_normalized: str


@dataclass(frozen=True)
class CatalogRecommendation:
    source_sku: str
    target_sku: str
    kind: str
    score: int
    confidence: str
    reason: str
    method: str
    source_catalog_code: str
    source_catalog_page: int
    target_catalog_code: str
    target_catalog_page: int
    catalog_evidence: str


def normalize_catalog_code(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def candidate_catalog_codes(product: CatalogProduct) -> list[str]:
    values = [product.record.alias_code, product.record.code]
    if product.record.code.startswith("5BE"):
        values.append(product.record.code[3:])
    values.extend(re.findall(r"\b(?:[A-Z]{1,8}[A-Z0-9]*\d[A-Z0-9]*|\d{4,14})\b", product.title.upper()))
    result: list[str] = []
    for value in values:
        code = normalize_catalog_code(value)
        if len(code) >= 4 and code not in result:
            result.append(code)
    return result


def find_index_pdf_pages(document: fitz.Document) -> list[int]:
    result: list[int] = []
    for pdf_page in range(max(1, len(document) - 20), len(document) + 1):
        text = norm(document[pdf_page - 1].get_text("text"))
        if INDEX_PAGE_MARKER in text or ("spis indeksow" in text and "nr zamowienia" in text):
            result.append(pdf_page)
    return result


def build_catalog_page_index(document: fitz.Document, target_codes: set[str]) -> dict[str, int]:
    index: dict[str, int] = {}
    for pdf_page in find_index_pdf_pages(document):
        words = document[pdf_page - 1].get_text("words")
        for word in words:
            code = normalize_catalog_code(word[4])
            if code not in target_codes:
                continue
            page_candidates: list[tuple[float, int]] = []
            for page_word in words:
                value = str(page_word[4])
                if not value.isdigit():
                    continue
                if page_word[0] <= word[2] or page_word[0] - word[2] > 50:
                    continue
                if abs(page_word[1] - word[1]) > 2.5:
                    continue
                printed_page = int(value)
                if 1 <= printed_page < 619:
                    page_candidates.append((page_word[0], printed_page))
            if page_candidates:
                index[code] = min(page_candidates)[1]
    return index


def page_lines(page: fitz.Page) -> list[tuple[float, float, str]]:
    grouped: dict[tuple[int, int], list[tuple[Any, ...]]] = defaultdict(list)
    for word in page.get_text("words"):
        grouped[(int(word[5]), int(word[6]))].append(word)
    lines: list[tuple[float, float, str]] = []
    for words in grouped.values():
        words.sort(key=lambda item: item[0])
        lines.append((min(item[0] for item in words), min(item[1] for item in words), " ".join(str(item[4]) for item in words)))
    return sorted(lines, key=lambda item: (item[1], item[0]))


def line_contains_code(line: str, code: str) -> bool:
    tokens = [normalize_catalog_code(token) for token in line.split()]
    for start in range(len(tokens)):
        joined = ""
        for end in range(start, min(start + 6, len(tokens))):
            joined += tokens[end]
            if joined == code:
                return True
            if not code.startswith(joined):
                break
    return False


def entry_context(page: fitz.Page, code: str) -> str:
    lines = page_lines(page)
    positions = [index for index, (_x, _y, line) in enumerate(lines) if line_contains_code(line, code)]
    if not positions:
        return ""
    position = positions[-1]
    selected = [compact(line) for _x, _y, line in lines[max(0, position - 5): position + 2]]
    return " | ".join(line for line in selected if line)[:1200]


def build_product_evidence(
    document: fitz.Document,
    products: list[CatalogProduct],
    page_index: dict[str, int],
) -> tuple[dict[str, CatalogEvidence], list[dict[str, str]]]:
    result: dict[str, CatalogEvidence] = {}
    coverage: list[dict[str, str]] = []
    text_cache: dict[int, str] = {}
    for product in products:
        matched_code = next((code for code in candidate_catalog_codes(product) if code in page_index), "")
        if not matched_code:
            coverage.append({
                "sku": product.sku,
                "title": product.title,
                "status": "BRAK W INDEKSIE KATALOGU",
                "catalog_code": "",
                "printed_page": "",
            })
            continue
        printed_page = page_index[matched_code]
        pdf_page = printed_page + 1
        if not 1 <= pdf_page <= len(document):
            coverage.append({
                "sku": product.sku,
                "title": product.title,
                "status": "NIEPRAWIDŁOWA STRONA KATALOGU",
                "catalog_code": matched_code,
                "printed_page": str(printed_page),
            })
            continue
        page = document[pdf_page - 1]
        if printed_page not in text_cache:
            text_cache[printed_page] = norm(page.get_text("text"))
        context = entry_context(page, matched_code)
        evidence = CatalogEvidence(
            catalog_code=matched_code,
            printed_page=printed_page,
            pdf_page=pdf_page,
            context=context,
            page_text_normalized=text_cache[printed_page],
        )
        result[product.sku] = evidence
        coverage.append({
            "sku": product.sku,
            "title": product.title,
            "status": "DOPASOWANO",
            "catalog_code": matched_code,
            "printed_page": str(printed_page),
        })
    return result, coverage


def configuration_confirmed(source: CatalogProduct, target: CatalogProduct) -> bool:
    source_text = norm(source.title)
    target_text = norm(target.title)
    source_rotary = "obrotow" in source_text or "pokret" in source_text
    target_rotary = "obrotow" in target_text or "pokret" in target_text
    if source_rotary != target_rotary:
        return False
    source_keyed = "na klucz" in source_text or "kluczyk" in source_text
    target_keyed = "na klucz" in target_text or "kluczyk" in target_text
    if source_keyed != target_keyed:
        return False
    if source.configuration_count and target.configuration_count:
        return source.configuration_count == target.configuration_count
    return True


def page_confirms_series(product: CatalogProduct, evidence: CatalogEvidence) -> bool:
    return any(series and series in evidence.page_text_normalized for series in product.series)


def page_confirms_function(product: CatalogProduct, evidence: CatalogEvidence) -> bool:
    phrases = {
        "single_switch": ("lacznik 1 klawiszowy", "klawisze pojedyncze", "klawisz pojedynczy"),
        "stair_switch": ("schodowy", "uniwersalny"),
        "cross_switch": ("krzyzowy",),
        "double_switch": ("lacznik 2 klawiszowy", "klawisze podwojne", "klawisz podwojny"),
        "button": ("przycisk", "dzwonek"),
        "shutter_control": ("zaluzj",),
        "dimmer": ("sciemniacz",),
        "power_socket": ("gniazdo", "schuko"),
        "data": ("rj45", "uae", "komputer"),
        "tv_sat": ("anten", "sat", "rtv"),
        "usb": ("usb",),
        "hdmi": ("hdmi",),
        "audio_video": ("audio", "video", "glosnik", "cinch"),
        "thermostat": ("temperatur", "termostat"),
        "motion_sensor": ("czujnik ruch",),
    }.get(product.function, ())
    return bool(phrases and any(phrase in evidence.page_text_normalized for phrase in phrases))


def catalog_filter_candidate(
    source: CatalogProduct,
    target: CatalogProduct,
    source_evidence: CatalogEvidence,
    target_evidence: CatalogEvidence,
) -> CatalogRecommendation | None:
    base = cross_candidate(source, target)
    if not base:
        return None
    if not page_confirms_series(source, source_evidence):
        return None
    if target.role != "mechanism" and not page_confirms_series(target, target_evidence):
        return None

    if "mechanizm dopasowany" in base.reason:
        if not configuration_confirmed(source, target):
            return None
        if not page_confirms_function(source, source_evidence) or not page_confirms_function(target, target_evidence):
            return None
        method = "KOMPATYBILNOŚĆ FUNKCJI POTWIERDZONA W KATALOGU"
        proof = (
            f"Katalog s. {source_evidence.printed_page}: element {source_evidence.catalog_code} odpowiada funkcji "
            f"{source.function}; katalog s. {target_evidence.printed_page}: mechanizm {target_evidence.catalog_code} "
            f"ma zgodną funkcję i krotność."
        )
    elif "ramka" in base.reason or "element pasujący do ramki" in base.reason:
        method = "ZGODNOŚĆ SYSTEMOWA SERII I KOLORU"
        proof = (
            f"Katalog s. {source_evidence.printed_page} i {target_evidence.printed_page}: oba elementy są ujęte "
            f"w zgodnym systemie serii; XLSX potwierdza ten sam Kolor Producenta."
        )
    elif "uszczelniając" in base.reason:
        method = "ZGODNOŚĆ IP44 POTWIERDZONA W KATALOGU"
        proof = (
            f"Katalog s. {source_evidence.printed_page} i {target_evidence.printed_page}: produkt oraz zestaw "
            f"uszczelniający są ujęte w zgodnej serii IP44."
        )
    else:
        return None

    return CatalogRecommendation(
        base.source_sku,
        base.target_sku,
        base.kind,
        base.score,
        base.confidence,
        base.reason,
        method,
        source_evidence.catalog_code,
        source_evidence.printed_page,
        target_evidence.catalog_code,
        target_evidence.printed_page,
        proof,
    )


def catalog_up_sell_candidate(
    source: CatalogProduct,
    target: CatalogProduct,
    source_evidence: CatalogEvidence,
    target_evidence: CatalogEvidence,
) -> CatalogRecommendation | None:
    base = up_sell_candidate(source, target)
    if not base:
        base = safe_relaxed_up_sell_candidate(source, target)
    if not base:
        return None
    if not page_confirms_series(source, source_evidence) or not page_confirms_series(target, target_evidence):
        return None
    if not page_confirms_function(source, source_evidence) or not page_confirms_function(target, target_evidence):
        return None
    proof = (
        f"Katalog s. {source_evidence.printed_page}: wersja bazowa {source_evidence.catalog_code}; "
        f"katalog s. {target_evidence.printed_page}: wersja {target_evidence.catalog_code} zachowuje funkcję i dodaje: "
        f"{base.reason.split(':', 1)[-1].strip()}."
    )
    return CatalogRecommendation(
        base.source_sku,
        base.target_sku,
        base.kind,
        base.score,
        base.confidence,
        base.reason,
        "UPGRADE POTWIERDZONY NA KARTACH KATALOGOWYCH",
        source_evidence.catalog_code,
        source_evidence.printed_page,
        target_evidence.catalog_code,
        target_evidence.printed_page,
        proof,
    )


def power_socket_kind(product: CatalogProduct) -> str:
    text = norm(product.title)
    if any(token in text for token in ("vga", "hdmi", "usb", "rj45", "rj11", "telefon", "anten", "glosnik", "audio", "video")):
        return ""
    if "schuko" in text:
        return "schuko"
    if "bez uziem" in text:
        return "bez_uziemienia"
    if "z uziem" in text or "uziemion" in text:
        return "z_uziemieniem"
    return ""


def data_socket_signature(product: CatalogProduct) -> tuple[bool, bool] | None:
    text = norm(product.title)
    if "rj45" not in text and "uae" not in text:
        return None
    return ("rj45" in text or "uae" in text, "rj11" in text or "telefon" in text)


def safe_relaxed_up_sell_candidate(source: CatalogProduct, target: CatalogProduct) -> Recommendation | None:
    """Broader, but still family-safe, upgrade rules for sockets and RJ45."""
    if source.sku == target.sku or target.record.is_old_duplicate or same_variant_group(source, target):
        return None
    if source.role in {"frame", "seal", "mounting", "other"} or source.role != target.role:
        return None
    if source.function == "other" or source.function != target.function:
        return None
    if not same_design_family(source, target) or not colors_compatible(source, target):
        return None
    if source.configuration_count != target.configuration_count:
        return None

    source_values = source.improvements
    target_values = target.improvements
    if not all(target_value >= source_value for source_value, target_value in zip(source_values, target_values)):
        return None
    improved = [index for index, (source_value, target_value) in enumerate(zip(source_values, target_values)) if target_value > source_value]
    if not improved:
        return None

    source_text = norm(source.title)
    target_text = norm(target.title)
    if source.function == "power_socket":
        if power_socket_kind(source) == "" or power_socket_kind(source) != power_socket_kind(target):
            return None
        if any(index not in {0, 1, 2} for index in improved):
            return None
        protected_features = (
            ("pokryw",),
            ("pole opis", "polem opis"),
            ("samozacisk",),
            ("kontroln",),
        )
        for alternatives in protected_features:
            if any(feature in source_text for feature in alternatives) and not any(feature in target_text for feature in alternatives):
                return None
        if improved == [0] and any(token in source_text for token in ("lamp", "led", "podswietl")):
            return None
    elif source.function == "data":
        if data_socket_signature(source) is None or data_socket_signature(source) != data_socket_signature(target):
            return None
        if any(index not in {3, 4} for index in improved):
            return None
    else:
        return None

    labels = [
        "podświetlenie",
        "wyższy stopień IP",
        "przesłony torów prądowych",
        "wyższa kategoria sieciowa",
        "ekranowanie",
        "większe możliwości/moc",
    ]
    reason = "bezpieczne ulepszenie tej samej rodziny katalogowej: " + ", ".join(labels[index] for index in improved)
    return Recommendation(source.sku, target.sku, "Up-sell", 95 + len(improved), "WYSOKA", reason)


def line_code_candidates(line: str) -> list[str]:
    tokens = [normalize_catalog_code(token) for token in line.split()]
    result: list[str] = []
    for start in range(len(tokens)):
        joined = ""
        for end in range(start, min(start + 5, len(tokens))):
            token = tokens[end]
            if not token:
                continue
            joined += token
            if len(joined) >= 4 and joined not in result:
                result.append(joined)
    return result


def explicit_catalog_references(
    document: fitz.Document,
    evidence_by_sku: dict[str, CatalogEvidence],
    products_by_sku: dict[str, CatalogProduct],
) -> list[tuple[str, str, int, str, str]]:
    code_to_skus: dict[str, list[str]] = defaultdict(list)
    code_to_printed_pages: dict[str, set[int]] = defaultdict(set)
    page_to_home_codes: dict[int, set[str]] = defaultdict(set)
    for sku, evidence in evidence_by_sku.items():
        code_to_skus[evidence.catalog_code].append(sku)
        code_to_printed_pages[evidence.catalog_code].add(evidence.printed_page)
        page_to_home_codes[evidence.printed_page].add(evidence.catalog_code)

    relations: set[tuple[str, str, int, str, str]] = set()
    known_codes = sorted(code_to_skus, key=len, reverse=True)
    for printed_page, home_codes in page_to_home_codes.items():
        pdf_page = printed_page + 1
        if not 1 <= pdf_page <= len(document):
            continue
        page = document[pdf_page - 1]
        lines = page_lines(page)
        optional_headers = [y for _x, y, line in lines if "wspolpracuje z" in norm(line)]
        if not optional_headers:
            continue

        home_occurrences: list[tuple[float, str]] = []
        for _x, y, line in lines:
            for code in home_codes:
                if line_contains_code(line, code):
                    home_occurrences.append((y, code))
        home_occurrences = sorted(set(home_occurrences))
        if not home_occurrences:
            continue

        for x, y, line in lines:
            if x < 410:
                continue
            if not any(0 <= y - header_y <= 90 for header_y in optional_headers):
                continue
            raw_candidates = line_code_candidates(line)
            referenced_codes: set[str] = set()
            for raw in raw_candidates:
                if raw in code_to_skus and raw not in home_codes:
                    referenced_codes.add(raw)
                if len(raw) >= 4:
                    referenced_codes.update(
                        code for code in known_codes
                        if code not in home_codes and code.startswith(raw) and len(code) - len(raw) <= 3
                    )
            # The catalogue often shortens a family code (for example
            # ``10107 ..``) and prints the target page in the next column.
            # Use that page number to avoid expanding the prefix to products
            # from a different series that happen to share the same prefix.
            target_pages = {
                int(other_line)
                for other_x, other_y, other_line in lines
                if other_x > x
                and abs(other_y - y) <= 2
                and re.fullmatch(r"[0-9]{1,3}", compact(other_line))
                and 1 <= int(other_line) <= len(document)
            }
            if target_pages:
                referenced_codes = {
                    code for code in referenced_codes
                    if code_to_printed_pages[code].intersection(target_pages)
                }
            if not referenced_codes:
                continue
            following = [(home_y, code) for home_y, code in home_occurrences if 0 < home_y - y <= 150]
            if not following:
                continue
            nearest_y = min(home_y for home_y, _code in following)
            source_codes = {code for home_y, code in home_occurrences if nearest_y <= home_y <= nearest_y + 40}
            for source_code in source_codes:
                for target_code in referenced_codes:
                    for source_sku in code_to_skus[source_code]:
                        for target_sku in code_to_skus[target_code]:
                            if source_sku == target_sku:
                                continue
                            relations.add((
                                source_sku,
                                target_sku,
                                printed_page,
                                source_code,
                                target_code,
                            ))
    return sorted(relations)


def explicit_recommendations(
    references: Iterable[tuple[str, str, int, str, str]],
    products_by_sku: dict[str, CatalogProduct],
    evidence_by_sku: dict[str, CatalogEvidence],
) -> list[CatalogRecommendation]:
    result: list[CatalogRecommendation] = []
    seen: set[tuple[str, str]] = set()
    for source_sku, target_sku, page, source_code, target_code in references:
        for left, right in ((source_sku, target_sku), (target_sku, source_sku)):
            if (left, right) in seen:
                continue
            source = products_by_sku[left]
            target = products_by_sku[right]
            if source.record.is_old_duplicate or target.record.is_old_duplicate or same_variant_group(source, target):
                continue
            if source.role == "frame" or target.role == "frame":
                frame = source if source.role == "frame" else target
                other = target if source.role == "frame" else source
                seal_like = other.role == "seal" or "uszczel" in norm(other.title)
                if not seal_like:
                    if not series_compatible(frame, other) or not colors_compatible(frame, other):
                        continue
            source_evidence = evidence_by_sku[left]
            target_evidence = evidence_by_sku[right]
            result.append(CatalogRecommendation(
                left,
                right,
                "Cross-sell",
                125,
                "KATALOG - JAWNE POWIĄZANIE",
                "katalog wskazuje produkt jako opcjonalnie współpracujący",
                "JAWNE ODNIESIENIE „WSPÓŁPRACUJE Z”",
                source_evidence.catalog_code,
                source_evidence.printed_page,
                target_evidence.catalog_code,
                target_evidence.printed_page,
                f"Katalog s. {page}: w bloku produktu {source_code} wskazano {target_code} w sekcji „Współpracuje z opcjonalnie”.",
            ))
            seen.add((left, right))
    return result


def select_catalog_recommendations(
    products: list[CatalogProduct],
    evidence_by_sku: dict[str, CatalogEvidence],
    explicit: list[CatalogRecommendation],
    max_cross_sells: int = 4,
    max_up_sells: int = 3,
) -> tuple[dict[str, list[CatalogRecommendation]], dict[str, list[CatalogRecommendation]]]:
    products_by_sku = {product.sku: product for product in products}
    eligible = [product for product in products if product.sku in evidence_by_sku and not product.record.is_old_duplicate]
    explicit_by_source: dict[str, list[CatalogRecommendation]] = defaultdict(list)
    for recommendation in explicit:
        explicit_by_source[recommendation.source_sku].append(recommendation)

    cross: dict[str, list[CatalogRecommendation]] = {}
    up: dict[str, list[CatalogRecommendation]] = {}
    for source in eligible:
        candidates = list(explicit_by_source.get(source.sku, []))
        for target in eligible:
            candidate = catalog_filter_candidate(source, target, evidence_by_sku[source.sku], evidence_by_sku[target.sku])
            if candidate:
                candidates.append(candidate)
        unique: dict[str, CatalogRecommendation] = {}
        for candidate in candidates:
            current = unique.get(candidate.target_sku)
            if current is None or candidate.score > current.score:
                unique[candidate.target_sku] = candidate
        ordered = sorted(unique.values(), key=lambda item: (-item.score, products_by_sku[item.target_sku].title, item.target_sku))
        chosen: list[CatalogRecommendation] = []
        seen_diversity: set[tuple[str, str]] = set()
        for candidate in ordered:
            key = diversity_key(products_by_sku[candidate.target_sku])
            if candidate.score < 120 and key in seen_diversity:
                continue
            chosen.append(candidate)
            seen_diversity.add(key)
            if len(chosen) >= max_cross_sells:
                break
        if chosen:
            cross[source.sku] = chosen

        cross_targets = {item.target_sku for item in chosen}
        up_candidates: list[CatalogRecommendation] = []
        for target in eligible:
            candidate = catalog_up_sell_candidate(source, target, evidence_by_sku[source.sku], evidence_by_sku[target.sku])
            if candidate and candidate.target_sku not in cross_targets:
                up_candidates.append(candidate)
        up_candidates.sort(key=lambda item: (-item.score, products_by_sku[item.target_sku].title, item.target_sku))
        if up_candidates:
            up[source.sku] = up_candidates[:max_up_sells]
    return cross, up


def write_catalog_report(
    path: Path,
    products: list[CatalogProduct],
    coverage: list[dict[str, str]],
    cross: dict[str, list[CatalogRecommendation]],
    up: dict[str, list[CatalogRecommendation]],
    validation: dict[str, bool],
    encoding_repairs: list[dict[str, Any]],
    source_path: Path,
    pdf_path: Path,
    output_path: Path,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Podsumowanie"
    rows = [
        ("Plik źródłowy XLSX", str(source_path.resolve())),
        ("Katalog źródłowy PDF", str(pdf_path.resolve())),
        ("Plik do importu", str(output_path.resolve())),
        ("Produkty", len(products)),
        ("Produkty dopasowane do strony katalogowej", sum(row["status"] == "DOPASOWANO" for row in coverage)),
        ("Produkty bez strony katalogowej", sum(row["status"] != "DOPASOWANO" for row in coverage)),
        ("Produkty z cross-sell", len(cross)),
        ("Powiązania cross-sell", sum(len(items) for items in cross.values())),
        ("Jawne powiązania „Współpracuje z”", sum(item.method.startswith("JAWNE") for items in cross.values() for item in items)),
        ("Produkty z up-sell", len(up)),
        ("Powiązania up-sell", sum(len(items) for items in up.values())),
        ("Stare duplikaty pominięte", sum(product.record.is_old_duplicate for product in products)),
        ("Naprawione wartości kodowania", len(encoding_repairs)),
    ]
    summary.append(["Metryka", "Wartość"])
    for row in rows:
        summary.append(row)

    products_by_sku = {product.sku: product for product in products}
    relations = workbook.create_sheet("Powiązania katalogowe")
    relations.append([
        "Rodzaj", "SKU źródłowe", "Tytuł źródłowy", "Kod katalogowy źródła", "Strona źródła",
        "SKU docelowe", "Tytuł docelowy", "Kod katalogowy celu", "Strona celu", "Ocena", "Pewność",
        "Metoda", "Uzasadnienie", "Dowód katalogowy",
    ])
    for kind, mapping in (("Cross-sell", cross), ("Up-sell", up)):
        for source_sku in sorted(mapping):
            for item in mapping[source_sku]:
                relations.append([
                    kind,
                    source_sku,
                    products_by_sku[source_sku].title,
                    item.source_catalog_code,
                    item.source_catalog_page,
                    item.target_sku,
                    products_by_sku[item.target_sku].title,
                    item.target_catalog_code,
                    item.target_catalog_page,
                    item.score,
                    item.confidence,
                    item.method,
                    item.reason,
                    item.catalog_evidence,
                ])

    coverage_sheet = workbook.create_sheet("Pokrycie katalogu")
    coverage_sheet.append(["SKU", "Tytuł", "Status", "Kod katalogowy", "Strona drukowana"])
    for row in coverage:
        coverage_sheet.append([row["sku"], row["title"], row["status"], row["catalog_code"], row["printed_page"]])

    validation_sheet = workbook.create_sheet("Walidacja")
    validation_sheet.append(["Test", "Wynik"])
    for name, passed in validation.items():
        validation_sheet.append([name, "OK" if passed else "BŁĄD"])

    repairs_sheet = workbook.create_sheet("Naprawy kodowania")
    repairs_sheet.append(["Wiersz", "Kolumna", "Wartość przed", "Wartość po"])
    for repair in encoding_repairs:
        repairs_sheet.append([repair["row"], repair["column"], repair["before"], repair["after"]])

    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=BLUE)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            width = min(max(len(compact(cell.value)) for cell in column) + 2, 70)
            sheet.column_dimensions[column[0].column_letter].width = max(width, 12)
    workbook.save(path)
    workbook.close()


def build_catalog_cross_up_sell(
    input_path: Path,
    pdf_path: Path,
    output_path: Path,
    report_path: Path,
    max_cross_sells: int = 4,
    max_up_sells: int = 3,
) -> dict[str, Any]:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("Plik wynikowy musi być inny niż źródłowy.")
    source_hash = sha256(input_path)
    source_workbook = load_workbook(input_path, read_only=True, data_only=False)
    try:
        source_sheet = source_workbook["Worksheet"]
        records = load_records(source_sheet, find_headers(source_sheet))
    finally:
        source_workbook.close()
    mark_old_duplicate_products(records)
    products = [CatalogProduct(record) for record in records]
    products_by_sku = {product.sku: product for product in products}

    document = fitz.open(pdf_path)
    try:
        target_codes = {code for product in products for code in candidate_catalog_codes(product)}
        page_index = build_catalog_page_index(document, target_codes)
        evidence_by_sku, coverage = build_product_evidence(document, products, page_index)
        references = explicit_catalog_references(document, evidence_by_sku, products_by_sku)
        explicit = explicit_recommendations(references, products_by_sku, evidence_by_sku)
        cross, up = select_catalog_recommendations(
            products,
            evidence_by_sku,
            explicit,
            max_cross_sells=max_cross_sells,
            max_up_sells=max_up_sells,
        )
    finally:
        document.close()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(input_path, output_path)
    workbook = load_workbook(output_path, data_only=False)
    encoding_repairs: list[dict[str, Any]] = []
    try:
        sheet = workbook["Worksheet"]
        headers = [cell.value for cell in sheet[1]]
        up_column = headers.index("Up-Sells") + 1
        cross_column = headers.index("Cross-Sells") + 1
        for product in products:
            sheet.cell(product.record.excel_row, up_column).value = "|".join(
                item.target_sku for item in up.get(product.sku, [])
            )
            sheet.cell(product.record.excel_row, cross_column).value = "|".join(
                item.target_sku for item in cross.get(product.sku, [])
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                before = str(cell.value or "")
                if before not in KNOWN_TEXT_REPAIRS:
                    continue
                cell.value = KNOWN_TEXT_REPAIRS[before]
                encoding_repairs.append({
                    "row": cell.row,
                    "column": headers[cell.column - 1],
                    "before": before,
                    "after": cell.value,
                })
        workbook.save(output_path)
    finally:
        workbook.close()

    base_cross = {sku: [Recommendation(item.source_sku, item.target_sku, item.kind, item.score, item.confidence, item.reason) for item in items] for sku, items in cross.items()}
    base_up = {sku: [Recommendation(item.source_sku, item.target_sku, item.kind, item.score, item.confidence, item.reason) for item in items] for sku, items in up.items()}
    validation = validate_recommendations(products, base_cross, base_up, max_cross_sells, max_up_sells)
    validation["every_relation_has_catalog_pages"] = all(
        item.source_catalog_page > 0 and item.target_catalog_page > 0
        for mapping in (cross, up) for items in mapping.values() for item in items
    )
    validation["source_file_unchanged"] = sha256(input_path) == source_hash
    validation["all_non_target_cells_preserved"] = compare_preserved_cells(input_path, output_path, "Worksheet")

    reopened = load_workbook(output_path, read_only=True, data_only=False)
    try:
        encoding_violations = [
            (sheet.title, compact(value), mojibake_markers(value))
            for sheet in reopened.worksheets
            for row in sheet.iter_rows(values_only=True)
            for value in row
            if mojibake_markers(value)
        ]
        validation["workbook_polish_text_encoding_intact"] = not encoding_violations
    finally:
        reopened.close()
    if not all(validation.values()):
        raise AssertionError(json.dumps({"validation": validation, "encoding": encoding_violations[:20]}, ensure_ascii=False, indent=2))

    write_catalog_report(
        report_path,
        products,
        coverage,
        cross,
        up,
        validation,
        encoding_repairs,
        input_path,
        pdf_path,
        output_path,
    )
    return {
        "products": len(products),
        "catalog_matched_products": len(evidence_by_sku),
        "catalog_missing_products": len(products) - len(evidence_by_sku),
        "products_with_cross_sell": len(cross),
        "cross_sell_relations": sum(len(items) for items in cross.values()),
        "explicit_catalog_relations_selected": sum(item.method.startswith("JAWNE") for items in cross.values() for item in items),
        "products_with_up_sell": len(up),
        "up_sell_relations": sum(len(items) for items in up.values()),
        "old_duplicates_skipped": sum(product.record.is_old_duplicate for product in products),
        "encoding_repairs": len(encoding_repairs),
        "output": str(output_path.resolve()),
        "report": str(report_path.resolve()),
        "validation": validation,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Buduje katalogowo potwierdzony cross-sell i up-sell Hager/Berker.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--pdf", type=Path, default=DEFAULT_PDF)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--max-cross-sells", type=int, default=4)
    parser.add_argument("--max-up-sells", type=int, default=3)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    metrics = build_catalog_cross_up_sell(
        args.input,
        args.pdf,
        args.output,
        args.report,
        max_cross_sells=args.max_cross_sells,
        max_up_sells=args.max_up_sells,
    )
    print(json.dumps(metrics, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
