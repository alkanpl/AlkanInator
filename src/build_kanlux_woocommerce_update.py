from __future__ import annotations

import argparse
import json
import re
from copy import copy
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from export_to_baselinker_csv import (
    hermetic_fluorescent_title,
    normalize_feature_name,
    normalize_feature_value,
    product_type_from_title,
    product_types_agree,
    reclassify_max_bulb_power_feature,
    split_multi_color_feature,
)
from optimize_titles import normalize_seo_title_terms, normalize_title_uniqueness_key
from utils import color_stem, compact_spaces, ensure_dir, normalize_header, slugify


DEFAULT_INPUT = "input/KanluxWoo.xlsx"
DEFAULT_FEATURES = "output/kanlux_baselinker_attributes_update_2026-06-11.csv"
# Tytuly tez z atrybutowego CSV (ma kolumne name + wszystkie produkty, takze te spoza
# referencji Baselinkera) - dzieki temu produkty "tylko Woo" dostaja tytul z pipeline.
DEFAULT_TITLES = "output/kanlux_baselinker_attributes_update_2026-06-11.csv"
DEFAULT_OUTPUT = "output/KanluxWoo_poprawione_atrybuty_2026-06-11.xlsx"
DEFAULT_REPORTS = "reports/kanlux_woocommerce_update_2026-06-11"
ATTRIBUTE_PREFIX = "Atrybut Produktu: "
PRODUCT_ATTRIBUTES_COLUMN = "Product Attributes"
# Aktualnie pelny kolor producenta zostaje tekstem typu "Bialy / Czarny".
MULTI_VALUE_FEATURES: set[str] = set()


def woo_attribute_cell_value(feature_name: str, value: str) -> str:
    if feature_name in MULTI_VALUE_FEATURES:
        return re.sub(r"\s*/\s*", "|", value)
    return value


def normalize_sku(value: Any) -> str:
    text = compact_spaces(str(value or "")).upper()
    text = re.sub(r"/(?:KANL?|KANLUX)$", "", text)
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def php_string(value: str) -> str:
    return f's:{len(value.encode("utf-8"))}:"{value}";'


# Sluga taksonomii pa_ wyjatkowo skracamy dla wybranych atrybutow (na zyczenie sklepu).
ATTRIBUTE_SLUG_OVERRIDES = {
    "Maksymalna moc źródła światła": "max-moc-zrodla",
}


def serialize_product_attributes(feature_names: list[str]) -> str:
    items: list[str] = []
    for position, feature_name in enumerate(feature_names):
        taxonomy = f"pa_{ATTRIBUTE_SLUG_OVERRIDES.get(feature_name, slugify(feature_name))}"
        items.append(
            php_string(taxonomy)
            + "a:6:{"
            + php_string("name")
            + php_string(taxonomy)
            + php_string("value")
            + php_string("")
            + php_string("position")
            + f"i:{position};"
            + php_string("is_visible")
            + "i:1;"
            + php_string("is_variation")
            + "i:1;"
            + php_string("is_taxonomy")
            + "i:1;"
            + "}"
        )
    return f"a:{len(items)}:{{{''.join(items)}}}"


def read_feature_updates(path: str | Path) -> dict[str, dict[str, str]]:
    updates = pd.read_csv(path, sep=";", dtype=str, keep_default_na=False, encoding="utf-8-sig")
    result: dict[str, dict[str, str]] = {}
    for _, row in updates.iterrows():
        sku = normalize_sku(row.get("sku", ""))
        features = json.loads(str(row.get("features", "{}")))
        if sku:
            result[sku] = {str(name): compact_spaces(str(value)) for name, value in features.items() if compact_spaces(str(value))}
    return result


def read_title_updates(path: str | Path | None) -> dict[str, str]:
    if not path:
        return {}
    update_path = Path(path)
    if not update_path.exists():
        return {}
    updates = pd.read_csv(update_path, sep=";", dtype=str, keep_default_na=False, encoding="utf-8-sig")
    if "sku" not in updates.columns or "name" not in updates.columns:
        return {}
    result: dict[str, str] = {}
    for _, row in updates.iterrows():
        sku = normalize_sku(row.get("sku", ""))
        name = compact_spaces(str(row.get("name", "")))
        if sku and name:
            result[sku] = name
    return result


def title_with_verified_product_type(current_title: str, features: dict[str, str], verified_title: str = "") -> str:
    title = compact_spaces(current_title)
    product_type = compact_spaces(features.get("Typ produktu", ""))
    current_type = product_type_from_title(title)
    verified_type = product_type_from_title(verified_title)
    if (
        product_type
        and current_type
        and verified_title
        and verified_type
        and not product_types_agree(current_type, product_type)
        and product_types_agree(verified_type, product_type)
    ):
        title = verified_title
    elif verified_title and title_has_feature_conflict(title, verified_title, features):
        title = verified_title
    return hermetic_fluorescent_title(normalize_seo_title_terms(title), features)


def title_has_feature_conflict(title: str, verified_title: str, features: dict[str, str]) -> bool:
    return (
        title_mounting_conflicts_with_feature(title, verified_title, features)
        or title_missing_verified_identity(title, verified_title)
        or title_word_feature_missing(title, verified_title, features, "ksztalt")
        or title_word_feature_missing(title, verified_title, features, "kolor")
        or title_word_feature_missing(title, verified_title, features, "trzonek")
        or title_word_feature_missing(title, verified_title, features, "czujnik ruchu")
        or any(
        title_token_conflicts_with_feature(title, verified_title, feature_value, pattern)
        for feature_value, pattern in [
            (feature_by_normalized_name(features, "stopien ochrony ip"), r"\bIP\s*(\d{2})\b"),
            (feature_by_normalized_name(features, "moc w"), r"\b(\d+(?:[,.]\d+)?)\s*W\b"),
            (feature_by_normalized_name(features, "temperatura barwowa k"), r"\b(\d{4})\s*K\b"),
            (feature_by_normalized_name(features, "strumien swietlny lm"), r"\b(\d+(?:-\d+)?)\s*lm\b"),
            (feature_by_normalized_name(features, "wymiary mm"), r"\b(\d+(?:[xX]\d+){1,2})\s*(?:mm|cm)?\b"),
        ]
    )
    )


def title_missing_verified_identity(title: str, verified_title: str) -> bool:
    current = normalize_header(title)
    verified = normalize_header(verified_title)
    if not current or not verified:
        return False
    for term in ["asymetryczny", "symetryczny"]:
        if term in verified and term not in current:
            return True
    protected_suffixes = ["sky v", "sky", "pt", "pro", "hi", "aio"]
    protected_prefixes = ["adtr", "fl agor", "varso", "daba"]
    for prefix in protected_prefixes:
        for suffix in protected_suffixes:
            phrase = f"{prefix} {suffix}"
            if phrase in verified and phrase not in current:
                return True
    return False


def title_mounting_conflicts_with_feature(title: str, verified_title: str, features: dict[str, str]) -> bool:
    current = mounting_from_title(title)
    expected = normalize_header(feature_by_normalized_name(features, "sposob montazu"))
    verified = mounting_from_title(verified_title)
    return bool(current and expected and current != expected and verified == expected)


def mounting_from_title(title: str) -> str:
    normalized = normalize_header(title)
    if "podtynkowy" in normalized:
        return "podtynkowy"
    if "natynkowy" in normalized:
        return "natynkowy"
    return ""


def title_word_feature_missing(title: str, verified_title: str, features: dict[str, str], normalized_name: str) -> bool:
    feature_value = feature_by_normalized_name(features, normalized_name)
    if normalized_name == "kolor":
        # Kolor porownujemy po rdzeniu (szary/szara -> szar), niezaleznie od formy.
        stem = color_stem(feature_value)
        if stem:
            return stem in normalize_header(verified_title) and stem not in normalize_header(title)
    expected = normalized_title_words(feature_value)
    if not expected:
        return False
    current = normalize_header(title)
    verified = normalize_header(verified_title)
    return any(word in verified and word not in current for word in expected)


def normalized_title_words(value: str) -> set[str]:
    ignored = {"z", "ze", "do", "bez", "nie", "tak", "na", "w", "i"}
    return {
        word
        for word in re.split(r"\s+", normalize_header(value))
        if len(word) > 1 and word not in ignored
    }


def feature_by_normalized_name(features: dict[str, str], normalized_name: str) -> str:
    for name, value in features.items():
        if normalize_header(str(name)) == normalized_name:
            return compact_spaces(str(value))
    return ""


def title_token_conflicts_with_feature(title: str, verified_title: str, feature_value: str, pattern: str) -> bool:
    current = normalized_title_tokens(title, pattern)
    expected = normalized_title_tokens(feature_value, pattern)
    if not expected:
        return False
    verified = normalized_title_tokens(verified_title, pattern)
    if not expected <= verified:
        return False
    if not current:
        return True
    return not expected <= current


def normalized_title_tokens(text: str, pattern: str) -> set[str]:
    tokens: set[str] = set()
    for match in re.findall(pattern, compact_spaces(str(text)), flags=re.IGNORECASE):
        value = match[0] if isinstance(match, tuple) else match
        tokens.add(value.replace(",", ".").replace(" ", "").upper())
    return tokens


def header_columns(worksheet: Any) -> dict[str, int]:
    return {
        compact_spaces(str(worksheet.cell(1, column).value or "")): column
        for column in range(1, worksheet.max_column + 1)
    }


def existing_features_by_sku(worksheet: Any) -> dict[str, dict[str, str]]:
    headers = header_columns(worksheet)
    sku_column = headers["SKU"]
    attribute_columns = [
        (column, header.removeprefix(ATTRIBUTE_PREFIX))
        for header, column in headers.items()
        if header.startswith(ATTRIBUTE_PREFIX)
    ]
    result: dict[str, dict[str, str]] = {}
    for row_number in range(2, worksheet.max_row + 1):
        sku = normalize_sku(worksheet.cell(row_number, sku_column).value)
        if not sku:
            continue
        features: dict[str, str] = {}
        for column, raw_name in attribute_columns:
            raw_value = worksheet.cell(row_number, column).value
            if raw_value is None:
                continue
            value = compact_spaces(str(raw_value))
            if not value or value.lower() in {"nan", "none", "null"}:
                continue
            feature_name = normalize_feature_name(raw_name)
            normalized_value = normalize_feature_value(value, feature_name)
            if normalized_value:
                features[feature_name] = normalized_value
        reclassify_max_bulb_power_feature(features)
        split_multi_color_feature(features)
        result[sku] = features
    return result


def add_missing_attribute_columns(
    worksheet: Any,
    feature_names: list[str],
) -> tuple[dict[str, int], list[str]]:
    headers = header_columns(worksheet)
    product_attributes_column = headers[PRODUCT_ATTRIBUTES_COLUMN]
    existing = {
        normalize_header(header.removeprefix(ATTRIBUTE_PREFIX)): column
        for header, column in headers.items()
        if header.startswith(ATTRIBUTE_PREFIX)
    }
    missing = [name for name in feature_names if normalize_header(name) not in existing]
    if missing:
        worksheet.insert_cols(product_attributes_column, amount=len(missing))
        template_column = product_attributes_column - 1
        for offset, feature_name in enumerate(missing):
            column = product_attributes_column + offset
            source = worksheet.cell(1, template_column)
            target = worksheet.cell(1, column, f"{ATTRIBUTE_PREFIX}{feature_name}")
            if source.has_style:
                target._style = copy(source._style)
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.number_format = source.number_format
            target.protection = copy(source.protection)
            worksheet.column_dimensions[get_column_letter(column)].width = worksheet.column_dimensions[
                get_column_letter(template_column)
            ].width
    headers = header_columns(worksheet)
    attribute_columns = {
        normalize_header(header.removeprefix(ATTRIBUTE_PREFIX)): column
        for header, column in headers.items()
        if header.startswith(ATTRIBUTE_PREFIX)
    }
    return attribute_columns, missing


def title_duplicate_key(title: str) -> str:
    stripped = re.sub(r"\s*Kanlux\s+[\w/.\-]+\s*$", "", compact_spaces(title), flags=re.IGNORECASE).strip()
    return normalize_title_uniqueness_key(stripped)


def disambiguate_duplicate_titles(
    worksheet: Any,
    title_column: int,
    sku_column: int,
    title_updates: dict[str, str],
) -> int:
    """Gdy kilka wierszy ma identyczna nazwe (po odcieciu kodu), bierze rozrozniajaca
    nazwe z pipeline'u (title_updates z CSV baselinkera). Stary tytul legacy gubil
    wymiary/warianty/czujnik, ktore pipeline juz policzyl."""
    titles: dict[int, str] = {}
    skus: dict[int, str] = {}
    rows_by_key: dict[str, list[int]] = {}
    for row_number in range(2, worksheet.max_row + 1):
        title = compact_spaces(str(worksheet.cell(row_number, title_column).value or ""))
        if not title:
            continue
        titles[row_number] = title
        skus[row_number] = normalize_sku(worksheet.cell(row_number, sku_column).value)
        rows_by_key.setdefault(title_duplicate_key(title), []).append(row_number)
    changed = 0
    for group_key, row_numbers in rows_by_key.items():
        if not group_key or len(row_numbers) < 2:
            continue
        for row_number in row_numbers:
            candidate = normalize_seo_title_terms(compact_spaces(title_updates.get(skus[row_number], "")))
            if not candidate or candidate == titles[row_number]:
                continue
            # Pipeline nie rozroznia tego produktu - zostawiamy duplikat do scalenia w sklepie.
            if title_duplicate_key(candidate) == group_key:
                continue
            worksheet.cell(row_number, title_column).value = candidate
            changed += 1
    return changed


def update_workbook(
    input_path: str | Path,
    features_path: str | Path,
    output_path: str | Path,
    titles_path: str | Path | None = None,
) -> dict[str, Any]:
    feature_updates = read_feature_updates(features_path)
    title_updates = read_title_updates(titles_path)
    workbook = load_workbook(input_path)
    worksheet = workbook[workbook.sheetnames[0]]
    existing_features = existing_features_by_sku(worksheet)
    feature_names = sorted(
        {
            name
            for feature_set in [*feature_updates.values(), *existing_features.values()]
            for name in feature_set
        },
        key=normalize_header,
    )
    attribute_columns, added_columns = add_missing_attribute_columns(worksheet, feature_names)
    headers = header_columns(worksheet)
    sku_column = headers["SKU"]
    title_column = headers["Title"]
    product_attributes_column = headers[PRODUCT_ATTRIBUTES_COLUMN]
    all_attribute_columns = sorted(attribute_columns.values())

    matched = 0
    normalized_existing = 0
    retitled = 0
    changed_rows: list[dict[str, Any]] = []
    for row_number in range(2, worksheet.max_row + 1):
        sku = normalize_sku(worksheet.cell(row_number, sku_column).value)
        features = feature_updates.get(sku)
        if features is None:
            features = existing_features.get(sku, {})
            normalized_existing += 1
            update_source = "EXISTING_WOO_NORMALIZED"
        else:
            matched += 1
            update_source = "FRESH_PARAMETERS"
        for column in all_attribute_columns:
            worksheet.cell(row_number, column).value = None
        for feature_name, value in features.items():
            column = attribute_columns[normalize_header(feature_name)]
            worksheet.cell(row_number, column).value = woo_attribute_cell_value(feature_name, value)
        worksheet.cell(row_number, product_attributes_column).value = serialize_product_attributes(list(features))
        title = compact_spaces(str(worksheet.cell(row_number, title_column).value or ""))
        new_title = title_with_verified_product_type(title, features, title_updates.get(sku, ""))
        if new_title != title:
            worksheet.cell(row_number, title_column).value = new_title
            retitled += 1
        changed_rows.append(
            {
                "row": row_number,
                "sku": worksheet.cell(row_number, sku_column).value,
                "attributes_count": len(features),
                "update_source": update_source,
            }
        )

    duplicate_disambiguated = disambiguate_duplicate_titles(worksheet, title_column, sku_column, title_updates)

    output = Path(output_path)
    ensure_dir(output.parent)
    workbook.save(output)
    return {
        "input_rows": worksheet.max_row - 1,
        "feature_update_rows": len(feature_updates),
        "title_update_rows": len(title_updates),
        "matched_rows": matched,
        "unmatched_updates": len(feature_updates) - matched,
        "normalized_existing_rows": normalized_existing,
        "retitled_rows": retitled,
        "duplicate_disambiguated": duplicate_disambiguated,
        "updated_rows": matched + normalized_existing,
        "original_columns": worksheet.max_column - len(added_columns),
        "output_columns": worksheet.max_column,
        "added_attribute_columns": added_columns,
        "changed_rows": changed_rows,
    }


def write_reports(reports_dir: str | Path, summary: dict[str, Any]) -> None:
    reports = ensure_dir(reports_dir)
    serializable = {key: value for key, value in summary.items() if key != "changed_rows"}
    (reports / "summary.json").write_text(json.dumps(serializable, ensure_ascii=False, indent=2), encoding="utf-8")
    pd.DataFrame(summary["changed_rows"]).to_csv(
        reports / "updated_rows.csv",
        index=False,
        encoding="utf-8-sig",
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Aktualizuje atrybuty Kanlux w formacie eksportu WooCommerce.")
    parser.add_argument("--input", default=DEFAULT_INPUT)
    parser.add_argument("--features", default=DEFAULT_FEATURES)
    parser.add_argument("--titles", default=DEFAULT_TITLES)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    parser.add_argument("--reports-dir", default=DEFAULT_REPORTS)
    args = parser.parse_args()

    summary = update_workbook(args.input, args.features, args.output, args.titles)
    write_reports(args.reports_dir, summary)
    print(json.dumps({key: value for key, value in summary.items() if key != "changed_rows"}, ensure_ascii=True))


if __name__ == "__main__":
    main()
