from __future__ import annotations

import csv
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "configs" / "hager_attribute_deduplication.yaml"
HAGER_PATH = ROOT / "output" / "Aparatura-Modulowa-Hager_nazwy_atrybuty_v2.xlsx"
WOO_KNOWLEDGE_PATH = ROOT / "dictionaries" / "woocommerce_catalog_knowledge.yaml"
WOO_EXPORT_PATH = ROOT / "archived_input_files" / "wszystko.csv"
OUTPUT_PATH = ROOT / "output" / "Hager_audyt_powielen_atrybutow_WooCommerce.xlsx"

ATTR_PREFIX = "Atrybut Produktu: "

NAVY = "17365D"
BLUE = "366092"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "EEF5FB"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
ORANGE = "FCE4D6"

THIN_GRAY = Side(style="thin", color="B7B7B7")
GRID_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", clean(value).casefold())
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = text.replace("ł", "l").replace("²", "2").replace("δ", "d")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def populated(value: Any) -> bool:
    return clean(value) != ""


def unique(values: Iterable[Any]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        item = clean(value)
        if not item or item in seen:
            continue
        seen.add(item)
        result.append(item)
    return result


def joined(values: Iterable[Any], limit: int = 12) -> str:
    items = unique(values)
    suffix = f"; +{len(items) - limit}" if len(items) > limit else ""
    return "; ".join(items[:limit]) + suffix


def numeric_tokens(value: Any) -> list[float]:
    text = clean(value).replace(",", ".")
    return [float(token) for token in re.findall(r"\d+(?:\.\d+)?", text)]


def pole_count(value: Any) -> int | None:
    text = norm(value).replace(" ", "")
    if not text:
        return None
    if text in {"1", "1p", "pojedyncze"}:
        return 1
    if text in {"2", "2p", "1p+n", "1pn"}:
        return 2
    if text in {"3", "3p"}:
        return 3
    if text in {"4", "4p", "3p+n", "3pn", "3+1/4p", "3+1"}:
        return 4
    numbers = numeric_tokens(value)
    return int(numbers[0]) if len(numbers) == 1 and numbers[0].is_integer() else None


def normalize_contacts(value: Any) -> str:
    text = norm(value).replace(" ", "").upper()
    translations = {"4/0": "4NO", "2/0": "2NO", "1/1": "1NO+1NC", "0/4": "4NC", "3/1": "3NO+1NC"}
    return translations.get(text, text)


def equivalent(new_attribute: str, new_value: Any, old_attribute: str, old_value: Any) -> bool:
    if not populated(new_value) or not populated(old_value):
        return False
    new_key = norm(new_attribute)
    old_key = norm(old_attribute)
    if "biegun" in new_key and "biegun" in old_key:
        return pole_count(new_value) == pole_count(old_value)
    if "styk" in new_key and "styk" in old_key:
        return normalize_contacts(new_value) == normalize_contacts(old_value)
    if "napiecie cewki" in new_key and old_key == "napiecie v":
        new_numbers = numeric_tokens(new_value)
        old_numbers = numeric_tokens(old_value)
        return bool(new_numbers) and new_numbers[0] in old_numbers
    if "przekroj szyny" in new_key and "przekroj poprzeczny" in old_key:
        new_numbers = numeric_tokens(new_value)
        old_numbers = numeric_tokens(old_value)
        return bool(new_numbers and old_numbers) and new_numbers[0] == old_numbers[0]
    return norm(new_value) == norm(old_value)


def load_hager() -> tuple[list[str], list[dict[str, Any]]]:
    workbook = openpyxl.load_workbook(HAGER_PATH, read_only=True, data_only=True)
    worksheet = workbook["Worksheet"]
    headers = [clean(value) for value in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    rows: list[dict[str, Any]] = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        if not any(populated(value) for value in values):
            continue
        rows.append({headers[index]: values[index] for index in range(min(len(headers), len(values)))})
    workbook.close()
    return headers, rows


def header_lookup(headers: list[str]) -> dict[str, str]:
    return {norm(header.removeprefix(ATTR_PREFIX)): header for header in headers if header.startswith(ATTR_PREFIX)}


def load_woo_knowledge() -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    with WOO_KNOWLEDGE_PATH.open("r", encoding="utf-8") as handle:
        knowledge = yaml.safe_load(handle)
    attributes: dict[str, dict[str, Any]] = {}
    for item in knowledge["top_attributes"]:
        key = norm(item["attribute"])
        previous = attributes.get(key)
        if previous is None or int(item.get("products_count", 0) or 0) > int(previous.get("products_count", 0) or 0):
            attributes[key] = item
    return attributes, knowledge["summary"]


def load_global_flags() -> dict[str, Counter[str]]:
    result: dict[str, Counter[str]] = defaultdict(Counter)
    with WOO_EXPORT_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            for slot in range(1, 31):
                name = clean(row.get(f"Nazwa atrybutu {slot}"))
                if name:
                    result[norm(name)][clean(row.get(f"Atrybut {slot}: globalny"))] += 1
    return result


def classify(decision: dict[str, Any]) -> str:
    verdict = clean(decision["verdict"])
    if clean(decision.get("existing_slug")):
        return "UŻYĆ ISTNIEJĄCEGO GLOBALNEGO"
    if norm(decision["attribute"]) == "kompatybilnosc":
        return "UŻYĆ ISTNIEJĄCEGO LOKALNEGO"
    if verdict.startswith("NIE TWORZYĆ"):
        return "SCALIĆ Z INNYM NOWYM"
    return "NOWY GLOBALNY"


def set_title(ws, title: str, subtitle: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    ws.cell(1, 1, title)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).font = Font(color=WHITE, bold=True, size=16)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.cell(2, 1).font = Font(color=NAVY, italic=True, size=10)
    ws.cell(2, 1).alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 34


def style_header(ws, row: int, end_col: int) -> None:
    for col in range(1, end_col + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = GRID_BORDER
    ws.row_dimensions[row].height = 36


def style_rows(ws, start_row: int, end_row: int, end_col: int) -> None:
    for row in range(start_row, end_row + 1):
        fill = PatternFill("solid", fgColor=WHITE if row % 2 == 0 else PALE_BLUE)
        for col in range(1, end_col + 1):
            cell = ws.cell(row, col)
            cell.fill = fill
            cell.border = GRID_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def set_widths(ws, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def finalize_table_sheet(ws, header_row: int, end_col: int, end_row: int) -> None:
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(end_col)}{max(header_row, end_row)}"
    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = f"1:{header_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def evidence_for(
    decision: dict[str, Any],
    rows: list[dict[str, Any]],
    headers_by_name: dict[str, str],
    woo_attributes: dict[str, dict[str, Any]],
    global_flags: dict[str, Counter[str]],
) -> dict[str, Any]:
    attribute = clean(decision["attribute"])
    existing = clean(decision.get("existing_attribute"))
    new_header = headers_by_name.get(norm(attribute), "")
    old_header = headers_by_name.get(norm(existing), "") if existing else ""
    new_values = [row.get(new_header) for row in rows if new_header and populated(row.get(new_header))]
    old_values_same = [row.get(old_header) for row in rows if old_header and populated(row.get(old_header))]
    both_rows = [row for row in rows if new_header and old_header and populated(row.get(new_header)) and populated(row.get(old_header))]
    compatible = sum(1 for row in both_rows if equivalent(attribute, row.get(new_header), existing, row.get(old_header)))
    pairs = joined((f"{clean(row.get(new_header))} ↔ {clean(row.get(old_header))}" for row in both_rows), limit=8)
    woo = woo_attributes.get(norm(existing), {}) if existing else {}
    flags = global_flags.get(norm(existing), Counter())
    global_status = "GLOBALNY" if flags.get("1", 0) >= flags.get("0", 0) and flags.get("1", 0) else ("LOKALNY" if flags.get("0", 0) else "BRAK / NOWY")
    return {
        "new_count": len(new_values),
        "new_values": joined(new_values),
        "existing_count_hager": len(old_values_same),
        "both_count": len(both_rows),
        "compatible_count": compatible,
        "pairs": pairs,
        "woo_products": int(woo.get("products_count", 0) or 0),
        "woo_values": clean(woo.get("top_values")),
        "woo_global": global_status,
    }


def write_decisions_sheet(
    wb: Workbook,
    decisions: list[dict[str, Any]],
    evidence: dict[str, dict[str, Any]],
) -> None:
    ws = wb.create_sheet("Decyzje")
    set_title(
        ws,
        "Audyt 23 nowych atrybutów Hager",
        "Werdykt uwzględnia nazwę, wartości w pełnym WooCommerce, wartości na tych samych SKU Hager oraz terminologię producenta.",
        16,
    )
    headers = [
        "Nowy atrybut", "Grupa decyzji", "Werdykt", "Istniejący odpowiednik Woo", "Istniejący slug",
        "Docelowa nazwa", "Pewność", "Produkty Hager z wartością", "Wartości Hager",
        "Produkty Woo z odpowiednikiem", "Wartości odpowiednika w Woo", "Status odpowiednika",
        "Uzasadnienie", "Normalizacja", "Źródła Hager", "Decyzja użytkownika",
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(5, col, header)
    style_header(ws, 5, len(headers))

    for row_index, decision in enumerate(decisions, 6):
        item = evidence[norm(decision["attribute"])]
        values = [
            decision["attribute"], classify(decision), decision["verdict"], decision.get("existing_attribute", ""),
            decision.get("existing_slug", ""), decision["final_name"], decision["confidence"], item["new_count"],
            item["new_values"], item["woo_products"], item["woo_values"], item["woo_global"], decision["rationale"],
            decision["normalization"], "\n".join(decision.get("source_urls", [])), "",
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)

    end_row = 5 + len(decisions)
    style_rows(ws, 6, end_row, len(headers))
    group_colors = {
        "NOWY GLOBALNY": GREEN,
        "UŻYĆ ISTNIEJĄCEGO GLOBALNEGO": YELLOW,
        "UŻYĆ ISTNIEJĄCEGO LOKALNEGO": ORANGE,
        "SCALIĆ Z INNYM NOWYM": LIGHT_BLUE,
    }
    for row in range(6, end_row + 1):
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=group_colors.get(clean(ws.cell(row, 2).value), GRAY))
        ws.cell(row, 7).fill = PatternFill("solid", fgColor=GREEN if ws.cell(row, 7).value == "WYSOKA" else YELLOW)
        if ws.cell(row, 12).value == "GLOBALNY":
            ws.cell(row, 12).fill = PatternFill("solid", fgColor=GREEN)
        elif ws.cell(row, 12).value == "LOKALNY":
            ws.cell(row, 12).fill = PatternFill("solid", fgColor=ORANGE)
    set_widths(ws, {1: 38, 2: 29, 3: 35, 4: 34, 5: 34, 6: 45, 7: 12, 8: 14, 9: 46, 10: 15, 11: 52, 12: 15, 13: 58, 14: 58, 15: 54, 16: 20})
    finalize_table_sheet(ws, 5, len(headers), end_row)


def write_value_evidence_sheet(wb: Workbook, decisions: list[dict[str, Any]], evidence: dict[str, dict[str, Any]]) -> None:
    ws = wb.create_sheet("Dowody z wartości")
    set_title(
        ws,
        "Dowody na powielenia i konflikty wartości",
        "Porównanie dotyczy tych samych produktów Hager oraz pełnego słownika 10 948 produktów WooCommerce.",
        12,
    )
    headers = [
        "Nowy atrybut", "Porównany atrybut", "Nowe wartości Hager", "Wartości istniejące w Woo",
        "Nowe pole: liczba SKU", "Istniejące pole w Hager: SKU", "Oba pola na tym samym SKU",
        "Zgodne / równoważne wartości", "Zgodność wśród obu pól", "Przykładowe pary na SKU",
        "Wniosek", "Pewność",
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(5, col, header)
    style_header(ws, 5, len(headers))
    compared = [decision for decision in decisions if clean(decision.get("existing_attribute"))]
    for row_index, decision in enumerate(compared, 6):
        item = evidence[norm(decision["attribute"])]
        conclusion = decision["verdict"]
        values = [
            decision["attribute"], decision.get("existing_attribute", ""), item["new_values"], item["woo_values"],
            item["new_count"], item["existing_count_hager"], item["both_count"], item["compatible_count"],
            f"=IFERROR(H{row_index}/G{row_index},0)", item["pairs"], conclusion, decision["confidence"],
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
    end_row = 5 + len(compared)
    style_rows(ws, 6, end_row, len(headers))
    for row in range(6, end_row + 1):
        ws.cell(row, 9).number_format = "0.0%"
        ws.cell(row, 11).fill = PatternFill("solid", fgColor=YELLOW if "NIE TWORZYĆ" in clean(ws.cell(row, 11).value) else LIGHT_BLUE)
    ws.conditional_formatting.add(
        f"I6:I{end_row}",
        CellIsRule(operator="greaterThanOrEqual", formula=["0.8"], fill=PatternFill("solid", fgColor=GREEN)),
    )
    set_widths(ws, {1: 38, 2: 38, 3: 46, 4: 55, 5: 14, 6: 16, 7: 16, 8: 16, 9: 14, 10: 54, 11: 39, 12: 12})
    finalize_table_sheet(ws, 5, len(headers), end_row)


def write_issues_sheet(wb: Workbook, issues: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("Problemy danych")
    set_title(
        ws,
        "Błędy i niejednoznaczności wykryte podczas audytu",
        "To nie są wyłącznie kwestie nazewnictwa. Wskazane wartości należy poprawić przed finalnym importem.",
        8,
    )
    headers = ["SKU", "Atrybut", "Obecna wartość", "Proponowana wartość", "Waga", "Powód", "Źródło", "Status korekty"]
    for col, header in enumerate(headers, 1):
        ws.cell(5, col, header)
    style_header(ws, 5, len(headers))
    for row_index, issue in enumerate(issues, 6):
        values = [
            issue["sku"], issue["attribute"], issue["current_value"], issue["proposed_value"], issue["severity"],
            issue["reason"], issue.get("source_url", ""), "DO POPRAWY",
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
    end_row = 5 + len(issues)
    style_rows(ws, 6, end_row, len(headers))
    for row in range(6, end_row + 1):
        severity = clean(ws.cell(row, 5).value)
        ws.cell(row, 5).fill = PatternFill("solid", fgColor=RED if severity in {"BŁĄD", "BŁĄD MODELU"} else YELLOW)
        ws.cell(row, 8).fill = PatternFill("solid", fgColor=YELLOW)
        ws.row_dimensions[row].height = 48
    set_widths(ws, {1: 25, 2: 42, 3: 34, 4: 48, 5: 16, 6: 62, 7: 56, 8: 18})
    finalize_table_sheet(ws, 5, len(headers), end_row)


def write_woo_dictionary_sheet(
    wb: Workbook,
    decisions: list[dict[str, Any]],
    woo_attributes: dict[str, dict[str, Any]],
    global_flags: dict[str, Counter[str]],
) -> None:
    ws = wb.create_sheet("Istniejące Woo")
    set_title(
        ws,
        "Istniejące atrybuty WooCommerce istotne dla audytu",
        "Dane pochodzą z lokalnego słownika pełnego katalogu oraz eksportu WooCommerce. Pokazano faktyczne wartości, a nie tylko nazwy.",
        8,
    )
    headers = ["Atrybut Woo", "Produkty", "Różne wartości", "Najczęstsze wartości", "Globalny/lokalny", "Wystąpienia globalne", "Wystąpienia lokalne", "Powiązany nowy atrybut"]
    for col, header in enumerate(headers, 1):
        ws.cell(5, col, header)
    style_header(ws, 5, len(headers))
    names: dict[str, set[str]] = defaultdict(set)
    for decision in decisions:
        existing = clean(decision.get("existing_attribute"))
        if existing:
            names[norm(existing)].add(clean(decision["attribute"]))
    for row_index, key in enumerate(sorted(names, key=lambda item: clean(woo_attributes.get(item, {}).get("attribute", item))), 6):
        item = woo_attributes.get(key, {})
        flags = global_flags.get(key, Counter())
        status = "GLOBALNY" if flags.get("1", 0) >= flags.get("0", 0) and flags.get("1", 0) else ("LOKALNY" if flags.get("0", 0) else "BRAK W SŁOWNIKU")
        values = [
            item.get("attribute", key), int(item.get("products_count", 0) or 0), int(item.get("unique_values", 0) or 0),
            item.get("top_values", ""), status, flags.get("1", 0), flags.get("0", 0), "; ".join(sorted(names[key])),
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
    end_row = 5 + len(names)
    style_rows(ws, 6, end_row, len(headers))
    for row in range(6, end_row + 1):
        status = clean(ws.cell(row, 5).value)
        ws.cell(row, 5).fill = PatternFill("solid", fgColor=GREEN if status == "GLOBALNY" else ORANGE)
    set_widths(ws, {1: 40, 2: 14, 3: 15, 4: 75, 5: 18, 6: 16, 7: 16, 8: 45})
    finalize_table_sheet(ws, 5, len(headers), end_row)


def write_method_sheet(wb: Workbook, woo_summary: dict[str, Any]) -> None:
    ws = wb.create_sheet("Metodyka i źródła")
    set_title(ws, "Metodyka audytu", "Zakres porównania i reguły użyte do odróżnienia duplikatu od podobnej wartości.", 6)
    headers = ["Obszar", "Źródło / reguła", "Zakres", "Dlaczego użyte", "Ograniczenie", "Wniosek"]
    for col, header in enumerate(headers, 1):
        ws.cell(5, col, header)
    style_header(ws, 5, len(headers))
    rows = [
        ("Pełny WooCommerce", "dictionaries/woocommerce_catalog_knowledge.yaml", f"{woo_summary.get('rows', 0)} produktów, {woo_summary.get('attribute_names', 0)} nazw atrybutów", "Nazwy, liczność i najczęstsze rzeczywiste wartości.", "Słownik jest migawką lokalnego eksportu.", "Podstawa do wykrycia równoważnych atrybutów."),
        ("Eksport WooCommerce", "archived_input_files/wszystko.csv", "30 slotów atrybutów na produkt", "Rozróżnienie globalnych i lokalnych atrybutów.", "W tym eksporcie komórki wartości są puste; wartości pochodzą ze słownika wiedzy.", "Potwierdza status taksonomii."),
        ("Hager v2", "output/Aparatura-Modulowa-Hager_nazwy_atrybuty_v2.xlsx", "153 produkty, 23 nowe atrybuty", "Porównanie wartości na tych samych SKU i wykrycie konfliktów modelu.", "Dane są przed finalną korektą.", "Wykryto również błędy wartości."),
        ("Nazwy producenta", "Oficjalne karty produktowe hager.com/pl", "MCB/RCD/SPD/styczniki/szyny/bezpieczniki", "Weryfikacja terminologii i wartości technicznych.", "Nazwy Hager bywają skrótowe i zależne od rodziny.", "Proponowane etykiety są czytelne dla sklepu i zgodne technicznie."),
        ("Reguła duplikatu", "Nazwa + wartości + użycie na tych samych SKU", "Nie opierano się wyłącznie na podobieństwie tekstu.", "Tak/Nie lub 230 może wystąpić w wielu niezależnych parametrach.", "Wspólna wartość bez wspólnego znaczenia nie oznacza duplikatu."),
        ("Reguła sklepu", "Jeden filtr na jedną decyzję zakupową", "Powiązane pola techniczne scalano tylko, gdy dwa filtry byłyby mylące.", "Czasem karta producenta rozróżnia więcej pól niż potrzebuje sklep.", "Stąd konsolidacja liczby/układu biegunów i Inc/Icc."),
    ]
    for row_index, values in enumerate(rows, 6):
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
    end_row = 5 + len(rows)
    style_rows(ws, 6, end_row, len(headers))
    for row in range(6, end_row + 1):
        ws.row_dimensions[row].height = 55
    set_widths(ws, {1: 24, 2: 55, 3: 38, 4: 52, 5: 46, 6: 52})
    finalize_table_sheet(ws, 5, len(headers), end_row)


def write_summary_sheet(wb: Workbook, decision_count: int, issue_count: int) -> None:
    ws = wb.active
    ws.title = "Podsumowanie"
    set_title(
        ws,
        "Hager — audyt powieleń atrybutów WooCommerce",
        "Wynik ponownego sprawdzenia nazw i rzeczywistych wartości. Ten raport zastępuje wcześniejsze założenie, że prawie wszystkie 23 pola są nowymi globalnymi atrybutami.",
        8,
    )
    cards = [
        ("Sprawdzone nowe pola", "=COUNTA('Decyzje'!A6:A100)"),
        ("Nowe globalne po scaleniu", '=COUNTIF(\'Decyzje\'!B6:B100,"NOWY GLOBALNY")'),
        ("Użyć istniejących globalnych", '=COUNTIF(\'Decyzje\'!B6:B100,"UŻYĆ ISTNIEJĄCEGO GLOBALNEGO")'),
        ("Nie tworzyć osobno", '=COUNTA(\'Decyzje\'!B6:B100)-COUNTIF(\'Decyzje\'!B6:B100,"NOWY GLOBALNY")'),
        ("Istniejący atrybut lokalny", '=COUNTIF(\'Decyzje\'!B6:B100,"UŻYĆ ISTNIEJĄCEGO LOKALNEGO")'),
        ("Scalić z innym nowym", '=COUNTIF(\'Decyzje\'!B6:B100,"SCALIĆ Z INNYM NOWYM")'),
        ("Problemy danych", "=COUNTA('Problemy danych'!A6:A100)"),
        ("Atrybuty z wysoką pewnością", '=COUNTIF(\'Decyzje\'!G6:G100,"WYSOKA")'),
    ]
    for index, (label, formula) in enumerate(cards):
        row = 4 if index < 4 else 7
        col = 1 + (index % 4) * 2
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.cell(row, col, label)
        ws.cell(row, col).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(row, col).font = Font(color=WHITE, bold=True, size=10)
        ws.cell(row, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        ws.cell(row + 1, col, formula)
        ws.cell(row + 1, col).fill = PatternFill("solid", fgColor=PALE_BLUE)
        ws.cell(row + 1, col).font = Font(color=NAVY, bold=True, size=20)
        ws.cell(row + 1, col).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 32
        ws.row_dimensions[row + 1].height = 34

    ws.merge_cells("A11:H11")
    ws["A11"] = "Najważniejsza zmiana decyzji"
    ws["A11"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A11"].font = Font(color=WHITE, bold=True, size=12)
    ws.merge_cells("A12:H14")
    ws["A12"] = (
        "Nie należy tworzyć 23 osobnych nowych atrybutów. Po porównaniu wartości rekomendacja to 16 nowych globalnych pól, "
        "wykorzystanie 5 istniejących globalnych taksonomii, użycie 1 istniejącego pola lokalnego oraz scalenie Icc z Inc. "
        "Najpierw trzeba poprawić wskazane błędy danych."
    )
    ws["A12"].fill = PatternFill("solid", fgColor=YELLOW)
    ws["A12"].font = Font(color=NAVY, bold=True, size=11)
    ws["A12"].alignment = Alignment(vertical="center", wrap_text=True)

    ws.merge_cells("A16:H16")
    ws["A16"] = "Pewne powielenia / pola do ponownego użycia"
    ws["A16"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A16"].font = Font(color=WHITE, bold=True)
    duplicates = [
        ("Układ biegunów", "wykorzystać pa_liczba-biegunow i ujednolicić wartości do 1P, 1P+N, 3P+N itd."),
        ("Typ prądu różnicowego", "poprawić istniejący pa_typ-wylacznika; zostawić AC/A/A-HI, a rodziny przenieść do Typu produktu."),
        ("Napięcie cewki", "na obecnym etapie wykorzystać Napięcie [V]; nowe pole jest niepełne dla 230 V AC / 110 V DC."),
        ("Konfiguracja styków", "wykorzystać pa_ilosc-stykow-zwiernych-rozwiernych i zmienić etykietę oraz format wartości."),
        ("Przekrój szyny", "wykorzystać pa_znamionowy-przekroj-poprzeczny — zgodność wartości na 7/7 produktach."),
        ("Kompatybilność", "użyć istniejącego lokalnego Kompatybilny z; nie tworzyć globalnych terminów ze zdań."),
        ("Icc i Inc", "utworzyć jeden atrybut Znamionowy warunkowy prąd zwarciowy [kA]."),
    ]
    for row_index, (name, action) in enumerate(duplicates, 17):
        ws.cell(row_index, 1, name)
        ws.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=8)
        ws.cell(row_index, 2, action)
        for col in range(1, 9):
            cell = ws.cell(row_index, col)
            cell.fill = PatternFill("solid", fgColor=WHITE if row_index % 2 else PALE_BLUE)
            cell.border = GRID_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row_index, 1).font = Font(bold=True, color=NAVY)
        ws.row_dimensions[row_index].height = 28

    ws.merge_cells("A26:H26")
    ws["A26"] = "Krytyczne poprawki przed importem"
    ws["A26"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A26"].font = Font(color=WHITE, bold=True)
    critical = [
        "SPA931: Uc 255 V → 335 V.",
        "SPA931: Układ biegunów 4P → 3P+N.",
        "EPN510/EPN520: Napięcie cewki 230 jest niepełne; sterowanie obejmuje też 110 V DC.",
        "SA463/SA480: „Styk pomocniczy” nie jest konfiguracją styków NO/NC.",
        "Typ Wyłącznika nie może jednocześnie przechowywać AC/A oraz Nadprądowy/Różnicowoprądowy.",
    ]
    for row_index, note in enumerate(critical, 27):
        ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=8)
        ws.cell(row_index, 1, note)
        ws.cell(row_index, 1).fill = PatternFill("solid", fgColor=RED if row_index in {27, 28} else ORANGE)
        ws.cell(row_index, 1).border = GRID_BORDER
        ws.cell(row_index, 1).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row_index].height = 24

    set_widths(ws, {1: 22, 2: 16, 3: 22, 4: 16, 5: 22, 6: 16, 7: 24, 8: 16})
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    assert decision_count == 23
    assert issue_count >= 6


def main() -> None:
    with CONFIG_PATH.open("r", encoding="utf-8") as handle:
        config = yaml.safe_load(handle)
    decisions = config["decisions"]
    issues = config["data_issues"]
    headers, hager_rows = load_hager()
    headers_by_name = header_lookup(headers)
    woo_attributes, woo_summary = load_woo_knowledge()
    global_flags = load_global_flags()

    evidence = {
        norm(decision["attribute"]): evidence_for(decision, hager_rows, headers_by_name, woo_attributes, global_flags)
        for decision in decisions
    }

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.properties.creator = "AlkanInator"
    wb.properties.title = "Hager — audyt powieleń atrybutów WooCommerce"

    write_decisions_sheet(wb, decisions, evidence)
    write_value_evidence_sheet(wb, decisions, evidence)
    write_issues_sheet(wb, issues)
    write_woo_dictionary_sheet(wb, decisions, woo_attributes, global_flags)
    write_method_sheet(wb, woo_summary)
    write_summary_sheet(wb, len(decisions), len(issues))

    order = ["Podsumowanie", "Decyzje", "Dowody z wartości", "Problemy danych", "Istniejące Woo", "Metodyka i źródła"]
    wb._sheets = [wb[name] for name in order]
    wb.active = 0
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)

    groups = Counter(classify(decision) for decision in decisions)
    print(f"OUTPUT={OUTPUT_PATH}")
    print(f"PRODUCTS_HAGER={len(hager_rows)}")
    print(f"DECISIONS={len(decisions)}")
    print(f"GROUPS={dict(groups)}")
    print(f"DATA_ISSUES={len(issues)}")


if __name__ == "__main__":
    main()
