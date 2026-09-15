from __future__ import annotations

import re
from collections import OrderedDict
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "configs" / "hager_new_woocommerce_attributes.yaml"
ORIGINAL_PATH = ROOT / "input" / "Aparatura-Modulowa-Hager.xlsx"
ENRICHED_PATH = ROOT / "output" / "Aparatura-Modulowa-Hager_nazwy_atrybuty_v2.xlsx"
PRIORITY_PATH = ROOT / "output" / "Hager_aparatura_kategorie_atrybuty_priorytety.xlsx"
OUTPUT_PATH = ROOT / "output" / "Hager_raport_nowe_atrybuty_vs_WooCommerce.xlsx"

ATTR_PREFIX = "Atrybut Produktu: "
PRIORITY_RANK = {"P0": 0, "P1": 1, "P2": 2}

NAVY = "17365D"
BLUE = "366092"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "EEF5FB"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
DARK_GRAY = "666666"
ORANGE = "FCE4D6"

THIN_GRAY = Side(style="thin", color="B7B7B7")
GRID_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def populated(value: Any) -> bool:
    return clean(value) != ""


def unique_join(values: Iterable[Any], limit: int = 5) -> str:
    found: list[str] = []
    seen: set[str] = set()
    for value in values:
        item = clean(value)
        if not item or item in seen:
            continue
        seen.add(item)
        found.append(item)
        if len(found) >= limit:
            break
    return "; ".join(found)


def load_sheet_rows(path: Path, sheet_name: str | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    raw_headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [clean(value) for value in raw_headers]
    rows: list[dict[str, Any]] = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        if not any(populated(value) for value in values):
            continue
        rows.append({headers[index]: values[index] for index in range(min(len(headers), len(values)))})
    workbook.close()
    return headers, rows


def load_matrix_rows(path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["Kategorie i atrybuty"]
    headers = [clean(value) for value in next(worksheet.iter_rows(min_row=5, max_row=5, values_only=True))]
    rows: list[dict[str, Any]] = []
    for values in worksheet.iter_rows(min_row=6, values_only=True):
        if not any(populated(value) for value in values):
            continue
        rows.append({headers[index]: values[index] for index in range(min(len(headers), len(values)))})
    workbook.close()
    return rows


def product_map(rows: list[dict[str, Any]]) -> OrderedDict[str, dict[str, Any]]:
    result: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for row in rows:
        sku = clean(row.get("SKU"))
        if sku:
            result[sku] = row
    return result


def parse_declared_taxonomies(rows: list[dict[str, Any]]) -> set[str]:
    declared: set[str] = set()
    pattern = re.compile(r's:\d+:"(pa_[^"]+)";a:6:')
    for row in rows:
        declared.update(pattern.findall(clean(row.get("Product Attributes"))))
    return declared


def count_values(rows: Iterable[dict[str, Any]], header: str) -> int:
    return sum(1 for row in rows if populated(row.get(header)))


def set_title(ws, title: str, subtitle: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, title)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(color=WHITE, bold=True, size=16)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    cell = ws.cell(2, 1, subtitle)
    cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    cell.font = Font(color=NAVY, italic=True, size=10)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 34


def style_header(ws, row: int, start_col: int, end_col: int) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = GRID_BORDER
    ws.row_dimensions[row].height = 34


def style_data_range(ws, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    for row in range(start_row, end_row + 1):
        fill = PatternFill("solid", fgColor=WHITE if row % 2 == 0 else PALE_BLUE)
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row, col)
            cell.fill = fill
            cell.border = GRID_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def set_widths(ws, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def add_autofilter_and_freeze(ws, header_row: int, end_col: int, end_row: int) -> None:
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(end_col)}{max(header_row, end_row)}"
    ws.freeze_panes = f"A{header_row + 1}"


def write_new_attributes_sheet(
    wb: Workbook,
    configured: list[dict[str, Any]],
    enriched_rows: list[dict[str, Any]],
    declared_taxonomies: set[str],
    matrix_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    ws = wb.create_sheet("Nowe atrybuty")
    set_title(
        ws,
        "Nowe atrybuty Hager względem WooCommerce",
        "Nowy atrybut = kolumna nieobecna w pierwotnym eksporcie WooCommerce. Pokrycie liczone dla aktualnych 153 produktów Hager.",
        20,
    )

    headers = [
        "Atrybut",
        "Proponowany slug",
        "Taxonomia WooCommerce",
        "Był w pierwotnym Woo",
        "Nowa kolumna w Hager",
        "Zadeklarowany w Product Attributes",
        "Rekomendacja",
        "Najwyższy priorytet",
        "Rola / zastosowanie",
        "Produkty z wartością",
        "Wszystkie produkty",
        "Pokrycie",
        "Liczba różnych wartości",
        "Kategorie z użyciem",
        "Liczba kategorii",
        "Przykładowe wartości",
        "Uzasadnienie rekomendacji",
        "Wymagane działanie",
        "Status techniczny",
        "Decyzja użytkownika",
    ]
    for index, header in enumerate(headers, 1):
        ws.cell(5, index, header)
    style_header(ws, 5, 1, len(headers))

    total = len(enriched_rows)
    report_rows: list[dict[str, Any]] = []
    for row_index, item in enumerate(configured, 6):
        attribute = clean(item["attribute"])
        header = ATTR_PREFIX + attribute
        slug = clean(item["slug"])
        taxonomy = f"pa_{slug}"
        values = [row.get(header) for row in enriched_rows if populated(row.get(header))]
        distinct = sorted({clean(value) for value in values})
        related = [row for row in matrix_rows if clean(row.get("Atrybut")) == attribute]
        active_related = [row for row in related if int(row.get("Produkty z wartością") or 0) > 0]
        priorities = [clean(row.get("Priorytet")) for row in related if clean(row.get("Priorytet")) in PRIORITY_RANK]
        highest = min(priorities, key=lambda value: PRIORITY_RANK[value]) if priorities else "—"
        categories = sorted({clean(row.get("Ścieżka kategorii")) for row in active_related})
        roles = unique_join((row.get("Rola") for row in related), limit=4) or "Opis techniczny"
        recommendation = clean(item["recommendation"])
        declared = taxonomy in declared_taxonomies
        action = (
            f"Utworzyć globalny atrybut „{attribute}” ze slugiem „{slug}”, utworzyć terminy i dodać mapowanie do importu."
            if recommendation == "UTWORZYĆ GLOBALNY"
            else "Najpierw zbudować kontrolowany słownik krótkich wartości; do tego czasu importować jako pole lokalne lub opisowe."
        )
        status = "GOTOWY DO UTWORZENIA" if recommendation == "UTWORZYĆ GLOBALNY" else "WYMAGA NORMALIZACJI"

        values_out = [
            attribute,
            slug,
            taxonomy,
            "Nie",
            "Tak",
            "Tak" if declared else "Nie",
            recommendation,
            highest,
            roles,
            len(values),
            total,
            f"=IFERROR(J{row_index}/K{row_index},0)",
            len(distinct),
            "\n".join(categories),
            len(categories),
            unique_join(distinct, limit=6),
            clean(item["reason"]),
            action,
            status,
            "",
        ]
        for col_index, value in enumerate(values_out, 1):
            ws.cell(row_index, col_index, value)
        report_rows.append(
            {
                "attribute": attribute,
                "slug": slug,
                "taxonomy": taxonomy,
                "recommendation": recommendation,
                "priority": highest,
                "count": len(values),
                "categories": categories,
                "declared": declared,
            }
        )

    end_row = 5 + len(configured)
    style_data_range(ws, 6, end_row, 1, len(headers))
    for row in range(6, end_row + 1):
        ws.cell(row, 12).number_format = "0.0%"
        priority = clean(ws.cell(row, 8).value)
        ws.cell(row, 8).fill = PatternFill("solid", fgColor={"P0": RED, "P1": YELLOW, "P2": GREEN}.get(priority, GRAY))
        recommendation = clean(ws.cell(row, 7).value)
        ws.cell(row, 7).fill = PatternFill("solid", fgColor=GREEN if recommendation == "UTWORZYĆ GLOBALNY" else YELLOW)
        ws.cell(row, 6).fill = PatternFill("solid", fgColor=RED if ws.cell(row, 6).value == "Nie" else GREEN)
        ws.cell(row, 19).fill = PatternFill("solid", fgColor=GREEN if "GOTOWY" in clean(ws.cell(row, 19).value) else YELLOW)

    decision_validation = DataValidation(type="list", formula1='"ZAAKCEPTOWAĆ,ODROCZYĆ,ODRZUCIĆ"', allow_blank=True)
    decision_validation.error = "Wybierz wartość z listy."
    decision_validation.errorTitle = "Nieprawidłowa decyzja"
    decision_validation.prompt = "Opcjonalna decyzja wdrożeniowa."
    decision_validation.promptTitle = "Decyzja"
    ws.add_data_validation(decision_validation)
    decision_validation.add(f"T6:T{end_row}")

    ws.conditional_formatting.add(
        f"L6:L{end_row}",
        CellIsRule(operator="lessThan", formula=["0.5"], fill=PatternFill("solid", fgColor=YELLOW)),
    )
    set_widths(
        ws,
        {
            1: 37, 2: 24, 3: 28, 4: 14, 5: 15, 6: 19, 7: 27, 8: 13, 9: 24, 10: 13,
            11: 12, 12: 11, 13: 14, 14: 43, 15: 12, 16: 38, 17: 42, 18: 48, 19: 24, 20: 20,
        },
    )
    add_autofilter_and_freeze(ws, 5, len(headers), end_row)
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A5:T{end_row}"
    ws.print_title_rows = "1:5"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return report_rows


def write_category_usage_sheet(wb: Workbook, new_names: set[str], matrix_rows: list[dict[str, Any]]) -> int:
    ws = wb.create_sheet("Użycie w kategoriach")
    set_title(
        ws,
        "Użycie nowych atrybutów w kategoriach Hager",
        "Matryca pokazuje, gdzie atrybut jest potrzebny, jaki ma priorytet i jak wygląda jego pokrycie w aktualnym pliku.",
        12,
    )
    headers = [
        "Atrybut", "Ścieżka kategorii", "Status kategorii", "Priorytet", "Wymagane", "W nazwie",
        "Rola", "Produkty z wartością", "Produkty w kategorii", "Pokrycie", "Przykładowe wartości", "Status pokrycia",
    ]
    for index, header in enumerate(headers, 1):
        ws.cell(5, index, header)
    style_header(ws, 5, 1, len(headers))

    selected = [row for row in matrix_rows if clean(row.get("Atrybut")) in new_names]
    selected.sort(key=lambda row: (clean(row.get("Atrybut")), PRIORITY_RANK.get(clean(row.get("Priorytet")), 9), clean(row.get("Ścieżka kategorii"))))
    for row_index, source in enumerate(selected, 6):
        values = [
            source.get("Atrybut"), source.get("Ścieżka kategorii"), source.get("Status kategorii"), source.get("Priorytet"),
            source.get("Wymagane"), source.get("W nazwie"), source.get("Rola"), source.get("Produkty z wartością"),
            source.get("Produkty w kategorii"), source.get("Pokrycie"), source.get("Przykładowe wartości"), source.get("Status pokrycia"),
        ]
        for col_index, value in enumerate(values, 1):
            ws.cell(row_index, col_index, value)

    end_row = 5 + len(selected)
    style_data_range(ws, 6, end_row, 1, len(headers))
    for row in range(6, end_row + 1):
        ws.cell(row, 10).number_format = "0.0%"
        priority = clean(ws.cell(row, 4).value)
        ws.cell(row, 4).fill = PatternFill("solid", fgColor={"P0": RED, "P1": YELLOW, "P2": GREEN}.get(priority, GRAY))
        if clean(ws.cell(row, 12).value) != "OK":
            ws.cell(row, 12).fill = PatternFill("solid", fgColor=YELLOW)
    set_widths(ws, {1: 38, 2: 50, 3: 18, 4: 12, 5: 12, 6: 12, 7: 24, 8: 15, 9: 15, 10: 12, 11: 42, 12: 20})
    add_autofilter_and_freeze(ws, 5, len(headers), end_row)
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = "1:5"
    return len(selected)


def write_existing_enriched_sheet(
    wb: Workbook,
    original_headers: list[str],
    original_products: OrderedDict[str, dict[str, Any]],
    enriched_products: OrderedDict[str, dict[str, Any]],
) -> int:
    ws = wb.create_sheet("Istniejące uzupełnione")
    set_title(
        ws,
        "Istniejące atrybuty WooCommerce uzupełnione w pliku Hager",
        "Te atrybuty nie są nowe dla WooCommerce. Zwiększyła się wyłącznie liczba produktów z wartością wśród 96 SKU obecnych w obu plikach.",
        11,
    )
    headers = [
        "Atrybut", "Przed: produkty z wartością", "Po: te same 96 SKU", "Przyrost", "Pokrycie przed",
        "Pokrycie po (96 SKU)", "Po: wszystkie 153 SKU", "Pokrycie po (153 SKU)", "Przykładowe wartości po",
        "Klasyfikacja", "Działanie",
    ]
    for index, header in enumerate(headers, 1):
        ws.cell(5, index, header)
    style_header(ws, 5, 1, len(headers))

    common_skus = [sku for sku in original_products if sku in enriched_products]
    rows_out: list[tuple[Any, ...]] = []
    for header in original_headers:
        if not header.startswith(ATTR_PREFIX):
            continue
        attribute = header[len(ATTR_PREFIX):]
        before = sum(1 for sku in common_skus if populated(original_products[sku].get(header)))
        after_same = sum(1 for sku in common_skus if populated(enriched_products[sku].get(header)))
        if after_same <= before:
            continue
        after_all = sum(1 for row in enriched_products.values() if populated(row.get(header)))
        samples = unique_join((row.get(header) for row in enriched_products.values()), limit=6)
        rows_out.append((attribute, before, after_same, after_same - before, after_all, samples))
    rows_out.sort(key=lambda item: (-int(item[3]), clean(item[0])))

    common_count = len(common_skus)
    all_count = len(enriched_products)
    for row_index, (attribute, before, after_same, increase, after_all, samples) in enumerate(rows_out, 6):
        values = [
            attribute, before, after_same, increase,
            f"=IFERROR(B{row_index}/{common_count},0)",
            f"=IFERROR(C{row_index}/{common_count},0)",
            after_all,
            f"=IFERROR(G{row_index}/{all_count},0)",
            samples,
            "ISTNIEJĄCY ATRYBUT — UZUPEŁNIONO WARTOŚCI",
            "Nie tworzyć ponownie; zachować istniejący slug i tylko zaktualizować wartości produktów.",
        ]
        for col_index, value in enumerate(values, 1):
            ws.cell(row_index, col_index, value)

    end_row = 5 + len(rows_out)
    if rows_out:
        style_data_range(ws, 6, end_row, 1, len(headers))
        for row in range(6, end_row + 1):
            for col in (5, 6, 8):
                ws.cell(row, col).number_format = "0.0%"
            ws.cell(row, 4).fill = PatternFill("solid", fgColor=GREEN)
            ws.cell(row, 10).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    set_widths(ws, {1: 36, 2: 16, 3: 16, 4: 12, 5: 14, 6: 16, 7: 16, 8: 17, 9: 44, 10: 32, 11: 52})
    add_autofilter_and_freeze(ws, 5, len(headers), end_row)
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = "1:5"
    return len(rows_out)


def write_control_columns_sheet(wb: Workbook, control_columns: list[str]) -> int:
    ws = wb.create_sheet("Kolumny kontrolne")
    set_title(
        ws,
        "Nowe kolumny kontrolne — nie są atrybutami produktu",
        "Pola służą do audytu źródeł i jakości danych. Należy wyłączyć je z mapowania atrybutów WooCommerce.",
        5,
    )
    headers = ["Kolumna", "Typ", "Cel", "Import do WooCommerce", "Rekomendacja"]
    for index, header in enumerate(headers, 1):
        ws.cell(5, index, header)
    style_header(ws, 5, 1, len(headers))

    purposes = {
        "Źródło": "Identyfikacja dokumentu lub strony, z której pochodzi parametr.",
        "Strona katalogu": "Numer strony katalogu producenta.",
        "Adres źródła": "Link do karty produktu lub dokumentu źródłowego.",
        "Pewność": "Ocena wiarygodności przypisania danych.",
        "Status": "Status walidacji rekordu: OK, WARNING lub DO_SPRAWDZENIA.",
        "Uwagi": "Komentarz roboczy dotyczący braków albo niejednoznaczności.",
    }
    for row_index, column in enumerate(control_columns, 6):
        values = [column, "META / KONTROLA", purposes.get(column, "Pole kontrolne procesu."), "NIE", "Wykluczyć z importu atrybutów; zachować w pliku roboczym i raportach."]
        for col_index, value in enumerate(values, 1):
            ws.cell(row_index, col_index, value)
    end_row = 5 + len(control_columns)
    style_data_range(ws, 6, end_row, 1, len(headers))
    for row in range(6, end_row + 1):
        ws.cell(row, 4).fill = PatternFill("solid", fgColor=RED)
    set_widths(ws, {1: 25, 2: 20, 3: 55, 4: 23, 5: 58})
    add_autofilter_and_freeze(ws, 5, len(headers), end_row)
    ws.sheet_view.showGridLines = False
    return len(control_columns)


def write_deployment_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Wdrożenie")
    set_title(
        ws,
        "Plan wdrożenia nowych atrybutów do WooCommerce",
        "Kolejność minimalizuje ryzyko powstania duplikatów, pustych taksonomii i niekontrolowanych wariantów wartości.",
        6,
    )
    headers = ["Krok", "Działanie", "Zakres", "Warunek zakończenia", "Ryzyko", "Status"]
    for index, header in enumerate(headers, 1):
        ws.cell(5, index, header)
    style_header(ws, 5, 1, len(headers))
    steps = [
        (1, "Zatwierdzić nazwy i slugi", "23 nowe atrybuty", "Brak kolizji z istniejącymi slugami; slug uzgodniony przed utworzeniem.", "Zmiana sluga po imporcie jest kosztowna.", "DO DECYZJI"),
        (2, "Utworzyć globalne atrybuty", "22 pozycje oznaczone UTWORZYĆ GLOBALNY", "Każdy atrybut widoczny w Produkty > Atrybuty.", "Nie tworzyć kolumn kontrolnych jako atrybutów.", "DO WYKONANIA"),
        (3, "Zbudować kontrolowane terminy", "Wartości pokazane w raporcie i pełne wartości z pliku Hager", "Brak duplikatów różniących się zapisem, spacją lub jednostką.", "Niespójne wartości pogorszą filtry.", "DO WYKONANIA"),
        (4, "Znormalizować Kompatybilność", "1 atrybut opisowy", "Powstał krótki słownik nazw rodzin/serii albo decyzja o pozostawieniu pola lokalnego.", "Zdania opisowe tworzą bezużyteczne terminy globalne.", "DO WYKONANIA"),
        (5, "Uzupełnić mapowanie importu", "Product Attributes / taxonomie pa_*", "Nowe kolumny wskazują właściwe globalne taxonomie.", "Obecnie żaden z 23 nowych slugów nie jest zadeklarowany.", "DO WYKONANIA"),
        (6, "Wykonać import próbny", "Po 1–2 SKU z każdej kategorii", "Atrybuty są widoczne, filtrowalne i nie dublują istniejących pól.", "Pełny import bez próby utrudnia cofnięcie błędów.", "DO WYKONANIA"),
        (7, "Wykonać pełny import i audyt", "153 produkty Hager", "Liczba produktów i pokrycie odpowiadają raportowi; wyjątki mają status DO_SPRAWDZENIA.", "Należy zachować oryginalny plik wejściowy bez zmian.", "DO WYKONANIA"),
    ]
    for row_index, values in enumerate(steps, 6):
        for col_index, value in enumerate(values, 1):
            ws.cell(row_index, col_index, value)
    end_row = 5 + len(steps)
    style_data_range(ws, 6, end_row, 1, len(headers))
    for row in range(6, end_row + 1):
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row, 6).fill = PatternFill("solid", fgColor=YELLOW)
        ws.row_dimensions[row].height = 48
    set_widths(ws, {1: 9, 2: 34, 3: 34, 4: 53, 5: 46, 6: 18})
    add_autofilter_and_freeze(ws, 5, len(headers), end_row)
    ws.sheet_view.showGridLines = False


def write_summary_sheet(
    wb: Workbook,
    new_count: int,
    existing_enriched_count: int,
    control_count: int,
    product_count: int,
    category_count: int,
) -> None:
    ws = wb.active
    ws.title = "Podsumowanie"
    set_title(
        ws,
        "Raport nowych atrybutów Hager względem WooCommerce",
        "Porównanie pierwotnego eksportu WooCommerce z uzupełnionym plikiem Hager. Raport nie modyfikuje danych wejściowych.",
        8,
    )

    cards = [
        ("Nowe atrybuty produktowe", "=COUNTA('Nowe atrybuty'!A6:A100)"),
        ("Rekomendowane globalne", '=COUNTIF(\'Nowe atrybuty\'!G6:G100,"UTWORZYĆ GLOBALNY")'),
        ("Do normalizacji", '=COUNTIF(\'Nowe atrybuty\'!G6:G100,"ZNORMALIZOWAĆ PRZED GLOBALNYM")'),
        ("Już zadeklarowane w Product Attributes", '=COUNTIF(\'Nowe atrybuty\'!F6:F100,"Tak")'),
        ("Istniejące atrybuty uzupełnione", "=COUNTA('Istniejące uzupełnione'!A6:A500)"),
        ("Kolumny kontrolne", "=COUNTA('Kolumny kontrolne'!A6:A50)"),
        ("Produkty w aktualnym pliku", f"={product_count}"),
        ("Kategorie docelowe", f"={category_count}"),
    ]
    for index, (label, formula) in enumerate(cards):
        row = 4 if index < 4 else 7
        col = 1 + (index % 4) * 2
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        label_cell = ws.cell(row, col, label)
        label_cell.fill = PatternFill("solid", fgColor=BLUE)
        label_cell.font = Font(color=WHITE, bold=True, size=10)
        label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        value_cell = ws.cell(row + 1, col, formula)
        value_cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
        value_cell.font = Font(color=NAVY, bold=True, size=20)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 32
        ws.row_dimensions[row + 1].height = 34

    ws.merge_cells("A11:H11")
    ws["A11"] = "Najważniejszy wniosek"
    ws["A11"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A11"].font = Font(color=WHITE, bold=True, size=12)
    ws["A11"].alignment = Alignment(vertical="center")
    ws.merge_cells("A12:H14")
    ws["A12"] = (
        f"W pliku dodano {new_count} nowych atrybutów produktowych, lecz żaden z proponowanych slugów nie jest jeszcze "
        "zadeklarowany w polu Product Attributes. Przed importem trzeba utworzyć globalne atrybuty WooCommerce, "
        "przygotować terminy i uzupełnić mapowanie. „Kompatybilność” wymaga najpierw normalizacji wartości."
    )
    ws["A12"].fill = PatternFill("solid", fgColor=YELLOW)
    ws["A12"].font = Font(color=NAVY, bold=True, size=11)
    ws["A12"].alignment = Alignment(vertical="center", wrap_text=True)

    ws.merge_cells("A16:H16")
    ws["A16"] = "Jak czytać raport"
    ws["A16"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A16"].font = Font(color=WHITE, bold=True)
    guidance = [
        ("Nowe atrybuty", "Pełna lista 23 nowych pól, slugi, pokrycie, priorytet oraz rekomendacja."),
        ("Użycie w kategoriach", "Powiązanie nowych atrybutów z kategoriami P0/P1/P2."),
        ("Istniejące uzupełnione", "Pola obecne już w WooCommerce, którym tylko dodano brakujące wartości."),
        ("Kolumny kontrolne", "Pola audytowe, których nie wolno mapować jako atrybuty produktu."),
        ("Wdrożenie", "Bezpieczna kolejność tworzenia atrybutów, terminów i importu próbnego."),
    ]
    for row_index, (sheet, description) in enumerate(guidance, 17):
        ws.cell(row_index, 1, sheet)
        ws.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=8)
        ws.cell(row_index, 2, description)
        for col in range(1, 9):
            cell = ws.cell(row_index, col)
            cell.fill = PatternFill("solid", fgColor=WHITE if row_index % 2 else PALE_BLUE)
            cell.border = GRID_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row_index, 1).font = Font(bold=True, color=NAVY)
        ws.row_dimensions[row_index].height = 26

    ws.merge_cells("A24:H24")
    ws["A24"] = "Źródła lokalne porównania"
    ws["A24"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A24"].font = Font(color=WHITE, bold=True)
    sources = [
        "input/Aparatura-Modulowa-Hager.xlsx — pierwotny eksport WooCommerce (96 produktów)",
        "output/Aparatura-Modulowa-Hager_nazwy_atrybuty_v2.xlsx — aktualny plik Hager (153 produkty)",
        "output/Hager_aparatura_kategorie_atrybuty_priorytety.xlsx — kategorie, priorytety i pokrycie",
        "configs/hager_new_woocommerce_attributes.yaml — proponowane slugi i decyzje wdrożeniowe",
    ]
    for row_index, source in enumerate(sources, 25):
        ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=8)
        ws.cell(row_index, 1, source)
        ws.cell(row_index, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(row_index, 1).border = GRID_BORDER
        ws.cell(row_index, 1).fill = PatternFill("solid", fgColor=WHITE if row_index % 2 else PALE_BLUE)

    set_widths(ws, {1: 20, 2: 16, 3: 20, 4: 16, 5: 20, 6: 16, 7: 22, 8: 16})
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    # Sanity checks remain in code so a schema drift fails loudly during generation.
    assert new_count == 23, f"Expected 23 new attributes, got {new_count}"
    assert control_count == 6, f"Expected 6 control columns, got {control_count}"
    assert product_count == 153, f"Expected 153 products, got {product_count}"
    assert existing_enriched_count >= 1, "Expected at least one enriched existing attribute"


def main() -> None:
    with CONFIG_PATH.open("r", encoding="utf-8") as handle:
        config = yaml.safe_load(handle)

    original_headers, original_rows = load_sheet_rows(ORIGINAL_PATH)
    enriched_headers, enriched_rows = load_sheet_rows(ENRICHED_PATH)
    matrix_rows = load_matrix_rows(PRIORITY_PATH)

    original_attr_names = {header[len(ATTR_PREFIX):] for header in original_headers if header.startswith(ATTR_PREFIX)}
    enriched_attr_names = {header[len(ATTR_PREFIX):] for header in enriched_headers if header.startswith(ATTR_PREFIX)}
    actual_new_names = enriched_attr_names - original_attr_names
    configured = config["attributes"]
    configured_names = {clean(item["attribute"]) for item in configured}
    if actual_new_names != configured_names:
        missing = sorted(actual_new_names - configured_names)
        extra = sorted(configured_names - actual_new_names)
        raise ValueError(f"Config and workbook schema differ. Missing in config: {missing}; extra in config: {extra}")

    slugs = [clean(item["slug"]) for item in configured]
    if len(slugs) != len(set(slugs)):
        raise ValueError("Duplicate proposed WooCommerce slugs in config")
    if any(len(slug) > 28 for slug in slugs):
        raise ValueError("A proposed WooCommerce attribute slug exceeds 28 characters")

    declared_taxonomies = parse_declared_taxonomies(enriched_rows)
    original_products = product_map(original_rows)
    enriched_products = product_map(enriched_rows)

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.properties.creator = "AlkanInator"
    wb.properties.title = "Raport nowych atrybutów Hager względem WooCommerce"
    wb.properties.subject = "Audyt schematu atrybutów i rekomendacje wdrożeniowe"

    new_report_rows = write_new_attributes_sheet(wb, configured, enriched_rows, declared_taxonomies, matrix_rows)
    write_category_usage_sheet(wb, configured_names, matrix_rows)
    existing_count = write_existing_enriched_sheet(wb, original_headers, original_products, enriched_products)
    control_count = write_control_columns_sheet(wb, list(config["control_columns"]))
    write_deployment_sheet(wb)

    category_count = len({clean(row.get("Ścieżka kategorii")) for row in matrix_rows if populated(row.get("Ścieżka kategorii"))})
    write_summary_sheet(
        wb,
        new_count=len(new_report_rows),
        existing_enriched_count=existing_count,
        control_count=control_count,
        product_count=len(enriched_products),
        category_count=category_count,
    )

    preferred_order = ["Podsumowanie", "Nowe atrybuty", "Użycie w kategoriach", "Istniejące uzupełnione", "Kolumny kontrolne", "Wdrożenie"]
    wb._sheets = [wb[name] for name in preferred_order]
    wb.active = 0
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)

    print(f"OUTPUT={OUTPUT_PATH}")
    print(f"PRODUCTS={len(enriched_products)}")
    print(f"NEW_ATTRIBUTES={len(new_report_rows)}")
    print(f"DECLARED_NEW_TAXONOMIES={sum(1 for row in new_report_rows if row['declared'])}")
    print(f"EXISTING_ENRICHED={existing_count}")
    print(f"CONTROL_COLUMNS={control_count}")


if __name__ == "__main__":
    main()
