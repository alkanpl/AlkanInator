"""Uzupelnia eksport Woo produktow Kontakt Simon: tagi wariantow, Up-Sells,
Cross-Sells, produkty kompatybilne, nowe zdjecia i opisy.

Wejscie to eksport WP All Export (id, Title, Sku, Image ..., Up-Sells,
Cross-Sells, _alkan_compatible_product_ids, Product Tags, Content). Wynik to
kopia tego eksportu z wypelnionymi kolumnami - do importu WP All Import
(opisy z FAQ JSON-LD musza isc przez WP All Import, nie przez Baselinkera).

- Product Tags: tag grupy wariantow per symbol bazowy (WPC Linked Variation).
- Up-Sells / Cross-Sells / kompatybilne: reguly z `dictionaries/simon_woo_relations.yaml`,
  ograniczone do produktow z pliku (SKU i ID musza istniec w sklepie).
- Zdjecia (`--images-dir`): oryginal z API e-katalogu (1000x1000 PNG) ->
  JPG 1000x1000 na bialym tle, ta sama obrobka co przy wysylce do Baselinkera.
  W kolumnach obrazu laduje nazwa pliku do wgrania w
  `wp-content/uploads/wpallimport/files/`, a alt i tytul obrazu = nazwa produktu.
- Content (`--html-dir`): gotowe opisy `<slug SKU>.html` napisane recznie.

Uzycie:
    py src/build_simon_woo_relations.py --input input/Simon-24.xlsx \
        --html-dir reports/descriptions/simon_24/html \
        --images-dir output/Simon-24_zdjecia_2026-10-28 \
        --output output/Simon-24_WOO_IMPORT_2026-10-28.xlsx \
        --report reports/Simon-24_relacje_kontrola_2026-10-28.xlsx
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

import openpyxl
import requests
import yaml

from baselinker_api.images import _encoded_image
from simon_api import load_api_key, normalize_symbol

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RULES = ROOT / "dictionaries" / "simon_woo_relations.yaml"
DEFAULT_API_CACHE = ROOT / "cache" / "simon_api" / "articles"
COLOR_SEGMENT = re.compile(r"^(AB)?\d{2}A?$")
# Grafika "brak zdjecia" ze strony Kontakt-Simon (patrz build_simon_baselinker_from_api).
PLACEHOLDER_IMAGE_SHA256 = {"dce9f01b351674251e60ff33f46c2a1f29042f62ee7e24c5cee5b5de49ec9edf"}
IMAGE_COLUMNS_CLEARED = ("Image Path", "Image ID", "Image Caption", "Image Description")


@dataclass
class Product:
    row: int
    id: str
    sku: str
    title: str
    base: str
    color: str
    role: str = ""
    tag: str = ""
    up_sells: list[str] = field(default_factory=list)
    cross_sells: list[str] = field(default_factory=list)
    compatible: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def symbol_parts(sku: str) -> tuple[str, str]:
    """('NW6/2.01/X', '11') z 'NW6/2.01/X/11/KON'; pusty kolor dla 'NUV2/KON'."""
    symbol = normalize_symbol(sku)
    head, _, tail = symbol.rpartition("/")
    if head and COLOR_SEGMENT.match(tail):
        return head, tail
    return symbol, ""


def plain_color(color: str) -> str:
    """Wersja z klapka transparentna (49A) ma ten sam kolor co 49."""
    return color[:-1] if color.endswith("A") else color


def slugify(text: str) -> str:
    text = text.replace("ł", "l").replace("Ł", "L")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")


def sku_slug(sku: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", sku).strip("_")


def group_tag(title: str, rules: dict) -> str:
    core = title
    for pattern in rules["tag_strip_patterns"]:
        core = re.sub(pattern, "", core, flags=re.IGNORECASE)
    core = re.sub(r"\s+", " ", core).strip()
    return f"{core} {rules['series_tag_suffix']}"


def role_for(base: str, rules: dict) -> dict:
    for role in rules["roles"]:
        if re.search(role["match"], base):
            return role
    return {"name": "", "up_sells": [], "cross_sells": []}


class Catalog:
    def __init__(self, products: list[Product], rules: dict) -> None:
        self.products = products
        self.rules = rules
        self.by_base: dict[str, list[Product]] = defaultdict(list)
        for product in products:
            self.by_base[product.base].append(product)
        order = {code: index for index, code in enumerate(rules["color_order"])}
        for members in self.by_base.values():
            members.sort(key=lambda p: (order.get(plain_color(p.color), 99), p.color))

    def resolve(self, source: Product, symbol: str) -> list[Product]:
        """Produkty symbolu bazowego w kolorze zrodla; 'X*' albo zrodlo bez koloru = kazdy kolor."""
        any_color = symbol.endswith("*") or not source.color
        members = self.by_base.get(symbol.rstrip("*"), [])
        if any_color:
            return list(members)
        wanted = plain_color(source.color)
        exact = [p for p in members if p.color == source.color]
        return exact or [p for p in members if plain_color(p.color) == wanted]


def pick(source: Product, symbols: list[str], catalog: Catalog, limit: int, exclude: set[str]) -> list[str]:
    result: list[str] = []
    for symbol in symbols:
        if symbol.rstrip("*") == source.base:
            continue
        for target in catalog.resolve(source, symbol):
            if target.sku == source.sku or target.sku in exclude or target.sku in result:
                continue
            result.append(target.sku)
            # Z jednego symbolu jeden produkt - chyba ze to relacja "dowolny kolor".
            if not (symbol.endswith("*") or not source.color):
                break
        if len(result) >= limit:
            break
    return result[:limit]


def build_relations(products: list[Product], rules: dict) -> None:
    catalog = Catalog(products, rules)
    limits = rules["limits"]

    tags_by_base: dict[str, Counter] = defaultdict(Counter)
    for product in products:
        tags_by_base[product.base][group_tag(product.title, rules)] += 1
    for product in products:
        counts = tags_by_base[product.base]
        product.tag = counts.most_common(1)[0][0]
        if len(counts) > 1:
            product.notes.append(f"Rozne rdzenie nazw w grupie {product.base}: {dict(counts)}; tag = najczestszy.")

    for product in products:
        role = role_for(product.base, rules)
        product.role = role["name"]
        up_symbols = rules["frames"] if role["up_sells"] == "frames" else role["up_sells"]
        product.up_sells = pick(product, up_symbols, catalog, limits["up_sells"], set())
        if role["up_sells"] and not product.up_sells:
            product.notes.append("Brak produktow Up-Sell w tym kolorze w pliku (np. brak ramek w kolorze).")
        product.cross_sells = pick(product, role["cross_sells"], catalog, limits["cross_sells"], set(product.up_sells))

    # Relacje producenta - lustrzane, bez pozycji obecnych juz w Up-Sells karty.
    pairs: dict[str, list[str]] = defaultdict(list)
    for product in products:
        for symbol in rules.get("compatible", {}).get(product.base, []):
            for target in catalog.resolve(product, symbol):
                pairs[product.sku].append(target.sku)
                pairs[target.sku].append(product.sku)
    by_sku = {p.sku: p for p in products}
    for sku, targets in pairs.items():
        product = by_sku[sku]
        seen: list[str] = []
        for target in targets:
            if target not in seen and target not in product.up_sells and target != sku:
                seen.append(target)
        product.compatible = seen[: limits["compatible"]]


def fetch_api_image(url: str, api_key: str) -> bytes:
    response = requests.get(url, headers={"X-ApiKey": api_key, "User-Agent": "Mozilla/5.0"}, timeout=60)
    response.raise_for_status()
    return response.content


def image_file_name(product: Product, index: int) -> str:
    title = product.title
    code = normalize_symbol(product.sku)
    stem = title[: -len(code)].strip() if title.endswith(code) else title
    words = slugify(stem)
    if len(words) > 80:
        words = words[:81].rsplit("-", 1)[0]
    name = f"{words}-{slugify(code)}"
    return f"{name}-{index + 1}.jpg" if index else f"{name}.jpg"


def build_images(products: list[Product], cache_dir: Path, images_dir: Path, api_key: str) -> dict[str, list[str]]:
    images_dir.mkdir(parents=True, exist_ok=True)
    result: dict[str, list[str]] = {}
    for product in products:
        cache_file = cache_dir / (normalize_symbol(product.sku).replace("/", "_") + ".json")
        if not cache_file.exists():
            product.notes.append(f"Brak produktu w cache API ({cache_file.name}) - zdjecie bez zmian.")
            continue
        api_product = json.loads(cache_file.read_text(encoding="utf-8"))
        names: list[str] = []
        for index, image in enumerate(api_product.get("images") or []):
            name = image_file_name(product, index)
            target = images_dir / name
            if not target.exists():
                try:
                    body = fetch_api_image(image["url"], api_key)
                except requests.RequestException as exc:
                    product.notes.append(f"Nie pobrano zdjecia {image['url']}: {exc}")
                    continue
                if hashlib.sha256(body).hexdigest() in PLACEHOLDER_IMAGE_SHA256:
                    product.notes.append("Zdjecie z API to grafika 'brak zdjecia' - pominiete.")
                    continue
                encoded, warnings = _encoded_image(BytesIO(body), image["url"])
                product.notes.extend(warnings)
                if not encoded:
                    continue
                target.write_bytes(base64.b64decode(encoded.removeprefix("data:")))
            names.append(name)
        if names:
            result[product.sku] = names
        else:
            product.notes.append("Brak zdjecia z API - zostaje obecne zdjecie sklepu.")
    return result


def read_products(sheet) -> tuple[list[str], list[Product]]:
    header = [str(cell.value or "") for cell in sheet[1]]
    col = {name: index for index, name in enumerate(header)}
    products: list[Product] = []
    for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        sku = str(row[col["Sku"]] or "").strip()
        if not sku:
            continue
        base, color = symbol_parts(sku)
        products.append(Product(
            row=row_index,
            id=str(row[col["id"]] or "").strip(),
            sku=sku,
            title=str(row[col["Title"]] or "").strip(),
            base=base,
            color=color,
        ))
    return header, products


def write_import(args, header: list[str], products: list[Product], images: dict[str, list[str]], html_dir: Path | None) -> dict[str, int]:
    workbook = openpyxl.load_workbook(args.input)
    sheet = workbook.worksheets[0]
    col = {name: index + 1 for index, name in enumerate(header)}
    stats = Counter()
    for product in products:
        def put(column: str, value: str) -> None:
            if column in col:
                sheet.cell(product.row, col[column]).value = value

        put("Product Tags", product.tag)
        put("Up-Sells", "|".join(product.up_sells))
        put("Cross-Sells", "|".join(product.cross_sells))
        if "_alkan_compatible_product_ids" in col:
            ids = [p.id for p in products if p.sku in product.compatible]
            put("_alkan_compatible_product_ids", ",".join(ids))
        names = images.get(product.sku)
        if names:
            put("Image URL", "|".join(names))
            put("Image Filename", "|".join(names))
            put("Image Featured", names[0])
            put("Image Title", "|".join(product.title for _ in names))
            put("Image Alt Text", "|".join(product.title for _ in names))
            for column in IMAGE_COLUMNS_CLEARED:
                put(column, "")
            stats["zdjecia"] += 1
        if html_dir is not None:
            path = html_dir / f"{sku_slug(product.sku)}.html"
            if path.exists():
                put("Content", path.read_text(encoding="utf-8").strip())
                stats["opisy"] += 1
            else:
                product.notes.append("Brak gotowego opisu HTML - Content bez zmian.")
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    return dict(stats)


def load_variant_colors(rules: dict) -> dict[str, str]:
    """Kod koloru -> wartosc atrybutu osi grupy WPC (11 -> Biały)."""
    group = rules["variant_group"]
    source = yaml.safe_load((ROOT / group["colors_source"]).read_text(encoding="utf-8"))
    return {code: str(spec[group["color_field"]]) for code, spec in source["colors"].items()}


def variant_groups(products: list[Product], rules: dict, colors: dict[str, str]) -> list[list]:
    """Wiersze arkusza "Grupy wariantow": tag, grupa WPC, os i wartosci osi per SKU."""
    group = rules["variant_group"]
    attribute = group["attribute"]
    members: dict[str, list[Product]] = defaultdict(list)
    for p in products:
        members[p.tag].append(p)
    rows = []
    for tag, items in sorted(members.items()):
        values = {p.sku: colors.get(plain_color(p.color), "") for p in items}
        if not any(p.color for p in items):
            # Produkt bez koloru (NUV2) nie ma wariantow - nie trafia do grupy WPC.
            rows.append([tag, slugify(tag), len(items), "-", "-", "", "OK",
                         "Produkt bez koloru - tagu nie dopisuje sie do grupy WPC.", "|".join(values)])
            continue
        notes = []
        missing = [sku for sku, value in values.items() if not value]
        if missing:
            notes.append(f"Brak wartosci atrybutu {attribute} dla: {', '.join(missing)}.")
        repeated = [value for value, count in Counter(v for v in values.values() if v).items() if count > 1]
        if repeated:
            notes.append(f"Powtorzony {attribute} ({', '.join(repeated)}) - grupa potrzebuje dodatkowej osi.")
        if len(items) == 1:
            notes.append("Jeden kolor w sklepie - przelacznik pojawi sie po dodaniu kolejnego koloru.")
        status = "DO_SPRAWDZENIA" if missing or repeated else "OK"
        rows.append([tag, slugify(tag), len(items), group["name"], attribute,
                     " | ".join(f"{sku}: {value or '-'}" for sku, value in values.items()),
                     status, " ".join(notes), "|".join(values)])
    return rows


def write_report(path: Path, products: list[Product], rules: dict) -> None:
    by_sku = {p.sku: p for p in products}
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Relacje"
    sheet.append(["id", "Sku", "Title", "Rola", "Product Tags", "Up-Sells", "Cross-Sells",
                  "Kompatybilne (SKU)", "Status", "Uwagi"])
    for p in products:
        status = "OK" if not p.notes else "WARNING"
        sheet.append([p.id, p.sku, p.title, p.role, p.tag, "|".join(p.up_sells), "|".join(p.cross_sells),
                      "|".join(p.compatible), status, " ".join(p.notes)])

    groups = workbook.create_sheet("Grupy wariantow")
    groups.append(["Tag", "Slug tagu", "Liczba produktow", "Grupa WPC (dopisz tag)", "Atrybut wariantu",
                   "Wartosc atrybutu per SKU", "Status", "Uwagi", "SKU"])
    for row in variant_groups(products, rules, load_variant_colors(rules)):
        groups.append(row)

    readable = workbook.create_sheet("Podglad relacji")
    readable.append(["Sku", "Title", "Relacja", "Sku celu", "Title celu"])
    for p in products:
        for kind, skus in (("Up-Sell", p.up_sells), ("Cross-Sell", p.cross_sells), ("Kompatybilny", p.compatible)):
            for sku in skus:
                readable.append([p.sku, p.title, kind, sku, by_sku[sku].title])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", required=True, help="Eksport WP All Export (XLSX).")
    parser.add_argument("--output", required=True, help="Kopia eksportu do WP All Import.")
    parser.add_argument("--report", required=True, help="Skoroszyt kontrolny.")
    parser.add_argument("--rules", default=str(DEFAULT_RULES))
    parser.add_argument("--api-cache", default=str(DEFAULT_API_CACHE))
    parser.add_argument("--images-dir", default="", help="Katalog na JPG 1000x1000; bez niego zdjecia zostaja bez zmian.")
    parser.add_argument("--html-dir", default="", help="Katalog z opisami <slug SKU>.html; bez niego Content bez zmian.")
    args = parser.parse_args()

    rules = yaml.safe_load(Path(args.rules).read_text(encoding="utf-8"))
    sheet = openpyxl.load_workbook(args.input, read_only=True).worksheets[0]
    header, products = read_products(sheet)
    build_relations(products, rules)
    images: dict[str, list[str]] = {}
    if args.images_dir:
        images = build_images(products, Path(args.api_cache), Path(args.images_dir), load_api_key())
    stats = write_import(args, header, products, images, Path(args.html_dir) if args.html_dir else None)
    write_report(Path(args.report), products, rules)
    warnings = sum(1 for p in products if p.notes)
    print(f"Produkty: {len(products)}, grupy wariantow: {len({p.tag for p in products})}, "
          f"zdjecia: {stats.get('zdjecia', 0)}, opisy: {stats.get('opisy', 0)}, z uwagami: {warnings}")
    print(f"Import: {args.output}")
    print(f"Kontrola: {args.report}")


if __name__ == "__main__":
    main()
