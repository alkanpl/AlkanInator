from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
import re
import unicodedata

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
SOURCE_WORKBOOK = ROOT / "output" / "Hager_nazwy_po_decyzjach_2026-08-07.xlsx"
AUDIT_CONFIG = ROOT / "configs" / "hager_attribute_deduplication.yaml"
IMPLEMENTATION_CONFIG = ROOT / "configs" / "hager_woocommerce_implementation.yaml"
OUTPUT = ROOT / "output" / "Hager_atrybuty_i_kategorie_do_wdrozenia_2026-08-07_v2.xlsx"

NAVY = "17365D"
BLUE = "2F75B5"
TEAL = "008C95"
LIGHT_BLUE = "DCE6F1"
LIGHT_TEAL = "DDEBF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_YELLOW = "FFF2CC"
LIGHT_RED = "FCE4D6"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
TEXT = "1F1F1F"
GRID = "D9E2F3"


def load_yaml(path: Path) -> dict:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def clean_slug(value: str) -> str:
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    value = re.sub(r"[^a-zA-Z0-9]+", "-", value.lower()).strip("-")
    return value[:28]


def unique(values):
    return sorted({str(value).strip() for value in values if value not in (None, "")})


def joined(values, limit=8):
    values = unique(values)
    if len(values) <= limit:
        return ", ".join(values)
    return ", ".join(values[:limit]) + f" (+{len(values) - limit})"


def sheet_data():
    workbook = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=True)

    product_rows = list(workbook["Nazwy po decyzjach"].iter_rows(min_row=6, values_only=True))
    product_lookup = {
        row[1]: {
            "type": row[7],
            "source_category": row[8],
            "target_category": row[9],
        }
        for row in product_rows
        if row[1]
    }

    worksheet_rows = list(workbook["Worksheet"].iter_rows(values_only=True))
    headers = list(worksheet_rows[0])
    sku_index = headers.index("SKU")
    attribute_data = {}
    for column_index, header in enumerate(headers):
        prefix = "Atrybut Produktu: "
        if not isinstance(header, str) or not header.startswith(prefix):
            continue
        attribute = header[len(prefix) :]
        entries = []
        for row in worksheet_rows[1:]:
            value = row[column_index]
            if value in (None, ""):
                continue
            sku = row[sku_index]
            entries.append((sku, value))
        attribute_data[attribute] = entries

    return product_rows, product_lookup, attribute_data


def attribute_stats(source_attributes, product_lookup, attribute_data):
    entries = []
    for attribute in source_attributes:
        entries.extend((attribute, sku, value) for sku, value in attribute_data.get(attribute, []))
    skus = {sku for _, sku, _ in entries}
    values = [value for _, _, value in entries]
    types = [product_lookup.get(sku, {}).get("type") for sku in skus]
    categories = [product_lookup.get(sku, {}).get("target_category") for sku in skus]
    return {
        "count": len(skus),
        "values": joined(values, 12),
        "types": joined(types, 6),
        "categories": joined(categories, 5),
    }


def rationale_for(source_attributes, audit_lookup):
    rationales = [audit_lookup[a]["rationale"] for a in source_attributes if a in audit_lookup]
    return " ".join(dict.fromkeys(rationales))


def normalization_for(source_attributes, audit_lookup):
    normalizations = [audit_lookup[a]["normalization"] for a in source_attributes if a in audit_lookup]
    return " ".join(dict.fromkeys(normalizations))


def add_title(ws, title, subtitle, end_column):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    ws.cell(1, 1, title)
    ws.cell(1, 1).font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).font = Font(name="Aptos", size=10, color="44546A")
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 32


def add_table(ws, start_row, headers, rows, name):
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(start_row, col, header)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row_number, row in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row_number, col, value)
            cell.font = Font(name="Aptos", size=9, color=TEXT)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_number].height = 52
    last_row = start_row + len(rows)
    last_col = len(headers)
    table = Table(displayName=name, ref=f"A{start_row}:{ws.cell(last_row, last_col).coordinate}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)
    ws.freeze_panes = f"A{start_row + 1}"
    ws.row_dimensions[start_row].height = 36
    return last_row


def add_status_validation(ws, column_letter, first_row, last_row):
    validation = DataValidation(
        type="list",
        formula1='"DO WYKONANIA,W TRAKCIE,GOTOWE,NIE DOTYCZY"',
        allow_blank=False,
    )
    validation.error = "Wybierz status z listy."
    validation.errorTitle = "Nieprawidłowy status"
    ws.add_data_validation(validation)
    validation.add(f"{column_letter}{first_row}:{column_letter}{last_row}")
    fills = {
        "GOTOWE": LIGHT_GREEN,
        "W TRAKCIE": LIGHT_YELLOW,
        "DO WYKONANIA": LIGHT_RED,
        "NIE DOTYCZY": LIGHT_GRAY,
    }
    for status, color in fills.items():
        ws.conditional_formatting.add(
            f"{column_letter}{first_row}:{column_letter}{last_row}",
            FormulaRule(formula=[f'{column_letter}{first_row}="{status}"'], fill=PatternFill("solid", fgColor=color)),
        )


def set_widths(ws, widths):
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def build():
    audit = load_yaml(AUDIT_CONFIG)
    config = load_yaml(IMPLEMENTATION_CONFIG)
    audit_lookup = {item["attribute"]: item for item in audit["decisions"]}
    product_rows, product_lookup, attribute_data = sheet_data()

    attribute_rows = []
    for order, item in enumerate(sorted(config["attribute_actions"], key=lambda x: (x["priority"], x["target_name"])), start=1):
        source_attributes = item["source_attributes"]
        stats = attribute_stats(source_attributes, product_lookup, attribute_data)
        existing_name = item.get("existing_name", "")
        existing_slug = item.get("existing_slug", "")
        target_slug = item.get("target_slug") or clean_slug(item["target_name"])
        attribute_rows.append([
            order,
            item["priority"],
            item["action"],
            existing_name,
            existing_slug,
            item["target_name"],
            target_slug,
            " + ".join(source_attributes),
            stats["count"],
            stats["values"],
            stats["types"],
            stats["categories"],
            normalization_for(source_attributes, audit_lookup),
            rationale_for(source_attributes, audit_lookup),
            "DO WYKONANIA",
        ])

    reuse_rows = []
    for item in config["reuse_without_new_attribute"]:
        stats = attribute_stats(item["source_attributes"], product_lookup, attribute_data)
        reuse_rows.append([
            item["action"],
            " + ".join(item["source_attributes"]),
            item["existing_name"],
            item["target_name"],
            stats["count"],
            item.get("display_values", stats["values"]),
            stats["types"],
            item["note"],
        ])

    source_target_counts = Counter((row[8], row[9]) for row in product_rows)
    source_target_types = defaultdict(set)
    source_target_skus = defaultdict(list)
    for row in product_rows:
        key = (row[8], row[9])
        source_target_types[key].add(row[7])
        source_target_skus[key].append(row[1])

    category_rows = []
    for order, item in enumerate(sorted(config["category_actions"], key=lambda x: (x["priority"], x["target_path"])), start=1):
        target = item["target_path"]
        affected = sum(count for (_, target_path), count in source_target_counts.items() if target_path == target)
        source_groups = [source for (source, target_path) in source_target_counts if target_path == target]
        category_rows.append([
            order,
            item["priority"],
            item["action"],
            item["existing_path"],
            target,
            affected,
            joined(source_groups, 10),
            item["note"],
            "DO WYKONANIA",
        ])

    move_rows = []
    for (source, target), count in sorted(source_target_counts.items(), key=lambda x: (x[0][1], x[0][0])):
        if source == target:
            continue
        action = "ZMIANA NAZWY KATEGORII" if any(
            item["action"] == "ZMIEŃ NAZWĘ" and item["existing_path"] == source and item["target_path"] == target
            for item in config["category_actions"]
        ) else "PRZENIEŚ / SCAL"
        move_rows.append([
            action,
            source,
            target,
            count,
            joined(source_target_types[(source, target)], 8),
            ", ".join(source_target_skus[(source, target)]),
            "DO WYKONANIA",
        ])

    wb = Workbook()
    summary = wb.active
    summary.title = "Podsumowanie"
    add_title(
        summary,
        "Hager — atrybuty i kategorie do wdrożenia",
        "Lista czynności w WooCommerce po audycie powieleń i decyzjach kategorii. Obejmuje 153 produkty Hager.",
        8,
    )
    summary["A4"] = "CO TRZEBA ZROBIĆ"
    summary["A4"].font = Font(name="Aptos", size=11, bold=True, color=NAVY)
    summary.merge_cells("A4:H4")
    summary_cards = [
        ("Atrybuty globalne do dodania", '=COUNTIF(\'Atrybuty do wdrożenia\'!C6:C200,"DODAJ GLOBALNY*")'),
        ("Atrybuty istniejące do przemianowania", '=COUNTIF(\'Atrybuty do wdrożenia\'!C6:C200,"ZMIEŃ NAZWĘ*")'),
        ("Kategorie do dodania", '=COUNTIF(\'Kategorie do wdrożenia\'!C6:C100,"DODAJ")'),
        ("Kategorie do przemianowania", '=COUNTIF(\'Kategorie do wdrożenia\'!C6:C100,"ZMIEŃ NAZWĘ")'),
    ]
    for index, (label, formula) in enumerate(summary_cards):
        start_col = 1 + index * 2
        summary.merge_cells(start_row=6, start_column=start_col, end_row=6, end_column=start_col + 1)
        summary.merge_cells(start_row=7, start_column=start_col, end_row=8, end_column=start_col + 1)
        summary.cell(6, start_col, label)
        summary.cell(7, start_col, formula)
        for row in (6, 7, 8):
            for col in (start_col, start_col + 1):
                summary.cell(row, col).fill = PatternFill("solid", fgColor=LIGHT_TEAL if row == 6 else WHITE)
                summary.cell(row, col).border = Border(
                    left=Side(style="thin", color=GRID), right=Side(style="thin", color=GRID),
                    top=Side(style="thin", color=GRID), bottom=Side(style="thin", color=GRID)
                )
        summary.cell(6, start_col).font = Font(name="Aptos", size=9, bold=True, color=NAVY)
        summary.cell(6, start_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        summary.cell(7, start_col).font = Font(name="Aptos Display", size=22, bold=True, color=TEAL)
        summary.cell(7, start_col).alignment = Alignment(horizontal="center", vertical="center")
    summary.row_dimensions[6].height = 34
    summary.row_dimensions[7].height = 28
    summary.row_dimensions[8].height = 12

    summary["A10"] = "Najważniejsza zasada"
    summary["A10"].font = Font(name="Aptos", size=11, bold=True, color=WHITE)
    summary["A10"].fill = PatternFill("solid", fgColor=TEAL)
    summary.merge_cells("A10:H10")
    summary["A11"] = (
        "Prowadzić równolegle Układ biegunów (1P, 1P+N, 2P, 3P, 3P+N, 4P) oraz Liczbę Faz (1 lub 3). "
        "Typ Wyłącznika zachować bez zmiany nazwy. Inc oraz Icc mają trafić do jednego nowego atrybutu."
    )
    summary.merge_cells("A11:H12")
    summary["A11"].alignment = Alignment(wrap_text=True, vertical="center")
    summary["A11"].fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
    summary["A11"].font = Font(name="Aptos", size=10, color=TEXT)

    summary["A14"] = "Zalecana kolejność"
    summary["A14"].font = Font(name="Aptos", size=11, bold=True, color=NAVY)
    steps = [
        "1. Najpierw przemianuj istniejące atrybuty i kategorie — zachowasz ich dotychczasowe terminy oraz przypisania.",
        "2. Dodaj 17 nowych globalnych atrybutów z nazwami i slugami z arkusza Atrybuty do wdrożenia.",
        "3. Wykonaj scalenia i przeniesienia produktów zgodnie z arkuszem Przeniesienia produktów.",
        "4. Na końcu przypisz wartości z pliku produktowego Hager i sprawdź filtry na froncie sklepu.",
    ]
    for row_index, step in enumerate(steps, start=15):
        summary.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=8)
        summary.cell(row_index, 1, step)
        summary.cell(row_index, 1).alignment = Alignment(wrap_text=True, vertical="center")
        summary.cell(row_index, 1).fill = PatternFill("solid", fgColor=WHITE if row_index % 2 else LIGHT_GRAY)
        summary.cell(row_index, 1).font = Font(name="Aptos", size=10, color=TEXT)
        summary.row_dimensions[row_index].height = 26
    set_widths(summary, {column: 14 for column in "ABCDEFGH"})
    summary.freeze_panes = "A4"

    attributes = wb.create_sheet("Atrybuty do wdrożenia")
    add_title(
        attributes,
        "Atrybuty do dodania lub przemianowania",
        "Priorytet 1 = wykonać przed importem. Slug bez prefiksu pa_; przy zmianie nazwy istniejącego atrybutu slug pozostaje bez zmian.",
        15,
    )
    attribute_headers = [
        "Kolejność", "Priorytet", "Akcja", "Obecna nazwa Woo", "Obecny slug",
        "Docelowa nazwa", "Docelowy slug", "Pola źródłowe Hager", "Produkty Hager",
        "Wartości", "Typy produktów", "Kategorie docelowe", "Normalizacja wartości",
        "Dlaczego", "Status",
    ]
    attribute_last = add_table(attributes, 5, attribute_headers, attribute_rows, "tblHagerAttributes")
    add_status_validation(attributes, "O", 6, attribute_last)
    set_widths(attributes, {
        "A": 10, "B": 9, "C": 25, "D": 31, "E": 30, "F": 40, "G": 30,
        "H": 41, "I": 13, "J": 28, "K": 38, "L": 45, "M": 46, "N": 55, "O": 17,
    })
    reuse = wb.create_sheet("Użyj istniejących")
    add_title(
        reuse,
        "Atrybuty istniejące do użycia lub uzupełnienia",
        "Te pola pozostają pod obecną nazwą w WooCommerce. Liczba Faz jest uzupełniana równolegle z nowym atrybutem Układ biegunów.",
        8,
    )
    reuse_headers = [
        "Akcja", "Pole Hager", "Istniejący atrybut Woo", "Używana nazwa", "Produkty Hager",
        "Wartości Hager", "Typy produktów", "Instrukcja",
    ]
    reuse_last = add_table(reuse, 5, reuse_headers, reuse_rows, "tblHagerReuse")
    set_widths(reuse, {"A": 20, "B": 30, "C": 31, "D": 25, "E": 14, "F": 34, "G": 45, "H": 72})

    categories = wb.create_sheet("Kategorie do wdrożenia")
    add_title(
        categories,
        "Kategorie do dodania lub przemianowania",
        "Końcowa struktura po decyzjach użytkownika. Nie twórz osobnych małych kategorii wymienionych w arkuszu Przeniesienia produktów.",
        9,
    )
    category_headers = [
        "Kolejność", "Priorytet", "Akcja", "Obecna ścieżka", "Docelowa ścieżka",
        "Produkty po zmianie", "Scalane grupy źródłowe", "Instrukcja", "Status",
    ]
    category_last = add_table(categories, 5, category_headers, category_rows, "tblHagerCategories")
    add_status_validation(categories, "I", 6, category_last)
    set_widths(categories, {"A": 10, "B": 9, "C": 18, "D": 52, "E": 56, "F": 19, "G": 72, "H": 62, "I": 17})

    moves = wb.create_sheet("Przeniesienia produktów")
    add_title(
        moves,
        "Przeniesienia i scalenia kategorii",
        "Lista grup produktowych, których kategoria źródłowa różni się od kategorii końcowej. SKU ułatwiają kontrolę po migracji.",
        7,
    )
    move_headers = [
        "Sposób", "Kategoria źródłowa", "Kategoria docelowa", "Liczba produktów",
        "Typy produktów", "SKU", "Status",
    ]
    move_last = add_table(moves, 5, move_headers, move_rows, "tblHagerMoves")
    add_status_validation(moves, "G", 6, move_last)
    set_widths(moves, {"A": 26, "B": 55, "C": 58, "D": 16, "E": 45, "F": 78, "G": 17})

    sources = wb.create_sheet("Źródła i uwagi")
    add_title(sources, "Źródła i zakres", "Pliki wykorzystane do przygotowania listy wdrożeniowej.", 4)
    source_rows = [
        ["Decyzje użytkownika", "archived_input_files/input_2026-08-07/Hager_plan_kategorii_atrybutow_i_audyt_uzupelniony_decyzje.xlsx", "Kategorie, łączenie grup i Liczba Faz"],
        ["Nazwy i kategorie końcowe", "output/Hager_nazwy_po_decyzjach_2026-08-07.xlsx", "153 produkty po zastosowaniu decyzji"],
        ["Audyt powieleń", "configs/hager_attribute_deduplication.yaml", "Porównanie nazw i wartości z WooCommerce"],
        ["Plan wdrożenia", "configs/hager_woocommerce_implementation.yaml", "Końcowe akcje, priorytety, nazwy i slugi"],
    ]
    add_table(sources, 5, ["Rodzaj", "Plik", "Zakres"], source_rows, "tblHagerSources")
    set_widths(sources, {"A": 28, "B": 100, "C": 62})
    sources["A11"] = "Uwaga"
    sources["A11"].font = Font(name="Aptos", bold=True, color=WHITE)
    sources["A11"].fill = PatternFill("solid", fgColor=TEAL)
    sources.merge_cells("A11:C11")
    sources["A12"] = (
        "Zmiana nazwy globalnego atrybutu WooCommerce nie wymaga zmiany jego sluga. "
        "Przed zmianą nazwy sprawdź jednak, czy ten sam globalny atrybut nie jest używany w innych działach sklepu; "
        "jeżeli jest, przetestuj etykietę na tych produktach."
    )
    sources.merge_cells("A12:C14")
    sources["A12"].alignment = Alignment(wrap_text=True, vertical="center")
    sources["A12"].fill = PatternFill("solid", fgColor=LIGHT_YELLOW)

    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5
        ws.sheet_view.zoomScale = 85 if ws.title != "Podsumowanie" else 100

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
