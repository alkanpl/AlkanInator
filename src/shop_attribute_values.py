"""Kontrola wartosci atrybutow wzgledem terminow, ktore sklep Woo juz ma.

Kazda wartosc atrybutu globalnego (pa_*) spoza slownika sklepu tworzy w Woo
NOWY termin - przy imporcie przez Baselinkera i przez WP All Import. Tak
powstaly w partii Simon kaszmir terminy "Antenowe, Komputerowe" (zamiast dwoch
terminow Antenowe|Komputerowe), "Kaszmirowy" obok "Kaszmir" i cale nowe
slowniki "Typ zaciskow" z danych ETIM.

Zrodla slownika:
- ``dictionaries/woo_attribute_terms_snapshot.tsv`` - zrzut terminow z bazy
  sklepu (``scripts/dump_woo_attribute_terms.py``),
- opcjonalnie eksport sklepu (WP All Export) - wartosci produktow spoza
  sprawdzanej partii (baza lokalna bywa starsza niz produkcja).

Reguly (``dictionaries/woo_attribute_value_rules.yaml``): mapowanie wartosci
dostawcy na terminy sklepu, wartosci zatwierdzone przez uzytkownika i klucz
rodzenstwa (ten sam produkt w innym kolorze).

Kolejnosc decyzji dla kazdej wartosci (czesci po ``|``):
1. termin sklepu (dokladnie albo ten sam slug/wielkosc liter - Woo i tak
   podepnie istniejacy termin),
2. ``value_map`` z regul (``null`` = usun wartosc),
3. wartosc z przecinkami, ktorej kazda czesc jest terminem -> terminy po ``|``,
4. ``approved_new_values`` - nowa wartosc zatwierdzona przez uzytkownika,
5. wartosc rodzenstwa w sklepie (wiekszosc, tylko istniejacy termin),
6. NOWA WARTOSC: raport ``DO_SPRAWDZENIA``; przy poprawianiu usuwana, chyba ze
   ``keep_new``.

Dodatkowo sprawdzane jest powiazanie filtra Typ produktu -> Podtyp produktu
(podtyp przypiety w sklepie do innego typu nie pokaze sie pod typem produktu).
"""

from __future__ import annotations

import csv
import difflib
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

import yaml

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SNAPSHOT = ROOT / "dictionaries" / "woo_attribute_terms_snapshot.tsv"
DEFAULT_RULES = ROOT / "dictionaries" / "woo_attribute_value_rules.yaml"
WOO_SEPARATOR = "|"
ATTRIBUTE_COLUMN_RE = re.compile(r"^atrybut produktu:\s*(.+)$", re.IGNORECASE)
TYP = "Typ produktu"
PODTYP = "Podtyp produktu"

KIND_EXISTING = "termin sklepu"
KIND_SAME_SLUG = "ten sam termin (inny zapis)"
KIND_MAPPED = "zmapowana na termin sklepu"
KIND_SPLIT = "rozbita na terminy sklepu"
KIND_APPROVED = "nowa wartosc zatwierdzona"
KIND_SIBLING = "wartosc rodzenstwa w sklepie"
KIND_REMOVED = "usunieta (brak odpowiednika)"
KIND_NEW = "NOWA WARTOSC"
KIND_NEW_TECHNICAL = "nowa wartosc techniczna (zostaje)"
KIND_NO_DICTIONARY = "atrybut bez slownika w sklepie"
KIND_TYPE_LINK = "podtyp z innego typu"
KIND_NO_PARENT = "podtyp bez typu w filtrze"

STATUS_BY_KIND = {
    KIND_SAME_SLUG: "OK",
    KIND_MAPPED: "WARNING",
    KIND_SPLIT: "WARNING",
    KIND_APPROVED: "OK",
    KIND_SIBLING: "WARNING",
    KIND_REMOVED: "WARNING",
    KIND_NEW: "DO_SPRAWDZENIA",
    KIND_NEW_TECHNICAL: "WARNING",
    KIND_NO_DICTIONARY: "WARNING",
    KIND_TYPE_LINK: "DO_SPRAWDZENIA",
    KIND_NO_PARENT: "WARNING",
}

_POLISH = str.maketrans({"ł": "l", "Ł": "l", "ø": "o", "Ø": "o", "ß": "ss"})


def cell_text(value: Any) -> str:
    """Tekst komorki jak w Woo: liczby calkowite bez ``.0``, bez spacji na brzegach."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def slugify(value: str) -> str:
    """Przyblizenie ``sanitize_title`` z WordPressa (po nim Woo rozpoznaje istniejacy termin)."""
    text = unicodedata.normalize("NFKD", value.translate(_POLISH))
    text = "".join(char for char in text if not unicodedata.combining(char)).lower()
    text = text.replace(".", "-")
    text = re.sub(r"[^a-z0-9 _-]", "", text)
    text = re.sub(r"[\s-]+", "-", text)
    return text.strip("-")


@dataclass
class ShopVocabulary:
    terms: dict[str, set[str]] = field(default_factory=dict)
    by_slug: dict[str, dict[str, str]] = field(default_factory=dict)
    parents: dict[str, dict[str, str]] = field(default_factory=dict)
    sources: list[str] = field(default_factory=list)

    def has_attribute(self, attribute: str) -> bool:
        return attribute in self.terms

    def add(self, attribute: str, term: str, parent: str = "") -> None:
        if not term:
            self.terms.setdefault(attribute, set())
            return
        self.terms.setdefault(attribute, set()).add(term)
        slugs = self.by_slug.setdefault(attribute, {})
        slugs.setdefault(slugify(term), term)
        slugs.setdefault(term.casefold(), term)
        if parent:
            self.parents.setdefault(attribute, {})[term] = parent

    def canonical(self, attribute: str, value: str) -> str | None:
        """Istniejacy termin, do ktorego Woo przypnie ``value`` (albo None = nowy termin)."""
        if value in self.terms.get(attribute, set()):
            return value
        slugs = self.by_slug.get(attribute, {})
        return slugs.get(value.casefold()) or slugs.get(slugify(value))

    def similar(self, attribute: str, value: str) -> str | None:
        names = sorted(self.terms.get(attribute, set()))
        lowered = {name.casefold(): name for name in names}
        match = difflib.get_close_matches(value.casefold(), list(lowered), n=1, cutoff=0.6)
        return lowered[match[0]] if match else None

    def parent(self, attribute: str, term: str) -> str | None:
        return self.parents.get(attribute, {}).get(term)


def load_vocabulary(path: Path = DEFAULT_SNAPSHOT) -> ShopVocabulary:
    vocabulary = ShopVocabulary(sources=[str(path)])
    with path.open("r", encoding="utf-8", newline="") as stream:
        lines = (line for line in stream if not line.startswith("#"))
        for row in csv.DictReader(lines, delimiter="\t"):
            vocabulary.add(row["attribute"], row.get("term") or "", row.get("parent_term") or "")
    return vocabulary


def load_rules(path: Path = DEFAULT_RULES) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as stream:
        return yaml.safe_load(stream) or {}


def is_identifier(attribute: str, rules: dict[str, Any]) -> bool:
    """Atrybut-identyfikator (Symbol): kazdy produkt ma wlasny termin, nie sprawdzamy."""
    return attribute in set(rules.get("identifier_attributes") or [])


def is_technical(attribute: str, rules: dict[str, Any]) -> bool:
    """Atrybut techniczny (liczba z jednostka): nowa wartosc zostaje, ale jest raportowana."""
    if attribute in set(rules.get("technical_attributes") or []):
        return True
    return any(re.search(pattern, attribute) for pattern in rules.get("technical_attribute_patterns") or [])


def attribute_columns(header: Iterable[str]) -> dict[str, str]:
    """Kolumna eksportu Woo -> nazwa atrybutu (``Atrybut produktu: Kolor`` -> ``Kolor``)."""
    result: dict[str, str] = {}
    for column in header:
        match = ATTRIBUTE_COLUMN_RE.match(column or "")
        if match:
            result[column] = match.group(1).strip()
    return result


def find_column(header: Iterable[str], *names: str) -> str | None:
    lowered = {str(column).lower(): column for column in header if column}
    for name in names:
        if name.lower() in lowered:
            return lowered[name.lower()]
    return None


def row_features(row: dict[str, Any], columns: dict[str, str]) -> dict[str, str]:
    features: dict[str, str] = {}
    for column, attribute in columns.items():
        value = cell_text(row.get(column))
        if value:
            features[attribute] = value
    return features


def sibling_key(sku: str, rules: dict[str, Any]) -> str:
    settings = rules.get("sibling_key") or {}
    value = (sku or "").strip()
    for key in ("strip_suffix_regex", "color_suffix_regex"):
        pattern = settings.get(key)
        if pattern:
            value = re.sub(pattern, "", value)
    return value


@dataclass
class ShopExport:
    """Produkty sklepu spoza sprawdzanej partii: dopelniaja slownik i daja rodzenstwo."""

    siblings: dict[str, list[dict[str, str]]] = field(default_factory=dict)
    rows: int = 0


def extend_from_export(
    vocabulary: ShopVocabulary,
    header: list[str],
    rows: list[dict[str, Any]],
    rules: dict[str, Any],
    exclude_skus: set[str] | None = None,
    source: str = "",
    export: ShopExport | None = None,
) -> ShopExport:
    export = export or ShopExport()
    columns = attribute_columns(header)
    sku_column = find_column(header, "Sku", "SKU")
    exclude = exclude_skus or set()
    for row in rows:
        sku = cell_text(row.get(sku_column)) if sku_column else ""
        if sku and sku in exclude:
            continue
        features = row_features(row, columns)
        for attribute, value in features.items():
            for part in value.split(WOO_SEPARATOR):
                part = part.strip()
                if part and vocabulary.has_attribute(attribute):
                    vocabulary.add(attribute, part)
        if sku:
            export.siblings.setdefault(sibling_key(sku, rules), []).append(features)
        export.rows += 1
    if source:
        vocabulary.sources.append(source)
    return export


@dataclass
class ValueCheck:
    sku: str
    title: str
    attribute: str
    value: str
    new_value: str | None
    kind: str
    status: str
    note: str = ""


def _rule_lookup(table: dict[str, Any], value: str) -> tuple[bool, Any]:
    if value in table:
        return True, table[value]
    lowered = value.casefold()
    for source, target in table.items():
        if str(source).casefold() == lowered:
            return True, target
    return False, None


def _approved(rules: dict[str, Any], extra: dict[str, set[str]] | None, attribute: str, value: str) -> bool:
    approved = {str(item) for item in ((rules.get("approved_new_values") or {}).get(attribute) or [])}
    approved |= (extra or {}).get(attribute, set())
    return value in approved


def _sibling_value(vocabulary: ShopVocabulary, attribute: str, siblings: list[dict[str, str]] | None) -> str | None:
    if not siblings:
        return None
    counts = Counter(features[attribute] for features in siblings if features.get(attribute))
    if not counts:
        return None
    value, _ = counts.most_common(1)[0]
    if all(vocabulary.canonical(attribute, part) for part in value.split(WOO_SEPARATOR)):
        return value
    return None


def resolve_part(
    attribute: str,
    part: str,
    vocabulary: ShopVocabulary,
    rules: dict[str, Any],
    siblings: list[dict[str, str]] | None = None,
    extra_approved: dict[str, set[str]] | None = None,
) -> tuple[str | None, str, str]:
    """Zwraca (wartosc docelowa albo None, rodzaj decyzji, uwaga) dla jednej czesci wartosci."""
    existing = vocabulary.canonical(attribute, part)
    if existing == part:
        return part, KIND_EXISTING, ""
    if existing:
        return existing, KIND_SAME_SLUG, f"w sklepie zapisany jako '{existing}'"

    found, target = _rule_lookup((rules.get("value_map") or {}).get(attribute) or {}, part)
    if found:
        if target is None:
            return None, KIND_REMOVED, "value_map: brak odpowiednika w sklepie"
        target = str(target)
        missing = [item for item in target.split(WOO_SEPARATOR) if not vocabulary.canonical(attribute, item)]
        note = "value_map" + (f"; UWAGA: tez spoza sklepu: {', '.join(missing)}" if missing else "")
        return target, KIND_MAPPED, note

    pieces = [piece for piece in re.split(r"\s*,\s*", part) if piece]
    if len(pieces) > 1:
        canon = [vocabulary.canonical(attribute, piece) for piece in pieces]
        if all(canon):
            return WOO_SEPARATOR.join(dict.fromkeys(canon)), KIND_SPLIT, "przecinek tworzy jeden nowy termin; w sklepie to osobne terminy"  # type: ignore[arg-type]

    if _approved(rules, extra_approved, attribute, part):
        return part, KIND_APPROVED, "approved_new_values"

    sibling = _sibling_value(vocabulary, attribute, siblings)
    if sibling:
        return sibling, KIND_SIBLING, f"rodzenstwo w sklepie ({len(siblings or [])} szt.)"

    similar = vocabulary.similar(attribute, part)
    note = f"podobny termin sklepu: '{similar}'" if similar else "brak podobnego terminu w sklepie"
    if is_technical(attribute, rules):
        return part, KIND_NEW_TECHNICAL, f"utworzy nowy termin; {note}"
    return None, KIND_NEW, note


def check_features(
    sku: str,
    title: str,
    features: dict[str, str],
    vocabulary: ShopVocabulary,
    rules: dict[str, Any],
    siblings: list[dict[str, str]] | None = None,
    keep_new: bool = False,
    extra_approved: dict[str, set[str]] | None = None,
    attributes: Iterable[str] | None = None,
) -> tuple[dict[str, str], list[ValueCheck]]:
    """Sprawdza (i poprawia) wartosci cech. Zwraca nowe cechy i liste ustalen.

    ``attributes`` zaweza kontrole do wybranych atrybutow (np. tych, ktore narzedzie zmienilo).
    """
    only = set(attributes) if attributes is not None else None
    fixed: dict[str, str] = {}
    checks: list[ValueCheck] = []
    for attribute, raw in features.items():
        value = cell_text(raw)
        if (only is not None and attribute not in only) or not value or is_identifier(attribute, rules):
            fixed[attribute] = raw
            continue
        if not vocabulary.has_attribute(attribute):
            fixed[attribute] = raw
            parts = [item.strip() for item in value.split(WOO_SEPARATOR) if item.strip()]
            if all(_approved(rules, extra_approved, attribute, part) for part in parts):
                continue
            checks.append(ValueCheck(sku, title, attribute, value, value, KIND_NO_DICTIONARY, STATUS_BY_KIND[KIND_NO_DICTIONARY],
                                     "atrybutu nie ma w zrzucie sklepu - wartosci nie da sie sprawdzic"))
            continue
        resolved: list[str] = []
        for part in [item.strip() for item in value.split(WOO_SEPARATOR) if item.strip()]:
            target, kind, note = resolve_part(attribute, part, vocabulary, rules, siblings, extra_approved)
            if kind == KIND_NEW and keep_new:
                target = part
            if kind != KIND_EXISTING:
                checks.append(ValueCheck(sku, title, attribute, part, target, kind, STATUS_BY_KIND[kind], note))
            if target:
                resolved.extend(item for item in target.split(WOO_SEPARATOR) if item not in resolved)
        if resolved:
            fixed[attribute] = WOO_SEPARATOR.join(resolved)
    checks.extend(check_type_links(sku, title, fixed, vocabulary))
    return fixed, checks


def check_type_links(sku: str, title: str, features: dict[str, str], vocabulary: ShopVocabulary) -> list[ValueCheck]:
    """Podtyp produktu musi byc przypiety w filtrze sklepu do Typu produktu tego produktu."""
    podtyp = cell_text(features.get(PODTYP))
    if not podtyp or not vocabulary.parents.get(PODTYP):
        return []
    types = {part.strip() for part in cell_text(features.get(TYP)).split(WOO_SEPARATOR) if part.strip()}
    checks: list[ValueCheck] = []
    for part in [item.strip() for item in podtyp.split(WOO_SEPARATOR) if item.strip()]:
        if not vocabulary.canonical(PODTYP, part):
            continue
        parent = vocabulary.parent(PODTYP, vocabulary.canonical(PODTYP, part) or part)
        if parent is None:
            checks.append(ValueCheck(sku, title, PODTYP, part, part, KIND_NO_PARENT, STATUS_BY_KIND[KIND_NO_PARENT],
                                     "termin nie ma powiazania z typem (ext_pf_cross_parent_term_id) - filtr go nie pokaze"))
        elif parent not in types:
            have = ", ".join(sorted(types)) or "brak typu"
            checks.append(ValueCheck(sku, title, PODTYP, part, part, KIND_TYPE_LINK, STATUS_BY_KIND[KIND_TYPE_LINK],
                                     f"w sklepie podtyp nalezy do typu '{parent}', produkt ma: {have}"))
    return checks


# --------------------------------------------------------------------------- raporty
def new_value_summary(checks: list[ValueCheck]) -> list[dict[str, Any]]:
    """Wartosci, ktore utworzylyby nowy termin, zgrupowane po (atrybut, wartosc)."""
    grouped: dict[tuple[str, str], dict[str, Any]] = {}
    for check in checks:
        if check.kind in (KIND_EXISTING, KIND_SAME_SLUG, KIND_TYPE_LINK, KIND_NO_PARENT):
            continue
        entry = grouped.setdefault((check.attribute, check.value), {
            "atrybut": check.attribute, "wartosc": check.value, "decyzja": check.kind, "status": check.status,
            "nowa_wartosc": check.new_value, "uwagi": check.note, "produktow": 0, "sku": [],
        })
        entry["produktow"] += 1
        entry["sku"].append(check.sku)
    order = {"DO_SPRAWDZENIA": 0, "WARNING": 1, "OK": 2}
    return sorted(grouped.values(), key=lambda item: (order.get(item["status"], 3), item["atrybut"], -item["produktow"]))


def summary_dict(checks: list[ValueCheck]) -> dict[str, Any]:
    grouped = new_value_summary(checks)
    unresolved = [item for item in grouped if item["decyzja"] == KIND_NEW]
    return {
        "zrodlo_slownika": "dictionaries/woo_attribute_terms_snapshot.tsv",
        "ustalen": len(checks),
        "wedlug_decyzji": dict(Counter(check.kind for check in checks).most_common()),
        "nowe_wartosci_bez_odpowiednika": len(unresolved),
        "niezgodnosci_typ_podtyp": sum(1 for check in checks if check.kind == KIND_TYPE_LINK),
        "wartosci": [{key: value for key, value in item.items() if key != "sku"} | {"sku": item["sku"][:12]} for item in grouped],
        "typ_podtyp": [check.__dict__ for check in checks if check.kind in (KIND_TYPE_LINK, KIND_NO_PARENT)],
    }


def write_check_sheets(workbook: Any, checks: list[ValueCheck], prefix: str = "") -> None:
    """Dopisuje arkusze 'Nowe wartosci' (z kolumna Decyzja) i 'Kontrola wartosci' do skoroszytu openpyxl."""
    from openpyxl.styles import Font, PatternFill

    fills = {"WARNING": PatternFill("solid", fgColor="FFF2CC"), "DO_SPRAWDZENIA": PatternFill("solid", fgColor="F8CBAD")}
    grouped = workbook.create_sheet(f"{prefix}Nowe wartosci")
    grouped.append(["Decyzja uzytkownika", "Atrybut", "Wartosc z pliku", "Co zrobiono", "Wartosc w imporcie",
                     "Status", "Produktow", "Uwagi", "SKU"])
    for item in new_value_summary(checks):
        grouped.append(["", item["atrybut"], item["wartosc"], item["decyzja"], item["nowa_wartosc"] or "(usunieta)",
                        item["status"], item["produktow"], item["uwagi"], ", ".join(item["sku"][:15])])
        fill = fills.get(item["status"])
        if fill:
            grouped.cell(grouped.max_row, 6).fill = fill
    detail = workbook.create_sheet(f"{prefix}Kontrola wartosci")
    detail.append(["SKU", "Tytul", "Atrybut", "Wartosc z pliku", "Wartosc w imporcie", "Co zrobiono", "Status", "Uwagi"])
    for check in checks:
        detail.append([check.sku, check.title, check.attribute, check.value, check.new_value or "(usunieta)",
                       check.kind, check.status, check.note])
        fill = fills.get(check.status)
        if fill:
            detail.cell(detail.max_row, 7).fill = fill
    for sheet, widths in ((grouped, [18, 26, 40, 30, 32, 16, 10, 60, 60]), (detail, [24, 60, 26, 40, 32, 30, 16, 60])):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for index, width in enumerate(widths):
            sheet.column_dimensions[chr(ord("A") + index)].width = width
        sheet.freeze_panes = "A2"


def group_checks_by_sku(checks: list[ValueCheck]) -> dict[str, list[ValueCheck]]:
    result: dict[str, list[ValueCheck]] = defaultdict(list)
    for check in checks:
        result[check.sku].append(check)
    return result
