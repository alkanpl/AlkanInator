"""Uzupelnia tagi wariantowe nowych produktow Simon na podstawie rodzenstwa.

Sklep laczy warianty wtyczka WPC Linked Variation po tagu produktu. Nowy kolor
istniejacego produktu (np. kaszmir /42, /142) ma dostac ten sam tag, co
pozostale kolory tego samego symbolu bazowego, zeby istniejaca grupa wariantow
objela takze nowy kolor.

Wejscie to pelny eksport WP All Export (XLSX). Wynik jest pelna kopia pliku -
zmienia sie wylacznie kolumna ``Product Tags`` w nowych produktach.

Uzycie:
    py src/fill_variant_tags_from_siblings.py \
        --input "input/Kaszmir Simon/Simon-55-i-55-Go-i-54.xlsx" \
        --color-codes 42,142 \
        --output output/Simon-55-i-55-Go-i-54_TAGI_KASZMIR.xlsx \
        --control output/Simon-55-i-55-Go-i-54_TAGI_KASZMIR_kontrola.xlsx \
        --report reports/simon_tagi_kaszmir_report.json
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from fix_simon_attributes_to_shop_model import DEFAULT_DICTIONARY, load_dictionary, symbol_base

ROOT = Path(__file__).resolve().parents[1]
TAG_SEPARATOR = "|"
SKU_COLUMN = "Sku"
TAGS_COLUMN = "Product Tags"
TITLE_COLUMN = "Title"
STATUS_COLUMN = "Status"


@dataclass
class TagDecision:
    sku: str
    title: str
    status: str
    tags: list[str] = field(default_factory=list)
    source: str = ""
    sibling_count: int = 0
    sibling_tag_counts: dict[str, int] = field(default_factory=dict)
    variant_groups: list[str] = field(default_factory=list)
    note: str = ""


def split_tags(value: Any) -> list[str]:
    return [part.strip() for part in str(value or "").split(TAG_SEPARATOR) if part.strip()]


def tag_slug(tag: str) -> str:
    text = tag.replace("ł", "l").replace("Ł", "L")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", text)).strip("-")


def color_code(sku: str, config: dict[str, Any]) -> str | None:
    value = (sku or "").strip()
    suffix = (config.get("symbol_base") or {}).get("strip_supplier_suffix")
    if suffix and value.endswith(suffix):
        value = value[: -len(suffix)]
    match = re.search(r"/(\d{2,3})$", value)
    return match.group(1) if match else None


def load_variant_groups(path: Path | None) -> dict[str, list[str]]:
    """Slug tagu -> lista opisow opublikowanych grup wariantow (nazwa + osie)."""
    groups: dict[str, list[str]] = defaultdict(list)
    if not path or not path.exists():
        return groups
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        parts = line.split("\t")
        if len(parts) < 6 or parts[1] != "publish":
            continue
        for slug in parts[3].split(","):
            if slug.strip():
                groups[slug.strip()].append(f"{parts[0]} [{parts[5]}]")
    return groups


def choose_tags(
    sibling_tags: list[list[str]], base: str, config: dict[str, Any]
) -> tuple[list[str], str, dict[str, int]]:
    settings = config.get("variant_tags") or {}
    ignored = set(settings.get("ignored_tags") or [])
    counts: Counter[str] = Counter()
    tagged = 0
    for tags in sibling_tags:
        useful = [tag for tag in tags if tag not in ignored]
        if useful:
            tagged += 1
            counts.update(set(useful))
    if not sibling_tags:
        return [], "brak rodzenstwa w pliku", {}
    if not counts:
        return [], "rodzenstwo nie ma tagow wariantowych", {}
    override = (settings.get("overrides") or {}).get(base)
    if override:
        return [override], "rozstrzygniecie ze slownika (overrides)", dict(counts)
    majority = [tag for tag, count in counts.items() if count * 2 > tagged]
    if majority:
        unanimous = all(counts[tag] == tagged for tag in majority) and len(counts) == len(majority)
        source = "jednoznaczny tag rodzenstwa" if unanimous else "wiekszosc rodzenstwa"
        return sorted(majority, key=lambda tag: -counts[tag]), source, dict(counts)
    return [], "rodzenstwo rozbite na rowne grupy - wymaga decyzji", dict(counts)


def decide_tags(
    rows: list[dict[str, Any]], color_codes: set[str], config: dict[str, Any], groups: dict[str, list[str]]
) -> tuple[list[TagDecision], list[dict[str, Any]]]:
    siblings: dict[str, list[dict[str, Any]]] = defaultdict(list)
    targets: list[dict[str, Any]] = []
    for row in rows:
        sku = str(row.get(SKU_COLUMN) or "")
        if not sku:
            continue
        if color_code(sku, config) in color_codes:
            targets.append(row)
        else:
            siblings[symbol_base(sku, config)].append(row)

    decisions: list[TagDecision] = []
    anomalies: list[dict[str, Any]] = []
    for row in targets:
        sku = str(row[SKU_COLUMN])
        base = symbol_base(sku, config)
        family = siblings.get(base, [])
        decision = TagDecision(
            sku=sku,
            title=str(row.get(TITLE_COLUMN) or ""),
            status=str(row.get(STATUS_COLUMN) or ""),
            sibling_count=len(family),
        )
        existing = split_tags(row.get(TAGS_COLUMN))
        if existing:
            decision.tags = existing
            decision.source = "produkt mial juz tagi - bez zmian"
        else:
            decision.tags, decision.source, decision.sibling_tag_counts = choose_tags(
                [split_tags(item.get(TAGS_COLUMN)) for item in family], base, config
            )
        for tag in decision.tags:
            decision.variant_groups += groups.get(tag_slug(tag), [])
        if decision.tags and not decision.variant_groups and groups:
            decision.note = "tag nie ma grupy wariantow w lokalnej kopii sklepu - sprawdz na produkcji"
        decisions.append(decision)

        if decision.tags and family:
            ignored = set((config.get("variant_tags") or {}).get("ignored_tags") or [])
            for item in family:
                item_tags = [tag for tag in split_tags(item.get(TAGS_COLUMN)) if tag not in ignored]
                if set(item_tags) != set(decision.tags):
                    anomalies.append(
                        {
                            "sku": str(item.get(SKU_COLUMN)),
                            "status": str(item.get(STATUS_COLUMN) or ""),
                            "tagi_obecne": TAG_SEPARATOR.join(item_tags),
                            "tagi_zachowane": TAG_SEPARATOR.join(
                                tag for tag in split_tags(item.get(TAGS_COLUMN)) if tag in ignored
                            ),
                            "tag_rodziny": TAG_SEPARATOR.join(decision.tags),
                            "problem": "brak tagu wariantowego" if not item_tags else "inny tag niz reszta rodziny",
                        }
                    )
    return decisions, anomalies


def fixed_sibling_tags(anomaly: dict[str, Any]) -> str:
    """Tagi po naprawie: zachowane tagi niewariantowe + tag rodziny."""
    kept = split_tags(anomaly.get("tagi_zachowane"))
    return TAG_SEPARATOR.join(kept + [tag for tag in split_tags(anomaly["tag_rodziny"]) if tag not in kept])


def process_workbook(
    input_path: Path,
    output_path: Path,
    color_codes: set[str],
    config: dict[str, Any],
    groups: dict[str, list[str]],
    fix_sibling_anomalies: bool = False,
) -> tuple[list[TagDecision], list[dict[str, Any]], dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(input_path)
    sheet = workbook.worksheets[0]
    header = [cell.value for cell in sheet[1]]
    sku_index = header.index(SKU_COLUMN)
    tags_index = header.index(TAGS_COLUMN)
    rows: list[dict[str, Any]] = []
    row_numbers: dict[str, int] = {}
    for number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        values = {}
        for index, cell in enumerate(cells):
            name = header[index]
            if name is not None and name not in values:
                values[name] = cell.value
        rows.append(values)
        if cells[sku_index].value:
            row_numbers[str(cells[sku_index].value)] = number

    decisions, anomalies = decide_tags(rows, color_codes, config, groups)
    written = 0
    for decision in decisions:
        if decision.tags and decision.source != "produkt mial juz tagi - bez zmian":
            sheet.cell(row_numbers[decision.sku], tags_index + 1).value = TAG_SEPARATOR.join(decision.tags)
            written += 1
    fixed = 0
    touched = {decision.sku for decision in decisions}
    for anomaly in anomalies:
        anomaly["naprawiono"] = False
        if fix_sibling_anomalies:
            anomaly["tagi_po_naprawie"] = fixed_sibling_tags(anomaly)
            sheet.cell(row_numbers[anomaly["sku"]], tags_index + 1).value = anomaly["tagi_po_naprawie"]
            anomaly["naprawiono"] = True
            touched.add(anomaly["sku"])
            fixed += 1
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    validation = validate_output(input_path, output_path, tags_index, touched, sku_index)
    validation["tags_written"] = written
    validation["sibling_tags_fixed"] = fixed
    return decisions, anomalies, validation


def validate_output(
    input_path: Path, output_path: Path, tags_index: int, target_skus: set[str], sku_index: int
) -> dict[str, Any]:
    from openpyxl import load_workbook

    source = list(load_workbook(input_path, read_only=True).worksheets[0].iter_rows(values_only=True))
    result = list(load_workbook(output_path, read_only=True).worksheets[0].iter_rows(values_only=True))
    foreign_changes = 0
    for before, after in zip(source, result):
        for index, (old, new) in enumerate(zip(before, after)):
            if old == new:
                continue
            if index == tags_index and str(before[sku_index]) in target_skus:
                continue
            foreign_changes += 1
    mojibake = sum(
        1
        for row in result
        for value in row
        if isinstance(value, str) and any(marker in value for marker in ("Ä…", "Å‚", "Ã³", "Ĺ‚", "�"))
    )
    source_mojibake = sum(
        1
        for row in source
        for value in row
        if isinstance(value, str) and any(marker in value for marker in ("Ä…", "Å‚", "Ã³", "Ĺ‚", "�"))
    )
    return {
        "rows_match": len(source) == len(result),
        "header_match": source[0] == result[0],
        "changes_outside_scope": foreign_changes,
        "new_mojibake_cells": mojibake - source_mojibake,
        "passed": len(source) == len(result) and source[0] == result[0] and foreign_changes == 0 and mojibake == source_mojibake,
    }


def write_control_workbook(path: Path, decisions: list[TagDecision], anomalies: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tagi nowych produktow"
    sheet.append(["SKU", "Nazwa", "Status", "Nadany tag", "Zrodlo decyzji", "Rodzenstwa", "Tagi rodzenstwa (ile szt.)", "Grupa wariantow w sklepie [osie]", "Uwaga"])
    warn = PatternFill("solid", fgColor="FFF2CC")
    missing = PatternFill("solid", fgColor="F8CBAD")
    for item in decisions:
        sheet.append([
            item.sku, item.title, item.status, TAG_SEPARATOR.join(item.tags), item.source, item.sibling_count,
            "; ".join(f"{tag} ({count})" for tag, count in item.sibling_tag_counts.items()),
            "; ".join(item.variant_groups), item.note,
        ])
        row = sheet.max_row
        if not item.tags:
            sheet.cell(row, 4).fill = missing
        elif item.source != "jednoznaczny tag rodzenstwa" or item.note:
            sheet.cell(row, 5).fill = warn
    extra = workbook.create_sheet("Anomalie rodzenstwa")
    extra.append(["SKU", "Status", "Tagi wariantowe przed", "Tag rodziny", "Problem", "Naprawiono", "Tagi po naprawie"])
    for anomaly in anomalies:
        extra.append([
            anomaly["sku"], anomaly["status"], anomaly["tagi_obecne"], anomaly["tag_rodziny"], anomaly["problem"],
            "TAK" if anomaly.get("naprawiono") else "NIE", anomaly.get("tagi_po_naprawie", ""),
        ])
    for worksheet, widths in ((sheet, [22, 60, 10, 55, 36, 11, 60, 60, 50]), (extra, [22, 10, 55, 55, 32, 12, 70])):
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        for column, width in zip("ABCDEFGHI", widths):
            worksheet.column_dimensions[column].width = width
        worksheet.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--color-codes", required=True, help="Koncowki koloru nowych produktow, np. 42,142.")
    parser.add_argument("--control", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--dictionary", type=Path, default=DEFAULT_DICTIONARY)
    parser.add_argument(
        "--fix-sibling-anomalies",
        action="store_true",
        help="Nadaje tag rodziny takze ISTNIEJACYM produktom z blednym/brakujacym tagiem (tylko po zgodzie uzytkownika).",
    )
    args = parser.parse_args(argv)
    if args.output.resolve() == args.input.resolve():
        parser.error("Plik wyjsciowy nie moze nadpisywac wejsciowego.")

    config = load_dictionary(args.dictionary)
    snapshot = (config.get("variant_tags") or {}).get("variant_groups_snapshot")
    groups = load_variant_groups(ROOT / snapshot if snapshot else None)
    codes = {code.strip() for code in args.color_codes.split(",") if code.strip()}
    decisions, anomalies, validation = process_workbook(
        args.input, args.output, codes, config, groups, fix_sibling_anomalies=args.fix_sibling_anomalies
    )

    sources = Counter(decision.source for decision in decisions)
    report = {
        "input": str(args.input),
        "output": str(args.output),
        "color_codes": sorted(codes),
        "new_products": len(decisions),
        "tagged": sum(1 for decision in decisions if decision.tags),
        "without_tag": [
            {"sku": decision.sku, "title": decision.title, "reason": decision.source}
            for decision in decisions
            if not decision.tags
        ],
        "sources": dict(sources),
        "sibling_anomalies": anomalies,
        "validation": validation,
        "decisions": [decision.__dict__ for decision in decisions],
    }
    if args.control:
        write_control_workbook(args.control, decisions, anomalies)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        f"Nowych produktow: {report['new_products']}, z tagiem: {report['tagged']}, "
        f"bez tagu: {len(report['without_tag'])}, anomalie rodzenstwa: {len(anomalies)} "
        f"(naprawione: {validation['sibling_tags_fixed']}), "
        f"walidacja: {'OK' if validation['passed'] else 'BLAD'} -> {args.output}"
    )
    return 0 if validation["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
