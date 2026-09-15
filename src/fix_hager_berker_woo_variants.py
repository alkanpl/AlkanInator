from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = (
    ROOT
    / "output"
    / "Hager-Berker-Gniazdka_WOO_ALL_IMPORT_CROSS_UP_SELL_PO_AUDYCIE_PDF_TYTULY_STARYCH_SKU.xlsx"
)
DEFAULT_OUTPUT = (
    ROOT
    / "output"
    / "Hager-Berker-Gniazdka_WOO_IMPORT_WARIANTY_POPRAWIONE.xlsx"
)
DEFAULT_REPORT = (
    ROOT
    / "reports"
    / "Hager-Berker-Gniazdka_WOO_IMPORT_WARIANTY_POPRAWIONE_report.json"
)

TAG_REMOVALS = {
    # Isolated products in otherwise valid multi-axis groups.
    "10411606/HAG": "Wariant - Puszka natynkowa B.3",
    "WL9016/HAG": "Wariant - Ramka antyczny Passion",
    "WL5010/HAG": "Wariant - Ramka Lumina2",
    "WDS5100/HAG": "Wariant - Mechanizm gniazda",
    "10116019/HAG": "Wariant - Ramka z polem opisowym Q.1",
    "10116189/HAG": "Wariant - Ramka lakier",
    "WWD1100PW/HAG": "Wariant - Ramka z uszczelką IP44",
    # Entire two-product groups differing on more than one technical axis.
    "WL2120/HAG": "Wariant - Gniazdo komputerowe podwójne RJ45 Lumina2",
    "WL2150/HAG": "Wariant - Gniazdo komputerowe podwójne RJ45 Lumina2",
    "WDF410866/HAG": "Wariant - Mechanizm podwójnego gniazda komputerowego UAE RJ45 ekranowany",
    "WDF410766/HAG": "Wariant - Mechanizm podwójnego gniazda komputerowego UAE RJ45 ekranowany",
    "WDE303808/HAG": "Wariant - Łącznik schodowy podwójny 2-klawiszowy mechanizm zaciski śrubowe",
    "WDEM3036/HAG": "Wariant - Łącznik schodowy podwójny 2-klawiszowy mechanizm zaciski śrubowe",
}

SHUTTERS_NO = {
    "WTS1100BG/HAG",
    "WTS1100WG/HAG",
    "47157006/HAG",
    "47157009/HAG",
    "47157013/HAG",
    "47157014/HAG",
    "47157015/HAG",
    "WL1052/HAG",
    "WL1053/HAG",
    "WLS1100WG/HAG",
    "WLS1100BG/HAG",
}

SHUTTERS_YES = {
    "WTS1200WG/HAG",
    "WTS1200BG/HAG",
    "47357006/HAG",
    "47357009/HAG",
    "WL1060/HAG",
    "WL1062/HAG",
    "WL1063/HAG",
    "WLS1200WG/HAG",
    "WLS1200BG/HAG",
}

WLS1200_SKUS = {"WLS1200WG/HAG", "WLS1200BG/HAG"}
MOJIBAKE_MARKERS = ("\ufffd", "Ă", "Ĺ", "Ã", "Â", "â€")
ALLOWED_CHANGED_COLUMNS = {
    "Product Tags",
    "Atrybut Produktu: Przesłony Torów Prądowych",
    "Atrybut Produktu: Typ Zacisków",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def split_tags(value: object) -> list[str]:
    return [tag.strip() for tag in re.split(r"\s*\|\s*", str(value or "")) if tag.strip()]


def json_value(value: object) -> object:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def variant_tags(value: object) -> list[str]:
    return [tag for tag in split_tags(value) if tag.casefold().startswith("wariant -")]


def audit_woo_variants(input_path: Path, output_path: Path) -> dict[str, object]:
    if not input_path.exists():
        raise FileNotFoundError(f"Brak pliku wejściowego: {input_path}")
    if not output_path.exists():
        raise FileNotFoundError(f"Brak pliku wynikowego: {output_path}")

    source_hash_before = sha256(input_path)
    output_hash_before = sha256(output_path)
    # Zwykły tryb jest tu celowo szybszy od read_only: audyt porównuje wiele
    # komórek po współrzędnych, a losowy dostęp w arkuszu strumieniowym jest kosztowny.
    input_book = load_workbook(input_path, read_only=False, data_only=False)
    output_book = load_workbook(output_path, read_only=False, data_only=False)
    try:
        input_sheet_names = input_book.sheetnames
        output_sheet_names = output_book.sheetnames
        sheet_names_match = input_sheet_names == output_sheet_names
        if not sheet_names_match or len(input_book.worksheets) != 1:
            raise ValueError(
                "Audyt wymaga jednego, zgodnego arkusza w pliku wejściowym i wynikowym."
            )

        input_sheet = input_book.active
        output_sheet = output_book.active
        dimensions_match = (
            input_sheet.max_row == output_sheet.max_row
            and input_sheet.max_column == output_sheet.max_column
        )
        if not dimensions_match:
            raise ValueError("Wymiary arkusza wejściowego i wynikowego są różne.")

        input_headers = [cell.value for cell in input_sheet[1]]
        output_headers = [cell.value for cell in output_sheet[1]]
        headers_match = input_headers == output_headers
        if not headers_match:
            raise ValueError("Nagłówki arkusza wejściowego i wynikowego są różne.")

        header_names = [str(value or "").strip() for value in output_headers]
        headers = {
            name: index
            for index, name in enumerate(header_names, start=1)
            if name
        }
        required = {"SKU", *ALLOWED_CHANGED_COLUMNS}
        missing = required - set(headers)
        if missing:
            raise KeyError(f"Brak wymaganych kolumn: {sorted(missing)}")

        sku_column = headers["SKU"]
        tag_column = headers["Product Tags"]
        shutters_column = headers["Atrybut Produktu: Przesłony Torów Prądowych"]
        terminals_column = headers["Atrybut Produktu: Typ Zacisków"]

        input_rows: dict[str, int] = {}
        output_rows: dict[str, int] = {}
        duplicate_input_skus: list[str] = []
        duplicate_output_skus: list[str] = []
        for row in range(2, output_sheet.max_row + 1):
            input_sku = str(input_sheet.cell(row, sku_column).value or "").strip().upper()
            output_sku = str(output_sheet.cell(row, sku_column).value or "").strip().upper()
            if input_sku:
                if input_sku in input_rows:
                    duplicate_input_skus.append(input_sku)
                input_rows[input_sku] = row
            if output_sku:
                if output_sku in output_rows:
                    duplicate_output_skus.append(output_sku)
                output_rows[output_sku] = row

        changes: list[dict[str, object]] = []
        changed_columns: Counter[str] = Counter()
        unexpected_changes: list[dict[str, object]] = []
        changed_products: set[str] = set()
        for row in range(1, output_sheet.max_row + 1):
            sku = str(output_sheet.cell(row, sku_column).value or "").strip().upper()
            for column in range(1, output_sheet.max_column + 1):
                before = input_sheet.cell(row, column).value
                after = output_sheet.cell(row, column).value
                if before == after:
                    continue
                header = header_names[column - 1] or f"Kolumna {column}"
                change = {
                    "cell": output_sheet.cell(row, column).coordinate,
                    "sku": sku,
                    "column": header,
                    "before": json_value(before),
                    "after": json_value(after),
                }
                changes.append(change)
                changed_columns[header] += 1
                if sku:
                    changed_products.add(sku)
                if header not in ALLOWED_CHANGED_COLUMNS:
                    unexpected_changes.append(change)

        required_skus = set(TAG_REMOVALS) | SHUTTERS_NO | SHUTTERS_YES | WLS1200_SKUS
        missing_required_skus = sorted(required_skus - set(output_rows))
        tag_removal_errors: list[str] = []
        shutters_errors: list[str] = []
        terminals_errors: list[str] = []
        expected_changed_cells: set[str] = set()

        for sku, removed_tag in sorted(TAG_REMOVALS.items()):
            row = output_rows.get(sku)
            input_row = input_rows.get(sku)
            if row is None or input_row is None:
                tag_removal_errors.append(f"{sku}: brak SKU")
                continue
            before_tags = split_tags(input_sheet.cell(input_row, tag_column).value)
            after_tags = split_tags(output_sheet.cell(row, tag_column).value)
            expected_after = [tag for tag in before_tags if tag != removed_tag]
            expected_changed_cells.add(output_sheet.cell(row, tag_column).coordinate)
            if removed_tag not in before_tags:
                tag_removal_errors.append(f"{sku}: źródło nie zawierało tagu {removed_tag}")
            if after_tags != expected_after:
                tag_removal_errors.append(f"{sku}: niezgodny zestaw tagów po korekcie")

        for sku in sorted(SHUTTERS_NO | SHUTTERS_YES):
            row = output_rows.get(sku)
            if row is None:
                shutters_errors.append(f"{sku}: brak SKU")
                continue
            expected = "Tak" if sku in SHUTTERS_YES else "Nie"
            actual = output_sheet.cell(row, shutters_column).value
            expected_changed_cells.add(output_sheet.cell(row, shutters_column).coordinate)
            if actual != expected:
                shutters_errors.append(f"{sku}: oczekiwano {expected}, otrzymano {actual}")

        for sku in sorted(WLS1200_SKUS):
            row = output_rows.get(sku)
            if row is None:
                terminals_errors.append(f"{sku}: brak SKU")
                continue
            actual = output_sheet.cell(row, terminals_column).value
            expected_changed_cells.add(output_sheet.cell(row, terminals_column).coordinate)
            if actual != "Samozaciski":
                terminals_errors.append(f"{sku}: oczekiwano Samozaciski, otrzymano {actual}")

        actual_changed_cells = {str(change["cell"]) for change in changes}
        changes_limited_to_expected_cells = actual_changed_cells <= expected_changed_cells

        input_variant_tags: set[str] = set()
        output_variant_tags: set[str] = set()
        input_variant_products = 0
        output_variant_products = 0
        suspicious_cells: list[str] = []
        for row in range(2, output_sheet.max_row + 1):
            before_tags = variant_tags(input_sheet.cell(row, tag_column).value)
            after_tags = variant_tags(output_sheet.cell(row, tag_column).value)
            if before_tags:
                input_variant_products += 1
                input_variant_tags.update(before_tags)
            if after_tags:
                output_variant_products += 1
                output_variant_tags.update(after_tags)
            for column in range(1, output_sheet.max_column + 1):
                value = output_sheet.cell(row, column).value
                if isinstance(value, str) and any(marker in value for marker in MOJIBAKE_MARKERS):
                    suspicious_cells.append(output_sheet.cell(row, column).coordinate)
                    if len(suspicious_cells) >= 20:
                        break
            if len(suspicious_cells) >= 20:
                break

        validation = {
            "sheet_names_match": sheet_names_match,
            "dimensions_match": dimensions_match,
            "headers_match": headers_match,
            "product_rows_match": len(input_rows) == len(output_rows),
            "sku_order_and_values_match": list(input_rows) == list(output_rows),
            "no_duplicate_input_skus": not duplicate_input_skus,
            "no_duplicate_output_skus": not duplicate_output_skus,
            "all_required_skus_present": not missing_required_skus,
            "all_requested_variant_tags_removed": not tag_removal_errors,
            "all_shutters_values_valid": not shutters_errors,
            "all_wls1200_terminal_values_valid": not terminals_errors,
            "changes_limited_to_expected_cells": changes_limited_to_expected_cells,
            "no_changes_outside_allowed_columns": not unexpected_changes,
            "no_mojibake_detected": not suspicious_cells,
        }
        validation["passed"] = all(validation.values())

        report: dict[str, object] = {
            "report_type": "Audyt korekty wariantów Hager/Berker dla importu WooCommerce",
            "generated_at": datetime.now().astimezone().isoformat(),
            "input_file": str(input_path),
            "input_sha256": source_hash_before,
            "output_file": str(output_path),
            "output_sha256": output_hash_before,
            "output_modified_at": datetime.fromtimestamp(output_path.stat().st_mtime).astimezone().isoformat(),
            "sheets": output_sheet_names,
            "rows_with_sku": len(output_rows),
            "worksheet_rows_including_header": output_sheet.max_row,
            "worksheet_columns": output_sheet.max_column,
            "variant_groups_before": len(input_variant_tags),
            "variant_groups_after": len(output_variant_tags),
            "variant_products_before": input_variant_products,
            "variant_products_after": output_variant_products,
            "removed_variant_tags": len(TAG_REMOVALS),
            "validated_shutters_values": len(SHUTTERS_NO | SHUTTERS_YES),
            "filled_wls1200_terminal_values": len(WLS1200_SKUS),
            "changed_cells": len(changes),
            "changed_products": len(changed_products),
            "changes_by_column": dict(sorted(changed_columns.items())),
            "missing_required_skus": missing_required_skus,
            "duplicate_input_skus": sorted(set(duplicate_input_skus)),
            "duplicate_output_skus": sorted(set(duplicate_output_skus)),
            "tag_removal_errors": tag_removal_errors,
            "shutters_errors": shutters_errors,
            "terminal_errors": terminals_errors,
            "unexpected_changes": unexpected_changes,
            "suspicious_encoding_cells": suspicious_cells,
            "validation": validation,
            "changes": changes,
        }
    finally:
        input_book.close()
        output_book.close()

    if sha256(input_path) != source_hash_before:
        raise RuntimeError("Plik wejściowy został nieoczekiwanie zmieniony podczas audytu.")
    if sha256(output_path) != output_hash_before:
        raise RuntimeError("Plik wynikowy został nieoczekiwanie zmieniony podczas audytu.")
    return report


def write_report(report: dict[str, object], report_path: Path, overwrite: bool = False) -> None:
    if report_path.exists() and not overwrite:
        raise FileExistsError(f"Raport już istnieje: {report_path}")
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report["report_file"] = str(report_path)
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def update_woo_variants(input_path: Path, output_path: Path) -> dict[str, object]:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("Plik wyjściowy musi być inną kopią niż plik wejściowy.")
    if output_path.exists():
        raise FileExistsError(f"Plik wyjściowy już istnieje: {output_path}")

    source_hash_before = sha256(input_path)
    workbook = load_workbook(input_path)
    sheet = workbook.active
    headers = {
        str(cell.value).strip(): cell.column
        for cell in sheet[1]
        if cell.value is not None and str(cell.value).strip()
    }
    required = {
        "SKU",
        "Product Tags",
        "Atrybut Produktu: Przesłony Torów Prądowych",
        "Atrybut Produktu: Typ Zacisków",
    }
    missing = required - set(headers)
    if missing:
        raise KeyError(f"Brak wymaganych kolumn: {sorted(missing)}")

    sku_column = headers["SKU"]
    tag_column = headers["Product Tags"]
    shutters_column = headers["Atrybut Produktu: Przesłony Torów Prądowych"]
    terminals_column = headers["Atrybut Produktu: Typ Zacisków"]
    row_by_sku: dict[str, int] = {}
    for row in range(2, sheet.max_row + 1):
        sku = str(sheet.cell(row, sku_column).value or "").strip().upper()
        if sku:
            if sku in row_by_sku:
                raise ValueError(f"Powtórzone SKU: {sku}")
            row_by_sku[sku] = row

    required_skus = set(TAG_REMOVALS) | SHUTTERS_NO | SHUTTERS_YES | WLS1200_SKUS
    missing_skus = sorted(required_skus - set(row_by_sku))
    if missing_skus:
        raise KeyError(f"Brak SKU wymaganych do korekty: {missing_skus}")

    changes: list[dict[str, object]] = []
    for sku, tag_to_remove in sorted(TAG_REMOVALS.items()):
        row = row_by_sku[sku]
        cell = sheet.cell(row, tag_column)
        before_tags = split_tags(cell.value)
        if tag_to_remove not in before_tags:
            raise ValueError(f"SKU {sku} nie ma oczekiwanego tagu: {tag_to_remove}")
        after_tags = [tag for tag in before_tags if tag != tag_to_remove]
        cell.value = " | ".join(after_tags)
        changes.append(
            {
                "sku": sku,
                "column": "Product Tags",
                "before": " | ".join(before_tags),
                "after": cell.value,
            }
        )

    for sku in sorted(SHUTTERS_NO | SHUTTERS_YES):
        row = row_by_sku[sku]
        cell = sheet.cell(row, shutters_column)
        expected = "Tak" if sku in SHUTTERS_YES else "Nie"
        before = str(cell.value or "")
        cell.value = expected
        if before != expected:
            changes.append(
                {
                    "sku": sku,
                    "column": "Atrybut Produktu: Przesłony Torów Prądowych",
                    "before": before,
                    "after": expected,
                }
            )

    for sku in sorted(WLS1200_SKUS):
        row = row_by_sku[sku]
        cell = sheet.cell(row, terminals_column)
        before = str(cell.value or "")
        cell.value = "Samozaciski"
        if before != "Samozaciski":
            changes.append(
                {
                    "sku": sku,
                    "column": "Atrybut Produktu: Typ Zacisków",
                    "before": before,
                    "after": "Samozaciski",
                }
            )

    for sku, removed_tag in TAG_REMOVALS.items():
        row = row_by_sku[sku]
        if removed_tag in split_tags(sheet.cell(row, tag_column).value):
            raise AssertionError(f"Nie usunięto tagu {removed_tag} z {sku}")
    for sku in SHUTTERS_NO:
        if sheet.cell(row_by_sku[sku], shutters_column).value != "Nie":
            raise AssertionError(f"Niepoprawna wartość przesłon dla {sku}")
    for sku in SHUTTERS_YES:
        if sheet.cell(row_by_sku[sku], shutters_column).value != "Tak":
            raise AssertionError(f"Niepoprawna wartość przesłon dla {sku}")
    for sku in WLS1200_SKUS:
        if sheet.cell(row_by_sku[sku], terminals_column).value != "Samozaciski":
            raise AssertionError(f"Niepoprawny typ zacisków dla {sku}")

    suspicious: list[str] = []
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and any(marker in cell.value for marker in MOJIBAKE_MARKERS):
                suspicious.append(cell.coordinate)
                if len(suspicious) >= 20:
                    break
        if len(suspicious) >= 20:
            break
    if suspicious:
        raise ValueError(f"Wykryto możliwe uszkodzenie kodowania: {suspicious}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

    source_hash_after = sha256(input_path)
    if source_hash_after != source_hash_before:
        raise RuntimeError("Plik wejściowy został nieoczekiwanie zmieniony.")

    return {
        "input": str(input_path),
        "output": str(output_path),
        "removed_variant_tags": len(TAG_REMOVALS),
        "validated_shutters_values": len(SHUTTERS_NO | SHUTTERS_YES),
        "filled_wls1200_terminal_values": len(WLS1200_SKUS),
        "changed_cells": len(changes),
        "source_unchanged": True,
        "source_sha256": source_hash_after,
        "changes": changes,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Tworzy kopię importu Woo z konserwatywnie poprawionymi wariantami Hager/Berker."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument(
        "--report-only",
        action="store_true",
        help="Audytuje istniejący plik wynikowy bez jego ponownego generowania.",
    )
    parser.add_argument(
        "--overwrite-report",
        action="store_true",
        help="Pozwala zastąpić istniejący raport, bez zmiany plików XLSX.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    input_path = args.input.resolve()
    output_path = args.output.resolve()
    report_path = args.report.resolve()
    if report_path in {input_path, output_path}:
        raise ValueError("Raport musi być osobnym plikiem JSON.")
    if report_path.exists() and not args.overwrite_report:
        raise FileExistsError(f"Raport już istnieje: {report_path}")
    if not args.report_only:
        update_woo_variants(input_path, output_path)
    result = audit_woo_variants(input_path, output_path)
    write_report(result, report_path, overwrite=args.overwrite_report)
    print(json.dumps(result, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
