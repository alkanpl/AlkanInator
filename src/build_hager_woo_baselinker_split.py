from __future__ import annotations

from copy import copy
import csv
import json
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
FINAL_WORKBOOK = ROOT / "output" / "Hager_nazwy_po_decyzjach_2026-08-07.xlsx"
SUPPLIER_WEB = ROOT / "input" / "Hager aparatura na www.xlsx"
IMPLEMENTATION_CONFIG = ROOT / "configs" / "hager_woocommerce_implementation.yaml"

WOO_OUTPUT = ROOT / "output" / "Hager_Woo_istniejace_96_produkty_2026-08-07.xlsx"
BASELINKER_OUTPUT = ROOT / "output" / "Hager_Baselinker_poza_Woo_57_produkty_2026-08-07.csv"
CONTROL_OUTPUT = ROOT / "output" / "Hager_Baselinker_poza_Woo_57_kontrola_2026-08-07.xlsx"
REPORT_OUTPUT = ROOT / "reports" / "hager_woo_baselinker_split_2026-08-07" / "summary.json"

NAVY = "17365D"
BLUE = "2F75B5"
TEAL = "008C95"
LIGHT_BLUE = "DCE6F1"
LIGHT_GREEN = "E2F0D9"
LIGHT_YELLOW = "FFF2CC"
WHITE = "FFFFFF"
GRID = "D9E2F3"
TEXT = "1F1F1F"


def load_yaml(path: Path) -> dict:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def normalize_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Tak" if value else "Nie"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_final_data():
    workbook = load_workbook(FINAL_WORKBOOK, read_only=True, data_only=True)

    worksheet_rows = list(workbook["Worksheet"].iter_rows(values_only=True))
    worksheet_headers = list(worksheet_rows[0])
    worksheet_index = {str(value): index for index, value in enumerate(worksheet_headers) if value is not None}

    name_rows = list(workbook["Nazwy po decyzjach"].iter_rows(min_row=6, values_only=True))
    names = {
        row[1]: {
            "code": row[2],
            "name": row[5],
            "product_type": row[7],
            "source_category": row[8],
            "target_category": row[9],
            "status": row[14],
            "source": row[15],
            "source_url": row[16],
        }
        for row in name_rows
        if row[1]
    }
    return worksheet_headers, worksheet_rows[1:], worksheet_index, names


def load_supplier_data():
    workbook = load_workbook(SUPPLIER_WEB, read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    headers = list(rows[0])
    index = {str(value): position for position, value in enumerate(headers) if value is not None}
    result = {}
    for row in rows[1:]:
        sku = normalize_value(row[index["Kod"]])
        if not sku:
            continue
        result[sku] = {
            "product_id": row[index["Numer Id"]],
            "www": row[index["www"]],
            "price_net": row[index["Nowa cena internetowa netto"]],
            "supplier_name": " ".join(
                part for part in [normalize_value(row[index["Nazwa"]]), normalize_value(row[index["Nazwa dodatkowa"]])] if part
            ),
        }
    return result


def build_attribute_mapping(config: dict):
    mapping = {}
    for item in config["attribute_actions"]:
        target = item["target_name"]
        for source in item["source_attributes"]:
            mapping[source] = target
        if item.get("existing_name"):
            mapping[item["existing_name"]] = target
    for item in config["reuse_without_new_attribute"]:
        target = item["target_name"]
        for source in item["source_attributes"]:
            mapping.setdefault(source, target)
        if item.get("existing_name"):
            mapping[item["existing_name"]] = target
    return mapping


def build_features(row, headers, names_data, attribute_mapping):
    prefix = "Atrybut Produktu: "
    features = {}
    sku = row[3]

    producer = "Hager"
    manufacturer_data = "HAGER, Fabryczna 10 43-100 Tychy, office@hager.pl"
    features["Producent"] = producer
    features["Dane producenta"] = manufacturer_data
    features["Typ produktu"] = names_data[sku]["product_type"]

    excluded = {
        "Producent", "Dane Producenta", "Typ Produktu",
        "Liczba Biegunów", "Typ Prądu Różnicowego",
        "Napięcie Cewki [V]", "Kompatybilność",
    }
    for index, header in enumerate(headers):
        if not isinstance(header, str) or not header.startswith(prefix):
            continue
        source_name = header[len(prefix) :]
        value = normalize_value(row[index])
        if not value or source_name in excluded:
            continue
        target_name = attribute_mapping.get(source_name, source_name)
        if target_name in {"Producent", "Dane producenta", "Typ produktu"}:
            continue
        features.setdefault(target_name, value)

    # Jawne preferencje po audycie duplikatów.
    direct_preference = {
        "Typ Wyłącznika": "Typ Wyłącznika",
        "Napięcie [V]": "Napięcie [V]",
        "Liczba Faz": "Liczba Faz",
        "Układ Biegunów": "Układ biegunów",
        "Ilość Styków Zwiernych/rozwiernych": "Konfiguracja styków (NO/NC)",
        "Znamionowy Przekrój Poprzeczny": "Przekrój poprzeczny [mm²]",
    }
    header_index = {str(header): index for index, header in enumerate(headers) if header is not None}
    for source_name, target_name in direct_preference.items():
        header = prefix + source_name
        if header not in header_index:
            continue
        value = normalize_value(row[header_index[header]])
        if value:
            features[target_name] = value

    return features


def copy_worksheet(source, target):
    for row in source.iter_rows():
        for cell in row:
            new_cell = target[cell.coordinate]
            new_cell.value = cell.value
            if cell.has_style:
                new_cell._style = copy(cell._style)
                new_cell.number_format = cell.number_format
                new_cell.alignment = copy(cell.alignment)
                new_cell.protection = copy(cell.protection)
    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key].width = dimension.width
        target.column_dimensions[key].hidden = dimension.hidden
    for key, dimension in source.row_dimensions.items():
        target.row_dimensions[key].height = dimension.height
        target.row_dimensions[key].hidden = dimension.hidden
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines


def add_control_title(ws, title, subtitle, end_column):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    ws.cell(1, 1, title)
    ws.cell(1, 1).font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).font = Font(name="Aptos", size=10, color="44546A")
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 34


def add_table(ws, start_row, headers, rows, name):
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(start_row, column, header)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row_number, row in enumerate(rows, start=start_row + 1):
        for column, value in enumerate(row, start=1):
            cell = ws.cell(row_number, column, value)
            cell.font = Font(name="Aptos", size=9, color=TEXT)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_number].height = 42
    last_row = start_row + len(rows)
    table = Table(displayName=name, ref=f"A{start_row}:{ws.cell(last_row, len(headers)).coordinate}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)
    ws.freeze_panes = f"A{start_row + 1}"
    return last_row


def build_woo_file(existing_rows, headers, names, attribute_mapping):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Worksheet"

    core_headers = ["id", "Title", "Content", "SKU", "Product Type", "Parent Product ID"]
    core_indices = [headers.index(header) for header in core_headers]
    woo_name_overrides = {
        "Dane producenta": "Dane Producenta",
        "Typ produktu": "Typ Produktu",
    }
    row_attributes = []
    attribute_names = set()
    for row in existing_rows:
        attributes = build_features(row, headers, names, attribute_mapping)
        attributes = {woo_name_overrides.get(key, key): value for key, value in attributes.items()}
        row_attributes.append(attributes)
        attribute_names.update(attributes)

    priority_order = [
        "Producent", "Dane Producenta", "Typ Produktu", "Symbol", "Seria", "Rodzaj", "Typ",
        "Charakterystyka Wyzwalania", "Typ Wyłącznika", "Układ biegunów", "Liczba Faz",
        "Prąd Znamionowy [A]", "Napięcie [V]", "Liczba Modułów",
    ]
    ordered_attributes = [name for name in priority_order if name in attribute_names]
    ordered_attributes.extend(sorted(attribute_names - set(ordered_attributes), key=str.casefold))
    output_headers = core_headers + [f"Atrybut Produktu: {name}" for name in ordered_attributes] + ["Kategoria po decyzjach"]

    for column, header in enumerate(output_headers, start=1):
        cell = worksheet.cell(1, column, header)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[1].height = 42

    for row_number, (row, attributes) in enumerate(zip(existing_rows, row_attributes), start=2):
        values = [row[index] for index in core_indices]
        values.extend(attributes.get(name, "") for name in ordered_attributes)
        values.append(names[row[headers.index("SKU")]]["target_category"])
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row_number, column, value)
            cell.font = Font(name="Aptos", size=9, color=TEXT)
            cell.alignment = Alignment(wrap_text=column in (2, 3, len(values)), vertical="top")
        worksheet.row_dimensions[row_number].height = 36

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{worksheet.cell(worksheet.max_row, worksheet.max_column).coordinate}"
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 76
    worksheet.column_dimensions["C"].width = 50
    worksheet.column_dimensions["D"].width = 18
    worksheet.column_dimensions["E"].width = 14
    worksheet.column_dimensions["F"].width = 18
    for column in range(7, worksheet.max_column):
        worksheet.column_dimensions[worksheet.cell(1, column).column_letter].width = 27
    worksheet.column_dimensions[worksheet.cell(1, worksheet.max_column).column_letter].width = 58

    control = workbook.create_sheet("Kontrola importu")
    add_control_title(
        control,
        "Hager — produkty istniejące w WooCommerce",
        "Plik zawiera wyłącznie produkty z niepustym ID Woo. Arkusz Worksheet jest przeznaczony do aktualizacji nazw, atrybutów i kategorii.",
        7,
    )
    control_rows = []
    for row in existing_rows:
        sku = row[headers.index("SKU")]
        control_rows.append([
            row[headers.index("id")], sku, row[headers.index("Title")], names[sku]["target_category"],
            names[sku]["product_type"], names[sku]["status"], "WOO — AKTUALIZACJA",
        ])
    add_table(
        control, 5,
        ["ID Woo", "SKU", "Nazwa końcowa", "Kategoria końcowa", "Typ produktu", "Status danych", "Przeznaczenie"],
        control_rows, "tblHagerWooControl",
    )
    widths = {"A": 13, "B": 18, "C": 72, "D": 62, "E": 42, "F": 18, "G": 24}
    for column, width in widths.items():
        control.column_dimensions[column].width = width
    control.sheet_view.zoomScale = 85
    workbook.save(WOO_OUTPUT)


def build_baselinker_files(new_rows, headers, names, supplier, attribute_mapping):
    csv_rows = []
    control_rows = []
    for row in new_rows:
        sku = row[headers.index("SKU")]
        source = supplier[sku]
        features = build_features(row, headers, names, attribute_mapping)
        features_json = json.dumps(features, ensure_ascii=False, separators=(",", ":"))
        csv_rows.append({
            "product_id": source["product_id"],
            "name": names[sku]["name"],
            "sku": sku,
            "manufacturer_name": "Hager",
            "category": names[sku]["target_category"],
            "features": features_json,
        })
        control_rows.append([
            source["product_id"], sku, names[sku]["name"], names[sku]["target_category"],
            source["price_net"], len(features), features_json, names[sku]["source_url"],
            names[sku]["status"], "BASELINKER — AKTUALIZACJA",
        ])

    with BASELINKER_OUTPUT.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["product_id", "name", "sku", "manufacturer_name", "category", "features"],
            delimiter=";",
            quoting=csv.QUOTE_MINIMAL,
        )
        writer.writeheader()
        writer.writerows(csv_rows)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Kontrola Baselinker"
    add_control_title(
        worksheet,
        "Hager — produkty poza WooCommerce",
        "57 produktów ma ID Baselinker z pliku Hager aparatura na www, ale nie ma ID Woo. CSV aktualizuje tylko nazwę, producenta, kategorię i cechy — nie czyści EAN, opisów ani zdjęć.",
        10,
    )
    add_table(
        worksheet, 5,
        ["ID Baselinker", "SKU", "Nazwa końcowa", "Kategoria końcowa", "Cena netto źródłowa", "Liczba cech", "Features JSON", "Źródło", "Status danych", "Przeznaczenie"],
        control_rows, "tblHagerBaselinkerControl",
    )
    widths = {
        "A": 16, "B": 18, "C": 72, "D": 62, "E": 20,
        "F": 14, "G": 100, "H": 65, "I": 18, "J": 26,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    for row_number in range(6, worksheet.max_row + 1):
        worksheet.cell(row_number, 5).number_format = "0.00"
    worksheet.sheet_view.zoomScale = 80

    notes = workbook.create_sheet("Uwagi do importu")
    add_control_title(notes, "Instrukcja importu Baselinker", "Zakres kolumn został celowo ograniczony, aby nie nadpisywać brakującymi danymi pól już istniejących w Baselinkerze.", 4)
    note_rows = [
        ["Dopasowanie", "product_id + sku", "Wszystkie 57 wierszy ma oba identyfikatory."],
        ["Aktualizowane", "name, manufacturer_name, category, features", "Dane końcowe po decyzjach Hager."],
        ["Nieaktualizowane", "ean, description, images_urls", "Brak tych danych w przekazanych plikach; kolumny pominięto, aby ich nie wyczyścić."],
        ["Cena", "nie jest importowana", "Cena netto ze źródła jest tylko w arkuszu kontrolnym."],
    ]
    add_table(notes, 5, ["Zakres", "Pole", "Wyjaśnienie"], note_rows, "tblHagerBaselinkerNotes")
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 50
    notes.column_dimensions["C"].width = 90
    workbook.save(CONTROL_OUTPUT)
    return csv_rows


def validate(existing_rows, new_rows, headers, names, supplier, csv_rows):
    sku_index = headers.index("SKU")
    id_index = headers.index("id")
    existing_skus = [row[sku_index] for row in existing_rows]
    new_skus = [row[sku_index] for row in new_rows]
    errors = []

    if len(existing_rows) != 96:
        errors.append(f"Oczekiwano 96 produktów Woo, jest {len(existing_rows)}")
    if len(new_rows) != 57:
        errors.append(f"Oczekiwano 57 produktów Baselinker, jest {len(new_rows)}")
    if any(row[id_index] in (None, "") for row in existing_rows):
        errors.append("Plik Woo zawiera produkt bez ID Woo")
    if any(row[id_index] not in (None, "") for row in new_rows):
        errors.append("Grupa Baselinker zawiera produkt z ID Woo")
    if set(existing_skus) & set(new_skus):
        errors.append("SKU występuje w obu grupach")
    if len(set(existing_skus + new_skus)) != len(existing_skus) + len(new_skus):
        errors.append("Wykryto duplikaty SKU")
    if any(sku not in supplier or supplier[sku]["product_id"] in (None, "") for sku in new_skus):
        errors.append("Produkt Baselinker nie ma źródłowego product_id")
    if any(not names[sku]["name"] or not names[sku]["target_category"] for sku in existing_skus + new_skus):
        errors.append("Brak nazwy lub kategorii końcowej")
    if any(sku.split("/")[0] not in names[sku]["name"] for sku in existing_skus + new_skus):
        errors.append("Nazwa końcowa nie zawiera kodu produktu")
    for item in csv_rows:
        try:
            features = json.loads(item["features"])
        except json.JSONDecodeError:
            errors.append(f"Nieprawidłowy JSON cech: {item['sku']}")
            continue
        if not features:
            errors.append(f"Brak cech: {item['sku']}")

    return {
        "source_products": len(existing_rows) + len(new_rows),
        "woo_existing": len(existing_rows),
        "baselinker_outside_woo": len(new_rows),
        "baselinker_with_product_id": sum(bool(supplier[sku]["product_id"]) for sku in new_skus),
        "missing_ean_in_supplied_sources": len(new_rows),
        "missing_images_in_supplied_sources": len(new_rows),
        "baselinker_columns_intentionally_omitted": ["ean", "description", "images_urls"],
        "duplicate_skus": len(existing_skus) + len(new_skus) - len(set(existing_skus + new_skus)),
        "errors": errors,
        "status": "OK" if not errors else "ERROR",
    }


def build():
    headers, rows, index, names = load_final_data()
    supplier = load_supplier_data()
    config = load_yaml(IMPLEMENTATION_CONFIG)
    attribute_mapping = build_attribute_mapping(config)

    existing_rows = [row for row in rows if row[index["id"]] not in (None, "")]
    new_rows = [row for row in rows if row[index["id"]] in (None, "")]

    build_woo_file(existing_rows, headers, names, attribute_mapping)
    csv_rows = build_baselinker_files(new_rows, headers, names, supplier, attribute_mapping)
    report = validate(existing_rows, new_rows, headers, names, supplier, csv_rows)
    REPORT_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    REPORT_OUTPUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if report["errors"]:
        raise RuntimeError("; ".join(report["errors"]))
    print(json.dumps({
        "woo": str(WOO_OUTPUT),
        "baselinker": str(BASELINKER_OUTPUT),
        "control": str(CONTROL_OUTPUT),
        "report": report,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
