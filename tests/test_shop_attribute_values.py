from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from audit_shop_attribute_values import main as audit_main  # noqa: E402
from shop_attribute_values import (  # noqa: E402
    KIND_APPROVED,
    KIND_NEW,
    KIND_NEW_TECHNICAL,
    KIND_SAME_SLUG,
    KIND_SIBLING,
    KIND_SPLIT,
    KIND_TYPE_LINK,
    ShopVocabulary,
    check_features,
    load_rules,
    load_vocabulary,
    slugify,
)


def vocabulary() -> ShopVocabulary:
    vocab = ShopVocabulary()
    for attribute, terms in {
        "Typ produktu": ["Gniazda", "Łączniki", "Inne"],
        "Seria": ["Simon 55", "Simon 54"],
        "Kolor": ["Kaszmir", "Biały"],
        "Napięcie [V]": ["230", "250"],
        "Symbol": ["TW1.01/11"],
        "Typ zacisków": ["Gwintowe", "Bezgwintowe"],
    }.items():
        for term in terms:
            vocab.add(attribute, term)
    for term, parent in {"Antenowe": "Gniazda", "Komputerowe": "Gniazda", "Pojedyncze": "Łączniki"}.items():
        vocab.add("Podtyp produktu", term, parent=parent)
    return vocab


class SlugTest(unittest.TestCase):
    def test_slug_like_wordpress(self) -> None:
        self.assertEqual(slugify("Łączniki"), "laczniki")
        self.assertEqual(slugify("IP 20"), "ip-20")
        self.assertEqual(slugify("Antenowe, Komputerowe"), "antenowe-komputerowe")

    def test_same_slug_is_existing_term(self) -> None:
        vocab = vocabulary()
        self.assertEqual(vocab.canonical("Kolor", "kaszmir"), "Kaszmir")
        self.assertIsNone(vocab.canonical("Kolor", "Kaszmirowy"))


class CheckFeaturesTest(unittest.TestCase):
    def setUp(self) -> None:
        self.vocab = vocabulary()
        self.rules = {"technical_attribute_patterns": [r"\[[^\]]+\]$"], "identifier_attributes": ["Symbol"],
                      "approved_new_values": {"Seria": ["Simon 55 Line"]}}

    def check(self, features: dict[str, str], **kwargs):
        return check_features("X/142/KON", "tytul", features, self.vocab, self.rules, **kwargs)

    def test_comma_value_is_split_into_existing_terms(self) -> None:
        fixed, checks = self.check({"Typ produktu": "Gniazda", "Podtyp produktu": "Antenowe, Komputerowe"})
        self.assertEqual(fixed["Podtyp produktu"], "Antenowe|Komputerowe")
        self.assertEqual(checks[0].kind, KIND_SPLIT)

    def test_case_variant_uses_shop_spelling(self) -> None:
        fixed, checks = self.check({"Kolor": "kaszmir"})
        self.assertEqual(fixed["Kolor"], "Kaszmir")
        self.assertEqual(checks[0].kind, KIND_SAME_SLUG)

    def test_new_value_removed_and_reported(self) -> None:
        fixed, checks = self.check({"Kolor": "Kaszmirowy"})
        self.assertNotIn("Kolor", fixed)
        self.assertEqual((checks[0].kind, checks[0].status), (KIND_NEW, "DO_SPRAWDZENIA"))
        fixed, _ = self.check({"Kolor": "Kaszmirowy"}, keep_new=True)
        self.assertEqual(fixed["Kolor"], "Kaszmirowy")

    def test_sibling_value_replaces_new_value(self) -> None:
        fixed, checks = self.check({"Seria": "Akcesoria do wielu serii"}, siblings=[{"Seria": "Simon 55"}] * 3)
        self.assertEqual(fixed["Seria"], "Simon 55")
        self.assertEqual(checks[0].kind, KIND_SIBLING)

    def test_approved_value_stays(self) -> None:
        fixed, checks = self.check({"Seria": "Simon 55 Line"})
        self.assertEqual(fixed["Seria"], "Simon 55 Line")
        self.assertEqual((checks[0].kind, checks[0].status), (KIND_APPROVED, "OK"))

    def test_technical_new_value_stays_with_warning(self) -> None:
        fixed, checks = self.check({"Napięcie [V]": "14"})
        self.assertEqual(fixed["Napięcie [V]"], "14")
        self.assertEqual((checks[0].kind, checks[0].status), (KIND_NEW_TECHNICAL, "WARNING"))

    def test_identifier_is_not_checked(self) -> None:
        fixed, checks = self.check({"Symbol": "TW1.01/142"})
        self.assertEqual(fixed["Symbol"], "TW1.01/142")
        self.assertEqual(checks, [])

    def test_subtype_of_other_type_is_flagged(self) -> None:
        _, checks = self.check({"Typ produktu": "Inne", "Podtyp produktu": "Pojedyncze"})
        self.assertEqual([check.kind for check in checks], [KIND_TYPE_LINK])
        _, checks = self.check({"Typ produktu": "Łączniki", "Podtyp produktu": "Pojedyncze"})
        self.assertEqual(checks, [])

    def test_attributes_filter_limits_check(self) -> None:
        fixed, checks = self.check({"Kolor": "Kaszmirowy", "Seria": "Simon 55"}, attributes=["Seria"])
        self.assertEqual(fixed["Kolor"], "Kaszmirowy")
        self.assertEqual(checks, [])


class SnapshotAndRulesTest(unittest.TestCase):
    """Utrwala slownik sklepu i mapowania z partii Simon kaszmir (2026-10-12)."""

    def test_snapshot_has_shop_terms_and_filter_links(self) -> None:
        vocab = load_vocabulary()
        self.assertEqual(vocab.canonical("Kolor producenta", "Kaszmir"), "Kaszmir")
        self.assertEqual(vocab.canonical("Typ zacisków", "Bezgwintowe - IDC"), "Bezgwintowe - IDC")
        self.assertEqual(vocab.parent("Podtyp produktu", "Antenowe"), "Gniazda")
        self.assertIsNone(vocab.canonical("Podtyp produktu", "Antenowe, Komputerowe"))

    def test_kaszmir_supplier_values_map_to_shop_terms(self) -> None:
        vocab, rules = load_vocabulary(), load_rules()
        features = {
            "Kolor producenta": "Kaszmirowy",
            "Typ zacisków": "zacisk przebijający izolację",
            "Materiał": "tworzywo sztuczne, metal, bezhalogenowe",
            "Podtyp produktu": "Wtyczkowe, USB",
            "Typ produktu": "Gniazda",
        }
        fixed, checks = check_features("TGZ1ZUSB2CA.01/142/KON", "", features, vocab, rules)
        self.assertEqual(fixed["Kolor producenta"], "Kaszmir")
        self.assertEqual(fixed["Typ zacisków"], "Bezgwintowe - IDC")
        self.assertEqual(fixed["Materiał"], "Tworzywo sztuczne|Metal")
        self.assertEqual(fixed["Podtyp produktu"], "Wtyczkowe|USB")
        self.assertFalse([check for check in checks if check.status == "DO_SPRAWDZENIA"])


class AuditCliTest(unittest.TestCase):
    def test_fixes_copy_of_woo_export_and_lists_orphaned_terms(self) -> None:
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "partia.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["id", "Title", "Sku", "Atrybut produktu: Typ produktu", "Atrybut produktu: Podtyp produktu",
                          "Atrybut produktu: Typ zacisków", "Atrybut produktu: Klasa energetyczna"])
            sheet.append([1, "Gniazdo F + RJ45", "TASFRJ45.01/142/KON", "Gniazda", "Antenowe, Komputerowe", "wtyk", "nie podlega"])
            sheet.append([2, "Łącznik", "TW1.01/142/KON", "Łączniki", "Pojedyncze", "szybkozłącza", None])
            workbook.save(source)
            output = tmp_path / "poprawione.xlsx"
            report = tmp_path / "raport.json"
            code = audit_main(["--input", str(source), "--output", str(output),
                               "--control", str(tmp_path / "kontrola.xlsx"), "--report", str(report)])
            self.assertEqual(code, 0)
            rows = list(load_workbook(output).worksheets[0].iter_rows(values_only=True))
            self.assertEqual(rows[1][4], "Antenowe|Komputerowe")
            self.assertIsNone(rows[1][5])
            self.assertIsNone(rows[1][6])
            self.assertEqual(rows[2][5], "Bezgwintowe")
            data = json.loads(report.read_text(encoding="utf-8"))
            orphans = {(item["atrybut"], item["termin"]) for item in data["terminy_do_usuniecia"]}
            self.assertIn(("Podtyp produktu", "Antenowe, Komputerowe"), orphans)
            self.assertIn(("Typ zacisków", "szybkozłącza"), orphans)
            sheets = load_workbook(tmp_path / "kontrola.xlsx").sheetnames
            self.assertIn("Nowe wartosci", sheets)
            self.assertIn("Terminy do usuniecia w Woo", sheets)


if __name__ == "__main__":
    unittest.main()
