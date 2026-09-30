from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook, load_workbook  # noqa: E402

from build_hager_berker_compatible_products import (  # noqa: E402
    META_KEY,
    STATUS_OK,
    STATUS_REVIEW,
    STATUS_WARNING,
    Relation,
    build_code_index,
    fill_workbook,
    load_rules,
    minimal_record,
    parse_existing_ids,
    select_compatible,
    series_fallback,
    validate,
    with_reverse,
)
from build_hager_berker_cross_upsell import CatalogProduct  # noqa: E402
from build_hager_berker_variants import Record  # noqa: E402


def product(
    sku: str,
    title: str,
    product_type: str,
    series: str = "B.3",
    color: str = "Biały",
    multiplicity: str = "",
    alias_code: str = "",
    tags: str = "",
    old_duplicate: bool = False,
) -> CatalogProduct:
    record = Record(
        excel_row=2,
        product_id=str(abs(hash(sku)) % 100000),
        title=title,
        sku=sku,
        code=sku.split("/", 1)[0],
        product_type=product_type,
        series=series,
        primary_series=series.split("|", 1)[0],
        color=color,
        manufacturer_color_before=color,
        manufacturer_color=color,
        material="",
        multiplicity=multiplicity,
        tags_before=tags,
        core=title,
        family="inne",
        frame=product_type == "Ramki",
        color_detail=color,
        material_detail="",
        finish_detail="",
        orientation="",
        network_category="",
        screened="",
        grounding="",
        terminal_type="",
        voltage="",
        nominal_current="",
        pole_count="",
        shutters="",
        detection_range="",
        output_count="",
        alias_code=alias_code,
    )
    record.series_key = series.split("|", 1)[0].lower()
    record.is_old_duplicate = old_duplicate
    return CatalogProduct(record)


def accessory(source: str, target: str) -> Relation:
    return Relation(source, target, "accessories", False, "test")


class CompatibleProductsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.rules = load_rules()

    def select(self, products, relations, upsells=None):
        return select_compatible(products, with_reverse(relations), self.rules, upsells)

    def test_relation_works_in_both_directions(self) -> None:
        frame = product("F1/HAG", "Ramka 1-krotna biała", "Ramki", multiplicity="1x")
        key = product("K1/HAG", "Klawisz pojedynczy biały", "Klawisze")
        selected = self.select([frame, key], [accessory("F1/HAG", "K1/HAG")])
        self.assertEqual(["K1/HAG"], [item.target_sku for item in selected["F1/HAG"]])
        self.assertEqual(["F1/HAG"], [item.target_sku for item in selected["K1/HAG"]])
        self.assertIn("relacja odwrotna", selected["K1/HAG"][0].source_label)

    def test_matching_color_wins_and_other_color_is_dropped(self) -> None:
        frame = product("F1/HAG", "Ramka 1-krotna biała", "Ramki", multiplicity="1x")
        white = product("K1/HAG", "Klawisz pojedynczy biały", "Klawisze")
        black = product("K2/HAG", "Klawisz pojedynczy czarny", "Klawisze", color="Czarny")
        selected = self.select(
            [frame, white, black], [accessory("F1/HAG", "K1/HAG"), accessory("F1/HAG", "K2/HAG")],
        )
        self.assertEqual(["K1/HAG"], [item.target_sku for item in selected["F1/HAG"]])
        self.assertEqual(STATUS_OK, selected["F1/HAG"][0].status)

    def test_color_without_own_frame_falls_back_with_warning(self) -> None:
        socket = product("S1/HAG", "Gniazdo SCHUKO pomarańczowe", "Gniazda", color="Pomarańczowy")
        frame = product("F1/HAG", "Ramka 1-krotna biała", "Ramki", multiplicity="1x")
        selected = self.select([socket, frame], [accessory("F1/HAG", "S1/HAG")])
        self.assertEqual("F1/HAG", selected["S1/HAG"][0].target_sku)
        self.assertEqual(STATUS_WARNING, selected["S1/HAG"][0].status)

    def test_mechanism_is_not_filtered_by_color(self) -> None:
        key = product("K1/HAG", "Klawisz pojedynczy czarny", "Klawisze", color="Czarny")
        mechanism = product("M1/HAG", "Łącznik pojedynczy mechanizm", "Mechanizmy", series="one.platform", color="")
        selected = self.select([key, mechanism], [Relation("K1/HAG", "M1/HAG", "mandatory", False, "test")])
        self.assertEqual(STATUS_OK, selected["K1/HAG"][0].status)
        self.assertEqual("BMEcat: element obowiązkowy", selected["K1/HAG"][0].source_label)

    def test_plain_product_ranks_before_niche_version(self) -> None:
        frame = product("F1/HAG", "Ramka 1-krotna biała", "Ramki", multiplicity="1x")
        nema = product("S1/HAG", "Gniazdo z uziemieniem USA NEMA białe", "Gniazda")
        schuko = product("S2/HAG", "Gniazdo SCHUKO białe", "Gniazda")
        selected = self.select(
            [frame, nema, schuko], [accessory("F1/HAG", "S1/HAG"), accessory("F1/HAG", "S2/HAG")],
        )
        self.assertEqual(["S2/HAG", "S1/HAG"], [item.target_sku for item in selected["F1/HAG"]])

    def test_limits_and_series_diversity_of_frames(self) -> None:
        key = product("K1/HAG", "Klawisz pojedynczy biały", "Klawisze", series="B.3|B.7")
        frames = [
            product(f"F{series}{index}/HAG", f"Ramka 1-krotna biała {index}", "Ramki", series=series, multiplicity="1x")
            for series in ("B.3", "B.7")
            for index in range(6)
        ]
        selected = self.select([key, *frames], [accessory(frame.sku, "K1/HAG") for frame in frames])
        chosen = selected["K1/HAG"]
        self.assertEqual(self.rules["max_per_kind"] * 2, len(chosen))
        self.assertEqual({"b.3", "b.7"}, {item.target_sku[1:4].lower() for item in chosen})

    def test_variants_upsells_and_old_duplicates_are_excluded(self) -> None:
        frame = product("F1/HAG", "Ramka 1-krotna biała", "Ramki", multiplicity="1x", tags="Wariant ramka")
        sibling = product("F2/HAG", "Ramka 2-krotna biała", "Ramki", multiplicity="2x", tags="Wariant ramka")
        upsell = product("K1/HAG", "Klawisz pojedynczy biały", "Klawisze")
        old = product("K2/HAG", "Klawisz podwójny biały", "Klawisze", old_duplicate=True)
        kept = product("K3/HAG", "Przycisk biały", "Przyciski")
        products = [frame, sibling, upsell, old, kept]
        upsells = {"F1/HAG": {"K1/HAG"}}
        selected = self.select(
            products, [accessory("F1/HAG", sku) for sku in ("F2/HAG", "K1/HAG", "K2/HAG", "K3/HAG")], upsells,
        )
        self.assertEqual(["K3/HAG"], [item.target_sku for item in selected["F1/HAG"]])
        self.assertNotIn("K2/HAG", selected)
        self.assertTrue(all(validate(products, selected, upsells, self.rules["max_items"]).values()))

    def test_code_index_maps_berker_alias_and_skips_old_duplicate(self) -> None:
        new = product("WNA683/HAG", "Adapter natynkowy (Berker 6019303505)", "Inne", alias_code="6019303505")
        old = product("6019303505/HAG", "Adapter natynkowy", "Inne", old_duplicate=True)
        index = build_code_index([new, old])
        self.assertEqual(["WNA683/HAG"], index["6019303505"])
        self.assertEqual(["WNA683/HAG"], index["WNA683"])

    def test_series_fallback_is_review_only_and_skips_surface_mounted(self) -> None:
        dimmer = product("D1/HAG", "Ściemniacz obrotowy biały", "Ściemniacz", series="Lumina")
        surface = product("D2/HAG", "Łącznik natynkowy IP55 biały", "Łączniki", series="Lumina")
        frame = product("F1/HAG", "Ramka 1-krotna biała", "Ramki", series="Lumina", multiplicity="1x")
        other = product("F2/HAG", "Ramka 1-krotna czarna", "Ramki", series="Lumina", color="Czarny", multiplicity="1x")
        proposals = series_fallback([dimmer, surface, frame, other], set(), self.rules)
        self.assertEqual(["F1/HAG"], [item.target_sku for item in proposals["D1/HAG"]])
        self.assertEqual(STATUS_REVIEW, proposals["D1/HAG"][0].status)
        self.assertNotIn("D2/HAG", proposals)


class FillExportTests(unittest.TestCase):
    def test_parse_existing_ids_reads_php_array_and_plain_list(self) -> None:
        self.assertEqual(["20648"], parse_existing_ids("a:1:{i:0;i:20648;}"))
        self.assertEqual(["19592", "19596"], parse_existing_ids('a:2:{i:0;s:5:"19592";i:1;i:19596;}'))
        self.assertEqual(["12", "34"], parse_existing_ids("12, 34"))
        self.assertEqual([], parse_existing_ids(None))

    def test_minimal_record_detects_berker_alias_for_old_duplicates(self) -> None:
        record = minimal_record(2, "10", "Łącznik świecznikowy Hager one.platform WDEM3035 (Berker 533035)", "WDEM3035/HAG")
        self.assertEqual("533035", record.alias_code)
        self.assertEqual("WDEM3035", record.code)

    def test_fill_keeps_other_cells_and_manual_ids_first(self) -> None:
        frame = product("F1/HAG", "Ramka 1-krotna biała", "Ramki", multiplicity="1x")
        key = product("K1/HAG", "Klawisz pojedynczy biały", "Klawisze")
        manual = product("K2/HAG", "Klawisz podwójny biały", "Klawisze")
        frame.record.product_id, key.record.product_id, manual.record.product_id = "1", "2", "3"
        rules = load_rules()
        selected = select_compatible(
            [frame, key, manual], with_reverse([accessory("F1/HAG", "K1/HAG"), accessory("F1/HAG", "K2/HAG")]), rules,
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            source, target = Path(temp_dir) / "export.xlsx", Path(temp_dir) / "filled.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["id", "Title", "Sku", "Cross-Sells", META_KEY])
            sheet.append([1, frame.title, "F1/HAG", "K1/HAG", None])
            sheet.append([2, key.title, "K1/HAG", None, "a:1:{i:0;i:999;}"])
            sheet.append([4, "Produkt spoza katalogu", "X1/HAG", None, None])
            workbook.save(source)

            metrics = fill_workbook(source, target, [frame, key, manual], selected, rules["max_items"])
            rows = [list(row) for row in load_workbook(target).active.iter_rows(values_only=True)]

        self.assertTrue(metrics["fill_other_cells_unchanged"])
        self.assertEqual(["K1/HAG"], metrics["fill_rows_with_manual_value_kept"])
        self.assertEqual([1, frame.title, "F1/HAG", "K1/HAG", "2,3"], rows[1])
        self.assertEqual("999,1", rows[2][4])
        self.assertIsNone(rows[3][4])


if __name__ == "__main__":
    unittest.main()
