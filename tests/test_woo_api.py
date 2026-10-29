from __future__ import annotations

import sys
import tempfile
from datetime import datetime
from io import BytesIO
from pathlib import Path
from unittest import TestCase

from openpyxl import load_workbook
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from shop_attribute_values import load_vocabulary
from woo_api.knowledge import write_knowledge
from woo_api.push import PushOptions, image_problems, make_changeset_id, prepare_rows, run_push
from woo_api.workbook import merge_columns, read_products_xlsx, write_products_xlsx


COLUMNS = ["id", "Title", "SKU", "Width", "Product Attributes", "Atrybut Produktu: Kolor"]


def product_row(product_id: int, sku: str, title: str = "Gniazdo", width: str = "") -> dict[str, object]:
    return {"id": product_id, "Title": title, "SKU": sku, "Width": width, "Product Attributes": "a:0:{}", "Atrybut Produktu: Kolor": "Biały"}


class FakeBridge:
    """Sklep, w którym zmienia się tylko Title innego niż "Gniazdo", a SKU "ZLE" daje błąd."""

    def __init__(self) -> None:
        self.calls: list[dict[str, object]] = []

    def sync(self, rows, *, dry_run=True, allow_new_terms=False, columns=None, changeset=""):
        self.calls.append({"rows": rows, "dry_run": dry_run, "columns": columns, "changeset": changeset})
        products = []
        for row in rows:
            errors = ["Shipping Class: nie ma klasy"] if row.get("SKU") == "ZLE" else []
            changes = [] if row.get("Title") == "Gniazdo" else [{"column": "Title", "old": "Gniazdo", "new": row.get("Title")}]
            status = "error" if errors else ("changed" if changes else "unchanged")
            if status == "changed" and not dry_run:
                status = "applied"
            products.append({"row": row["_row"], "id": int(row["id"]), "sku": row["SKU"], "title": "", "status": status, "changes": changes, "errors": errors, "warnings": []})
        summary = {
            "rows": len(products),
            "changed": sum(p["status"] in ("changed", "applied") for p in products),
            "applied": sum(p["status"] == "applied" for p in products),
            "errors": sum(bool(p["errors"]) for p in products),
            "warnings": 0,
            "changes": sum(len(p["changes"]) for p in products),
        }
        return {"changeset": "" if dry_run else changeset, "summary": summary, "products": products, "new_terms": [], "written_columns": ["Title"], "read_only_columns": ["id", "SKU"], "ignored_columns": []}


def image_bytes(size=(1000, 1000), fmt="JPEG", mode="RGB", color=(255, 255, 255)) -> bytes:
    buffer = BytesIO()
    Image.new(mode, size, color).save(buffer, fmt)
    return buffer.getvalue()


class ImageBridge(FakeBridge):
    """Sklep, w którym każdy produkt ma dziś zdjęcie stare.jpg, a Image URL z pliku je zmienia."""

    def sync(self, rows, *, dry_run=True, allow_new_terms=False, columns=None, changeset=""):
        result = super().sync(rows, dry_run=dry_run, allow_new_terms=allow_new_terms, columns=columns, changeset=changeset)
        for row, product in zip(rows, result["products"]):
            if row.get("Image URL") and row["Image URL"] != "https://sklep/stare.jpg":
                product["changes"].append({"column": "Image URL", "old": "https://sklep/stare.jpg", "new": row["Image URL"]})
                product["status"] = "applied" if not dry_run else "changed"
        result["summary"]["changed"] = sum(p["status"] in ("changed", "applied") for p in result["products"])
        return result


IMAGES = {
    "https://x/dobre.jpg": image_bytes(),
    "https://x/male.png": image_bytes((500, 500), "PNG", "RGBA", (0, 0, 0, 0)),
    "https://x/szare.jpg": image_bytes(color=(200, 200, 200)),
}


class ImageCheckTests(TestCase):
    def test_image_problems(self) -> None:
        self.assertEqual(image_problems(IMAGES["https://x/dobre.jpg"]), ([], []))
        errors, _ = image_problems(IMAGES["https://x/male.png"])
        self.assertEqual(errors, ["format PNG, a ma być JPG", "500x500, a ma być 1000x1000", "tryb kolorów RGBA, a ma być RGB"])
        self.assertEqual(image_problems(IMAGES["https://x/szare.jpg"])[0], [])
        self.assertTrue(image_problems(IMAGES["https://x/szare.jpg"])[1])
        self.assertEqual(image_problems(b"<html>")[0], ["to nie jest plik graficzny"])

    def _push(self, temp_dir: str, urls: dict[int, str], options: PushOptions):
        columns = COLUMNS + ["Image URL"]
        rows = [{**product_row(index, f"S{index}", title="Nowy"), "Image URL": url} for index, url in urls.items()]
        path = Path(temp_dir) / "zdjecia.xlsx"
        write_products_xlsx(path, columns, rows)
        bridge = ImageBridge()
        fetched: list[str] = []

        def fetch(url: str) -> bytes:
            fetched.append(url)
            return IMAGES[url]

        result = run_push(bridge, path, options, report_dir=temp_dir, log=lambda _: None, image_fetcher=fetch)
        return bridge, result, fetched

    def test_bad_image_blocks_product_even_with_skip_errors(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            urls = {1: "https://x/dobre.jpg", 2: "https://sklep/stare.jpg|https://x/male.png"}
            bridge, result, fetched = self._push(temp_dir, urls, PushOptions(apply=True, skip_errors=True))
            # Zdjecie juz przypisane produktowi nie jest pobierane.
            self.assertEqual(sorted(fetched), ["https://x/dobre.jpg", "https://x/male.png"])
            self.assertEqual(result.dry_run["summary"]["errors"], 1)
            writes = [call for call in bridge.calls if not call["dry_run"]]
            self.assertEqual([row["id"] for row in writes[0]["rows"]], ["1"])

    def test_bad_image_blocks_apply_without_skip_errors(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            _, result, _ = self._push(temp_dir, {1: "https://x/male.png"}, PushOptions(apply=True))
            self.assertTrue(result.blocked)
            problems = load_workbook(result.report_paths[0])["Problemy"]
            self.assertIn("500x500", problems.cell(2, 6).value)

    def test_check_can_be_skipped(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            _, result, fetched = self._push(temp_dir, {1: "https://x/male.png"}, PushOptions(check_images=False))
            self.assertEqual(fetched, [])
            self.assertEqual(result.dry_run["summary"]["errors"], 0)


class WorkbookTests(TestCase):
    def test_roundtrip_keeps_wp_all_export_layout(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "produkty.xlsx"
            write_products_xlsx(path, COLUMNS, [product_row(10, "A/SIM", width="8.5"), product_row(11, "B/SIM")])

            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["Worksheet"])
            self.assertEqual(workbook["Worksheet"]["A2"].value, 10)
            self.assertIsNone(workbook["Worksheet"]["D3"].value)

            headers, rows = read_products_xlsx(path)
            self.assertEqual(headers, COLUMNS)
            self.assertEqual(rows[0]["_row"], 2)
            self.assertEqual(rows[0]["id"], "10")
            self.assertEqual(rows[0]["Width"], "8.5")
            self.assertEqual(rows[1]["Width"], "")

    def test_read_decodes_excel_escaped_characters(self) -> None:
        # Excel saves "\r" in cell text as "_x000D_"; reading must give back the real character.
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "produkty.xlsx"
            write_products_xlsx(path, ["id", "Content"], [{"id": 10, "Content": "Akapit._x000D_\n<h2>Nagłówek</h2>"}])

            _, rows = read_products_xlsx(path)
            self.assertEqual(rows[0]["Content"], "Akapit.\r\n<h2>Nagłówek</h2>")

    def test_merge_columns_puts_attributes_last_alphabetically(self) -> None:
        merged = merge_columns([["id", "Title", "Atrybut Produktu: Typ"], ["id", "SKU", "Atrybut Produktu: Kolor"]])
        self.assertEqual(merged, ["id", "Title", "SKU", "Atrybut Produktu: Kolor", "Atrybut Produktu: Typ"])


class PrepareRowsTests(TestCase):
    def test_drops_read_only_columns_and_keeps_identity(self) -> None:
        rows = [{"_row": 2, **{key: str(value) for key, value in product_row(10, "A/SIM").items()}}]
        payload = prepare_rows(COLUMNS, rows, [])
        self.assertNotIn("Product Attributes", payload[0])
        self.assertEqual(payload[0]["SKU"], "A/SIM")

    def test_columns_limit_what_is_sent(self) -> None:
        rows = [{"_row": 2, **{key: str(value) for key, value in product_row(10, "A/SIM").items()}}]
        payload = prepare_rows(COLUMNS, rows, ["Width"])
        self.assertEqual(set(payload[0]), {"_row", "id", "SKU", "Width"})

    def test_rejects_duplicates_and_unknown_columns(self) -> None:
        rows = [{"_row": 2, "id": "10", "SKU": "A"}, {"_row": 3, "id": "10", "SKU": "A"}]
        with self.assertRaisesRegex(ValueError, "wierszu 2"):
            prepare_rows(["id", "SKU"], rows, [])
        with self.assertRaisesRegex(ValueError, "Width"):
            prepare_rows(["id", "SKU"], rows[:1], ["Width"])


class PushTests(TestCase):
    def _file(self, temp_dir: str, rows: list[dict[str, object]]) -> Path:
        path = Path(temp_dir) / "zmiany.xlsx"
        write_products_xlsx(path, COLUMNS, rows)
        return path

    def test_dry_run_never_writes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            bridge = FakeBridge()
            path = self._file(temp_dir, [product_row(10, "A", title="Nowy")])
            result = run_push(bridge, path, PushOptions(), report_dir=Path(temp_dir) / "raporty", log=lambda _: None)
            self.assertEqual([call["dry_run"] for call in bridge.calls], [True])
            self.assertIsNone(result.applied)
            self.assertTrue(result.report_paths[0].is_file())

    def test_apply_is_blocked_by_errors_unless_skipped(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._file(temp_dir, [product_row(10, "A", title="Nowy"), product_row(11, "ZLE")])

            bridge = FakeBridge()
            blocked = run_push(bridge, path, PushOptions(apply=True), report_dir=temp_dir, log=lambda _: None)
            self.assertTrue(blocked.blocked)
            self.assertFalse(any(not call["dry_run"] for call in bridge.calls))

            bridge = FakeBridge()
            result = run_push(
                bridge, path, PushOptions(apply=True, skip_errors=True), report_dir=temp_dir, log=lambda _: None,
                now=datetime(2026, 10, 21, 15, 30),
            )
            writes = [call for call in bridge.calls if not call["dry_run"]]
            self.assertEqual(len(writes), 1)
            self.assertEqual([row["id"] for row in writes[0]["rows"]], ["10"])
            self.assertEqual(result.changeset, "ai-20261021-153000-zmiany")
            self.assertEqual(writes[0]["changeset"], result.changeset)

    def test_rows_are_sent_in_batches(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            bridge = FakeBridge()
            path = self._file(temp_dir, [product_row(index, f"S{index}") for index in range(1, 6)])
            run_push(bridge, path, PushOptions(batch_size=2), report_dir=temp_dir, log=lambda _: None)
            self.assertEqual([len(call["rows"]) for call in bridge.calls], [2, 2, 1])

    def test_changeset_id_is_ascii(self) -> None:
        self.assertEqual(make_changeset_id("output/Łączniki Simon 54.xlsx", datetime(2026, 1, 2, 3, 4, 5)), "ai-20260102-030405-Laczniki-Simon-54")


class KnowledgeTests(TestCase):
    def test_snapshots_keep_the_database_dump_format(self) -> None:
        knowledge = {
            "site_url": "https://alkanlocal.test:8443/",
            "attributes": [
                {"label": "Typ produktu", "taxonomy": "pa_typ-produktu", "terms": [{"name": "Gniazda", "slug": "gniazda", "count": 5, "parent_term": ""}]},
                {"label": "Podtyp produktu", "taxonomy": "pa_podtyp-produktu", "terms": [{"name": "Antenowe", "slug": "antenowe", "count": 2, "parent_term": "Gniazda"}]},
                {"label": "Pusty", "taxonomy": "pa_pusty", "terms": []},
            ],
            "categories": [{"term_id": 2, "parent_id": 1, "name": "Simon 54", "slug": "simon-54", "count": 3}, {"term_id": 1, "parent_id": 0, "name": "Gniazdka", "slug": "gniazdka", "count": 9}],
            "tags": [{"term_id": 7, "name": "Gniazdo RTV", "slug": "gniazdo-rtv", "count": 4}],
            "shipping_classes": [{"term_id": 8, "name": "Standard", "slug": "standard", "count": 100}],
            "shipping_within": [],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            paths = write_knowledge(knowledge, Path(temp_dir), Path(temp_dir) / "cache" / "knowledge.json")

            vocabulary = load_vocabulary(paths["attributes"])
            self.assertTrue(vocabulary)
            lines = paths["attributes"].read_text(encoding="utf-8").splitlines()
            self.assertTrue(lines[0].startswith("# "))
            self.assertEqual(lines[2], "attribute\ttaxonomy\tterm\tslug\tcount\tparent_term")
            self.assertIn("Podtyp produktu\tpa_podtyp-produktu\tAntenowe\tantenowe\t2\tGniazda", lines)
            self.assertIn("Pusty\tpa_pusty\t\t\t0\t", lines)

            categories = paths["categories"].read_text(encoding="utf-8").splitlines()
            self.assertEqual(categories[3:], ["1\t0\tGniazdka\tgniazdka\t9", "2\t1\tSimon 54\tsimon-54\t3"])
            self.assertIn("product_shipping_class\t8\tStandard\tstandard\t100", paths["shop_terms"].read_text(encoding="utf-8"))
            self.assertTrue(paths["full"].is_file())
