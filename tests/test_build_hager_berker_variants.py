from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from build_hager_berker_variants import (  # noqa: E402
    Record,
    assign_variant_groups,
    mojibake_markers,
)


def frame_record(
    code: str,
    multiplicity: str,
    orientation: str,
    manufacturer_color: str = "Biały aksamit",
) -> Record:
    return Record(
        excel_row=2,
        product_id=code,
        title=f"Ramka {multiplicity} z polem opisowym; biały, aksamit; Q.1 {code}",
        sku=f"{code}/HAG",
        code=code,
        product_type="Ramka",
        series="Q.1",
        primary_series="Q.1",
        color="Biały",
        manufacturer_color_before=manufacturer_color,
        manufacturer_color=manufacturer_color,
        material="",
        multiplicity=multiplicity,
        tags_before="",
        core="ramka z polem opisowym",
        family="ramki",
        frame=True,
        color_detail="biały aksamit",
        material_detail="",
        finish_detail="aksamit",
        orientation=orientation,
        network_category="",
        screened="",
        grounding="",
        terminal_type="",
        voltage="",
        nominal_current="",
        pole_count="",
        shutters="",
        detection_range="",
        output_count="",
        alias_code="",
        series_key="q system",
    )


class HagerVariantPluginCompatibilityTests(unittest.TestCase):
    def test_two_products_differing_on_two_axes_are_not_linked(self) -> None:
        records = [
            frame_record("10116019", "1x", "Bez orientacji"),
            frame_record("10126019", "2x", "Pionowa"),
        ]

        accepted, rejected = assign_variant_groups(records)

        self.assertEqual(accepted, [])
        self.assertEqual(len(rejected), 1)
        self.assertTrue(all(record.variant_tag == "" for record in records))
        self.assertTrue(all(record.decision == "DO WERYFIKACJI" for record in records))
        self.assertIn("więcej niż jednej osi", records[0].reason)

    def test_two_products_differing_on_one_axis_are_linked(self) -> None:
        records = [
            frame_record("10116019", "1x", "Bez orientacji", "Biały aksamit"),
            frame_record("10116086", "1x", "Bez orientacji", "Antracyt aksamit"),
        ]

        accepted, rejected = assign_variant_groups(records)

        self.assertEqual(len(accepted), 1)
        self.assertEqual(rejected, [])
        self.assertEqual(accepted[0]["axes"], ["Kolor producenta"])
        self.assertTrue(all(record.variant_tag for record in records))

    def test_isolated_one_fold_frame_is_excluded_but_switchable_grid_remains(self) -> None:
        isolated = frame_record("10116019", "1x", "Bez orientacji")
        switchable = [
            frame_record("10126019", "2x", "Pionowa"),
            frame_record("10226019", "2x", "Pozioma"),
            frame_record("10136019", "3x", "Pionowa"),
            frame_record("10236019", "3x", "Pozioma"),
        ]

        accepted, rejected = assign_variant_groups([isolated, *switchable])

        self.assertEqual(len(accepted), 1)
        self.assertEqual(rejected, [])
        self.assertEqual({record.code for record in accepted[0]["records"]}, {record.code for record in switchable})
        self.assertEqual(accepted[0]["axes"], ["Krotność", "Orientacja"])
        self.assertEqual(isolated.variant_tag, "")
        self.assertEqual(isolated.decision, "DO WERYFIKACJI")
        self.assertIn("nieprzełączalny", isolated.reason)
        self.assertTrue(all(record.variant_tag for record in switchable))


class PolishEncodingTests(unittest.TestCase):
    def test_mojibake_detector_accepts_polish_and_rejects_corruption(self) -> None:
        correct = "Łącznik schodowy — śrubowe zaciski; żółty i biały"
        corrupted = "ÅÄ…cznik schodowy â€” Å›rubowe zaciski"

        self.assertEqual(mojibake_markers(correct), [])
        self.assertTrue(mojibake_markers(corrupted))

    def test_generator_source_is_valid_utf8_without_mojibake(self) -> None:
        source = (ROOT / "src" / "build_hager_berker_variants.py").read_text(encoding="utf-8")
        markers_start = source.index("MOJIBAKE_MARKERS = (")
        markers_end = source.index("\n)\n", markers_start) + len("\n)\n")
        source_without_marker_definitions = source[:markers_start] + source[markers_end:]

        self.assertIn('record.terminal_type = "Śrubowe"', source)
        self.assertEqual(mojibake_markers(source_without_marker_definitions), [])

    def test_polish_text_round_trips_through_xlsx(self) -> None:
        expected = "Wariant - Łącznik schodowy — śrubowe zaciski; żółty"
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "polskie_znaki.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = expected
            workbook.save(path)
            workbook.close()

            saved = load_workbook(path, read_only=True)
            try:
                self.assertEqual(saved.active["A1"].value, expected)
            finally:
                saved.close()


if __name__ == "__main__":
    unittest.main()
