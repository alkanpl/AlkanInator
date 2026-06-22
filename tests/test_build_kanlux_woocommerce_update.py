from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from build_kanlux_woocommerce_update import (  # noqa: E402
    existing_features_by_sku,
    normalize_sku,
    serialize_product_attributes,
    update_workbook,
)
from openpyxl import Workbook, load_workbook  # noqa: E402


class KanluxWooCommerceUpdateTests(unittest.TestCase):
    def test_normalizes_kanlux_sku_suffix(self) -> None:
        self.assertEqual(normalize_sku("29002/KAN"), "29002")
        self.assertEqual(normalize_sku("29002/KANLUX"), "29002")

    def test_serializes_product_attributes_in_input_format(self) -> None:
        result = serialize_product_attributes(["Producent", "Stopień ochrony [IP]"])

        self.assertTrue(result.startswith("a:2:{"))
        self.assertIn('s:12:"pa_producent";', result)
        self.assertIn('s:21:"pa_stopien-ochrony-ip";', result)
        self.assertIn('s:8:"position";i:0;', result)
        self.assertIn('s:8:"position";i:1;', result)
        self.assertTrue(result.endswith("}"))

    def test_existing_dimensions_are_normalized_to_millimeters(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["SKU", "Atrybut Produktu: Długość", "Atrybut Produktu: Wysokość"])
        worksheet.append(["TEST/KAN", "1.2 m", "8.5 cm"])

        result = existing_features_by_sku(worksheet)

        self.assertEqual(result["TEST"]["Długość"], "1200 mm")
        self.assertEqual(result["TEST"]["Wysokość"], "85 mm")


    def test_updates_title_when_verified_title_matches_feature_product_type(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_path = root / "woo.xlsx"
            features_path = root / "features.csv"
            titles_path = root / "titles.csv"
            output_path = root / "out.xlsx"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["SKU", "Title", "Atrybut Produktu: Typ Produktu", "Product Attributes"])
            worksheet.append(["39232/KAN", "Zasilacz BLINGO AIO 40W Kanlux 39232", "", ""])
            workbook.save(input_path)

            features_path.write_text(
                'product_id;sku;features\n'
                '158300;39232/KAN;"{""Typ produktu"":""Panel LED""}"\n',
                encoding="utf-8-sig",
            )
            titles_path.write_text(
                "product_id;name;sku\n"
                "158300;Panel LED natynkowy BLINGO AIO 40W Kanlux 39232;39232/KAN\n",
                encoding="utf-8-sig",
            )

            summary = update_workbook(input_path, features_path, output_path, titles_path)
            output_workbook = load_workbook(output_path, read_only=True)
            output_sheet = output_workbook.active
            output_title = output_sheet.cell(2, 2).value
            output_features = existing_features_by_sku(output_sheet)
            output_workbook.close()

            self.assertEqual(summary["retitled_rows"], 1)
            self.assertEqual(
                output_title,
                "Panel LED natynkowy BLINGO AIO 40W Kanlux 39232",
            )
            self.assertEqual(output_features["39232"]["Typ produktu"], "Panel LED")

    def test_updates_title_when_existing_title_has_stale_ip(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_path = root / "woo.xlsx"
            features_path = root / "features.csv"
            titles_path = root / "titles.csv"
            output_path = root / "out.xlsx"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["SKU", "Title", "Atrybut Produktu: Typ Produktu", "Atrybut Produktu: Stopien ochrony [IP]", "Product Attributes"])
            worksheet.append(["39240/KAN", "Panel LED podtynkowy BLINGO U 4000K neutralne IP12 Kanlux 39240", "", "", ""])
            workbook.save(input_path)

            features_path.write_text(
                'product_id;sku;features\n'
                '158142;39240/KAN;"{""Typ produktu"":""Panel LED"",""Stopien ochrony [IP]"":""IP 54|IP 20""}"\n',
                encoding="utf-8-sig",
            )
            titles_path.write_text(
                "product_id;name;sku\n"
                "158142;Panel LED podtynkowy BLINGO U 4000K neutralne IP54 / IP20 Kanlux 39240;39240/KAN\n",
                encoding="utf-8-sig",
            )

            summary = update_workbook(input_path, features_path, output_path, titles_path)
            output_workbook = load_workbook(output_path, read_only=True)
            output_sheet = output_workbook.active
            output_title = output_sheet.cell(2, 2).value
            output_workbook.close()

            self.assertEqual(summary["retitled_rows"], 1)
            self.assertEqual(
                output_title,
                "Panel LED podtynkowy BLINGO U 4000K neutralne IP54 / IP20 Kanlux 39240",
            )

    def test_updates_title_when_existing_title_has_stale_mounting(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_path = root / "woo.xlsx"
            features_path = root / "features.csv"
            titles_path = root / "titles.csv"
            output_path = root / "out.xlsx"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["SKU", "Title", "Atrybut Produktu: Typ Produktu", "Atrybut Produktu: Sposob montazu", "Product Attributes"])
            worksheet.append(["39172/KAN", "Panel LED podtynkowy BLINGO U 34W 4000K neutralne Kanlux 39172", "", "", ""])
            workbook.save(input_path)

            features_path.write_text(
                'product_id;sku;features\n'
                '158147;39172/KAN;"{""Typ produktu"":""Panel LED"",""Sposob montazu"":""Natynkowy""}"\n',
                encoding="utf-8-sig",
            )
            titles_path.write_text(
                "product_id;name;sku\n"
                "158147;Panel LED natynkowy BLINGO U 34W 4000K neutralne Kanlux 39172;39172/KAN\n",
                encoding="utf-8-sig",
            )

            summary = update_workbook(input_path, features_path, output_path, titles_path)
            output_workbook = load_workbook(output_path, read_only=True)
            output_sheet = output_workbook.active
            output_title = output_sheet.cell(2, 2).value
            output_workbook.close()

            self.assertEqual(summary["retitled_rows"], 1)
            self.assertEqual(
                output_title,
                "Panel LED natynkowy BLINGO U 34W 4000K neutralne Kanlux 39172",
            )

    def test_updates_title_when_existing_title_misses_title_bearing_attribute(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_path = root / "woo.xlsx"
            features_path = root / "features.csv"
            titles_path = root / "titles.csv"
            output_path = root / "out.xlsx"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["SKU", "Title", "Atrybut Produktu: Typ Produktu", "Atrybut Produktu: Wymiary [mm]", "Product Attributes"])
            worksheet.append(["70523/KAN", "Oprawa kanałowa MILO E27 IP54 szara Kanlux 70523", "", "", ""])
            workbook.save(input_path)

            features_path.write_text(
                'product_id;sku;features\n'
                '70523;70523/KAN;"{""Typ produktu"":""Oprawa kanałowa"",""Wymiary [mm]"":""165x100""}"\n',
                encoding="utf-8-sig",
            )
            titles_path.write_text(
                "product_id;name;sku\n"
                "70523;Oprawa kanałowa MILO E27 165x100mm IP54 szara Kanlux 70523;70523/KAN\n",
                encoding="utf-8-sig",
            )

            summary = update_workbook(input_path, features_path, output_path, titles_path)
            output_workbook = load_workbook(output_path, read_only=True)
            output_sheet = output_workbook.active
            output_title = output_sheet.cell(2, 2).value
            output_workbook.close()

            self.assertEqual(summary["retitled_rows"], 1)
            self.assertIn("165x100mm", output_title)

    def test_updates_title_when_existing_title_misses_model_identity(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_path = root / "woo.xlsx"
            features_path = root / "features.csv"
            titles_path = root / "titles.csv"
            output_path = root / "out.xlsx"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["SKU", "Title", "Atrybut Produktu: Typ Produktu", "Product Attributes"])
            worksheet.append(["39755/KAN", "Ramka do paneli ADTR 60x60 cm Kanlux 39755", "", ""])
            workbook.save(input_path)

            features_path.write_text(
                'product_id;sku;features\n'
                '39755;39755/KAN;"{""Typ produktu"":""Ramka do paneli"",""Wymiary [mm]"":""600 x 600""}"\n',
                encoding="utf-8-sig",
            )
            titles_path.write_text(
                "product_id;name;sku\n"
                "39755;Ramka do paneli ADTR SKY-V 60x60 cm kwadratowa złoto-brązowy Kanlux 39755;39755/KAN\n",
                encoding="utf-8-sig",
            )

            summary = update_workbook(input_path, features_path, output_path, titles_path)
            output_workbook = load_workbook(output_path, read_only=True)
            output_sheet = output_workbook.active
            output_title = output_sheet.cell(2, 2).value
            output_workbook.close()

            self.assertEqual(summary["retitled_rows"], 1)
            self.assertIn("ADTR SKY-V", output_title)

    def test_updates_title_when_existing_agor_title_misses_beam_type(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_path = root / "woo.xlsx"
            features_path = root / "features.csv"
            titles_path = root / "titles.csv"
            output_path = root / "out.xlsx"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["SKU", "Title", "Atrybut Produktu: Typ Produktu", "Product Attributes"])
            worksheet.append(["38423/KAN", "Naświetlacz LED FL AGOR 100W 13500lm 4000K IP65 Kanlux 38423", "", ""])
            workbook.save(input_path)

            features_path.write_text(
                'product_id;sku;features\n'
                '38423;38423/KAN;"{""Typ produktu"":""Naświetlacz LED"",""Moc [W]"":""100""}"\n',
                encoding="utf-8-sig",
            )
            titles_path.write_text(
                "product_id;name;sku\n"
                "38423;Naświetlacz LED FL AGOR PRO asymetryczny 100W 13500lm 4000K IP65 Kanlux 38423;38423/KAN\n",
                encoding="utf-8-sig",
            )

            summary = update_workbook(input_path, features_path, output_path, titles_path)
            output_workbook = load_workbook(output_path, read_only=True)
            output_sheet = output_workbook.active
            output_title = output_sheet.cell(2, 2).value
            output_workbook.close()

            self.assertEqual(summary["retitled_rows"], 1)
            self.assertIn("asymetryczny", output_title)


if __name__ == "__main__":
    unittest.main()
