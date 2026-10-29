"""Dane techniczne narzedzi Högert z hoegert.com po kodzie - do rozstrzygania powtorzen nazw w Ekonomie.

GTV sprzedaje Högert pod kodami HT... i w Ekonomie wiele kodow ma identyczna nazwe ("Palnik gazowy" x6).
Strona producenta tez ma te same nazwy, ale karta (`hoegert.com/oferta/<slug>/`, znajdowana wyszukiwarka
`?s=<kod>`) podaje "Dane techniczne" (Moc [kW], Rodzaj gazu...) i "Charakterystyke" - tym sie roznia.

Wynik: `cache/hoegert_web/<kod>.json` = {"code", "found", "name", "url", "params", "features"} oraz
(z `--review`) skoroszyt: grupy powtorzen z arkusza "Powtorzenia" i parametry, ktore roznia sie w grupie -
do decyzji uzytkownika, jaki atrybut dopisac do nazwy (nie dobieramy go sami).

    py scripts/fetch_hoegert_web.py --report reports/Ekonom-nazwy-Gtv_kontrola_2026-10-15_v2.xlsx \
        --review reports/Ekonom-Hogert_powtorzenia_2026-10-15.xlsx
"""
from __future__ import annotations

import argparse
import html
import json
import re
import time
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / "cache" / "hoegert_web"
SEARCH = "https://www.hoegert.com/?s={}"


def get(url: str, tries: int = 3) -> str:
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    for attempt in range(tries):
        try:
            with urllib.request.urlopen(req, timeout=60) as resp:
                return resp.read().decode("utf-8", "replace")
        except TimeoutError:  # strona bywa wolna - ponawiamy
            if attempt == tries - 1:
                raise
            time.sleep(2)
    return ""


def page_lines(raw: str) -> list[str]:
    text = re.sub(r"<script.*?</script>|<style.*?</style>", "", raw, flags=re.S)
    text = html.unescape(re.sub(r"<[^>]+>", "\n", text))
    return [" ".join(line.split()) for line in text.splitlines() if line.strip()]


def cell(text: str) -> str:
    return " ".join(html.unescape(re.sub(r"<br\s*/?>", ", ", re.sub(r"<(?!br)[^>]+>", " ", text))).split()).strip(", ")


def parse_product(raw: str) -> tuple[str, dict[str, str], list[str]]:
    """Nazwa, tabela "Dane techniczne" (wiersze <td>etykieta</td><td>wartosc</td>) i lista "Charakterystyka"."""
    lines = page_lines(raw)
    name = lines[0].split(" - Högert")[0] if lines else ""
    params: dict[str, str] = {}
    i = raw.rfind("Dane techniczne</h3>")
    table = re.search(r"<table.*?</table>", raw[i:], flags=re.S) if i >= 0 else None
    for label, value in re.findall(r"<tr>\s*<td>(.*?)</td>\s*<td>(.*?)</td>\s*</tr>", table.group(0) if table else "",
                                   flags=re.S):
        key = cell(label)
        if key and key != "Producent":  # wiersz producenta (GTV Poland) jest na kazdej karcie
            params[key] = cell(value)
    features: list[str] = []
    chars = [k for k, line in enumerate(lines) if line == "Charakterystyka"]
    if chars:
        k = chars[-1] + 1
        while k < len(lines) and lines[k] not in ("Do pobrania", "Dane techniczne"):
            features.append(lines[k].rstrip(";."))
            k += 1
    return name, params, features


def fetch(code: str, url: str = "") -> dict:
    """Karta po kodzie; `url` z cache pomija wyszukiwarke (ponowne parsowanie)."""
    links = [url] if url else list(dict.fromkeys(re.findall(
        r'href="(https://(?:www\.)?hoegert\.com/oferta/[^"#?]+)"', get(SEARCH.format(urllib.parse.quote(code))))))
    for link in links[:5]:
        raw = get(link)
        # karta musi miec dokladnie ten kod (SKU) - wyszukiwarka zwraca tez podobne produkty
        if any(line == code for line in page_lines(raw)):
            name, params, features = parse_product(raw)
            return {"code": code, "found": True, "name": name, "url": link, "params": params, "features": features}
    return {"code": code, "found": False, "name": "", "url": "", "params": {}, "features": []}


def duplicate_groups(path: Path) -> list[tuple[int, str, str, str]]:
    """(grupa, kod, nazwa w Ekonomie, nowa nazwa) - tylko grupy z kodami Högert HT..."""
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb["Powtorzenia"].iter_rows(values_only=True))
    head = rows[0]
    ig, ic, io, inn = head.index("Grupa"), head.index("Kod towaru"), head.index("Stara nazwa"), head.index("Nowa nazwa")
    return [(r[ig], str(r[ic]).split("/")[0], r[io], r[inn]) for r in rows[1:] if str(r[ic]).startswith("HT")]


def write_review(path: Path, groups: list[tuple[int, str, str, str]], recs: dict[str, dict]) -> int:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Powtorzenia Hogert"
    ws.append(["Grupa", "Kod", "Nazwa w Ekonomie", "Nowa nazwa", "Nazwa na hoegert.com", "Czym sie rozni (hoegert.com)",
               "Karta", "Decyzja uzytkownika"])
    by_group: dict[int, list[tuple[str, str, str]]] = {}
    for g, code, old, new in groups:
        by_group.setdefault(g, []).append((code, old, new))
    resolved = 0
    for g, members in by_group.items():
        recs_g = [recs.get(code) or {} for code, _, _ in members]
        # 1) nazwa na stronie producenta bywa dokladniejsza niz w Ekonomie ("... 11,5 kg", "... i LCD")
        names = [r.get("name", "") if r.get("found") else "" for r in recs_g]
        common_words = set.intersection(*(set(n.split()) for n in names)) if all(names) else set()
        name_diffs = [" ".join(w for w in n.split() if w not in common_words) for n in names]
        # 2) dane techniczne - tylko prawdziwe etykiety (na czesci stron w tabele wpada blok kontaktowy)
        def clean(params: dict) -> dict:
            return {k: v for k, v in params.items() if re.match(r"[A-Za-zĄĆĘŁŃÓŚŹŻąćęłńóśźż]", k)
                    and not re.search(r"\+48|B2B|BHP|@", k + v) and len(k) < 60}
        params_g = [clean(r.get("params", {})) for r in recs_g]
        keys = {k for prm in params_g for k in prm}
        diff_keys = [k for k in sorted(keys) if len({prm.get(k, "") for prm in params_g}) > 1]
        feats = [set(r.get("features", [])) for r in recs_g]
        common = set.intersection(*feats) if feats else set()
        diffs = []
        for nd, prm, r in zip(name_diffs, params_g, recs_g):
            if all(names) and len(set(names)) == len(names):
                parts = [f"nazwa: {nd or '(wersja podstawowa)'}"]
            else:
                parts = [f"{k}: {prm.get(k, '-')}" for k in diff_keys] or sorted(set(r.get("features", [])) - common)[:3]
            diffs.append("; ".join(parts))
        resolved += len(set(diffs)) == len(diffs) and all(diffs)
        for (code, old, new), r, d in zip(members, recs_g, diffs):
            ws.append([g, code, old, new, r.get("name", "") if r.get("found") else "brak na hoegert.com", d,
                       r.get("url", ""), ""])
    for col, width in zip("ABCDEFGH", (7, 16, 45, 45, 45, 70, 40, 25)):
        ws.column_dimensions[col].width = width
    wb.save(path)
    return resolved


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--report", type=Path, help="skoroszyt kontrolny GTV (arkusz Powtorzenia)")
    parser.add_argument("--review", type=Path, help="zapisz zestawienie roznic w grupach powtorzen")
    parser.add_argument("--code", nargs="*", default=[])
    parser.add_argument("--refresh", action="store_true")
    parser.add_argument("--reparse", action="store_true", help="pobierz ponownie tylko karty z cache (bez wyszukiwarki)")
    parser.add_argument("--workers", type=int, default=4)
    args = parser.parse_args()

    groups = duplicate_groups(args.report) if args.report else []
    codes = list(dict.fromkeys(args.code + [code for _, code, _, _ in groups]))
    CACHE.mkdir(parents=True, exist_ok=True)

    def one(code: str) -> dict | None:
        path = CACHE / f"{re.sub(r'[^A-Za-z0-9._-]', '_', code)}.json"
        old = json.loads(path.read_text(encoding="utf-8")) if path.exists() else None
        if old and not args.refresh and not (args.reparse and old.get("found")):
            return old
        try:
            rec = fetch(code, old["url"] if args.reparse and old and old.get("found") else "")
        except Exception as exc:  # siec / zmiana strony - ponowimy nastepnym razem
            print(f"{code}: blad {exc}")
            return None
        path.write_text(json.dumps(rec, ensure_ascii=False, indent=1), encoding="utf-8")
        time.sleep(0.3)
        return rec

    with ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
        recs = {code: rec for code, rec in zip(codes, pool.map(one, codes)) if rec}
    found = sum(r["found"] for r in recs.values())
    print(f"Kody: {len(codes)}  na hoegert.com: {found}  brak: {len(recs) - found}  bledy: {len(codes) - len(recs)}")
    if args.review and groups:
        resolved = write_review(args.review, groups, recs)
        print(f"Grupy: {len({g for g, *_ in groups})}, rozroznialne danymi ze strony: {resolved}")
        print(f"Zestawienie: {args.review}")


if __name__ == "__main__":
    main()
