"""Usuwa wzmianki o producencie z tytulow i opisow kabli (eksport WP All Export).

Jednorazowa poprawka z 2026-10-12: producent kabli w sklepie to zawsze
"Polski Producent", wiec opisy i tytuly nie moga wymieniac Elektrokabel, NKT,
Radpol, Bitner/BiTLAN ani Helukabel/SOLARFLEX-X, a opisy nie podaja kodow
katalogowych producenta. Kazde zdanie jest przeredagowane recznie (lista
EDITS) - zwykle wyciecie nazwy zostawialo zdania typu "YDY 5X1 750V marki to".

Zmienia tylko kolumny Title i Content. Wynik: kopia eksportu w output/ i
skoroszyt kontrolny w reports/. Przerywa, gdy ktoregos fragmentu nie znajdzie.

    py scripts/strip_cable_manufacturer_mentions.py
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

NB = "\xa0"
FI = '<strong style="font-family: inherit;">'

# Wiersze specyfikacji z nazwa albo kodem producenta - usuwane w calosci.
SPEC_LINE = re.compile(
    r"\n[ \t]*<li>(?:- )?(?:<strong>)?"
    r"(?:Indeks producenta|Numer katalogowy|Indeks/kod produktu|Producent|Marka)"
    r":(?:</strong>)?[^<\n]*</li>"
)

# id produktu -> lista (stary fragment, nowy fragment); stary musi wystapic.
EDITS: dict[int, list[tuple[str, str]]] = {
    26636: [  # 8PYDYP215
        ("Przewody elektryczne płaskie firmy Elektrokabel - efektywne", "Przewody elektryczne płaskie YDYp - efektywne"),
        ("YDYp 2x1,5 750V firmy Elektrokabel zaprojektowany", "YDYp 2x1,5 750V zaprojektowany"),
        ("Przewód YDYp 2x1,5 firmy Elektrokabel jest idealny", "Przewód YDYp 2x1,5 jest idealny"),
    ],
    26642: [  # 8PYDYP34750
        ("YDYp ŻO 3x4 750V Elektrokabel zapewnia", "YDYp ŻO 3x4 750V zapewnia"),
        ("YDYp żo 3x4 750V Elektrokabel - kompaktowe", "YDYp żo 3x4 750V - kompaktowe"),
        ("YDYp żo 3x4 750V Elektrokabel to wysokiej", "YDYp żo 3x4 750V to wysokiej"),
        (f"Przewód{NB} płaski YDYp 3x4 750V Elektrokabel jest", "Przewód płaski YDYp 3x4 750V jest"),
    ],
    26644: [  # 8PYDYP425750
        ("YDYp z żo 4x2,5 Elektrokabel - innowacyjne", "YDYp z żo 4x2,5 - innowacyjne"),
        ("<h2>Przewód płaski YDYp z żo 4x2,5 Elektrokabel</h2>", "<h2>Przewód płaski YDYp z żo 4x2,5</h2>"),
        ("Przewód płaski YDYp z zo 4x2,5 Elektrokabel dzięki płaskiemu kształtowi, przewód ten doskonale",
         "Dzięki płaskiemu kształtowi przewód płaski YDYp z żo 4x2,5 doskonale"),
    ],
    26654: [  # 8PYDY5025750
        ("Przewód YDY żo 5x2,5 Elektrokabel to", "Przewód YDY żo 5x2,5 to"),
        ("Przewód YDY 5x2,5 Elektrokabel cechuje", "Przewód YDY 5x2,5 cechuje"),
        ("YDY ŻO 5x2,5 750V firmy Elektrokabel pozwala", "YDY ŻO 5x2,5 750V pozwala"),
    ],
    26658: [  # 8PYDY506750
        ("Przewód YDY 5x6 żo Elektrokabel", "Przewód YDY 5x6 żo"),  # naglowek i akapit
    ],
    26664: [  # 8PYDY2025
        ("450/750V biały Elektrokabel</strong>", "450/750V biały</strong>"),
        (f"Przewód YDY{NB} 2x2,5 450/750V biały Elektrokabel jest", "Przewód YDY 2x2,5 450/750V biały jest"),
        ("Firma Elektrokabel cieszy się dobrą reputacją i zaufaniem w branży elektrycznej. "
         "Przewody tej marki są znane z wysokiej jakości wykonania, trwałości i bezpieczeństwa. ", ""),
        ("450/750V biały Elektrokabel spełni", "450/750V biały spełni"),
    ],
    26672: [  # 8PYDY304
        ("Firma Elektrokabel, jako renomowany producent przewodów elektrycznych, jest znana ze swojej "
         "dbałości o jakość i bezpieczeństwo swoich produktów. Kable YDY tego producenta spełniają",
         "Kable YDY spełniają"),
        ("<h2>Przewód YDY 3x4 750V biały Elektrokabel</h2>", "<h2>Przewód YDY 3x4 750V biały</h2>"),
    ],
    26674: [  # 8PYDY306
        ("<h2>Przewód instalacyjny YDY żo 3x6 biały Elektrokabel</h2>", "<h2>Przewód instalacyjny YDY żo 3x6 biały</h2>"),
        ("Przewód 3x6 marki Elektrokabel, do której należy przewodów elektrycznnych z żyłą ochronną , "
         "cieszą się uznaniem profesjonalistów w branży elektrycznej. Dzięki ich wysokiej jakości wykonania "
         "i zgodności z najnowszymi normami branżowymi, przewody YDY żo tej marki są wyborem godnym zaufania, "
         "bez względu na rodzaj instalacji czy środowisko, ten kabel",
         "Przewody YDY żo 3x6 z żyłą ochronną cieszą się uznaniem profesjonalistów w branży elektrycznej. "
         "Dzięki wysokiej jakości wykonania i zgodności z najnowszymi normami branżowymi są wyborem godnym "
         "zaufania - bez względu na rodzaj instalacji czy środowisko ten kabel"),
    ],
    26684: [("0,6/1kV Elektrokabel</b>", "0,6/1kV</b>")],  # 8PYKY510
    26691: [("0,6/1kV NKT to", "0,6/1kV to")],  # 8PYKY506/NKT
    39298: [  # mufa Radpol
        ("(K-ZZ) produkowana przez RADPOL jest", "(K-ZZ) jest"),
        ("<strong>Producent:</strong> RADPOL\n\n", ""),
    ],
    104802: [  # 8PYDYP24 - resztki po wycietej nazwie
        ("YDYp 2x4 750V firmy niezastąpionym", "YDYp 2x4 750V niezastąpionym"),
        ("<h2>Przewód płaski YDYp 2x4 firmy - skuteczne", "<h2>Przewód płaski YDYp 2x4 - skuteczne"),
        ("Przewód płaski YDYp 2x4 firmy dla kompaktowych instalacji elektrycznych, które wymagają wysokiej "
         "wydajności i niezawodności. nalazł",
         "Przewód płaski YDYp 2x4 to rozwiązanie dla kompaktowych instalacji elektrycznych, które wymagają "
         "wysokiej wydajności i niezawodności. Znalazł"),
    ],
    104816: [  # 112194015/NKT
        (" Kabel ten jest wyprodukowany przez renomowanego producenta NKT, który jest znany ze swojej "
         "doskonałej jakości i niezawodności.", ""),
    ],
    109617: [("Decydując się na produkt tej marki inwestujesz", "Decydując się na ten kabel, inwestujesz")],
    127369: [("YDY 5X1 750V marki to", "YDY 5X1 750V to")],
    127371: [("YDY 5X1 750V marki to", "YDY 5X1 750V to")],
    127395: [("YDY 4X1 750V marki to", "YDY 4X1 750V to")],
    129279: [("YKYżo 3x1,5 od producenta to", "YKYżo 3x1,5 to")],
    129655: [("500V biały marki to", "500V biały to")],
    129675: [("5x6 500 V marki to", "5x6 500 V to")],
    129691: [("500 V czarny marki to", "500 V czarny to")],
    129695: [("500 V czarny marki Elektrokabel to", "500 V czarny to")],
    129697: [
        ("500 V czarny marki Elektrokabel to", "500 V czarny to"),
        ("500 V czarny Elektrokabel 100 m</h2>", "500 V czarny 100 m</h2>"),
    ],
    147084: [  # TI0006 karton
        ("200MHz 305m Bitner TI0006</strong>", "200MHz 305m</strong>"),
        ("postawić na model Bitner TI0006?", "postawić na kabel kat. 5e?"),
        ("Kabel UTP kategorii 5e marki Bitner został", "Kabel UTP kategorii 5e został"),
        ("Model TI0006 spełnia", "Kabel spełnia"),
    ],
    147086: [  # TI0044 karton
        ("kat. 6 350MHz Bitner TI0044</strong>", "kat. 6 350MHz</strong>"),
        ("Produkt renomowanego polskiego producenta BITNER gwarantuje", "Kabel gwarantuje"),
        ("kabel UTP kat. 6 Bitner TI0044 do", "kabel UTP kat. 6 do"),
        ("Kabel kategorii 6 U/UTP Bitner TI0044 zaprojektowano", "Kabel kategorii 6 U/UTP zaprojektowano"),
    ],
    147551: [  # TI0006 1m
        ("(24AWG) BiTLAN kat. 5e 200MHz 1m TI0006</strong>", "(24AWG) kat. 5e 200MHz 1m</strong>"),
        ("U/UTP 4x2x0,5 BiTLAN kat. 5e?", "U/UTP 4x2x0,5 kat. 5e?"),
        ("Kabel U/UTP kat. 5e BiTLAN zapewnia", "Kabel U/UTP kat. 5e zapewnia"),
    ],
    147553: [("(24AWG) BiTLAN kat. 5e 200MHz to", "(24AWG) kat. 5e 200MHz to")],
    147555: [("(24AWG) BiTLAN kat.5e 200MHz</strong>", "(24AWG) kat.5e 200MHz</strong>")],
    147557: [  # TI0044 1m
        ("(23AWG) BiTLAN kat. 6 350MHz TI0044</strong>", "(23AWG) kat. 6 350MHz</strong>"),
        (f"U/UTP{NB}</strong>BiTLAN{FI}{NB}kat. 6 350Hz</strong>", f"U/UTP{NB}</strong>{FI}kat. 6 350MHz</strong>"),
        ("Kabel U/UTP BiTLAN kat. 6 marki Bitner zaprojektowany", "Kabel U/UTP kat. 6 zaprojektowany"),
    ],
    147559: [  # TI0048 1m
        ("(23AWG) BiTLAN kat. 6 350MHz</strong>", "(23AWG) kat. 6 350MHz</strong>"),
        ("kat. 6 BiTLAN?</strong>", "kat. 6?</strong>"),
    ],
    147561: [  # TI0054 1m
        ("kat. 6 BiTLAN 4x2x0,54 (23AWG)</strong>", "kat. 6 4x2x0,54 (23AWG)</strong>"),
        (f"U/UTPf{NB}BiTLAN</strong>", "U/UTPf</strong>"),
    ],
    147563: [  # TI0060 1m
        ("(24AWG) BiTLAN kat. 5e 200MHz</strong>", "(24AWG) kat. 5e 200MHz</strong>"),
        (f"U/UTP{NB}</strong>BiTLAN{FI}{NB}kat.5e", f"U/UTP{NB}</strong>{FI}kat.5e"),
        ("kategorii 5e marki Bitner został", "kategorii 5e został"),
    ],
    147589: [  # TI0007 karton
        ("BiTLAN kat. 5e 200MHz TI0007</strong>", "kat. 5e 200MHz</strong>"),
        ("kabel F/UTP BiTLAN kat. 5e 200 MHz?", "kabel F/UTP kat. 5e 200 MHz?"),
        ("Kabel F/UTP ekranowany BiTLAN zapewnia", "Kabel F/UTP ekranowany zapewnia"),
    ],
    147592: [  # TI0012 305m
        ("4x2x0,5 BiTLAN 200MHz</strong>", "4x2x0,5 200MHz</strong>"),
        ("U/UTPf kat.5e BiTLAN?", "U/UTPf kat.5e?"),
    ],
    147594: [  # TI0048 305m
        ("(23AWG) BiTLAN kat. 6 350MHz</strong>", "(23AWG) kat. 6 350MHz</strong>"),
        (f"F/UTP </strong>BiTLAN{FI}P kat. 6", f"F/UTP </strong>{FI}kat. 6"),
    ],
    147596: [  # TI0055 305m
        ("kat. 6 BiTLAN outdoor 350MHz</strong>", "kat. 6 outdoor 350MHz</strong>"),
        ("kat. 6 BiTLAN outdoor?", "kat. 6 outdoor?"),
        ("Kabel BiTLAN F/UTPf cat. 6", "Kabel F/UTPf kat. 6"),
    ],
    157722: [  # Helukabel - nazwa handlowa SOLARFLEX-X (rowniez w FAQ JSON-LD)
        ("SOLARFLEX-X H1Z2Z2-K", "H1Z2Z2-K"),
        ("przewód SOLARFLEX-X?", "przewód H1Z2Z2-K?"),
    ],
    157724: [("SOLARFLEX-X H1Z2Z2-K", "H1Z2Z2-K")],
    165423: [  # 172153016C0100/NKT
        (f"przewód płaski{NB} 3x6 firmy NKT,", "przewód płaski YDYp 3x6,"),
        ("750V firmy NKT", "750V"),  # naglowek, akapit, podsumowanie
    ],
    165425: [  # 5907702810553/EKB
        ("750V biały firmy Elektrokabel", "750V biały"),
        ("Przewód YDY 5x1 Elektrokabel - niezawodny", "Przewód YDY 5x1 - niezawodny"),
        ("Przewód YDY 5x1 750V Elektrokabel to", "Przewód YDY 5x1 750V to"),
    ],
    165427: [  # 5907702810492/EKB
        ("<p>Przewody elektryczne Elektrokabel to produkt, zapewniający", "<p>Przewody elektryczne YDY to produkt zapewniający"),
        ("ochronną 4x1 Elektrokabel mogą", "ochronną 4x1 mogą"),
    ],
    165446: [  # 5907702810430/EKB
        ("<strong>Przewód YDY 3x1 firmy Elektrokabel</strong>", "<strong>Przewód YDY 3x1</strong>"),
        ("przewód 3x1 Elektrokabel posiada", "przewód 3x1 posiada"),
        ("<p>Firma Elektrokabel cieszy się renomą i uznaniem w branży elektrycznej. Jej produkty są znane "
         "z wysokiej jakości wykonania, bezpieczeństwa i niezawodności. Przewód YDY 3x1", "<p>Przewód YDY 3x1"),
    ],
    165448: [  # 5907702810386/EKB
        ("żo 2x1,5 Elektrokabel to", "żo 2x1,5 to"),
        ("750V firmy Elektrokabel to", "750V to"),
    ],
    165490: [  # 5907702810096/EKB
        ("<h2>Przewód płaski YDYp 2x1 750V Elektrokabel to", "<h2>Przewód płaski YDYp 2x1 750V to"),
        ("<p>Przewód YDYp 2x1 750V Elektrokabel dla", "<p>Przewód YDYp 2x1 750V to rozwiązanie dla"),
    ],
    165492: [("750V firmy Elektrokabel", "750V")],  # 5907702810157/EKB, 4 wystapienia
    165496: [  # 172153019C0100/NKT
        (" Jest on produkowany przez cenionego producenta, NKT, który słynie z solidności i profesjonalizmu.", ""),
    ],
    165506: [  # 5907702810379/EKB
        ("450/ 750 firmy Elektrokabel</strong>", "450/ 750</strong>"),
        (f"<strong>Zalety przewodów 2X1 Elektrokabel{NB}</strong>", "<strong>Zalety przewodów YDY 2X1</strong>"),
    ],
}

# Tytuly: wycinane tylko nazwy producenta i jego linii, kody na koncu zostaja.
TITLE_EDITS: list[tuple[str, str]] = [
    (" Elektrokabel", ""),
    (" 0,6/1kV NKT", " 0,6/1kV"),
    (" Bitner TI", " TI"),
    ("BiTLAN ", ""),
    ("SOLARFLEX-X ", ""),
    (" Helukabel ", " "),
]

# Kontrola koncowa: nic z tego nie moze zostac w Title/Content.
FORBIDDEN = re.compile(
    r"elektrokabel|bitner|bitlan|\bnkt\b|radpol|zamel|heluk|solarflex|elpar|elpio|anticor|\btfk\b"
    r"|\bmarki\b|\bmarka\b|\bfirmy\b|\bfirma\b|od producenta|indeks producenta|numer katalogowy"
    r"|indeks/kod|\bT0096\b|\bT0138\b",
    re.I,
)
# Neutralne zwroty o producencie, ktore zostaja.
ALLOWED_PRODUCER = re.compile(r"(zaleceniami|wytycznych) producenta|Seria producenta", re.I)


def apply_edits(content: str, pid: int, log: list) -> str:
    for old, new in EDITS.get(pid, []):
        count = content.count(old)
        if not count:
            raise SystemExit(f"[{pid}] nie znaleziono fragmentu: {old[:80]!r}")
        content = content.replace(old, new)
        log.append(("Content", old, new, count))
    for match in SPEC_LINE.findall(content):
        log.append(("Content (specyfikacja)", match.strip(), "", 1))
    content = SPEC_LINE.sub("", content)
    return content


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--input", default="input/Wszystkie-kable-przeglad-konfekcji.xlsx")
    parser.add_argument("--output", default="output/Kable-bez-producenta_2026-10-12.xlsx")
    parser.add_argument("--report", default="reports/Kable-bez-producenta_kontrola_2026-10-12.xlsx")
    args = parser.parse_args()

    if Path(args.output).resolve() == Path(args.input).resolve():
        raise SystemExit("Wyjscie nie moze nadpisac pliku wejsciowego.")

    wb = openpyxl.load_workbook(args.input)
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    c_id, c_sku, c_title, c_content = col["id"], col["Sku"], col["Title"], col["Content"]

    changes, leftovers, seen = [], [], set()
    for row in ws.iter_rows(min_row=2):
        pid = row[c_id].value
        sku = row[c_sku].value
        title_cell, content_cell = row[c_title], row[c_content]

        title = title_cell.value or ""
        new_title = title
        for old, new in TITLE_EDITS:
            new_title = new_title.replace(old, new)
        if new_title != title:
            title_cell.value = new_title
            changes.append((pid, sku, "Title", title, new_title, 1))

        content = content_cell.value or ""
        log: list = []
        new_content = apply_edits(content, pid, log)
        if new_content != content:
            content_cell.value = new_content
            seen.add(pid)
            changes.extend((pid, sku, *entry) for entry in log)

        for field, text in (("Title", new_title), ("Content", new_content)):
            plain = re.sub(r"<[^>]+>", " ", text)
            for m in FORBIDDEN.finditer(plain):
                leftovers.append((pid, sku, field, plain[max(0, m.start() - 80):m.end() + 80].strip()))
            for m in re.finditer(r"producent", plain, re.I):
                window = plain[max(0, m.start() - 20):m.end() + 20]
                if not ALLOWED_PRODUCER.search(window) and not FORBIDDEN.search(window):
                    leftovers.append((pid, sku, field, plain[max(0, m.start() - 80):m.end() + 80].strip()))

    missing = set(EDITS) - seen
    if missing:
        raise SystemExit(f"Produkty z listy EDITS bez zmian w pliku: {sorted(missing)}")

    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)

    rep = openpyxl.Workbook()
    sh = rep.active
    sh.title = "Zmiany"
    sh.append(["id", "Sku", "Pole", "Przed", "Po", "Wystapien", "Status"])
    for pid, sku, field, old, new, count in changes:
        sh.append([pid, sku, field, old, new, count, "OK"])
    sh2 = rep.create_sheet("Pozostale wzmianki")
    sh2.append(["id", "Sku", "Pole", "Kontekst", "Status"])
    for item in leftovers:
        sh2.append([*item, "DO_SPRAWDZENIA"])
    for sheet, widths in ((sh, [9, 26, 22, 80, 80, 10, 10]), (sh2, [9, 26, 10, 120, 16])):
        for i, w in enumerate(widths):
            sheet.column_dimensions[chr(65 + i)].width = w
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for r in sheet.iter_rows(min_row=2):
            for cell in r:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
    Path(args.report).parent.mkdir(parents=True, exist_ok=True)
    rep.save(args.report)

    products = {c[0] for c in changes}
    titles = sum(1 for c in changes if c[2] == "Title")
    print(f"Zmienione produkty: {len(products)} (tytuly: {titles}, opisy: {len(seen)})")
    print(f"Pozostale wzmianki: {len(leftovers)}")
    print(f"Wynik: {args.output}")
    print(f"Raport: {args.report}")


if __name__ == "__main__":
    main()
