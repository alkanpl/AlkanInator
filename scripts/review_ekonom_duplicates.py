"""Przeglad grup powtorzen nazw do Ekonoma: co jeszcze odroznia produkty o tej samej nazwie.

Czyta arkusz "Powtorzenia" ze skoroszytu kontrolnego `build_ekonom_names.py` i dla kazdej grupy
zbiera dowody z katalogow PDF (tekst z cache `cache/ekonom_pdf/<dostawca>/`, ten sam parser co
generator - `package_from_pdfs`):

- rozne nazwy katalogowe -> rozne produkty; nazwa katalogowa podpowiada wyroznik (rozmiar globe,
  zarnik, kolor, kat) - decyzja uzytkownika, potem regula w YAML,
- ta sama nazwa katalogowa -> prawdopodobnie inne opakowanie / zdublowana kartoteka,
- brak w katalogach -> potrzebne dane producenta.

Ceny nie sa dowodem (decyzja uzytkownika 2026-10-15). Marke z opakowania, multipaki i barwe
z katalogu rozstrzyga juz generator - tu zostaje to, czego nie umie.

    py scripts/review_ekonom_duplicates.py --supplier ledvance \
        --report reports/Ekonom-nazwy-Ledvance_kontrola_2026-10-15_v2.xlsx \
        --output reports/Ekonom-Ledvance_przeglad_powtorzen_2026-10-15_v2.xlsx
"""
from __future__ import annotations

import argparse
import collections
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from build_ekonom_names import package_from_pdfs  # noqa: E402
from utils import load_yaml  # noqa: E402

IDENTICAL = "identyczna nazwa w Ekonomie"
V_DIFF = "rozne produkty wg katalogu"
V_SAME_ROW = "ta sama nazwa katalogowa"
V_PARTIAL = "czesciowo w katalogach"
V_NONE = "brak w katalogach"
PROPOSALS = {
    V_DIFF: "Dodac wyroznik z nazwy katalogowej (rozmiar, ksztalt, kolor, kat, zarnik) - decyzja uzytkownika",
    V_SAME_ROW: "Ten sam produkt w katalogu - inne opakowanie albo zdublowana kartoteka, do sprawdzenia",
    V_PARTIAL: "Czesc EAN-ow bez katalogu - porownac recznie z nazwami katalogowymi obok",
    V_NONE: "Brak dowodow - potrzebne dane producenta (EAN, pelna nazwa, opakowanie)",
}


def norm(text: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", text.upper())


def classify(eans: list[str], pkg: dict[str, dict[str, Any]]) -> tuple[str, str]:
    names = [pkg.get(e, {}).get("name", "") for e in eans]
    detail = " || ".join(f"{e}: {n or '-'}" for e, n in zip(eans, names))
    if not any(names):
        return V_NONE, detail
    if not all(names):
        return V_PARTIAL, detail
    return (V_SAME_ROW if len({norm(n) for n in names}) == 1 else V_DIFF), detail


def write_workbook(path: Path, groups: list[dict[str, Any]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Podsumowanie"
    ws.append(["Werdykt", "Stara nazwa", "Grup", "Produktow", "Propozycja"])
    counts: collections.Counter = collections.Counter()
    sizes: collections.Counter = collections.Counter()
    for g in groups:
        key = (g["werdykt"], "identyczna" if g["kategorie"] == IDENTICAL else "rozna")
        counts[key] += 1
        sizes[key] += g["ile"]
    for (verdict, old), n in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])):
        ws.append([verdict, old, n, sizes[(verdict, old)], PROPOSALS[verdict]])

    ws = wb.create_sheet("Grupy")
    ws.append(["Grupa", "Ile", "Typ", "Nowa nazwa", "Czym roznia sie stare nazwy", "Werdykt", "Dowod",
               "Propozycja", "Decyzja uzytkownika"])
    for g in groups:
        ws.append([g["grupa"], g["ile"], g["typ"], g["nowa_nazwa"], g["kategorie"], g["werdykt"], g["dowod"],
                   PROPOSALS[g["werdykt"]], ""])

    ws = wb.create_sheet("Produkty")
    ws.append(["Grupa", "Werdykt", "Kod towaru", "EAN", "Stara nazwa", "Nowa nazwa", "Nazwa katalogowa",
               "Katalog"])
    for g in groups:
        for p in g["produkty"]:
            ws.append([g["grupa"], g["werdykt"], p["kod"], p["ean"], p["stara"], g["nowa_nazwa"],
                       p["katalog"], p["plik"]])

    widths = {"Podsumowanie": [32, 12, 8, 10, 90], "Grupy": [7, 5, 14, 60, 40, 30, 100, 60, 30],
              "Produkty": [7, 30, 20, 15, 60, 60, 60, 50]}
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for idx, width in enumerate(widths[sheet.title], start=1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--supplier", required=True)
    parser.add_argument("--report", required=True, type=Path, help="skoroszyt kontrolny build_ekonom_names.py")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--rules", type=Path, default=ROOT / "dictionaries" / "ekonom_name_rules.yaml")
    args = parser.parse_args()

    rules = load_yaml(args.rules)["suppliers"][args.supplier]
    cache = ROOT / "cache" / "ekonom_pdf" / args.supplier
    texts = {p.name: [line.strip() for line in p.read_text(encoding="utf-8").split("\n")]
             for p in sorted(cache.glob("*.txt"))}
    pkg = package_from_pdfs(texts, rules.get("package_brand", {}).get("catalog_columns", {}),
                            rules.get("package_catalogs", []))

    wb = openpyxl.load_workbook(args.report, read_only=True)
    rows = list(wb["Powtorzenia"].iter_rows(values_only=True))
    header = rows[0]
    by_group: dict[int, list[dict[str, Any]]] = collections.defaultdict(list)
    for values in rows[1:]:
        r = dict(zip(header, values))
        by_group[r["Grupa"]].append(r)

    groups = []
    for g, rs in sorted(by_group.items()):
        eans = [str(r["EAN"]) for r in rs]
        verdict, detail = classify(eans, pkg)
        groups.append({
            "grupa": g, "ile": len(rs), "typ": rs[0]["Typ"], "nowa_nazwa": rs[0]["Nowa nazwa"],
            "kategorie": rs[0]["Kategorie roznic w grupie"], "werdykt": verdict, "dowod": detail,
            "produkty": [{"kod": r["Kod towaru"], "ean": e, "stara": r["Stara nazwa"],
                          "katalog": pkg.get(e, {}).get("name", ""),
                          "plik": pkg.get(e, {}).get("src", "")} for r, e in zip(rs, eans)],
        })
    write_workbook(args.output, groups)
    for verdict, n in collections.Counter(g["werdykt"] for g in groups).most_common():
        print(f"{n:4d}  {verdict}")
    print(f"Przeglad: {args.output}")


if __name__ == "__main__":
    main()
