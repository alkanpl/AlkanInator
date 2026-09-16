"""Zrzuca fakty z briefow description_codex_brief_*.xlsx do jednego pliku txt.

Narzedzie pomocnicze do czytania: NIE generuje tresci opisow.
Konsola Windows uzywa cp1250, dlatego wynik idzie do pliku w utf-8.
"""
from __future__ import annotations

import argparse
import pathlib

import openpyxl

SKIP_FACTS = {
    "ID_frame",
    "Nazwa producenta",
    "Adres producenta",
    "E-mail producenta",
    "Typ opakowania zbiorczego",
    "Wysokosc opakowania zbiorczego",
    "Szerokosc opakowania zbiorczego",
    "Dlugosc opakowania zbiorczego",
    "Dlugosc opakowania jednostkowego",
    "Wysokość opakowania zbiorczego",
    "Szerokość opakowania zbiorczego",
    "Długość opakowania zbiorczego",
    "Długość opakowania jednostkowego",
    "Paleta kolorystyczna",
}


def sheet_rows(wb, title):
    if title not in wb.sheetnames:
        return []
    ws = wb[title]
    return [
        [("" if c is None else str(c)).strip() for c in row]
        for row in ws.iter_rows(values_only=True)
    ]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--briefs-dir", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    briefs = sorted(pathlib.Path(args.briefs_dir).glob("description_codex_brief_*.xlsx"))
    out: list[str] = []
    for path in briefs:
        wb = openpyxl.load_workbook(path, read_only=True)
        produkt = sheet_rows(wb, "Produkt")
        header, values = produkt[0], produkt[1]
        info = dict(zip(header, values))
        out.append("=" * 70)
        out.append(f"BRIEF: {path.name}")
        out.append(f"SKU: {info.get('sku','')}")
        out.append(f"EAN: {info.get('ean','')}")
        out.append(f"NAZWA: {info.get('name','')}")
        out.append(f"FRAZA: {info.get('seo_keyword','')}")
        out.append(f"KATEGORIA: {info.get('category','')}")
        for title, label in (
            ("Product facts", "FAKTY"),
            ("Compatibility facts", "KOMPATYBILNOSC"),
            ("Rejected facts", "ODRZUCONE"),
        ):
            rows = sheet_rows(wb, title)[1:]
            rows = [r for r in rows if r and r[0] and r[0] not in SKIP_FACTS and r[1]]
            if not rows:
                continue
            out.append(f"-- {label}:")
            for r in rows:
                out.append(f"   {r[0]}: {r[1]}")
        out.append("")
        wb.close()

    pathlib.Path(args.output).write_text("\n".join(out), encoding="utf-8")
    print(f"OK: {len(briefs)} briefow -> {args.output}")


if __name__ == "__main__":
    main()
