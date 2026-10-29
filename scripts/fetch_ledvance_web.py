"""Pobiera z ledvance.com dane produktu po EAN - cache dla nazw do Ekonoma.

Dwa zrodla, oba zapisywane w `cache/ledvance_web/<EAN>.json` (czyta je `build_ekonom_names.py`):

- karta `/en/gtin/<EAN>` (`gtin_page`) - dokladnie ten GTIN: rodzina, nazwa produktu, kod
  zamowieniowy, sztuk w opakowaniu (Folding box) i w kartonie. Dziala bez zapory antybotowej.
- podpowiedzi wyszukiwarki `/en-int/search/autocomplete/<id>?query=<EAN>` (`products`) - tylko gdy
  karty GTIN nie ma; dopasowanie bywa nieostre, generator sprawdza moc/barwe/gwint.

Kazdy EAN pytamy raz (`--refresh` pyta ponownie).

    py scripts/fetch_ledvance_web.py --report reports/Ekonom-nazwy-Ledvance_kontrola_2026-10-15_v4.xlsx
    py scripts/fetch_ledvance_web.py --ean 4058075603035 4099854547638
    py scripts/fetch_ledvance_web.py --missing-lumen reports/Ekonom-nazwy-Ledvance_kontrola_2026-10-20.xlsx --workers 4

Z karty zapisujemy tez CRI, strumien (lm), moc i barwe - generator uzupelnia nimi brakujace lm (2026-10-20).
"""
from __future__ import annotations

import argparse
import html
import json
import re
import time
from concurrent.futures import ThreadPoolExecutor
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / "cache" / "ledvance_web"
ENDPOINT = "https://www.ledvance.com/en-int/search/autocomplete/34958?query={}"
GTIN_PAGE = "https://www.ledvance.com/en/gtin/{}"


def get(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=40) as resp:
        return resp.read()


def fetch_autocomplete(ean: str) -> list[dict]:
    data = json.loads(get(ENDPOINT.format(urllib.parse.quote(ean))))
    return [{"title": r.get("title", ""), "description": r.get("description", ""), "url": r.get("url", "")}
            for r in data.get("results", []) if r.get("context") == "product"]


def values(m: re.Match | None, unit: str) -> list[float]:
    """"6650 lm / 5900 lm" -> [6650.0, 5900.0]; "1,200 lm" to 1200 (przecinek tysiecy), "2,5 W" to 2.5."""
    if not m:
        return []
    out = []
    for v in re.findall(rf"([\d.,]+)\s+{unit}", m.group(1)):
        out.append(float(v.replace(",", "") if unit == "lm" else v.replace(",", ".")))
    return out


def parse_gtin_page(raw: str, ean: str) -> dict | None:
    """Karta produktu jako tekst "pole | wartosc | ..."; None gdy strona nie zna GTIN."""
    text = re.sub(r"<script.*?</script>|<style.*?</style>", "", raw, flags=re.S)
    text = html.unescape(re.sub(r"(\s*\|\s*)+", " | ", re.sub(r"<[^>]+>", " | ", text)))
    head = re.search(r"PRODUCT INFORMATION \| ([^|]+?) \| ([^|]+?) \|", text)
    if not head or ean not in text:
        return None
    pack = re.search(re.escape(ean) + r" \| ([^|]+?) \| (\d+) \|", text)
    carton = re.search(r"\| Shipping box \| (\d+) \|", text)
    order = re.search(r"Order reference \| ([^|]+?) \|", text)
    # CRI z karty: "LED MR16 DIM S" = 97, "DIM P" = 90 - sama litera w nazwie tego nie mowi (2026-10-20)
    cri = re.search(r"Color rendering index Ra \| [^0-9|]{0,3}(\d{2,3}) \|", text)
    # oprawy z Power Select maja kilka stopni: "55 W / 49 W / 43 W" i "6650 lm / 5900 lm / 5250 lm" (pary w kolejnosci)
    lumens = values(re.search(r"\| Luminous flux \| ([\d.,]+\s+lm(?:\s*/\s*[\d.,]+\s+lm)*) \|", text), "lm")
    watts = values(re.search(r"\| Nominal wattage \| ([\d.,]+\s+W(?:\s*/\s*[\d.,]+\s+W)*) \|", text), "W")  # "250  W"
    kelvin = re.search(r"\| Color temperature \| ([\d.,]+) K \|", text)
    # kat swiecenia opraw: "60 °", "40 ° x 90 °" - zastepuje kody optyki N / W / VW / OV w nazwie (2026-10-20)
    beam = re.search(r"\| Beam angle \| ([^|]+?) \|", text)
    beam_deg = re.findall(r"(\d+(?:[.,]\d+)?)\s*°", beam.group(1)) if beam else []
    return {
        "family": head.group(1).strip(),
        "name": re.sub(r"\s+", " ", head.group(2)).strip(),
        "order_reference": order.group(1).strip() if order else "",
        "package": pack.group(1).strip() if pack else "",
        "pieces": int(pack.group(2)) if pack else None,
        "carton": int(carton.group(1)) if carton else None,
        "cri": int(cri.group(1)) if cri else None,
        "lumen": lumens[0] if lumens else None,
        "power": watts[0] if watts else None,
        "lumens": lumens,
        "powers": watts,
        "cct": int(float(kelvin.group(1).replace(",", ""))) if kelvin else None,
        "beam": "x".join(d.replace(",", ".") for d in beam_deg) + "°" if beam_deg else None,
        "url": GTIN_PAGE.format(ean),
    }


def fetch(ean: str) -> dict:
    rec = {"ean": ean, "fetched": time.strftime("%Y-%m-%d"), "gtin_page": None, "products": []}
    try:
        rec["gtin_page"] = parse_gtin_page(get(GTIN_PAGE.format(ean)).decode("utf-8", "ignore"), ean)
    except urllib.error.HTTPError as exc:
        if exc.code != 404:
            raise
    if not rec["gtin_page"]:
        rec["products"] = fetch_autocomplete(ean)
    return rec


def duplicate_eans(report: Path) -> list[str]:
    wb = openpyxl.load_workbook(report, read_only=True)
    rows = list(wb["Powtorzenia"].iter_rows(values_only=True))
    col = rows[0].index("EAN")
    return sorted({str(r[col]) for r in rows[1:] if r[col]})


def missing_lumen_eans(report: Path) -> list[str]:
    """EAN-y z arkusza "Nazwy" z brakiem lm (kolumna Braki)."""
    wb = openpyxl.load_workbook(report, read_only=True)
    rows = list(wb["Nazwy"].iter_rows(values_only=True))
    ean, braki = rows[0].index("EAN"), rows[0].index("Braki")
    return sorted({str(r[ean]) for r in rows[1:] if r[ean] and "lumen" in str(r[braki] or "")})


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--report", type=Path, help="skoroszyt kontrolny - pytamy o EAN-y z arkusza Powtorzenia")
    parser.add_argument("--ean", nargs="*", default=[])
    parser.add_argument("--missing-beam", action="store_true",
                        help="pobierz ponownie karty bez kata swiecenia (pole beam dodane 2026-10-20)")
    parser.add_argument("--ean-file", type=Path, help="plik z EAN-ami (jeden w wierszu) - gdy lista jest za dluga na --ean")
    parser.add_argument("--refresh", action="store_true")
    parser.add_argument("--missing-lumen", type=Path, help="skoroszyt kontrolny - EAN-y bez lm (arkusz Nazwy)")
    parser.add_argument("--delay", type=float, default=0.4)
    parser.add_argument("--workers", type=int, default=1, help="rownolegle zapytania (4 dziala, wiecej nie testowane)")
    args = parser.parse_args()

    eans = list(args.ean) + (args.ean_file.read_text(encoding="utf-8").split() if args.ean_file else []) + (duplicate_eans(args.report) if args.report else [])         + (missing_lumen_eans(args.missing_lumen) if args.missing_lumen else [])
    eans = list(dict.fromkeys(eans))
    CACHE.mkdir(parents=True, exist_ok=True)
    stats = {"pages": 0, "found": 0, "asked": 0}

    def one(ean: str) -> None:
        path = CACHE / f"{ean}.json"
        rec = json.loads(path.read_text(encoding="utf-8")) if path.exists() else None
        page = (rec or {}).get("gtin_page")
        # stary cache: bez karty GTIN albo karta bez CRI / lm (pola dodane 2026-10-20)
        stale = rec is None or "gtin_page" not in rec or (page and ("cri" not in page or "lumens" not in page)) \
            or (args.missing_beam and page and "beam" not in page)
        if stale or args.refresh:
            try:
                new = fetch(ean)
            except Exception as exc:  # siec / limit - nie zapisujemy, ponowimy przy nastepnym uruchomieniu
                print(f"{ean}: blad {exc}")
                return
            if rec and rec.get("gtin_verified"):
                new["gtin_verified"] = rec["gtin_verified"]
            rec = new
            path.write_text(json.dumps(rec, ensure_ascii=False, indent=1), encoding="utf-8")
            stats["asked"] += 1
            time.sleep(args.delay)
        stats["pages"] += bool(rec.get("gtin_page"))
        stats["found"] += bool(rec.get("gtin_page") or rec.get("products"))

    with ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
        list(pool.map(one, eans))
    pages, found, asked = stats["pages"], stats["found"], stats["asked"]
    print(f"EAN: {len(eans)}, nowe zapytania: {asked}, karta GTIN: {pages}, cokolwiek na ledvance.com: {found}")


if __name__ == "__main__":
    main()
