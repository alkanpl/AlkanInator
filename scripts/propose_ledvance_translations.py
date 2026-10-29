"""Arkusz propozycji EN->PL dla surowych nazw Ledvance do Ekonoma (pkt 3 decyzji uzytkownika 2026-10-21).

Bierze z arkusza "Nazwy" skoroszytu kontrolnego `build_ekonom_names.py` nazwy z co najmniej dwoma surowymi slowami
angielskimi (WIELKIMI LITERAMI) i stosuje do nich slownik `G`. Wynik to PROPOZYCJE do akceptacji uzytkownika
(arkusze "Nazwy" i "Slownik fraz", kolumna "Decyzja") - nic nie trafia do nazw samo. Zaakceptowane frazy przenosi
sie do `dictionaries/ekonom_name_rules.yaml` (sekcja `ledvance`) i pokrywa testem.

    py scripts/propose_ledvance_translations.py reports/Ekonom-nazwy-Ledvance_kontrola_2026-10-21_v20.xlsx \\
        reports/Ekonom-Ledvance_tlumaczenia_EN-PL_do_akceptacji_2026-10-21.xlsx

Slownik: (warunek na nazwe, wzorzec, zamiana, opis). Warunek "" = kazda nazwa.
"""
import re
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill

COLOR = {"BK": "czarny", "WT": "biały", "GY": "szary"}
COLOR_F = {"BK": "czarna", "WT": "biała", "GY": "szara"}

G = [
    # --- Office Linear: laczniki ksztaltowe ---
    ("Office Linear", r"\s*\bSHAPE CONNECTORS\b", "", "SHAPE CONNECTORS - rzeczownik 'Łącznik' juz mowi"),
    ("^Łącznik do Office Linear", r"\bSF\b", "natynkowy", "SF = surface (natynkowy)"),
    ("Office Linear", r"\bL-CON\b", "narożny L", "L-CON"),
    ("Office Linear", r"\bT-CON\b", "trójnik T", "T-CON"),
    ("Office Linear", r"\bX-CON\b", "krzyżowy X", "X-CON"),
    ("Office Linear", r"\bY-CON\b", "typu Y", "Y-CON"),
    ("^Łącznik do Office Linear", r"\bPRISM\b", "pryzmatyczny", "PRISM = klosz pryzmatyczny"),
    # --- TruSys ---
    ("TruSys", r"\b([LTX])-NODE \1(\d\d)-NODE\b",
     lambda m: {"L": "narożny", "T": "trójnik", "X": "krzyżowy"}[m.group(1)] + f" {m.group(1)}{m.group(2)}", "L/T/X-NODE"),
    ("", r"\s*\bFEED-IN CONNECTOR\b", "", "FEED-IN CONNECTOR - rzeczownik 'Zasilanie szyny' juz mowi"),
    ("^Zaślepka", r"\bProfil BLIND COVER\b", "profilu", "Profil BLIND COVER"),
    ("^Zaślepka", r"\s*\bBLIND COVER\b", "", "BLIND COVER - rzeczownik 'Zaślepka' juz mowi"),
    ("^Zaślepka", r"\bEND CAP LOW ENERGY RAIL\b", "niskiej szyny zasilającej", "END CAP LOW ENERGY RAIL"),
    ("^Zaślepka", r"\bEND CAP ENERGY RAIL\b", "szyny zasilającej", "END CAP ENERGY RAIL"),
    ("^Zaślepka", r"\bEND CAP LOW RAIL\b", "niskiej szyny", "END CAP LOW RAIL"),
    ("^Zaślepka", r"\bEND CAP RAIL\b", "szyny", "END CAP RAIL"),
    ("^Uchwyt", r"\bGRID CEILING\b", "do sufitu rastrowego", "GRID CEILING"),
    ("^TruSys FLEX SENSOR TRACK", r"^.*?(\d{3,4}) 8P WT Ledvance$", r"Szyna z czujnikiem do TruSys Flex \1 8P biała Ledvance",
     "SENSOR TRACK"),
    ("^TruSys FLEX T5 TRACK", r"^.*\b(ST|DT) WT Ledvance$", r"Szyna T5 \1 do TruSys Flex biała Ledvance", "T5 TRACK"),
    ("^Oprawa TRUSYS PERFORMANCE SPOT TRACK", r"^.*$", "Szyna Spot Track do TruSys Performance 1500 5P Ledvance", "SPOT TRACK"),
    ("^Oprawa TRUSYS FLEX SPOT TRACK", r"^.*$", "Szyna Spot Track do TruSys Flex 1500 biała Ledvance", "SPOT TRACK"),
    # --- Spot Combo: pierscienie ---
    ("^Pierścień", r"\bFIRE RATED\b", "przeciwpożarowy", "FIRE RATED"),
    ("^Pierścień", r"\bADJUST ADJ\b", "uchylny", "ADJUST ADJ"),
    ("^Pierścień", r"\bFIX\b", "stały", "FIX"),
    ("^Pierścień", r"\bROUND\b", "okrągły", "ROUND (m)"),
    ("^Pierścień", r"\bSQUARE\b", "kwadratowy", "SQUARE (m)"),
    # --- Bulkhead: oslony ---
    ("^Osłona", r"\bROUND EYELID\b", "okrągła z daszkiem", "ROUND EYELID"),
    ("^Osłona", r"\bSURFACE EYELID\b", "natynkowa z daszkiem", "SURFACE EYELID"),
    ("^Osłona", r"\bSURFACE RING\b", "natynkowa z pierścieniem", "SURFACE RING"),
    ("^Oprawa BULKHEAD COMBO ROUND SENSOR", r"^.*$", "Moduł czujnika do Bulkhead Combo okrągłej Ledvance",
     "BULKHEAD COMBO SENSOR"),
    # --- ramki / uchwyty ---
    ("^Ramka CANOPY COMPACT SURFACE", r"^.*$", "Ramka natynkowa do Canopy Compact Ledvance", "CANOPY COMPACT SURFACE"),
    ("^Ramka SURFACE MOUNT", r"\bSURFACE MOUNT\b", "natynkowa", "SURFACE MOUNT"),
    ("^Uchwyt CANOPY COMPACT U BRACKET", r"^.*$", "Uchwyt typu U do Canopy Compact Ledvance", "U BRACKET"),
    ("^Uchwyt Klipsy ECO CLASS", r"^.*$", "Klipsy montażowe natynkowe do Eco Class Ledvance", "SURFACE MOUNT CLIPS"),
    ("^Uchwyt do High Bay HEAVY DUTY ARM", r"^.*$", "Ramię wzmocnione do High Bay Ledvance", "HEAVY DUTY ARM"),
    ("^Uchwyt do High Bay FLEX ACCESSORIES CHAIN", r"^Uchwyt do High Bay FLEX ACCESSORIES CHAIN", "Łańcuch do High Bay Flex",
     "CHAIN"),
    ("^Uchwyt do High Bay FLEX ACCESSORIES SURFACE", r"^.*$", "Uchwyt natynkowy do High Bay Flex Ledvance", "SURFACE"),
    ("^Pilot do High Bay ACCESSORIES CONV BOX", r"^.*$", "Moduł Conversion Box do High Bay Ledvance", "CONV BOX (to nie pilot)"),
    ("^Uchwyt HIGHBAY DALI-2", r"^.*$", "Uchwyt sufitowy czujnika HighBay DALI-2 LS/PD Ledvance", "CEILING MOUNTING"),
    ("^Adapter montażowy czujnika SENSOR CEILING MOUNTING KIT", r"^.*$", "Adapter montażowy czujnika LS/PD-6 na sufit Ledvance",
     "CEILING MOUNTING KIT"),
    ("^Adapter Redukcja ECO AREA POLE REDUCTION", r"^.*$", "Adapter redukcyjny słupa do Eco Area Ledvance", "POLE REDUCTION"),
    ("^Oprawa FLOODLIGHT TRAVERSE", r"^.*$", "Trawers do naświetlaczy 1/2 szt. na rurę 42/60/76 mm Ledvance", "TRAVERSE"),
    ("^Oprawa URBAN CLASSIC GLASS SET GLASS FROSTED", r"^Oprawa URBAN CLASSIC GLASS SET GLASS FROSTED",
     "Klosz szklany matowy do Urban Classic", "GLASS SET FROSTED"),
    ("", r"\bZHAGA TOP&BOTTOM\b", "gniazdo Zhaga góra/dół", "ZHAGA TOP&BOTTOM"),
    ("^Odbłyśnik", r"\bSYM(\d+)DEG\b", r"\1°", "SYM15DEG"),
    ("^Piktogram", r"\bLEFT RIGHT\b", "strzałka lewo/prawo", "LEFT RIGHT"),
    # --- odciazki / zasilacze ---
    ("^Odciążka", r"\bOT CABLE CLAMP ([A-Z])-STYLE\b", r"do zasilaczy OT typ \1", "OT CABLE CLAMP x-STYLE"),
    ("^Odciążka przewodu Zestaw CLAMP DUO", r"Zestaw CLAMP DUO", "Clamp Duo", "Zestaw CLAMP DUO"),
    ("^Zasilacz", r"\bDIP-SWITCH LINEAR\b", "liniowy z przełącznikiem DIP", "DIP-SWITCH LINEAR"),
    ("^Zasilacz LED 500W FLOODLIGHT MAST", r"^.*$", "Zasilacz LED 500W do naświetlaczy na maszt Ledvance",
     "FLOODLIGHT MAST POWER SUPPLY"),
    ("^Zasilacz LED 500W FOR STRIP VOLTAGE", r"^.*$", "Zasilacz LED 500W do taśm High Voltage 230V Ledvance", "FOR STRIP VOLTAGE"),
    ("^Zasilacz LED CONNECTOR FOR P2", r"^.*$", "Złączka zasilająca do taśm LED P2 Ledvance", "CONNECTOR FOR P2 (to nie zasilacz)"),
    ("^Zasilacz NFC SCANNER", r"^.*$", "Skaner NFC do programowania zasilaczy Ledvance", "NFC SCANNER (to nie zasilacz)"),
    ("^OPTOTRONIC 4DIM", r"^.*$", "Zasilacz LED Optotronic 4DIM NFC IP20 Ledvance", "brak reguly"),
    # --- tasmy ---
    ("^Taśma LED STRIP P HIGH VOLTAGE", r"^Taśma LED STRIP P HIGH VOLTAGE 230V 1200 8(\d)(\d) 50M IP66",
     lambda m: f"Taśma LED 230V 50m {m.group(1)}{m.group(2)}00K IP66 High Voltage", "STRIP HIGH VOLTAGE"),
    ("^Złączka do taśm LED CONNECTOR FOR LED STRIP TO LED STRIP", r"CONNECTOR FOR LED STRIP TO LED STRIP (P2 100) P",
     r"taśma-taśma \1", "CONNECTOR STRIP TO STRIP"),
    ("^Złączka do taśm LED CONNECTOR FOR AC MAINS DIRECT", r"CONNECTOR FOR AC MAINS DIRECT (P2 300) P", r"zasilająca 230V \1",
     "CONNECTOR FOR AC MAINS DIRECT"),
    ("^Wtyczki zasilające POWER PLUGS", r"^.*\b(EU|UK) PLUG \1 Ledvance$", r"Wtyczka zasilająca \1 do taśm LED 230V Ledvance",
     "POWER PLUGS FOR LED STRIP"),
    ("^Uchwyt do taśm LED MONTAZOWY MOUNTING BRACKETS", r"MONTAZOWY MOUNTING BRACKETS SMB", "montażowy SMB", "MOUNTING BRACKETS"),
    ("^Profil do taśm LED LF-LTS-MB MOUNTING BRACKET", r"^.*$", "Uchwyt montażowy do profilu LF-LTS Ledvance",
     "MOUNTING BRACKET (to uchwyt)"),
    ("", r"\bSLIM TRACK\b", "Slim Track", "SLIM TRACK - nazwa wlasna"),
    ("^Złączka do taśm LED LINEAR INDIVILED LIGHTLINE CONNECTOR", r"LINEAR INDIVILED LIGHTLINE CONNECTOR",
     "do Linear IndiviLED Lightline", "INDIVILED LIGHTLINE CONNECTOR"),
    ("^Oprawa FLEX INDOOR AUDIO USB COLOR SENSOR", r"^.*$",
     "Taśma LED 2m RGB USB z czujnikiem dźwięku i koloru Ledvance Flex Audio", "FLEX AUDIO COLOR SENSOR"),
    # --- tuby, lampy, UVC ---
    ("^Tuba", r"\bEM HIGH TEMPERATURE\b", "EM do wysokich temperatur", "HIGH TEMPERATURE"),
    ("^ST8V-1.2M", r"^.*$", "Tuba LED T8 120cm 16,2W 4000K EM Osram Substitube Value", "brak reguly"),
    ("^Zapłonnik ST 171 SAFETY TRY 25", r"^.*$", "Zapłonnik ST 171 Safety Ledvance op. 25 szt.",
     "SAFETY TRY 25 (TRY = tacka 25 szt.)"),
    ("^Latarka", r"\s*\bFLASHL HEADL RECH\b", "", "FLASHL HEADL RECH - jest juz 'czołowa akumulatorowa'"),
    ("^Oprawa WORKLIGHTS BATTERY", r"^Oprawa WORKLIGHTS BATTERY (\S+) (\S+) Ledvance$",
     r"Lampa robocza LED akumulatorowa \1 \2 Ledvance Worklights", "WORKLIGHTS BATTERY"),
    ("^Oprawa do UVC \\d+mm LINEAR HOUSING", r"^Oprawa do UVC (\d+mm) LINEAR HOUSING \d+ 1xTUBE UVC SENS WT",
     r"Oprawa UVC \1 na 1 tubę z czujnikiem biała", "LINEAR HOUSING 1xTUBE SENS"),
    ("^Pudełko do sterylizacji", r"^.*$", "Pudełko do sterylizacji UVC LED akumulatorowe Ledvance", "STERILIZATION BOX"),
    ("^Oczyszczacz powietrza UVC HEPA AIR PURIFIER", r"^.*$", "Oczyszczacz powietrza UVC z filtrem HEPA Ledvance", "AIR PURIFIER"),
    ("^Wymienne filtry do UVC AIR PURIFIER", r"UVC AIR PURIFIER", "oczyszczacza UVC", "AIR PURIFIER"),
    ("^Oprawa GU10 Tuba CEILING HOUSING", r"^Oprawa GU10 Tuba CEILING HOUSING (BK|WT)40X1EUEBELLX Ledvance$",
     lambda m: f"Oprawa sufitowa tuba GU10 {COLOR_F[m.group(1)]} Bellalux", "CEILING HOUSING (BELLX = Bellalux)"),
    ("^Oprawa UNDERCABINET SOCKET FL KIT830DUSB", r"^Oprawa UNDERCABINET SOCKET FL KIT830DUSB(BK|WT)",
     lambda m: f"Oprawa podszafkowa LED z gniazdem USB 3000K ściemnialna {COLOR_F[m.group(1)]}", "UNDERCABINET SOCKET"),
    ("^Oprawa STREETLIGHT AREA SMALL RV20ST", r"^.*$", "Oprawa uliczna LED opt. RV20ST Ledvance Streetlight Area Small",
     "bez mocy"),
    ("^SL FLEX SM RV35ST P 25W 727 WAL", r"^.*$",
     "Oprawa uliczna LED 25W 2700K opt. RV35ST aluminium Ledvance Streetlight Flex Small", "brak reguly"),
    # --- awaryjne ---
    ("^Oprawa EMERGENCY CONVERSION BOX TEMPERATURE SENSOR", r"^.*$",
     "Czujnik temperatury do modułu awaryjnego Conversion Box Ledvance", "CONVERSION BOX TEMPERATURE SENSOR"),
    ("^Konwerter do opraw EMERGENCY CONVERSION BOX", r"^.*$", "Moduł awaryjny Conversion Box 3W 3h Li-Ion autotest Ledvance",
     "EMERGENCY CONVERSION BOX"),
    ("^Oprawa DAMP PROOF BATTEN CBO EL BAT", r"^.*$", "Akumulator 6,4V 1,8Ah do Damp Proof / Batten / Linear Combo Ledvance",
     "CBO EL BAT"),
    ("^Oprawa DAMP PROOF BATTEN CBO EL CG", r"^.*$", "Moduł awaryjny 3h autotest do Damp Proof / Batten / Linear Combo Ledvance",
     "CBO EL CG 3H AT"),
    ("^Oprawa DAMP PROOF BATTEN CBO SENSOR", r"^.*$",
     "Czujnik mikrofalowy ruchu i światła do Damp Proof / Batten / Linear Combo Ledvance", "CBO SENSOR LS PD MW"),
    # --- Office Linear moduly ---
    ("^Oprawa OFFICE LINEAR (SF V )?SENSOR", r"^Oprawa OFFICE LINEAR (?:SF V )?SENSOR (BK|WT|GY) Ledvance$",
     lambda m: f"Moduł czujnika do Office Linear {COLOR[m.group(1)]} Ledvance", "OFFICE LINEAR SENSOR"),
    ("^Oprawa OFFICE LINEAR (SF V )?UPLIGHT MODULE", r"^.*$", "Moduł światła górnego do Office Linear Ledvance", "UPLIGHT MODULE"),
    # --- sterowanie DALI ---
    ("^Przycisk DALI-2 PUSH BUTTON COUPLER", r"^.*$", "Sprzęgacz przycisków DALI-2 Ledvance", "PUSH BUTTON COUPLER"),
    ("^Sterownik PCU FLUSH MOUNT", r"PCU FLUSH MOUNT", "podtynkowy PCU", "FLUSH MOUNT"),
    ("^Kompaktowa jednostka sterująca DALI-2 FLUSH MOUNT DALI-2", r"DALI-2 FLUSH MOUNT DALI-2", "podtynkowa DALI-2", "FLUSH MOUNT"),
    ("^Kompaktowa jednostka sterująca DALI-2 REAL-TIME-CLOCK DALI-2", r"DALI-2 REAL-TIME-CLOCK DALI-2",
     "DALI-2 z zegarem czasu rzeczywistego", "REAL-TIME-CLOCK"),
    ("^Ręczna jednostka sterująca SELECT DALI-2 SELECT", r"SELECT DALI-2 SELECT", "Select DALI-2", "SELECT ... SELECT"),
    ("^Ręczna jednostka sterująca SELECT DALI-2 EXCHANGEABLE COVER", r"SELECT DALI-2 EXCHANGEABLE COVER",
     "Select DALI-2 z wymienną nakładką", "EXCHANGEABLE COVER"),
    ("^Ręczna jednostka sterująca SELECT DALI-2 TW", r"SELECT DALI-2", "Select DALI-2", "SELECT"),
    ("^Ręczna jednostka sterująca TOUCH DALI-2", r"TOUCH DALI-2", "dotykowa DALI-2", "TOUCH"),
    ("^Czujnik ruchu i światła DALI-2 LIGHT/PRESENCE", r"LIGHT/PRESENCE-?(\d*) Infrared Ceiling (Integration|Mounting)",
     lambda m: f"LS/PD{('-' + m.group(1)) if m.group(1) else ''} na podczerwień "
               + ("do zabudowy w suficie" if m.group(2) == "Integration" else "natynkowy na sufit"),
     "LIGHT/PRESENCE Infrared Ceiling"),
    ("^MCU GROUP", r"^MCU GROUP", "Jednostka sterująca MCU Group", "brak reguly"),
    ("^MCU SELECT", r"^MCU SELECT", "Jednostka sterująca MCU Select", "brak reguly"),
    ("^Panel dotykowy LC RF TOUCH PANEL", r"\s*TOUCH PANEL", "", "TOUCH PANEL - rzeczownik juz mowi"),
    ("^Sterownik DALI Bluetooth Control DALI ACU BT CONTROL", r"^.*$", "Sterownik DALI ACU Bluetooth Ledvance", "BT CONTROL x2"),
    # --- nazwy wlasne WIELKIMI -> wielkosc liter producenta ---
    ("^Wkład oprawy LINEAR MAGNETIC BATTEN", r"LINEAR MAGNETIC BATTEN", "Linear Magnetic Batten", "nazwa wlasna"),
    ("^Oprawa LN MAG BATTEN T5", r"^Oprawa LN MAG BATTEN T5 (\d+) Ledvance Linear Magnetic Batten T5$",
     r"Wkład oprawy Linear Magnetic Batten T5 \1 Ledvance", "LN MAG BATTEN"),
    ("^Żarówka halogenowa HALOSTAR STARLITE", r"HALOSTAR STARLITE", "Halostar Starlite", "nazwa wlasna"),
]

# skroty techniczne i nazwy, ktore zostaja WIELKIMI - nie sa "surowym angielskim"
KEEP_UPPER = {"LED", "DALI", "RGBW", "RGB", "CCT", "UGR", "COB", "SMD", "IP", "TW", "HF", "EM", "PCR", "CRI", "USB", "PIR",
              "HCL", "UVC", "SMART", "WIFI", "OSRAM", "DIM", "AC", "DC", "ECG", "EVG", "KVG", "HQL", "HQI", "NAV", "HCI",
              "HID", "SPD", "IK", "PAR", "MR", "AR", "ON-OFF", "ZHAGA", "TRUSYS", "CPS", "TPA", "TPB", "PLUS", "PRO", "SET",
              "DA", "NFL", "OPT", "DULUX"}


def raw_rows(report):
    ws = openpyxl.load_workbook(report, read_only=True)["Nazwy"]
    it = ws.iter_rows(values_only=True)
    h = next(it)
    for r in it:
        d = dict(zip(h, r))
        name = str(d["Nowa nazwa"])
        words = [w for w in re.findall(r"[A-Z][A-Z-]{2,}", name) if w not in KEEP_UPPER and not re.search(r"\d", w)]
        if len(words) >= 2:
            yield str(d["EAN"]), d["Typ"], d["Status"], d["Stara nazwa"], name, " ".join(words)


names_rows, used = [], {}
for ean, typ, status, old, new, words in raw_rows(sys.argv[1]):
    prop, hits = new, []
    for i, (cond, pat, repl, desc) in enumerate(G):
        if cond and not re.search(cond, prop):
            continue
        out = re.sub(pat, repl, prop)
        if out != prop:
            hits.append(desc)
            used.setdefault(i, []).append(ean)
            prop = " ".join(out.split())
    names_rows.append((ean, typ, old, new, prop if prop != new else "", "; ".join(hits)))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Nazwy"
ws.append(["EAN", "Typ", "Stara nazwa (Ekonom)", "Nazwa v20", "Propozycja PL", "Uzyte frazy",
           "Decyzja (tak / nie / inna nazwa)"])
for r in sorted(names_rows, key=lambda r: (r[4] == "", r[1], r[3])):
    ws.append(list(r) + ["" if r[4] else "bez propozycji - kod modelu / nazwa wlasna"])
ws2 = wb.create_sheet("Slownik fraz")
ws2.append(["Fraza", "Warunek (nazwa)", "Zamiana", "Ile nazw", "Przyklad EAN", "Decyzja (tak / nie / inna)"])
for i in sorted(used):
    cond, pat, repl, desc = G[i]
    ws2.append([desc, cond.lstrip("^"), repl if isinstance(repl, str) else "(regula - patrz arkusz Nazwy)",
                len(used[i]), used[i][0], ""])
for w in (ws, ws2):
    for c in w[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    w.freeze_panes = "A2"
    w.auto_filter.ref = w.dimensions
    for col in w.columns:
        w.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col[:200]) + 2, 80)
wb.save(sys.argv[2])
n = sum(1 for r in names_rows if r[4])
print(f"nazw: {len(names_rows)}, z propozycja: {n}, bez: {len(names_rows) - n}, regul uzytych: {len(used)}/{len(G)}")
print("nieuzyte:", [G[i][3] for i in range(len(G)) if i not in used])
