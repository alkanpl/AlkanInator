"""Sklada paczke zadania dla zewnetrznego czata: N produktow + kontrakt + wzorzec.

Po co: opisy dla kolejnych SKU pisze zewnetrzny model, a repo dostarcza mu
komplet danych i regul w jednym pliku do wklejenia. Dzieki temu model nie musi
czytac XLSX ani zgadywac konwencji, a wynik da sie zwalidowac tymi samymi
skryptami co zawsze.

Paczka zawiera:
  - reguly z wygenerowanego promptu briefu (jedno zrodlo prawdy, bez przepisywania),
  - jeden gotowy opis jako wzorzec stylu,
  - dla kazdego SKU: nazwe, fraze kluczowa, kategorie i fakty produktowe.

Skrypt niczego nie pisze za model i nie tworzy prozy - tylko przepakowuje briefy.
"""
from __future__ import annotations

import argparse
import pathlib
import re
import sys

import openpyxl

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / 'src'))

from generate_product_descriptions import (  # noqa: E402
    build_codex_description_prompt,
    load_default_seo_knowledge,
)


def sheet_rows(wb, title: str) -> list[list[str]]:
    if title not in wb.sheetnames:
        return []
    return [
        [("" if c is None else str(c)).strip() for c in row]
        for row in wb[title].iter_rows(values_only=True)
    ]


SKIP_FACTS = {
    "ID_frame", "Nazwa producenta", "Adres producenta", "E-mail producenta",
    "Typ opakowania zbiorczego", "Wysokość opakowania zbiorczego",
    "Szerokość opakowania zbiorczego", "Długość opakowania zbiorczego",
    "Długość opakowania jednostkowego", "Paleta kolorystyczna",
}



def product_type(wb) -> str:
    """Wartosc pola "Typ produktu" - walidator wymaga jej w naglowku <h2>."""
    for row in sheet_rows(wb, "Product facts")[1:]:
        if row and row[0] == "Typ produktu" and len(row) > 1:
            return row[1]
    return ""


def brief_block(path: pathlib.Path, index: int) -> str:
    wb = openpyxl.load_workbook(path, read_only=True)
    produkt = sheet_rows(wb, "Produkt")
    info = dict(zip(produkt[0], produkt[1]))
    lines = [
        f"### {index}. {info.get('sku', '')}",
        f"NAZWA PRODUKTU (uzyj doslownie w <strong> na poczatku): {info.get('name', '')}",
        f"FRAZA KLUCZOWA: {info.get('seo_keyword', '')}",
        f"TYP PRODUKTU (musi wystapic w <h2>): {product_type(wb)}",
        f"KATEGORIA: {info.get('category', '')}",
        f"EAN: {info.get('ean', '')}",
        "FAKTY PRODUKTOWE:",
    ]
    for label, block in (("FAKTY PRODUKTOWE", "Product facts"),
                         ("KOMPATYBILNOSC", "Compatibility facts"),
                         ("NIE UZYWAJ JAKO CECH PRODUKTU", "Rejected facts")):
        rows = [r for r in sheet_rows(wb, block)[1:] if r and r[0] and r[0] not in SKIP_FACTS and len(r) > 1 and r[1]]
        if not rows:
            continue
        if label != "FAKTY PRODUKTOWE":
            lines.append(f"{label}:")
        for row in rows:
            lines.append(f"  - {row[0]}: {row[1]}")
    wb.close()
    return "\n".join(lines)


# Prompt briefu powstal dla agenta pracujacego w repo. Zewnetrzny czat nie ma
# dostepu do plikow, wiec linie o sciezkach, Pythonie i review XLSX trzeba wyciac
# - inaczej model dostaje polecenia, ktorych nie da sie wykonac.
_REPO_ONLY = re.compile(
    r"(prompts[\\/]|src[\\/]|C:\\\\Users|\.xlsx|review|Python może|py \w|--brief-input"
    r"|--review-|walidator|Walidacja gotowego)",
    re.I,
)
# Ta linia niesie fraze kluczowa JEDNEGO produktu - w paczce jest ich wiele.
_PER_PRODUCT_KEYWORD = re.compile(r"^- Fraza kluczowa dla tego produktu to:")


def rules_from_brief(path: pathlib.Path) -> str:
    # Reguly bierzemy z AKTUALNEGO kodu, a nie z pliku briefu: briefy bywaja
    # starsze niz kontrakt i nie zawieraja np. sekcji FAQ czy formatu list.
    prompt = build_codex_description_prompt("", [], [], [], "", load_default_seo_knowledge())
    # Sekcje PRODUCT_NAME/PRODUCT_FACTS dotycza jednego SKU - paczka ma wlasne.
    prompt = re.split(r"\nPRODUCT_NAME:", prompt)[0]

    kept: list[str] = []
    for line in prompt.splitlines():
        if _PER_PRODUCT_KEYWORD.match(line.strip()):
            kept.append("- Fraza kluczowa jest podana osobno przy KAZDYM produkcie w sekcji PRODUKTY.")
            continue
        if _REPO_ONLY.search(line):
            continue
        kept.append(line)

    text = "\n".join(kept)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--briefs-dir", required=True)
    parser.add_argument("--done-dir", default="", help="Katalog gotowych opisow - te SKU sa pomijane.")
    parser.add_argument("--example", default="", help="Plik HTML uzyty jako wzorzec stylu.")
    parser.add_argument("--size", type=int, default=15)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    briefs = sorted(pathlib.Path(args.briefs_dir).glob("description_codex_brief_*.xlsx"))
    done = set()
    if args.done_dir:
        done = {p.stem for p in pathlib.Path(args.done_dir).glob("*.html")}
    todo = [p for p in briefs if p.stem.removeprefix("description_codex_brief_") not in done]
    batch = todo[: args.size]
    if not batch:
        raise SystemExit("Brak produktow do zrobienia - wszystkie maja juz opis.")

    parts = [
        "# ZADANIE: opisy HTML SEO do sklepu Alkan",
        "",
        f"W tej paczce jest {len(batch)} produktow. Napisz osobny opis dla KAZDEGO.",
        "Kazdy opis pisz od zera pod konkretny produkt. Nie uzywaj wspolnego szablonu",
        "akapitow ani tych samych zdan w innej kolejnosci miedzy produktami.",
        "",
        "Wynik oddaj jako jeden plik HTML, gdzie kazdy opis jest w bloku:",
        '<div class="item"><div class="sku">SLUG</div> ...opis... </div>',
        "SLUG to identyfikator produktu podany nizej przy kazdym SKU.",
        "",
        "## REGULY",
        "",
        rules_from_brief(batch[0]),
        "",
        "## UWAGI Z POPRZEDNICH PACZEK",
        "",
        "1. Naglowek <h2> MUSI zawierac fraze kluczowa podana przy danym SKU.",
        "   To wystarczy - nie trzeba dodatkowo wciskac mnogiej nazwy kategorii.",
        "   Naglowek ma brzmiec jak normalne polskie zdanie. Jesli zeby spelnic",
        "   regule trzeba zlamac gramatyke, regula jest zle zrozumiana.",
        "   ZLE:  'Jak sterownik przyciskowy daje sterowniki Wi-Fi dla dwoch wyjsc?'",
        "   ZLE:  'Jak gniazdo glosnikowe miesci sie wsrod gniazdka multimedialnego?'",
        "   DOBRZE: 'Jak sterownik przyciskowy obsluguje dwa wyjscia z jednego modulu?'",
        "",
        "2. NIE pisz klientowi o 'briefie', 'materiale zrodlowym' ani 'danych",
        "   wejsciowych'. Zamiast tego: 'Producent podaje...', 'Karta produktu...'.",
        "",
        "3. Unikaj konstrukcji z listy zakazanych: 'dobrze wpisuje sie',",
        "   'ulatwia planowanie', 'jasno wskazuje', 'calego zestawu'.",
        "",
    ]

    if args.example:
        example = pathlib.Path(args.example).read_text(encoding="utf-8")
        parts += [
            "## WZORZEC STYLU I STRUKTURY",
            "",
            "Ponizszy opis pokazuje oczekiwany poziom konkretu, uklad sekcji i format FAQ.",
            "Nie przenos z niego zadnych faktow do innych produktow.",
            "",
            "```html",
            example.strip(),
            "```",
            "",
        ]

    parts += ["## PRODUKTY", ""]
    for index, path in enumerate(batch, 1):
        slug = path.stem.removeprefix("description_codex_brief_")
        parts.append(f"SLUG: {slug}")
        parts.append(brief_block(path, index))
        parts.append("")

    out = pathlib.Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(parts), encoding="utf-8")
    print(f"OK: paczka {len(batch)} produktow -> {out}")
    print(f"    zostalo poza paczka: {len(todo) - len(batch)}")


if __name__ == "__main__":
    main()
